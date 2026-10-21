"""Fixtures for built-in batch 6's section ops — insert_section_break, remove_section_break and set_page_setup's layout
(docs/BUILTIN_OPS_BATCH6_DESIGN.md §8). Every paragraph is found by its text in the tests, never by index.

``build_sections()`` — a four-section specification:

* section 0, the cover: "SECTION 21 13 13", the title, "PROJECT REVISION HISTORY" and a revision table; its own
  default header ("PROJECT MANUAL" and a logo picture only it uses), its own default footer ("21 13 13 - {PAGE}") and,
  with a different first page, its own first-page header ("COVER");
* section 1, the body: PART 1 and PART 2 (Heading 1) with articles, a two-row table and "2.02 SPRINKLERS" after it; the
  page numbers restart at 1; its own default header ("BODY HEADER") and footer ("BODY FOOTER - {PAGE}");
* section 2, a landscape appendix ("APPENDIX A - SCHEDULE", a table, a note): no references of its own (it shows the
  body's header and footer);
* section 3, the end (the body's last sectPr): PART 3 and "END OF SECTION"; it names the body's header part itself
  (two sections share one header part) and has its own footer ("END FOOTER").

``build_single()`` — one section, its own header and footer and (``title_page``) a different first page with its own
first-page header: "INTRODUCTION", three articles, "APPENDIX A", "APPENDIX B", each followed by a paragraph.

``build_page_break()`` — one section whose appendix follows a paragraph holding only a manual page break.
"""
from __future__ import annotations

import io

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Twips
from lxml import etree

from make_fixtures import _page_field, _set_pgnum_start, patch_docx, png

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W = "{%s}" % W_NS
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
MAIN = "word/document.xml"


def save(doc) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _table(doc, rows):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            t.cell(r, c).text = text
    return t


def _new_section(doc, restart=None):
    """A next-page section after the last; references, the first page and the restart are not inherited from the
    section before (python-docx clones its sectPr)."""
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    sec.header.is_linked_to_previous = True
    sec.footer.is_linked_to_previous = True
    sec.different_first_page_header_footer = False
    for old in sec._sectPr.findall(qn("w:pgNumType")):
        sec._sectPr.remove(old)
    if restart is not None:
        _set_pgnum_start(sec, restart)
    return sec


def build_sections() -> bytes:
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(10)
    cover = doc.sections[0]
    cover.header.is_linked_to_previous = False
    p = cover.header.paragraphs[0]
    p.add_run("PROJECT MANUAL ")
    p.add_run().add_picture(io.BytesIO(png((20, 120, 20))), width=Pt(12))
    cover.footer.is_linked_to_previous = False
    p = cover.footer.paragraphs[0]
    p.add_run("21 13 13 - ")
    _page_field(p)
    cover.different_first_page_header_footer = True
    cover.first_page_header.is_linked_to_previous = False
    cover.first_page_header.paragraphs[0].add_run("COVER")
    doc.add_paragraph("SECTION 21 13 13")
    doc.add_paragraph("WET-PIPE SPRINKLER SYSTEMS")
    doc.add_paragraph("PROJECT REVISION HISTORY")
    _table(doc, [("Rev", "Date"), ("0", "2026-01-01")])

    body = _new_section(doc, restart=1)
    body.header.is_linked_to_previous = False
    body.header.paragraphs[0].add_run("BODY HEADER")
    body.footer.is_linked_to_previous = False
    p = body.footer.paragraphs[0]
    p.add_run("BODY FOOTER - ")
    _page_field(p)
    doc.add_paragraph("PART 1 - GENERAL", style="Heading 1")
    doc.add_paragraph("1.01\tSUMMARY")
    doc.add_paragraph("A.\tThis section includes wet-pipe sprinkler systems.")
    doc.add_paragraph("PART 2 - PRODUCTS", style="Heading 1")
    doc.add_paragraph("2.01\tPIPING")
    _table(doc, [("Pipe", "Schedule 40"), ("Fittings", "Grooved")])
    doc.add_paragraph("2.02\tSPRINKLERS")
    doc.add_paragraph("A.\tSprinklers shall be quick response.")

    appendix = _new_section(doc)
    appendix.orientation = WD_ORIENT.LANDSCAPE
    appendix.page_width, appendix.page_height = appendix.page_height, appendix.page_width
    doc.add_paragraph("APPENDIX A - SCHEDULE")
    _table(doc, [("Mark", "Type"), ("S1", "Pendent")])
    doc.add_paragraph("Notes: see the drawings.")

    end = _new_section(doc)
    end.orientation = WD_ORIENT.PORTRAIT
    end.page_width, end.page_height = end.page_height, end.page_width
    end.footer.is_linked_to_previous = False
    end.footer.paragraphs[0].add_run("END FOOTER")
    doc.add_paragraph("PART 3 - EXECUTION", style="Heading 1")
    doc.add_paragraph("3.01\tINSTALLATION")
    doc.add_paragraph("END OF SECTION")

    # the last section names the body's header part itself: two sections share one header part (python-docx's Section
    # objects all wrap the body's last sectPr as it grows: the body's own is the second in document order)
    body_ref = doc.sections[1]._sectPr.find(qn("w:headerReference"))
    shared = OxmlElement("w:headerReference")
    shared.set(qn("w:type"), "default")
    shared.set(qn("r:id"), body_ref.get(qn("r:id")))
    doc.sections[3]._sectPr.insert(0, shared)
    return save(doc)


def build_single(title_page: bool = True) -> bytes:
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(10)
    sec = doc.sections[0]
    sec.header.is_linked_to_previous = False
    sec.header.paragraphs[0].add_run("SPEC HEADER")
    sec.footer.is_linked_to_previous = False
    p = sec.footer.paragraphs[0]
    p.add_run("Page ")
    _page_field(p)
    if title_page:
        sec.different_first_page_header_footer = True
        sec.first_page_header.is_linked_to_previous = False
        sec.first_page_header.paragraphs[0].add_run("FIRST PAGE")
    doc.add_paragraph("INTRODUCTION", style="Heading 1")
    for k in range(1, 4):
        doc.add_paragraph(f"1.0{k}\tArticle {k}. " + "The work of this article is described here. " * 6)
    doc.add_paragraph("APPENDIX A", style="Heading 1")
    doc.add_paragraph("Appendix A text. " * 10)
    doc.add_paragraph("APPENDIX B", style="Heading 1")
    doc.add_paragraph("Appendix B text. " * 10)
    return save(doc)


def build_page_break(text_after: bool = False) -> bytes:
    """One section: "INTRODUCTION", three paragraphs of body text, a paragraph holding only a manual page break (the
    usual way an appendix is put on its own page; ``text_after``: the break followed by "Continued." in the same
    paragraph), "APPENDIX A" and its text."""
    doc = Document()
    doc.add_paragraph("INTRODUCTION", style="Heading 1")
    for _ in range(3):
        doc.add_paragraph("Body text. " * 30)
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)
    if text_after:
        p.add_run("Continued.")
    doc.add_paragraph("APPENDIX A", style="Heading 1")
    doc.add_paragraph("Appendix text. " * 20)
    return save(doc)


def build_columns_text() -> bytes:
    """One section of plain text long enough to fill two columns of a page (the LibreOffice column check)."""
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(10)
    doc.add_paragraph("COLUMN TEST", style="Heading 1")
    for k in range(1, 25):
        doc.add_paragraph(f"Paragraph {k}: " + "Sprinkler piping shall be installed as shown. " * 3)
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Twips(1440)
    return save(doc)


# ---- surgery the tests apply -------------------------------------------------------------------------------------

def _x(xml: str):
    return etree.fromstring(f'<w:x xmlns:w="{W_NS}" xmlns:r="{R_NS}">{xml}</w:x>')[0]


def text_of(p) -> str:
    """A paragraph's words, a tab in a run read as a tab."""
    return "".join("\t" if el.tag == W + "tab" else (el.text or "") for el in p.iter(W + "t", W + "tab")
                   if el.tag == W + "t" or el.getparent().tag == W + "r")


def body_paragraph(root, text: str):
    return next(p for p in root.iter(W + "p") if text_of(p).startswith(text))


def sect_prs(root):
    body = root.find(W + "body")
    out = [sp for p in body.iter(W + "p") for sp in p.findall(f"{W}pPr/{W}sectPr")]
    return out + [body.find(W + "sectPr")]


def with_section_xml(data: bytes, section: int, xml: str) -> bytes:
    """``data`` with ``xml`` (sectPr children) appended to section ``section``'s sectPr, each where the schema puts it
    (before a w:sectPrChange; the fixtures' sectPr children are in schema order)."""
    order = ["headerReference", "footerReference", "footnotePr", "endnotePr", "type", "pgSz", "pgMar", "paperSrc",
             "pgBorders", "lnNumType", "pgNumType", "cols", "formProt", "vAlign", "noEndnote", "titlePg",
             "textDirection", "bidi", "rtlGutter", "docGrid", "printerSettings", "sectPrChange"]

    def edit(root):
        sp = sect_prs(root)[section]
        for el in etree.fromstring(f'<w:x xmlns:w="{W_NS}" xmlns:r="{R_NS}">{xml}</w:x>'):
            name = etree.QName(el).localname
            for old in sp.findall(W + name):
                sp.remove(old)
            rank = order.index(name)
            later = [c for c in sp if etree.QName(c).localname in order and order.index(etree.QName(c).localname) > rank]
            if later:
                later[0].addprevious(el)
            else:
                sp.append(el)
    return patch_docx(data, {MAIN: edit})


def with_tracked_mark(data: bytes, text: str, kind: str = "ins") -> bytes:
    """The paragraph starting with ``text`` with its mark a pending tracked insertion (or deletion)."""
    def edit(root):
        p = body_paragraph(root, text)
        ppr = p.find(W + "pPr")
        if ppr is None:
            ppr = etree.SubElement(p, W + "pPr")
            p.insert(0, ppr)
        rpr = ppr.find(W + "rPr")
        if rpr is None:
            rpr = etree.SubElement(ppr, W + "rPr")
        rpr.append(_x(f'<w:{kind} w:id="901" w:author="Jack" w:date="2026-09-23T08:00:00Z"/>'))
    return patch_docx(data, {MAIN: edit})


def in_content_control(data: bytes, text: str) -> bytes:
    """The paragraph starting with ``text`` wrapped in a block content control."""
    def edit(root):
        p = body_paragraph(root, text)
        sdt = _x("<w:sdt><w:sdtPr/><w:sdtContent/></w:sdt>")
        p.addprevious(sdt)
        sdt.find(W + "sdtContent").append(p)
    return patch_docx(data, {MAIN: edit})


def with_textbox(data: bytes, after: str, text: str) -> bytes:
    """A text box holding a paragraph ``text`` in a run of the paragraph starting with ``after``."""
    def edit(root):
        p = body_paragraph(root, after)
        p.append(_x('<w:r><w:pict><w:txbxContent><w:p><w:r><w:t>%s</w:t></w:r></w:p></w:txbxContent></w:pict></w:r>'
                    % text))
    return patch_docx(data, {MAIN: edit})


def with_section_change(data: bytes, section: int) -> bytes:
    """Section ``section``'s sectPr with a pending tracked change of its properties (an older page size)."""
    return with_section_xml(data, section, '<w:sectPrChange w:id="902" w:author="Jack" w:date="2026-09-23T08:00:00Z">'
                                           '<w:sectPr><w:pgSz w:w="11906" w:h="16838"/></w:sectPr></w:sectPrChange>')
