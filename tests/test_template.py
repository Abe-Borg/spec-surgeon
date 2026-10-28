"""Phase 2: apply_header_footer_template end to end on the template fixture (make_fixtures.build_template) and
build() targets — both slot conventions, tab rescaling, pinned properties, style import, old parts removed — the
refusals, and planted defects: wrong outputs AND wrong implementations shared by plan and apply, each caught."""
from __future__ import annotations

import copy
import hashlib
import io
import shutil
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import pytest
from lxml import etree

from make_fixtures import (TEMPLATE_FILLS, WRAPPING_FILLS, _styleref_field, build, build_rich, build_template,
                           build_wrapping_template)
from spec_surgeon.batch import run_batch
from spec_surgeon.compiler import _Batch, check_submission
from spec_surgeon.editspec import apply_spec, load_spec, plan_spec
from spec_surgeon.ooxml import block_full_text, paragraphs
from spec_surgeon.ops import JobException, get_op, merge_fragments
from spec_surgeon.ops import apply_header_footer_template as tmpl_module
from spec_surgeon.ops.apply_header_footer_template import ApplyHeaderFooterTemplate
from spec_surgeon.ops.effective import AmbiguousToggle, Resolver, theme_fonts
from spec_surgeon.ops.selectors import StyleSheet
from spec_surgeon.ops.slots import extract_fill
from spec_surgeon.package import CT_PART, DocxPackage, W, serialize_xml
from spec_surgeon.verify import run_suite
from spec_surgeon.verify.reader import VPackage
from spec_surgeon.verify.template import Eff, _v_field_runs, fill_value, slots_of
from test_adaptive_messy import _build_messy_inline_spec, mutate_xml

OP = "apply_header_footer_template"
SPEC = Path(__file__).resolve().parent.parent / "examples" / "specs" / "apply-header-footer-template.json"
WP_DOCPR = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}docPr"
STAMP = TEMPLATE_FILLS["ISSUED FOR PERMIT"]["value"]


@pytest.fixture(scope="module")
def template(tmp_path_factory) -> Path:
    p = tmp_path_factory.mktemp("template") / "00_00_00_TEMPLATE.docx"
    p.write_bytes(build_template())
    return p


@pytest.fixture(scope="module")
def target() -> bytes:
    return build("21 05 00", "COMMON WORK RESULTS FOR FIRE SUPPRESSION")


def params(tpl: Path, **kw) -> dict:
    d = {"template": str(tpl), "template_sections": {"cover": 0, "body": 1},
         "slot_fills": copy.deepcopy(TEMPLATE_FILLS)}
    d.update(kw)
    return d


def run(src: bytes, prm: dict, apply_with=None):
    """Plan on the source, apply on a working copy (``apply_with`` swaps in a wrong implementation), verify."""
    pkg = DocxPackage(src)
    frag = get_op(OP).plan(pkg, {"sections": "all"}, prm)
    plan = merge_fragments([frag])
    work = pkg.copy()
    (apply_with or get_op(OP)).apply(work, frag)
    out = work.output_bytes()
    return plan, out, run_suite(src, out, plan.to_dict(), "src.docx", "out.docx")


def failed(rec):
    return [c.id for c in rec.checks if c.status == "FAIL"]


def mutate(docx: bytes, part: str, fn) -> bytes:
    zin = zipfile.ZipFile(io.BytesIO(docx))
    root = etree.fromstring(zin.read(part))
    fn(root)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        for info in zin.infolist():
            z.writestr(info, serialize_xml(root) if info.filename == part else zin.read(info.filename))
    return buf.getvalue()


def texts(docx: bytes, part: str):
    return [block_full_text(p) for p in paragraphs(DocxPackage(docx).xml(part))]


def direct_tabs(docx: bytes, part: str, block: int):
    p = paragraphs(DocxPackage(docx).xml(part))[block]
    return [(t.get(W + "val"), int(t.get(W + "pos"))) for t in p.iter(W + "tab") if t.getparent().tag == W + "tabs"]


def run_with_text(root, needle):
    return [r for r in root.iter(W + "r") if any(needle in (t.text or "") for t in r.iter(W + "t"))][0]


# ---- end to end ------------------------------------------------------------------------------------

def test_template_applied_to_a_two_section_spec(template, target):
    t_md5 = hashlib.md5(template.read_bytes()).hexdigest()
    plan, out, rec = run(target, params(template))
    assert rec.passed, rec.to_markdown()
    assert plan.mode == "untracked"                                   # the literal stamp is user-supplied words
    assert set(plan.members_removed) == {"word/header1.xml", "word/footer1.xml", "word/footer2.xml"}
    assert plan.members_added == ["word/header2.xml", "word/_rels/header2.xml.rels", "word/media/image1.png",
                                  "word/footer3.xml", "word/header3.xml", "word/footer4.xml"]
    o = DocxPackage(out)
    secs = o.sections()
    assert (secs[0].header_refs, secs[0].footer_refs) == ({"default": "word/header2.xml"}, {"default": "word/footer3.xml"})
    assert (secs[1].header_refs, secs[1].footer_refs) == ({"default": "word/header3.xml"}, {"default": "word/footer4.xml"})
    assert not o.has("word/header1.xml") and not o.has("word/footer1.xml")
    # the fills: extracted from the spec itself (number from its footer, title after "SECTION nn nn nn")
    assert texts(out, "word/header2.xml")[0] == "MICROSOFT — AVC11 DATA CENTER" + STAMP
    assert texts(out, "word/footer3.xml") == ["SECTION 21 05 00"]       # one highlighted slot over three runs
    assert texts(out, "word/header3.xml") == ["21 05 00 – COMMON WORK RESULTS FOR FIRE SUPPRESSION" + STAMP]
    body_footer = texts(out, "word/footer4.xml")
    assert body_footer[0].startswith("21 05 00 - ") and "PAGE" in body_footer[0]
    assert body_footer[1:3] == ["© Owner", "21 05 00"]                 # the slot inside the footer table
    # slot runs: one run per slot, its highlight gone, the template's character style kept (and imported)
    for part in plan.members_added:
        if part.endswith(".xml") and "_rels" not in part:
            assert next(o.xml(part).iter(W + "highlight"), None) is None, part
    stamp = run_with_text(o.xml("word/header3.xml"), STAMP)
    assert stamp.find(f"{W}rPr/{W}rStyle").get(W + "val") == "HFStamp"
    assert [s.get(W + "styleId") for s in o.xml("word/styles.xml").findall(W + "style")].count("HFStamp") == 1
    assert "HFStamp" in " ".join(plan.notes)
    # tabs: the cover header's right stop follows the right margin (9360 → 8640); the style's stops are cleared
    assert sorted(direct_tabs(out, "word/header2.xml", 0)) == [("clear", 4680), ("clear", 9360), ("right", 8640)]
    # the body header's style stops are rescaled and pinned: center scales, right keeps its margin distance
    assert sorted(direct_tabs(out, "word/header3.xml", 0)) == [("center", 4320), ("clear", 4680), ("clear", 9360),
                                                               ("right", 8640)]
    # pinned: the footer renders at the template's 7 pt although the spec's Footer style says 10 pt
    for needle in ("21 05 00", " - "):                      # the slot's fill run and the template's own run
        sz = run_with_text(o.xml("word/footer4.xml"), needle).find(f"{W}rPr/{W}sz")
        assert sz is not None and sz.get(W + "val") == "14"
    # the logo: a new media part, a fresh drawing id, the table-width limitation flagged
    ids = [e.get("id") for part in o.text_parts() for e in o.xml(part).iter(WP_DOCPR)]
    assert len(ids) == len(set(ids)) == 1
    assert any("table widths are not rescaled" in f for f in plan.flags)
    assert any("literal fill(s) ['ISSUED FOR PERMIT']" in f for f in plan.flags)
    # the body (revision-history table included) is untouched, and so is the template file
    assert texts(out, "word/document.xml") == texts(target, "word/document.xml")
    assert hashlib.md5(template.read_bytes()).hexdigest() == t_md5
    v10 = rec.get("V10").summary
    assert "hf_template_refs" in v10 and "styles_added" in v10 and "template_parts word/header3.xml" in v10
    assert "baseline" in rec.get("V5a").summary or rec.get("V5a").status == "PASS"


def test_pending_revision_in_an_old_footer_is_predicted_and_flagged(template):
    src = build("21 13 13", "WET-PIPE SPRINKLER SYSTEMS", footer_revision=True)
    plan, out, rec = run(src, params(template))
    assert rec.passed, rec.to_markdown()
    assert plan.rev_delta["ins"] == -1
    assert any("Jeremy" in f and "discards" in f for f in plan.flags)
    assert texts(out, "word/header3.xml") == ["21 13 13 – WET-PIPE SPRINKLER SYSTEMS" + STAMP]


def test_narrow_target_gets_tabs_rescaled_to_its_own_width(template):
    src = build("21 13 16", "DRY-PIPE SPRINKLER SYSTEMS", side_margin_twips=2160)     # text width 7920
    plan, out, rec = run(src, params(template))
    assert rec.passed, rec.to_markdown()
    assert sorted(direct_tabs(out, "word/header2.xml", 0)) == [("clear", 4680), ("clear", 9360), ("right", 7920)]
    assert ("center", 3960) in direct_tabs(out, "word/header3.xml", 0)              # 4680 × 7920 / 9360
    assert "tabs 9360→7920" in rec.get("V10").summary


def test_rescaling_off_keeps_the_template_stops(template, target):
    plan, out, rec = run(target, params(template, rescale_tabs=False))
    assert rec.passed, rec.to_markdown()
    assert sorted(direct_tabs(out, "word/header2.xml", 0)) == [("clear", 4680), ("right", 9360)]


def test_one_template_section_for_all_sections_shares_one_copy(template, target):
    prm = params(template, template_sections={"all": 1})
    del prm["slot_fills"]["XX XX XX"]                      # the cover footer's slot is not used any more
    plan, out, rec = run(target, prm)
    assert rec.passed, rec.to_markdown()
    secs = DocxPackage(out).sections()
    assert secs[0].header_refs == secs[1].header_refs == {"default": "word/header2.xml"}
    assert secs[0].footer_refs == secs[1].footer_refs == {"default": "word/footer3.xml"}
    assert not any(m.startswith("word/media/") for m in plan.members_added)        # the logo is cover-only


def test_footers_only_leaves_the_headers_alone(template, target):
    prm = params(template, kinds=["footer"])
    for unused in ("SECTION_TITLE", "ISSUED FOR PERMIT"):
        del prm["slot_fills"][unused]
    plan, out, rec = run(target, prm)
    assert rec.passed, rec.to_markdown()
    assert plan.mode == "formatting"                       # no literal fill left
    secs = DocxPackage(out).sections()
    assert secs[0].header_refs == {"default": "word/header1.xml"}                    # untouched
    assert set(plan.members_removed) == {"word/footer1.xml", "word/footer2.xml"}
    assert texts(out, "word/header1.xml") == texts(target, "word/header1.xml")


def test_header_footer_distance_follows_the_template(template):
    src = build("21 22 00", "CLEAN-AGENT FIRE-EXTINGUISHING SYSTEMS", header_distance_twips=1440)
    plan, out, rec = run(src, params(template))
    assert rec.passed, rec.to_markdown()
    for sp in DocxPackage.sect_prs(DocxPackage(out).xml("word/document.xml")):
        assert sp.find(W + "pgMar").get(W + "header") == "720"
    assert any("distance" in n for n in plan.notes) and "section_props" in rec.get("V10").summary
    plan, out, rec = run(src, params(template, align_header_footer_distance=False))
    assert rec.passed, rec.to_markdown()                    # not changed, not asserted — V7 is what sees it (below)
    for sp in DocxPackage.sect_prs(DocxPackage(out).xml("word/document.xml")):
        assert sp.find(W + "pgMar").get(W + "header") == "1440"


# ---- refusals ----------------------------------------------------------------------------------------

def _tpl_variant(template: Path, tmp_path: Path, part: str, fn) -> Path:
    p = tmp_path / "variant.docx"
    p.write_bytes(mutate(template.read_bytes(), part, fn))
    return p


def _plan(src, prm, select=None):
    return get_op(OP).plan(DocxPackage(src), select or {"sections": "all"}, prm)


def _ins_wrap(root):
    r = root.find(f".//{W}r")
    ins = etree.Element(W + "ins", {W + "id": "1", W + "author": "Jack", W + "date": "2026-01-01T00:00:00Z"})
    r.addprevious(ins)
    ins.append(r)


def _overlapping_slot(root):
    r = run_with_text(root, "ISSUED FOR PERMIT")
    r.find(W + "t").text = "{{STAMP}} ISSUED"


def _tab_in_slot(root):
    r = run_with_text(root, "ISSUED FOR PERMIT")
    r.find(W + "t").addprevious(etree.Element(W + "tab"))


def _bookmark(root):
    p = root.find(f".//{W}p")
    p.append(etree.Element(W + "bookmarkStart", {W + "id": "0", W + "name": "x"}))


@pytest.mark.parametrize("change, code", [
    (lambda p: p["slot_fills"].pop("SECTION_TITLE"), "unfilled-slot"),
    (lambda p: p["slot_fills"].update(PROJECT={"rule": "literal", "value": "AVC11"}), "unknown-slot"),
    (lambda p: p["slot_fills"].update(SECTION_NUMBER={"rule": "regex", "source": "document", "pattern": r"^(\d+)\.01"}),
     "ambiguous"),
    (lambda p: p["slot_fills"].update(SECTION_NUMBER={"rule": "regex", "source": "document", "pattern": r"^DIVISION (\d+)"}),
     "not-found"),
    (lambda p: p["slot_fills"].update(SECTION_TITLE={"rule": "paragraph_after", "anchor_regex": r"\d\.01"}), "ambiguous"),
    (lambda p: p["slot_fills"].update(SECTION_NUMBER={"rule": "regex", "source": "footer.first", "pattern": "x"}),
     "not-found"),
    (lambda p: p["slot_fills"].update(SECTION_NUMBER={"rule": "guess"}), "bad-params"),
    (lambda p: p["slot_fills"].update(SECTION_NUMBER={"rule": "literal", "value": "   "}), "not-found"),
    (lambda p: p.update(template_sections={"cover": 0, "body": 5}), "bad-params"),
    (lambda p: p.update(template_sections={"first": 0}), "bad-params"),
    (lambda p: p.update(template="no such template.docx"), "bad-params"),
])
def test_refusals(template, target, change, code):
    prm = params(template)
    change(prm)
    with pytest.raises(JobException) as e:
        _plan(target, prm)
    assert e.value.code == code, e.value.reason


def test_refuses_a_partial_section_selection(template, target):
    with pytest.raises(JobException) as e:
        _plan(target, params(template), {"sections": [0]})
    assert e.value.code == "bad-params"


def test_refuses_to_switch_even_odd_pages_silently(template):
    with pytest.raises(JobException) as e:                  # build_rich has evenAndOddHeaders on, the template off
        _plan(build_rich(), params(template))
    assert e.value.code == "even-odd"


@pytest.mark.parametrize("part, fn, code", [
    ("word/header2.xml", _ins_wrap, "template-revisions"),
    ("word/header2.xml", _overlapping_slot, "slot-overlap"),
    ("word/header2.xml", _tab_in_slot, "slot-structure"),
    ("word/footer1.xml", _bookmark, "structure"),
])
def test_refuses_template_constructs_it_does_not_handle(template, target, tmp_path, part, fn, code):
    variant = _tpl_variant(template, tmp_path, part, fn)
    prm = params(template, template=str(variant))
    if code == "slot-overlap":
        prm["slot_fills"]["STAMP"] = {"rule": "literal", "value": "x"}
    with pytest.raises(JobException) as e:
        _plan(target, prm)
    assert e.value.code == code, e.value.reason


def test_template_changed_between_plan_and_apply_is_refused(template, target, tmp_path):
    t = tmp_path / "T.docx"
    t.write_bytes(template.read_bytes())
    pkg = DocxPackage(target)
    frag = get_op(OP).plan(pkg, {"sections": "all"}, params(template, template=str(t)))
    t.write_bytes(mutate(t.read_bytes(), "word/header2.xml", _tab_in_slot))
    with pytest.raises(JobException) as e:
        get_op(OP).apply(pkg.copy(), frag)
    assert e.value.code == "internal"


# ---- planted defects -------------------------------------------------------------------------------

def test_typed_fill_in_the_output_is_caught_by_v2_and_v10(template, target):
    plan, out, rec = run(target, params(template))
    assert rec.passed

    def typo(root):
        t = run_with_text(root, "COMMON WORK").find(W + "t")
        t.text = t.text.replace("COMMON WORK", "COMMON WORKS")
    bad = mutate(out, "word/header3.xml", typo)
    rec = run_suite(target, bad, plan.to_dict(), "s", "o")
    assert "V2" in failed(rec) and "V10" in failed(rec)


def test_fill_typed_by_plan_and_apply_alike_is_caught_by_v10(template, target, monkeypatch):
    """The transform's fill extraction is wrong in plan AND apply (so V2 sees exactly the predicted text): the
    title of another spec is placed. V10 re-extracts every fill from the source by its rule and disagrees."""
    real = tmpl_module.extract_fill
    monkeypatch.setattr(tmpl_module, "extract_fill",
                        lambda pkg, name, rule: "WET-PIPE SPRINKLER SYSTEMS" if name == "SECTION_TITLE" else real(pkg, name, rule))
    plan, out, rec = run(target, params(template))
    assert rec.get("V2").status == "PASS"
    assert failed(rec) == ["V10"] and "re-extracted" in rec.get("V10").summary


class _NoContentTypeOverride(ApplyHeaderFooterTemplate):
    """Wrong: the new body footer is added without its content-type override (typed application/xml by default)."""
    def apply(self, pkg, frag):
        super().apply(pkg, frag)
        root = pkg.xml(CT_PART)
        for o in list(root):
            if o.get("PartName") == "/word/footer4.xml":
                root.remove(o)
        pkg.write_xml(CT_PART, root)


def test_missing_content_type_override_is_caught_by_v1(template, target):
    _, _, rec = run(target, params(template), apply_with=_NoContentTypeOverride())
    assert "V1" in failed(rec) and "content type of word/footer4.xml" in rec.get("V1").summary


def test_tabs_not_rescaled_by_plan_and_apply_alike_is_caught(template, target, monkeypatch):
    monkeypatch.setattr(tmpl_module, "rescale_tabs", lambda tabs, a, b: tuple(tabs))
    _, out, rec = run(target, params(template))
    assert ("right", 9360) in direct_tabs(out, "word/header2.xml", 0)
    assert failed(rec) == ["V10"] and "effective paragraph properties" in rec.get("V10").summary


def test_unpinned_font_size_by_plan_and_apply_alike_is_caught(template, target, monkeypatch):
    """The pinning diff forgets sz in plan AND apply: the copied footer would render at the spec's 10 pt."""
    monkeypatch.setattr(tmpl_module, "RUN_KEYS", tuple(k for k in tmpl_module.RUN_KEYS if k != "sz"))
    _, _, rec = run(target, params(template))
    assert failed(rec) == ["V10"] and "effective run properties" in rec.get("V10").summary


class _DropsPins(ApplyHeaderFooterTemplate):
    """Wrong: apply skips the pins of the body footer."""
    def apply(self, pkg, frag):
        for c in frag.transform_data["copies"]:
            if c["part"] == "word/footer4.xml":
                c["pins"] = {}
        super().apply(pkg, frag)


def test_pins_dropped_by_apply_are_caught(template, target):
    _, _, rec = run(target, params(template), apply_with=_DropsPins())
    assert "V10" in failed(rec) and "word/footer4.xml" in rec.get("V10").summary


class _NoDocPrRenumbering(ApplyHeaderFooterTemplate):
    def apply(self, pkg, frag):
        for c in frag.transform_data["copies"]:
            c["docpr_start"] = 1
        super().apply(pkg, frag)


def test_drawing_id_collision_is_caught(template):
    src = build("21 05 13", "COMMON MOTOR REQUIREMENTS", body_picture=True)
    plan, out, rec = run(src, params(template))
    assert rec.passed, rec.to_markdown()
    _, _, rec = run(src, params(template), apply_with=_NoDocPrRenumbering())
    assert "V10" in failed(rec) and "docPr" in rec.get("V10").summary


def test_imported_style_altered_is_caught(template, target):
    plan, out, rec = run(target, params(template))

    def alter(root):
        st = [s for s in root.findall(W + "style") if s.get(W + "styleId") == "HFStamp"][0]
        st.find(f".//{W}rFonts").set(W + "ascii", "Arial")
    rec = run_suite(target, mutate(out, "word/styles.xml", alter), plan.to_dict(), "s", "o")
    assert "V10" in failed(rec) and "HFStamp differs" in rec.get("V10").summary


class _KeepsOldParts(ApplyHeaderFooterTemplate):
    """Wrong: the old headers/footers' references and rels go but the parts stay in the package."""
    def apply(self, pkg, frag):
        frag.transform_data["removal"]["members"] = []
        super().apply(pkg, frag)


def test_old_parts_left_behind_are_caught_by_v1(template, target):
    _, _, rec = run(target, params(template), apply_with=_KeepsOldParts())
    assert "V1" in failed(rec) and "ORPHANED" in rec.get("V1").summary and "word/header1.xml" in rec.get("V1").summary


def test_highlight_left_on_a_fill_is_caught(template, target, monkeypatch):
    real = tmpl_module.pin_rpr
    def leave_highlight(rpr, pins):
        real(rpr, pins)
        if "color" in pins:  # clear-mode slot decoration is pinned after carving the fill run
            etree.SubElement(rpr, W + "highlight").set(W + "val", "yellow")
    monkeypatch.setattr(tmpl_module, "pin_rpr", leave_highlight)
    _, _, rec = run(target, params(template))
    assert failed(rec) == ["V10"] and "slot highlight must go" in rec.get("V10").summary


class _SwapsCoverAndBody(ApplyHeaderFooterTemplate):
    def apply(self, pkg, frag):
        nr = frag.transform_data["new_refs"]
        nr["0"], nr["1"] = nr["1"], nr["0"]
        super().apply(pkg, frag)


def test_cover_and_body_swapped_is_caught(template, target):
    _, _, rec = run(target, params(template), apply_with=_SwapsCoverAndBody())
    assert "V10" in failed(rec) and "hf_template_refs" in rec.get("V10").summary


def test_template_edited_after_planning_fails_v10(template, target, tmp_path):
    t = tmp_path / "T.docx"
    t.write_bytes(template.read_bytes())
    plan, out, rec = run(target, params(template, template=str(t)))
    assert rec.passed, rec.to_markdown()
    t.write_bytes(mutate(t.read_bytes(), "word/header2.xml", _tab_in_slot))
    rec = run_suite(target, out, plan.to_dict(), "s", "o")
    assert "V10" in failed(rec) and "template file changed" in rec.get("V10").summary


# ---- the two resolvers agree ------------------------------------------------------------------------

def test_transform_and_verifier_effective_resolvers_agree(template):
    """ops/effective.Resolver and verify/template.Eff implement one definition (CLAUDE.md) separately: on every
    paragraph, paragraph mark and run of the template, a target and the rich fixture they give the same answer
    (or both refuse an ambiguous toggle)."""
    checked = 0
    for data in (template.read_bytes(), build("21 05 00", "A"), build_rich()):
        pkg, vpkg = DocxPackage(data), VPackage(data)
        res, eff = Resolver(StyleSheet(pkg), theme_fonts(pkg)), Eff(vpkg)
        for part in pkg.text_parts():
            for p, vp in zip(paragraphs(pkg.xml(part)), vpkg.xml(part).iter(W + "p")):
                assert res.paragraph(p) == eff.para(vp)
                assert res.paragraph(p, direct=False) == eff.para(vp, with_direct=False)
                pairs = [(p.find(f"{W}pPr/{W}rPr"), vp.find(f"{W}pPr/{W}rPr"))]
                pairs += list(zip([r.find(W + "rPr") for r in p.iter(W + "r")],
                                  [r.find(W + "rPr") for r in vp.iter(W + "r")]))
                for rpr, vrpr in pairs:
                    try:
                        a = res.run(rpr, p)
                    except AmbiguousToggle:
                        with pytest.raises(ValueError):
                            eff.run(vrpr, vp)
                        continue
                    b = eff.run(vrpr, vp)
                    assert {k: v for k, v in a.items() if k not in ("cs", "szCs")} == b
                    checked += 1
    assert checked > 100


# ---- the EditSpec example through the batch ---------------------------------------------------------

def test_example_spec_runs_in_a_batch_and_skips_the_template(tmp_path):
    spec = load_spec(SPEC)
    tpl = Path(spec["ops"][0]["params"]["template"])
    assert tpl.is_file()
    src = tmp_path / "in"
    src.mkdir()
    (src / "21 05 00 – COMMON WORK.docx").write_bytes(build("21 05 00", "COMMON WORK"))
    (src / "21 13 13 – WET-PIPE.docx").write_bytes(build("21 13 13", "WET-PIPE", footer_revision=True))
    (src / "21 13 16 – DRY-PIPE.docx").write_bytes(build("21 13 16", "DRY-PIPE", side_margin_twips=2160))
    shutil.copy(tpl, src / "00 00 00 – TEMPLATE copy.docx")          # sits in the project folder: skipped by md5
    batch = run_batch([str(src)], "edit-spec", {}, tmp_path / "out", workers=1, renderer="none", do_render=False,
                      spec=spec)
    assert len(batch.files) == 3 and batch.skipped == [str(src / "00 00 00 – TEMPLATE copy.docx")]
    for f in batch.files:
        assert f.status == "PASS", f.record.to_markdown() if f.record else f.exception


def test_hf_only_copy_keeps_only_section_breaks(template, tmp_path):
    """The V7 reference documents: the body is three empty pages per section (pageBreakBefore), every section
    nextPage with numbering restarted at 1 and otherwise as it was, every other member byte-identical."""
    from spec_surgeon.verify.geometry import hf_only_copy
    hf, n = hf_only_copy(template, tmp_path / "hf.docx")
    src, out = VPackage(template.read_bytes()), VPackage(hf.read_bytes())
    assert src.names == out.names and n == 2
    assert all(src.read(m) == out.read(m) for m in src.names if m != "word/document.xml")
    body = out.xml("word/document.xml").find(W + "body")
    assert not list(body.iter(W + "t")) and not list(body.iter(W + "tbl"))
    rest = lambda sp: [etree.tostring(c) for c in sp if c.tag not in (W + "type", W + "pgNumType")]
    pairs = list(zip(src.xml("word/document.xml").iter(W + "sectPr"), body.iter(W + "sectPr")))
    assert len(pairs) == n
    for a, b in pairs:
        assert rest(a) == rest(b)
        assert b.find(W + "type").get(W + "val") == "nextPage" and b.find(W + "pgNumType").get(W + "start") == "1"
    assert len(list(body.iter(W + "pageBreakBefore"))) == 2 * n


def _pdf(path: Path, pages):
    """A PDF whose pages carry the given (text, baseline y) words — a stand-in for a headers/footers-only render."""
    import pymupdf
    doc = pymupdf.open()
    for words in pages:
        pg = doc.new_page(width=612, height=792)
        for text, y in words:
            pg.insert_text((72, y), text, fontsize=8)
    doc.save(str(path))
    return path


def test_v7_template_pages_are_paired_not_searched(tmp_path):
    """Codex review on PR #2: with a "best template page" search, a blank template variant (a cover's first page
    with titlePg and no first-page header) matched every output page with no problems, so a displaced header
    passed. Pages are paired by section and variant now, and a blank template region requires a blank output."""
    from spec_surgeon.verify.geometry import compare_template_bands
    cover, body = [("COVER", 45), ("SECTION", 760)], [("BODY", 45), ("FOOTER", 760)]
    tpl = _pdf(tmp_path / "t.pdf", [[], cover, cover, body, body, body])          # cover page 1 blank
    good = _pdf(tmp_path / "good.pdf", [[], cover, cover, body, body, body])
    res = compare_template_bands(tpl, good, [0, 1])
    assert not res.failed_pages and [p.template_page for p in res.pages] == [1, 2, 3, 4, 5, 6]
    moved = [("BODY", 81), ("FOOTER", 760)]                                       # header 36pt lower
    res = compare_template_bands(tpl, _pdf(tmp_path / "moved.pdf", [[], cover, cover, moved, moved, moved]), [0, 1])
    assert res.failed_pages == [4, 5, 6] and "no words where the template's header is" in res.pages[3].problems[0]
    res = compare_template_bands(tpl, _pdf(tmp_path / "extra.pdf", [cover, cover, cover, body, body, body]), [0, 1])
    assert res.failed_pages == [1] and "the template shows no header on this page" in res.pages[0].problems[0]
    # only the templated kinds are compared: a footers-only spec keeps the document's own headers
    res = compare_template_bands(tpl, _pdf(tmp_path / "own.pdf", [[], cover, cover, moved, moved, moved]), [0, 1],
                                 kinds=["footer"])
    assert not res.failed_pages
    # pages that cannot be paired fail rather than pass
    res = compare_template_bands(tpl, _pdf(tmp_path / "short.pdf", [[], cover, cover, body, body]), [0, 1])
    assert res.failed_pages and "cannot be paired" in res.pages[0].problems[0]


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")
def test_v7_template_reference_rendered(tmp_path):
    """V7 against the template's headers/footers rendered alone. With the distances aligned the regions match
    (the body footer included, though it is taller than the bottom margin). With alignment off each document keeps
    its own distances — a header 1" from the edge (template 0.5"), a header 1pt lower, a footer 0.25" higher — and
    the reference takes them from the document, so each copy passes with its header and footer anchored where the
    document has them (the deep-dive stress test's D4: these copies, exactly as asked, were each held on V7)."""
    spec = load_spec(SPEC)
    src = tmp_path / "in"
    src.mkdir()
    (src / "a – header 1in.docx").write_bytes(build("21 22 00", "CLEAN-AGENT", header_distance_twips=1440))
    batch = run_batch([str(src)], "edit-spec", {}, tmp_path / "aligned", workers=1, renderer="libreoffice", spec=spec)
    rec = batch.files[0].record
    assert rec.passed and rec.get("V7").status == "PASS", rec.to_markdown()
    assert "rendered alone" in rec.get("V7").summary
    assert all(p["anchors"] >= 4 for p in rec.get("V7").data["pages"])           # header AND footer anchored
    (src / "b – header 1pt lower.docx").write_bytes(build("21 22 00", "CLEAN-AGENT", header_distance_twips=740))
    (src / "c – footer 0.75in.docx").write_bytes(build("21 22 00", "CLEAN-AGENT", footer_distance_twips=1080))
    off = copy.deepcopy(spec)
    off["ops"][0]["params"]["align_header_footer_distance"] = False
    batch = run_batch([str(src)], "edit-spec", {}, tmp_path / "unaligned", workers=1, renderer="libreoffice", spec=off)
    assert len(batch.files) == 3
    for f in batch.files:
        rec = f.record
        assert rec.passed and rec.get("V7").status == "PASS", rec.to_markdown()
        assert "at its own text widths and its own header/footer distances" in rec.get("V7").summary
        assert all(p["anchors"] >= 4 for p in rec.get("V7").data["pages"]), rec.get("V7").summary


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")
def test_v7_copies_are_named_short_whatever_the_documents_name(tmp_path):
    """The owner's 1.10.0 run: the headers/footers-only copy of a document with an 84-character name, under a
    recovery attempt's folders, reached 261 characters and could not be written on Windows, so its V7 could not be
    made. The copies are numbered, never named after the document or the template (the template reference is made
    per document, from its plan), and V7 still pairs each file with its own."""
    spec = load_spec(SPEC)
    src = tmp_path / "in"
    src.mkdir()
    name = "32_FAA_AC-150-5370-10H_Standard-Specifications-for-Construction-of-Airports.docx"
    (src / name).write_bytes(build("21 22 00", "CLEAN-AGENT"))
    batch = run_batch([str(src)], "edit-spec", {}, tmp_path / "run", workers=1, renderer="libreoffice", spec=spec)
    rec = batch.files[0].record
    assert rec.passed and rec.get("V7").status == "PASS", rec.to_markdown()
    render = tmp_path / "run" / "render"
    assert sorted(p.name for p in (render / "output-hf").glob("*.docx")) == ["000 headers-footers.docx"]
    assert sorted(p.name for p in (render / "template").glob("*.docx")) == ["000 headers-footers.docx"]


def _title_page_on_cover(root):
    sp = root.find(f".//{W}pPr/{W}sectPr")                 # section 0's break
    etree.SubElement(sp, W + "titlePg")


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")
def test_v7_blank_first_page_variant_does_not_hide_a_displaced_header(template, tmp_path):
    """The template's cover has titlePg and no first-page header/footer (its first page renders blank). Aligned,
    the output's blank first page pairs with the template's and V7 passes; with the header distance set wrong by plan
    and apply alike the displaced headers fail V7 on every page that shows one; with the distance not aligned the
    document's own distance is the reference's and V7 passes, the blank first page still paired."""
    variant = _tpl_variant(template, tmp_path, "word/document.xml", _title_page_on_cover)
    spec = load_spec(SPEC)
    spec["ops"][0]["params"]["template"] = str(variant)
    src = tmp_path / "in"
    src.mkdir()
    (src / "a.docx").write_bytes(build("21 22 00", "CLEAN-AGENT", header_distance_twips=1440))
    batch = run_batch([str(src)], "edit-spec", {}, tmp_path / "aligned", workers=1, renderer="libreoffice", spec=spec)
    rec = batch.files[0].record
    assert rec.passed, rec.to_markdown()
    pages = rec.get("V7").data["pages"]
    assert pages[0]["anchors"] == 0 and not pages[0]["problems"]                  # blank ↔ blank
    with pytest.MonkeyPatch.context() as mp:
        _distances_set_wrong(mp, "1440")
        batch = run_batch([str(src)], "edit-spec", {}, tmp_path / "wrong", workers=1, renderer="libreoffice",
                          spec=spec)
    rec = batch.files[0].record
    assert failed(rec) == ["V7"], rec.to_markdown()
    bad = [p["page"] for p in rec.get("V7").data["pages"] if p["problems"]]
    assert bad == [2, 3, 4, 5, 6]
    off = copy.deepcopy(spec)
    off["ops"][0]["params"]["align_header_footer_distance"] = False
    batch = run_batch([str(src)], "edit-spec", {}, tmp_path / "unaligned", workers=1, renderer="libreoffice", spec=off)
    rec = batch.files[0].record
    assert rec.passed and rec.get("V7").status == "PASS", rec.to_markdown()
    pages = rec.get("V7").data["pages"]
    assert pages[0]["anchors"] == 0 and not pages[0]["problems"]


def _wrapping_spec(tpl: Path, **params) -> dict:
    return {"request": "make the header/footer look like this template", "summary": "the template's header/footer",
            "assumptions": [], "questions": [], "applies_to": {"or": "all"}, "mode": "formatting",
            "ops": [{"op": OP, "select": {"sections": "all"},
                     "params": {"template": str(tpl), "template_sections": {"all": 0},
                                "slot_fills": copy.deepcopy(WRAPPING_FILLS), **params}}],
            "review": {"sample_files": 2, "geometry_reference": "template"}}


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")
def test_v7_template_reference_carries_each_documents_fills_and_text_width(tmp_path, monkeypatch):
    """Live QA finding F08: V7 compared the filled header with the unfilled template, so a fill shorter than its
    placeholder let a word up onto an earlier line and a correct copy was held for an 11.50 pt "shift". The template
    reference is now the template as each document's plan makes it show: its fills in the slots, laid at its own
    text width. Over a header that wraps: short fills pull words up (a), a long title pushes them down within the
    template's own lines (b), a narrower page wraps the line elsewhere (c) — each fails against the unfilled template
    and every copy passes, every word of its header (the fills among them) anchored at the reference's height."""
    from spec_surgeon import render as render_mod
    from spec_surgeon.verify import geometry as geom
    from spec_surgeon.verify import template_layout
    tpl = tmp_path / "wrapping template.docx"
    tpl.write_bytes(build_wrapping_template())
    src = tmp_path / "in"
    src.mkdir()
    (src / "a.docx").write_bytes(build("21 13 13", "SPRINKLERS"))
    (src / "b.docx").write_bytes(build("21 05 00", "COMMON WORK RESULTS FOR FIRE SUPPRESSION"))
    (src / "c.docx").write_bytes(build("21 13 13", "WET-PIPE SPRINKLER SYSTEMS", side_margin_twips=2520))
    spec = _wrapping_spec(tpl)
    batch = run_batch([str(src)], "edit-spec", {}, tmp_path / "run", workers=1, renderer="libreoffice", spec=spec)
    for f in batch.files:
        assert f.record.passed and f.record.get("V7").status == "PASS", f.record.to_markdown()
        assert all(p["anchors"] >= 30 for p in f.record.get("V7").data["pages"]), f.record.get("V7").summary
        assert "as this document's plan fills them" in f.record.get("V7").summary
    # the unfilled template at its own width is the old reference: every one of the three is a reflow it fails
    old_hf, _ = geom.hf_only_copy(tpl, tmp_path / "old" / "template.docx")
    old_pdf = render_mod.render_pdfs([old_hf], tmp_path / "old", "libreoffice", 1)[old_hf]
    for k, name in enumerate("abc"):
        out_pdf = tmp_path / "run" / "render" / "output-hf" / f"{k:03d} headers-footers.pdf"
        old = geom.compare_template_bands(old_pdf, out_pdf, [0, 0])
        assert old.failed_pages == [1, 2, 3, 4, 5, 6] and "moved up to 11.5" in old.pages[0].problems[0], name
    # the reference laid at the template's width instead of the document's: the narrower page fails again
    monkeypatch.setattr(template_layout, "_width", lambda sp: 8640)
    batch = run_batch([str(src)], "edit-spec", {}, tmp_path / "template width", workers=1, renderer="libreoffice",
                      spec=spec)
    assert {Path(f.source).name: failed(f.record) for f in batch.files} == {"a.docx": [], "b.docx": [],
                                                                           "c.docx": ["V7"]}


def _distances_set_wrong(monkeypatch, header: str) -> None:
    """A planted defect shared by plan and apply: the op sets every section's header distance to ``header`` twips —
    where it aligns them, in place of the template's; where it was told to leave them, anyway — and asserts what it
    did, so V10 passes and only V7, from its own reading of the template and the source, can see it."""
    real = ApplyHeaderFooterTemplate.plan

    def wrong_distance(self, pkg, select, params):
        frag = real(self, pkg, select, params)
        main = pkg.main_part()
        for i in range(len(DocxPackage.sect_prs(pkg.xml(main)))):
            frag.transform_data["dist_changes"].setdefault(str(i), {})["header"] = header
            asserted = [a for a in frag.property_assertions if a["kind"] == "section_props" and a["sections"] == [i]
                        and "pgMar" in a["elements"]]
            if asserted:
                asserted[0]["elements"]["pgMar"]["set"]["header"] = header
            else:
                frag.property_assertions.append({"kind": "section_props", "part": main, "sections": [i],
                                                 "elements": {"pgMar": {"set": {"header": header}}},
                                                 "expected_sections": 1})
                frag.touches.append((main, f"sectPr:{i}", "pgMar"))
        return frag
    monkeypatch.setattr(ApplyHeaderFooterTemplate, "plan", wrong_distance)


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")
def test_v7_template_reference_holds_the_vertical_layout_the_plan_asks_for(tmp_path, monkeypatch):
    """The reference takes the document's fills and text width, and its vertical layout only where the plan keeps the
    document's distances (the deep-dive stress test's D4). Aligned, a header distance set wrong by plan and apply
    alike (1000 twips, not the template's 720: only V7 fails, the XML is exactly as planned) is not where the
    template's header is, on every page. Not aligned, a header left 1" from the edge passes, at the document's own
    distance; and one that plan and apply move to 720 anyway (V10 passes: the plan asserts it) is not where the
    document's header was, and only V7 fails — the reference reads the distance from the source, never from the
    plan or the copy."""
    tpl = tmp_path / "wrapping template.docx"
    tpl.write_bytes(build_wrapping_template())
    src = tmp_path / "in"
    src.mkdir()
    (src / "a.docx").write_bytes(build("21 13 13", "SPRINKLERS", header_distance_twips=1440))
    _distances_set_wrong(monkeypatch, "1000")
    batch = run_batch([str(src)], "edit-spec", {}, tmp_path / "wrong", workers=1, renderer="libreoffice",
                      spec=_wrapping_spec(tpl))
    rec = batch.files[0].record
    assert failed(rec) == ["V7"], rec.to_markdown()
    assert [p["page"] for p in rec.get("V7").data["pages"] if p["problems"]] == [1, 2, 3, 4, 5, 6]
    monkeypatch.undo()
    keep = _wrapping_spec(tpl, align_header_footer_distance=False)
    batch = run_batch([str(src)], "edit-spec", {}, tmp_path / "unaligned", workers=1, renderer="libreoffice",
                      spec=keep)
    rec = batch.files[0].record
    assert rec.passed and rec.get("V7").status == "PASS", rec.to_markdown()
    assert "its own header/footer distances" in rec.get("V7").summary
    _distances_set_wrong(monkeypatch, "720")
    batch = run_batch([str(src)], "edit-spec", {}, tmp_path / "moved anyway", workers=1, renderer="libreoffice",
                      spec=keep)
    rec = batch.files[0].record
    assert failed(rec) == ["V7"], rec.to_markdown()
    assert [p["page"] for p in rec.get("V7").data["pages"] if p["problems"]] == [1, 2, 3, 4, 5, 6]


EXAMPLES = Path(__file__).resolve().parent.parent / "examples" / "specs"
PREPARED_SPEC = EXAMPLES / "apply-header-footer-template.json"


def test_single_section_spec_plans_with_a_cover_and_body_template_and_the_batch_compiles(tmp_path):
    spec = load_spec(PREPARED_SPEC)
    tpl = Path(spec["ops"][0]["params"]["template"])
    single = DocxPackage(build("21 98 00", "VARIANT", single_section=True))
    plan, frags = plan_spec(single, spec)                       # was: unknown-slot (body rules unused here)
    out = apply_spec(single, frags).output_bytes()
    rec = run_suite(single.data, out, plan.to_dict(), "src.docx", "out.docx")
    assert rec.passed, rec.to_markdown()
    assert [a["part"] for a in plan.property_assertions if a["kind"] == "template_parts"]
    folder = tmp_path / "specs"
    folder.mkdir()
    (folder / "21 13 13 – WET-PIPE SPRINKLER SYSTEMS.docx").write_bytes(build("21 13 13", "WET-PIPE SPRINKLER SYSTEMS"))
    (folder / "SINGLE SECTION.docx").write_bytes(single.data)
    batch = _Batch(sorted(folder.glob("*.docx")), tpl)
    assert len(batch.prof["patterns"]) == 2
    assert check_submission(spec, batch, [spec["request"], "ISSUED FOR CONSTRUCTION  2026-02-15"]) == []
    # a rule naming a slot in none of the mapped template sections is still refused, on every file
    spec["ops"][0]["params"]["slot_fills"]["NOWHERE"] = {"rule": "literal", "value": "ISSUED FOR CONSTRUCTION  2026-02-15"}
    for pkg in (single, DocxPackage(build("21 13 13", "WET-PIPE SPRINKLER SYSTEMS"))):
        with pytest.raises(JobException) as e:
            plan_spec(pkg, spec)
        assert e.value.code == "unknown-slot" and "NOWHERE" in e.value.reason
    errs = check_submission(spec, batch, [spec["request"], "ISSUED FOR CONSTRUCTION  2026-02-15"])
    assert errs and "unknown-slot" in errs[0]


def test_fills_keep_the_documents_own_characters_on_both_sides():
    def nbsp_title(root):
        for t in root.find(W + "body").iter(W + "t"):
            if t.text and "WET-PIPE" in t.text:
                t.text = t.text.replace("WET-PIPE SPRINKLER", "WET-PIPE SPRINKLER"); return
    messy = DocxPackage(mutate_xml(_build_messy_inline_spec("21 13 13", "WET-PIPE SPRINKLER SYSTEMS"), "word/document.xml", nbsp_title))
    rx = {"rule": "regex", "source": "document", "pattern": r"^SECTION\s+\d{2}\s+\d{2}\s+\d{2}\s*[-–—]\s*(.+)$"}
    t, v = extract_fill(messy, "T", rx), fill_value(VPackage(messy.data), "T", rx)
    assert t == v == "WET-PIPE SPRINKLER SYSTEMS"
    clean = DocxPackage(mutate_xml(build("21 13 13", "WET-PIPE SPRINKLER SYSTEMS"), "word/document.xml", nbsp_title))
    pa = {"rule": "paragraph_after", "source": "document", "anchor_regex": r"^SECTION \d{2} \d{2} \d{2}$"}
    assert extract_fill(clean, "T", pa) == fill_value(VPackage(clean.data), "T", pa) == "WET-PIPE SPRINKLER SYSTEMS"
    # consensus is still judged with whitespace collapsed: two anchors whose titles differ only in spacing agree
    def two_anchors(root):
        body = root.find(W + "body"); pars = body.findall(W + "p")
        p = etree.Element(W + "p"); etree.SubElement(etree.SubElement(p, W + "r"), W + "t").text = "SECTION 21 13 13"
        p2 = etree.Element(W + "p"); etree.SubElement(etree.SubElement(p2, W + "r"), W + "t").text = "WET-PIPE  SPRINKLER SYSTEMS"
        pars[6].addnext(p); p.addnext(p2)
    two = DocxPackage(mutate_xml(clean.data, "word/document.xml", two_anchors))
    assert extract_fill(two, "T", pa) == fill_value(VPackage(two.data), "T", pa) == "WET-PIPE SPRINKLER SYSTEMS"


def _spec(mode, ops):
    return {"request": "x", "summary": "x", "assumptions": [], "questions": [], "applies_to": {"or": "all"},
            "mode": mode, "ops": ops}


def _template_with_field_slots(tmp_path: Path) -> Path:
    doc = Document()
    s0 = doc.sections[0]
    h = s0.header
    h.is_linked_to_previous = False
    p = h.paragraphs[0]
    p.add_run("SECTION {{SECTION_NUMBER}} – ")
    _styleref_field(p, "{{SECTION_TITLE}}")                 # the title typed as a STYLEREF field's cached result
    p = h.add_paragraph()
    fs = OxmlElement("w:fldSimple"); fs.set(qn("w:instr"), " DOCPROPERTY Title ")
    r = OxmlElement("w:r"); rpr = OxmlElement("w:rPr"); hl = OxmlElement("w:highlight"); hl.set(qn("w:val"), "yellow")
    rpr.append(hl); r.append(rpr); t = OxmlElement("w:t"); t.text = "SECTION_TITLE"; r.append(t); fs.append(r)
    p._p.append(fs)                                         # a highlighted slot inside a simple field
    f = s0.footer
    f.is_linked_to_previous = False
    f.paragraphs[0].add_run("{{SECTION_NUMBER}} - page")
    doc.add_paragraph("TEMPLATE BODY")
    tpl = tmp_path / "field_slots_template.docx"
    buf = io.BytesIO()
    doc.save(buf)
    tpl.write_bytes(buf.getvalue())
    return tpl


def test_a_slot_shown_by_a_field_is_refused_on_both_sides(tmp_path):
    tpl = _template_with_field_slots(tmp_path)
    fills = {"SECTION_NUMBER": {"rule": "regex", "source": "footer.default", "pattern": "^(\\d{2} \\d{2} \\d{2})"},
             "SECTION_TITLE": {"rule": "paragraph_after", "source": "document",
                               "anchor_regex": "^SECTION \\d{2} \\d{2} \\d{2}$"}}
    pkg = DocxPackage(build("21 13 13", "WET-PIPE SPRINKLER SYSTEMS"))
    for markers in (["braces"], ["highlight"]):
        spec = _spec("untracked", [{"op": "apply_header_footer_template", "select": {"sections": "all"},
                                    "params": {"template": str(tpl), "template_sections": {"cover": 0, "body": 0},
                                               "slot_markers": markers, "slot_fills": fills}}])
        with pytest.raises(JobException) as e:
            plan_spec(pkg, spec)
        assert e.value.code == "slot-in-field" and "SECTION_TITLE" in e.value.reason
    root = VPackage(tpl.read_bytes()).xml("word/header1.xml")
    pars = list(root.iter(W + "p"))
    for par, markers in ((pars[0], ["braces"]), (pars[1], ["highlight"])):
        with pytest.raises(ValueError, match="inside a field"):
            slots_of(par, markers, ["yellow"], None, _v_field_runs(root))
        assert "SECTION_TITLE" in [s["name"] for s in slots_of(par, markers, ["yellow"], None, None)]
