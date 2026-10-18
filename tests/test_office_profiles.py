"""Portable profiles remain data; reviewed targets stay tied to original bytes."""
from __future__ import annotations

import copy
import json
import os
import zipfile

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE

from spec_surgeon.office_profiles import (ProfileError, create_profile, execute_review, export_profile,
    import_profile, load_profile, load_review, prepare_review, preview_review, review_text, save_review,
    validate_review)


def setup_profile(tmp_path, *, text=None, mapping=None, policies=(), direct_formatting="keep", tracked=False):
    source, template = tmp_path / "source.docx", tmp_path / "template.docx"
    doc = Document()
    for line in text or ["PART 1 - GENERAL", "1.1 SUMMARY", "A. Scope"]:
        doc.add_paragraph(line)
    doc.save(source)
    tpl = Document()
    for name in ("Office Part", "Office Article", "Office Paragraph"):
        style = tpl.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = "Arial"
    tpl.save(template)
    profile = create_profile(tmp_path / "profile", template, "Office standard", mapping or {
        "part": "Office Part", "article": "Office Article", "paragraph": "Office Paragraph"},
        policies=policies, direct_formatting=direct_formatting, tracked=tracked)
    return source, template, profile


def test_portable_round_trip_pins_original_assets_and_never_changes_references(tmp_path):
    source, template, profile = setup_profile(tmp_path)
    originals = [source.read_bytes(), template.read_bytes()]
    exported = export_profile(profile, tmp_path / "portable.zip")
    imported = import_profile(exported, tmp_path / "moved" / "profile")
    assert imported.digest == profile.digest
    assert imported.template.read_bytes() == template.read_bytes()
    assert not str(tmp_path) in json.dumps(imported.manifest)
    template.unlink()
    review = prepare_review(imported, [source])
    assert review["status"] == "PROPOSED"
    assert source.read_bytes() == originals[0]
    assert imported.template.read_bytes() == originals[1]


def test_bundle_detects_tampering_undeclared_files_and_symlinks(tmp_path):
    _, _, profile = setup_profile(tmp_path)
    original = profile.template.read_bytes()
    profile.template.write_bytes(original + b"tampered")
    with pytest.raises(ProfileError, match="changed"):
        load_profile(profile.root)
    profile.template.write_bytes(original)
    (profile.root / "payload.py").write_text("raise RuntimeError('never execute')")
    with pytest.raises(ProfileError, match="undeclared"):
        load_profile(profile.root)
    (profile.root / "payload.py").unlink()
    asset_path = profile.template
    asset_path.unlink()
    asset_path.symlink_to(tmp_path / "template.docx")
    with pytest.raises(ProfileError, match="symlink"):
        load_profile(profile.root)


@pytest.mark.parametrize("name", ["../escaped.docx", "/absolute.docx", "assets/../../escaped.docx",
                                  "assets\\escaped.docx", "assets/C:escaped.docx"])
def test_import_rejects_path_traversal_before_publication(tmp_path, name):
    archive = tmp_path / "unsafe.zip"
    with zipfile.ZipFile(archive, "w") as z:
        z.writestr(name, b"untrusted")
    destination = tmp_path / "imported"
    with pytest.raises(ProfileError):
        import_profile(archive, destination)
    assert not destination.exists()
    assert not list(tmp_path.glob(".office-import-*"))


def test_duplicate_zip_members_and_duplicate_manifest_keys_refused(tmp_path):
    archive = tmp_path / "unsafe.zip"
    with pytest.warns(UserWarning):
        with zipfile.ZipFile(archive, "w") as z:
            z.writestr("manifest.json", "{}")
            z.writestr("manifest.json", "{}")
    with pytest.raises(ProfileError, match="Duplicate"):
        import_profile(archive, tmp_path / "out")
    _, _, profile = setup_profile(tmp_path)
    (profile.root / "manifest.json").write_text('{"name":"a","name":"b"}')
    with pytest.raises(ProfileError, match="Duplicate JSON"):
        load_profile(profile.root)


def test_failed_creation_is_atomic_and_never_replaces_existing_destination(tmp_path):
    source, template, profile = setup_profile(tmp_path)
    with pytest.raises(ProfileError, match="already exists"):
        create_profile(profile.root, template, "replacement", {"part": "Heading 1"})
    with pytest.raises(ProfileError, match="roles"):
        create_profile(tmp_path / "bad-profile", template, "invalid", {"bogus": "Heading 1"})
    assert not (tmp_path / "bad-profile").exists()
    assert not list(tmp_path.glob(".office-profile-*"))
    assert source.exists() and load_profile(profile.root).name == "Office standard"


def test_unresolved_requires_explicit_per_original_choice_and_prints_every_disposition(tmp_path):
    source, _, profile = setup_profile(tmp_path, text=["PART 1 - GENERAL", "This prose has no structural marker."])
    review = prepare_review(profile, [source])
    assert review["status"] == "NEEDS_INPUT"
    assert [p["paragraph"] for p in review["documents"][0]["unresolved"]] == [1]
    with pytest.raises(ProfileError, match="needs input"):
        validate_review(review)
    chosen = prepare_review(profile, [source], choices={str(source): {1: "exclude"}})
    assert chosen["status"] == "PROPOSED"
    assert "explicitly excluded" in review_text(chosen)
    assert "This prose has no structural marker." in review_text(chosen)
    classified = prepare_review(profile, [source], choices={str(source): {1: "article"}})
    assert classified["documents"][0]["targets"][1]["decision"] == "manual"
    assert classified["digest"] != chosen["digest"]


def test_excluded_table_paragraph_cannot_be_promoted_by_manual_choice(tmp_path):
    source, _, profile = setup_profile(tmp_path)
    doc = Document(source)
    doc.add_table(1, 1).cell(0, 0).text = "1.2 TABLE HEADING"
    doc.save(source)
    review = prepare_review(profile, [source])
    excluded = next(p for p in review["documents"][0]["analysis"]["paragraphs"] if "TABLE HEADING" in p["text"])
    assert excluded["disposition"] == "excluded"
    with pytest.raises(ProfileError, match="unsupported region"):
        prepare_review(profile, [source], choices={str(source): {excluded["paragraph"]: "article"}})


def test_multiple_exact_targets_preview_and_execute_with_independent_final_proof(tmp_path):
    source, template, profile = setup_profile(tmp_path)
    original, reference = source.read_bytes(), template.read_bytes()
    review = prepare_review(profile, [source])
    targets = review["documents"][0]["targets"]
    assert [t["paragraph"] for t in targets] == [0, 1, 2]
    assert [e["select"]["paragraphs"]["index_range"] for e in review["documents"][0]["contract"]["effects"]] == [[0, 1], [1, 2], [2, 3]]
    previews = preview_review(review, tmp_path / "preview", renderer="none", do_render=False)
    sample = previews[str(source)].samples[0]
    assert sample.passed, sample.outcomes
    assert sample.checks["OFFICE_OUTCOMES"] == "PASS"
    results = execute_review(review, tmp_path / "run", approval=review["digest"], previews=previews,
        preview_approvals={key: value.token for key, value in previews.items()}, workers=1, renderer="none", do_render=False)
    result = results[0].files[0]
    assert result.status == "PASS", result.reason
    assert result.outcomes[-1]["kind"] == "approved_goal"
    assert result.outcomes[-1]["status"] == "PASS"
    output = Document(result.delivered_to)
    assert [p.style.name for p in output.paragraphs] == ["Office Part", "Office Article", "Office Paragraph"]
    assert source.read_bytes() == original and template.read_bytes() == reference
    with pytest.raises(ProfileError, match="consumed"):
        execute_review(review, tmp_path / "second-run", approval=review["digest"], previews=previews,
            preview_approvals={key: value.token for key, value in previews.items()}, workers=1, renderer="none", do_render=False)


def test_changed_source_profile_scope_and_review_refuse_approval(tmp_path):
    source, _, profile = setup_profile(tmp_path)
    review = prepare_review(profile, [source])
    validate_review(review, approval=review["digest"])
    with pytest.raises(ProfileError, match="exact"):
        validate_review(review, approval="different")
    with pytest.raises(ProfileError, match="scope"):
        validate_review(review, files=[])
    changed = copy.deepcopy(review)
    changed["documents"][0]["targets"][0]["style"] = "Normal"
    with pytest.raises(ProfileError, match="review changed"):
        validate_review(changed)
    before = source.stat()
    os.utime(source, ns=(before.st_atime_ns, before.st_mtime_ns + 1_000_000_000))
    assert source.stat().st_mtime_ns != before.st_mtime_ns
    with pytest.raises(ProfileError, match="original document changed"):
        validate_review(review)
    os.utime(source, ns=(before.st_atime_ns, before.st_mtime_ns))
    manifest = profile.manifest
    manifest["name"] = "Changed office"
    (profile.root / "manifest.json").write_text(json.dumps(manifest))
    with pytest.raises(ProfileError, match="profile changed"):
        validate_review(review)


def test_preview_tokens_and_visual_findings_are_separate_from_review_approval(tmp_path):
    source, _, profile = setup_profile(tmp_path)
    review = prepare_review(profile, [source])
    previews = preview_review(review, tmp_path / "preview", renderer="none", do_render=False)
    with pytest.raises(ProfileError, match="Explicit approval"):
        execute_review(review, tmp_path / "run", approval="", previews=previews, preview_approvals={})
    with pytest.raises(ProfileError, match="explicitly approved preview"):
        execute_review(review, tmp_path / "run", approval=review["digest"], previews=previews, preview_approvals={})
    class Blocking:
        def blocking(self):
            return ["Look at page 1"]
    previews[str(source)].review = Blocking()
    with pytest.raises(ProfileError, match="visual findings"):
        execute_review(review, tmp_path / "run", approval=review["digest"], previews=previews,
            preview_approvals={key: value.token for key, value in previews.items()}, renderer="none", do_render=False)
    assert not (tmp_path / "run").exists()


def test_output_input_overlap_and_saved_review_do_not_grant_authority(tmp_path):
    source, _, profile = setup_profile(tmp_path)
    review = prepare_review(profile, [source])
    save_review(review, tmp_path / "review.json")
    assert load_review(tmp_path / "review.json") == review
    with pytest.raises(ProfileError, match="original document"):
        save_review(review, source)
    with pytest.raises(ProfileError, match="outside the profile"):
        save_review(review, profile.root / "review.json")
    with pytest.raises(ValueError):
        preview_review(review, tmp_path, renderer="none", do_render=False)
    with pytest.raises(ProfileError, match="outside the office"):
        preview_review(review, profile.root / "preview", renderer="none", do_render=False)
    with pytest.raises(ProfileError, match="reference assets"):
        prepare_review(profile, [profile.template])


def test_styled_roles_and_page_setup_compose_but_incompatible_policy_is_visible(tmp_path):
    source, _, profile = setup_profile(tmp_path, policies=[{"type": "page_setup", "margins": {"left": 1440}}])
    review = prepare_review(profile, [source])
    assert review["status"] == "PROPOSED"
    previews = preview_review(review, tmp_path / "preview", renderer="none", do_render=False)
    assert previews[str(source)].samples[0].passed
    manifest = profile.manifest
    manifest["policies"] = [{"type": "proofing_language", "params": {"language": "en-US"}}]
    (profile.root / "manifest.json").write_text(json.dumps(manifest))
    with pytest.raises(ProfileError, match="separate reviewed applications"):
        prepare_review(profile.root, [source])


def test_style_owned_policy_for_multiple_roles_preserves_unowned_inline_emphasis(tmp_path):
    source, _, profile = setup_profile(tmp_path, direct_formatting="replace_style_owned")
    doc = Document(source)
    doc.paragraphs[0].runs[0].font.name = "Courier New"
    doc.paragraphs[0].runs[0].italic = True
    doc.save(source)
    review = prepare_review(profile, [source])
    assert review["status"] == "PROPOSED", review["documents"][0]["reason"]
    previews = preview_review(review, tmp_path / "preview", renderer="none", do_render=False)
    sample = previews[str(source)].samples[0]
    assert sample.passed, sample.outcomes
    changed = Document(sample.output)
    assert changed.paragraphs[0].runs[0].font.name is None
    assert changed.paragraphs[0].runs[0].italic is True


def test_failed_preview_independent_outcome_cannot_be_executed(tmp_path):
    from spec_surgeon import office_profiles
    source, _, profile = setup_profile(tmp_path)
    review = prepare_review(profile, [source])
    # A deliberately wrong but valid recipe must never redefine the goal.
    changed = copy.deepcopy(review)
    changed["documents"][0]["chain"]["stages"][0]["spec"]["ops"][0]["params"]["style"] = "Heading 1"
    office_profiles._seal(changed)
    previews = preview_review(changed, tmp_path / "preview", renderer="none", do_render=False)
    sample = previews[str(source)].samples[0]
    assert not sample.passed and sample.checks["OFFICE_OUTCOMES"] == "FAIL"
    with pytest.raises(ProfileError, match="preview failed"):
        execute_review(changed, tmp_path / "run", approval=changed["digest"], previews=previews,
            preview_approvals={key: value.token for key, value in previews.items()}, renderer="none", do_render=False)


def test_policy_only_profile_does_not_require_unrelated_semantic_roles(tmp_path):
    source, template, _ = setup_profile(tmp_path, text=["Unnumbered prose need not be classified to set proofing language."])
    profile = create_profile(tmp_path / "language", template, "Office language", {},
        policies=[{"type": "proofing_language", "params": {"language": "en-US"}}])
    review = prepare_review(profile, [source])
    assert review["status"] == "PROPOSED", review["documents"][0]["reason"]
    assert review["documents"][0]["analysis"]["paragraphs"][0]["disposition"] == "unresolved"
    assert review["documents"][0]["unresolved"] == []
    previews = preview_review(review, tmp_path / "language-preview", renderer="none", do_render=False)
    assert previews[str(source)].samples[0].passed


@pytest.mark.parametrize("tracked", [False, True])
def test_consecutive_targets_share_stages_without_crossing_excluded_or_unmapped_gaps(tmp_path, tracked):
    text = ["PART 1 - GENERAL", "1.1 SUMMARY", "A. First clause", "B. Second clause",
            "C. Explicitly excluded clause", "D. Fourth clause", "E. Fifth clause", "F. Sixth clause",
            "1.2 UNMAPPED ARTICLE", "A. Last clause"]
    source, template, profile = setup_profile(tmp_path, text=text, mapping={"paragraph": "Office Paragraph"},
                                              direct_formatting="replace_style_owned", tracked=tracked)
    doc = Document(source)
    for paragraph in doc.paragraphs:
        paragraph.runs[0].font.name = "Courier New"
        paragraph.runs[0].italic = True
    doc.save(source)
    original, reference = source.read_bytes(), template.read_bytes()
    review = prepare_review(profile, [source], choices={str(source): {4: "exclude"}})
    row = review["documents"][0]
    assert review["status"] == "PROPOSED", row["reason"]
    assert [target["paragraph"] for target in row["targets"]] == [2, 3, 5, 6, 7, 9]
    assert [effect["select"]["paragraphs"]["index_range"] for effect in row["contract"]["effects"]] == [[2, 4], [5, 8], [9, 10]]
    assert [effect["expected_count"] for effect in row["contract"]["effects"]] == [2, 3, 1]
    assert len(row["chain"]["stages"]) == 3
    assert "explicitly excluded" in review_text(review)
    previews = preview_review(review, tmp_path / "preview", renderer="none", do_render=False)
    sample = previews[str(source)].samples[0]
    assert sample.passed and sample.checks["OFFICE_OUTCOMES"] == "PASS", sample.outcomes
    results = execute_review(review, tmp_path / "run", approval=review["digest"], previews=previews,
        preview_approvals={key: value.token for key, value in previews.items()}, workers=1,
        renderer="none", do_render=False)
    result = results[0].files[0]
    assert result.status == "PASS", result.reason
    assert result.outcomes[-1]["status"] == "PASS"
    output = Document(result.delivered_to)
    assert [paragraph.text for paragraph in output.paragraphs] == text
    for index, paragraph in enumerate(output.paragraphs):
        selected = index in {2, 3, 5, 6, 7, 9}
        assert paragraph.style.name == ("Office Paragraph" if selected else "Normal")
        assert paragraph.runs[0].font.name == (None if selected else "Courier New")
        assert paragraph.runs[0].italic is True
    assert source.read_bytes() == original and template.read_bytes() == reference
