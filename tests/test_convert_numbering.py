"""convert_numbering and V10's ``converted_numbering`` kind (built-in batch 5, item 33 ★ — docs/BUILTIN_OPS_BATCH5_DESIGN.md
§1): end to end on the architect's outline (build_numbering) converted to the example file's scheme (PRT / ART / PR1 / PR2
/ PR3, imported with its list), plain and tracked, kept and cleared formatting, from the main list and from named
styles, level_map; a held scheme; the refusals (an ilvl-only own numPr, a shared list, a level beyond the scheme, a
template that is no scheme); the flags (typed numbers, a REF field, written references, shifts of other paragraphs);
wrong OUTPUTS and wrong IMPLEMENTATIONS shared by plan and apply, each caught by V10; plan-time conflicts (rule 9); the
template as an input (rule 12); the worked example."""
from __future__ import annotations

import copy
import io
import zipfile
from pathlib import Path
from typing import Any, Dict, List, Tuple

import pytest
from docx import Document
from lxml import etree

from make_fixtures import SCHEME_LEVELS, build_numbering, build_scheme_example, patch_docx
from spec_surgeon.catalog import COMMON_ASSUMPTIONS, op_catalog
from spec_surgeon.editspec import apply_spec, load_spec, plan_spec, spec_input_files, validate_spec
from spec_surgeon.ops import JobException, get_op, merge_fragments
from spec_surgeon.ops import convert_numbering as cn
from spec_surgeon.ops import lists as tl
from spec_surgeon.ops.selectors import StyleSheet
from spec_surgeon.package import DocxPackage, W, serialize_xml
from spec_surgeon.report import FAILURE_EXPLANATIONS
from spec_surgeon.verify import run_suite

ROOT = Path(__file__).resolve().parent.parent
MAIN = "word/document.xml"
NUMBERING = "word/numbering.xml"
AUTHOR = "Abraham Borg"
W_DECL = f'xmlns:w="{W[1:-1]}"'


def key(p) -> str:
    """w:t (and a run's w:tab as \\t) whose nearest paragraph is ``p`` — the tests' name for a paragraph."""
    out = []
    for el in p.iter(W + "t", W + "tab"):
        q = el.getparent()
        while q is not None and q.tag != W + "p":
            q = q.getparent()
        if q is p and (el.tag == W + "t" or el.getparent().tag == W + "r"):
            out.append(el.text or "" if el.tag == W + "t" else "\t")
    return "".join(out)


def find(root, text: str):
    return next(p for p in root.iter(W + "p") if key(p) == text)


def _x(xml: str):
    """One element written as XML (the w: prefix undeclared)."""
    return etree.fromstring(f"<w:x {W_DECL}>{xml}</w:x>")[0]


def _no_ilvl_only(root):
    p = find(root, "Product data.")
    numpr = p.find(f"{W}pPr/{W}numPr")
    numpr.getparent().remove(numpr)


def _tyco(root):
    """'Tyco.' gets its own numPr (the same list, explicitly), an indent, a centring and a bold run."""
    p = find(root, "Tyco.")
    ppr = p.find(W + "pPr")
    ppr.insert(1, _x('<w:numPr><w:ilvl w:val="2"/><w:numId w:val="10"/></w:numPr>'))
    ppr.append(_x('<w:ind w:left="720"/>'))
    ppr.append(_x('<w:jc w:val="center"/>'))
    r = p.find(W + "r")
    rpr = r.find(W + "rPr")
    if rpr is None:
        rpr = etree.SubElement(r, W + "rPr")
        r.insert(0, rpr)
    rpr.append(_x("<w:b/>"))


_CACHE: Dict[str, bytes] = {}


def fixture(name: str) -> bytes:
    if name not in _CACHE:
        base = build_numbering(shared_restart=False)
        _CACHE[name] = {
            "raw": lambda: base,
            "clean": lambda: patch_docx(base, {MAIN: _no_ilvl_only}),
            "variant": lambda: patch_docx(base, {MAIN: lambda r: (_no_ilvl_only(r), _tyco(r))}),
            "shared": lambda: patch_docx(build_numbering(shared_restart=True), {MAIN: _no_ilvl_only}),
        }[name]()
    return _CACHE[name]


@pytest.fixture()
def tpl(tmp_path) -> str:
    p = tmp_path / "NUMBERING SCHEME.docx"
    p.write_bytes(build_scheme_example())
    return str(p)


def run_ops(data: bytes, specs: List[Tuple[str, dict, dict]]):
    pkg = DocxPackage(data)
    frags = [get_op(o).plan(pkg, sel, par) for o, sel, par in specs]
    plan = merge_fragments(frags)
    work = pkg.copy()
    for f in frags:
        get_op(f.op).apply(work, f)
    out = work.output_bytes()
    return plan, out, run_suite(data, out, plan.to_dict(), "src.docx", "out.docx")


def converting(data: bytes, tpl: str = None, **params):
    if tpl is not None:
        params["template"] = tpl
    params.setdefault("levels", list(SCHEME_LEVELS))
    return run_ops(data, [("convert_numbering", {}, {"author": AUTHOR, **params})])


def failed(rec) -> List[str]:
    return [c.id for c in rec.checks if c.status == "FAIL"]


def v10(rec) -> str:
    c = next(c for c in rec.checks if c.id == "V10")
    return c.summary + " | " + " | ".join(c.data.get("errors", []))


def numbers(data: bytes) -> Dict[str, Any]:
    """paragraph name → its rendered number (or (issue,) when not evaluated), by the transform's model."""
    pkg = DocxPackage(data)
    root = pkg.xml(MAIN)
    pars = list(root.iter(W + "p"))
    out = {}
    for x in tl.count(root, StyleSheet(pkg), tl.ListDefs.load(pkg)):
        k = key(pars[x.index])
        while k in out:
            k += " (again)"
        out[k] = x.text if x.evaluated else (x.issue,)
    return out


def pstyle(p):
    ps = p.find(f"{W}pPr/{W}pStyle")
    return ps.get(W + "val") if ps is not None else None


CONVERTED = {
    "GENERAL": "PART 1 -", "SUMMARY": "1.1", "Section includes wet-pipe sprinkler systems.": "A.",
    "Related requirements:": "B.", 'Division 21 Section "Common Work Results for Fire Suppression".': "1.",
    "Hangers.": "a.", "Supports.": "b.", 'Division 28 Section "Fire Detection and Alarm".': "2.", "": "C.",
    "Deferred submittal.": "D.", "REFERENCES": "1.2", "NFPA 13.": "A.", "Hidden requirement.": "B.",
    "UL listings.": "C.", "SUBMITTALS": "1.3", "Product data.": "A.", "Shop drawings.": "B.",
    "Calculations, as specified in Article 1.02.": "C.", "Comply with Article 1.02.": "D.", "PRODUCTS": "PART 2 -",
    "MANUFACTURERS": "2.1", "Tyco.": "A.", "Viking.": "B.", "MATERIALS": "2.2", "Steel pipe.": "A.", "PIPING": "2.3",
    "Black steel.": "A.", "VALVES": "2.4", "Gate valves.": "A.", "B.\tCheck valves.": "B.", "EXECUTION": "PART 3 -",
    "INSTALLATION": "3.1", "Install per NFPA 13.": "A.", "Cell requirement.": "B.", "After the table.": "C.",
}
OTHER_LISTS = {"Bullet one": "•", "Bullet two": "•", "Appendix one": "APPENDIX III", "Zed": "Z.",
               "Legal one": "3.26.1", "Double A": "AA.", "Override sub": ("override",),
               "Zed again": ("skipped",)}        # level 1 skipped by "Legal three" (lists.skipped_below)
STYLE_OF = {"GENERAL": "PRT", "SUMMARY": "ART", "Section includes wet-pipe sprinkler systems.": "PR1",
            'Division 21 Section "Common Work Results for Fire Suppression".': "PR2", "Hangers.": "PR3",
            "UL listings.": "PR1", "Cell requirement.": "PR1"}


# ---- end to end -------------------------------------------------------------------------------------------------------

def test_the_architects_outline_becomes_the_example_files_scheme(tpl):
    """The starred case: every paragraph of the main list takes its level's style from the example file (imported with
    its list, new ids above the file's own) and the scheme's list numbers it — "1.01" becomes "1.1", the letters and
    the PARTs stay — while the other lists keep their numbers; plain by default."""
    src = fixture("clean")
    plan, out, rec = converting(src, tpl)
    assert rec.passed, rec.to_markdown()
    got = numbers(out)
    assert {k: got[k] for k in CONVERTED} == CONVERTED
    assert {k: got[k] for k in OTHER_LISTS} == OTHER_LISTS
    root = DocxPackage(out).xml(MAIN)
    for name, sid in STYLE_OF.items():
        assert pstyle(find(root, name)) == sid, name
    a = next(x for x in plan.property_assertions if x["kind"] == "converted_numbering")
    imp = next(x for x in plan.property_assertions if x["kind"] == "styles_imported")
    assert imp["imported"] == SCHEME_LEVELS and a["dest_num"] == int(imp["num_map"]["10"])
    assert a["dest_num"] > 23 and len(a["converted"]) == 35 and a["expected_restyled"] == 35
    assert sorted(plan.intended_parts) == sorted([MAIN, "word/styles.xml", NUMBERING])
    assert not plan.rev_delta.get("pPrChange") and plan.geometry_reference == "none"
    assert find(root, "Unnumbered note under a numbered style.").find(f"{W}pPr/{W}pStyle").get(W + "val") == "ArchPara"
    # the flags: the double number, the stale REF field, the written reference to 1.02
    flags = " | ".join(plan.flags)
    assert "‘B. Check valves.’" in flags and "REF _Ref1001" in flags and "'1.02' (now '1.2')" in flags
    assert "Calculations, as specified in Article 1" in flags
    # the hidden paragraph: Word counts it (the prediction follows Word), LibreOffice's proof does not
    assert any("hidden" in f and "LibreOffice" in f and "‘Hidden requirement.’" in f for f in plan.flags)


def test_tracked_conversion_with_cleared_formatting(tpl):
    """Tracked: one w:pPrChange per changed paragraph (the snapshot before the restyle, the numPr removal and the
    clear), one w:rPrChange per cleared run; 'Tyco.' loses its own numPr, its indent, its centring and its bold."""
    src = fixture("variant")
    plan, out, rec = converting(src, tpl, tracked=True, direct_formatting="clear")
    assert rec.passed, rec.to_markdown()
    a = next(x for x in plan.property_assertions if x["kind"] == "converted_numbering")
    assert plan.rev_delta["pPrChange"] == a["expected_paragraphs"] == 35 and plan.rev_delta["rPrChange"] == 1
    p = find(DocxPackage(out).xml(MAIN), "Tyco.")
    ppr = p.find(W + "pPr")
    assert [etree.QName(c).localname for c in ppr] == ["pStyle", "pPrChange"]
    snap = ppr.find(f"{W}pPrChange/{W}pPr")
    assert [etree.QName(c).localname for c in snap] == ["pStyle", "numPr", "spacing", "ind", "jc"] or \
        {"numPr", "ind", "jc", "pStyle"} <= {etree.QName(c).localname for c in snap}
    assert p.find(f"{W}r/{W}rPr/{W}b") is None and p.find(f"{W}r/{W}rPr/{W}rPrChange") is not None
    assert numbers(out)["Tyco."] == "A."


def test_a_plain_conversion_keeps_direct_formatting_and_removes_own_numbering(tpl):
    plan, out, rec = converting(fixture("variant"), tpl)
    assert rec.passed, rec.to_markdown()
    p = find(DocxPackage(out).xml(MAIN), "Tyco.")
    assert [etree.QName(c).localname for c in p.find(W + "pPr")] == ["pStyle", "ind", "jc"]
    assert p.find(f"{W}r/{W}rPr/{W}b") is not None
    a = next(x for x in plan.property_assertions if x["kind"] == "converted_numbering")
    assert [c for c in a["changed"] if c[3]] and (MAIN, str(a["changed"][0][0]), "pStyle") in \
        get_op("convert_numbering").plan(DocxPackage(fixture("variant")), {},
                                         {"levels": SCHEME_LEVELS, "template": tpl}).touches


def test_from_named_styles_and_a_level_map(tpl):
    """from.styles picks the paragraphs by style per level; a three-level scheme takes the deeper levels through
    level_map (sub-paragraphs become paragraphs: the counts change and are flagged)."""
    styles = [["Arch Part"], ["ArchArt"], ["ArchPara", "ArchParaAlt"], ["ArchSub"], ["ArchSub2"]]
    plan, out, rec = converting(fixture("clean"), tpl, **{"from": {"styles": styles}})
    assert rec.passed, rec.to_markdown()
    assert {k: numbers(out)[k] for k in CONVERTED} == CONVERTED
    plan, out, rec = converting(fixture("clean"), tpl, levels=SCHEME_LEVELS[:3], level_map={"3": 2, "4": 2})
    assert rec.passed, rec.to_markdown()
    got = numbers(out)
    assert got['Division 21 Section "Common Work Results for Fire Suppression".'] == "C." and got["Hangers."] == "D."
    assert any("count differently than before" in f for f in plan.flags)


def test_a_held_scheme_without_template(tpl):
    """The document's own outline as the scheme (no template, nothing imported): only 'UL listings.' (ArchParaAlt, level 2
    through its basedOn) is restyled; the numbering part is unchanged and every number stays."""
    levels = ["ArchPart", "ArchArt", "ArchPara", "ArchSub", "ArchSub2"]
    src = fixture("clean")
    plan, out, rec = converting(src, levels=levels)
    assert rec.passed, rec.to_markdown()
    a = plan.property_assertions[0]
    assert a["changed"] == [[a["changed"][0][0], 2, True, False]] and a["dest_num"] == 10
    assert len(plan.property_assertions) == 1 and NUMBERING not in plan.intended_parts
    assert numbers(out) == numbers(src)
    assert any("already as asked, left alone: 34" in n for n in plan.notes)


def test_a_paragraph_another_list_numbers_converted_into_a_held_list_shifts_the_others_flagged(tmp_path):
    """The example file itself as the document, with a paragraph numbered by another list before its first PART: from
    Normal paragraphs into the held scheme's level 0 — the conversion moves every later PART and article by one; it is
    predicted, verified and flagged by name."""
    doc = Document(io.BytesIO(build_scheme_example()))
    p = doc.paragraphs[1].insert_paragraph_before("Introduction")
    numpr = etree.fromstring(f'<w:numPr {W_DECL}><w:ilvl w:val="0"/><w:numId w:val="1"/></w:numPr>')
    p._p.get_or_add_pPr().append(numpr)
    buf = io.BytesIO()
    doc.save(buf)
    src = buf.getvalue()
    plan, out, rec = converting(src, **{"from": {"styles": [["Normal"]]}})
    assert rec.passed, rec.to_markdown()
    got = numbers(out)
    assert got["Introduction"] == "PART 1 -" and got["GENERAL"] == "PART 2 -" and got["SUMMARY"] == "2.1"
    assert any("not converted change their number" in f and "‘GENERAL’ 'PART 1 -' → 'PART 2 -'" in f
               for f in plan.flags)


# ---- refusals ---------------------------------------------------------------------------------------------------------

def test_an_ilvl_only_own_numbering_refuses(tpl):
    with pytest.raises(JobException) as e:
        get_op("convert_numbering").plan(DocxPackage(fixture("raw")), {}, {"levels": SCHEME_LEVELS, "template": tpl})
    assert e.value.code == "numbering-unsupported" and "Product data." in e.value.reason


def test_a_shared_list_refuses_and_a_held_shared_scheme_too(tpl):
    """build_numbering(True): nums 10 and 11 share abstractNum 10 — the source numbers cannot be read (Word and
    LibreOffice differ), and the document's own list as the destination is not clean either."""
    with pytest.raises(JobException) as e:
        get_op("convert_numbering").plan(DocxPackage(fixture("shared")), {}, {"levels": SCHEME_LEVELS, "template": tpl})
    assert e.value.code == "numbering-unsupported" and "shared list definition" in e.value.reason
    with pytest.raises(JobException) as e:
        get_op("convert_numbering").plan(DocxPackage(fixture("shared")), {},
                                         {"levels": ["ArchPart", "ArchArt", "ArchPara", "ArchSub", "ArchSub2"]})
    assert e.value.code == "numbering-unsupported" and "destination list" in e.value.reason


def test_a_level_beyond_the_scheme_refuses_and_the_scheme_must_fit(tpl):
    with pytest.raises(JobException) as e:
        get_op("convert_numbering").plan(DocxPackage(fixture("clean")), {}, {"levels": SCHEME_LEVELS[:3], "template": tpl})
    assert e.value.code == "scheme-depth"
    with pytest.raises(JobException) as e:                       # no template: the file has no PRT
        get_op("convert_numbering").plan(DocxPackage(fixture("clean")), {}, {"levels": SCHEME_LEVELS})
    assert e.value.code == "scheme"


@pytest.mark.parametrize("params, code", [
    ({}, "bad-params"),
    ({"levels": []}, "bad-params"),
    ({"levels": ["PRT", "PRT"]}, "bad-params"),
    ({"levels": SCHEME_LEVELS, "existing": "replace"}, "bad-params"),           # existing without a template
    ({"levels": SCHEME_LEVELS, "from": "all"}, "bad-params"),
    ({"levels": SCHEME_LEVELS, "from": {"styles": [[]]}}, "bad-params"),
    ({"levels": SCHEME_LEVELS, "level_map": {"0": 7}}, "bad-params"),
    ({"levels": SCHEME_LEVELS, "level_map": {"x": 1}}, "bad-params"),
    ({"levels": SCHEME_LEVELS, "direct_formatting": "wipe"}, "bad-params"),
    ({"levels": SCHEME_LEVELS, "tracked": "yes"}, "bad-params"),
    ({"levels": SCHEME_LEVELS, "colour": "red"}, "bad-params"),
])
def test_parameters(params, code):
    with pytest.raises(JobException) as e:
        get_op("convert_numbering").plan(DocxPackage(fixture("clean")), {}, params)
    assert e.value.code == code
    with pytest.raises(JobException) as e:
        get_op("convert_numbering").plan(DocxPackage(fixture("clean")), {"paragraphs": "all"},
                                         {"levels": SCHEME_LEVELS})
    assert e.value.code == "bad-params"


def test_template_refusals_are_bad_params(tpl):
    """The example file is the same for every document: a level it lacks, the wrong order, a bullet level — fatal."""
    for levels in (["PRT", "ART", "NoSuch"], ["ART", "PRT"], ["PR1"]):
        with pytest.raises(JobException) as e:
            get_op("convert_numbering").plan(DocxPackage(fixture("clean")), {}, {"levels": levels, "template": tpl})
        assert e.value.code == "bad-params", levels


def test_nothing_numbered_is_a_noted_no_op_unless_allow_missing_false(tpl):
    src = build_scheme_example()
    doc = Document()
    doc.add_paragraph("Nothing numbered here.")
    buf = io.BytesIO()
    doc.save(buf)
    plan, out, rec = converting(buf.getvalue(), tpl)
    assert rec.passed and any("nothing numbered to convert" in n for n in plan.notes)
    with pytest.raises(JobException) as e:
        get_op("convert_numbering").plan(DocxPackage(buf.getvalue()), {},
                                         {"levels": SCHEME_LEVELS, "template": tpl, "allow_missing": False})
    assert e.value.code == "not-found"
    assert src


# ---- wrong OUTPUTS are caught -----------------------------------------------------------------------------------------

def mutate(docx: bytes, part: str, fn) -> bytes:
    zin = zipfile.ZipFile(io.BytesIO(docx))
    root = etree.fromstring(zin.read(part))
    fn(root)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        for info in zin.infolist():
            z.writestr(info, serialize_xml(root) if info.filename == part else zin.read(info.filename))
    return buf.getvalue()


def _broken(tpl, fn, data=None, **params) -> str:
    data = data or fixture("variant")
    plan, out, rec = converting(data, tpl, **params)
    assert rec.passed, rec.to_markdown()
    bad = mutate(out, MAIN, fn)
    rec2 = run_suite(data, bad, plan.to_dict(), "src.docx", "out.docx")
    assert "V10" in failed(rec2)
    return v10(rec2)


def test_a_wrong_levels_style_fails(tpl):
    def fn(root):
        find(root, "SUMMARY").find(f"{W}pPr/{W}pStyle").set(W + "val", "PR1")
    msg = _broken(tpl, fn)
    assert "converted_numbering" in msg and ("paragraph properties" in msg or "numbers" in msg)


def test_an_own_numbering_left_in_place_fails(tpl):
    def fn(root):
        ppr = find(root, "Tyco.").find(W + "pPr")
        ppr.insert(1, _x('<w:numPr><w:ilvl w:val="2"/><w:numId w:val="10"/></w:numPr>'))
    msg = _broken(tpl, fn)
    assert "output's numbers are not the predicted" in msg or "paragraph properties" in msg


def test_a_tracked_change_with_a_wrong_snapshot_fails(tpl):
    def fn(root):
        snap = find(root, "SUMMARY").find(f"{W}pPr/{W}pPrChange/{W}pPr")
        etree.SubElement(snap, W + "keepLines")
    assert "snapshot" in _broken(tpl, fn, tracked=True)


def test_a_cleared_property_left_in_place_fails(tpl):
    def fn(root):
        find(root, "Tyco.").find(W + "pPr").append(_x('<w:jc w:val="center"/>'))
    assert "converted_numbering" in _broken(tpl, fn, direct_formatting="clear")


# ---- wrong IMPLEMENTATIONS shared by plan and apply are caught --------------------------------------------------------

def test_an_off_by_one_level_map_is_caught(tpl, monkeypatch):
    """A level map that sends the paragraphs (level 2) one level down, the same in plan and apply: the output is
    self-consistent (the plan's own prediction matches it), V1–V9 pass; the verifier's own levels do not."""
    monkeypatch.setattr(cn, "clean_level_map", lambda lm, n: {2: 3})
    plan, out, rec = converting(fixture("clean"), tpl)
    assert failed(rec) == ["V10"] and "converted_numbering" in v10(rec)


def test_a_prediction_that_forgets_restarts_is_caught(tpl, monkeypatch):
    """A transform-side count that never restarts a lower level (the plan's predicted numbers carry on "D., E." under
    the next article): the output is right, the plan's prediction is not — V10 compares it with its own."""
    monkeypatch.setattr(tl, "restarts", lambda restart, L, k: False)
    plan, out, rec = converting(fixture("clean"), tpl)
    assert failed(rec) == ["V10"] and "predicted numbers are not the verifier's" in v10(rec)


def test_a_clear_that_keeps_the_indent_is_caught(tpl, monkeypatch):
    """A shared clear set without w:ind (plan and apply agree; the indent stays) — the verifier's own default set
    clears it."""
    real = cn.ClearSettings
    monkeypatch.setattr(cn, "ClearSettings", lambda: real(para_names=real().para_names - {"ind"}))
    plan, out, rec = converting(fixture("variant"), tpl, direct_formatting="clear")
    assert failed(rec) == ["V10"] and "converted_numbering" in v10(rec)


def test_a_dropped_import_assertion_is_caught(tpl):
    plan, out, rec = converting(fixture("clean"), tpl)
    d = plan.to_dict()
    d["property_assertions"] = [a for a in d["property_assertions"] if a["kind"] != "styles_imported"]
    rec2 = run_suite(fixture("clean"), out, d, "src.docx", "out.docx")
    assert not rec2.passed and "converted_numbering" in v10(rec2)


# ---- plan-time conflicts (rule 9) -------------------------------------------------------------------------------------

def _spec(*ops, mode="formatting"):
    return {"request": "r", "summary": "s", "mode": mode, "ops": list(ops)}


def test_conflicts_and_what_plans_beside(tpl):
    conv = {"op": "convert_numbering", "params": {"levels": SCHEME_LEVELS, "template": tpl}}
    src = fixture("clean")
    for other, mode in (
            ({"op": "replace_text", "params": {"find": "Viking", "replace": "Victaulic"}}, "content"),
            ({"op": "apply_style", "select": {"paragraphs": {"text_regex": "^Tyco"}}, "params": {"style": "Normal"}},
             "formatting"),
            ({"op": "map_styles", "params": {"map": [{"from": "ArchPara", "to": "Normal"}]}}, "formatting"),
            ({"op": "convert_typed_numbers", "params": {"levels": SCHEME_LEVELS, "template": tpl, "tracked": False}},
             "untracked")):
        for ops in ((conv, other), (other, conv)):
            with pytest.raises(JobException) as e:
                plan_spec(DocxPackage(src), _spec(*ops, mode=mode))
            assert e.value.code == "conflict", other["op"]
    # a font change of the runs plans beside it, and both verify
    font = {"op": "set_font", "params": {"font": "Arial", "only_where_defined": False}}
    for ops in ((conv, font), (font, conv)):
        plan, frags = plan_spec(DocxPackage(src), _spec(*ops))
        out = apply_spec(DocxPackage(src), frags).output_bytes()
        rec = run_suite(src, out, plan.to_dict(), "src.docx", "out.docx")
        assert rec.passed, rec.to_markdown()


def test_touches_declare_the_count_reads(tpl):
    frag = get_op("convert_numbering").plan(DocxPackage(fixture("clean")), {}, {"levels": SCHEME_LEVELS, "template": tpl})
    t = set(frag.touches)
    assert {(NUMBERING, "*", "verify"), (MAIN, "*", "verify-props:pStyle"), (MAIN, "*", "verify-props:numPr"),
            (MAIN, "scope", "verify"), (NUMBERING, "*", "add")} <= t
    assert ("word/styles.xml", "style:PRT", "add") in t and ("word/styles.xml", "style:ArchPara", "uses-") in t
    assert any(a == "verify:numPr" for _, _, a in t)


# ---- the template is an input (rule 12) -------------------------------------------------------------------------------

def test_a_template_changed_after_planning_is_refused_and_fails_v10(tpl):
    src = fixture("clean")
    pkg = DocxPackage(src)
    frag = get_op("convert_numbering").plan(pkg, {}, {"levels": SCHEME_LEVELS, "template": tpl})
    plan = merge_fragments([frag])
    work = pkg.copy()
    get_op("convert_numbering").apply(work, frag)
    out = work.output_bytes()

    def touch(root):
        extra = etree.SubElement(root, W + "style")
        extra.set(W + "type", "paragraph")
        extra.set(W + "styleId", "AddedLater")
    Path(tpl).write_bytes(mutate(build_scheme_example(), "word/styles.xml", touch))
    with pytest.raises(JobException):
        get_op("convert_numbering").apply(pkg.copy(), frag)
    rec = run_suite(src, out, plan.to_dict(), "src.docx", "out.docx")
    assert not rec.passed and "template" in v10(rec)


# ---- the worked example, the catalog --------------------------------------------------------------------------------

def test_the_worked_example_plans_and_verifies():
    spec = load_spec(ROOT / "examples" / "specs" / "convert-numbering-to-scheme.json")
    validate_spec(spec)
    tpl = (ROOT / "examples" / "templates" / "NUMBERING SCHEME.docx").resolve()
    assert Path(spec["ops"][0]["params"]["template"]) == tpl and tpl in spec_input_files(spec)
    src = fixture("clean")
    plan, frags = plan_spec(DocxPackage(src), spec)
    out = apply_spec(DocxPackage(src), frags).output_bytes()
    rec = run_suite(src, out, plan.to_dict(), "src.docx", "out.docx")
    assert rec.passed, rec.to_markdown()
    assert numbers(out)["REFERENCES"] == "1.2"


def test_catalog_and_explanations():
    cat = op_catalog()
    assert "convert_numbering" in cat["ops"] and COMMON_ASSUMPTIONS["convert_numbering"]
    for code in ("scheme", "scheme-depth", "numbering-unsupported", "existing-revisions"):
        assert code in FAILURE_EXPLANATIONS, code
