"""Unsupported reusable settings fail at publication, before a source is selected."""
from __future__ import annotations

import copy
import io
import itertools

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE

from make_fixtures import SCHEME_LEVELS, build_scheme_example
from spec_surgeon.office_profiles import (ProfileError, create_profile, import_profile, prepare_review,
                                         preview_review, save_review)


POLICIES = {
    "page_setup": {"type": "page_setup", "margins": {"left": 1440}},
    "header_footer_template": {"type": "header_footer_template", "params": {
        "template_sections": {"all": 0}, "slot_fills": {}}},
    "table_format": {"type": "table_format", "params": {"cant_split": True}},
    "table_style": {"type": "table_style", "params": {"style": "Table Grid"}},
    "proofing_language": {"type": "proofing_language", "params": {"language": "en-US"}},
    "numbering": {"type": "numbering", "params": {"levels": list(SCHEME_LEVELS)}},
}


@pytest.fixture
def template(tmp_path):
    doc = Document(io.BytesIO(build_scheme_example()))
    doc.styles.add_style("Office Part", WD_STYLE_TYPE.PARAGRAPH).font.name = "Arial"
    path = tmp_path / "office.docx"
    doc.save(path)
    return path


@pytest.mark.parametrize("policy", list(POLICIES))
@pytest.mark.parametrize("direct", ["keep", "replace_style_owned"])
def test_roles_reject_unsupported_policy_combinations_during_creation(tmp_path, template, policy, direct):
    destination = tmp_path / "profile"
    permitted = policy == "page_setup" or (direct == "keep" and policy in {"table_format", "table_style"})
    if permitted:
        create_profile(destination, template, "Office", {"part": "Office Part"},
                       policies=[POLICIES[policy]], direct_formatting=direct)
    else:
        with pytest.raises(ProfileError, match="separate"):
            create_profile(destination, template, "Office", {"part": "Office Part"},
                           policies=[POLICIES[policy]], direct_formatting=direct)
        assert not destination.exists()


@pytest.mark.parametrize("first,second", list(itertools.combinations(POLICIES, 2)))
def test_policy_pairs_have_an_explicit_supported_subset(tmp_path, template, first, second):
    pair = {first, second}
    permitted = "page_setup" in pair or pair in (
        {"table_format", "table_style"}, {"table_format", "proofing_language"})
    destination = tmp_path / "profile"
    if permitted:
        create_profile(destination, template, "Office", {}, policies=[POLICIES[first], POLICIES[second]])
    else:
        with pytest.raises(ProfileError, match="separate"):
            create_profile(destination, template, "Office", {}, policies=[POLICIES[first], POLICIES[second]])
        assert not destination.exists()


def test_combined_roles_tables_and_page_policy_preview_proves_original_to_final(tmp_path, template):
    doc = Document()
    doc.add_paragraph("PART 1 - GENERAL")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Unchanged schedule words"
    source = tmp_path / "source.docx"
    doc.save(source)
    profile = create_profile(tmp_path / "profile", template, "Office", {"part": "Office Part"},
                             policies=[POLICIES[k] for k in ("page_setup", "table_style", "table_format")])
    review = prepare_review(profile, [source])
    assert review["status"] == "PROPOSED", review
    previews = preview_review(review, tmp_path / "preview", renderer="none", do_render=False)
    sample = previews[str(source)].samples[0]
    assert sample.passed and sample.checks["OFFICE_OUTCOMES"] == "PASS", sample.outcomes
    assert any("model visual review was not run" in warning for warning in previews[str(source)].warnings)


@pytest.mark.parametrize("name,reason", [("Not In Template", "paragraph style defined"),
                                         ("Normal", "default style"), ("Strong", "paragraph style defined"),
                                         (SCHEME_LEVELS[0], "numbered style")])
def test_unusable_role_styles_are_rejected_before_publication(tmp_path, template, name, reason):
    with pytest.raises(ProfileError, match=reason):
        create_profile(tmp_path / "profile", template, "Office", {"part": name})
    assert not (tmp_path / "profile").exists()


@pytest.mark.parametrize("kind", list(POLICIES))
def test_duplicate_policy_types_are_rejected(tmp_path, template, kind):
    with pytest.raises(ProfileError, match="at most one policy"):
        create_profile(tmp_path / "profile", template, "Office", {}, policies=[POLICIES[kind], POLICIES[kind]])


def test_table_style_and_numbering_templates_are_validated_before_publication(tmp_path, template):
    for kind, params in (("table_style", {"style": "Unknown Table"}), ("numbering", {"levels": ["Office Part"]})):
        with pytest.raises(ProfileError):
            create_profile(tmp_path / kind, template, "Office", {}, policies=[{"type": kind, "params": params}])
        assert not (tmp_path / kind).exists()


def test_row_only_table_composition_and_template_distance_exceptions_remain_narrow(tmp_path, template):
    whole_table = {"type": "table_format", "params": {"width": {"pct": 100}}}
    with pytest.raises(ProfileError, match="table style plus"):
        create_profile(tmp_path / "tables", template, "Office", {}, policies=[POLICIES["table_style"], whole_table])
    page = {"type": "page_setup", "margins": {"header": 720}}
    with pytest.raises(ProfileError, match="distance alignment conflicts"):
        create_profile(tmp_path / "headers", template, "Office", {}, policies=[POLICIES["header_footer_template"], page])
    headers = copy.deepcopy(POLICIES["header_footer_template"])
    headers["params"]["align_header_footer_distance"] = False
    create_profile(tmp_path / "explicit-distances", template, "Office", {}, policies=[headers, page])


def test_review_cannot_replace_original_reference_or_unrelated_json(tmp_path, template):
    profile = create_profile(tmp_path / "profile", template, "Office", {}, policies=[POLICIES["page_setup"]])
    source = tmp_path / "source.docx"
    Document().save(source)
    review = prepare_review(profile, [source])
    before = template.read_bytes()
    with pytest.raises(ProfileError, match=".json file"):
        save_review(review, template)
    assert template.read_bytes() == before
    unrelated = tmp_path / "config.json"
    unrelated.write_text('{"important": true}')
    with pytest.raises(ProfileError, match="valid office review JSON"):
        save_review(review, unrelated)
    assert unrelated.read_text() == '{"important": true}'
    saved = tmp_path / "review.json"
    save_review(review, saved)
    save_review(review, saved)


def test_bad_archive_is_a_profile_error_and_staging_is_removed(tmp_path):
    archive = tmp_path / "invalid.zip"
    archive.write_bytes(b"not a zip")
    with pytest.raises(ProfileError, match="archive cannot be read"):
        import_profile(archive, tmp_path / "profile")
    assert not (tmp_path / "profile").exists()
    assert not list(tmp_path.glob(".office-import-*"))
