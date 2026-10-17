"""Built-in batch 5's review round (docs/BUILTIN_OPS_BATCH5_DESIGN.md, "Built"): every finding reproduced first, each
fix pinned by a test that fails without it."""
from __future__ import annotations

from lxml import etree

import pytest

from make_fixtures import build, build_rich, patch_docx
from spec_surgeon import package as package_module
from spec_surgeon.editspec import apply_spec, plan_spec
from spec_surgeon.ooxml import SECTPR_ORDER, insert_in_order
from spec_surgeon.ops import JobException
from spec_surgeon.package import DocxPackage, W
from spec_surgeon.verify import run_suite

MAIN = "word/document.xml"


def _spec(*ops, mode="formatting"):
    return {"request": "r", "summary": "r", "assumptions": [], "questions": [], "applies_to": {"or": "all"},
            "mode": mode, "ops": list(ops)}


COVER_UNNUMBERED = {"op": "fix_page_numbering", "params": {"option": "cover-unnumbered"}}


def _cover_title_off(data: bytes) -> bytes:
    """The cover section's w:titlePg written as <w:titlePg w:val="0"/> — off (CT_OnOff), though present."""
    def edit(root):
        sp = DocxPackage.sect_prs(root)[0]
        for old in sp.findall(W + "titlePg"):
            sp.remove(old)
        tp = etree.Element(W + "titlePg")
        tp.set(W + "val", "0")
        insert_in_order(sp, tp, SECTPR_ORDER)
    return patch_docx(data, {MAIN: edit})


# ---- a titlePg written off is off --------------------------------------------------------------------------------------

def test_cover_unnumbered_reads_a_title_page_switched_off_as_off():
    """build_rich's cover has a first-page footer (footer1) and a default one (footer2) its body shares. With
    <w:titlePg w:val="0"/> the cover shows footer2 — shared, so the op refuses — never footer1, which nobody sees."""
    pkg = DocxPackage(_cover_title_off(build_rich()))
    assert [s.title_pg for s in pkg.sections()][0] is False
    with pytest.raises(JobException) as e:
        plan_spec(pkg, _spec(COVER_UNNUMBERED))
    assert e.value.code == "shared-part" and "word/footer2.xml" in e.value.reason


def test_cover_unnumbered_with_title_page_off_edits_the_footer_the_cover_shows_and_verifies():
    data = _cover_title_off(build("21 13 13", "WET-PIPE SPRINKLER SYSTEMS"))
    pkg = DocxPackage(data)
    plan, frags = plan_spec(pkg, _spec(COVER_UNNUMBERED))
    assert {"kind": "cover_unnumbered", "footer": pkg.sections()[0].footer_refs["default"]} in [
        {k: a[k] for k in ("kind", "footer")} for a in plan.property_assertions if a["kind"] == "cover_unnumbered"]
    out = apply_spec(pkg, frags).output_bytes()
    assert run_suite(data, out, plan.to_dict(), "a.docx", "b.docx").passed


def test_v10_catches_a_plan_that_edits_a_footer_the_cover_does_not_show(monkeypatch):
    """Planted: the old reading (titlePg present = on) shared by plan and apply — the op edits the cover's first-page
    footer while the cover shows its default one. V10's own reading of what the cover shows fails it."""
    real = DocxPackage.sections

    def presence_only(self):
        out = real(self)
        for s, sp in zip(out, DocxPackage.sect_prs(self.xml(self.main_part()))):
            s.title_pg = sp.find(W + "titlePg") is not None
        return out

    data = _cover_title_off(build_rich())
    pkg = DocxPackage(data)
    monkeypatch.setattr(package_module.DocxPackage, "sections", presence_only)
    plan, frags = plan_spec(pkg, _spec(COVER_UNNUMBERED))
    assert plan.intended_parts == ["word/footer1.xml"]                  # the footer nobody sees
    out = apply_spec(pkg, frags).output_bytes()
    monkeypatch.setattr(package_module.DocxPackage, "sections", real)
    rec = run_suite(data, out, plan.to_dict(), "a.docx", "b.docx")
    assert not rec.passed and "cover_unnumbered" in rec.get("V10").summary


# ---- remove_header_footer reads the title pages it judges inheritance by -------------------------------------------------

def test_removing_first_page_parts_beside_a_different_first_page_is_a_plan_time_conflict():
    """It used to plan and PASS with section 1's first page showing no header or footer at all, while set_page_setup's
    note said it showed header1 — which the same spec removed."""
    remove = {"op": "remove_header_footer", "select": {"sections": [0], "types": ["first"]}, "params": {}}
    first = {"op": "set_page_setup", "select": {"sections": [1]}, "params": {"different_first_page": True}}
    for ops in ((remove, first), (first, remove)):
        with pytest.raises(JobException) as e:
            plan_spec(DocxPackage(build_rich()), _spec(*ops))
        assert e.value.code == "conflict" and "titlePg" in e.value.reason
    # alone, each still plans
    plan_spec(DocxPackage(build_rich()), _spec(remove))
    plan_spec(DocxPackage(build_rich()), _spec(first))


# ---- the landscape example on a one-section piping spec ----------------------------------------------------------------

def _one_section_spec(heading: str = "") -> bytes:
    import io
    from docx import Document
    doc = Document()
    for text in ("SECTION 21 13 13", "WET-PIPE SPRINKLER SYSTEMS", "PART 2 - PRODUCTS",
                 "A.\tSteel pipe: ASTM A 53, Schedule 40, black steel.", "PART 3 - EXECUTION") + ((heading,) if heading else ()):
        doc.add_paragraph(text)
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text, t.cell(1, 0).text, t.cell(1, 1).text = "Pipe size", "Max spacing", "1 in", "12 ft"
    doc.add_paragraph("END OF SECTION 21 13 13")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_the_landscape_example_leaves_schedule_40_pipe_alone():
    """It chose sections by '(?i)\\bschedule\\b' + a table: a one-section piping spec ('Schedule 40') went landscape
    as a whole, every check passing and nothing flagged."""
    from pathlib import Path
    from spec_surgeon.editspec import load_spec
    spec = load_spec(Path(__file__).resolve().parent.parent / "examples" / "specs" / "landscape-schedule-sections.json")
    plan, frags = plan_spec(DocxPackage(_one_section_spec()), spec)
    assert plan.intended_parts == [] and any("nothing" in n or "no section" in n for n in plan.notes)


def test_a_content_selection_that_chooses_every_section_is_flagged():
    data = _one_section_spec("SPRINKLER HEAD SCHEDULE")
    op = {"op": "set_page_setup", "select": {"sections": {"containing": {"text_regex": "SCHEDULE$"}, "with_tables": True}},
          "params": {"orientation": "landscape"}}
    plan, frags = plan_spec(DocxPackage(data), _spec(op))
    assert any("chose every section" in f and "whole document" in f for f in plan.flags)
    out = apply_spec(DocxPackage(data), frags).output_bytes()
    assert run_suite(data, out, plan.to_dict(), "a.docx", "b.docx").passed
    plan2, _ = plan_spec(DocxPackage(data), _spec({**op, "select": {"sections": [0]}}))
    assert not any("chose every section" in f for f in plan2.flags)        # an index chooses on purpose


# ---- the selector reads an empty pStyle alike on both sides --------------------------------------------------------------

def _empty_pstyle_in_section(data, section, val=""):
    from spec_surgeon.ops.set_tab_stops import paragraph_sections

    def edit(root):
        for p, s in zip(list(root.iter(W + "p")), paragraph_sections(root)):
            if s != section:
                continue
            ppr = p.find(W + "pPr")
            if ppr is None:
                ppr = etree.Element(W + "pPr")
                p.insert(0, ppr)
            if ppr.find(W + "pStyle") is None:
                ps = etree.Element(W + "pStyle")
                if val is not None:
                    ps.set(W + "val", val)
                ppr.insert(0, ps)
    return patch_docx(data, {MAIN: edit})


@pytest.mark.parametrize("val", ["", None])
def test_an_empty_pstyle_names_no_style_on_both_sides(val):
    """It names no style (batch 4's definition): the transform read it so, the verifier as the default style — a section
    chosen by style_in differed and a valid plan failed V10."""
    from make_fixtures import build_schedule
    from spec_surgeon.ops.selectors import resolve_sections
    from spec_surgeon.verify import props as vprops
    from spec_surgeon.verify.reader import STYLES_REL, VPackage
    data = _empty_pstyle_in_section(build_schedule(), 2, val)
    for sel in ({"containing": {"style_in": ["Normal"]}}, {"containing": {"has_bold": True}},
                {"containing": {"is_all_caps": True}}):
        v = VPackage(data)
        styles = vprops.Styles(v.xml(v.related(STYLES_REL)))
        assert resolve_sections(DocxPackage(data), sel, content=True) == vprops._v_resolve_sections(v, sel, styles), sel
    op = {"op": "set_page_setup", "select": {"sections": {"containing": {"style_in": ["Normal"]}}},
          "params": {"orientation": "landscape"}}
    plan, frags = plan_spec(DocxPackage(data), _spec(op))
    out = apply_spec(DocxPackage(data), frags).output_bytes()
    assert run_suite(data, out, plan.to_dict(), "a.docx", "b.docx").passed


@pytest.mark.parametrize("pred", [{"style_in": "Normal"}, {"style_in": [1]}, {"index_range": ["a", 3]},
                                  {"index_range": [0, True]}])
def test_a_malformed_paragraph_selector_is_refused(pred):
    """A bare string was read as a set of characters on one side and a substring test on the other; non-numbers
    crashed."""
    op = {"op": "set_page_setup", "select": {"sections": {"containing": pred}}, "params": {"orientation": "landscape"}}
    with pytest.raises(JobException) as e:
        plan_spec(DocxPackage(build_rich()), _spec(op))
    assert e.value.code == "bad-params"


# ---- the minor items ---------------------------------------------------------------------------------------------------

def test_with_tables_no_longer_refuses_a_text_edit_beside_it():
    """No text edit moves a table: with_tables used to declare a text read of the whole main part."""
    from make_fixtures import build_schedule
    from spec_surgeon.ops.selectors import section_reads
    assert section_reads(DocxPackage(build_schedule()), {"with_tables": True}) == []
    assert ("word/document.xml", "*", "verify:text") in section_reads(
        DocxPackage(build_schedule()), {"containing": {"text_regex": "X"}})


def test_a_set_page_setup_assertion_without_its_select_fails_v10():
    data = build("21 13 13", "WET-PIPE SPRINKLER SYSTEMS")
    op = {"op": "set_page_setup", "select": {"sections": [1]}, "params": {"orientation": "landscape"}}
    plan, frags = plan_spec(DocxPackage(data), _spec(op))
    out = apply_spec(DocxPackage(data), frags).output_bytes()
    d = plan.to_dict()
    assert run_suite(data, out, d, "a.docx", "b.docx").passed
    for a in d["property_assertions"]:
        a.pop("select", None)
    rec = run_suite(data, out, d, "a.docx", "b.docx")
    assert not rec.passed and "carries no select" in rec.get("V10").summary


def test_a_pending_section_change_is_flagged():
    def edit(root):
        sp = DocxPackage.sect_prs(root)[1]
        chg = etree.SubElement(sp, W + "sectPrChange")
        chg.set(W + "id", "901"); chg.set(W + "author", "Jack")
        etree.SubElement(chg, W + "sectPr")
    data = patch_docx(build("21 13 13", "WET-PIPE SPRINKLER SYSTEMS"), {MAIN: edit})
    op = {"op": "set_page_setup", "select": {"sections": [1]}, "params": {"orientation": "landscape"}}
    plan, _ = plan_spec(DocxPackage(data), _spec(op))
    assert any("sectPrChange" in f for f in plan.flags)


def test_recovery_says_different_first_page_has_no_final_check():
    from spec_surgeon import recovery
    spec = _spec({"op": "set_page_setup", "select": {"sections": [0]}, "params": {"different_first_page": True}})
    _, notes = recovery._goals(spec)
    assert any("different-first-page" in n for n in notes)


# ---- merging an op with many repeated touches is not quadratic ----------------------------------------------------------

def test_repeated_touches_merge_in_linear_time():
    """One conversion over 2,000 paragraphs declared 16,000 touches, 6,000 distinct; the merge took 16 s."""
    import time
    from spec_surgeon.ops.base import PlanFragment, check_conflicts
    a = PlanFragment(job="a", mode="formatting", params={}, intended_parts=[], op="a", op_index=0)
    b = PlanFragment(job="b", mode="formatting", params={}, intended_parts=[], op="b", op_index=1)
    a.touches = [("word/styles.xml", "style:Normal", "verify:pPr")] * 3000 + [("word/document.xml", "1", "rPr")]
    b.touches = [("word/styles.xml", "style:Normal", "verify:pPr")] * 3000 + [("word/document.xml", "2", "rPr")]
    t = time.perf_counter()
    check_conflicts([a, b])
    assert time.perf_counter() - t < 2.0
    b.touches.append(("word/styles.xml", "style:Normal", "pPr"))         # still a conflict when there is one
    with pytest.raises(JobException) as e:
        check_conflicts([a, b])
    assert e.value.code == "conflict"
