"""Fixtures for remove_comments and remove_highlight (built-in batch 2, docs/BUILTIN_OPS_BATCH2_DESIGN.md §2–§3).

``build_commented()`` — a spec with comment threads as Word 2016+ writes them (comments.xml with w14:paraId, the
commentsExtended / commentsIds / commentsExtensible side parts): Jeremy's comment 0 with Jack's reply 1 (a resolved
thread); Jack's comment 2 with Jeremy's reply 3 (open); Jeremy's comment 4 spanning paragraphs C and D; Abe's comment 5
whose text holds a pending tracked insertion; Jeremy's comment 6 whose range starts between two paragraphs (a body-level
mark). ``side=False`` leaves the side parts out (Word 2010 or older); ``textbox=True`` puts a mark of comment 2 into a
text box; ``duplicate=True`` gives comment 3 comment 2's id. ``build_word_threads()`` moves the replies' marks where
Word writes them; ``build_classic_comments()`` is Word 2007's (no paragraph ids, no side parts) — built-in batch 6.

``build_highlighted()`` — yellow, green and cyan runs, a darkYellow run, an explicit ``none``, a highlighted paragraph
mark, a highlighted run in a header, inside a pending deletion and in a text box (both copies), bold highlighted text, a
run already carrying a tracked formatting change, and a character style that highlights.
"""
from __future__ import annotations

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

from make_fixtures import W, W_NS, _textbox_run, add_parts, docx_bytes, patch_docx

W14 = "http://schemas.microsoft.com/office/word/2010/wordml"
W15 = "http://schemas.microsoft.com/office/word/2012/wordml"
CID = "http://schemas.microsoft.com/office/word/2016/wordml/cid"
CEX = "http://schemas.microsoft.com/office/word/2018/wordml/cex"
DOC = "word/document.xml"
COMMENTS = "word/comments.xml"
EXTENDED_REL = "http://schemas.microsoft.com/office/2011/relationships/commentsExtended"
IDS_REL = "http://schemas.microsoft.com/office/2016/09/relationships/commentsIds"
EXTENSIBLE_REL = "http://schemas.microsoft.com/office/2018/08/relationships/commentsExtensible"
CT = "application/vnd.openxmlformats-officedocument.wordprocessingml.{}+xml"

# comment id → (parent id or None, done)
THREADS = {0: (None, True), 1: (0, True), 2: (None, False), 3: (2, False), 4: (None, False), 5: (None, False),
           6: (None, False)}


def para_id(cid: int, k: int = 0) -> str:
    return f"{0x1A000000 + cid * 16 + k:08X}"


def durable_id(cid: int) -> str:
    return f"{0x2B000000 + cid:08X}"


def _side_parts():
    ex = etree.Element("{%s}commentsEx" % W15, nsmap={"w15": W15})
    ids = etree.Element("{%s}commentsIds" % CID, nsmap={"w16cid": CID})
    cex = etree.Element("{%s}commentsExtensible" % CEX, nsmap={"w16cex": CEX})
    for cid, (parent, done) in THREADS.items():
        e = etree.SubElement(ex, "{%s}commentEx" % W15)
        e.set("{%s}paraId" % W15, para_id(cid))
        if parent is not None:
            e.set("{%s}paraIdParent" % W15, para_id(parent))
        e.set("{%s}done" % W15, "1" if done else "0")
        i = etree.SubElement(ids, "{%s}commentId" % CID)
        i.set("{%s}paraId" % CID, para_id(cid)); i.set("{%s}durableId" % CID, durable_id(cid))
        x = etree.SubElement(cex, "{%s}commentExtensible" % CEX)
        x.set("{%s}durableId" % CEX, durable_id(cid)); x.set("{%s}dateUtc" % CEX, "2026-10-01T12:00:00Z")
    ser = lambda r: etree.tostring(r, xml_declaration=True, encoding="UTF-8", standalone=True)
    return {"word/commentsExtended.xml": ser(ex), "word/commentsIds.xml": ser(ids),
            "word/commentsExtensible.xml": ser(cex)}


def build_commented(side: bool = True, textbox: bool = False, duplicate: bool = False) -> bytes:
    doc = Document()
    doc.add_paragraph("SECTION 21 13 13 - WET-PIPE SPRINKLER SYSTEMS")
    p = doc.add_paragraph("A.\tSprinklers shall be ")
    r = p.add_run("listed and approved")
    p.add_run(".")
    doc.add_comment(r, text="Which listing?", author="Jeremy", initials="JS")                  # 0
    doc.add_comment(r, text="UL and FM.", author="Jack", initials="JK")                         # 1, reply to 0
    p = doc.add_paragraph("B.\tPipe shall be ")
    r = p.add_run("black steel")
    p.add_run(", Schedule 40.")
    doc.add_comment(r, text="Galvanized for dry systems?", author="Jack", initials="JK")        # 2
    doc.add_comment(r, text="Yes, add a note.", author="Jeremy", initials="JS")                 # 3, reply to 2
    p = doc.add_paragraph("C.\tHangers shall be ")
    r = p.add_run("listed")
    doc.add_comment(r, text="This spans C and D.", author="Jeremy", initials="JS")             # 4, moved below
    doc.add_paragraph("D.\tSeismic bracing per NFPA 13.")
    p = doc.add_paragraph("E.\tProvide an ")
    r = p.add_run("inspector's test connection")
    doc.add_comment(r, text="Check this location", author="Abe", initials="AB")                # 5
    p = doc.add_paragraph("F.\tA comment that starts between paragraphs.")
    doc.add_comment(p.runs[0], text="Body-level start", author="Jeremy", initials="JS")        # 6, moved below
    p = doc.add_paragraph("G.\tNo comment here.")
    if textbox:
        _textbox_run(p, "Box text")
    data = docx_bytes(doc)

    def fix_document(root):
        pars = list(root.iter(W + "p"))
        by_text = lambda s: next(q for q in pars if "".join(t.text or "" for t in q.iter(W + "t")).startswith(s))
        c, d = by_text("C."), by_text("D.")
        for el in list(c):                                          # comment 4: its end and reference go into D
            if el.tag in (W + "commentRangeEnd",) or (el.tag == W + "r" and el.find(W + "commentReference") is not None):
                c.remove(el)
                d.append(el)
        f = by_text("F.")
        start = next(el for el in f if el.tag == W + "commentRangeStart")   # comment 6 starts at body level
        f.remove(start)
        f.addprevious(start)
        if textbox:                                                  # a reference mark of comment 2 in the text box
            for box_p in [q for q in pars if q.getparent().tag == W + "txbxContent"]:
                run = etree.SubElement(box_p, W + "r")
                ref = etree.SubElement(run, W + "commentReference")
                ref.set(W + "id", "2")

    def fix_comments(root):
        for c in root.findall(W + "comment"):
            cid = int(c.get(W + "id"))
            if duplicate and cid == 3:
                c.set(W + "id", "2")
            for k, cp in enumerate(c.iter(W + "p")):
                cp.set("{%s}paraId" % W14, para_id(cid, k))
            if cid == 5:                                             # a pending insertion inside the comment's text
                last = list(c.iter(W + "p"))[-1]
                run = etree.SubElement(last, W + "r")
                t = etree.SubElement(run, W + "t"); t.text = " — and its valve"
                ins = etree.Element(W + "ins")
                ins.set(W + "id", "900"); ins.set(W + "author", "Abe"); ins.set(W + "date", "2026-10-02T09:00:00Z")
                run.addprevious(ins); ins.append(run)
    data = patch_docx(data, {DOC: fix_document, COMMENTS: fix_comments})
    if side:
        parts = _side_parts()
        data = add_parts(data, parts,
                         rels=[("word/_rels/document.xml.rels", "rId90", EXTENDED_REL, "commentsExtended.xml"),
                               ("word/_rels/document.xml.rels", "rId91", IDS_REL, "commentsIds.xml"),
                               ("word/_rels/document.xml.rels", "rId92", EXTENSIBLE_REL, "commentsExtensible.xml")],
                         overrides={"word/commentsExtended.xml": CT.format("commentsExtended"),
                                    "word/commentsIds.xml": CT.format("commentsIds"),
                                    "word/commentsExtensible.xml": CT.format("commentsExtensible")})
    return data


def build_word_threads(side: bool = True) -> bytes:
    """``build_commented()`` with each reply's marks where Word writes them (built-in batch 6, §9): the reply's range
    start right after its parent's, and its range end and reference after the parent's reference — start0 start1 text
    end0 ref0 end1 ref1 (python-docx nests the reply inside its parent's marks instead)."""
    def word_order(root):
        for p in root.iter(W + "p"):
            kids = list(p)
            for parent, reply in (("0", "1"), ("2", "3")):
                end = next((e for e in kids if e.tag == W + "commentRangeEnd" and e.get(W + "id") == reply), None)
                if end is None:
                    continue
                ref = end.getnext()                                   # the reply's reference run
                parent_ref = next(e for e in kids if e.tag == W + "r" and e.find(W + "commentReference") is not None
                                  and e.find(W + "commentReference").get(W + "id") == parent)
                p.remove(end)
                p.remove(ref)
                parent_ref.addnext(ref)
                parent_ref.addnext(end)
    return patch_docx(build_commented(side=side), {DOC: word_order})


def build_classic_comments() -> bytes:
    """Comments as Word 2007 writes them (built-in batch 6, §9): no paragraph ids, no side parts — the replies' threads
    and the resolved state must be created with them."""
    def no_ids(root):
        for p in root.iter(W + "p"):
            p.attrib.pop("{%s}paraId" % W14, None)
    return patch_docx(build_commented(side=False), {COMMENTS: no_ids})


def build_plain(text: str = "Nothing to remove here.") -> bytes:
    doc = Document()
    doc.add_paragraph(text)
    return docx_bytes(doc)


# ---- highlighting ------------------------------------------------------------------------------------------------------

def _hl(run, color: str) -> None:
    rpr = run._r.get_or_add_rPr()
    h = OxmlElement("w:highlight"); h.set(qn("w:val"), color)
    rpr.append(h)


def build_highlighted() -> bytes:
    from docx.enum.style import WD_STYLE_TYPE
    doc = Document()
    st = doc.styles.add_style("Marker", WD_STYLE_TYPE.CHARACTER)
    st.font.highlight_color = WD_COLOR_INDEX.YELLOW
    doc.sections[0].header.paragraphs[0].add_run("HEADER ")
    _hl(doc.sections[0].header.paragraphs[0].add_run("DRAFT ISSUE"), "yellow")
    p = doc.add_paragraph("A.\tSprinklers shall be ")
    _hl(p.add_run("listed"), "yellow")
    p.add_run(" and ")
    r = p.add_run("approved")
    r.bold = True
    _hl(r, "green")
    p.add_run(".")
    p = doc.add_paragraph("B.\tPipe shall be ")
    _hl(p.add_run("Schedule 40"), "cyan")
    p.add_run(" and ")
    _hl(p.add_run("black"), "darkYellow")
    p.add_run(" steel; ")
    _hl(p.add_run("no highlight"), "none")
    p = doc.add_paragraph("C.\tA paragraph whose mark is highlighted.")
    mark = OxmlElement("w:rPr"); h = OxmlElement("w:highlight"); h.set(qn("w:val"), "yellow"); mark.append(h)
    p._p.get_or_add_pPr().append(mark)
    p = doc.add_paragraph("D.\tDeleted highlighted words: ")
    r = p.add_run("old wording")
    _hl(r, "yellow")
    p = doc.add_paragraph("E.\tStyled: ")
    p.add_run("marked by its style").style = "Marker"
    p = doc.add_paragraph("F.\tFormatted earlier: ")
    r = p.add_run("tracked bold")
    r.bold = True
    _hl(r, "yellow")
    p = doc.add_paragraph("G.\tA text box: ")
    _textbox_run(p, "Box words")
    data = docx_bytes(doc)

    def fix(root):
        pars = list(root.iter(W + "p"))
        by_text = lambda s: next(q for q in pars if "".join(t.text or "" for t in q.iter(W + "t")).startswith(s))
        d = by_text("D.")
        run = [r for r in d.findall(W + "r")][-1]                    # a pending deletion of highlighted words
        for t in run.findall(W + "t"):
            t.tag = W + "delText"
        dl = etree.Element(W + "del")
        dl.set(W + "id", "700"); dl.set(W + "author", "Jack"); dl.set(W + "date", "2026-10-03T08:00:00Z")
        run.addprevious(dl); dl.append(run)
        f = by_text("F.")
        rpr = f.findall(W + "r")[-1].find(W + "rPr")                  # an earlier tracked formatting change
        chg = etree.SubElement(rpr, W + "rPrChange")
        chg.set(W + "id", "701"); chg.set(W + "author", "Jeremy"); chg.set(W + "date", "2026-10-03T08:30:00Z")
        etree.SubElement(chg, W + "rPr")
        for box in root.iter(W + "txbxContent"):                      # the text box's words highlighted (both copies)
            for r in box.iter(W + "r"):
                rp = r.find(W + "rPr")
                if rp is None:
                    rp = etree.Element(W + "rPr"); r.insert(0, rp)
                h = etree.SubElement(rp, W + "highlight"); h.set(W + "val", "magenta")
    return patch_docx(data, {DOC: fix})


def build_comment_in_deletion() -> bytes:
    """Jeremy's comment on "black steel ", which Jack then deleted with Track Changes on: the comment's range marks and
    its reference sit inside Jack's ``w:del``, as Word keeps them. Accepting that deletion alone would take the
    comment's anchor away (``marks-in-change``); with the comment removed first, the deletion accepts cleanly."""
    from fx_revisions import JACK, wrap
    doc = Document()
    doc.add_paragraph("SECTION 21 13 13 - WET-PIPE SPRINKLER SYSTEMS")
    p = doc.add_paragraph("A. Pipe shall be ")
    r = p.add_run("black steel ")
    doc.add_comment(r, text="Why black?", author="Jeremy", initials="JS")
    p.add_run("Schedule 40.")
    kids = list(p._p)
    start = next(k for k in kids if k.tag == W + "commentRangeStart")
    end = next(k for k in kids if k.tag == W + "commentRangeEnd")
    ref = next(k for k in kids if k.tag == W + "r" and k.find(W + "commentReference") is not None)
    wrap(r, "del", JACK)
    deletion = r._r.getparent()
    for t in r._r.iter(W + "t"):
        t.tag = W + "delText"
    deletion.insert(0, start)
    deletion.append(end)
    deletion.append(ref)
    return docx_bytes(doc)
