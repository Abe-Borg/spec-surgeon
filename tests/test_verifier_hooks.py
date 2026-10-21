"""The verifier's shared hooks (docs/BUILTIN_OPS_BATCH2_DESIGN.md §1): structure-only declarations (V2/V3), the
V10 hooks (verified parts and members, expected removals outside paragraphs, the member residual over settings.xml and
the document-properties parts), and the both-sides definitions every batch-2 op relies on — fields, effective hidden and
colour, block text and structure — each compared transform side against verifier side."""
from __future__ import annotations

import copy

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

from make_fixtures import build_batch1, build_rich, docx_bytes, patch_docx
from spec_surgeon.ooxml import block_struct_counts, block_text_without, paragraphs
from spec_surgeon.ops.fields import fields_of
from spec_surgeon.ops.markup import Look, color_in, color_names
from spec_surgeon.ops.selectors import StyleSheet
from spec_surgeon.package import DocxPackage
from spec_surgeon.verify import props
from spec_surgeon.verify.fields import v_fields
from spec_surgeon.verify.markup import v_color, v_color_in, v_color_names, v_hidden, v_mark_hidden
from spec_surgeon.verify.props import Styles, element_path
from spec_surgeon.verify.reader import VPackage
from spec_surgeon.verify.suite import run_suite
from spec_surgeon.verify.textcheck import blocks

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W = "{%s}" % W_NS
DOC = "word/document.xml"


def _plan(**kw):
    base = {"job": "t", "mode": "formatting", "intended_parts": [DOC], "diffs": [], "rev_delta": {"ins": 0, "del": 0},
            "property_assertions": []}
    base.update(kw)
    return base


def _commented() -> bytes:
    doc = Document()
    doc.add_paragraph("Alpha paragraph")
    p = doc.add_paragraph("Beta paragraph")
    r = OxmlElement("w:r"); ref = OxmlElement("w:commentReference"); ref.set(qn("w:id"), "7"); r.append(ref)
    p._p.append(r)
    doc.add_paragraph("Gamma paragraph")
    return docx_bytes(doc)


def _without_reference(data: bytes) -> bytes:
    def fn(root):
        for ref in list(root.iter(W + "commentReference")):
            run = ref.getparent()
            run.getparent().remove(run)
    return patch_docx(data, {DOC: fn})


def _check(src, out, plan, cid):
    return run_suite(src, out, plan, "s.docx", "o.docx").get(cid)


# ---- structure-only declarations ---------------------------------------------------------------------------------------

def test_a_structure_only_declaration_covers_its_own_block_and_nothing_else():
    src = _commented()
    out = _without_reference(src)
    bi = next(b.index for b in blocks(VPackage(src).xml(DOC)) if b.text == "Beta paragraph")
    good = {"part": DOC, "block_index": bi, "before": "Beta paragraph", "after": "Beta paragraph", "struct_after": {}}
    plan = _plan(diffs=[good])
    assert _check(src, out, plan, "V2").status == "PASS"
    assert _check(src, out, plan, "V3").status == "PASS"
    # undeclared: V3 sees the reference mark go
    assert _check(src, out, _plan(), "V3").status == "FAIL"
    # declared for another block: not this one's declaration
    assert _check(src, out, _plan(diffs=[{**good, "block_index": bi - 1}]), "V3").status == "FAIL"
    # declared counts wrong
    assert _check(src, out, _plan(diffs=[{**good, "struct_after": {"commentReference": 1}}]), "V3").status == "FAIL"


def test_a_structure_only_declaration_whose_change_did_not_happen_fails():
    src = _commented()
    bi = next(b.index for b in blocks(VPackage(src).xml(DOC)) if b.text == "Beta paragraph")
    plan = _plan(diffs=[{"part": DOC, "block_index": bi, "before": "Beta paragraph", "after": "Beta paragraph",
                         "struct_after": {}}])
    rec = run_suite(src, src, plan, "s.docx", "o.docx")
    assert rec.get("V3").status == "FAIL"
    assert "declared structure change not found" in str(rec.get("V3").data)
    assert any("declared structure change not found" in x for _, _, _, d in rec.failures() for x in d)


def test_a_structure_only_block_stays_in_the_residual(monkeypatch):
    """A structure-only declaration satisfies V2/V3 but never exempts its paragraph from V10: the kind that removed the
    mark checks the paragraph item by item (verified_blocks), or the residual compares it — so another change in the
    same paragraph (an attribute, a bookmark id) cannot hide behind the declaration (Codex on #27)."""
    src = _commented()

    def fn(root, also=False):
        for ref in list(root.iter(W + "commentReference")):
            p = ref.getparent().getparent()
            p.remove(ref.getparent())
            if also:
                p.set(W + "rsidR", "00BADBAD")
    bi = next(b.index for b in blocks(VPackage(src).xml(DOC)) if b.text == "Beta paragraph")
    plan = _plan(diffs=[{"part": DOC, "block_index": bi, "before": "Beta paragraph", "after": "Beta paragraph",
                         "struct_after": {}}])
    hidden = patch_docx(src, {DOC: lambda r: fn(r, also=True)})
    rec = run_suite(src, hidden, plan, "s.docx", "o.docx")
    assert rec.get("V3").status == "PASS" and rec.get("V10").status == "FAIL"
    # with no kind checking the paragraph, even the declared change alone is the residual's to see (fail closed)
    clean = patch_docx(src, {DOC: fn})
    assert run_suite(src, clean, plan, "s.docx", "o.docx").get("V10").status == "FAIL"

    def kind(ctx, a, errs):                          # a kind that checked the paragraph item by item takes it out
        ctx.verified_blocks.setdefault(DOC, {})[bi] = bi
        return "test: paragraph checked item by item"
    monkeypatch.setitem(props.KINDS, "test_items", kind)
    plan["property_assertions"] = [{"kind": "test_items"}]
    assert run_suite(src, clean, plan, "s.docx", "o.docx").get("V10").status == "PASS"


def test_a_text_change_hiding_behind_a_structure_only_declaration_still_fails_v2():
    src = _commented()

    def fn(root):
        for t in root.iter(W + "t"):
            if t.text == "Beta paragraph":
                t.text = "Beta paragraph!"
        for ref in list(root.iter(W + "commentReference")):
            ref.getparent().getparent().remove(ref.getparent())
    out = patch_docx(src, {DOC: fn})
    bi = next(b.index for b in blocks(VPackage(src).xml(DOC)) if b.text == "Beta paragraph")
    plan = _plan(diffs=[{"part": DOC, "block_index": bi, "before": "Beta paragraph", "after": "Beta paragraph",
                         "struct_after": {}}])
    assert _check(src, out, plan, "V2").status == "FAIL"


# ---- the V10 hooks ----------------------------------------------------------------------------------------------------

def _settings_off(data: bytes) -> bytes:
    def fn(root):
        for el in list(root):
            if el.tag == W + "trackRevisions":
                root.remove(el)
    return patch_docx(data, {"word/settings.xml": fn})


def test_the_member_residual_fails_an_unasserted_settings_change(monkeypatch):
    src = build_rich()
    out = _settings_off(src)
    plan = _plan(intended_parts=["word/settings.xml"])
    rec = run_suite(src, out, plan, "s.docx", "o.docx")
    assert rec.get("V1").status == "PASS"                     # intended: V1 alone would let it through
    assert rec.get("V10").status == "FAIL" and "word/settings.xml: unpredicted change" in rec.get("V10").summary

    def kind(ctx, a, errs):
        ctx.member_strip.setdefault("word/settings.xml", set()).add(W + "trackRevisions")
        return "test: trackRevisions verified"
    monkeypatch.setitem(props.KINDS, "test_settings", kind)
    plan["property_assertions"] = [{"kind": "test_settings"}]
    assert run_suite(src, out, plan, "s.docx", "o.docx").get("V10").status == "PASS"
    # the registration covers that element only: another settings change still fails
    def also(root):
        for el in list(root):
            if el.tag == W + "trackRevisions":
                root.remove(el)
        el = etree.SubElement(root, W + "updateFields")
    worse = patch_docx(src, {"word/settings.xml": also})
    assert run_suite(src, worse, plan, "s.docx", "o.docx").get("V10").status == "FAIL"


def test_a_core_properties_change_is_covered_or_fails():
    src = build_rich()

    def fn(root):
        title = root.find("{http://purl.org/dc/elements/1.1/}title")
        if title is None:
            title = etree.SubElement(root, "{http://purl.org/dc/elements/1.1/}title")
        title.text = "Changed"
    out = patch_docx(src, {"docProps/core.xml": fn})
    rec = run_suite(src, out, _plan(intended_parts=["docProps/core.xml"]), "s.docx", "o.docx")
    assert rec.get("V10").status == "FAIL" and "docProps/core.xml" in rec.get("V10").summary


def _body_bookmark() -> bytes:
    doc = Document()
    doc.add_paragraph("First")
    doc.add_paragraph("Second")
    data = docx_bytes(doc)

    def fn(root):
        body = root.find(W + "body")
        bs = etree.Element(W + "bookmarkStart"); bs.set(W + "id", "5"); bs.set(W + "name", "between")
        be = etree.Element(W + "bookmarkEnd"); be.set(W + "id", "5")
        body.insert(1, bs)
        body.insert(2, be)
    return patch_docx(data, {DOC: fn})


def test_expected_removed_strips_exactly_the_named_elements(monkeypatch):
    src = _body_bookmark()

    def drop_marks(root, also=False):
        body = root.find(W + "body")
        for el in list(body):
            if el.tag in (W + "bookmarkStart", W + "bookmarkEnd"):
                body.remove(el)
        if also:
            body.find(W + "p").set(W + "rsidR", "00ABCDEF")
    out = patch_docx(src, {DOC: drop_marks})
    plan = _plan()
    assert run_suite(src, out, plan, "s.docx", "o.docx").get("V10").status == "FAIL"

    def kind(ctx, a, errs):
        root = ctx.view(DOC)[0].root
        ctx.expected_removed[DOC] = [element_path(el) for el in root.find(W + "body")
                                     if el.tag in (W + "bookmarkStart", W + "bookmarkEnd")]
        return "test: marks expected gone"
    monkeypatch.setitem(props.KINDS, "test_gone", kind)
    plan["property_assertions"] = [{"kind": "test_gone"}]
    assert run_suite(src, out, plan, "s.docx", "o.docx").get("V10").status == "PASS"
    worse = patch_docx(src, {DOC: lambda r: drop_marks(r, also=True)})
    assert run_suite(src, worse, plan, "s.docx", "o.docx").get("V10").status == "FAIL"
    # an element still there that was expected gone: the source copy loses it, the output keeps it — a difference
    assert run_suite(src, src, plan, "s.docx", "o.docx").get("V10").status == "FAIL"


def test_an_expected_removal_inside_a_paragraph_is_refused(monkeypatch):
    src = _commented()
    out = _without_reference(src)

    def kind(ctx, a, errs):
        root = ctx.view(DOC)[0].root
        ctx.expected_removed[DOC] = [element_path(next(root.iter(W + "commentReference")))]
        return "test"
    monkeypatch.setitem(props.KINDS, "test_inside", kind)
    rec = run_suite(src, out, _plan(property_assertions=[{"kind": "test_inside"}]), "s.docx", "o.docx")
    assert rec.get("V10").status == "FAIL" and "not an element outside the paragraphs" in rec.get("V10").summary


def test_verified_parts_and_members_leave_the_residual(monkeypatch):
    src = _commented()
    out = _without_reference(src)
    plan = _plan()
    assert run_suite(src, out, plan, "s.docx", "o.docx").get("V10").status == "FAIL"

    def kind(ctx, a, errs):
        ctx.verified_parts.add(DOC)
        return "test: compared whole elsewhere"
    monkeypatch.setitem(props.KINDS, "test_whole", kind)
    plan["property_assertions"] = [{"kind": "test_whole"}]
    assert run_suite(src, out, plan, "s.docx", "o.docx").get("V10").status == "PASS"


# ---- both-sides definitions agree ----------------------------------------------------------------------------------

FIELD_XML = f"""<w:document xmlns:w="{W_NS}"><w:body>
<w:p><w:r><w:t>A </w:t></w:r><w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText> DOCPROPERTY Title </w:instrText></w:r>
<w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:t>Spec</w:t></w:r><w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>
<w:p><w:fldSimple w:instr=" DATE \\@ &quot;yyyy&quot; "><w:r><w:t>2026</w:t></w:r></w:fldSimple></w:p>
<w:p><w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText> IF </w:instrText></w:r>
<w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText>MERGEFIELD x</w:instrText></w:r><w:r><w:fldChar w:fldCharType="end"/></w:r>
<w:r><w:instrText> = "y" "a" "b"</w:instrText></w:r><w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:t>a</w:t></w:r>
<w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>
<w:p><w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText> TOC \\o </w:instrText></w:r><w:r><w:fldChar w:fldCharType="separate"/></w:r>
<w:r><w:t>One</w:t></w:r></w:p>
<w:p><w:r><w:t>Two</w:t></w:r><w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>
<w:p><w:r><w:fldChar w:fldCharType="begin"><w:ffData/></w:fldChar></w:r><w:r><w:instrText> FORMTEXT </w:instrText></w:r>
<w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:t>x</w:t></w:r><w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>
<w:p><w:ins w:id="1" w:author="J" w:date="2026-01-01T00:00:00Z"><w:r><w:fldChar w:fldCharType="begin"/></w:r>
<w:r><w:instrText> PAGE </w:instrText></w:r><w:r><w:fldChar w:fldCharType="end"/></w:r></w:ins></w:p>
<w:p><w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText>=2+2</w:instrText></w:r><w:r><w:fldChar w:fldCharType="end"/></w:r>
<w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText xml:space="preserve"> HYPERLINK "https://a.b" \\l x </w:instrText></w:r>
<w:r><w:fldChar w:fldCharType="separate"/></w:r><w:hyperlink><w:r><w:t>link</w:t></w:r></w:hyperlink><w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>
</w:body></w:document>"""


def test_the_two_field_parsers_agree():
    root = etree.fromstring(FIELD_XML.encode())
    t, v = fields_of(root), v_fields(root)
    assert len(t) == len(v) == 9
    for a, b in zip(t, v):
        assert a.type == b.type and a.instruction == b.instruction
        assert a.refusal() == b.refusal()
        assert [id(x) for x in a.markers()] == [id(x) for x in b.markers()]
        assert (a.simple is None) == (b.simple is None)
        assert (a.parent is None) == (b.parent is None) and a.in_parent_code == b.in_parent_code
    assert [f.type for f in t] == ["DOCPROPERTY", "DATE", "IF", "MERGEFIELD", "TOC", "FORMTEXT", "PAGE", "=", "HYPERLINK"]
    assert [f.refusal() for f in t] == [None, None, "nested-field", None, "field-span", "form-field",
                                        "existing-revisions", None, None]


def _looks() -> bytes:
    from docx.enum.style import WD_STYLE_TYPE
    doc = Document()
    st = doc.styles.add_style("Note", WD_STYLE_TYPE.PARAGRAPH)
    st.font.hidden = True
    st.font.color.rgb = None
    cs = doc.styles.add_style("Blue", WD_STYLE_TYPE.CHARACTER)
    cs.font.hidden = True
    p = doc.add_paragraph("plain ")
    r = p.add_run("direct hidden"); r.font.hidden = True
    r = p.add_run("direct shown"); r.font.hidden = False
    p = doc.add_paragraph("styled note", style="Note")
    p.add_run("both hidden").style = "Blue"                         # paragraph style and character style: ambiguous
    p = doc.add_paragraph()
    r = p.add_run("blue"); r.font.color.rgb = __import__("docx.shared", fromlist=["RGBColor"]).RGBColor(0, 0, 0xFF)
    r = p.add_run("theme")
    rpr = r._r.get_or_add_rPr(); c = OxmlElement("w:color"); c.set(qn("w:val"), "4472C4")
    c.set(qn("w:themeColor"), "accent1"); rpr.append(c)
    return docx_bytes(doc)


@pytest.mark.parametrize("build", [_looks, build_rich, build_batch1])
def test_the_two_hidden_and_colour_readers_agree(build):
    data = build()
    pkg, vp = DocxPackage(data), VPackage(data)
    look = Look(StyleSheet(pkg))
    styles = Styles(vp.xml(vp.related("http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles")))
    wanted = ["0000FF", "accent1", "auto"]
    seen = set()
    for part in pkg.text_parts():
        root = pkg.xml(part)
        for p in paragraphs(root):
            assert look.mark_hidden(p) == v_mark_hidden(p, styles)
            for r in p.iter(W + "r"):
                rpr = r.find(W + "rPr")
                h = look.hidden(rpr, p)
                assert h == v_hidden(rpr, p, styles)
                c = look.color(rpr, p)
                assert c == v_color(rpr, p, styles)
                assert color_in(c, color_names(wanted)) == v_color_in(c, v_color_names(wanted))
                seen.add(h)
    if build is _looks:
        assert seen == {True, False, None}                           # direct, off, ambiguous all met


@pytest.mark.parametrize("build", [build_rich, build_batch1, _commented])
def test_block_text_and_structure_agree_with_the_verifier(build):
    data = build()
    pkg, vp = DocxPackage(data), VPackage(data)
    for part in pkg.text_parts():
        pars = paragraphs(pkg.xml(part))
        vb = blocks(vp.xml(part))
        assert [block_text_without(p) for p in pars] == [b.text for b in vb]
        assert [block_struct_counts(p) for p in pars] == [b.struct for b in vb]


def test_block_text_without_leaves_out_exactly_the_named_elements():
    root = etree.fromstring(FIELD_XML.encode())
    p = paragraphs(root)[0]
    f = fields_of(root)[0]
    assert block_text_without(p) == "A  DOCPROPERTY Title Spec"
    assert block_text_without(p, f.markers()) == "A Spec"
    assert block_struct_counts(p) == {"fldChar": 3} and block_struct_counts(p, f.markers()) == {}
