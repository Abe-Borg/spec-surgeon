"""Safe mode's buttons (``recipes``): every button builds a fixed spec from its form
and the attachments, the spec passes the one gate every submission passes, and the window's preview → approve → run →
verify carries it out with no AI call — no visual check, no recovery — on clean and dirty documents alike. What a
button cannot do on a document is said beside the copies, and can be handed to Default mode."""
from __future__ import annotations

import re
from pathlib import Path

import pytest
from docx import Document
from docx.shared import Pt
from lxml import etree

from fx_removal import build_spec_notes
from make_fixtures import DIRTY_KINDS, build, build_dirty, build_reference_spec, build_rich, build_sized, build_template
from spec_surgeon import compiler, layouts, llm, preview, recipes, vision
from spec_surgeon.document_task import SAFE_NO_CONVERSATION, WORK_TODO, DocumentTask
from spec_surgeon.editspec import plan_spec
from spec_surgeon.fingerprint import fingerprint
from spec_surgeon.gui import Session
from spec_surgeon.ops import hidden as hidden_ops
from spec_surgeon.ops.selectors import StyleSheet
from spec_surgeon.package import FOOTER_REL, DocxPackage, W
from spec_surgeon.settings import MODES, Settings, load_settings, save_settings


@pytest.fixture(scope="module")
def folder(tmp_path_factory) -> Path:
    d = tmp_path_factory.mktemp("recipes")
    for num, title in (("21 13 13", "WET-PIPE SPRINKLER SYSTEMS"), ("21 13 16", "DRY-PIPE SPRINKLER SYSTEMS")):
        (d / f"{num} – {title}.docx").write_bytes(build(num, title))
    (d / "tpl").mkdir()
    (d / "tpl" / "00_00_00_TEMPLATE.docx").write_bytes(build_template())
    (d / "ref").mkdir()
    (d / "ref" / "08 71 00 - DOOR HARDWARE (ACME).docx").write_bytes(build_reference_spec())
    return d


def docs(folder: Path):
    return sorted(folder.glob("*.docx"))


def quiet(_m):
    pass


def built_and_gated(recipe_id, inputs, files, template=None, references=()):
    ctx = recipes.Context(list(files), template, list(references))
    built = recipes.build(recipe_id, inputs, ctx)
    outcome, errors = compiler.compose_submission(built.spec, compiler._Batch(files, template, input_files=list(references)),
                                                  [built.request, *built.sources])
    return built, outcome, errors


def plain_template(path: Path) -> Path:
    """A Word file with no header or footer at all: nothing to copy."""
    doc = Document()
    doc.add_paragraph("SECTION 01 00 00")
    doc.add_paragraph("GENERAL")
    doc.save(path)
    return path


# ---- the catalog ----------------------------------------------------------------------------------------------------

def test_every_button_is_in_a_group_with_a_form_the_panel_can_render():
    seen = set()
    for item in recipes.catalog():
        assert item["group"] in recipes.GROUPS and item["id"] not in seen
        seen.add(item["id"])
        assert item["label"] and item["summary"].endswith(".")
        for f in item["fields"]:
            assert f["kind"] in {"text", "lines", "pairs", "choice", "check", "picture", "table", "slots", "roles"}
            if f["kind"] == "choice":
                assert f["default"] in {o["value"] for o in f["options"]}
    assert seen == set(recipes.BY_ID)
    # the owner's list (2026-10-23): first set and second set, nothing dropped back in; and the body button; and the
    # third set (2026-10-28: N1–N12 of the Safe mode proposal, with Accept widened to reject); Save clean PDFs left
    # for the top bar's Export → PDF (Clean), in every mode (2026-10-29); Remove specifier notes, one press with no
    # form (M2, 2026-10-29)
    assert seen == {"headers_footers_template", "remove_headers_footers", "watermark", "swap_logo", "font_template",
                    "find_replace", "accept_changes", "remove_comments", "remove_highlight", "remove_hidden_text",
                    "remove_empty_paragraphs", "name_files", "title_property", "track_changes_off", "body_formatting",
                    "set_font", "page_setup", "issue_stamp", "end_of_section", "nonbreaking_spaces",
                    "clean_spaces", "quotes_dashes", "remove_hyperlinks", "freeze_fields", "remove_unused_styles",
                    "author_company", "proofing_language", "remove_notes"}


def test_tracking_is_offered_only_where_the_operation_can_track_and_is_off_by_default():
    with_check = {r.id for r in recipes.RECIPES if any(f.name == "tracked" for f in r.fields)}
    assert with_check == {"find_replace", "remove_hidden_text", "remove_empty_paragraphs", "issue_stamp",
                          "end_of_section", "nonbreaking_spaces", "clean_spaces", "quotes_dashes"}
    assert recipes.TRACKED.default is False


# ---- the template kinds ---------------------------------------------------------------------------------------------

def test_a_prepared_template_fills_number_and_title_from_each_document_and_asks_for_the_rest(folder):
    kind = recipes.template_kind(folder / "tpl" / "00_00_00_TEMPLATE.docx")
    assert kind["kind"] == "prepared" and kind["sections"] == 2
    assert kind["typed_slots"] == ["ISSUED FOR PERMIT"]
    assert set(kind["filled_slots"]) == {"XX XX XX", "SECTION_NUMBER", "SECTION_TITLE"}


def test_an_ordinary_spec_is_a_reference_template(folder):
    kind = recipes.template_kind(folder / "ref" / "08 71 00 - DOOR HARDWARE (ACME).docx")
    assert kind["kind"] == "reference" and set(kind["reference_slots"]) == {"SECTION_NUMBER", "SECTION_TITLE"}


def test_a_template_without_headers_or_footers_is_unusable_with_the_reason(tmp_path):
    kind = recipes.template_kind(plain_template(tmp_path / "plain.docx"))
    assert kind["kind"] == "unusable" and "no header or footer" in kind["reason"]
    assert recipes.template_kind(None)["kind"] == "none"


def test_slot_names_for_the_number_and_title(folder):
    assert [recipes.slot_role(n) for n in ("XX XX XX", "SECTION_NUMBER", "section number", "{{TITLE}}"[2:-2],
                                           "ISSUED FOR PERMIT")] == ["number", "number", "number", "title", ""]
    # the owner's template (2026-10-28): "00 00 00" and "Specification Name" were asked for in the form, and typing
    # the placeholder back in left it in every copy's footer
    assert [recipes.slot_role(n) for n in ("00 00 00", "000000", "xx.xx.xx", "## ## ##", "00 00 00.00", "Spec No.",
                                           "[SECTION NUMBER]")] == ["number"] * 7
    assert [recipes.slot_role(n) for n in ("Specification Name", "Section_Name", "Spec Title", "<SECTION TITLE>",
                                           "«Specification Section Name»")] == ["title"] * 5
    assert [recipes.slot_role(n) for n in ("Project Name", "NAME", "DATE", "00 00", "0000", "XX XX XX XX",
                                           "100% IFR PACKAGE", "21 05 17")] == [""] * 8


def named_slot_template(path: Path) -> Path:
    """The owner's prepared template as it was laid out (2026-10-28): one section, a header stamp the user types, and a
    footer with the title placeholder "Specification Name" over the number placeholder "00 00 00" and the page."""
    from make_fixtures import _highlight, _page_field
    doc = Document()
    header = doc.sections[0].header.paragraphs[0]
    header.add_run("MICROSOFT CONFIDENTIAL\t")
    _highlight(header.add_run("100% IFR PACKAGE"))
    footer = doc.sections[0].footer
    _highlight(footer.paragraphs[0].add_run("Specification Name"))
    number = footer.add_paragraph()
    _highlight(number.add_run("00 00 00"))
    number.add_run(" - ")
    _page_field(number)
    doc.add_paragraph("SECTION 00 00 00")
    doc.add_paragraph("SPECIFICATION NAME")
    doc.save(path)
    return path


def test_a_template_with_00_00_00_and_specification_name_fills_both_from_each_document(folder, tmp_path,
                                                                                         monkeypatch):
    template = named_slot_template(tmp_path / "00_00_00_SPECIFICATION HEADER AND FOOTER TEMPLATE.docx")
    kind = recipes.template_kind(template)
    assert kind["kind"] == "prepared" and kind["typed_slots"] == ["100% IFR PACKAGE"]
    assert kind["slot_roles"] == {"Specification Name": "title", "00 00 00": "number"}
    monkeypatch.setattr(llm, "client", lambda *a, **k: (_ for _ in ()).throw(AssertionError("no AI")))
    task = DocumentTask(Session(out_root=tmp_path / "out", visual_check=False), render=False, mode="safe")
    task.attach(docs(folder))
    task.attach([template], template=True)
    assert task.template_info()["typed_slots"] == ["100% IFR PACKAGE"]          # what the form asks for
    view = task.perform("recipe", recipe="headers_footers_template",
                        inputs={"slots": {"100% IFR PACKAGE": "100% IFR"}})
    assert view["stage"] == "preview", [m["text"] for m in task.messages]
    fills = task.session.spec["ops"][0]["params"]["slot_fills"]
    assert fills["00 00 00"] == recipes.REFERENCE_FILLS["SECTION_NUMBER"]
    assert fills["Specification Name"] == recipes.REFERENCE_FILLS["SECTION_TITLE"]
    task.perform("primary", token=view["token"])
    assert len(task.results) == 2 and all(r.status == "PASS" for r in task.session.last_batch.files)
    for original, output in task.results.items():
        pkg = DocxPackage.from_path(Path(output))
        footers = [r.resolved for r in pkg.rels(pkg.main_part()).values() if r.type == FOOTER_REL]
        shown = " | ".join("".join(t.text or "" for t in p.iter(W + "t"))
                           for part in footers for p in pkg.xml(part).iter(W + "p"))
        number, title = Path(original).stem.split(" – ")
        assert f"{title} | {number} - 1" in shown and "00 00 00" not in shown and "Specification Name" not in shown


# ---- every button through the gate and the engine -------------------------------------------------------------------

CASES = {
    "headers_footers_template": {"slots": {"ISSUED FOR PERMIT": "ISSUED FOR CONSTRUCTION"}},
    "remove_headers_footers": {},
    "watermark": {"text": "NOT FOR CONSTRUCTION"},
    "find_replace": {"pairs": [{"find": "sprinkler", "replace": "fire sprinkler"}]},
    "accept_changes": {},
    "remove_comments": {},
    "remove_highlight": {},
    "remove_hidden_text": {},
    "remove_notes": {},
    "remove_empty_paragraphs": {},
    "name_files": {},
    "title_property": {"pattern": "SECTION {number} - {title}"},
    "track_changes_off": {},
    "font_template": {},
    "set_font": {"font": "Times New Roman", "size": "11"},
    "page_setup": {"paper": "letter", "margin_top": "1", "margin_bottom": "1", "margin_left": "1.25",
                   "margin_right": "1.25", "header_footer": "0.5"},
    "issue_stamp": {"current": ["ISSUED FOR CONSTRUCTION  2026-01-01"], "text": "ISSUED FOR BID 2026-11-20"},
    "end_of_section": {"text": "END OF SECTION {number}"},
    "nonbreaking_spaces": {"units": "psi, gpm", "section_numbers": False},
    "clean_spaces": {},
    "quotes_dashes": {"quotes": "curly"},
    "remove_hyperlinks": {},
    "freeze_fields": {},
    "remove_unused_styles": {"which": "all"},
    "author_company": {"author": "A. Designer", "company": "ACME Engineering"},
    "proofing_language": {"language": "en-US"},
}


# The buttons whose copies are rendered and compared here (V6/V7 on real pages, where a renderer is installed): the
# ones that move what a page shows or check it against a reference of their own. The others change words, properties
# or markup in place; their ops' own tests render them (tests/test_*.py, "libreoffice_opens_and_renders").
RENDERED = {"font_template", "headers_footers_template", "page_setup", "remove_headers_footers", "watermark"}


@pytest.mark.parametrize("recipe_id", [pytest.param(r, marks=pytest.mark.render) if r in RENDERED else r
                                       for r in sorted(CASES)])
def test_each_button_previews_and_runs_to_verified_copies_without_any_ai_call(recipe_id, folder, tmp_path,
                                                                               monkeypatch):
    def never(*a, **k):
        raise AssertionError("a Safe-mode button makes no AI call")
    monkeypatch.setattr(vision, "review_samples", never)
    monkeypatch.setattr(llm, "client", never)
    monkeypatch.setattr(preview, "render_samples", lambda *a, **k: None)   # the pages, without a renderer
    template = folder / ("ref/08 71 00 - DOOR HARDWARE (ACME).docx" if recipe_id == "font_template"
                         else "tpl/00_00_00_TEMPLATE.docx")
    task = DocumentTask(Session(out_root=tmp_path / "out"), render=True, mode="safe")   # key "present" (conftest)
    task.attach(docs(folder))
    task.attach([template], template=True)
    before = {p: p.read_bytes() for p in docs(folder)}
    view = task.perform("recipe", recipe=recipe_id, inputs=CASES[recipe_id])
    assert view["stage"] == "preview", [m["text"] for m in task.messages]
    assert task.session.model_free and view["action"] == "Create edited copies"
    assert view["messages"][-1]["role"] == "user"           # the button's sentence, as what the user asked
    view = task.perform("primary", token=view["token"])
    assert view["stage"] == "done" and len(task.results) == 2, view["left_untouched"]
    assert view["left_untouched"] == []
    assert {p: p.read_bytes() for p in before} == before     # originals untouched
    for record in task.session.last_batch.files:
        assert record.status == "PASS"
    assert task.history and task.generation == 1             # the next button builds on the copies


def test_the_button_sentence_is_what_a_later_conversation_hears(folder, tmp_path):
    task = DocumentTask(Session(out_root=tmp_path / "out", visual_check=False), render=False, mode="safe")
    task.attach(docs(folder))
    view = task.perform("recipe", recipe="remove_highlight", inputs={})
    task.perform("primary", token=view["token"])
    assert any("highlight" in line.lower() for line in task._history_in_scope())


def test_dirty_documents_plan_or_are_left_untouched_with_a_reason_never_half_edited(tmp_path, monkeypatch):
    monkeypatch.setattr(llm, "client", lambda *a, **k: (_ for _ in ()).throw(AssertionError("no AI")))
    files = []
    for kind in DIRTY_KINDS:
        path = tmp_path / f"dirty-{kind}.docx"
        path.write_bytes(build_dirty(kind))
        files.append(path)
    task = DocumentTask(Session(out_root=tmp_path / "out", visual_check=False), render=False, mode="safe")
    task.attach(files)
    view = task.perform("recipe", recipe="title_property", inputs={"pattern": "SECTION {number} - {title}"})
    assert view["stage"] == "preview"
    untouched = {i["name"] for i in view["set_aside"]}
    assert "dirty-empty_body.docx" in untouched                # no SECTION line to read: set aside, with its reason
    assert all(i["reason"] for i in view["set_aside"])
    task.perform("primary", token=view["token"])
    left = {e["name"]: e["reason"] for e in task.view()["left_untouched"]}
    assert "dirty-empty_body.docx" in left and left["dirty-empty_body.docx"]
    for record in task.session.last_batch.files:
        assert record.status in {"PASS", "EXCEPTION"}            # a copy passes every check, or none is made


def test_an_unusable_form_is_said_once_and_nothing_is_previewed(folder, tmp_path):
    task = DocumentTask(Session(out_root=tmp_path / "out", visual_check=False), render=False, mode="safe")
    task.attach(docs(folder))
    view = task.perform("recipe", recipe="find_replace", inputs={"pairs": [{"find": "owner", "replace": "Owner"}]})
    assert view["stage"] == "empty" and task.session.spec is None
    assert "changes only capitals" in task.messages[-1]["text"]
    view = task.perform("recipe", recipe="headers_footers_template", inputs={})
    assert "Attach your template first" in task.messages[-1]["text"] and task.session.spec is None
    view = task.perform("recipe", recipe="watermark", inputs={"action": "add", "text": ""})
    assert "watermark's words" in task.messages[-1]["text"]


def test_a_template_with_nothing_to_fill_is_refused_with_default_mode_offered(folder, tmp_path):
    files = docs(folder)
    with pytest.raises(recipes.RecipeError, match="Default mode"):
        recipes.build("headers_footers_template", {}, recipes.Context(files, plain_template(tmp_path / "p.docx")))


def test_a_prepared_template_needs_its_typed_values(folder):
    files = docs(folder)
    template = folder / "tpl" / "00_00_00_TEMPLATE.docx"
    with pytest.raises(recipes.RecipeError, match="ISSUED FOR PERMIT"):
        recipes.build("headers_footers_template", {"slots": {}}, recipes.Context(files, template))
    built, outcome, errors = built_and_gated("headers_footers_template",
                                             {"slots": {"ISSUED FOR PERMIT": "ISSUED FOR BID"}}, files, template)
    assert errors == [] and built.spec["mode"] == "untracked" and built.sources == ["ISSUED FOR BID"]
    fills = built.spec["ops"][0]["params"]["slot_fills"]
    assert fills["ISSUED FOR PERMIT"] == {"rule": "literal", "value": "ISSUED FOR BID"}
    assert fills["XX XX XX"] == recipes.REFERENCE_FILLS["SECTION_NUMBER"]


def test_a_reference_template_fills_only_the_values_its_headers_show(folder):
    files = docs(folder)
    template = folder / "ref" / "08 71 00 - DOOR HARDWARE (ACME).docx"
    built, outcome, errors = built_and_gated("headers_footers_template", {}, files, template)
    params = built.spec["ops"][0]["params"]
    assert errors == [] and params["slot_markers"] == ["reference"] and built.spec["mode"] == "formatting"
    assert set(params["slot_fills"]) == {"SECTION_NUMBER", "SECTION_TITLE"} and built.sources == []


# ---- the words a button places are the user's, and only those -------------------------------------------------------

def test_typed_words_are_the_only_words_a_button_may_place(folder):
    files = docs(folder)
    built, _, errors = built_and_gated("find_replace", {"pairs": [{"find": "sprinkler", "replace": "fire sprinkler"},
                                                                  {"find": "SYSTEMS", "replace": ""}]}, files)
    assert errors == [] and built.sources == ["fire sprinkler"]
    built, _, errors = built_and_gated("watermark", {"text": "DRAFT"}, files)
    assert errors == [] and built.sources == ["DRAFT"]
    # a spec whose placed words were not typed is refused by the gate (the same gate the conversation passes)
    built.spec["ops"][0]["params"]["text"] = "PRELIMINARY"
    _, errors = compiler.compose_submission(built.spec, compiler._Batch(files, None), [built.request, *built.sources])
    assert errors and "never gave them verbatim" in " ".join(errors)


def test_a_skip_is_accepted_only_as_the_user_unticked_it(folder):
    """Backlog F1, rule 1: a replace_table.skip entry passes the gate only when the window recorded it from the user's
    unticks (its key); a spec that writes one on its own is refused, in every mode."""
    from spec_surgeon.ops import get_op
    files = docs(folder)
    built = recipes.build("find_replace", {"pairs": [{"find": "sprinkler", "replace": "fire sprinkler"}],
                                           "whole_word": False, "match_case": False},
                          recipes.Context(files, None, []))
    first = files[0]
    op = built.spec["ops"][0]
    listed = get_op("replace_table").plan(DocxPackage.from_path(first), op.get("select") or {}, op["params"])
    occ = next(a for a in listed.property_assertions if a["kind"] == "replaced_text")["occurrences"]
    assert occ, "the fixture holds the word"
    o = occ[0]
    skip = {"file": str(first), "part": o["part"], "paragraph": o["paragraph"], "start": o["start"], "end": o["end"],
            "find": o["find"]}
    spec = {**built.spec, "ops": [{**op, "params": {**op["params"], "skip": [skip]}}]}
    batch = compiler._Batch(files, None)
    sources = [built.request, *built.sources]
    _, errors = compiler.compose_submission(spec, batch, sources)
    assert errors and "did not untick" in errors[0]
    _, errors = compiler.compose_submission(spec, batch, sources, skips=[compiler.skip_record_key({**skip, "start": 99})])
    assert errors and "did not untick" in errors[0]
    _, errors = compiler.compose_submission(spec, batch, sources, unguarded=True, shown=["anything"])
    assert errors and "did not untick" in errors[0]              # Obliterate lifts no part of this
    outcome, errors = compiler.compose_submission(spec, batch, sources, skips=[compiler.skip_record_key(skip)])
    assert errors == [] and outcome["ops"][0]["params"]["skip"] == [skip]
    # a malformed entry is validate_spec's refusal, in its own words
    bad = {**built.spec, "ops": [{**op, "params": {**op["params"], "skip": [{**skip, "offset": 1}]}}]}
    _, errors = compiler.compose_submission(bad, batch, sources, skips=[compiler.skip_record_key(skip)])
    assert errors and "offset" in errors[0]


def test_the_tracked_checkbox_sets_the_mode_and_clean_copies_are_the_default(folder):
    files = docs(folder)
    built, _, _ = built_and_gated("remove_hidden_text", {}, files)
    assert built.spec["mode"] == "untracked" and built.spec["ops"][0]["params"]["tracked"] is False
    built, _, errors = built_and_gated("remove_hidden_text", {"tracked": True}, files)
    assert errors == [] and built.spec["mode"] == "content" and built.spec["ops"][0]["params"]["tracked"] is True
    built, _, errors = built_and_gated("find_replace", {"pairs": [{"find": "sprinkler", "replace": "x"}],
                                                        "tracked": True}, files)
    assert errors == [] and built.spec["mode"] == "content"


# ---- the font button: each document at its own body size ------------------------------------------------------------

def font_fixtures(tmp_path):
    """Two documents whose body text differs in size, built differently (two structure patterns): ``nine`` (Century
    Gothic 9 pt body, 8 pt headers and notes, a 12 pt note) and ``eleven`` (Arial 11 pt body, its revision table set
    in 9 pt); the template's body is Calibri 11 pt."""
    nine = tmp_path / "nine.docx"
    nine.write_bytes(build_sized())
    eleven = tmp_path / "eleven.docx"
    eleven.write_bytes(build("21 13 16", "DRY-PIPE SPRINKLER SYSTEMS", body_size=11, table_size=9))
    template = tmp_path / "ref.docx"
    template.write_bytes(build_reference_spec(body_font="Calibri", body_size=11))
    return nine, eleven, template


def pattern(path: Path) -> str:
    return fingerprint(DocxPackage.from_path(path))["pattern_id"]


def no_ai(monkeypatch):
    def never(*a, **k):
        raise AssertionError("a Safe-mode button makes no AI call")
    monkeypatch.setattr(vision, "review_samples", never)
    monkeypatch.setattr(llm, "client", never)
    monkeypatch.setattr(preview, "render_samples", lambda *a, **k: None)


def run_size(root, text: str):
    """The direct ``w:sz`` of the first run whose text holds ``text``."""
    for r in root.iter(W + "r"):
        if text in "".join(t.text or "" for t in r.findall(W + "t")):
            sz = r.find(f"{W}rPr/{W}sz")
            return sz.get(W + "val") if sz is not None else None
    raise AssertionError(f"no run {text!r}")


def test_the_font_button_plans_each_document_at_its_own_body_size(tmp_path):
    """The stress-test finding: attached beside a 9 pt-body document, an 11 pt-body document had its 9 pt table text
    changed too, because one spec carried every body size of the batch. Each document's plan now carries its own body
    size alone, and the plan on a document is the same whether it is attached alone or with another."""
    nine, eleven, template = font_fixtures(tmp_path)
    alone = {p: recipes.build("font_template", {}, recipes.Context([p], template)) for p in (nine, eleven)}
    for p, size in ((nine, 9.0), (eleven, 11.0)):
        assert not alone[p].gated and alone[p].left_out == {}
        assert [o["params"]["size_from"] for o in alone[p].spec["ops"]] == [[size]] * 3
        assert alone[p].spec["assertions"][0]["size_from"] == [size]
    built = recipes.build("font_template", {}, recipes.Context([nine, eleven], template))
    assert built.gated and built.spec is None and built.left_out == {}
    plan = built.outcome
    assert layouts.is_layout_plan(plan) and len(plan["layouts"]) == 2
    assert plan["request"] == "Match the font to the template ref.docx (Calibri 11 pt)" == built.request
    by_size = {}
    for layout in plan["layouts"]:
        sizes = {tuple(o["params"]["size_from"]) for o in layout["spec"]["ops"]}
        sizes.add(tuple(layout["spec"]["assertions"][0]["size_from"]))
        assert len(sizes) == 1, sizes                            # one body size per layout: every op and the check
        by_size[sizes.pop()[0]] = layout
    assert by_size[9.0]["pattern_ids"] == [pattern(nine)] and by_size[11.0]["pattern_ids"] == [pattern(eleven)]
    assert by_size[9.0]["label"].startswith("body text 9 pt: nine.docx")
    assert "body size (9 pt) changes size" in " ".join(by_size[9.0]["spec"]["assumptions"])
    assert "body size (11 pt) changes size" in " ".join(by_size[11.0]["spec"]["assumptions"])
    for p, size in ((nine, 9.0), (eleven, 11.0)):
        # the layout's spec is the document's own single-document spec, the routing pin aside
        assert {k: v for k, v in by_size[size]["spec"].items() if k != "applies_to"} == alone[p].spec
        # and the plan on the document is the same alone and together (a formatting plan has no text diffs: its
        # change is in its parts, its per-op summaries and its property assertions)
        pkg = DocxPackage.from_path(p)
        together, _ = plan_spec(pkg, layouts.layout_spec(plan, by_size[size]))
        single, _ = plan_spec(pkg, alone[p].spec)
        planned = [(x.intended_parts, x.ops, x.property_assertions, x.diffs, x.members_added, x.members_removed,
                    x.rev_delta, x.expected_structural_delta) for x in (together, single)]
        assert planned[0] == planned[1]
        assert set(single.intended_parts) >= ({pkg.main_part()} if p == nine else set())   # nine's 9 pt runs change


def test_a_mixed_batch_runs_each_document_at_its_own_size_and_the_other_sizes_keep_theirs(tmp_path, monkeypatch):
    """``Session.run`` through the gate, no AI call: both documents are delivered and verified; the 11 pt-body
    document's 9 pt table text is still 9 pt (the old batch-wide plan took it to 11 pt), its body text is in the
    template's font; the 9 pt document's body goes to 11 pt while its 8 pt note style, its 8 pt header run and its 12
    pt note keep their sizes."""
    no_ai(monkeypatch)
    nine, eleven, template = font_fixtures(tmp_path)
    task = DocumentTask(Session(out_root=tmp_path / "out"), render=True, mode="safe")   # key "present" (conftest)
    task.attach([nine, eleven])
    task.attach([template], template=True)
    before = {p: p.read_bytes() for p in (nine, eleven)}
    view = task.perform("recipe", recipe="font_template", inputs={})
    assert view["stage"] == "preview" and view["set_aside"] == [] and task.session.model_free
    labels = [g["label"] for g in task._plan_view()["groups"]]
    assert labels[0].startswith("body text 9 pt") and labels[1].startswith("body text 11 pt")
    view = task.perform("primary", token=view["token"])
    assert view["stage"] == "done" and view["left_untouched"] == [] and len(task.results) == 2
    for record in task.session.last_batch.files:
        assert record.status == "PASS"
    assert {p: p.read_bytes() for p in before} == before
    out = {Path(src).name: DocxPackage.from_path(Path(copy)) for src, copy in task.results.items()}
    main = out["eleven.docx"].xml(out["eleven.docx"].main_part())
    table_sizes = [sz.get(W + "val") for tbl in main.iter(W + "tbl") for sz in tbl.iter(W + "sz")]
    assert table_sizes and set(table_sizes) == {"18"}                  # the 9 pt table text is still 9 pt
    normal = StyleSheet(out["eleven.docx"]).styles["Normal"]
    assert normal.find(f"{W}rPr/{W}rFonts").get(W + "ascii") == "Calibri"
    assert normal.find(f"{W}rPr/{W}sz").get(W + "val") == "22"
    sheet = StyleSheet(out["nine.docx"])
    assert sheet.styles["Normal"].find(f"{W}rPr/{W}sz").get(W + "val") == "22"
    assert sheet.styles["Normal"].find(f"{W}rPr/{W}rFonts").get(W + "ascii") == "Calibri"
    assert sheet.styles["SpecNote"].find(f"{W}rPr/{W}sz").get(W + "val") == "16"
    assert sheet.root.find(f"{W}docDefaults/{W}rPrDefault/{W}rPr/{W}sz").get(W + "val") == "22"
    main = out["nine.docx"].xml(out["nine.docx"].main_part())
    assert run_size(main, "Text pasted at 9 pt.") == "22" and run_size(main, "note deliberately 12 pt.") == "24"
    header = next(part for part in out["nine.docx"].text_parts() if "header" in part)
    assert run_size(out["nine.docx"].xml(header), "SECTION 21 13 13") == "16"


def test_documents_built_alike_at_two_body_sizes_take_the_size_most_of_them_have(tmp_path, monkeypatch):
    """A layout covers structure patterns, so two documents of one pattern cannot take two plans in one press: the
    size most of them have wins (the first attached decides a tie); the others wait under Left untouched with the
    reason, to be pressed alone — never planned at another document's body size."""
    no_ai(monkeypatch)
    a, b, c = (tmp_path / name for name in ("a.docx", "b.docx", "c.docx"))
    a.write_bytes(build_reference_spec("08 71 00", "DOOR HARDWARE", body_size=11))
    b.write_bytes(build_reference_spec("08 71 10", "DOOR HARDWARE SETS", body_size=9))
    c.write_bytes(build_reference_spec("08 71 20", "DOOR HARDWARE KEYING", body_size=9))
    template = tmp_path / "ref.docx"
    template.write_bytes(build_reference_spec(body_font="Calibri", body_size=11))
    assert pattern(a) == pattern(b) == pattern(c)
    built = recipes.build("font_template", {}, recipes.Context([a, b], template))
    assert not built.gated and built.spec["ops"][0]["params"]["size_from"] == [11.0]
    assert list(built.left_out) == [str(b.resolve())]
    reason = built.left_out[str(b.resolve())]
    assert reason.startswith("Not matched with the others: it is built like a.docx, whose body text is 11 pt where "
                             "its own is 9 pt")
    assert "Press the button with it alone" in reason
    built = recipes.build("font_template", {}, recipes.Context([a, b, c], template))
    assert built.spec["ops"][0]["params"]["size_from"] == [9.0] and list(built.left_out) == [str(a.resolve())]
    # through the window: b waits beside the preview with the reason; the run delivers a alone and b is untouched
    before = b.read_bytes()
    task = DocumentTask(Session(out_root=tmp_path / "out", visual_check=False), render=False, mode="safe")
    task.attach([a, b])
    task.attach([template], template=True)
    view = task.perform("recipe", recipe="font_template", inputs={})
    assert view["stage"] == "preview" and [(i["name"], i["reason"]) for i in view["set_aside"]] == [("b.docx", reason)]
    task.perform("primary", token=view["token"])
    view = task.view()
    assert list(task.results) == [str(a)] and b.read_bytes() == before
    assert [(e["name"], e["reason"]) for e in view["left_untouched"]] == [("b.docx", reason)]
    for record in task.session.last_batch.files:
        assert record.status == "PASS"


def test_a_document_with_no_measurable_body_size_takes_the_font_alone(tmp_path):
    nine, _, template = font_fixtures(tmp_path)
    empty = tmp_path / "empty.docx"
    empty.write_bytes(build_dirty("empty_body"))
    built = recipes.build("font_template", {}, recipes.Context([nine, empty], template))
    assert built.gated and built.left_out == {}
    plain = next(lay for lay in built.outcome["layouts"] if lay["label"].startswith("no one body size: empty.docx"))
    assert all("size" not in o["params"] and "size_from" not in o["params"] for o in plain["spec"]["ops"])
    assert "size" not in plain["spec"]["assertions"][0]
    sized = next(lay for lay in built.outcome["layouts"] if lay is not plain)
    assert sized["spec"]["ops"][0]["params"]["size_from"] == [9.0]
    alone = recipes.build("font_template", {}, recipes.Context([empty], template))     # the same font-only spec
    assert {k: v for k, v in plain["spec"].items() if k not in ("applies_to", "request")} == \
           {k: v for k, v in alone.spec.items() if k != "request"}


def test_a_template_with_no_one_body_font_is_refused(tmp_path, folder):
    doc = Document()
    for i, (font, size) in enumerate((("Arial", 10), ("Times New Roman", 12), ("Calibri", 11))):
        run = doc.add_paragraph().add_run(f"Body text in its own font number {i}")
        run.font.name, run.font.size = font, Pt(size)
    doc.save(tmp_path / "mixed.docx")
    with pytest.raises(recipes.RecipeError, match="no one body font"):
        recipes.build("font_template", {}, recipes.Context(docs(folder), tmp_path / "mixed.docx"))


def test_files_the_form_names_must_be_attached(folder, tmp_path):
    files = docs(folder)
    picture = tmp_path / "logo.png"
    picture.write_bytes(b"\x89PNG\r\n\x1a\n")
    with pytest.raises(recipes.RecipeError, match="files you attached"):
        recipes.build("swap_logo", {"image": str(tmp_path / "other.png")}, recipes.Context(files, None, [picture]))
    with pytest.raises(recipes.RecipeError, match="must be one of the files you attached"):
        recipes.build("find_replace", {"table": str(tmp_path / "pairs.csv")}, recipes.Context(files, None, []))


# ---- Safe mode in the task ------------------------------------------------------------------------------------------

def test_safe_mode_has_no_conversation_and_the_buttons_are_safe_modes(folder, tmp_path):
    task = DocumentTask(Session(out_root=tmp_path / "out", visual_check=False), render=False, mode="safe")
    task.attach(docs(folder))
    with pytest.raises(ValueError, match="Safe mode has no conversation"):
        task.perform("send", text="Make the fonts Arial")
    assert SAFE_NO_CONVERSATION
    task.set_mode("default")
    with pytest.raises(ValueError, match="Safe mode's"):
        task.perform("recipe", recipe="remove_highlight", inputs={})
    with pytest.raises(ValueError):
        task.set_mode("obliterate")


def test_changed_attachments_void_a_button_preview_and_say_so(folder, tmp_path):
    task = DocumentTask(Session(out_root=tmp_path / "out", visual_check=False), render=False, mode="safe")
    task.attach(docs(folder)[:1])
    view = task.perform("recipe", recipe="remove_highlight", inputs={})
    assert view["stage"] == "preview"
    task.attach(docs(folder)[1:])
    assert task.stage == "empty" and task.decisions == [] and task.session.spec is None
    assert "Press the button again" in task.messages[-1]["text"]
    assert task.view()["action"] == ""                       # nothing to continue without the conversation


def test_what_a_button_left_untouched_can_be_handed_to_default_mode(tmp_path):
    good = tmp_path / "21 13 13.docx"
    good.write_bytes(build("21 13 13", "WET-PIPE SPRINKLER SYSTEMS"))
    empty = tmp_path / "empty.docx"
    empty.write_bytes(build_dirty("empty_body"))
    task = DocumentTask(Session(out_root=tmp_path / "out", visual_check=False), render=False, mode="safe")
    task.attach([good, empty])
    view = task.perform("recipe", recipe="title_property", inputs={"pattern": "SECTION {number} - {title}"})
    with pytest.raises(ValueError, match="before handing"):
        task.hand_off([str(empty)])                          # a preview waits: create the copies first
    task.perform("primary", token=view["token"])
    view = task.view()
    assert [e["name"] for e in view["left_untouched"]] == ["empty.docx"]
    assert view["todo"] == []                                # a button's refusals stay in scope until handed over
    with pytest.raises(ValueError, match="not the ones"):
        task.hand_off([str(good)])
    task.hand_off([str(empty.resolve())])
    view = task.view()
    assert task.mode == "default" and view["mode"] == "default" and view["left_untouched"] == []
    assert [t["name"] for t in view["todo"]] == ["empty.docx"]
    entry = task.todo[str(empty.resolve())]
    assert entry["decisions"] == ["Set the document title property to “SECTION {number} - {title}”"]
    assert entry["why"] == "refused" and entry["reason"].startswith("The button could not change it: ")
    assert view["todo"][0]["reason"] == entry["reason"]      # the to-do card says why, as it does for a conversation's
    assert WORK_TODO in task.messages[-1]["text"]


def test_setting_every_document_aside_in_safe_mode_creates_nothing_and_starts_no_conversation(folder, tmp_path):
    task = DocumentTask(Session(out_root=tmp_path / "out", visual_check=False), render=False, mode="safe")
    files = docs(folder)
    task.attach(files)
    view = task.perform("recipe", recipe="remove_highlight", inputs={})
    task.perform("primary", token=view["token"], hold=[str(f) for f in files])
    view = task.view()
    assert task.mode == "safe" and task.stage == "empty" and task.decisions == [] and task.focus is None
    assert sorted(t["name"] for t in view["todo"]) == sorted(f.name for f in files)
    assert not task.results and "no copy was created" in task.messages[-1]["text"]


def test_a_removed_document_leaves_the_untouched_list_and_cannot_be_handed_over(tmp_path):
    good = tmp_path / "21 13 13.docx"
    good.write_bytes(build("21 13 13", "WET-PIPE SPRINKLER SYSTEMS"))
    empty = tmp_path / "empty.docx"
    empty.write_bytes(build_dirty("empty_body"))
    task = DocumentTask(Session(out_root=tmp_path / "out", visual_check=False), render=False, mode="safe")
    task.attach([good, empty])
    view = task.perform("recipe", recipe="title_property", inputs={"pattern": "SECTION {number} - {title}"})
    task.perform("primary", token=view["token"])
    assert [e["name"] for e in task.view()["left_untouched"]] == ["empty.docx"]
    stale = task.view()["left_untouched"][0]["path"]
    task.remove(empty)
    assert task.view()["left_untouched"] == []
    task.left_untouched = [{"path": stale, "name": "empty.docx", "why": "refused", "reason": "x", "request": "r"}]
    with pytest.raises(ValueError, match="not the ones"):    # a page out of date names a document no longer attached
        task.hand_off([stale])
    assert task.todo == {} and task.mode == "safe"


def test_a_default_plan_approved_in_safe_mode_takes_no_ai_step(folder, tmp_path, monkeypatch):
    task = DocumentTask(Session(out_root=tmp_path / "out", visual_check=False), render=False, mode="default")
    task.attach(docs(folder))
    _built, outcome, errors = built_and_gated("remove_highlight", {}, task.documents())
    assert not errors
    s = task.session
    task.decisions = ["Remove the highlighting"]
    s._set_outcome(spec=outcome, source="conversation")       # a plan the conversation made, previewed in Default
    task._settle(quiet)
    assert task.stage == "preview" and not s.model_free
    task.set_mode("safe")                                    # the preview stays; Safe mode makes no AI call

    def never(*a, **k):
        raise AssertionError("an AI step in Safe mode")
    monkeypatch.setattr(s, "resolve_remaining", never)
    monkeypatch.setattr(s, "resolve_next_layout", never)
    real_run = s.run

    def run_one_failing(*a, **k):
        real_run(*a, **k)
        s.last_batch.files[0].exception = {"code": "planted", "reason": "a planted failure"}   # Default recovers it
    monkeypatch.setattr(s, "run", run_one_failing)
    task.perform("primary", token=task.view()["token"])
    view = task.view()
    assert task.stage == "done" and task.todo == {}
    assert [(e["name"], e["reason"]) for e in view["left_untouched"]] == [(Path(s.last_batch.files[0].source).name,
                                                                           "a planted failure")]


def test_a_button_no_document_can_take_previews_nothing_to_create(tmp_path):
    empty = tmp_path / "empty.docx"
    empty.write_bytes(build_dirty("empty_body"))
    task = DocumentTask(Session(out_root=tmp_path / "out", visual_check=False), render=False, mode="safe")
    task.attach([empty])
    view = task.perform("recipe", recipe="name_files", inputs={})
    assert view["stage"] == "preview" and view["action"] == "" and view["set_aside"]
    assert "None of the documents can take this change" in task.messages[-1]["text"]


def test_a_dirty_rich_file_refused_by_the_template_is_left_untouched_not_guessed(tmp_path):
    rich = tmp_path / "rich.docx"
    rich.write_bytes(build_rich())
    plain = tmp_path / "plain.docx"
    plain.write_bytes(build("21 13 16", "DRY-PIPE SPRINKLER SYSTEMS"))
    template = tmp_path / "tpl.docx"
    template.write_bytes(build_template())
    task = DocumentTask(Session(out_root=tmp_path / "out", visual_check=False), render=False, mode="safe")
    task.attach([rich, plain])
    task.attach([template], template=True)
    view = task.perform("recipe", recipe="headers_footers_template", inputs={"slots": {"ISSUED FOR PERMIT": "BID"}})
    reasons = {i["name"]: i["reason"] for i in view["set_aside"]}
    assert "rich.docx" in reasons and "odd" in reasons["rich.docx"].lower()   # the even/odd pages setting differs


# ---- the mode is remembered; Obliterate never is -------------------------------------------------------------------

def test_the_window_opens_in_the_mode_last_used_and_safe_mode_first(tmp_path):
    assert Settings().mode == "safe" and MODES == ("safe", "default")
    path = tmp_path / "settings.json"
    save_settings(Settings(mode="default"), path)
    assert load_settings(path).mode == "default"
    path.write_text('{"mode": "obliterate"}', encoding="utf-8")
    assert load_settings(path).mode == "safe"


def test_the_mock_servers_buttons_are_the_apps():
    """frontend/mock/recipes.json is what the page is developed and tested against; it is the catalog as the app
    serves it (regenerate: json.dumps({"recipes": recipes.catalog(), "groups": list(recipes.GROUPS)}, indent=1))."""
    import json
    mock = Path(__file__).resolve().parent.parent / "frontend" / "mock" / "recipes.json"
    assert json.loads(mock.read_text(encoding="utf-8")) == {"recipes": recipes.catalog(), "groups": list(recipes.GROUPS)}


def test_a_sample_a_later_step_refused_is_left_untouched_beside_a_buttons_copies(tmp_path, monkeypatch):
    # A button's spec passes the gate that turns overlapping operations into ordered steps, so a later step can refuse a
    # sample. Create does not run that document (Session._without_refused)
    # and Left untouched names it with the step's reason, never a document dropped without a word.
    a, b = tmp_path / "21 13 13.docx", tmp_path / "21 13 16.docx"
    a.write_bytes(build("21 13 13", "WET-PIPE SPRINKLER SYSTEMS"))
    b.write_bytes(build("21 13 16", "DRY-PIPE SPRINKLER SYSTEMS"))
    before = b.read_bytes()
    real = preview.make_preview

    def later_step_refuses(*args, **kwargs):
        pv = real(*args, **kwargs)
        refused = next(sv for sv in pv.samples if Path(sv.source).name == b.name)
        refused.passed = None
        refused.exception = {"code": "numbering-unsupported", "kind": "refusal",
                             "reason": "stage 2 (fonts): a list defined through a numbering style"}
        return pv
    monkeypatch.setattr(preview, "make_preview", later_step_refuses)
    task = DocumentTask(Session(out_root=tmp_path / "out", visual_check=False), render=False, mode="safe")
    task.attach([a, b])
    view = task.perform("recipe", recipe="title_property", inputs={"pattern": "SECTION {number} - {title}"})
    assert view["stage"] == "preview" and [e["name"] for e in view["set_aside"]] == [b.name]
    task.perform("primary", token=view["token"])
    view = task.view()
    assert list(task.results) == [str(a)] and b.read_bytes() == before
    assert [(e["name"], e["reason"]) for e in view["left_untouched"]] == [
        (b.name, "a list defined through a numbering style")]


# ---- the third set (2026-10-28): each new button's own rules --------------------------------------------------------

def two_specs(tmp_path):
    files = []
    for num, title in (("21 13 13", "WET-PIPE SPRINKLER SYSTEMS"), ("21 13 16", "DRY-PIPE SPRINKLER SYSTEMS")):
        path = tmp_path / f"{num}.docx"
        path.write_bytes(build(num, title))
        files.append(path)
    return files


def created(tmp_path, monkeypatch, files, recipe_id, inputs):
    """The button pressed, previewed and created with no AI call; each copy by its source's name."""
    no_ai(monkeypatch)
    task = DocumentTask(Session(out_root=tmp_path / "out", visual_check=False), render=False, mode="safe")
    task.attach(files)
    view = task.perform("recipe", recipe=recipe_id, inputs=inputs)
    assert view["stage"] == "preview", [m["text"] for m in task.messages]
    view = task.perform("primary", token=view["token"])
    assert view["stage"] == "done"
    for record in task.session.last_batch.files:
        assert record.status == "PASS"
    return task, {Path(src).name: DocxPackage.from_path(out) for src, out in task.results.items()}


def shown_text(pkg, part):
    """The text a part shows: its w:t, never a pending deletion's w:delText."""
    return "".join(t.text or "" for t in pkg.xml(part).iter(W + "t"))


def body_paragraphs(pkg):
    body = pkg.xml(pkg.main_part()).find(W + "body")
    return [p for p in body if p.tag == W + "p"]


def paragraph_text(p):
    return "".join(t.text or "" for t in p.iter(W + "t"))


def simple_doc(path, lines, landscape=False):
    from docx.enum.section import WD_ORIENT
    doc = Document()
    for line in lines:
        doc.add_paragraph(line)
    if landscape:
        sec = doc.sections[0]
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = sec.page_height, sec.page_width
    doc.save(path)
    return path


def test_the_margins_sit_side_by_side_and_a_list_field_is_one_the_panel_renders():
    fields = {f["name"]: f for f in recipes.BY_ID["page_setup"].to_dict()["fields"]}
    assert [n for n, f in fields.items() if f.get("row") == "margins"] == [
        "margin_top", "margin_bottom", "margin_left", "margin_right"]
    assert all(fields[n]["default"] == "1" for n in ("margin_top", "margin_bottom", "margin_left", "margin_right"))
    current = recipes.BY_ID["issue_stamp"].to_dict()["fields"][0]
    assert current["kind"] == "lines" and current["required"] is True


def test_the_issue_stamp_is_replaced_wherever_it_reads_as_typed(tmp_path, monkeypatch):
    files = two_specs(tmp_path)
    old, new = "ISSUED FOR CONSTRUCTION  2026-01-01", "ISSUED FOR BID 2026-11-20"
    built, _, errors = built_and_gated("issue_stamp", {"current": [old, "ISSUED FOR PERMIT"], "text": new}, files)
    assert errors == [] and built.sources == [new] and built.spec["mode"] == "untracked"
    assert built.spec["ops"][0]["params"]["any_of"] == [old, "ISSUED FOR PERMIT"]
    _, copies = created(tmp_path, monkeypatch, files, "issue_stamp", {"current": [old], "text": new})
    for pkg in copies.values():
        headers = [p for p in pkg.text_parts() if "header" in p]
        assert any(new in shown_text(pkg, p) for p in headers)
        assert not any(old in shown_text(pkg, p) for p in pkg.text_parts())
    for inputs, said in (({"current": [], "text": new}, "Type the stamp as it reads now"),
                         ({"current": ["  "], "text": new}, "Type the stamp as it reads now"),
                         ({"current": [old], "text": ""}, "Type the new stamp"),
                         ({"current": [new], "text": new}, "already reads as the new stamp"),
                         ({"current": ["issued for bid"], "text": "ISSUED FOR BID"}, "changes only capitals"),
                         ({"current": [old, old], "text": new}, "listed twice")):
        with pytest.raises(recipes.RecipeError, match=said):
            recipes.build("issue_stamp", inputs, recipes.Context(files))


def test_end_of_section_ends_each_document_with_its_own_number_like_its_section_line(tmp_path, monkeypatch):
    files = two_specs(tmp_path)
    built, _, errors = built_and_gated("end_of_section", {"text": "END OF SECTION {number}"}, files)
    params = built.spec["ops"][0]["params"]
    assert errors == [] and built.sources == ["END OF SECTION {number}"]
    assert params["formatting"] == "copy" and params["number"] == recipes.REFERENCE_FILLS["SECTION_NUMBER"]
    assert not re.search(params["copy_from_regex"], "Section 230513 Common Motor Requirements")   # a reference
    assert re.search(params["copy_from_regex"], "SECTION 21 13 13")
    _, copies = created(tmp_path, monkeypatch, files, "end_of_section", {"text": "END OF SECTION {number}"})
    for name, pkg in copies.items():
        pars = [p for p in body_paragraphs(pkg) if paragraph_text(p).strip()]
        assert paragraph_text(pars[-1]) == f"END OF SECTION {Path(name).stem}"
        section_line = next(p for p in pars if paragraph_text(p).startswith("SECTION "))
        ppr = lambda p: etree.tostring(p.find(W + "pPr")) if p.find(W + "pPr") is not None else None  # noqa: E731
        assert ppr(pars[-1]) == ppr(section_line)
        # pressed again on the copy: the line is recognized and nothing more is added
        again, _ = plan_spec(pkg, built.spec)
        assert [a["expected_case"] for a in again.property_assertions if a["kind"] == "end_of_section"] == ["none"]
    # the line, in any capitals, is the one the operation recognizes by its default: a second press changes nothing
    from spec_surgeon.ops.end_of_section import DEFAULT_DETECT
    assert re.search(DEFAULT_DETECT, "End of Section 21 13 13") and recipes.END_LINE.match("end of section {number}")
    plain, _, errors = built_and_gated("end_of_section", {"text": "END OF SECTION {number}", "look": "plain"}, files)
    assert errors == [] and "formatting" not in plain.spec["ops"][0]["params"]
    # a line in other words would not be recognized in a document that already ends with it: a second one each press
    for text, said in (("END OF SECTION", "needs {number} once"), ("{number} END {number}", "needs {number} once"),
                       ("END OF {title} {number}", "Only {number}"), ("END OF {number}", "starts with END OF SECTION"),
                       ("SECTION {number} ENDS", "starts with END OF SECTION")):
        with pytest.raises(recipes.RecipeError, match=re.escape(said)):
            recipes.build("end_of_section", {"text": text}, recipes.Context(files))


def test_page_size_and_margins_leave_a_landscape_document_untouched_unless_only_margins_change(tmp_path, monkeypatch):
    files = two_specs(tmp_path)
    wide = simple_doc(tmp_path / "21 05 00 schedule.docx", ["SECTION 21 05 00", "SCHEDULE"], landscape=True)
    inputs = {"paper": "letter", "margin_top": "1", "margin_bottom": "1", "margin_left": "1.25",
              "margin_right": '1.25"', "header_footer": "0.5 in"}
    built = recipes.build("page_setup", inputs, recipes.Context([*files, wide]))
    params = built.spec["ops"][0]["params"]
    assert params == {"size": {"w": 12240, "h": 15840},
                      "margins": {"top": 1440, "bottom": 1440, "left": 1800, "right": 1800, "header": 720,
                                  "footer": 720}}
    assert list(built.left_out) == [str(wide.resolve())] and "landscape" in built.left_out[str(wide.resolve())]
    assert built.request == ("Set the paper to Letter (8.5 × 11 in) and the margins to 1 in top, 1 in bottom, 1.25 in "
                             "left and 1.25 in right")
    no_ai(monkeypatch)
    task = DocumentTask(Session(out_root=tmp_path / "out", visual_check=False), render=False, mode="safe")
    task.attach([*files, wide])
    view = task.perform("recipe", recipe="page_setup", inputs=inputs)
    view = task.perform("primary", token=view["token"])
    assert view["stage"] == "done" and len(task.results) == 2
    assert [e["name"] for e in view["left_untouched"]] == [wide.name]
    for out in task.results.values():
        pkg = DocxPackage.from_path(out)
        for sp in DocxPackage.sect_prs(pkg.xml(pkg.main_part())):
            assert sp.find(W + "pgSz").get(W + "w") == "12240" and sp.find(W + "pgMar").get(W + "left") == "1800"
            assert sp.find(W + "pgMar").get(W + "header") == "720"
    # margins alone go on every document, landscape pages included
    margins_only = recipes.build("page_setup", {**inputs, "paper": "keep", "header_footer": ""},
                                 recipes.Context([*files, wide]))
    assert margins_only.left_out == {} and "size" not in margins_only.spec["ops"][0]["params"]
    assert "header" not in margins_only.spec["ops"][0]["params"]["margins"]
    # every document landscape: nothing to preview, each one's reason
    alone = recipes.build("page_setup", inputs, recipes.Context([wide]))
    assert alone.gated and alone.outcome is None and list(alone.left_out) == [str(wide.resolve())]
    for value, said in (("one", "number of inches"), ("30", "more than 22 inches"), ("", "Type the top margin")):
        with pytest.raises(recipes.RecipeError, match=said):
            recipes.build("page_setup", {**inputs, "margin_top": value}, recipes.Context(files))


def test_change_font_needs_no_template_and_sizes_by_each_documents_own_body_size(tmp_path):
    nine, eleven, _ = font_fixtures(tmp_path)
    keep = recipes.build("set_font", {"font": "Arial", "size": "keep"}, recipes.Context([nine, eleven]))
    assert not keep.gated and keep.request == "Set the font to Arial"
    assert all("size" not in o["params"] for o in keep.spec["ops"])
    assert keep.spec["assumptions"][0] == "the font you chose is Arial"
    sized = recipes.build("set_font", {"font": "Arial", "size": "10"}, recipes.Context([nine, eleven]))
    assert sized.gated and sized.request == "Set the font to Arial 10 pt"
    plan = sized.outcome
    assert len(plan["layouts"]) == 2 and plan["assumptions"][0] == "the font you chose is Arial 10 pt"
    assert sorted(o["params"]["size_from"][0] for layout in plan["layouts"]
                  for o in layout["spec"]["ops"][:1]) == [9.0, 11.0]
    one = recipes.build("set_font", {"font": "Arial", "size": "10"}, recipes.Context([eleven]))
    assert [o["params"]["size_from"] for o in one.spec["ops"]] == [[11.0]] * 3
    for font, said in (("", "Type the font"), ("Arial; Wingdings", "a font's name"), ("A" * 40, "longer than 31")):
        with pytest.raises(recipes.RecipeError, match=said):
            recipes.build("set_font", {"font": font, "size": "10"}, recipes.Context([eleven]))


def test_reject_all_tracked_changes_takes_the_insertions_out_and_brings_the_deletions_back(tmp_path, monkeypatch):
    files = two_specs(tmp_path)
    built, _, errors = built_and_gated("accept_changes", {"action": "reject"}, files)
    assert errors == [] and built.request == "Reject all tracked changes"
    assert built.spec["ops"][0]["params"] == {"action": "reject"}
    assert recipes.build("accept_changes", {}, recipes.Context(files)).request == "Accept all tracked changes"
    _, copies = created(tmp_path, monkeypatch, files, "accept_changes", {"action": "reject"})
    for pkg in copies.values():
        body = shown_text(pkg, pkg.main_part())
        assert "black steel Schedule 40" in body and "UL listed" not in body
        assert next(pkg.xml(pkg.main_part()).iter(W + "ins"), None) is None


def test_numbers_keep_their_units_on_one_line_and_only_unit_symbols_are_taken(tmp_path, monkeypatch):
    doc = simple_doc(tmp_path / "21 13 13.docx", ["SECTION 21 13 13", "WET-PIPE SPRINKLER SYSTEMS",
                                                  "Test the piping at 200 psi for 2 hours.",
                                                  "Coordinate with Section 21 05 00."])
    _, copies = created(tmp_path, monkeypatch, [doc], "nonbreaking_spaces", {"units": "psi, gpm"})
    body = shown_text(copies[doc.name], copies[doc.name].main_part())
    assert "200 psi" in body and "21 13 13" in body and "21 05 00" in body
    built = recipes.build("nonbreaking_spaces", {"units": " psi, gpm,, °F , psi"}, recipes.Context([doc]))
    assert built.spec["ops"][0]["params"]["units"] == ["psi", "gpm", "°F"]
    assert built.request == "Keep numbers with psi, gpm and °F, and the parts of section numbers (21 13 13), on one line"
    for inputs, said in (({"units": "psi!"}, "not a unit symbol"), ({"units": "", "section_numbers": False},
                                                                    "at least one unit")):
        with pytest.raises(recipes.RecipeError, match=said):
            recipes.build("nonbreaking_spaces", inputs, recipes.Context([doc]))


def test_spaces_quotes_and_dashes_are_cleaned_as_chosen(tmp_path, monkeypatch):
    doc = simple_doc(tmp_path / "21 13 13.docx", ["SECTION 21 13 13", "WET-PIPE SPRINKLER SYSTEMS",
                                                  "The Owner’s  “insurer” -- FM Global.  "])
    _, copies = created(tmp_path, monkeypatch, [doc], "clean_spaces", {})
    assert "The Owner’s “insurer” -- FM Global." in shown_text(copies[doc.name], "word/document.xml")
    _, copies = created(tmp_path / "2", monkeypatch, [doc], "quotes_dashes", {"quotes": "straight", "dash": "en"})
    assert "The Owner's  \"insurer\" – FM Global." in shown_text(copies[doc.name], "word/document.xml")
    with pytest.raises(recipes.RecipeError, match="at least one change"):
        recipes.build("quotes_dashes", {"quotes": "keep"}, recipes.Context([doc]))


def test_fields_freeze_as_plain_text_and_page_numbers_never_do(tmp_path, monkeypatch):
    files = two_specs(tmp_path)
    built, _, errors = built_and_gated("freeze_fields", {"which": "references"}, files)
    assert errors == [] and built.spec["ops"][0]["params"]["types"] == ["REF", "NOTEREF"]
    _, copies = created(tmp_path, monkeypatch, files, "freeze_fields", {})
    for pkg in copies.values():
        codes = " ".join(t.text or "" for part in pkg.text_parts() for t in pkg.xml(part).iter(W + "instrText"))
        assert "PAGE" in codes and "REF" not in codes


def test_author_and_company_are_the_users_words_and_author_names_can_be_cleared(tmp_path, monkeypatch):
    files = two_specs(tmp_path)
    inputs = {"author": "A. Designer", "company": "ACME Engineering", "subject": ""}
    built, _, errors = built_and_gated("author_company", inputs, files)
    assert errors == [] and built.sources == ["A. Designer", "ACME Engineering"]
    assert built.spec["ops"][0]["params"] == {"set": {"author": "A. Designer", "company": "ACME Engineering"}}
    assert built.request == "Set Author to “A. Designer” and Company to “ACME Engineering”"
    _, copies = created(tmp_path, monkeypatch, files, "author_company", inputs)
    for pkg in copies.values():
        assert b"A. Designer" in pkg.read("docProps/core.xml") and b"ACME Engineering" in pkg.read("docProps/app.xml")
    cleared = recipes.build("author_company", {"company": "ACME", "clear_authors": True}, recipes.Context(files))
    assert cleared.spec["ops"][0]["params"]["clear"] == ["author", "last_modified_by"]
    assert cleared.request == "Set Company to “ACME” and clear the author names"
    for inputs, said in (({"author": "A", "clear_authors": True}, "not both"), ({}, "Type an author")):
        with pytest.raises(recipes.RecipeError, match=said):
            recipes.build("author_company", inputs, recipes.Context(files))


def test_the_proofing_language_changes_every_language_or_only_other_english(tmp_path):
    files = two_specs(tmp_path)
    built = recipes.build("proofing_language", {"language": "en-US", "only": "english"}, recipes.Context(files))
    params = built.spec["ops"][0]["params"]
    assert params["language"] == "en-US" and "en-GB" in params["replace_only"] and "en-US" not in params["replace_only"]
    every = recipes.build("proofing_language", {}, recipes.Context(files))
    assert every.spec["ops"][0]["params"] == {"language": "en-US"}
    with pytest.raises(recipes.RecipeError, match="choose one"):
        recipes.build("proofing_language", {"language": "fr-CA", "only": "english"}, recipes.Context(files))


def test_links_and_unused_styles_follow_the_form(tmp_path):
    files = two_specs(tmp_path)
    links = recipes.build("remove_hyperlinks", {"which": "all", "keep_style": True}, recipes.Context(files))
    assert links.spec["ops"][0]["params"] == {"which": "all", "keep_style": True}
    styles = recipes.build("remove_unused_styles", {"keep": "CMT, SpecifierNote, CMT"}, recipes.Context(files))
    assert styles.spec["ops"][0]["params"] == {"which": "custom", "keep": ["CMT", "SpecifierNote"]}
    assert styles.request == "Remove the unused styles, keeping “CMT” and “SpecifierNote”"


def test_the_safe_mode_note_folds_and_the_window_remembers_it(tmp_path, monkeypatch):
    """The note above Safe mode's buttons (what Safe mode is, and Default mode with an API key for more): Hide folds it
    to one line and "What is Safe mode?" opens it; the choice is kept in settings.json and carried by every view."""
    from fastapi.testclient import TestClient
    from spec_surgeon.web.server import TOKEN_HEADER, WindowState, create_app
    from test_web_server import TOKEN, FakeUpdates
    monkeypatch.setenv("SPEC_SURGEON_CONFIG_DIR", str(tmp_path / "config"))
    task = DocumentTask(Session(out_root=tmp_path / "out", visual_check=False), render=False, mode="safe")
    state = WindowState(task, Settings(), token=TOKEN, updates=FakeUpdates(), version="1.6.0")
    client = TestClient(create_app(state), headers={TOKEN_HEADER: TOKEN})
    assert Settings().safe_note_hidden is False
    assert client.get("/api/state").json()["safe_note_hidden"] is False
    assert client.get("/api/settings").json()["safe_note_hidden"] is False
    assert client.post("/api/settings", json={"safe_note_hidden": True}).json()["safe_note_hidden"] is True
    assert load_settings().safe_note_hidden is True and client.get("/api/state").json()["safe_note_hidden"] is True
    assert client.post("/api/settings", json={"safe_note_hidden": "yes"}).status_code == 400
    assert client.post("/api/settings", json={"safe_note_hidden": False}).json()["safe_note_hidden"] is False
    assert load_settings().safe_note_hidden is False


# ---- Remove specifier notes (M2): one press, no form ----------------------------------------------------------------

NOTES_GONE = ("Retain paragraph below if the project has a fire pump.", "Coordinate the alarm with the fire alarm Section.",
              "NOTE TO SPECIFIER: edit the list to suit the project.", "Retain or delete this article.",
              "Select one of the following paragraphs.",
              "[Editor's Note: verify the hazard classification with the Owner.]", "Edit to suit.",
              "Specifier: this whole paragraph is hidden.")
NOTES_STAY = ("SECTION 21 13 13 - WET-PIPE SPRINKLER SYSTEMS",
              "© 2026 Owner Data Centers LLC. All rights reserved. Confidential.", "PART 1 - GENERAL",
              "A.\tProvide a fire pump.", "Coordinate the work with the Owner's insurer.",
              "Retain records of all tests required in Paragraph 1.6.", "ARTICLE 1.2 - SUBMITTALS",
              "Select one of the listed manufacturers.",
              "Provide two [Verify quantity with Owner] spare sprinklers per type.",
              "Provide <Insert manufacturer> valves.", "System pressure: [TBD].",
              "Hydrant flow test date to be determined.", "Provide two [retain or delete] spare heads per cabinet.",
              "Provide galvanized pipe in the preaction zones.", "Pipe shall be Schedule 40.",
              "END OF SECTION 21 13 13")


def spec_notes(tmp_path, *field_notes):
    files = []
    for i, field_note in enumerate(field_notes or ("",)):
        path = tmp_path / f"21 13 1{i}.docx"
        path.write_bytes(build_spec_notes(field_note=field_note))
        files.append(path)
    return files


def notes_op(built):
    return built.spec["stages"][1]["spec"]["ops"][0]


def body_texts(pkg):
    out = []
    for p in body_paragraphs(pkg):
        text = "".join((t.text or "") if t.tag == W + "t" else "\t" for t in p.iter(W + "t", W + "tab"))
        if text:
            out.append(text)
    return out


def test_remove_specifier_notes_takes_the_notes_and_hidden_text_and_leaves_everything_else(tmp_path, monkeypatch):
    files = spec_notes(tmp_path)
    task, out = created(tmp_path, monkeypatch, files, "remove_notes", {})
    body = body_texts(out[files[0].name])
    assert body == list(NOTES_STAY), body                  # every note gone; every other paragraph, in order, as it was
    assert all(n not in " ".join(body) for n in NOTES_GONE) and "retain for 2-inch" not in " ".join(body)
    # clean copies: no tracked change was written
    xml = etree.tostring(out[files[0].name].xml(out[files[0].name].main_part()))
    assert b"<w:del " not in xml and b"<w:ins " not in xml


def test_remove_specifier_notes_has_no_form_and_says_what_it_leaves(folder, tmp_path):
    recipe = recipes.BY_ID["remove_notes"]
    assert recipe.group == "Clean up" and recipe.fields == () and recipe.needs == ()
    for stays in ("placeholders", "copyright lines", "headings"):
        assert stays in recipe.summary
    built, outcome, errors = built_and_gated("remove_notes", {}, spec_notes(tmp_path))
    assert not errors and built.sources == []              # it places no words
    ops = [st["spec"]["ops"][0] for st in built.spec["stages"]]      # two steps, hidden text first
    assert [o["op"] for o in ops] == ["remove_hidden_text", "remove_notes"]
    assert all(o["params"]["tracked"] is False for o in ops)
    assert all(st["spec"]["mode"] == "untracked" for st in built.spec["stages"])
    said = " ".join(built.spec["assumptions"])
    for stays in ("placeholders", "to be determined", "copyright and confidentiality lines", "headings",
                  "coloured or italic text", "tracked changes and comments stay"):
        assert stays in said
    assert "“CMT”, “CMTChild” and “Specifier Note”" in said   # the styles found, by the names Word shows
    assert set(ops[1]["params"]["rules"][0]["style_in"]) == {"CMT", "CMTChild", "SpecifierNote"}


def test_note_styles_are_found_by_name_and_base_and_one_used_otherwise_elsewhere_is_left_out(tmp_path):
    files = spec_notes(tmp_path)
    other = tmp_path / "other.docx"
    doc = Document()
    doc.styles.add_style("CMTChild", 1).base_style = doc.styles["Normal"]   # the same id, not a note here
    doc.add_paragraph("Provide hangers.", style="CMTChild")
    doc.save(other)
    built = recipes.build("remove_notes", {}, recipes.Context(files + [other]))
    assert notes_op(built)["params"]["rules"][0]["style_in"] == ["CMT", "SpecifierNote"]
    assert any("“CMTChild” is a note style in one document and not in another" in a
               for a in built.spec["assumptions"])
    plain = recipes.build("remove_notes", {}, recipes.Context([other]))
    assert notes_op(plain)["params"]["rules"] == [{"text_regex": hidden_ops.SPECIFIER_OPENINGS}]
    assert any("no note style" in a for a in plain.spec["assumptions"])


def test_a_hidden_note_holding_a_field_goes_and_a_visible_one_sets_only_its_document_aside(tmp_path, monkeypatch):
    files = spec_notes(tmp_path, "hidden", "visible")
    no_ai(monkeypatch)
    task = DocumentTask(Session(out_root=tmp_path / "out", visual_check=False), render=False, mode="safe")
    task.attach(files)
    view = task.perform("recipe", recipe="remove_notes", inputs={})
    assert view["stage"] == "preview"
    view = task.perform("primary", token=view["token"])
    assert view["stage"] == "done"
    assert [Path(src).name for src in task.results] == [files[0].name]
    assert [(e["name"], "holds a field" in e["reason"]) for e in view["left_untouched"]] == [(files[1].name, True)]
    for record in task.session.last_batch.files:
        assert record.status == "PASS"
    body = body_texts(DocxPackage.from_path(next(iter(task.results.values()))))
    assert body == list(NOTES_STAY) and "Retain with Section" not in " ".join(body)


@pytest.mark.parametrize("text", NOTES_GONE[:-1])
def test_every_note_the_fixture_holds_is_one_by_its_style_or_its_opening(text):
    from spec_surgeon.ops.patterns import compile_pattern
    from spec_surgeon.verify.patterns import v_search
    by_text = text not in ("Retain paragraph below if the project has a fire pump.",
                           "Coordinate the alarm with the fire alarm Section.", "Edit to suit.")
    found = compile_pattern(hidden_ops.SPECIFIER_OPENINGS, where="t").search(text) is not None
    assert found == by_text or (found and text.startswith("Retain paragraph below"))
    assert (v_search(hidden_ops.SPECIFIER_OPENINGS, text) is not None) == found      # the verifier reads it alike


@pytest.mark.parametrize("text", NOTES_STAY + (
    "*** END OF SECTION ***", "Retain the services of a testing agency.", "Verify dimensions in the field.",
    "Delete the existing valve and cap the pipe.", "Copyright 2024 Owner Inc.", "Select one-piece valves.",
    "Revise as required.", "Note: provide access panels.", "Editorial review is required before issue."))
def test_no_requirement_heading_placeholder_or_copyright_line_opens_as_a_note(text):
    from spec_surgeon.ops.patterns import compile_pattern
    from spec_surgeon.verify.patterns import v_search
    assert compile_pattern(hidden_ops.SPECIFIER_OPENINGS, where="t").search(text) is None
    assert v_search(hidden_ops.SPECIFIER_OPENINGS, text) is None


FURTHER = ("Cut every [Insert …], <Insert …>, [Verify …] and [TBD] placeholder, delete the copyright line, change "
           "“to be determined” to “scheduled for October 2026”, and delete the ARTICLE 1.2 heading.")
FURTHER_SPEC = {
    "request": FURTHER, "summary": "Cut the placeholders, the copyright line and a heading; edit one phrase",
    "mode": "untracked", "assumptions": ["clean copies"],
    "ops": [
        {"op": "replace_table", "params": {"pairs": [{"find": r"\s*[\[<](?:Insert|Verify|TBD)[^\]>]*[\]>]",
                                                      "replace": "", "regex": True, "match_case": False}],
                                           "tracked": False}},
        {"op": "replace_text", "params": {"find": "to be determined", "replace": "scheduled for October 2026",
                                          "tracked": False}},
        {"op": "delete_paragraphs", "select": {"paragraphs": {"text_regex": "©|All rights reserved"}},
         "params": {"tracked": False}},
        {"op": "delete_paragraphs", "select": {"paragraphs": {"text_regex": "^ARTICLE 1\\.2"}},
         "params": {"tracked": False}},
    ]}


@pytest.mark.parametrize("mode", ["default", "obliterate"])
def test_default_and_obliterate_clean_what_the_notes_button_leaves(tmp_path, monkeypatch, mode):
    """The button leaves placeholders, copyright lines, “to be determined” and headings; a Default or Obliterate request
    takes them, through the same gate (the one typed word traced to the request), preview and checked run."""
    from fake_claude import FakeClient, resp, tool
    files = spec_notes(tmp_path)
    no_ai(monkeypatch)
    safe_task = DocumentTask(Session(out_root=tmp_path / "out", visual_check=False), render=False, mode="safe")
    safe_task.attach(files)
    safe_task.perform("primary", token=safe_task.perform("recipe", recipe="remove_notes", inputs={})["token"])
    copy = Path(next(iter(safe_task.results.values())))
    client = FakeClient(resp(tool("submit_editspec", {"spec": FURTHER_SPEC}, "t1")))
    task = DocumentTask(Session(client=client, out_root=tmp_path / "out2", visual_check=False), render=False,
                        obliterate_allowed=True)
    if mode == "obliterate":
        task.set_mode("obliterate", acknowledged=True)
    task.attach([copy])
    view = task.perform("send", text=FURTHER)
    assert view["stage"] == "preview", [m["text"] for m in task.messages][-2:]
    view = task.perform("primary", token=view["token"])
    assert view["stage"] == "done"
    assert all(r.status == "PASS" for r in task.session.last_batch.files)
    body = body_texts(DocxPackage.from_path(next(iter(task.results.values()))))
    assert body == ["SECTION 21 13 13 - WET-PIPE SPRINKLER SYSTEMS", "PART 1 - GENERAL", "A.\tProvide a fire pump.",
                    "Coordinate the work with the Owner's insurer.",
                    "Retain records of all tests required in Paragraph 1.6.",
                    "Select one of the listed manufacturers.", "Provide two spare sprinklers per type.",
                    "Provide valves.", "System pressure:.", "Hydrant flow test date scheduled for October 2026.",
                    "Provide two [retain or delete] spare heads per cabinet.",
                    "Provide galvanized pipe in the preaction zones.", "Pipe shall be Schedule 40.",
                    "END OF SECTION 21 13 13"], body
