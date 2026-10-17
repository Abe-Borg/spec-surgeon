"""Approved workspace effects reach existing operations without weakening final proof."""
import copy
import hashlib

import pytest
from docx import Document

from fake_claude import FakeClient, resp, tool
from make_fixtures import SCHEME_LEVELS, build_numbering, build_scheme_example, build_typed
from spec_surgeon.campaign import create_campaign, load_campaign
from spec_surgeon.campaign_planner import inspect_goal_effect
from spec_surgeon.goals import (GoalError, NeedsInput, choose_goal, compile_goal_candidate,
                                goal_catalog, goal_templates, resolve_goal, validate_contract, verify_goal)
from spec_surgeon.recovery import RecoveryOptions
from test_goal_contracts import contract, execute


def source_doc(path, count=2):
    doc = Document()
    doc.add_paragraph("SECTION 21 13 13")
    for i in range(count):
        doc.add_paragraph(f"REQUIREMENT {i}")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "SCHEDULE"
    table.cell(1, 0).text = "Approved dimensions remain exact."
    doc.save(path)
    return path


def effect(kind, params, select=None, count="all", **extra):
    return {"id": kind, "type": kind, "params": params, "select": select or {}, "expected_count": count, **extra}


@pytest.mark.parametrize("kind,params,select", [
    ("paragraph_style", {"style": "Heading 1", "tracked": False}, {"paragraphs": {"text_regex": "^REQUIREMENT"}}),
    ("style_mapping", {"map": [{"from": "Normal", "to": "Body Text"}]}, {"scope": ["document"]}),
    ("proofing_language", {"language": "en-US"}, {}),
    ("table_format", {"cant_split": True, "repeat_header_rows": 1}, {}),
    ("table_style", {"style": "Light Shading Accent 1"}, {}),
])
def test_new_workspace_families_execute_and_verify_original_to_final(tmp_path, kind, params, select):
    source = source_doc(tmp_path / "spec.docx")
    goal = contract(effect(kind, params, select))
    binding = resolve_goal(source, goal)["binding"]
    assert binding["effects"][kind]["source_count"] > 0
    execute(source, goal, binding)
    assert kind in goal_catalog()["effects"]


@pytest.mark.parametrize("kind,builder", [("numbering", lambda: build_numbering(shared_restart=False)),
                                           ("typed_numbering", build_typed)])
def test_numbering_contract_pins_template_and_has_final_outcome_proof(tmp_path, kind, builder):
    source, template = tmp_path / "spec.docx", tmp_path / "scheme.docx"
    source.write_bytes(builder())
    template.write_bytes(build_scheme_example())
    pin = hashlib.sha256(template.read_bytes()).hexdigest()
    goal = contract(effect(kind, {"levels": list(SCHEME_LEVELS)}, template=str(template), template_sha256=pin))
    assert goal_templates(goal) == [template]
    binding = resolve_goal(source, goal)["binding"]
    assert binding["effects"][kind]["source_count"] > 1
    output = execute(source, goal, binding)
    template.write_bytes(source.read_bytes())
    with pytest.raises(GoalError, match="template bytes have changed"):
        compile_goal_candidate(source, goal, binding)
    assert not verify_goal(source, output, goal, binding)["passed"]


def test_all_matches_binds_each_original_count_and_cannot_be_forged(tmp_path):
    source = source_doc(tmp_path / "spec.docx", 3)
    goal = contract(effect("paragraph_style", {"style": "Heading 1", "tracked": False},
                           {"paragraphs": {"text_regex": "^REQUIREMENT"}}))
    binding = resolve_goal(source, goal)["binding"]
    assert binding["effects"]["paragraph_style"]["source_count"] == 3
    output = execute(source, goal, binding)
    forged = copy.deepcopy(binding)
    forged["effects"]["paragraph_style"]["source_count"] = 2
    with pytest.raises(GoalError, match="count differs"):
        compile_goal_candidate(source, goal, forged)
    assert not verify_goal(source, output, goal, forged)["passed"]
    changed = copy.deepcopy(goal)
    changed["effects"][0]["expected_count"] = 2
    result = resolve_goal(source, changed)
    assert result["status"] == "needs_input" and result["binding"] is None
    with pytest.raises(NeedsInput):
        compile_goal_candidate(source, changed)
    with pytest.raises(GoalError, match="cannot be changed"):
        choose_goal(source, changed, [], effect_id="paragraph_style")
    source_doc(source, 4)
    with pytest.raises(GoalError, match="source or approved goal changed"):
        compile_goal_candidate(source, goal, binding)


def test_original_target_count_includes_already_satisfied_paragraphs(tmp_path):
    source = source_doc(tmp_path / "spec.docx", 2)
    doc = Document(source)
    doc.paragraphs[1].style = "Heading 1"
    doc.save(source)
    goal = contract(effect("paragraph_style", {"style": "Heading 1", "tracked": False},
                           {"paragraphs": {"text_regex": "^REQUIREMENT"}}, count=2))
    execute(source, goal)


def test_contract_author_is_used_by_tracked_style_effect(tmp_path):
    source = source_doc(tmp_path / "spec.docx", 2)
    goal = contract(effect("paragraph_style", {"style": "Heading 1", "tracked": True},
                           {"paragraphs": {"text_regex": "^REQUIREMENT"}}, count=2))
    goal["author"] = "Approved Reviewer"
    output = execute(source, goal)
    from spec_surgeon.package import DocxPackage, W
    pkg = DocxPackage(output)
    try:
        assert {el.get(W + "author") for el in pkg.xml(pkg.main_part()).iter(W + "pPrChange")} == {"Approved Reviewer"}
    finally:
        pkg.close()


@pytest.mark.parametrize("patch", [
    {"expected_count": -1}, {"expected_count": True}, {"expected_count": "some"},
    {"params": {"style": "Heading 1", "template": "/tmp/unpinned.docx"}},
    {"template": "/tmp/unpinned.docx"}, {"params": {"style": "Heading 1", "invented_option": True}},
    {"select": {"paragraphs": {"text_regex": "["}}},
])
def test_operation_effect_rejects_invalid_or_unpinned_intent(patch):
    proposed = effect("paragraph_style", {"style": "Heading 1"})
    proposed.update(patch)
    with pytest.raises(GoalError):
        validate_contract(contract(proposed))


def test_unverified_compositions_are_refused_before_approval():
    with pytest.raises(GoalError, match="stable paragraph"):
        validate_contract(contract(effect("proofing_language", {"language": "en-US"}, count=1),
            {"id": "words", "type": "replace_text", "target": {"all": True},
             "find": "Owner", "replace": "Client", "expected_matches": 1}))
    with pytest.raises(GoalError, match="separate approved goals"):
        validate_contract(contract(effect("table_style", {"style": "Table Grid"}),
                                   effect("table_format", {"width": {"pct": 100}})))


@pytest.mark.parametrize("other", [
    {"id": "bold", "type": "format", "target": {"all": True}, "run": {"b": True}},
    effect("paragraph_style", {"style": "Heading 1"}),
    effect("style_mapping", {"map": [{"from": "Normal", "to": "Body Text"}]}),
    effect("table_style", {"style": "Table Grid"}),
])
def test_whole_language_property_proof_rejects_overlapping_families_before_approval(other):
    with pytest.raises(GoalError, match="proofing language.*separate approved goals"):
        validate_contract(contract(effect("proofing_language", {"language": "en-US"}), other))


def test_clear_direct_formatting_cannot_be_combined_with_direct_overrides():
    with pytest.raises(GoalError, match="clearing direct formatting"):
        validate_contract(contract(effect("paragraph_style", {"style": "Heading 1", "direct_formatting": "clear"}),
            {"id": "bold", "type": "format", "target": {"all": True}, "run": {"b": True}}))


def test_page_content_selector_and_first_page_false_are_valid_effects(tmp_path):
    source = source_doc(tmp_path / "spec.docx")
    doc = Document(source)
    doc.sections[0].different_first_page_header_footer = True
    doc.save(source)
    goal = contract({"id": "page", "type": "page_setup", "sections": {"with_tables": True},
                     "different_first_page": False})
    execute(source, goal)


def test_planner_count_inspection_is_read_only_scoped_and_source_pinned(tmp_path):
    source = source_doc(tmp_path / "spec.docx", 3)
    before = source.read_bytes()
    document = {"id": "doc1", "source": str(source), "source_sha256": hashlib.sha256(before).hexdigest()}
    proposed = effect("paragraph_style", {"style": "Heading 1"}, {"paragraphs": {"text_regex": "^REQUIREMENT"}})
    result = inspect_goal_effect([document], "doc1", proposed)
    assert result["source_count"] == 3 and source.read_bytes() == before
    assert proposed["expected_count"] == "all"
    with pytest.raises(ValueError, match="inside this project"):
        inspect_goal_effect([document], "foreign", proposed)
    proposed.update(template=str(tmp_path / "foreign-template.docx"), template_sha256="0" * 64)
    with pytest.raises(ValueError, match="explicitly attached"):
        inspect_goal_effect([document], "doc1", proposed)
    source_doc(source, 4)
    with pytest.raises(ValueError, match="source changed"):
        inspect_goal_effect([document], "doc1", effect("proofing_language", {"language": "en-US"}))


def test_mixed_count_workspace_plan_reload_and_execute_without_model(tmp_path):
    folder = tmp_path / "project"
    folder.mkdir()
    first = source_doc(folder / "21 13 13.docx", 2)
    second = source_doc(folder / "21 13 16.docx", 4)
    campaign = create_campaign([str(folder)], tmp_path / "workspace")
    pid = next(iter(campaign.projects))
    goal = contract(effect("paragraph_style", {"style": "Heading 1", "tracked": False},
                           {"paragraphs": {"text_regex": "^REQUIREMENT"}}))
    client = FakeClient(resp(tool("propose_plan", {"contract": goal,
                       "rationale": "Apply the approved heading style to each original REQUIREMENT paragraph."})))
    options = RecoveryOptions(render=False, use_memory=False)
    campaign.plan(pid, "Style the requirement headings", options=options, client=client)
    assert campaign.projects[pid]["proposed_plan"]["status"] == "PROPOSED"
    campaign.approve_plan(pid, campaign.projects[pid]["proposed_plan"]["id"])
    loaded = load_campaign(campaign.out_dir)
    no_model = FakeClient()
    loaded.run(options=options, client=no_model)
    assert not no_model.requests
    assert all(d["status"] == "READY" for d in loaded.documents.values()), loaded.to_text()
    assert first.exists() and second.exists()


def test_planning_tool_inspects_original_count_before_draft_without_editing(tmp_path):
    source = source_doc(tmp_path / "21 13 13.docx", 3)
    before = source.read_bytes()
    campaign = create_campaign([str(source)], tmp_path / "workspace")
    pid, document_id = next(iter(campaign.projects)), next(iter(campaign.documents))
    proposed = effect("paragraph_style", {"style": "Heading 1"}, {"paragraphs": {"text_regex": "^REQUIREMENT"}})
    client = FakeClient(resp(tool("inspect_goal_effect", {"document_id": document_id, "effect": proposed})),
                        resp(tool("propose_plan", {"contract": contract(proposed),
                             "rationale": "Three inspected original requirement headings receive the approved style."})))
    campaign.plan(pid, "Style requirement headings", options=RecoveryOptions(render=False, use_memory=False), client=client)
    assert campaign.projects[pid]["proposed_plan"]["status"] == "PROPOSED"
    assert source.read_bytes() == before and not campaign.projects[pid]["runs"]
    import json
    assert '"source_count": 3' in json.dumps(client.requests[-1]["messages"]).replace('\\"', '"')
