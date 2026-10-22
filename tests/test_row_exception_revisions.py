"""A tracked change of a row's table-property exceptions (``w:tr/w:tblPrEx/w:tblPrExChange``) is a revision like every
other: it carries a ``w:id``, V4 counts it (its before/after delta against the plan's, its id in the duplicate and
collision checks), V5c checks its grammar, and every op that takes one out predicts it in ``rev_delta`` from the source
(``ooxml.PROPERTY_REV_TAGS`` ↔ ``verify/suite.OTHER_REV``). resolve_tracked_changes accepts it (the ``w:tblPrEx`` keeps
the row's exceptions) and refuses to reject it, as it does a table's own property change — both sides
(ops/revisions.py ↔ verify/revisions.py). Each rule has a wrong output and a wrong implementation shared by plan and
apply."""
from __future__ import annotations

import io
import zipfile

import docx
import pytest
from lxml import etree

from spec_surgeon import ooxml
from spec_surgeon.ops import JobException, get_op, merge_fragments
from spec_surgeon.ops import replace_text as replace_module
from spec_surgeon.ops import revisions as rev_module
from spec_surgeon.package import DocxPackage, W, serialize_xml
from spec_surgeon.verify import run_suite

MAIN = "word/document.xml"
STAMP = 'w:author="R" w:date="2026-10-01T09:00:00Z"'
CHANGE = f'<w:tblPrExChange w:id="7" {STAMP}><w:tblPrEx/></w:tblPrExChange>'
CENTRED = '<w:jc w:val="center"/>'
SCHEDULE = {"first_row_regex": "^Tag"}
OLD_PROPERTY_TAGS = ("rPrChange", "pPrChange", "sectPrChange", "tblPrChange", "trPrChange", "tcPrChange",
                     "tblGridChange", "numberingChange")                  # the lists before w:tblPrExChange was counted


def build(exceptions: str = CENTRED + CHANGE, insertion: bool = False) -> bytes:
    """A schedule table whose first row starts with ``<w:tblPrEx>`` + ``exceptions``, then four paragraphs "The Owner
    shall review.". ``insertion``: an earlier paragraph holds Q's tracked insertion (w:id 9) — a second revision for
    resolve_tracked_changes, and an id above the exceptions' change, so the revision-id seed of this branch (which reads
    w:ins but not w:tblPrExChange) starts clear of 7."""
    d = docx.Document()
    t = d.add_table(rows=3, cols=2)
    for r, (a, b) in enumerate((("Tag", "Description"), ("FP-1", "Fire pump"), ("FP-2", "Jockey pump"))):
        t.cell(r, 0).text, t.cell(r, 1).text = a, b
    tr = t._tbl.findall(W + "tr")[0]
    tr.insert(0, etree.fromstring(f'<w:tblPrEx xmlns:w="{W[1:-1]}">{exceptions}</w:tblPrEx>'))
    if insertion:
        p = d.add_paragraph("1.1 SUMMARY")._p
        p.append(etree.fromstring(f'<w:ins xmlns:w="{W[1:-1]}" w:id="9" w:author="Q" w:date="2026-10-02T09:00:00Z">'
                                  f'<w:r><w:t xml:space="preserve"> (revised)</w:t></w:r></w:ins>'))
    for _ in range(4):
        d.add_paragraph("The Owner shall review.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def run(data: bytes, op: str, select: dict, params: dict):
    pkg = DocxPackage(data)
    frag = get_op(op).plan(pkg, select, params)
    plan = merge_fragments([frag])
    work = pkg.copy()
    get_op(op).apply(work, frag)
    out = work.output_bytes()
    return plan, out, run_suite(data, out, plan.to_dict(), "src.docx", "out.docx")


def rename(data: bytes, **params):
    return run(data, "replace_text", {}, {"find": "Owner", "replace": "Client", "tracked": True, "author": "A",
                                          **params})


def failed(rec):
    return [c.id for c in rec.checks if c.status == "FAIL"]


def body(data: bytes):
    return DocxPackage(data).xml(MAIN)


def mutate(data: bytes, fn) -> bytes:
    zin = zipfile.ZipFile(io.BytesIO(data))
    root = etree.fromstring(zin.read(MAIN))
    fn(root)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        for info in zin.infolist():
            z.writestr(info, serialize_xml(root) if info.filename == MAIN else zin.read(info.filename))
    return buf.getvalue()


def first_row_exceptions(data: bytes):
    tr = next(body(data).iter(W + "tr"))
    return [etree.QName(c).localname for c in tr.find(W + "tblPrEx")] if tr.find(W + "tblPrEx") is not None else None


def table_texts(data: bytes):
    return [["".join(t.text or "" for t in tc.iter(W + "t")) for tc in tr.findall(W + "tc")]
            for tr in body(data).iter(W + "tr")]


# ---- the transform's kinds and the verifier's: the same list --------------------------------------------------------

def test_both_sides_count_it_as_a_property_change():
    from spec_surgeon.editspec import REV_DELTA_KINDS
    from spec_surgeon.verify import suite
    el = etree.fromstring(f'<w:tblPrEx xmlns:w="{W[1:-1]}">{CHANGE}</w:tblPrEx>')[0]
    assert ooxml.revision_kind(el) == "tblPrExChange" and suite._rev_kind(el) == "tblPrExChange"
    assert "tblPrExChange" in REV_DELTA_KINDS and "tblPrExChange" in suite.PROPERTY_REV
    assert set(ooxml.PROPERTY_REV_TAGS) | set(ooxml.CELL_REV_TAGS) == set(suite.OTHER_REV)


# ---- V4: its id is in the duplicate and collision checks -------------------------------------------------------------

@pytest.mark.parametrize("seed, collides", [(0, ["7"]), (7, [])])
def test_a_new_revision_reusing_its_id_fails_v4(monkeypatch, seed, collides):
    # the shared defect: the seed that does not read w:tblPrExChange (plan and apply both use the ids it gives — the
    # plan records them in transform_data); the seed that reads it (7, R1's) starts after it
    monkeypatch.setattr(replace_module, "max_revision_id", lambda pkg: seed)
    src = build()
    plan, out, rec = rename(src)
    v4 = rec.get("V4")
    assert plan.rev_delta == {"ins": 4, "del": 4}
    if collides:
        # V4 by its counts; V10's paragraph check by its own id set (verify/edits.part_revision_ids): the insertion
        # numbered 7 reads as a revision the source already had, not the new one the edit places
        assert not rec.passed and failed(rec) == ["V4", "V10"], rec.to_markdown()
        assert v4.data["colliding_new_ids"] == collides
        assert v4.data["duplicate_ids_all_kinds"] == collides
        ids = [el.get(W + "id") for el in body(out).iter(W + "ins", W + "del", W + "tblPrExChange")]
        assert ids.count("7") == 2
    else:
        assert rec.passed, rec.to_markdown()
        assert v4.data["counts_before"]["tblPrExChange"] == v4.data["counts_after"]["tblPrExChange"] == 1


def test_a_revision_renumbered_onto_its_id_fails_v4(monkeypatch):
    # a wrong OUTPUT: the transform planned ids 8.. and wrote one of them as 7
    monkeypatch.setattr(replace_module, "max_revision_id", lambda pkg: 7)
    src = build()
    plan, out, rec = rename(src)
    assert rec.passed, rec.to_markdown()
    bad = mutate(out, lambda root: next(root.iter(W + "ins")).set(W + "id", "7"))
    rec = run_suite(src, bad, plan.to_dict(), "src.docx", "out.docx")
    assert "V4" in failed(rec) and rec.get("V4").data["colliding_new_ids"] == ["7"]


def test_a_dropped_change_is_an_unpredicted_delta(monkeypatch):
    # a wrong OUTPUT: the row's pending change vanished although nothing planned it (V4 counted nothing here before)
    monkeypatch.setattr(replace_module, "max_revision_id", lambda pkg: 7)
    src = build()
    plan, out, rec = rename(src)
    bad = mutate(out, lambda root: [c.getparent().remove(c) for c in list(root.iter(W + "tblPrExChange"))])
    rec = run_suite(src, bad, plan.to_dict(), "src.docx", "out.docx")
    assert "V4" in failed(rec)
    assert rec.get("V4").data["unpredicted"] == {"tblPrExChange": (-1, 0)}


@pytest.mark.parametrize("break_it, says", [
    (lambda c: c.attrib.pop(W + "author"), "w:tblPrExChange (id 7) missing w:id or w:author"),
    (lambda c: c.remove(c[0]), "w:tblPrExChange (id 7) must hold exactly one w:tblPrEx snapshot"),
])
def test_its_grammar_is_checked_against_the_source(monkeypatch, break_it, says):
    monkeypatch.setattr(replace_module, "max_revision_id", lambda pkg: 7)
    src = build()
    plan, out, rec = rename(src)
    assert rec.get("V5c").status == "PASS"
    bad = mutate(out, lambda root: break_it(next(root.iter(W + "tblPrExChange"))))
    rec = run_suite(src, bad, plan.to_dict(), "src.docx", "out.docx")
    assert rec.get("V5c").status == "FAIL" and any(says in e for e in rec.get("V5c").data["errors"])


# ---- ops that take one out predict it ------------------------------------------------------------------------------

@pytest.mark.parametrize("op, select", [
    ("delete_table_rows", {"tables": SCHEDULE, "rows": {"index_range": [0, 0]}}),
    ("delete_tables", {"tables": SCHEDULE}),
])
def test_an_untracked_deletion_predicts_the_change_it_takes(op, select):
    src = build()
    plan, out, rec = run(src, op, select, {"tracked": False})
    assert rec.passed, rec.to_markdown()
    assert plan.rev_delta["tblPrExChange"] == -1
    v4 = rec.get("V4").data
    assert v4["counts_before"] == {"tblPrExChange": 1} and v4["counts_after"] == {}
    assert next(body(out).iter(W + "tblPrExChange"), None) is None


def test_a_wrong_count_shared_by_plan_and_apply_fails_v4(monkeypatch):
    # the transform's kinds without w:tblPrExChange: the plan predicts no change of it, the row takes it away
    monkeypatch.setattr(ooxml, "PROPERTY_REV_TAGS", OLD_PROPERTY_TAGS)
    src = build()
    plan, out, rec = run(src, "delete_table_rows", {"tables": SCHEDULE, "rows": {"index_range": [0, 0]}},
                         {"tracked": False})
    assert "tblPrExChange" not in plan.rev_delta
    assert failed(rec) == ["V4"], rec.to_markdown()
    assert rec.get("V4").data["unpredicted"] == {"tblPrExChange": (-1, 0)}


def test_a_tracked_row_deletion_leaves_it_pending():
    src = build(insertion=True)
    plan, out, rec = run(src, "delete_table_rows", {"tables": SCHEDULE, "rows": {"index_range": [0, 0]}},
                         {"tracked": True, "author": "A"})
    assert rec.passed, rec.to_markdown()
    assert plan.rev_delta.get("tblPrExChange", 0) == 0
    assert first_row_exceptions(out) == ["jc", "tblPrExChange"]


@pytest.mark.parametrize("tracked", [False, True])
def test_a_row_copied_from_it_does_not_copy_the_change(tracked):
    src = build(insertion=True)
    plan, out, rec = run(src, "insert_table_rows", {"tables": SCHEDULE, "rows": {"index_range": [0, 0]}},
                         {"rows": [["FP-0", "Spare"]], "tracked": tracked, "author": "A"})
    assert rec.passed, rec.to_markdown()
    assert plan.rev_delta.get("tblPrExChange", 0) == 0
    trs = list(body(out).iter(W + "tr"))
    assert [etree.QName(c).localname for c in trs[1].find(W + "tblPrEx")] == ["jc"]
    assert first_row_exceptions(out) == ["jc", "tblPrExChange"]


def test_clearing_a_rows_exceptions_keeps_its_change_and_flags_it():
    src = build(CENTRED + '<w:shd w:val="clear" w:color="auto" w:fill="FF0000"/>' + CHANGE)
    plan, out, rec = run(src, "apply_table_style", {}, {"style": "Table Grid", "direct_formatting": "clear"})
    assert rec.passed, rec.to_markdown()
    assert plan.rev_delta.get("tblPrExChange", 0) == 0
    assert first_row_exceptions(out) == ["jc", "tblPrExChange"]
    assert any("1 changed table(s), row(s) or cell(s) already carry a pending tracked property change" in f
               for f in plan.flags), plan.flags


# ---- resolve_tracked_changes -----------------------------------------------------------------------------------------

def test_accepting_it_keeps_the_exceptions():
    src = build(insertion=True)
    plan, out, rec = run(src, "resolve_tracked_changes", {}, {"action": "accept"})
    assert rec.passed, rec.to_markdown()
    assert {k: v for k, v in plan.rev_delta.items() if v} == {"ins": -1, "tblPrExChange": -1}
    assert plan.property_assertions[0]["expected_gone"] == {MAIN: {"ins": 1, "tblPrExChange": 1}}
    assert first_row_exceptions(out) == ["jc"]


def test_accepting_the_only_thing_a_tblprex_holds_leaves_it_empty():
    src = build(CHANGE.replace("<w:tblPrEx/>", f"<w:tblPrEx>{CENTRED}</w:tblPrEx>"))
    plan, out, rec = run(src, "resolve_tracked_changes", {}, {"action": "accept"})
    assert rec.passed, rec.to_markdown()
    assert plan.rev_delta["tblPrExChange"] == -1
    assert first_row_exceptions(out) == []


def test_rejecting_it_is_done_in_word():
    with pytest.raises(JobException) as e:
        get_op("resolve_tracked_changes").plan(DocxPackage(build(insertion=True)), {}, {"action": "reject"})
    assert e.value.code == "revision-kind"
    # not chosen, it stays as it is: rejecting only Q's insertion
    src = build(insertion=True)
    plan, out, rec = run(src, "resolve_tracked_changes", {}, {"action": "reject", "authors": ["Q"]})
    assert rec.passed, rec.to_markdown()
    assert {k: v for k, v in plan.rev_delta.items() if v} == {"ins": -1}
    assert first_row_exceptions(out) == ["jc", "tblPrExChange"]


def test_accepting_with_the_old_kinds_shared_by_plan_and_apply_fails_v10(monkeypatch):
    # the op without w:tblPrExChange among its revisions: it accepts Q's insertion and leaves the row's change pending,
    # and predicts exactly that; the verifier resolves the source with its own kinds and finds the change still there
    old = tuple(W + t for t in OLD_PROPERTY_TAGS)
    monkeypatch.setattr(rev_module, "PROPERTY_CHANGE_TAGS", old)
    monkeypatch.setattr(rev_module, "REVISION_TAGS", rev_module.CONTENT_TAGS + old + rev_module.CELL_TAGS)
    src = build(insertion=True)
    plan, out, rec = run(src, "resolve_tracked_changes", {}, {"action": "accept"})
    assert first_row_exceptions(out) == ["jc", "tblPrExChange"]
    assert "V10" in failed(rec), rec.to_markdown()
    errors = " | ".join(rec.get("V10").data.get("errors", []))
    assert "is not the source with the chosen tracked changes accepted" in errors
    assert "revisions that go here" in errors
