"""Roles beside redlines (backlog C7). A list is unsettled only by a revision that can change a number: a paragraph
mark inserted, deleted or moved, a tracked paragraph-property change, a numbering change. A wording redline inside a
run changes no number, so its list keeps its roles, and the redlined paragraph takes part when it reads alike with its
changes accepted and with them rejected; when its two views read as different things it is excluded, alone, as any
excluded line is. Roles are planning evidence only; the verifier does not read them."""
from __future__ import annotations

import pytest
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from spec_surgeon.structure import STRUCTURAL_REVISION, VIEWS_DIFFER, _paragraph_text, analyze_structure, role_summary
from test_role_summary import selected
from test_structure import document, numbered, numbering, package

XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"


def wrapper(kind, rid):
    el = OxmlElement("w:" + kind)
    el.set(qn("w:id"), str(rid))
    el.set(qn("w:author"), "Reviewer")
    el.set(qn("w:date"), "2026-10-27T00:00:00Z")
    return el


def run(text, deleted=False):
    r = OxmlElement("w:r")
    t = OxmlElement("w:delText" if deleted else "w:t")
    t.set(XML_SPACE, "preserve")
    t.text = text
    r.append(t)
    return r


def revise(p, kind, text, rid, at_start=False):
    """``text`` inserted, deleted or moved (``kind``) inside the paragraph as a reviewer's tracked change, at its end
    or before its first run."""
    el = wrapper(kind, rid)
    el.append(run(text, deleted=kind in ("del", "moveFrom")))
    if at_start:
        p.runs[0]._r.addprevious(el)
    else:
        p._p.append(el)
    return el


def mark_revision(p, kind, rid):
    """The paragraph mark itself inserted, deleted or moved."""
    ppr = p._p.get_or_add_pPr()
    rpr = ppr.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        ppr.append(rpr)
    rpr.append(wrapper(kind, rid))


def property_change(p, rid, old_numbering=None):
    """A tracked paragraph-property change; ``old_numbering`` (level, list id) numbers the old properties."""
    change = wrapper("pPrChange", rid)
    old = OxmlElement("w:pPr")
    if old_numbering is not None:
        num = OxmlElement("w:numPr")
        for name, value in (("ilvl", old_numbering[0]), ("numId", old_numbering[1])):
            child = OxmlElement("w:" + name)
            child.set(qn("w:val"), str(value))
            num.append(child)
        old.append(num)
    change.append(old)
    p._p.get_or_add_pPr().append(change)


def run_format_change(p, rid):
    """The first run's formatting changed with tracking on."""
    change = wrapper("rPrChange", rid)
    change.append(OxmlElement("w:rPr"))
    p.runs[0]._r.get_or_add_rPr().append(change)


def listed(*clauses):
    """SECTION, its title, then one automatic list: PART 1 - GENERAL, 1.1 SUMMARY and the A., B., … clauses."""
    doc = document("SECTION 21 13 13", "SPRINKLERS")
    numbering(doc)
    for level, text in ((0, "GENERAL"), (1, "SUMMARY")) + tuple((2, c) for c in clauses):
        numbered(doc.add_paragraph(text), level)
    return doc


def test_a_wording_redline_inside_a_run_keeps_the_paragraph_and_its_list_in_roles():
    doc = listed("Install the products", "Later clause", "Last clause")
    p = doc.paragraphs[4]
    revise(p, "del", " promptly", 8)
    revise(p, "ins", " within a week", 9)
    rows = analyze_structure(package(doc))["paragraphs"]
    assert rows[4]["disposition"] == "classified" and rows[4]["role"] == "paragraph"
    assert rows[4]["reference_id"] == "1.1.A"
    assert rows[4]["text"] == "Install the products promptly within a week"        # what a selector's regex meets
    assert rows[4]["views"] == {"accept": "Install the products within a week",
                                "reject": "Install the products promptly"}
    assert rows[4]["content"] == "Install the products within a week"
    assert "wording redline: reads alike with its changes accepted and rejected" in rows[4]["evidence"]
    assert [r["reference_id"] for r in rows[5:]] == ["1.1.B", "1.1.C"]
    assert not any("pending" in r["reason"] for r in rows)
    assert all("views" not in r for r in rows if r is not rows[4])


def test_a_redline_that_reads_differently_in_its_two_views_is_excluded_alone():
    doc = listed("Inserted clause", "Later clause")        # the wording inserted whole: rejected, an empty item
    p = doc.paragraphs[4]
    ins = wrapper("ins", 8)
    ins.append(p.runs[0]._r)
    p._p.append(ins)
    rows = analyze_structure(package(doc))["paragraphs"]
    assert rows[4]["disposition"] == "excluded" and rows[4]["reason"] == VIEWS_DIFFER
    assert "views" not in rows[4] and rows[4]["text"] == "Inserted clause"
    assert rows[5]["disposition"] == "unresolved" and "earlier pending revisions" not in rows[5]["reason"]
    assert rows[5]["reason"].startswith("recognized marker lacks")     # the excluded line broke its context, as any


def test_a_typed_marker_taken_out_by_a_redline_reads_differently():
    doc = document("SECTION 21 13 13", "SPRINKLERS", "1.1 SUMMARY", "Supply approved products", "B. Next clause")
    revise(doc.paragraphs[3], "del", "A. ", 8, at_start=True)
    rows = analyze_structure(package(doc))["paragraphs"]
    assert rows[3]["text"] == "A. Supply approved products" and rows[3]["reason"] == VIEWS_DIFFER
    assert rows[4]["disposition"] == "unresolved"


@pytest.mark.parametrize("kind", ["ins", "del", "moveFrom", "moveTo", "pPrChange"])
def test_a_revision_that_can_change_a_number_excludes_the_paragraph_and_unsettles_its_list(kind):
    doc = listed("Changed clause", "Later clause")
    if kind == "pPrChange":
        property_change(doc.paragraphs[4], 8)
    else:
        mark_revision(doc.paragraphs[4], kind, 8)
    rows = analyze_structure(package(doc))["paragraphs"]
    assert rows[4]["disposition"] == "excluded" and rows[4]["reason"] == STRUCTURAL_REVISION
    assert rows[5]["disposition"] == "unresolved" and "earlier pending revisions" in rows[5]["reason"]


def test_a_tracked_property_change_that_took_a_paragraph_out_of_its_list_unsettles_the_list():
    doc = listed("First clause")
    property_change(doc.add_paragraph("Taken out of the list"), 8, old_numbering=(2, 77))
    numbered(doc.add_paragraph("Later clause"), 2)
    rows = analyze_structure(package(doc))["paragraphs"]
    assert rows[4]["reference_id"] == "1.1.A"
    assert rows[5]["disposition"] == "excluded" and rows[5]["reason"] == STRUCTURAL_REVISION
    assert rows[6]["disposition"] == "unresolved" and "earlier pending revisions" in rows[6]["reason"]


def test_a_formatting_change_unsettles_nothing():
    doc = listed("First clause", "Later clause")
    run_format_change(doc.paragraphs[4], 8)
    rows = analyze_structure(package(doc))["paragraphs"]
    assert [r["reference_id"] for r in rows[4:]] == ["1.1.A", "1.1.B"]
    assert rows[4]["views"] == {"accept": "First clause", "reject": "First clause"}


def test_the_two_views_of_nested_and_moved_wording():
    doc = document("Kept")
    p = doc.paragraphs[0]
    ins = revise(p, "ins", " added", 8)
    struck = wrapper("del", 9)                     # a deletion inside someone's insertion
    struck.append(run(" struck", deleted=True))
    ins.append(struck)
    revise(p, "moveFrom", " away", 10)
    revise(p, "moveTo", " here", 11)
    assert _paragraph_text(p._p) == "Kept added struck away here"
    assert _paragraph_text(p._p, "accept") == "Kept added here"
    assert _paragraph_text(p._p, "reject") == "Kept away"


def test_a_redlined_paragraph_counts_for_its_role_with_the_marker_read_from_the_accepted_view():
    doc = document("SECTION 21 13 13", "SPRINKLERS", "1.1 SUMMARY", "A. Supply approved products",
                   "1.2 REFERENCES", "A. NFPA 13")
    p = doc.paragraphs[3]
    revise(p, "del", " promptly", 8)
    revise(p, "ins", " at once", 9)
    pkg = package(doc)
    summary = role_summary(pkg)
    assert summary["roles"]["paragraph"]["count"] == 2
    assert summary["roles"]["paragraph"]["markers"] == ["A."]
    assert summary["roles"]["paragraph"]["examples"][0] == "A. Supply approved products promptly at once"
    selector = summary["selectors"]["paragraph"]
    assert selector == {"text_regex": r"^\s*[A-Z]\.(?:\s|$)"}
    assert selected(pkg, selector) == 2


def test_a_marker_the_engine_reads_otherwise_under_its_redline_gets_no_selector():
    """A regex selector is validated on the text the engine's selector meets (the kept and the deleted wording
    together): a marker's letter retyped under a redline reads alike in both views but not there, so the role gets
    no selector rather than one that misses the paragraph."""
    doc = document("SECTION 21 13 13", "SPRINKLERS", "1.1 SUMMARY", ". Supply approved products", "1.2 REFERENCES",
                   "A. NFPA 13")
    p = doc.paragraphs[3]
    revise(p, "del", "A", 8, at_start=True)
    revise(p, "ins", "A", 9, at_start=True)
    pkg = package(doc)
    rows = analyze_structure(pkg)["paragraphs"]
    assert rows[3]["role"] == "paragraph" and rows[3]["text"] == "AA. Supply approved products"
    summary = role_summary(pkg)
    assert summary["roles"]["paragraph"]["count"] == 2 and summary["roles"]["paragraph"]["markers"] == ["A."]
    assert "paragraph" not in summary["selectors"]


def test_a_redline_on_a_section_line_that_changes_what_follows_is_excluded():
    """Codex on #120: a title inserted on a bare SECTION line gives both views the same section reading, but with the
    change rejected the next uppercase line is the section's title and with it accepted it is not: the readings
    differ. A word changed inside a title the line carries in both views reads alike."""
    doc = document("SECTION 21 13 13", "SPRINKLERS", "1.1 SUMMARY")
    revise(doc.paragraphs[0], "ins", " WET-PIPE SPRINKLER SYSTEMS", 8)
    rows = analyze_structure(package(doc))["paragraphs"]
    assert rows[0]["disposition"] == "excluded" and rows[0]["reason"] == VIEWS_DIFFER
    assert rows[1]["role"] is None
    doc = document("SECTION 21 13 13 WET-PIPE SYSTEMS", "1.1 SUMMARY")
    revise(doc.paragraphs[0], "ins", " AND VALVES", 8)
    rows = analyze_structure(package(doc))["paragraphs"]
    assert rows[0]["role"] == "section_id" and rows[0]["content"] == "WET-PIPE SYSTEMS AND VALVES"
    assert rows[1]["role"] == "article"


@pytest.mark.parametrize("where", ["run", "mark"])
def test_a_formatting_change_that_hid_the_text_before_is_not_admitted(where):
    """Codex on #120: a tracked formatting change whose old properties hide the run (or the mark) reads alike in
    both views, but with the change rejected the text is invisible; the paragraph stays outside coverage, as hidden
    text is, and its list is not unsettled."""
    doc = listed("First clause", "Later clause")
    change = wrapper("rPrChange", 8)
    old = OxmlElement("w:rPr")
    old.append(OxmlElement("w:vanish"))
    change.append(old)
    p = doc.paragraphs[4]
    if where == "run":
        p.runs[0]._r.get_or_add_rPr().append(change)
    else:
        mark = OxmlElement("w:rPr")
        mark.append(change)
        p._p.get_or_add_pPr().append(mark)
    rows = analyze_structure(package(doc))["paragraphs"]
    assert rows[4]["disposition"] == "excluded" and rows[4]["reason"].startswith("hidden text")
    assert rows[5]["disposition"] == "unresolved" and "earlier pending revisions" not in rows[5]["reason"]
