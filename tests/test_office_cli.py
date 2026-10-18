"""The user approves actual preview evidence, independently of the outcome review."""
from __future__ import annotations

import io
import re
import zipfile

import fitz
import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE

from spec_surgeon.cli import main
from spec_surgeon.office_cli import preview_digest
from spec_surgeon.office_profiles import create_profile, load_review, prepare_review, preview_review
from spec_surgeon.ooxml import RevisionStamp
from spec_surgeon.package import W, parse_xml, serialize_xml


def documents(tmp_path):
    source, template = tmp_path / "source.docx", tmp_path / "template.docx"
    doc = Document()
    doc.add_paragraph("PART 1 - GENERAL")
    doc.save(source)
    doc = Document()
    doc.styles.add_style("Office Part", WD_STYLE_TYPE.PARAGRAPH).font.name = "Arial"
    doc.save(template)
    return source, template


def fixed_date(monkeypatch, date):
    monkeypatch.setattr(RevisionStamp, "now", classmethod(lambda cls, author, use_date_utc, convention="word-local-z":
        cls(author, date, date, use_date_utc)))


@pytest.mark.parametrize("tracked", [False, True])
def test_real_cli_create_review_preview_run_with_separate_current_approvals(tmp_path, capsys, monkeypatch, tracked):
    source, template = documents(tmp_path)
    original = source.read_bytes()
    profile, review_file, output = tmp_path / "profile", tmp_path / "review.json", tmp_path / "job"
    fixed_date(monkeypatch, "2026-01-01T00:00:00Z")
    args = ["profiles", "create", str(template), "--name", "Office", "--roles", '{"part":"Office Part"}',
            "--out", str(profile)] + (["--tracked"] if tracked else [])
    assert main(args) == 0
    assert main(["profiles", "review", str(profile), str(source), "--out", str(review_file)]) == 0
    review = load_review(review_file)
    capsys.readouterr()
    common = [str(review_file), "--out", str(output), "--renderer", "none", "--no-render"]
    assert main(["profiles", "preview", *common]) == 0
    printed = capsys.readouterr().out
    digest = re.search(r"Preview approval digest: ([0-9a-f]{64})", printed).group(1)
    # Regeneration at another time must not invalidate an otherwise identical
    # newly-generated formatting revision. Existing revisions stay literal.
    fixed_date(monkeypatch, "2026-02-01T12:34:56Z")
    assert main(["profiles", "run", *common, "--approval", review["digest"],
                 "--preview-approval", "wrong", "--workers", "1"]) == 2
    assert not (output / "run").exists()
    assert main(["profiles", "run", *common, "--approval", review["digest"],
                 "--preview-approval", digest, "--workers", "1"]) == 0, capsys.readouterr()
    delivered = list((output / "run" / "file-0001" / "outputs").glob("*.docx"))
    assert len(delivered) == 1
    assert Document(delivered[0]).paragraphs[0].style.name == "Office Part"
    assert source.read_bytes() == original


def test_preview_digest_detects_actual_document_checks_and_application_changes(tmp_path, monkeypatch):
    from spec_surgeon import office_cli
    source, template = documents(tmp_path)
    profile = create_profile(tmp_path / "profile", template, "Office", {"part": "Office Part"})
    review = prepare_review(profile, [source])
    previews = preview_review(review, tmp_path / "preview", renderer="none", do_render=False)
    first = preview_digest(review, previews, "none", False)
    sample = previews[str(source)].samples[0]
    data = sample.output.read_bytes()
    doc = Document(io.BytesIO(data))
    doc.paragraphs[0].runs[0].italic = True
    doc.save(sample.output)
    assert preview_digest(review, previews, "none", False) != first
    sample.output.write_bytes(data)
    sample.checks["NEW_CHECK"] = "FAIL"
    assert preview_digest(review, previews, "none", False) != first
    sample.checks.pop("NEW_CHECK")
    assert preview_digest(review, previews, "none", False) == first
    monkeypatch.setattr(office_cli, "_code_digest", lambda: "different-installed-build")
    assert preview_digest(review, previews, "none", False) != first


def test_preview_digest_checks_all_rendered_pages_but_ignores_pdf_metadata(tmp_path):
    source, template = documents(tmp_path)
    profile = create_profile(tmp_path / "profile", template, "Office", {"part": "Office Part"})
    review = prepare_review(profile, [source])
    previews = preview_review(review, tmp_path / "preview", renderer="none", do_render=False)
    sample = previews[str(source)].samples[0]
    pdf = tmp_path / "preview.pdf"
    def write_pdf(second_page, metadata):
        with fitz.open() as doc:
            for text in ["First page", second_page]:
                doc.new_page().insert_text((72, 72), text)
            doc.set_metadata({"creationDate": metadata})
            doc.save(pdf)
    write_pdf("Second page", "D:20260101000000Z")
    sample.output_pdf = pdf
    first = preview_digest(review, previews, "none", False)
    write_pdf("Second page", "D:20260201000000Z")
    assert preview_digest(review, previews, "none", False) == first
    write_pdf("Changed second page", "D:20260201000000Z")
    assert preview_digest(review, previews, "none", False) != first


def test_existing_revision_dates_are_never_normalized_out_of_preview_authority(tmp_path):
    from spec_surgeon.office_cli import _document_digest, _formatting_revision_ids
    source, _ = documents(tmp_path)
    with zipfile.ZipFile(source) as z:
        members = [(entry, z.read(entry)) for entry in z.infolist()]
    from lxml import etree
    def write(date):
        data = io.BytesIO()
        with zipfile.ZipFile(data, "w") as z:
            for entry, payload in members:
                if entry.filename == "word/document.xml":
                    root = parse_xml(payload)
                    p = next(root.iter(W + "p"))
                    ppr = etree.Element(W + "pPr")
                    change = etree.SubElement(ppr, W + "pPrChange", {W + "id": "5", W + "author": "Original author", W + "date": date})
                    etree.SubElement(change, W + "pPr")
                    p.insert(0, ppr)
                    payload = serialize_xml(root)
                z.writestr(entry, payload)
        source.write_bytes(data.getvalue())
    write("2025-01-01T00:00:00Z")
    protected = _formatting_revision_ids([source])
    original = _document_digest(source, protected)
    write("2025-02-01T00:00:00Z")
    assert _document_digest(source, protected) != original


def test_frozen_preview_approval_includes_embedded_application_code(tmp_path, monkeypatch):
    from spec_surgeon import office_cli
    executable = tmp_path / "SpecSurgeon.exe"
    executable.write_bytes(b"first embedded PYZ build")
    monkeypatch.setattr(office_cli.sys, "frozen", True, raising=False)
    monkeypatch.setattr(office_cli.sys, "executable", str(executable))
    first = office_cli._code_digest()
    executable.write_bytes(b"changed embedded PYZ build")
    assert office_cli._code_digest() != first
