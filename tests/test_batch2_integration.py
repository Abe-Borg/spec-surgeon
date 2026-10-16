"""Built-in batch 2 in the rest of the app (docs/BUILTIN_OPS_BATCH2_DESIGN.md): the words the new ops place are traced to
the user (compiler.placed_words — rule 13), the profiles show Claude what each clean-up op would find (profile.markup),
and the method's three rules read as Abraham changed them."""
from __future__ import annotations

import ast
from pathlib import Path

import pytest
from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import RGBColor

from make_fixtures import build_rich, docx_bytes, patch_docx
from spec_surgeon import SKILL_DIR
from spec_surgeon.compiler import placed_words, system_prompt, untraceable
from spec_surgeon.editspec import apply_spec, plan_spec
from spec_surgeon.ops import JobException
from spec_surgeon.ops import fields
from spec_surgeon.package import DocxPackage
from spec_surgeon.profile import markup, profile
from spec_surgeon.report import FAILURE_EXPLANATIONS
from spec_surgeon.verify import run_suite, xsd

OPS_DIR = Path(__file__).resolve().parent.parent / "src" / "spec_surgeon" / "ops"
BATCH2_MODULES = ["comments", "highlight", "revisions", "hidden", "empty_paragraphs", "watermark", "links", "fields",
                  "docprops", "markup", "removal"]


def test_the_words_batch2_places_are_traced_to_the_user():
    spec = {"request": "r", "summary": "s", "mode": "untracked", "ops": [
        {"op": "add_watermark", "params": {"text": "NOT FOR CONSTRUCTION"}},
        {"op": "set_document_properties", "params": {"set": {"title": "Fire Suppression", "company": "ACME"},
                                                     "clear": ["author"], "custom": {"Project Number": "P-1003"}}},
        {"op": "remove_comments", "params": {"which": "all"}},
        {"op": "resolve_tracked_changes", "params": {"action": "accept", "authors": ["Jeremy"]}},
        {"op": "set_document_state", "params": {"final": True}},
    ]}
    assert placed_words(spec) == ["NOT FOR CONSTRUCTION", "Fire Suppression", "ACME", "P-1003"]
    assert untraceable(spec, ["watermark NOT FOR CONSTRUCTION, title Fire Suppression, company ACME"]) == ["P-1003"]


def _marked() -> bytes:
    doc = Document()
    p = doc.add_paragraph("Plain ")
    r = p.add_run("hidden note"); r.font.hidden = True
    r = p.add_run("yellow"); r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    r = p.add_run("blue"); r.font.color.rgb = RGBColor(0, 0, 0xFF)
    doc.core_properties.title = "Spec title"
    doc.core_properties.author = "Jeremy"
    return docx_bytes(doc)


def test_the_profile_shows_what_the_clean_up_ops_would_find():
    m = markup(DocxPackage(_marked()))
    assert m["hidden_runs"] == 1 and m["highlight_colors"] == {"yellow": 1} and m["run_colors"] == {"0000FF": 1}
    assert m["properties"]["title"] == "Spec title" and m["properties"]["author"] == "Jeremy"
    rich = profile(DocxPackage(build_rich()))
    assert rich["markup"]["tracked_changes"]["Jeremy"]["ins"] == 2
    assert rich["markup"]["tracked_change_dates"] == ["2026-09-22", "2026-09-23"]


def test_claude_is_told_when_the_batch2_ops_may_be_used():
    prompt = system_prompt(chat=True)
    assert "resolve_tracked_changes accepts or rejects tracked changes ONLY when the user asks" in prompt
    assert "a watermark's words and a property's value are the user's exact" in prompt


def test_the_method_reads_as_abraham_changed_it():
    skill = (Path(SKILL_DIR) / "SKILL.md").read_text(encoding="utf-8")
    assert "on the model's own initiative" in skill and "resolve_tracked_changes" in skill
    assert "by the model. The user accepts and rejects in Word." not in skill
    elements = (Path(SKILL_DIR) / "references" / "element-types.md").read_text(encoding="utf-8")
    assert "never toggle it on your own" in elements and "Mark as Final" in elements


def _refusal_codes(module: str):
    """(the code literals of every JobException the module raises, the calls whose code is not a literal)."""
    tree = ast.parse((OPS_DIR / f"{module}.py").read_text(encoding="utf-8"))
    literal, other = set(), []
    for n in ast.walk(tree):
        if isinstance(n, ast.Call) and getattr(n.func, "id", None) == "JobException":
            code = n.args[1] if len(n.args) > 1 else next((k.value for k in n.keywords if k.arg == "code"), None)
            if isinstance(code, ast.Constant):
                literal.add(code.value)
            else:
                other.append((module, n.lineno))
    return literal, other


def test_every_refusal_code_of_the_batch2_ops_is_explained():
    codes, computed = set(), []
    for m in BATCH2_MODULES:
        literal, other = _refusal_codes(m)
        codes |= literal
        computed += other
    # the one computed code: links._field_refusal raises Field.refusal(), a key of fields.REFUSAL_TEXT
    assert [m for m, _ in computed] == ["links"]
    codes |= set(fields.REFUSAL_TEXT)
    codes.discard("internal")                                         # diagnostics.classify: an internal error
    assert {"revision-kind", "paragraph-merge", "partial-move", "marks-in-change", "watermark-exists", "no-header",
            "watermark-size", "field-span", "nested-field", "form-field", "no-part", "protected", "property-type",
            "malformed-properties", "tracked-format", "toggle-ambiguous", "field-result"} <= codes
    assert sorted(c for c in codes if c not in FAILURE_EXPLANATIONS) == []


@pytest.mark.parametrize("op", [
    {"op": "remove_highlight"},
    {"op": "resolve_tracked_changes", "params": {"action": "accept", "authors": ["Nobody"]}},
    {"op": "remove_hidden_text", "params": {"tracked": False}},
    {"op": "remove_notes", "params": {"rules": [{"style_in": ["NoSuchStyle"]}], "tracked": False},
     "select": {"scope": ["headers"]}},
    {"op": "remove_empty_paragraphs", "params": {"tracked": False}, "select": {"scope": ["headers"]}},
    {"op": "remove_hyperlinks"},
    {"op": "unlink_fields", "params": {"types": ["DOCPROPERTY"]}},
], ids=lambda o: o["op"])
def test_a_clean_up_op_beside_a_header_removal_is_a_plan_time_conflict(op):
    """Rule 9: each of these re-derives its decisions over every part in scope, so it declares (part, "scope",
    "verify") — beside an op that removes a header part it is refused at plan time, even when it has nothing to do in
    that header (it used to plan, and its V10 kind then failed on the missing part after apply)."""
    pkg = DocxPackage(build_rich())
    spec = {"request": "r", "summary": "s", "mode": "untracked" if op["op"] in (
        "resolve_tracked_changes", "remove_hidden_text", "remove_notes", "remove_empty_paragraphs") else "formatting",
        "ops": [{"op": "remove_header_footer", "select": {"sections": "all", "kinds": ["header"]}}, op]}
    with pytest.raises(JobException) as e:
        plan_spec(pkg, spec)
    assert e.value.code == "conflict" and "remove_header_footer" in str(e.value)
    alone = {**spec, "ops": [op]}
    plan, frags = plan_spec(pkg, alone)                    # on its own the op plans (a no-op where nothing matches)
    out = apply_spec(pkg, frags).output_bytes()
    rec = run_suite(pkg.data, out, plan.to_dict(), "src.docx", "out.docx")
    assert [c.id for c in rec.checks if c.status == "FAIL"] == []


def _verify(data: bytes, ops, mode="untracked", tamper=None):
    pkg = DocxPackage(data)
    plan, frags = plan_spec(pkg, {"request": "r", "summary": "s", "mode": mode, "ops": ops})
    out = apply_spec(pkg, frags).output_bytes()
    if tamper:
        out = tamper(out)
    return run_suite(pkg.data, out, plan.to_dict(), "src.docx", "out.docx")


def test_accepting_beside_another_op_on_a_part_it_does_not_resolve():
    """resolved_changes compares WHOLE only the parts it resolves; a part with nothing to resolve stays with the
    residual — so a set_font on footers the accept never touches verifies (it planned clean and failed V10 when the
    kind claimed every part in scope), and a change nobody checked in such a part still fails."""
    ops = [{"op": "resolve_tracked_changes", "params": {"action": "accept"}},
           {"op": "set_font", "select": {"scope": ["footers"]}, "params": {"font": "Arial"}}]
    rec = _verify(build_rich(), ops)
    assert [c.id for c in rec.checks if c.status == "FAIL"] == []

    def tamper(out: bytes) -> bytes:                  # the footer loses a character nobody asked to change
        return patch_docx(out, {"word/footer1.xml": lambda root: _drop_first_char(root)})
    bad = _verify(build_rich(), ops, tamper=tamper)
    assert "V10" in [c.id for c in bad.checks if c.status == "FAIL"]


def _drop_first_char(root):
    t = next(t for t in root.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t") if t.text)
    t.text = t.text[1:]


SETTINGS_HEAD = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                 '<w:settings xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">')


def test_v5a_judges_a_settings_part_against_the_source_in_schema_order():
    """libxml2 stops checking a sequence at the first child it does not expect: a w:trackRevisions placed first hides a
    w:zoom without its required w:percent. Turning track changes off then let V5a reach the zoom error and call it new
    on a correct output. The baseline also validates the source with its children in schema order — every error that
    copy shows is the source's own — and a really new error is still new."""
    src = (SETTINGS_HEAD + '<w:trackRevisions/><w:zoom w:val="bestFit"/><w:defaultTabStop w:val="720"/>'
           '</w:settings>').encode()
    off = (SETTINGS_HEAD + '<w:zoom w:val="bestFit"/><w:defaultTabStop w:val="720"/></w:settings>').encode()
    assert xsd.validate_part("word/settings.xml", src) == set()          # the zoom error is out of reach here
    assert xsd.validate_part("word/settings.xml", off) != set()          # and in reach once the element goes
    base, out, new = xsd.compare("word/settings.xml", src, off)
    assert out and new == set()
    broken = (SETTINGS_HEAD + '<w:zoom w:val="bestFit"/><w:defaultTabStop/></w:settings>').encode()
    assert xsd.compare("word/settings.xml", src, broken)[2] != set()     # a missing w:val is the output's own
    fine = (SETTINGS_HEAD + '<w:zoom w:percent="100"/><w:defaultTabStop w:val="720"/></w:settings>').encode()
    assert xsd.in_schema_order(fine) is None                              # in order: nothing more to validate
    assert xsd.compare("word/settings.xml", fine, broken.replace(b'w:val="bestFit"', b'w:percent="100"'))[2] != set()
    rec = _verify(build_rich(), [{"op": "set_document_state", "params": {"track_changes": "off"}}], mode="formatting")
    assert [c.id for c in rec.checks if c.status == "FAIL"] == []
