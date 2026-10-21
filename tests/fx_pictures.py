"""Fixtures for built-in batch 6's picture ops — insert_image, remove_images, resize_images
(docs/BUILTIN_OPS_BATCH6_DESIGN.md §5). Every paragraph is found by its text in the tests, never by index.

``build_pictures()``: a two-section specification — the body's 6.5-inch text width, then a 3.5-inch one —
* the header (a red logo, the same bytes as a picture in the body: one media member, two relationships) and a footer
  with no relationships at all;
* the body: a site plan (blue, alt text "Site plan") in its own paragraph, the same blue picture again in a table cell
  (one relationship shared by both drawings), a riser diagram (green) in the run that holds the text before it, the
  red logo, a floating picture (an anchor, wrapped square) and a picture with a hyperlink on it;
* ``Figure`` (a paragraph style, centred) and ``Emphasis Strong`` (a character style).

The ``picture files`` are PNG, JPEG, GIF and BMP headers and bodies written here (a resolution where asked); the
``mutate`` helpers plant the cases each op refuses or leaves alone.
"""
from __future__ import annotations

import io
import struct
import zipfile
import zlib
from typing import Callable, Dict, Optional, Tuple

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from lxml import etree


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W = "{%s}" % W_NS
WP = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
PIC = "{http://schemas.openxmlformats.org/drawingml/2006/picture}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
V = "{urn:schemas-microsoft-com:vml}"
O = "{urn:schemas-microsoft-com:office:office}"
MC = "{http://schemas.openxmlformats.org/markup-compatibility/2006}"
WP14 = "{http://schemas.microsoft.com/office/word/2010/wordprocessingDrawing}"
HYPERLINK_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink"

RED, BLUE, GREEN, GREY = (200, 30, 30), (30, 30, 200), (10, 200, 30), (90, 90, 90)


# ---- picture files ---------------------------------------------------------------------------------------------------

def _chunk(kind: bytes, data: bytes) -> bytes:
    return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data) & 0xFFFFFFFF)


def png(w: int, h: int, rgb=GREY, ppm: Optional[Tuple[int, int]] = None) -> bytes:
    """A solid w × h PNG; ``ppm`` (pixels per metre across, down) writes a pHYs chunk."""
    raw = b"".join(b"\x00" + bytes(rgb) * w for _ in range(h))
    out = b"\x89PNG\r\n\x1a\n" + _chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0))
    if ppm:
        out += _chunk(b"pHYs", struct.pack(">IIB", ppm[0], ppm[1], 1))
    return out + _chunk(b"IDAT", zlib.compress(raw)) + _chunk(b"IEND", b"")


def gif(w: int, h: int) -> bytes:
    """A GIF with a w × h logical screen (a two-colour table, one 1 × 1 frame)."""
    return (b"GIF89a" + struct.pack("<HH", w, h) + bytes([0x80, 0, 0]) + b"\x00\x00\x00\xff\xff\xff"
            + b"\x2c" + struct.pack("<HHHH", 0, 0, 1, 1) + b"\x00" + b"\x02\x02\x44\x01\x00" + b"\x3b")


def bmp(w: int, h: int, ppm: Tuple[int, int] = (0, 0), rgb=GREY) -> bytes:
    """A 24-bit bottom-up BMP (BITMAPINFOHEADER), ``ppm`` its pixels per metre (0: none)."""
    row = bytes(rgb[::-1]) * w
    row += b"\x00" * ((4 - len(row) % 4) % 4)
    pixels = row * h
    dib = struct.pack("<IiiHHIIiiII", 40, w, h, 1, 24, 0, len(pixels), ppm[0], ppm[1], 0, 0)
    return b"BM" + struct.pack("<IHHI", 14 + len(dib) + len(pixels), 0, 0, 14 + len(dib)) + dib + pixels


def jpeg(w: int, h: int, dpi: Optional[int] = None) -> bytes:
    """A real JPEG (PyMuPDF), its JFIF density set to ``dpi`` when given."""
    import pymupdf
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, w, h), False)
    pix.clear_with(180)
    if dpi:
        pix.set_dpi(dpi, dpi)
    return pix.tobytes("jpeg")


def turned_jpeg(w: int, h: int) -> bytes:
    """A JPEG's header up to its frame header whose Exif orientation (6) turns it a quarter turn."""
    def seg(marker: int, body: bytes) -> bytes:
        return b"\xff" + bytes([marker]) + struct.pack(">H", len(body) + 2) + body
    ifd = struct.pack("<H", 1) + struct.pack("<HHI", 0x0112, 3, 1) + struct.pack("<HH", 6, 0)
    return (b"\xff\xd8" + seg(0xE1, b"Exif\x00\x00II" + struct.pack("<HI", 42, 8) + ifd + b"\x00" * 4)
            + seg(0xC0, b"\x08" + struct.pack(">HH", h, w) + b"\x03\x01\x22\x00\x02\x11\x01\x03\x11\x01") + b"\xff\xd9")


SVG = b'<?xml version="1.0"?><svg xmlns="http://www.w3.org/2000/svg" width="10" height="10"/>'


# ---- the document ----------------------------------------------------------------------------------------------------

def save(doc) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _name(run, name: str, descr: Optional[str] = None) -> None:
    """The last drawing of a python-docx run: its wp:docPr name (and alt text)."""
    dp = list(run._r.iter(WP + "docPr"))[-1]
    dp.set("name", name)
    if descr is not None:
        dp.set("descr", descr)


def _float(run) -> None:
    """The run's inline picture made floating: a wp:anchor, square wrapping, beside the column."""
    inline = run._r.find(f"{W}drawing/{WP}inline")
    anchor = etree.Element(WP + "anchor", {"distT": "0", "distB": "0", "distL": "114300", "distR": "114300",
                                           "simplePos": "0", "relativeHeight": "251659264", "behindDoc": "0",
                                           "locked": "0", "layoutInCell": "1", "allowOverlap": "1"})
    etree.SubElement(anchor, WP + "simplePos", {"x": "0", "y": "0"})
    h = etree.SubElement(anchor, WP + "positionH", {"relativeFrom": "column"})
    etree.SubElement(h, WP + "posOffset").text = "0"
    v = etree.SubElement(anchor, WP + "positionV", {"relativeFrom": "paragraph"})
    etree.SubElement(v, WP + "posOffset").text = "0"
    anchor.append(inline.find(WP + "extent"))
    etree.SubElement(anchor, WP + "effectExtent", {"l": "0", "t": "0", "r": "0", "b": "0"})
    etree.SubElement(anchor, WP + "wrapSquare", {"wrapText": "bothSides"})
    for child in list(inline):
        anchor.append(child)
    inline.getparent().replace(inline, anchor)


def build_pictures() -> bytes:
    doc = Document()
    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)
    figure = styles.add_style("Figure", WD_STYLE_TYPE.PARAGRAPH)
    figure.base_style = styles["Normal"]
    figure.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    styles.add_style("Emphasis Strong", WD_STYLE_TYPE.CHARACTER).font.bold = True

    s0 = doc.sections[0]
    s0.page_width, s0.page_height = Inches(8.5), Inches(11)
    s0.left_margin = s0.right_margin = Inches(1)
    s0.gutter = 0
    hp = s0.header.paragraphs[0]
    hp.add_run("ACME DATA CENTERS ")
    run = hp.add_run()
    run.add_picture(io.BytesIO(png(8, 8, RED)), width=Pt(36))
    _name(run, "Logo 1")
    s0.footer.paragraphs[0].add_run("21 13 13 - SPRINKLERS")

    doc.add_paragraph("SECTION 21 13 13 - WET-PIPE SPRINKLER SYSTEMS")
    doc.add_paragraph("PART 1 - GENERAL")
    doc.add_paragraph("1.01\tSUMMARY")
    doc.add_paragraph("A.\tSee the site plan below.")
    run = doc.add_paragraph().add_run()
    run.add_picture(io.BytesIO(png(8, 8, BLUE)), width=Pt(72))
    _name(run, "Picture 2", "Site plan")
    p = doc.add_paragraph()
    p.add_run("B.\tThe riser diagram: ")
    run = p.add_run("riser ")
    run.add_picture(io.BytesIO(png(8, 4, GREEN)), width=Pt(48))                # the picture shares the text's run
    _name(run, "Riser diagram")
    t = doc.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "Valve room"
    run = t.cell(0, 1).paragraphs[0].add_run()
    run.add_picture(io.BytesIO(png(8, 8, BLUE)), width=Pt(36))                 # the site plan's bytes: one relationship
    _name(run, "Picture 4")
    doc.add_paragraph("C.\tThe owner's logo:")
    run = doc.add_paragraph().add_run()
    run.add_picture(io.BytesIO(png(8, 8, RED)), width=Pt(36))                  # the header logo's bytes
    _name(run, "Logo 2")
    p = doc.add_paragraph("D.\tA floating detail beside this text.")
    run = p.add_run()
    run.add_picture(io.BytesIO(png(16, 8, GREY)), width=Pt(96))
    _name(run, "Detail 1")
    _float(run)
    run = doc.add_paragraph().add_run()
    run.add_picture(io.BytesIO(png(4, 4, (250, 200, 0))), width=Pt(24))
    _name(run, "Linked badge")
    link_rid = run.part.relate_to("https://example.com/badge", HYPERLINK_REL, is_external=True)
    click = etree.SubElement(list(run._r.iter(WP + "docPr"))[-1], A + "hlinkClick")
    click.set(R + "id", link_rid)
    doc.add_paragraph("E.\tEnd of part 1.")

    s1 = doc.add_section(WD_SECTION.NEW_PAGE)
    s1.left_margin = s1.right_margin = Inches(2.5)                             # 3.5 inches between the margins
    doc.add_paragraph("PART 2 - PRODUCTS")
    doc.add_paragraph("2.01\tMATERIALS")
    doc.add_paragraph("A.\tPipe shall be black steel.")
    return save(doc)


# ---- mutations: the cases each op refuses or leaves alone ------------------------------------------------------------

def mutate(docx: bytes, part: str, fn: Callable) -> bytes:
    """``docx`` with ``part`` parsed, changed by ``fn(root)`` and written back (fixture surgery only)."""
    zin = zipfile.ZipFile(io.BytesIO(docx))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            if info.filename == part:
                root = etree.fromstring(data)
                fn(root)
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            zout.writestr(info, data)
    return out.getvalue()


def docpr(root, name: str):
    return next(d for d in root.iter(WP + "docPr") if d.get("name") == name)


def drawing_named(root, name: str):
    d = docpr(root, name)
    while d.tag != W + "drawing":
        d = d.getparent()
    return d


def text_of(p) -> str:
    """A paragraph's text, a w:tab read as a tab (python-docx writes the tabs of a text as w:tab)."""
    return "".join("\t" if e.tag == W + "tab" else e.text or "" for e in p.iter(W + "t", W + "tab")
                   if e.getparent().tag == W + "r")


def paragraph_with(root, text: str):
    return next(p for p in root.iter(W + "p") if text_of(p).startswith(text))


def turned(name: str) -> Callable:
    def fn(root):
        next(drawing_named(root, name).iter(A + "xfrm")).set("rot", "5400000")
    return fn


def cropped(name: str) -> Callable:
    def fn(root):
        next(drawing_named(root, name).iter(A + "blip")).addnext(etree.Element(A + "srcRect", l="25000"))
    return fn


def relative(name: str) -> Callable:
    """The floating picture's width follows the page (wp14:sizeRelH)."""
    def fn(root):
        anchor = drawing_named(root, name).find(WP + "anchor")
        rel = etree.SubElement(anchor, WP14 + "sizeRelH", relativeFrom="page")
        etree.SubElement(rel, WP14 + "pctWidth").text = "20000"
    return fn


def alternate(name: str) -> Callable:
    def fn(root):
        d = drawing_named(root, name)
        run = d.getparent()
        ac = etree.Element(MC + "AlternateContent")
        choice = etree.SubElement(ac, MC + "Choice", Requires="wp14")
        run.replace(d, ac)
        choice.append(d)
    return fn


def vml(text: str, rid: str) -> Callable:
    """A VML picture (Word 2003's) of relationship ``rid`` in the paragraph starting ``text``."""
    def fn(root):
        r = etree.SubElement(paragraph_with(root, text), W + "r")
        shape = etree.SubElement(etree.SubElement(r, W + "pict"), V + "shape", alt="old picture")
        etree.SubElement(shape, V + "imagedata", {R + "id": rid, O + "title": "Picture old"})
    return fn


def watermark(rid: str) -> Callable:
    """A picture watermark (Word's VML ``WordPictureWatermark``) of relationship ``rid`` in the header."""
    def fn(root):
        r = etree.SubElement(next(root.iter(W + "p")), W + "r")
        shape = etree.SubElement(etree.SubElement(r, W + "pict"), V + "shape", id="WordPictureWatermark123",
                                 style="position:absolute;z-index:-251657216")
        etree.SubElement(shape, V + "imagedata", {R + "id": rid, O + "title": "Logo watermark"})
    return fn


def in_insertion(name: str) -> Callable:
    """The picture's run inside someone's pending tracked insertion."""
    def fn(root):
        run = drawing_named(root, name).getparent()
        ins = etree.Element(W + "ins", {W + "id": "700", W + "author": "Jeremy", W + "date": "2026-09-22T10:15:00Z"})
        run.addprevious(ins)
        ins.append(run)
    return fn


def in_hyperlink(name: str, rid: str) -> Callable:
    """The picture's run inside a w:hyperlink of relationship ``rid`` (a clickable picture)."""
    def fn(root):
        run = drawing_named(root, name).getparent()
        link = etree.Element(W + "hyperlink", {R + "id": rid})
        run.addprevious(link)
        link.append(run)
    return fn


def format_change(name: str) -> Callable:
    """The picture's run bold, as someone's pending tracked formatting change (w:rPrChange)."""
    def fn(root):
        run = drawing_named(root, name).getparent()
        rpr = run.find(W + "rPr")
        if rpr is None:
            rpr = etree.Element(W + "rPr")
            run.insert(0, rpr)
        etree.SubElement(rpr, W + "b")
        change = etree.SubElement(rpr, W + "rPrChange", {W + "id": "800", W + "author": "Jeremy",
                                                         W + "date": "2026-09-22T10:15:00Z"})
        etree.SubElement(change, W + "rPr")
    return fn


def in_field(name: str) -> Callable:
    """The picture as the result of an INCLUDEPICTURE field (begin, code, separate — the picture's run — end)."""
    def fn(root):
        run = drawing_named(root, name).getparent()

        def r(xml: str):
            return etree.fromstring(f'<w:r xmlns:w="{W_NS}">{xml}</w:r>')
        run.addprevious(r('<w:fldChar w:fldCharType="begin"/>'))
        run.addprevious(r('<w:instrText xml:space="preserve"> INCLUDEPICTURE "site.png" \\* MERGEFORMAT </w:instrText>'))
        run.addprevious(r('<w:fldChar w:fldCharType="separate"/>'))
        run.addnext(r('<w:fldChar w:fldCharType="end"/>'))
    return fn


WPS_URI = "http://schemas.microsoft.com/office/word/2010/wordprocessingShape"


def in_textbox(text: str) -> Callable:
    """A text box (a DrawingML shape, no fallback) in the paragraph starting ``text``, holding a copy of the site plan
    (named "Boxed picture") in its own paragraph."""
    def fn(root):
        site = etree.fromstring(etree.tostring(drawing_named(root, "Picture 2")))
        dp = site.find(f"{WP}inline/{WP}docPr")
        dp.set("id", "950")
        dp.set("name", "Boxed picture")
        box = etree.fromstring(
            f'<w:r xmlns:w="{W_NS}" xmlns:wp="{WP[1:-1]}" xmlns:a="{A[1:-1]}" xmlns:wps="{WPS_URI}"><w:drawing>'
            f'<wp:inline><wp:extent cx="914400" cy="914400"/><wp:docPr id="951" name="Text Box 1"/><a:graphic>'
            f'<a:graphicData uri="{WPS_URI}"><wps:wsp><wps:txbx><w:txbxContent><w:p><w:r/></w:p></w:txbxContent>'
            f'</wps:txbx><wps:bodyPr/></wps:wsp></a:graphicData></a:graphic></wp:inline></w:drawing></w:r>')
        next(box.iter(W + "p")).find(W + "r").append(site)
        paragraph_with(root, text).append(box)
    return fn


def svg_blip(name: str, rid: str) -> Callable:
    def fn(root):
        blip = next(drawing_named(root, name).iter(A + "blip"))
        ext = etree.SubElement(etree.SubElement(blip, A + "extLst"), A + "ext",
                               uri="{96DAC541-7B7A-43D3-8B79-37D633B846F1}")
        etree.SubElement(ext, "{http://schemas.microsoft.com/office/drawing/2016/SVG/main}svgBlip",
                         {R + "embed": rid})
    return fn


def cell_setup(text: str, tcpr: Optional[str] = None, cell_mar: Optional[str] = None,
               style: Optional[str] = None) -> Callable:
    """The table cell holding the paragraph starting ``text``: its ``w:tcPr`` replaced by one holding ``tcpr`` (XML),
    its table's ``w:tblCellMar`` set to hold ``cell_mar`` (XML) and its ``w:tblStyle`` to ``style``, each when given."""
    def fn(root):
        tc = next(paragraph_with(root, text).iterancestors(W + "tc"))
        pr = next(tc.iterancestors(W + "tbl")).find(W + "tblPr")
        if tcpr is not None:
            if tc.find(W + "tcPr") is not None:
                tc.remove(tc.find(W + "tcPr"))
            tc.insert(0, etree.fromstring(f'<w:tcPr xmlns:w="{W_NS}">{tcpr}</w:tcPr>'))
        if cell_mar is not None:                        # CT_TblPr: … tblW, …, tblCellMar, tblLook
            pr.find(W + "tblLook").addprevious(etree.fromstring(f'<w:tblCellMar xmlns:w="{W_NS}">{cell_mar}'
                                                                f'</w:tblCellMar>'))
        if style is not None:
            pr.insert(0, etree.fromstring(f'<w:tblStyle xmlns:w="{W_NS}" w:val="{style}"/>'))
    return fn


def section_attrs(section: int, element: str, **values: Optional[str]) -> Callable:
    """Section ``section``'s ``w:pgSz`` or ``w:pgMar`` (``element``) gets the attributes ``values`` (None: removed)."""
    def fn(root):
        el = list(root.iter(W + "sectPr"))[section].find(W + element)
        for k, v in values.items():
            if v is None:
                el.attrib.pop(W + k, None)
            else:
                el.set(W + k, v)
    return fn


def line_spacing(style: Optional[str], rule: str) -> Callable:
    """styles.xml: the style ``style`` (None: the document defaults) gets the line spacing rule ``rule`` (its line 240,
    its spacing before and after kept)."""
    def fn(root):
        holder = (root.find(f"{W}docDefaults/{W}pPrDefault") if style is None
                  else next(s for s in root.iter(W + "style") if s.get(W + "styleId") == style))
        ppr = holder.find(W + "pPr")
        if ppr is None:                                 # CT_Style: … pPr, rPr, tblPr …
            ppr = etree.Element(W + "pPr")
            later = next((c for c in holder if c.tag in (W + "rPr", W + "tblPr", W + "trPr", W + "tcPr")), None)
            later.addprevious(ppr) if later is not None else holder.append(ppr)
        sp = ppr.find(W + "spacing")
        if sp is None:                                  # before the fixture styles' w:jc
            sp = etree.Element(W + "spacing")
            ppr.insert(0, sp)
        sp.set(W + "line", "240")
        sp.set(W + "lineRule", rule)
    return fn


def no_page_size(root) -> None:
    """Every section without w:pgSz (its size is Word's default, which the engine does not assume)."""
    for sz in list(root.iter(W + "pgSz")):
        sz.getparent().remove(sz)


def rels_of(docx: bytes, part: str) -> Dict[str, Tuple[str, str]]:
    """rId → (type's last word, target) of a part's relationships."""
    d, b = part.rsplit("/", 1)
    name = f"{d}/_rels/{b}.rels"
    z = zipfile.ZipFile(io.BytesIO(docx))
    if name not in z.namelist():
        return {}
    root = etree.fromstring(z.read(name))
    return {r.get("Id"): (r.get("Type").rsplit("/", 1)[-1], r.get("Target")) for r in root}
