"""Fixtures for move_paragraphs (built-in batch 6, F9 — docs/BUILTIN_OPS_BATCH6_DESIGN.md §10). Every paragraph is found
by its text in the tests, never by index.

``build_moving()`` — a specification in two sections, in reading order:

* "SECTION 21 13 13 - WET-PIPE SPRINKLER SYSTEMS" (with ``utc``: Jeremy's pending insertion " (rev)", w16du:dateUtc);
  with ``toc``, then a table of contents — a TOC field begun in "Summary ..... 1" (its instruction is block text too),
  showing "References ..... 1" and ended in "Submittals ..... 2";
* "PART 1 - GENERAL" (Heading 1); "1.1⇥SUMMARY" (Heading 2), "A.⇥Section includes wet-pipe sprinkler systems.";
* "1.2⇥REFERENCES", "A.⇥NFPA 13, 2025 edition." (with ``comment``: Jack's comment on "NFPA 13");
* "1.3⇥SUBMITTALS" — its words inside the bookmark _Ref_Submittals, w14:paraId 1A2B3C01 — then two List Number
  paragraphs (numbered by the style): "Product data for each sprinkler." (w14:paraId 1A2B3C02) and "Shop drawings,
  ⟨signed and sealed⟩." (the middle run bold, the paragraph indented 0.5 inch directly);
* "1.4⇥QUALITY ASSURANCE", "A.⇥Installer qualifications: licensed.", a 2×2 table ("Item" | "Requirement" /
  "Installer" | "Licensed"), "B.⇥Welding: AWS D10.9.";
* "1.5⇥WARRANTY" — with ``sdt`` followed by a block content control holding "Controlled text." (``data_bound``: bound
  to a custom XML part) — then "A.⇥Two years from Substantial Completion." (with ``hyperlink``: then a link "the
  warranty form");
* an empty paragraph carrying the first section's break (w:sectPr);
* "PART 2 - PRODUCTS" (Heading 1); "2.1⇥PIPE", "A.⇥Schedule 40 black steel." (with ``revision``: Jack's pending
  deletion " galvanized"; with ``textbox``: a text box run "Boxed note");
* "2.2⇥FITTINGS", "A.⇥Threaded or grooved, see " (with ``field``: then a REF field showing "21 10 00");
* table one (1×1, "TABLE ONE"); "Between the tables."; table two: its first cell holds "Cell first", "Cell second",
  "Cell third", its second cell "Other cell";
* "END OF SECTION".

The header holds "Header A", "Header B" and "Header C". ``split_bookmark``: the bookmark _Ref_Split starts in "1.2⇥
REFERENCES" and ends in "1.3⇥SUBMITTALS".

``with_article_bookmark`` puts a bookmark over article 1.2 whose marks sit between blocks or at a paragraph's start;
``with_comment_thread`` gives the comment three paragraphs and a thread entry naming the last.
"""
from __future__ import annotations

import io

from docx import Document
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from lxml import etree

from make_fixtures import _hyperlink, _ref_field, _rev, _textbox_run

W14 = "http://schemas.microsoft.com/office/word/2010/wordml"
W16DU = "http://schemas.microsoft.com/office/word/2023/wordml/word16du"


def save(doc) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _bookmark_around(paragraph, name: str, bid: str):
    """A bookmark around every run of the paragraph (as Word writes a cross-reference target)."""
    p = paragraph._p
    runs = p.findall(qn("w:r"))
    bs = OxmlElement("w:bookmarkStart"); bs.set(qn("w:id"), bid); bs.set(qn("w:name"), name)
    be = OxmlElement("w:bookmarkEnd"); be.set(qn("w:id"), bid)
    runs[0].addprevious(bs)
    runs[-1].addnext(be)


def _para_id(paragraph, value: str):
    paragraph._p.set("{%s}paraId" % W14, value)
    paragraph._p.set("{%s}textId" % W14, "77777777")


def _fld(paragraph, kind: str):
    fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), kind)
    paragraph.add_run()._r.append(fc)


def _toc(doc, entries):
    """A table of contents as Word writes one: a TOC field whose result runs over one paragraph per entry."""
    first = doc.add_paragraph()
    _fld(first, "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = ' TOC \\o "1-3" \\h '
    first.add_run()._r.append(it)
    _fld(first, "separate")
    first.add_run(entries[0])
    for text in entries[1:-1]:
        doc.add_paragraph(text)
    _fld(doc.add_paragraph(entries[-1]), "end")


def build_moving(*, utc: bool = False, toc: bool = False, comment: bool = False, hyperlink: bool = False,
                 sdt: bool = False, data_bound: bool = False, revision: bool = False, textbox: bool = False,
                 field: bool = False, split_bookmark: bool = False) -> bytes:
    doc = Document()
    hdr = doc.sections[0].header
    hdr.is_linked_to_previous = False
    hdr.paragraphs[0].text = "Header A"
    hdr.add_paragraph("Header B")
    hdr.add_paragraph("Header C")

    p = doc.add_paragraph("SECTION 21 13 13 - WET-PIPE SPRINKLER SYSTEMS")
    if utc:
        ins = _rev("ins", p.add_run(" (rev)"), 501, "Jeremy")
        ins.set("{%s}dateUtc" % W16DU, "2026-09-22T17:15:00Z")
    if toc:
        _toc(doc, ["Summary ..... 1", "References ..... 1", "Submittals ..... 2"])
    doc.add_paragraph("PART 1 - GENERAL", style="Heading 1")
    doc.add_paragraph("1.1\tSUMMARY", style="Heading 2")
    doc.add_paragraph("A.\tSection includes wet-pipe sprinkler systems.")
    references = doc.add_paragraph("1.2\tREFERENCES", style="Heading 2")
    p = doc.add_paragraph("A.\t")
    r = p.add_run("NFPA 13")
    p.add_run(", 2025 edition.")
    if comment:
        doc.add_comment(r, text="Confirm the edition.", author="Jack", initials="JK")
    submittals = doc.add_paragraph("1.3\tSUBMITTALS", style="Heading 2")
    _bookmark_around(submittals, "_Ref_Submittals", "11")
    _para_id(submittals, "1A2B3C01")
    _para_id(doc.add_paragraph("Product data for each sprinkler.", style="List Number"), "1A2B3C02")
    p = doc.add_paragraph("Shop drawings, ", style="List Number")
    p.add_run("signed and sealed").bold = True
    p.add_run(".")
    p.paragraph_format.left_indent = Inches(0.5)
    if split_bookmark:
        bs = OxmlElement("w:bookmarkStart"); bs.set(qn("w:id"), "12"); bs.set(qn("w:name"), "_Ref_Split")
        be = OxmlElement("w:bookmarkEnd"); be.set(qn("w:id"), "12")
        references._p.append(bs)
        submittals._p.append(be)
    doc.add_paragraph("1.4\tQUALITY ASSURANCE", style="Heading 2")
    doc.add_paragraph("A.\tInstaller qualifications: licensed.")
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "Item"; t.cell(0, 1).text = "Requirement"
    t.cell(1, 0).text = "Installer"; t.cell(1, 1).text = "Licensed"
    doc.add_paragraph("B.\tWelding: AWS D10.9.")
    doc.add_paragraph("1.5\tWARRANTY", style="Heading 2")
    if sdt:                                             # appended last: before the body's own (last) w:sectPr
        doc.element.body.find(qn("w:sectPr")).addprevious(_block_control("Controlled text.", data_bound))
    p = doc.add_paragraph("A.\tTwo years from Substantial Completion.")
    if hyperlink:
        p.add_run(" Use ")
        _hyperlink(p, "the warranty form")
    doc.add_section(WD_SECTION.CONTINUOUS)
    doc.add_paragraph("PART 2 - PRODUCTS", style="Heading 1")
    doc.add_paragraph("2.1\tPIPE", style="Heading 2")
    p = doc.add_paragraph("A.\tSchedule 40 black steel.")
    if revision:
        _rev("del", p.add_run(" galvanized"), 502, "Jack")
    if textbox:
        _textbox_run(p, "Boxed note")
    doc.add_paragraph("2.2\tFITTINGS", style="Heading 2")
    p = doc.add_paragraph("A.\tThreaded or grooved, see ")
    if field:
        _ref_field(p)
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "TABLE ONE"
    doc.add_paragraph("Between the tables.")
    t = doc.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "Cell first"
    t.cell(0, 0).add_paragraph("Cell second")
    t.cell(0, 0).add_paragraph("Cell third")
    t.cell(0, 1).text = "Other cell"
    doc.add_paragraph("END OF SECTION")
    return save(doc)


W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
W15 = "http://schemas.microsoft.com/office/word/2012/wordml"
MAIN = "word/document.xml"
COMMENTS = "word/comments.xml"


def _text(p) -> str:
    return "".join(t.text or "" for t in p.iter(W + "t"))


def _starting(root, text: str):
    return next(p for p in root.iter(W + "p") if _text(p).startswith(text))


def with_article_bookmark(data: bytes, first: str = "1.2", start: str = "in", end: str = "between") -> bytes:
    """A bookmark _Ref_Article (id 40) from the paragraph starting with ``first`` — inside it, right after its
    properties (``start="in"``), or right before it between blocks (``"between"``) — through "A.⇥NFPA 13, 2025
    edition.": its end right after that paragraph between blocks (``end="between"``), or at the start of "1.3⇥
    SUBMITTALS" (``"next"``, as Word ends a bookmark over whole paragraphs, their marks included)."""
    from make_fixtures import patch_docx

    def mark(root):
        head, nfpa = _starting(root, first), _starting(root, "A.NFPA")
        bs = etree.Element(W + "bookmarkStart"); bs.set(W + "id", "40"); bs.set(W + "name", "_Ref_Article")
        be = etree.Element(W + "bookmarkEnd"); be.set(W + "id", "40")
        if start == "in":
            head.insert(1 if head.find(W + "pPr") is not None else 0, bs)
        else:
            head.addprevious(bs)
        if end == "between":
            nfpa.addnext(be)
        else:
            nxt = _starting(root, "1.3SUBMITTALS")
            nxt.insert(1 if nxt.find(W + "pPr") is not None else 0, be)
    return patch_docx(data, {MAIN: mark})


def with_comment_thread(data: bytes, threaded: bool = True) -> bytes:
    """``build_moving(comment=True)`` with Jack's comment in three paragraphs — "Confirm the edition.", "Second note.",
    "Third note." (w14:paraId 2B000001–3) — and (``threaded``) a commentsExtended entry naming its last paragraph, as
    Word writes a comment's thread entry."""
    from make_fixtures import add_parts, patch_docx

    def three(root):
        c = root.find(W + "comment")
        for text in ("Second note.", "Third note."):
            p = etree.SubElement(c, W + "p")
            etree.SubElement(etree.SubElement(p, W + "r"), W + "t").text = text
        for k, p in enumerate(c.findall(W + "p")):
            p.set("{%s}paraId" % W14, f"2B00000{k + 1}")
            p.set("{%s}textId" % W14, "77777777")
    data = patch_docx(data, {COMMENTS: three})
    if not threaded:
        return data
    ex = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<w15:commentsEx xmlns:w15="{W15}">'
          f'<w15:commentEx w15:paraId="2B000003" w15:done="0"/></w15:commentsEx>').encode()
    return add_parts(data, {"word/commentsExtended.xml": ex},
                     rels=[("word/_rels/document.xml.rels", "rId891",
                            "http://schemas.microsoft.com/office/2011/relationships/commentsExtended",
                            "commentsExtended.xml")],
                     overrides={"word/commentsExtended.xml":
                                "application/vnd.openxmlformats-officedocument.wordprocessingml.commentsExtended+xml"})


def _block_control(text: str, data_bound: bool):
    w = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    binding = ('<w:dataBinding w:xpath="/root[1]/value[1]" w:storeItemID="{11111111-2222-3333-4444-555555555555}"/>'
               if data_bound else "")
    return etree.fromstring(f'<w:sdt xmlns:w="{w}"><w:sdtPr><w:alias w:val="Note"/>{binding}</w:sdtPr><w:sdtContent>'
                            f'<w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:sdtContent></w:sdt>')
