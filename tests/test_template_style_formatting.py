"""Style identities alone cannot authorize losing inherited template formatting."""
from io import BytesIO
import shutil

import pytest
from lxml import etree

from test_template import mutate, run
from test_template_styles import check_goal, effect, renamed, setup_case
from spec_surgeon.ops import JobException
from spec_surgeon.ooxml import STYLE_ORDER, insert_in_order
from spec_surgeon.package import DocxPackage, W
from spec_surgeon.verify.goal_templates import verify_template_effect
from spec_surgeon.verify.reader import VPackage


def style_properties(data, sid, xml):
    def edit(root):
        style = next(s for s in root if s.get(W + "styleId") == sid)
        prop = etree.fromstring(f'<wrapper xmlns:w="{W[1:-1]}">{xml}</wrapper>')[0]
        previous = style.find(prop.tag)
        if previous is not None:
            style.replace(previous, prop)
        else:
            insert_in_order(style, prop, STYLE_ORDER)
    return mutate(data, "word/styles.xml", edit)


@pytest.mark.parametrize("sid,xml", [
    ("Header", '<w:pPr><w:shd w:val="clear" w:fill="CCFFCC"/></w:pPr>'),
    ("Header", '<w:pPr><w:pBdr><w:bottom w:val="single" w:sz="18" w:color="FF0000"/></w:pBdr></w:pPr>'),
    ("Normal", '<w:pPr><w:shd w:val="clear" w:fill="CCFFCC"/></w:pPr>'),
    ("HeaderChar", '<w:rPr><w:u w:val="double"/></w:rPr>'),
    ("HeaderChar", '<w:rPr><w:shd w:val="clear" w:fill="CCFFCC"/></w:rPr>'),
    ("DefaultParagraphFont", '<w:rPr><w:u w:val="double"/></w:rPr>'),
    ("TableNormal", '<w:tcPr><w:shd w:val="clear" w:fill="CCFFCC"/></w:tcPr>'),
])
@pytest.mark.parametrize("styled_side", ["template", "destination"])
def test_unpinned_style_differences_are_isolated(tmp_path, sid, xml, styled_side):
    # Built-in DefaultParagraphFont/TableNormal ignore their properties in Word.
    # Custom default IDs exercise actual inherited formatting instead.
    custom = {"DefaultParagraphFont": "CustomDefaultChar", "TableNormal": "CustomDefaultTable"}
    template, source, values = setup_case(tmp_path, custom)
    source = renamed(source, custom)
    original_sid, sid = sid, custom.get(sid, sid)
    if styled_side == "template":
        template.write_bytes(style_properties(template.read_bytes(), sid, xml))
    else:
        source = style_properties(source, sid, xml)
    plan, output, record = run(source, values)
    assert record.passed, record.to_markdown()
    check_goal(source, output, template, values)
    assertion = next(a for a in plan.property_assertions if a["kind"] == "styles_added")
    mapped = assertion["style_map"][sid]
    assert mapped != sid
    before = DocxPackage(source).xml("word/styles.xml")
    after = DocxPackage(output).xml("word/styles.xml")
    old = next(s for s in before if s.get(W + "styleId") == sid)
    unchanged = next(s for s in after if s.get(W + "styleId") == sid)
    imported = next(s for s in after if s.get(W + "styleId") == mapped)
    wanted = next(s for s in DocxPackage(template.read_bytes()).xml("word/styles.xml")
                  if s.get(W + "styleId") == sid)
    assert etree.tostring(old, method="c14n") == etree.tostring(unchanged, method="c14n")
    for tag in ("pPr", "rPr", "tcPr"):
        actual_prop, expected_prop = imported.find(W + tag), wanted.find(W + tag)
        assert (etree.tostring(actual_prop, method="c14n", exclusive=True) if actual_prop is not None else None) == (
            etree.tostring(expected_prop, method="c14n", exclusive=True) if expected_prop is not None else None)
    assert imported.get(W + "default") is None
    copied = [a["part"] for a in plan.property_assertions if a["kind"] == "template_parts"]
    if original_sid == "DefaultParagraphFont":
        assert all(r.find(f"{W}rPr/{W}rStyle") is not None
                   for part in copied for r in DocxPackage(output).xml(part).iter(W + "r"))
    if original_sid == "Normal":
        assert all(p.find(f"{W}pPr/{W}pStyle") is not None
                   for part in copied for p in DocxPackage(output).xml(part).iter(W + "p"))
    if original_sid == "TableNormal":
        assert all(t.find(f"{W}tblPr/{W}tblStyle").get(W + "val") == mapped
                   for part in copied for t in DocxPackage(output).xml(part).iter(W + "tbl"))


@pytest.mark.parametrize("xml", [
    '<w:tblPr><w:tblBorders><w:bottom w:val="double" w:sz="18" w:color="FF0000"/></w:tblBorders></w:tblPr>',
    '<w:tblStylePr w:type="firstRow"><w:rPr><w:b/></w:rPr><w:tcPr><w:shd w:val="clear" w:fill="FFCC00"/></w:tcPr></w:tblStylePr>',
    '<w:rPr><w:sz w:val="36"/></w:rPr>',
])
def test_table_style_formatting_requires_complete_compatibility(tmp_path, xml):
    template, source, values = setup_case(tmp_path, {})
    def use_grid(root):
        table_properties = next(root.iter(W + "tblPr"))
        table_properties.insert(0, etree.Element(W + "tblStyle", {W + "val": "TableGrid"}))
    data = mutate(template.read_bytes(), "word/footer2.xml", use_grid)
    template.write_bytes(style_properties(data, "TableGrid", xml))
    plan, output, record = run(source, values)
    assert record.passed, record.to_markdown()
    check_goal(source, output, template, values)
    mapping = next(a for a in plan.property_assertions if a["kind"] == "template_parts")["style_map"]
    assert mapping["TableGrid"] != "TableGrid"
    table = next(s for s in DocxPackage(output).xml("word/styles.xml")
                 if s.get(W + "styleId") == mapping["TableGrid"])
    expected = etree.fromstring(f'<wrapper xmlns:w="{W[1:-1]}">{xml}</wrapper>')[0]
    assert etree.tostring(table.find(expected.tag), method="c14n", exclusive=True) == etree.tostring(expected, method="c14n", exclusive=True)


def test_nested_paragraph_style_cannot_rely_on_top_level_pins(tmp_path):
    template, source, values = setup_case(tmp_path, {})
    def textbox(root):
        outer = etree.SubElement(root, W + "p")
        content = etree.SubElement(etree.SubElement(etree.SubElement(outer, W + "r"), W + "pict"), W + "txbxContent")
        p = etree.SubElement(content, W + "p")
        etree.SubElement(etree.SubElement(p, W + "pPr"), W + "pStyle").set(W + "val", "Header")
        etree.SubElement(etree.SubElement(p, W + "r"), W + "t").text = "Nested template text"
    template.write_bytes(mutate(template.read_bytes(), "word/header1.xml", textbox))
    plan, output, record = run(source, values)
    assert record.passed, record.to_markdown()
    check_goal(source, output, template, values)
    mapping = next(a for a in plan.property_assertions if a["kind"] == "template_parts")["style_map"]
    assert mapping["Header"] != "Header"  # Header's source font size differs; nested content has no pins.


def test_disabled_import_does_not_allow_incompatible_name_reuse(tmp_path):
    template, source, values = setup_case(tmp_path, {})
    source = style_properties(source, "Header", '<w:pPr><w:shd w:val="clear" w:fill="FFCC00"/></w:pPr>')
    values["import_missing_styles"] = False
    with pytest.raises(JobException, match="no equivalent destination style"):
        run(source, values)


def test_verifiers_reject_a_transform_that_skips_formatting_compatibility(tmp_path, monkeypatch):
    import spec_surgeon.ops.template_styles as transform
    template, source, values = setup_case(tmp_path, {})
    source = style_properties(source, "Header", '<w:pPr><w:shd w:val="clear" w:fill="FFCC00"/></w:pPr>')
    monkeypatch.setattr(transform, "_formatting_signature", lambda *args: ())
    _, output, record = run(source, values)
    assert not record.passed and "input-derived mapping" in record.get("V10").summary
    before, after = VPackage(source), VPackage(output)
    try:
        assert verify_template_effect(before, after, effect(template, values))["errors"]
    finally:
        before.zip.close()
        after.zip.close()


def test_private_defaults_cannot_be_imported_without_binding_the_copied_content(tmp_path, monkeypatch):
    import spec_surgeon.ops.apply_header_footer_template as transform
    template, source, values = setup_case(tmp_path, {})
    template.write_bytes(style_properties(template.read_bytes(), "Normal", '<w:pPr><w:shd w:val="clear" w:fill="FFCC00"/></w:pPr>'))
    real = transform.remap_references
    monkeypatch.setattr(transform, "remap_references", lambda root, mapping, *args: real(root, mapping))
    _, output, record = run(source, values)
    assert not record.passed and record.get("V10").status == "FAIL"
    before, after = VPackage(source), VPackage(output)
    try:
        assert verify_template_effect(before, after, effect(template, values))["errors"]
    finally:
        before.zip.close()
        after.zip.close()


@pytest.mark.parametrize("sid,xml", [
    ("DefaultParagraphFont", '<w:rPr><w:u w:val="double"/><w:sz w:val="72"/></w:rPr>'),
    ("TableNormal", '<w:tcPr><w:shd w:val="clear" w:fill="CCFFCC"/></w:tcPr>'),
])
def test_word_inert_defaults_do_not_activate_formatting_or_dependencies(tmp_path, sid, xml):
    template, source, values = setup_case(tmp_path, {})
    data = style_properties(template.read_bytes(), sid, xml)
    template.write_bytes(style_properties(data, sid, '<w:basedOn w:val="IgnoredMissingBase"/>'))
    plan, output, record = run(source, values)
    assert record.passed, record.to_markdown()
    check_goal(source, output, template, values)
    mapping = next(a for a in plan.property_assertions if a["kind"] == "template_parts")["style_map"]
    assert mapping[sid] == sid
    assert "IgnoredMissingBase" not in mapping


def test_private_copy_of_word_inert_default_discards_inactive_children(tmp_path):
    template, source, values = setup_case(tmp_path, {})
    source = renamed(source, {"DefaultParagraphFont": "CustomDefaultChar"})
    source = style_properties(source, "CustomDefaultChar", '<w:rPr><w:u w:val="double"/></w:rPr>')
    def occupied(root):
        node = etree.SubElement(root, W + "style", {W + "type": "character", W + "styleId": "DefaultParagraphFont"})
        etree.SubElement(node, W + "name", {W + "val": "Occupied identity"})
    source = mutate(source, "word/styles.xml", occupied)
    data = style_properties(template.read_bytes(), "DefaultParagraphFont", '<w:rPr><w:u w:val="single"/></w:rPr>')
    template.write_bytes(style_properties(data, "DefaultParagraphFont", '<w:basedOn w:val="IgnoredMissingBase"/>'))
    plan, output, record = run(source, values)
    assert record.passed, record.to_markdown()
    check_goal(source, output, template, values)
    mapping = next(a for a in plan.property_assertions if a["kind"] == "template_parts")["style_map"]
    assert mapping["DefaultParagraphFont"] == "TemplateDefaultParagraphFont"
    imported = next(s for s in DocxPackage(output).xml("word/styles.xml")
                    if s.get(W + "styleId") == mapping["DefaultParagraphFont"])
    assert imported.find(W + "rPr") is None and imported.find(W + "basedOn") is None


@pytest.mark.render
@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")
def test_inherited_borders_and_shading_match_the_rendered_template(tmp_path):
    """Pixel evidence for the review regression and Word's inert built-in defaults."""
    from docx import Document
    from docx.shared import Inches, Pt
    import pymupdf

    from spec_surgeon.ooxml import PPR_ORDER
    from spec_surgeon.render import render_pdfs

    def document(is_template):
        doc = Document()
        doc.styles["Normal"].font.name = "Arial"
        doc.styles["Normal"].font.size = Pt(11)
        doc.add_paragraph("Preserve this body paragraph exactly.")
        section = doc.sections[0]
        section.header.paragraphs[0].text = "New project header" if is_template else "Old project header"
        section.footer.paragraphs[0].text = "New project footer" if is_template else "Old project footer"
        properties = doc.styles["Header"]._element.get_or_add_pPr()
        border = etree.Element(W + "pBdr")
        etree.SubElement(border, W + "bottom", {W + "val": "single", W + "sz": "18",
                                               W + "color": "0000FF" if is_template else "FF0000"})
        insert_in_order(properties, border, PPR_ORDER)
        insert_in_order(properties, etree.Element(W + "shd", {W + "val": "clear",
                        W + "fill": "CCFFCC" if is_template else "FFCCCC"}), PPR_ORDER)
        if is_template:
            doc.styles["Default Paragraph Font"].font.underline = True
            doc.styles["Default Paragraph Font"].font.size = Pt(36)
            table = section.header.add_table(rows=1, cols=2, width=Inches(6))
            table.cell(0, 0).text, table.cell(0, 1).text = "Template table left", "Template table right"
        tc = etree.Element(W + "tcPr")
        etree.SubElement(tc, W + "shd", {W + "val": "clear", W + "fill": "CCCCFF" if is_template else "FFFFCC"})
        insert_in_order(doc.styles["Normal Table"]._element, tc, STYLE_ORDER)
        data = BytesIO()
        doc.save(data)
        return data.getvalue()

    template, actual = tmp_path / "template.docx", tmp_path / "actual.docx"
    template.write_bytes(document(True))
    source = document(False)
    values = {"template": str(template), "template_sections": {"all": 0}, "slot_fills": {}}
    _, output, record = run(source, values)
    assert record.passed, record.to_markdown()
    check_goal(source, output, template, values)
    actual.write_bytes(output)
    pdfs = render_pdfs([template, actual], tmp_path / "pdfs", "libreoffice", workers=1)
    assert set(pdfs) == {template, actual}
    with pymupdf.open(pdfs[template]) as desired, pymupdf.open(pdfs[actual]) as result:
        assert len(desired) == len(result) == 1
        a, b = desired[0].get_pixmap(), result[0].get_pixmap()
        assert (a.width, a.height, a.samples) == (b.width, b.height, b.samples)
