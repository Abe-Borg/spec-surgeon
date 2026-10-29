"""Fixtures for built-in batch 2's removal ops — remove_hidden_text, remove_notes, remove_empty_paragraphs
(docs/BUILTIN_OPS_BATCH2_DESIGN.md §5, §6, §8); ``build_crossing_marks`` — bookmarks and a permission range running out
of what they delete into what stays. Every paragraph is found by its text in the tests, never by index."""
from __future__ import annotations

import copy
import io

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor
from lxml import etree

from make_fixtures import _hyperlink, _rev

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W = "{%s}" % W_NS
BLUE = RGBColor(0, 0, 0xFF)


def save(doc) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def hide(run, on: bool = True):
    run.font.hidden = on
    return run


def hide_mark(paragraph):
    rpr = paragraph._p.get_or_add_pPr().find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        paragraph._p.get_or_add_pPr().append(rpr)
    rpr.append(OxmlElement("w:vanish"))


def field(paragraph, instr: str, result=None, hidden=(), rpr_hidden_all: bool = False):
    """A complex field: begin, instruction, [separate, result], end — each its own run; ``hidden`` names the pieces
    (begin, instr, separate, result, end) whose run carries w:vanish (or all with ``rpr_hidden_all``)."""
    pieces = [("begin", "fldChar", "begin"), ("instr", "instrText", instr)]
    if result is not None:
        pieces += [("separate", "fldChar", "separate"), ("result", "t", result)]
    pieces.append(("end", "fldChar", "end"))
    for name, tag, val in pieces:
        r = paragraph.add_run()
        if tag == "fldChar":
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), val)
        else:
            el = OxmlElement(f"w:{tag}"); el.text = val; el.set(qn("xml:space"), "preserve")
        r._r.append(el)
        if rpr_hidden_all or name in hidden:
            hide(r)


def simple_field(paragraph, instr: str, result: str, hidden: bool = True):
    fs = OxmlElement("w:fldSimple"); fs.set(qn("w:instr"), instr)
    r = OxmlElement("w:r")
    if hidden:
        rpr = OxmlElement("w:rPr"); rpr.append(OxmlElement("w:vanish")); r.append(rpr)
    t = OxmlElement("w:t"); t.text = result; r.append(t)
    fs.append(r)
    paragraph._p.append(fs)


def tab_run(paragraph, hidden: bool = True):
    r = paragraph.add_run()
    r._r.append(OxmlElement("w:tab"))
    if hidden:
        hide(r)
    return r


def textbox_run(paragraph, text: str, hidden: bool = True):
    v = '<w:rPr><w:vanish/></w:rPr>' if hidden else ''
    xml = (
        f'<w:r xmlns:w="{W_NS}" xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
        'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
        'xmlns:v="urn:schemas-microsoft-com:vml">'
        '<mc:AlternateContent><mc:Choice Requires="wps"><w:drawing><wp:inline><wp:extent cx="914400" cy="457200"/>'
        '<wp:docPr id="901" name="Text Box 1"/><a:graphic><a:graphicData '
        'uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"><wps:wsp><wps:txbx><w:txbxContent>'
        f'<w:p><w:r>{v}<w:t>{text}</w:t></w:r></w:p></w:txbxContent></wps:txbx><wps:bodyPr/></wps:wsp></a:graphicData>'
        '</a:graphic></wp:inline></w:drawing></mc:Choice><mc:Fallback><w:pict><v:shape><v:textbox><w:txbxContent>'
        f'<w:p><w:r>{v}<w:t>{text}</w:t></w:r></w:p></w:txbxContent></v:textbox></v:shape></w:pict></mc:Fallback>'
        '</mc:AlternateContent></w:r>')
    paragraph._p.append(etree.fromstring(xml))


# ---- hidden text -----------------------------------------------------------------------------------------------------

def build_hidden(*, header: bool = True, cell: bool = True, insertion: bool = True, link: bool = True,
                 xe: bool = True) -> bytes:
    """Hidden text of every shape the op handles:

    * "Visible lead. ⟨secret words⟩ visible tail." — one directly hidden run;
    * "Plain text ⟨char-styled secret⟩" — hidden by the SecretChar character style;
    * a HiddenNote-styled paragraph (mark and text hidden by the style) — deleted;
    * an empty paragraph whose mark is hidden directly — deleted;
    * "Before⟨tab⟩After" with the tab run hidden — a structure-only change untracked;
    * "Index entry here" + a wholly hidden XE field (``xe``);
    * "Shown mark hidden" — a hidden mark with visible text: nothing goes;
    * "Inserted: ⟨kept⟩⟨hidden inside⟩" — a hidden run inside Jeremy's pending insertion (``insertion``);
    * "See ⟨link⟩" — a hidden run inside a hyperlink (``link``);
    * a table cell whose only paragraph is a HiddenNote paragraph (``cell``) — its mark must stay;
    * a header run "Header ⟨hidden header words⟩" (``header``).
    """
    doc = Document()
    st = doc.styles.add_style("HiddenNote", WD_STYLE_TYPE.PARAGRAPH)
    st.base_style = doc.styles["Normal"]
    st.font.hidden = True
    cs = doc.styles.add_style("SecretChar", WD_STYLE_TYPE.CHARACTER)
    cs.font.hidden = True
    if header:
        hp = doc.sections[0].header.paragraphs[0]
        hp.add_run("Header ")
        hide(hp.add_run("hidden header words"))
    doc.add_paragraph("SECTION 21 13 13 - WET-PIPE SPRINKLER SYSTEMS")
    p = doc.add_paragraph()
    p.add_run("Visible lead. ")
    hide(p.add_run("secret words"))
    p.add_run(" visible tail.")
    p = doc.add_paragraph()
    p.add_run("Plain text ")
    p.add_run("char-styled secret").style = "SecretChar"
    doc.add_paragraph("Specifier: this whole paragraph is hidden.", style="HiddenNote")
    p = doc.add_paragraph()
    hide_mark(p)                                                        # an empty paragraph with a hidden mark
    p = doc.add_paragraph()
    p.add_run("Before")
    tab_run(p)
    p.add_run("After")
    if xe:
        p = doc.add_paragraph()
        p.add_run("Index entry here")
        field(p, ' XE "Sprinkler" ', rpr_hidden_all=True)
    p = doc.add_paragraph("Shown mark hidden")
    hide_mark(p)
    if insertion:
        p = doc.add_paragraph()
        p.add_run("Inserted: ")
        r1 = p.add_run("kept ")
        r2 = hide(p.add_run("hidden inside"))
        ins = _rev("ins", r1, 301, "Jeremy")
        ins.append(r2._r)
    if link:
        p = doc.add_paragraph()
        p.add_run("See ")
        _hyperlink(p, "the portal")
        p._p.find(qn("w:hyperlink")).find(qn("w:r")).insert(0, _vanish_rpr())
        p.add_run(" now.")
    if cell:
        t = doc.add_table(rows=1, cols=2)
        t.cell(0, 0).text = "Cell text"
        cp = t.cell(0, 1).paragraphs[0]
        cp.style = "HiddenNote"
        cp.add_run("hidden cell note")
    doc.add_paragraph("END OF SECTION")
    return save(doc)


def _vanish_rpr():
    rpr = OxmlElement("w:rPr")
    rpr.append(OxmlElement("w:vanish"))
    return rpr


def build_hidden_case(kind: str) -> bytes:
    """One awkward case each: ambiguous (a run hidden by its paragraph style AND its character style), partial-field (a
    field whose result alone is hidden), textbox (hidden text in a text box), simple (a hidden simple field), plain (no
    hidden text at all), section (a hidden paragraph carrying a section break)."""
    doc = Document()
    doc.add_paragraph("Intro paragraph.")
    if kind == "ambiguous":
        st = doc.styles.add_style("HiddenPara", WD_STYLE_TYPE.PARAGRAPH)
        st.font.hidden = True
        cs = doc.styles.add_style("SecretChar", WD_STYLE_TYPE.CHARACTER)
        cs.font.hidden = True
        p = doc.add_paragraph("visible? ", style="HiddenPara")
        p.add_run("both levels").style = "SecretChar"
    elif kind == "partial-field":
        p = doc.add_paragraph("See ")
        field(p, " REF _Ref1 \\h ", "21 10 00", hidden=("result",))
    elif kind == "textbox":
        p = doc.add_paragraph("A box: ")
        textbox_run(p, "boxed secret", hidden=False)
        p2 = doc.add_paragraph("Also a box: ")
        textbox_run(p2, "hidden in a box", hidden=True)
    elif kind == "simple":
        p = doc.add_paragraph("Title: ")
        simple_field(p, " DOCPROPERTY Title ", "Spec title", hidden=True)
        p.add_run(" end.")
    elif kind == "section":
        p = doc.add_paragraph("hidden end of the first section")
        hide(p.runs[0])
        hide_mark(p)
        p._p.get_or_add_pPr().append(copy.deepcopy(doc.element.body.find(qn("w:sectPr"))))
        doc.add_paragraph("Body.")
    elif kind == "plain":
        doc.add_paragraph("Nothing hidden here.")
    return save(doc)


# ---- notes -----------------------------------------------------------------------------------------------------------

def build_notes(*, cell_note: bool = True, field_note: bool = False) -> bytes:
    """Specifier notes of every shape remove_notes handles:

    * "Retain paragraph below if …" in the CMT style (blue by the style) — a note by style (and by text);
    * "Coordinate with the Owner's insurer." — every run blue (direct) — a note by colour;
    * "Mixed: black and ⟨blue⟩" — not a colour note (one black run), its blue run an inline note;
    * "Pipe shall be ⟨[Verify size]⟩ Schedule 40." — an inline blue note;
    * "See ⟨the portal⟩ for data." — a blue hyperlink (the Hyperlink style is blue): never an inline note;
    * "Delete this paragraph when …" in the normal style, black — a note by text only;
    * a table cell whose only paragraph is a CMT note (``cell_note``) — kept, its words removed;
    * a CMT note holding a field (``field_note``) — the op must refuse.
    """
    doc = Document()
    cmt = doc.styles.add_style("CMT", WD_STYLE_TYPE.PARAGRAPH)
    cmt.base_style = doc.styles["Normal"]
    cmt.font.color.rgb = BLUE
    doc.add_paragraph("PART 1 - GENERAL")
    doc.add_paragraph("Retain paragraph below if the project has a fire pump.", style="CMT")
    doc.add_paragraph("A.\tProvide a fire pump.")
    p = doc.add_paragraph()
    r = p.add_run("Coordinate with the Owner's insurer.")
    r.font.color.rgb = BLUE
    p = doc.add_paragraph()
    p.add_run("Mixed: black and ")
    r = p.add_run("blue"); r.font.color.rgb = BLUE
    p = doc.add_paragraph()
    p.add_run("Pipe shall be ")
    r = p.add_run("[Verify size]"); r.font.color.rgb = BLUE
    p.add_run(" Schedule 40.")
    p = doc.add_paragraph()
    p.add_run("See ")
    _hyperlink(p, "the portal")
    link_run = p._p.find(qn("w:hyperlink")).find(qn("w:r"))
    rpr = OxmlElement("w:rPr"); c = OxmlElement("w:color"); c.set(qn("w:val"), "0000FF"); rpr.append(c)
    link_run.insert(0, rpr)
    p.add_run(" for data.")
    doc.add_paragraph("Delete this paragraph when the system is wet-pipe only.")
    if cell_note:
        t = doc.add_table(rows=1, cols=2)
        t.cell(0, 0).text = "Item"
        t.cell(0, 1).paragraphs[0].style = "CMT"
        t.cell(0, 1).paragraphs[0].add_run("Edit the value to suit.")
    if field_note:
        p = doc.add_paragraph("Retain with ", style="CMT")
        field(p, " REF _Ref9 \\h ", "21 10 00")
    doc.add_paragraph("END OF SECTION")
    return save(doc)


def build_last_section(kind: str, *, hidden_break: bool = False, field_in: str = "") -> bytes:
    """Three sections, each break on a paragraph of its own (empty, as Word and python-docx write it): "Cover", a break;
    "Body text", a break; then the last section — "NOTE: one" and "NOTE: two" and nothing else. ``kind`` "notes": plain
    paragraphs (notes by the text rule ``^NOTE``); "hidden": their runs and marks hidden (hidden paragraphs).
    ``hidden_break``: the second break's paragraph has a hidden mark too — a hidden paragraph that carries a section
    break, which stays. ``field_in`` ("one" or "two"): that paragraph ends with a REF field (it cannot be deleted)."""
    doc = Document()
    doc.add_paragraph("Cover")
    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.add_paragraph("Body text")
    doc.add_section(WD_SECTION.NEW_PAGE)
    if hidden_break:
        ppr = doc.paragraphs[-1]._p.get_or_add_pPr()
        ppr.find(qn("w:sectPr")).addprevious(_vanish_rpr())            # w:rPr before w:sectPr (CT_PPr's order)
    for word in ("one", "two"):
        p = doc.add_paragraph(f"NOTE: {word}")
        if field_in == word:
            field(p, " REF _Ref9 \\h ", "21 10 00", rpr_hidden_all=kind == "hidden")
        if kind == "hidden":
            hide(p.runs[0])
            hide_mark(p)
    return save(doc)


# ---- range marks running out of what goes ----------------------------------------------------------------------------

def _mark(paragraph, tag: str, mid: str, name: str = "", inside=None):
    """A range mark (``w:bookmarkStart`` with ``name``, ``w:bookmarkEnd``, ``w:permStart``, ``w:permEnd``) at the end of
    the paragraph — or of ``inside`` (an element of it)."""
    el = OxmlElement(f"w:{tag}"); el.set(qn("w:id"), mid)
    if name:
        el.set(qn("w:name"), name)
    if tag == "permStart":
        el.set(qn("w:edGrp"), "everyone")
    (paragraph._p if inside is None else inside).append(el)


def build_crossing_marks(kind: str, *, simple: bool = False, permission: bool = True) -> bytes:
    """Bookmarks and a permission range (unless ``permission`` is false) running out of the paragraphs a removal deletes
    into paragraphs that stay. ``kind`` "notes": the two paragraphs to delete are notes by the text rule ``^NOTE``;
    "hidden": their runs and marks are hidden. In reading order:

    * "Lead." — the start of "Back";
    * "NOTE: starts the bookmark." (deleted) — all of "Inner", then the start of "B", the start of "Wholly", and a
      permission range's start;
    * "Visible text." — the end of "B" and of the permission range;
    * "NOTE: ends two bookmarks." (deleted) — the end of "Wholly" (wholly in the deleted text: it goes) and of "Back";
    * "See ⟨REF B⟩ and ⟨REF Back⟩." — cross-references to the two bookmarks whose other end stays;
    * with ``simple``: "Title: ⟨a hidden DOCPROPERTY simple field⟩ end." whose field holds all of "InField" and the
      start of "FieldMark", whose end is in "After the field."; then "END OF SECTION"."""
    doc = Document()
    lead = doc.add_paragraph("Lead.")
    first = doc.add_paragraph("NOTE: starts the bookmark.")
    middle = doc.add_paragraph("Visible text.")
    second = doc.add_paragraph("NOTE: ends two bookmarks.")
    see = doc.add_paragraph("See ")
    field(see, " REF B \\h ", "Visible text.")
    see.add_run(" and ")
    field(see, " REF Back \\h ", "Lead.")
    see.add_run(".")
    _mark(lead, "bookmarkStart", "4", "Back")
    _mark(first, "bookmarkStart", "2", "Inner")
    _mark(first, "bookmarkEnd", "2")
    _mark(first, "bookmarkStart", "1", "B")
    _mark(first, "bookmarkStart", "3", "Wholly")
    _mark(middle, "bookmarkEnd", "1")
    if permission:
        _mark(first, "permStart", "5")
        _mark(middle, "permEnd", "5")
    _mark(second, "bookmarkEnd", "3")
    _mark(second, "bookmarkEnd", "4")
    if kind == "hidden":
        for p in (first, second):
            hide(p.runs[0])
            hide_mark(p)
    if simple:
        title = doc.add_paragraph("Title: ")
        simple_field(title, " DOCPROPERTY Title ", "Spec title", hidden=True)
        fs = title._p.find(qn("w:fldSimple"))
        _mark(title, "bookmarkStart", "7", "InField", inside=fs)
        _mark(title, "bookmarkEnd", "7", inside=fs)
        _mark(title, "bookmarkStart", "6", "FieldMark", inside=fs)
        fs.append(fs.find(qn("w:r")))                                  # the marks first, then the field's result
        title.add_run(" end.")
        after = doc.add_paragraph("After the field.")
        _mark(after, "bookmarkEnd", "6")
    doc.add_paragraph("END OF SECTION")
    return save(doc)


# ---- blank paragraphs ------------------------------------------------------------------------------------------------

def _bookmark(paragraph, name: str = "keep_me", bid: str = "7"):
    bs = OxmlElement("w:bookmarkStart"); bs.set(qn("w:id"), bid); bs.set(qn("w:name"), name)
    be = OxmlElement("w:bookmarkEnd"); be.set(qn("w:id"), bid)
    paragraph._p.append(bs)
    paragraph._p.append(be)


def _numbered(paragraph, num_id: str = "1"):
    ppr = paragraph._p.get_or_add_pPr()
    np_ = OxmlElement("w:numPr")
    il = OxmlElement("w:ilvl"); il.set(qn("w:val"), "0"); np_.append(il)
    ni = OxmlElement("w:numId"); ni.set(qn("w:val"), num_id); np_.append(ni)
    ppr.append(np_)


def build_blank(*, header_blanks: bool = True) -> bytes:
    """Blank paragraphs of every shape remove_empty_paragraphs meets, in reading order:

    "Para A", three blanks, "Para B", a spaces-only paragraph, an NBSP-only paragraph, "Para C", an empty numbered
    paragraph (numId 1), an empty paragraph of the NumberedStyle style (numbering from the style), an empty paragraph whose
    own numId 0 switches the style's numbering off, an empty paragraph with a bookmark, "Para D", a tab-only paragraph, a
    page-break-only paragraph, "Para E", table 1, one blank, table 2, "Para F", a blank carrying a section break, "Para
    G", a bordered blank (BoxStyle), "Para H", a cell holding "X" then a blank, a cell holding only a blank, two trailing
    blanks at the end of the body. ``header_blanks``: the header holds "Header text" and two blanks.
    """
    doc = Document()
    ns = doc.styles.add_style("NumberedStyle", WD_STYLE_TYPE.PARAGRAPH)
    _numbered_style(ns)
    box = doc.styles.add_style("BoxStyle", WD_STYLE_TYPE.PARAGRAPH)
    _boxed_style(box)
    if header_blanks:
        hdr = doc.sections[0].header
        hdr.paragraphs[0].add_run("Header text")
        hdr.add_paragraph("")
        hdr.add_paragraph("")
    doc.add_paragraph("Para A")
    for _ in range(3):
        doc.add_paragraph("")
    doc.add_paragraph("Para B")
    doc.add_paragraph("   ")
    doc.add_paragraph("  ")
    doc.add_paragraph("Para C")
    _numbered(doc.add_paragraph(""))
    doc.add_paragraph("", style="NumberedStyle")
    p = doc.add_paragraph("", style="NumberedStyle")
    _numbered(p, "0")                                                   # numbering switched off: blank
    _bookmark(doc.add_paragraph(""))
    doc.add_paragraph("Para D")
    p = doc.add_paragraph()
    tab_run(p, hidden=False)
    p = doc.add_paragraph()
    br = OxmlElement("w:br"); br.set(qn("w:type"), "page"); p.add_run()._r.append(br)
    doc.add_paragraph("Para E")
    t1 = doc.add_table(rows=1, cols=1)
    t1.cell(0, 0).text = "T1"
    doc.add_paragraph("")
    t2 = doc.add_table(rows=1, cols=1)
    t2.cell(0, 0).text = "T2"
    doc.add_paragraph("Para F")
    p = doc.add_paragraph("")
    p._p.get_or_add_pPr().append(copy.deepcopy(doc.element.body.find(qn("w:sectPr"))))
    doc.add_paragraph("Para G")
    doc.add_paragraph("", style="BoxStyle")
    doc.add_paragraph("Para H")
    t3 = doc.add_table(rows=1, cols=2)
    t3.cell(0, 0).text = "X"
    t3.cell(0, 0).add_paragraph("")
    doc.add_paragraph("")
    doc.add_paragraph("")
    return save(doc)


def build_blank_last_section() -> bytes:
    """Three sections, each break on a paragraph of its own (empty, as Word and python-docx write it): "Cover", a break;
    "Body text", a blank, a break; then the last section — three blank paragraphs and nothing else."""
    doc = Document()
    doc.add_paragraph("Cover")
    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.add_paragraph("Body text")
    doc.add_paragraph("")
    doc.add_section(WD_SECTION.NEW_PAGE)
    for _ in range(3):
        doc.add_paragraph("")
    return save(doc)


def _numbered_style(style):
    ppr = style.element.get_or_add_pPr()
    np_ = OxmlElement("w:numPr")
    ni = OxmlElement("w:numId"); ni.set(qn("w:val"), "1"); np_.append(ni)
    ppr.append(np_)


def _boxed_style(style):
    ppr = style.element.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    top = OxmlElement("w:top"); top.set(qn("w:val"), "single"); top.set(qn("w:sz"), "4"); top.set(qn("w:space"), "1")
    top.set(qn("w:color"), "auto")
    bdr.append(top)
    ppr.append(bdr)


def build_spec_notes(*, field_note: str = "") -> bytes:
    """A specification as a master leaves it, for Safe mode's Remove specifier notes button (recipes._remove_notes):

    * notes by style — "Retain paragraph below …" in ``CMT`` (blue by the style); "Coordinate the alarm …" in
      ``SpecifierNote`` (named "Specifier Note"); "Edit to suit." in ``CMTChild`` (based on CMT);
    * notes by their opening words, in the body style — "NOTE TO SPECIFIER: …", "Retain or delete this article.",
      "Select one of the following paragraphs.", "[Editor's Note: …]";
    * hidden text — a paragraph whose runs and mark are hidden, and a hidden run inside a requirement;
    * what stays — the headings (SECTION, PART, ARTICLE, END OF SECTION), requirements that open with a note's word
      ("Retain records …", "Select one of the listed manufacturers.", "Coordinate the work …"), placeholders
      ("[Verify quantity with Owner]", "<Insert manufacturer>", "[TBD]"), "to be determined", a copyright line, a
      red italic requirement, a requirement holding a note marker mid-sentence;
    * ``field_note``: a CMT note holding a REF field — "visible": remove_notes cannot delete it, so the document is
      set aside; "hidden" (its runs, the field's and its mark hidden, as a master hides its notes): the hidden-text
      step removes the field with the words and keeps the mark, and the notes step then deletes the empty note.
    """
    doc = Document()
    cmt = doc.styles.add_style("CMT", WD_STYLE_TYPE.PARAGRAPH)
    cmt.base_style = doc.styles["Normal"]
    cmt.font.color.rgb = BLUE
    child = doc.styles.add_style("CMTChild", WD_STYLE_TYPE.PARAGRAPH)
    child.base_style = cmt
    doc.styles.add_style("Specifier Note", WD_STYLE_TYPE.PARAGRAPH).base_style = doc.styles["Normal"]
    doc.add_paragraph("SECTION 21 13 13 - WET-PIPE SPRINKLER SYSTEMS")
    doc.add_paragraph("© 2026 Owner Data Centers LLC. All rights reserved. Confidential.")
    doc.add_paragraph("PART 1 - GENERAL")
    doc.add_paragraph("Retain paragraph below if the project has a fire pump.", style="CMT")
    doc.add_paragraph("A.\tProvide a fire pump.")
    doc.add_paragraph("Coordinate the alarm with the fire alarm Section.", style="Specifier Note")
    doc.add_paragraph("Coordinate the work with the Owner's insurer.")
    doc.add_paragraph("NOTE TO SPECIFIER: edit the list to suit the project.")
    doc.add_paragraph("Retain records of all tests required in Paragraph 1.6.")
    doc.add_paragraph("ARTICLE 1.2 - SUBMITTALS")
    doc.add_paragraph("Retain or delete this article.")
    doc.add_paragraph("Select one of the following paragraphs.")
    doc.add_paragraph("Select one of the listed manufacturers.")
    doc.add_paragraph("[Editor's Note: verify the hazard classification with the Owner.]")
    doc.add_paragraph("Edit to suit.", style="CMTChild")
    doc.add_paragraph("Provide two [Verify quantity with Owner] spare sprinklers per type.")
    doc.add_paragraph("Provide <Insert manufacturer> valves.")
    doc.add_paragraph("System pressure: [TBD].")
    doc.add_paragraph("Hydrant flow test date to be determined.")
    doc.add_paragraph("Provide two [retain or delete] spare heads per cabinet.")
    p = doc.add_paragraph()
    r = p.add_run("Provide galvanized pipe in the preaction zones.")
    r.font.color.rgb = RGBColor(0xFF, 0, 0)
    r.font.italic = True
    p = doc.add_paragraph()
    hide(p.add_run("Specifier: this whole paragraph is hidden."))
    hide_mark(p)
    p = doc.add_paragraph()
    p.add_run("Pipe shall be Schedule 40")
    hide(p.add_run(" (retain for 2-inch and smaller)"))
    p.add_run(".")
    if field_note:
        p = doc.add_paragraph(style="CMT")
        hidden = field_note == "hidden"
        r = p.add_run("Retain with Section ")
        field(p, " REF _Ref9 \\h ", "21 10 00", rpr_hidden_all=hidden)
        if hidden:
            hide(r)
            hide_mark(p)
    doc.add_paragraph("END OF SECTION 21 13 13")
    return save(doc)
