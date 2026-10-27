"""Delivery and approval capabilities stay bound to the inputs and run the user reviewed."""
from __future__ import annotations

import json
import os
import copy
import hashlib
from pathlib import Path

import pytest
from docx import Document

from make_fixtures import build, build_batch1, build_sized, build_style_template, build_styles, png
from spec_surgeon import batch as batch_mod, batch as batch_module, guard, preview, verify
from spec_surgeon.batch import FileResult, check_inputs_outside, load_set_aside, output_names, run_batch, verified_bytes
from spec_surgeon.chain import DRIFT_CODE, chain_from_specs, chain_key, run_chain
from spec_surgeon.goals import resolve_goal
from spec_surgeon.ooxml import block_full_text, paragraphs
from spec_surgeon.package import DocxPackage, W
from spec_surgeon.verify import CheckResult


def font(name="Courier New"):
    return {"request": "Set the font", "summary": "Set the font", "mode": "formatting", "ops": [
        {"op": "set_font", "select": {"scope": ["document"], "runs": "all"}, "params": {"font": name}}]}


def body(path):
    pkg = DocxPackage.from_path(path)
    return "".join(t.text or "" for t in pkg.xml(pkg.main_part()).iter(W + "t"))


@pytest.mark.parametrize("operation", ["template", "table", "picture"])
@pytest.mark.parametrize("surface", ["batch", "chain", "preview", "chain_preview"])
def test_results_overlap_refuses_every_input_before_clearing_or_writing_it(tmp_path, operation, surface):
    source = tmp_path / "source.docx"
    source.write_bytes(build_styles())
    out = tmp_path / "run"
    artifact = out / "outputs" / "source.docx"
    artifact.parent.mkdir(parents=True)
    original = build_style_template() if operation == "template" else b"external input must survive"
    artifact.write_bytes(original)
    if operation == "template":
        op = {"op": "apply_style", "params": {"style": "Heading 1", "template": str(artifact)}}
    elif operation == "table":
        op = {"op": "fill_variables", "params": {"table": {"file": str(artifact)}}}
    else:
        op = {"op": "swap_image", "params": {"image": str(artifact), "match": {"image": str(artifact)}}}
    spec = {"request": "Use the supplied input", "summary": "Use the input",
            "mode": "content" if operation == "table" else "formatting", "ops": [op]}
    chain = chain_from_specs([("edit", spec)])
    with pytest.raises(ValueError, match="input file.*inside the results folder"):
        if surface == "batch":
            run_batch([str(source)], "edit-spec", {}, out, spec=spec, workers=1, do_render=False)
        elif surface == "chain":
            run_chain([str(source)], chain, out, workers=1, do_render=False)
        elif surface == "preview":
            preview.make_preview(spec, [str(source)], out, do_render=False)
        else:
            preview.make_chain_preview(chain, [str(source)], out, do_render=False)
    assert artifact.read_bytes() == original


def test_reused_batch_results_contain_no_previous_successes_when_new_run_refuses(tmp_path):
    sources = [tmp_path / name for name in ("one.docx", "two.docx")]
    for source in sources:
        source.write_bytes(build_sized())
    out = tmp_path / "run"
    first = run_batch([str(p) for p in sources], "edit-spec", {}, out,
                      spec=font(), workers=1, do_render=False)
    assert first.counts()["PASS"] == 2
    stale_failed = out / "failed" / "earlier.docx"
    stale_failed.parent.mkdir()
    stale_failed.write_bytes(b"previous rejected result")
    refused = {"request": "Replace absent wording", "summary": "Replace wording", "mode": "content", "ops": [
        {"op": "replace_text", "params": {"find": "THIS PHRASE IS ABSENT", "replace": "X", "tracked": True}}]}
    second = run_batch([str(sources[0])], "edit-spec", {}, out,
                       spec=refused, workers=1, do_render=False)
    assert second.counts()["EXCEPTION"] == 1
    assert not list((out / "outputs").iterdir())
    assert not stale_failed.exists()


@pytest.mark.parametrize("chain_preview", [False, True])
@pytest.mark.parametrize("changed_input", ["source", "template"])
def test_saves_during_preview_never_approve_the_new_unpreviewed_input(tmp_path, monkeypatch,
                                                                    chain_preview, changed_input):
    source = tmp_path / "source.docx"
    source.write_bytes(build_styles() if changed_input == "template" else build_sized(title="OLD SOURCE"))
    if changed_input == "template":
        changed = tmp_path / "template.docx"
        changed.write_bytes(build_style_template())
        spec = {"request": "Import the heading style", "summary": "Import style", "mode": "formatting", "ops": [
            {"op": "apply_style", "select": {"paragraphs": {"text_regex": "PART"}},
             "params": {"style": "Heading 1", "template": str(changed), "existing": "replace", "tracked": False}}]}
    else:
        changed, spec = source, font()
    chain = chain_from_specs([("edit", spec)])
    key = chain_key(chain) if chain_preview else spec
    before = preview.approval_token(key, [source])
    before_inputs = preview.approval_inputs(key, [source])

    def user_saves_during_render(*args, **kwargs):
        changed.write_bytes(build_sized(title="NEW UNPREVIEWED SOURCE"))
        return None

    monkeypatch.setattr(preview, "render_samples", user_saves_during_render)
    maker = preview.make_chain_preview if chain_preview else preview.make_preview
    pv = maker(chain if chain_preview else spec, [str(source)], tmp_path / "preview", do_render=False)
    assert pv.samples[0].passed
    if changed_input == "source":
        assert "OLD SOURCE" in body(pv.samples[0].output)
    assert not pv.token
    assert pv.token_inputs == before_inputs
    assert preview.changed_since_preview(key, [source], pv.token, pv.token_inputs)
    assert any("changed while the preview" in warning for warning in pv.warnings)


def test_source_content_changes_void_approval_even_if_size_and_timestamp_are_restored(tmp_path):
    source = tmp_path / "source.docx"
    source.write_bytes(b"original")
    stat = source.stat()
    token, before = preview.approval_token(font(), [source]), preview.approval_inputs(font(), [source])
    source.write_bytes(b"changed!")
    os.utime(source, ns=(stat.st_atime_ns, stat.st_mtime_ns))
    assert preview.approval_token(font(), [source]) != token
    assert "content changed" in preview.changed_since_preview(font(), [source], token, before)[0]
    assert before["files"][str(source)][:2] == [stat.st_size, stat.st_mtime_ns]


def test_copy_chain_withholds_stale_output_after_original_is_saved_during_later_stage(tmp_path):
    source = tmp_path / "source.docx"
    source.write_bytes(build_sized(title="OLD SOURCE"))
    chain = chain_from_specs([("first", font()), ("second", font("Arial"))])
    state = {"stage": 1}

    def progress(line):
        if line.startswith("stage 2 of"):
            state["stage"] = 2
        if line.startswith("  done ") and state["stage"] == 2:
            source.write_bytes(build_sized(title="NEW USER SAVE"))

    result = run_chain([str(source)], chain, tmp_path / "run", workers=1, do_render=False, progress=progress)
    item = result.files[0]
    assert item.status == "EXCEPTION" and item.code == DRIFT_CODE
    assert "original source changed" in item.reason
    assert item.delivered_to is None
    assert not (result.out_dir / "outputs" / source.name).exists()
    assert "NEW USER SAVE" in body(source)


def margins(left):
    return {"request": "Set the left margin", "summary": "Set margin", "mode": "formatting", "ops": [
        {"op": "set_page_setup", "params": {"margins": {"left": left}}}]}


@pytest.mark.parametrize("surface", ["batch", "chain"])
@pytest.mark.parametrize("in_place", [False, True])
@pytest.mark.parametrize("satisfies_goal", [False, True])
def test_explicit_goal_is_checked_before_copy_or_in_place_delivery(tmp_path, surface, in_place, satisfies_goal):
    source = tmp_path / "source.docx"
    source.write_bytes(build_sized())
    original = source.read_bytes()
    goal = {"kind": "document_effects", "effects": [
        {"id": "margin", "type": "page_setup", "margins": {"left": 2880}}]}
    bindings = {str(source): resolve_goal(source, goal)["binding"]}
    spec = margins(2880 if satisfies_goal else 1440)
    out = tmp_path / "run"
    if surface == "batch":
        result = run_batch([str(source)], "edit-spec", {}, out, spec=spec, workers=1, do_render=False,
                           in_place=in_place, enforce_goals=True, goal_contract=goal, goal_bindings=bindings)
        assert result.files[0].record.get("GOAL").status == ("PASS" if satisfies_goal else "FAIL")
    else:
        # The first stage would modify the original even when the final goal fails if delivery happened stage by stage.
        chain = chain_from_specs([("first", margins(2160)), ("final", spec)])
        result = run_chain([str(source)], chain, out, workers=1, do_render=False,
                           in_place=in_place, enforce_goals=True, goal_contract=goal, goal_bindings=bindings)
        assert result.files[0].outcomes[-1]["kind"] == "approved_goal"
    assert result.goal_contract == goal and result.goal_bindings == bindings
    assert result.counts()["PASS" if satisfies_goal else "FAIL"] == 1
    if not satisfies_goal:
        assert result.files[0].delivered_to is None
        assert source.read_bytes() == original
        assert not (out / "outputs" / source.name).exists()
    elif in_place:
        assert result.files[0].delivered_to == str(source)
        assert source.read_bytes() != original
    else:
        assert (out / "outputs" / source.name).is_file()
        assert source.read_bytes() == original


@pytest.mark.parametrize("surface", ["batch", "chain", "preview", "chain_preview"])
@pytest.mark.parametrize("alias", ["symlink", "hardlink"])
def test_result_aliases_cannot_write_through_to_inputs(tmp_path, surface, alias):
    source = tmp_path / "inputs" / "source.docx"
    source.parent.mkdir()
    source.write_bytes(build_sized())
    original = source.read_bytes()
    out = tmp_path / "run"
    out.mkdir()
    artifacts = out / ("samples" if "preview" in surface else "outputs")
    try:
        if alias == "symlink":
            artifacts.symlink_to(source.parent, target_is_directory=True)
        else:
            artifacts.mkdir()
            os.link(source, artifacts / source.name)
    except OSError as exc:
        pytest.skip(f"filesystem aliases unavailable: {exc}")
    chain = chain_from_specs([("edit", font())])
    with pytest.raises(ValueError, match="symbolic link|hard link|junction"):
        if surface == "batch":
            run_batch([str(source)], "edit-spec", {}, out, spec=font(), workers=1, do_render=False)
        elif surface == "chain":
            run_chain([str(source)], chain, out, workers=1, do_render=False)
        elif surface == "preview":
            preview.make_preview(font(), [str(source)], out, do_render=False)
        else:
            preview.make_chain_preview(chain, [str(source)], out, do_render=False)
    assert source.read_bytes() == original


def test_copy_chain_rechecks_source_after_expensive_final_goal_verification(tmp_path, monkeypatch):
    source = tmp_path / "source.docx"
    source.write_bytes(build_sized(title="OLD SOURCE"))
    goal = {"kind": "document_effects", "effects": [
        {"id": "margin", "type": "page_setup", "margins": {"left": 2880}}]}
    binding = resolve_goal(source, goal)["binding"]
    real_check = batch_module.verify_delivery_goal

    def save_after_goal_proof(*args, **kwargs):
        result = real_check(*args, **kwargs)
        assert result["passed"], result
        source.write_bytes(build_sized(title="NEW SAVE DURING FINAL GOAL PROOF"))
        return result

    monkeypatch.setattr(batch_module, "verify_delivery_goal", save_after_goal_proof)
    result = run_chain([str(source)], chain_from_specs([("margin", margins(2880))]), tmp_path / "run",
                       workers=1, do_render=False, enforce_goals=True, goal_contract=goal, goal_bindings={str(source): binding})
    assert result.files[0].status == "EXCEPTION" and result.files[0].code == DRIFT_CODE
    assert result.files[0].delivered_to is None
    assert not (result.out_dir / "outputs" / source.name).exists()
    assert "NEW SAVE DURING FINAL GOAL PROOF" in body(source)


@pytest.mark.parametrize("chain_preview", [False, True])
@pytest.mark.parametrize("restore_before_samples", [False, True])
def test_preview_never_reapproves_a_source_restored_after_different_bytes_were_read(tmp_path, monkeypatch,
                                                                                 chain_preview, restore_before_samples):
    source = tmp_path / "source.docx"
    source.write_bytes(build_sized(title="APPROVED ORIGINAL A"))
    original, original_stat = source.read_bytes(), source.stat()
    spec = font()
    chain = chain_from_specs([("font", spec)])
    key = chain_key(chain) if chain_preview else spec
    original_token = preview.approval_token(key, [source])
    real_plan_batch = preview.plan_batch_chain if chain_preview else preview.plan_batch

    def restore():
        source.write_bytes(original)
        os.utime(source, ns=(original_stat.st_atime_ns, original_stat.st_mtime_ns))

    def read_other_version_for_plan(*args, **kwargs):
        source.write_bytes(build_sized(title="UNAPPROVED TRANSIENT B"))
        result = real_plan_batch(*args, **kwargs)
        if restore_before_samples:
            restore()
        return result

    def restore_before_final_identity_check(*args, **kwargs):
        restore()
        return None

    # a chain's every step is planned by plan_batch_chain (backlog Q2); a spec's by plan_batch
    monkeypatch.setattr(preview, "plan_batch_chain" if chain_preview else "plan_batch", read_other_version_for_plan)
    monkeypatch.setattr(preview, "render_samples", restore_before_final_identity_check)
    maker = preview.make_chain_preview if chain_preview else preview.make_preview
    pv = maker(chain if chain_preview else spec, [str(source)], tmp_path / "preview", do_render=False)
    assert pv.samples[0].passed
    assert ("APPROVED ORIGINAL A" if restore_before_samples else "UNAPPROVED TRANSIENT B") in body(pv.samples[0].output)
    assert preview.approval_token(key, [source]) == original_token
    assert not pv.token
    assert any("changed while the preview" in warning for warning in pv.warnings)


@pytest.mark.parametrize("surface", ["batch", "chain"])
@pytest.mark.parametrize("in_place", [False, True])
def test_transform_rejects_bytes_different_from_the_acquired_original_even_after_restore(tmp_path, monkeypatch,
                                                                                     surface, in_place):
    source = tmp_path / "source.docx"
    source.write_bytes(build_sized(title="ACQUIRED SOURCE A"))
    original, original_stat = source.read_bytes(), source.stat()
    real_acquire = batch_module.guard.acquire
    changed = False

    def save_between_acquisition_and_read(path):
        nonlocal changed
        identity = real_acquire(path)
        if not changed and str(path) == str(source):
            source.write_bytes(build_sized(title="UNACQUIRED SOURCE B"))
            changed = True
        return identity

    def restore_before_delivery(line):
        if line.startswith("  done "):
            source.write_bytes(original)
            os.utime(source, ns=(original_stat.st_atime_ns, original_stat.st_mtime_ns))

    monkeypatch.setattr(batch_module.guard, "acquire", save_between_acquisition_and_read)
    out = tmp_path / "run"
    if surface == "batch":
        result = run_batch([str(source)], "edit-spec", {}, out, spec=font(), workers=1, do_render=False,
                           in_place=in_place, progress=restore_before_delivery)
        assert result.files[0].exception["code"] == "source-drift"
    else:
        result = run_chain([str(source)], chain_from_specs([("font", font())]), out, workers=1, do_render=False,
                           in_place=in_place, progress=restore_before_delivery)
        assert result.files[0].code == "source-drift"
    assert result.counts()["EXCEPTION"] == 1
    assert source.read_bytes() == original
    assert not (out / "outputs" / source.name).exists()


@pytest.mark.parametrize("in_place", [False, True])
@pytest.mark.parametrize("changed", ["source", "output"])
def test_noop_output_still_requires_unchanged_source_and_verified_bytes(tmp_path, in_place, changed):
    source = tmp_path / "source.docx"
    source.write_bytes(build_sized())
    out = tmp_path / "run"
    spec = {"request": "Remove hidden text", "summary": "Remove hidden text", "mode": "content", "ops": [
        {"op": "remove_hidden_text", "select": {"scope": ["document"]}, "params": {"tracked": True}}]}

    def save_after_worker(line):
        if line.startswith("  done "):
            target = source if changed == "source" else out / "outputs" / source.name
            target.write_bytes(build_sized(title="CHANGED AFTER VERIFICATION"))

    result = run_batch([str(source)], "edit-spec", {}, out, spec=spec, workers=1, do_render=False,
                       in_place=in_place, progress=save_after_worker)
    assert result.files[0].noop
    assert result.counts()["FAIL"] == 1
    assert result.files[0].record.get("V9").status == "FAIL"
    assert not (out / "outputs" / source.name).exists()


def test_source_read_error_at_delivery_is_one_file_failure_not_a_batch_abort(tmp_path, monkeypatch):
    sources = [tmp_path / name for name in ("unavailable.docx", "readable.docx")]
    for source in sources:
        source.write_bytes(build_sized())
    real_drifted = batch_module.guard.drifted

    def source_share_disappears(identity):
        if identity.path == str(sources[0]):
            raise PermissionError("source share became unavailable")
        return real_drifted(identity)

    monkeypatch.setattr(batch_module.guard, "drifted", source_share_disappears)
    result = run_batch([str(p) for p in sources], "edit-spec", {}, tmp_path / "run",
                       spec=font(), workers=1, do_render=False)
    assert result.counts()["FAIL"] == 1 and result.counts()["PASS"] == 1
    assert not (result.out_dir / "outputs" / sources[0].name).exists()
    assert (result.out_dir / "outputs" / sources[1].name).exists()


def auxiliary_case(folder, kind):
    source = folder / "source.docx"
    source.write_bytes(build_styles() if kind == "template" else build_batch1())
    if kind == "table":
        auxiliary = folder / "PROJECT.json"
        original = json.dumps({"SITE_NAME": "APPROVED A", "OWNER_NAME": "Owner", "PROJECT NO": "123"}).encode()
        altered = json.dumps({"SITE_NAME": "UNAPPROVED B", "OWNER_NAME": "Owner", "PROJECT NO": "123"}).encode()
        spec = {"request": "Fill my variables", "summary": "Fill variables", "mode": "untracked", "ops": [
            {"op": "fill_variables", "params": {"table": {"file": str(auxiliary)}, "tracked": False}}]}
    elif kind == "picture":
        auxiliary = folder / "logo.png"
        original, altered = png((180, 20, 20), size=8), png((20, 180, 20), size=8)
        spec = {"request": "Swap my logo", "summary": "Swap logo", "mode": "formatting", "ops": [
            {"op": "swap_image", "params": {"image": str(auxiliary), "match": "all"}}]}
    else:
        auxiliary = folder / "template.docx"
        original = build_style_template()
        package = DocxPackage(original)
        styles = package.xml("word/styles.xml")
        heading = next(s for s in styles.iter(W + "style") if s.get(W + "styleId") == "Heading1")
        fonts = heading.find(W + "rPr/" + W + "rFonts")
        fonts.set(W + "ascii", "Courier New")
        fonts.set(W + "hAnsi", "Courier New")
        package.write_xml("word/styles.xml", styles)
        altered = package.output_bytes()
        spec = {"request": "Import my heading style", "summary": "Import heading", "mode": "formatting", "ops": [
            {"op": "apply_style", "select": {"paragraphs": {"text_regex": "PART"}},
             "params": {"style": "Heading 1", "template": str(auxiliary), "existing": "replace", "tracked": False}}]}
    auxiliary.write_bytes(original)
    return source, spec, auxiliary, original, altered


@pytest.mark.parametrize("kind", ["table", "template", "picture"])
@pytest.mark.parametrize("chain_preview", [False, True])
def test_preview_consumes_only_approved_auxiliary_snapshots_during_transient_saves(tmp_path, monkeypatch,
                                                                               kind, chain_preview):
    source, spec, auxiliary, original, altered = auxiliary_case(tmp_path, kind)
    original_source, stamp = source.read_bytes(), auxiliary.stat()
    chain = chain_from_specs([("edit", spec)])
    key = chain_key(chain) if chain_preview else spec
    token = preview.approval_token(key, [source])
    planner = "plan_batch_chain" if chain_preview else "plan_batch"      # a chain's every step, or the spec
    real_batch, real_plan, real_suite = getattr(preview, planner), preview.plan_spec, verify.run_suite

    def restore():
        auxiliary.write_bytes(original)
        os.utime(auxiliary, ns=(stamp.st_atime_ns, stamp.st_mtime_ns))

    def transient_batch(*args, **kwargs):
        auxiliary.write_bytes(altered)
        try:
            return real_batch(*args, **kwargs)
        finally:
            restore()

    def transient_plan(*args, **kwargs):
        auxiliary.write_bytes(altered)
        return real_plan(*args, **kwargs)

    def transient_suite(*args, **kwargs):
        try:
            return real_suite(*args, **kwargs)
        finally:
            restore()

    monkeypatch.setattr(preview, planner, transient_batch)
    monkeypatch.setattr(preview, "plan_spec", transient_plan)
    monkeypatch.setattr(verify, "run_suite", transient_suite)
    maker = preview.make_chain_preview if chain_preview else preview.make_preview
    pv = maker(chain if chain_preview else spec, [str(source)], tmp_path / "preview", do_render=False)
    assert pv.samples[0].passed and pv.token == token == preview.approval_token(key, [source])
    canonical, _ = real_plan(DocxPackage(original_source), spec)
    independent = real_suite(original_source, pv.samples[0].output.read_bytes(), canonical.to_dict(),
                             str(source), str(pv.samples[0].output))
    assert independent.passed, independent.to_markdown()
    assert pv.input_snapshots[0]["source"] == str(auxiliary)
    assert Path(pv.input_snapshots[0]["snapshot"]).read_bytes() == original
    assert pv.batch_plan.spec == spec


@pytest.mark.parametrize("kind", ["table", "template", "picture"])
@pytest.mark.parametrize("surface", ["batch", "chain"])
@pytest.mark.parametrize("in_place", [False, True])
@pytest.mark.parametrize("restored", [False, True])
def test_execution_freezes_auxiliary_inputs_and_refuses_persistent_saves_before_delivery(tmp_path, kind, surface,
                                                                                      in_place, restored):
    source, spec, auxiliary, original, altered = auxiliary_case(tmp_path, kind)
    original_source, stamp = source.read_bytes(), auxiliary.stat()

    def progress(line):
        if "transform+verify with" in line:
            auxiliary.write_bytes(altered)
        if restored and line.startswith("  done "):
            auxiliary.write_bytes(original)
            os.utime(auxiliary, ns=(stamp.st_atime_ns, stamp.st_mtime_ns))

    out = tmp_path / "run"
    if surface == "batch":
        result = run_batch([str(source)], "edit-spec", {}, out, spec=spec, workers=1, do_render=False,
                           progress=progress, in_place=in_place)
    else:
        result = run_chain([str(source)], chain_from_specs([("edit", spec)]), out, workers=1, do_render=False,
                           progress=progress, in_place=in_place)
        assert result.stages[0].batch.spec == spec
    assert result.input_snapshots[0]["source"] == str(auxiliary)
    assert Path(result.input_snapshots[0]["snapshot"]).read_bytes() == original
    if not restored:
        assert result.counts()["PASS"] == 0
        assert source.read_bytes() == original_source
        assert not (out / "outputs" / source.name).exists()
    else:
        assert result.counts()["PASS"] == 1
        canonical, _ = preview.plan_spec(DocxPackage(original_source), spec)
        independent = verify.run_suite(original_source, result.files[0].output.read_bytes(), canonical.to_dict(),
                                       str(source), str(result.files[0].output))
        assert independent.passed, independent.to_markdown()


@pytest.mark.parametrize("surface", ["batch", "chain", "preview", "chain_preview"])
@pytest.mark.parametrize("in_place", [False, True])
def test_modified_private_snapshot_cannot_authorize_delivery_or_preview(tmp_path, monkeypatch, surface, in_place):
    source, spec, auxiliary, original, altered = auxiliary_case(tmp_path, "table")
    original_source = source.read_bytes()
    out = tmp_path / "run"

    def corrupt_snapshot():
        snapshot = next(out.glob("input-snapshot-*/0/PROJECT.json"))
        snapshot.write_bytes(altered)

    def progress(line):
        if "transform+verify with" in line:
            corrupt_snapshot()

    if "preview" in surface:
        planner = "plan_batch_chain" if surface == "chain_preview" else "plan_batch"
        real_plan_batch = getattr(preview, planner)

        def after_acquisition(*args, **kwargs):
            corrupt_snapshot()
            return real_plan_batch(*args, **kwargs)

        monkeypatch.setattr(preview, planner, after_acquisition)
        maker = preview.make_chain_preview if surface == "chain_preview" else preview.make_preview
        result = maker(chain_from_specs([("edit", spec)]) if surface == "chain_preview" else spec,
                       [str(source)], out, do_render=False)
        assert not result.token
    elif surface == "batch":
        result = run_batch([str(source)], "edit-spec", {}, out, spec=spec, workers=1, do_render=False,
                           progress=progress, in_place=in_place)
        assert result.counts()["PASS"] == 0
    else:
        result = run_chain([str(source)], chain_from_specs([("edit", spec)]), out, workers=1, do_render=False,
                           progress=progress, in_place=in_place)
        assert result.counts()["PASS"] == 0
    assert auxiliary.read_bytes() == original
    assert source.read_bytes() == original_source
    assert not (out / "outputs" / source.name).exists()


def test_batch_public_recipe_remains_the_one_executed_after_caller_mutation(tmp_path):
    source = tmp_path / "source.docx"
    source.write_bytes(build_sized())
    spec = font()
    approved = copy.deepcopy(spec)

    def mutate_caller_recipe(line):
        if line.startswith("  done "):
            spec["ops"][0]["params"]["font"] = "Times New Roman"

    result = run_batch([str(source)], "edit-spec", {}, tmp_path / "run", spec=spec, workers=1, do_render=False,
                       progress=mutate_caller_recipe)
    assert result.counts()["PASS"] == 1
    assert result.spec == approved
    assert json.loads((result.out_dir / "batch.json").read_text(encoding="utf-8"))["spec"] == approved


@pytest.mark.parametrize("surface", ["batch", "chain"])
@pytest.mark.parametrize("in_place", [False, True])
def test_all_queued_sources_are_pinned_before_the_first_worker_or_progress_callback(tmp_path, surface, in_place):
    sources = [tmp_path / "01.docx", tmp_path / "02.docx"]
    for source in sources:
        source.write_bytes(build_sized())
    saved = build_sized(title="USER SAVE AFTER RUN START")

    def progress(line):
        if line.startswith("  done 01.docx"):
            sources[1].write_bytes(saved)

    out = tmp_path / "run"
    if surface == "batch":
        result = run_batch([str(p) for p in sources], "edit-spec", {}, out, spec=font(), workers=1,
                           do_render=False, in_place=in_place, progress=progress)
        assert result.files[1].exception["code"] == "source-drift"
    else:
        result = run_chain([str(p) for p in sources], chain_from_specs([("font", font())]), out, workers=1,
                           do_render=False, in_place=in_place, progress=progress)
        assert result.files[1].code == "source-drift"
    assert result.counts()["PASS"] == 1 and result.counts()["EXCEPTION"] == 1
    assert sources[1].read_bytes() == saved
    assert result.files[1].delivered_to is None
    assert not (out / "outputs" / sources[1].name).exists()


@pytest.mark.parametrize("surface", ["batch", "chain"])
@pytest.mark.parametrize("changed", ["source", "scope", "recipe", "auxiliary"])
def test_approved_inputs_close_frontend_handoff_before_any_delivery_cleanup(tmp_path, surface, changed):
    source, spec, auxiliary, original, altered = auxiliary_case(tmp_path, "table")
    chain = chain_from_specs([("edit", spec)])
    key = chain_key(chain) if surface == "chain" else spec
    approved = preview.approval_inputs(key, [source])
    inputs = [str(source)]
    if changed == "source":
        source.write_bytes(build_sized(title="SAVED AFTER APPROVAL"))
    elif changed == "scope":
        extra = tmp_path / "extra.docx"
        extra.write_bytes(build_sized())
        inputs.append(str(extra))
    elif changed == "recipe":
        (chain["stages"][0]["spec"] if surface == "chain" else spec)["summary"] = "Changed after approval"
    else:
        auxiliary.write_bytes(altered)
    out = tmp_path / "run"
    previous = out / "outputs" / "previous.docx"
    previous.parent.mkdir(parents=True)
    previous.write_bytes(b"previous delivery must remain intact when approval is refused")
    with pytest.raises(ValueError, match="approval|approved|preview|snapshot"):
        if surface == "batch":
            run_batch(inputs, "edit-spec", {}, out, spec=spec, workers=1, do_render=False, approved_inputs=approved)
        else:
            run_chain(inputs, chain, out, workers=1, do_render=False, approved_inputs=approved)
    assert previous.read_bytes() == b"previous delivery must remain intact when approval is refused"


# ---- an outcome contract is advisory unless a governed caller enforces it ---------------------------------------------

TITLE_REQUEST = "Make the FIRE SUPPRESSION section title bold in every document."
TITLE_CONTRACT = {"kind": "document_effects", "effects": [{"id": "title", "type": "format",
    "target": {"locators": [{"exact_text": "FIRE SUPPRESSION"}], "cardinality": 1}, "run": {"b": True}}]}


def document(path, *, title_style="Normal", body_style="Normal", duplicate=False):
    doc = Document()
    doc.add_paragraph("SECTION 21 13 13")
    title = doc.add_paragraph("FIRE SUPPRESSION", style=title_style)
    title.runs[0].bold = False
    if duplicate:
        doc.add_paragraph("FIRE SUPPRESSION")
    body = doc.add_paragraph("Owner shall retain all content.", style=body_style)
    body.runs[0].bold = False
    doc.save(path)
    return path


def recipe(predicate):
    return {"request": TITLE_REQUEST, "summary": "Bold the requested section titles", "mode": "formatting",
        "ops": [{"op": "set_properties", "select": {"paragraphs": predicate},
                 "params": {"run": {"b": True}, "tracked": False}}]}


@pytest.mark.parametrize("chain,in_place", [(False, False), (False, True), (True, False), (True, True)])
def test_an_unenforced_outcome_contract_cannot_withhold_verified_delivery(tmp_path, chain, in_place):
    source = document(tmp_path / "spec.docx")
    spec = recipe({"text_regex": "^Owner"})
    out = tmp_path / "run"
    kwargs = dict(workers=1, do_render=False, in_place=in_place,
                  goal_contract=TITLE_CONTRACT, goal_bindings={})  # even missing bindings cannot create a gate
    result = (run_chain([str(source)], chain_from_specs([("bold", spec)]), out, **kwargs) if chain else
              run_batch([str(source)], "edit-spec", {}, out, spec=spec, **kwargs))
    assert result.counts() == {"PASS": 1, "FAIL": 0, "EXCEPTION": 0, "PENDING": 0}
    assert result.goal_contract == TITLE_CONTRACT
    assert result.files[0].delivered_to
    record = result.files[0].results[-1].record if chain else result.files[0].record
    assert record.get("GOAL") is None


# ---- failed outputs never sit beside verified ones -------------------------------------------------------------------

AUTHOR = "Abraham Borg"


def test_failed_outputs_are_quarantined_and_explained(tmp_path, monkeypatch):
    real = batch_module.run_suite

    def planted_v2_failure(src, out, plan, s, o):
        rec = real(src, out, plan, s, o)
        rec.set(CheckResult("V2", "Text preservation", "FAIL", "planted"))
        return rec
    monkeypatch.setattr(batch_module, "run_suite", planted_v2_failure)
    d = tmp_path / "in"; d.mkdir()
    (d / "21 05 00 – CLEAN.docx").write_bytes(build("21 05 00", "COMMON WORK RESULTS"))
    spec = {"request": "replace FM Global with GRC", "summary": "s", "mode": "content",
            "ops": [{"op": "replace_text", "select": {"scope": ["document"]},
                     "params": {"find": "FM Global", "replace": "GRC", "author": AUTHOR}}]}
    out = tmp_path / "out"
    b = run_batch([str(d)], "edit-spec", {}, out, workers=1, renderer="none", do_render=False, spec=spec)
    assert b.counts()["FAIL"] == 1
    assert not list((out / "outputs").iterdir())
    assert [p.name for p in (out / "failed").iterdir()] == ["21 05 00 – CLEAN.docx"]
    f = b.files[0]
    assert f.output == out / "failed" / "21 05 00 – CLEAN.docx" and f.record.output == str(f.output)
    assert f.record.get("V9").status == "NA" and "failed/" in f.record.get("V9").summary
    item = load_set_aside(out)["files"][0]
    assert item["code"] == "verify-fail" and item["category"] == "Failed Verification (Not Delivered)"
    assert "failed/" in (out / "BATCH_REPORT.md").read_text(encoding="utf-8")


def test_v9_guard_failures_are_quarantined_too(tmp_path, monkeypatch):
    """Codex review on the v0.3.1 PR: V9 runs after the first quarantine check, so a drifted source or a refused
    in-place write-back used to leave a stale output in outputs/ (and record it as delivered)."""
    spec = {"request": "replace FM Global with GRC", "summary": "s", "mode": "content",
            "ops": [{"op": "replace_text", "select": {"scope": ["document"]},
                     "params": {"find": "FM Global", "replace": "GRC", "author": AUTHOR}}]}
    # non-in-place: the source drifts while the batch runs
    d = tmp_path / "drift"; d.mkdir()
    (d / "21 05 00 – A.docx").write_bytes(build("21 05 00", "A"))
    monkeypatch.setattr(batch_module.guard, "drifted", lambda identity: "mtime moved during the run")
    out = tmp_path / "out1"
    b = run_batch([str(d)], "edit-spec", {}, out, workers=1, renderer="none", do_render=False, spec=spec)
    f = b.files[0]
    assert f.status == "FAIL" and f.record.get("V9").status == "FAIL" and f.delivered_to is None
    assert not list((out / "outputs").iterdir()) and [p.name for p in (out / "failed").iterdir()] == ["21 05 00 – A.docx"]
    assert f.output == out / "failed" / "21 05 00 – A.docx"
    monkeypatch.undo()
    # in place: the guard refuses the write-back (file open in Word / drifted)
    d2 = tmp_path / "locked"; d2.mkdir()
    src = d2 / "21 05 00 – B.docx"; src.write_bytes(build("21 05 00", "B")); before = src.read_bytes()

    def refuse(identity, data):
        raise batch_module.guard.GuardError("Word lock file present: the document is open")
    monkeypatch.setattr(batch_module.guard, "write_back", refuse)
    out2 = tmp_path / "out2"
    b2 = run_batch([str(d2)], "edit-spec", {}, out2, workers=1, renderer="none", do_render=False, spec=spec, in_place=True)
    f2 = b2.files[0]
    assert f2.status == "FAIL" and f2.record.get("V9").status == "FAIL" and f2.delivered_to is None
    assert src.read_bytes() == before
    assert not list((out2 / "outputs").iterdir()) and [p.name for p in (out2 / "failed").iterdir()] == ["21 05 00 – B.docx"]
    assert "failed/" in f2.record.get("V9").summary and "failed/" in (out2 / "BATCH_REPORT.md").read_text(encoding="utf-8")


# ---- inputs that share a name or sit in the results; delivery bound to verified bytes --------------------------------

def _spec(mode, ops):
    return {"request": "x", "summary": "x", "assumptions": [], "questions": [], "applies_to": {"or": "all"},
            "mode": mode, "ops": ops}


REPLACE = _spec("content", [{"op": "replace_text", "params": {"find": "FM Global", "replace": "FM Approved"}}])


def _title(path: Path) -> str:
    return block_full_text(paragraphs(DocxPackage.from_path(path).xml("word/document.xml"))[1])


def test_outputs_and_records_stay_distinct_when_inputs_share_a_name(tmp_path):
    files = []
    for folder, num, title in (("a", "21 00 01", "ALPHA"), ("b", "21 00 02", "BRAVO")):
        d = tmp_path / folder
        d.mkdir()
        (d / "Spec.docx").write_bytes(build(num, title))
        files.append(d / "Spec.docx")
    (tmp_path / "c").mkdir()
    files.append(tmp_path / "c" / "Spec (1).docx")
    files[-1].write_bytes(build("21 00 03", "CHARLIE"))
    # a name an input really has is never handed to another input; the first keeps its own, the next gets a free one
    assert [output_names(files)[str(f)] for f in files] == ["Spec.docx", "Spec (2).docx", "Spec (1).docx"]
    assert list(output_names([Path("x/SPEC.docx"), Path("y/spec.docx")]).values()) == ["SPEC.docx", "spec (1).docx"]
    before = {f: _title(f) for f in files}
    b = run_batch([str(f) for f in files], "edit-spec", {}, tmp_path / "out", workers=1, renderer="none",
                  do_render=False, in_place=True, spec=REPLACE)
    assert [f.status for f in b.files] == ["PASS"] * 3
    assert [f.out_name for f in b.files] == ["Spec.docx", "Spec (2).docx", "Spec (1).docx"]
    assert [f.delivered_to for f in b.files] == [str(f) for f in files]
    assert {f: _title(f) for f in files} == before                 # every original is still its own document
    assert sorted(p.name for p in (tmp_path / "out" / "records").glob("*.record.json")) == \
        ["Spec (1).record.json", "Spec (2).record.json", "Spec.record.json"]
    assert len(list((tmp_path / "out" / "plans").glob("*.plan.json"))) == 3


def test_an_input_inside_the_results_folder_is_refused_before_anything_is_written(tmp_path):
    results = tmp_path / "results"
    src = results / "outputs" / "Spec.docx"
    src.parent.mkdir(parents=True)
    src.write_bytes(build("21 00 04", "DELTA"))
    md5 = hashlib.md5(src.read_bytes()).hexdigest()
    with pytest.raises(ValueError, match="inside the results folder"):
        run_batch([str(src)], "edit-spec", {}, results, workers=1, renderer="none", do_render=False, spec=REPLACE)
    assert hashlib.md5(src.read_bytes()).hexdigest() == md5
    assert sorted(p.name for p in results.iterdir()) == ["outputs"]       # nothing created
    check_inputs_outside([tmp_path / "in" / "x.docx"], results)             # outside: fine
    with pytest.raises(ValueError):
        check_inputs_outside([results / "deeper" / "x.docx"], results)


def test_a_verifier_crash_leaves_nothing_in_outputs(tmp_path, monkeypatch):
    f = tmp_path / "A.docx"
    f.write_bytes(build("21 00 05", "ECHO"))

    def boom(*a, **k):
        raise RuntimeError("injected verifier crash")
    monkeypatch.setattr(batch_mod, "run_suite", boom)
    b = run_batch([str(f)], "edit-spec", {}, tmp_path / "out", workers=1, renderer="none", do_render=False, spec=REPLACE)
    fr = b.files[0]
    # the engine's own failure, said as such (never "a corrupt file"): the stage, the line of code, the traceback
    assert fr.status == "EXCEPTION" and fr.exception["code"] == "internal-error" and "injected" in fr.exception["reason"]
    assert fr.exception["kind"] == "internal-error" and fr.exception["stage"] == "verify"
    assert "injected verifier crash" in fr.exception["traceback"] and fr.exception["where"].startswith("batch.py:")
    assert not list((tmp_path / "out" / "outputs").iterdir()) and not (tmp_path / "out" / "failed").exists()


def test_delivery_is_bound_to_the_bytes_the_suite_verified(tmp_path, monkeypatch):
    f = tmp_path / "A.docx"
    f.write_bytes(build("21 00 06", "FOXTROT"))
    md5 = hashlib.md5(f.read_bytes()).hexdigest()
    real = batch_mod._process_one

    def tampered(*a, **k):
        r = real(*a, **k)
        Path(r["output"]).write_bytes(build("99 99 99", "SOMETHING ELSE"))  # replaced after verification
        return r
    monkeypatch.setattr(batch_mod, "_process_one", tampered)
    for in_place in (True, False):
        out = tmp_path / ("in_place" if in_place else "copies")
        b = run_batch([str(f)], "edit-spec", {}, out, workers=1, renderer="none", do_render=False,
                      in_place=in_place, spec=REPLACE)
        fr = b.files[0]
        v9 = fr.record.get("V9")
        assert fr.status == "FAIL" and v9.status == "FAIL" and "not the one the suite verified" in v9.summary
        assert fr.delivered_to is None and (out / "failed" / "A.docx").is_file() and not (out / "outputs" / "A.docx").exists()
        assert hashlib.md5(f.read_bytes()).hexdigest() == md5                 # the original untouched
    fr2 = FileResult(source=f, output=f, output_md5="0" * 32)
    with pytest.raises(guard.GuardError, match="not the one the suite verified"):
        verified_bytes(fr2)
