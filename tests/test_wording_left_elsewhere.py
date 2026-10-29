"""A change kept to one place says where the same words remain.

From the owner's Division 21 run: the rename touched only the footer of 21 05 53, the first page still read
"Identification for Fire Suppression Systems", and no step said so. ``replace_text`` now flags each place its find
text (the same matching) also occurs and no replacement of the spec edits, and a literal header/footer fill of the
section number or title flags a document whose own value differs. Flags only: the predictions, the output and the
verification are what they were."""
from __future__ import annotations

import copy
import io
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from make_fixtures import TEMPLATE_FILLS, build, build_template
from spec_surgeon import compiler
from spec_surgeon.editspec import apply_spec, plan_spec
from spec_surgeon.ops import JobException
from spec_surgeon.ops.base import LEFT_ELSEWHERE
from spec_surgeon.package import DocxPackage
from spec_surgeon.planner import LEFT_WARNINGS, left_elsewhere_notes, left_elsewhere_warnings, plan_batch
from spec_surgeon.preview import make_preview
from spec_surgeon.verify import run_suite

TITLE = "Identification for Fire Suppression Systems"
NEW = "Identification for Fire-Suppression Piping and Equipment"


def styleref(paragraph, result: str) -> None:
    """A STYLEREF field whose cached result shows ``result``."""
    for kind, value in (("begin", None), ("code", ' STYLEREF "Title" '), ("separate", None), ("result", result),
                        ("end", None)):
        run = OxmlElement("w:r")
        if kind == "code":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = value
        elif kind == "result":
            el = OxmlElement("w:t")
            el.text = value
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        run.append(el)
        paragraph._p.append(run)


def spec_doc(*, header_field=False, footer=TITLE, body_extra=(), number="21 05 53") -> bytes:
    """21 05 53: its title on the first page, in the footer, and (``header_field``) in a header field's result."""
    doc = Document()
    sec = doc.sections[0]
    sec.footer.is_linked_to_previous = False
    for i, line in enumerate([footer] if isinstance(footer, str) else footer):
        (sec.footer.paragraphs[0] if i == 0 else sec.footer.add_paragraph()).add_run(line)
    if header_field:
        sec.header.is_linked_to_previous = False
        styleref(sec.header.paragraphs[0], TITLE)
    doc.add_paragraph(f"SECTION {number}")
    doc.add_paragraph(TITLE)
    doc.add_paragraph("PART 1 - GENERAL")
    for text in body_extra:
        doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def rename(scope=("footers",), find=TITLE, replace=NEW, **params) -> dict:
    return {"op": "replace_text", "select": {"scope": list(scope), **({"paragraphs": params.pop("paragraphs")}
                                                                        if "paragraphs" in params else {})},
            "params": {"find": find, "replace": replace, "tracked": True, **params}}


def spec_of(*ops, mode="content") -> dict:
    return {"request": f"Rename “{TITLE}” to “{NEW}”.", "summary": "rename", "mode": mode, "ops": list(ops)}


def left(plan) -> list:
    return [f.split(LEFT_ELSEWHERE, 1)[1] for f in plan.flags if LEFT_ELSEWHERE in f]


def test_a_footer_only_rename_says_the_first_page_keeps_the_old_title():
    src = spec_doc()
    pkg = DocxPackage(src)
    plan, frags = plan_spec(pkg, spec_of(rename()))
    assert left(plan) == [f"“{TITLE}” also occurs in the body (1×); this change leaves it as it is."]
    # a flag only: the output and the verification are those of the same edit without it
    out = apply_spec(DocxPackage(src), frags).output_bytes()
    assert run_suite(src, out, plan.to_dict(), "src.docx", "out.docx").passed
    assert [p.text for p in Document(io.BytesIO(out)).paragraphs][1] == TITLE


def test_renaming_both_places_leaves_nothing_to_say():
    pkg = DocxPackage(spec_doc())
    assert left(plan_spec(pkg, spec_of(rename(scope=("document", "footers"))))[0]) == []
    # two replacements of the same words, one per place: neither leaves the other's words behind
    plan, _ = plan_spec(DocxPackage(spec_doc()), spec_of(rename(), rename(scope=("document",))))
    assert left(plan) == []


def test_the_selection_and_the_first_occurrence_name_what_they_leave():
    src = spec_doc(body_extra=[f"The {TITLE} are listed in Part 2."])
    pkg = DocxPackage(src)
    plan, _ = plan_spec(pkg, spec_of(rename(scope=("document",), paragraphs={"text_regex": "^Identification"})))
    assert left(plan) == [f"“{TITLE}” also occurs in the body, outside the selected paragraphs (1×) and the footers "
                          f"(1×); this change leaves it as it is."]
    plan, _ = plan_spec(DocxPackage(src), spec_of(rename(scope=("document", "footers"), occurrences="first")))
    assert left(plan) == [f"“{TITLE}” also occurs in the body, beyond the one changed (1×) and the footers, beyond "
                          f"the one changed (1×); this change leaves it as it is."]


def test_a_field_result_is_counted_and_named():
    plan, _ = plan_spec(DocxPackage(spec_doc(header_field=True)), spec_of(rename()))
    assert left(plan) == [f"“{TITLE}” also occurs in the body (1×) and the headers (1×, 1 in a field's result); this "
                          f"change leaves it as it is."]


def test_the_same_matching_decides_what_counts():
    src = spec_doc(body_extra=["IDENTIFICATION FOR FIRE SUPPRESSION SYSTEMS", "NFPA 13R and NFPA 13D apply."],
                   footer=[TITLE, "NFPA 13 applies."])
    plan, _ = plan_spec(DocxPackage(src), spec_of(rename()))
    assert left(plan) == [f"“{TITLE}” also occurs in the body (1×); this change leaves it as it is."]   # case kept
    plan, _ = plan_spec(DocxPackage(src), spec_of(rename(match_case=False)))
    assert left(plan) == [f"“{TITLE}” also occurs in the body (2×); this change leaves it as it is."]
    plan, _ = plan_spec(DocxPackage(src), spec_of(rename(find="NFPA 13", replace="NFPA 13 (2025)")))
    assert left(plan) == []                                                  # 13R and 13D are not whole words


def test_a_rename_that_finds_nothing_in_its_scope_says_where_the_words_are():
    with pytest.raises(JobException) as e:
        plan_spec(DocxPackage(spec_doc(footer="21 05 53 - 1")), spec_of(rename()))
    assert e.value.code == "not-found"
    assert e.value.reason.endswith("; it occurs in the body (1×)")
    assert e.value.details["occurs_elsewhere"] == 1


@pytest.fixture(scope="module")
def template(tmp_path_factory) -> Path:
    p = tmp_path_factory.mktemp("template") / "00_00_00_TEMPLATE.docx"
    p.write_bytes(build_template())
    return p


def template_op(tpl: Path, title_rule: dict) -> dict:
    fills = copy.deepcopy(TEMPLATE_FILLS)
    fills["SECTION_TITLE"] = title_rule
    return {"op": "apply_header_footer_template", "select": {"sections": "all"},
            "params": {"template": str(tpl), "template_sections": {"cover": 0, "body": 1}, "slot_fills": fills}}


def test_a_literal_title_fill_says_the_documents_own_title_differs(template):
    src = build("21 05 00", "COMMON WORK RESULTS FOR FIRE SUPPRESSION")
    words = {"request": "Put COMMON WORK RESULTS in the header.", "summary": "template", "mode": "untracked"}
    plan, frags = plan_spec(DocxPackage(src), {**words, "ops": [template_op(template, {
        "rule": "literal", "value": "COMMON WORK RESULTS"})]})
    assert left(plan) == ["the headers/footers get the words given for the section title, “COMMON WORK RESULTS”, while "
                          "the document's own section title reads “COMMON WORK RESULTS FOR FIRE SUPPRESSION”; this "
                          "change leaves the document as it is."]
    out = apply_spec(DocxPackage(src), frags).output_bytes()
    assert run_suite(src, out, plan.to_dict(), "src.docx", "out.docx").passed
    # the same words in another case are not this flag's (C4 is about case); an extracted fill is the file's own
    same = plan_spec(DocxPackage(src), {**words, "ops": [template_op(template, {
        "rule": "literal", "value": "Common Work Results for Fire Suppression"})]})[0]
    assert left(same) == []
    extracted = plan_spec(DocxPackage(src), {**words, "mode": "untracked", "ops": [template_op(
        template, TEMPLATE_FILLS["SECTION_TITLE"])]})[0]
    assert left(extracted) == []
    # a chain whose regex filled the slot: its literal fallback placed nothing
    chain = plan_spec(DocxPackage(src), {**words, "ops": [template_op(template, {"rule": "first_match", "rules": [
        TEMPLATE_FILLS["SECTION_TITLE"], {"rule": "literal", "value": "COMMON WORK RESULTS"}]})]})[0]
    assert left(chain) == []


def test_every_file_reaches_the_preview_text_and_the_window(tmp_path):
    files = []
    for i in range(8):
        path = tmp_path / f"21 05 {50 + i}.docx"
        path.write_bytes(spec_doc(number=f"21 05 {50 + i}"))
        files.append(path)
    spec = spec_of(rename())
    bp = plan_batch([str(f) for f in files], spec)
    assert [name for name, _ in bp.left_elsewhere()] == [f.name for f in files]
    text = bp.to_text(per_file=False)                  # the summaries show one file per pattern; this, every file
    assert "WORDING LEFT ELSEWHERE" in text and all(f.name in text for f in files)
    warnings = left_elsewhere_warnings(bp)
    assert len(warnings) == LEFT_WARNINGS + 1
    assert warnings[0] == f"21 05 50.docx: “{TITLE}” also occurs in the body (1×); this change leaves it as it is."
    assert warnings[-1].startswith("2 more place(s) like these keep the old words")
    # the window's Worth knowing card: the same notes, by file, the count of the rest about no file
    notes = left_elsewhere_notes(bp)
    assert [f"{n['name']}: {n['text']}" if n["name"] else n["text"] for n in notes] == warnings
    assert notes[0] == {"name": "21 05 50.docx", "text": f"“{TITLE}” also occurs in the body (1×); this change leaves "
                                                         f"it as it is."}
    assert notes[-1]["name"] == "" and notes[-1]["text"].startswith("2 more place(s)")
    pv = make_preview(spec, [str(f) for f in files[:2]], tmp_path / "preview", do_render=False)
    assert f"21 05 51.docx: “{TITLE}” also occurs in the body (1×); this change leaves it as it is." in pv.warnings
    assert {"name": "21 05 51.docx", "text": f"“{TITLE}” also occurs in the body (1×); this change leaves it as it "
                                             f"is."} in pv.worth_knowing
    assert "rendering skipped: plan summaries and verification only" in pv.warnings     # a note the window keeps back
    assert not any("rendering skipped" in n["text"] for n in pv.worth_knowing)
    assert "WORDING LEFT ELSEWHERE" in pv.to_text()


def test_a_note_starts_with_a_capital_under_its_file_name(tmp_path):
    src = tmp_path / "21 05 53.docx"
    src.write_bytes(spec_doc())
    spec = spec_of(rename())
    bp = plan_batch([str(src)], spec)
    bp.left_elsewhere = lambda: [("21 05 53.docx", "the headers/footers get the words given for the section title")]
    bp.case_notes = lambda: []
    assert left_elsewhere_warnings(bp) == ["21 05 53.docx: the headers/footers get the words given for the section "
                                           "title"]
    assert left_elsewhere_notes(bp) == [{"name": "21 05 53.docx",
                                         "text": "The headers/footers get the words given for the section title"}]


def test_the_window_lists_the_notes_once_each_and_keeps_every_other_warning_back(tmp_path):
    from spec_surgeon.document_task import DocumentTask
    from spec_surgeon.gui import Session
    src = tmp_path / "21 05 53.docx"
    src.write_bytes(spec_doc())
    task = DocumentTask(Session(out_root=tmp_path / "out", visual_check=False), render=False, mode="safe")
    task.attach([src])
    view = task.perform("recipe", recipe="remove_highlight", inputs={})
    assert view["stage"] == "preview" and view["worth_knowing"] == []
    pv = task.session.outcome().preview
    note = {"name": "21 05 53.docx", "text": f"“{TITLE}” also occurs in the body (1×); this change leaves it as it is."}
    pv.worth_knowing += [note, dict(note), {"name": "", "text": "A file changed while this preview was being made."}]
    pv.warnings += ["preview rendered with libreoffice; Word's layout can differ"]
    view = task.view()
    assert view["worth_knowing"] == [note, {"name": "", "text": "A file changed while this preview was being made."}]
    assert "preview rendered with libreoffice; Word's layout can differ" in view["warnings"]
    assert task._samples_view()["worth_knowing"] == view["worth_knowing"]          # the samples frame carries the same


def test_the_prompt_says_to_name_them_with_an_offer():
    prompt = compiler.system_prompt(chat=True)
    assert "WORDING LEFT ELSEWHERE" in prompt and "never change them unasked" in prompt


# ---- from review (Codex on #80) ----

def test_past_forty_notes_every_file_is_still_named(tmp_path):
    files = []
    for i in range(45):
        path = tmp_path / f"21 {10 + i // 10:02d} {i % 10:02d}.docx"
        path.write_bytes(spec_doc(number=f"21 {10 + i // 10:02d} {i % 10:02d}"))
        files.append(path)
    text = plan_batch([str(f) for f in files], spec_of(rename())).to_text(per_file=False)
    assert "… and 5 more note(s) of the same kind, in: " in text
    assert all(f.name in text.split("WORDING LEFT ELSEWHERE", 1)[1] for f in files)


def test_an_allowed_miss_still_says_where_the_words_are():
    src = spec_doc(footer="21 05 53 - 1")
    plan, frags = plan_spec(DocxPackage(src), spec_of(rename(allow_missing=True)))
    assert left(plan) == [f"“{TITLE}” also occurs in the body (1×); this change leaves it as it is."]
    out = apply_spec(DocxPackage(src), frags).output_bytes()       # a no-op still applies as one
    assert run_suite(src, out, plan.to_dict(), "src.docx", "out.docx").passed


def test_the_rule_that_supplies_a_fill_decides_whether_it_was_typed(template):
    src = build("21 05 00", "COMMON WORK RESULTS FOR FIRE SUPPRESSION")
    words = {"request": "Put COMMON WORK RESULTS in the header.", "summary": "template", "mode": "untracked"}
    reads = {"rule": "regex", "source": "document", "pattern": r"^(COMMON WORK RESULTS)\b"}
    typed = {"rule": "literal", "value": "COMMON WORK RESULTS"}
    # the regex reads the same words as the literal fallback: the document supplied them, nothing was typed
    plan = plan_spec(DocxPackage(src), {**words, "ops": [template_op(template, {"rule": "first_match",
                                                                                "rules": [reads, typed]})]})[0]
    assert left(plan) == []
    # the regex finds nothing: the fallback's typed words fill the slot, and differ from the document's title
    misses = {"rule": "regex", "source": "document", "pattern": r"^(NOWHERE IN THIS FILE)"}
    plan = plan_spec(DocxPackage(src), {**words, "ops": [template_op(template, {"rule": "first_match",
                                                                                "rules": [misses, typed]})]})[0]
    assert len(left(plan)) == 1 and "“COMMON WORK RESULTS”" in left(plan)[0]


def _chain(*stages):
    return {"name": "rename", "stages": [{"name": f"step {i}", "spec": s} for i, s in enumerate(stages, 1)]}


def test_a_chain_names_what_its_later_steps_leave_and_drops_what_they_change(tmp_path):
    from spec_surgeon.preview import make_chain_preview
    folder = tmp_path / "docs"
    folder.mkdir()
    for n in ("21 05 53", "21 05 54"):
        (folder / f"{n}.docx").write_bytes(spec_doc(number=n))
    bold = {"request": "Bold PART 1.", "summary": "bold", "mode": "formatting",
            "ops": [{"op": "set_properties", "select": {"paragraphs": {"text_regex": "^PART 1"}},
                     "params": {"run": {"b": True}, "tracked": False}}]}
    # a later step renames the footer: every file says the body keeps the old title (every step is planned on every
    # file, backlog Q2), the note naming the step
    pv = make_chain_preview(_chain(bold, spec_of(rename())), [str(folder)], tmp_path / "p1", renderer="none",
                            do_render=False)
    later = [w for w in pv.warnings if "also occurs in the body" in w]
    assert later and all(w.endswith("this change leaves it as it is.") and "after step 2: " in w for w in later)
    assert {w.split(":")[0] for w in later} == {"21 05 53.docx", "21 05 54.docx"}
    assert not any("sample documents only" in w for w in pv.warnings)
    text = pv.batch_plan.to_text(per_file=False)
    assert "WORDING LEFT ELSEWHERE —" in text and text.count("after step 2: ") == 2
    # step 1 renames the footer, step 2 the body: step 1's note no longer holds and is not shown, step 2 leaves none
    pv = make_chain_preview(_chain(spec_of(rename()), spec_of(rename(scope=("document",)))), [str(folder)],
                            tmp_path / "p2", renderer="none", do_render=False)
    assert not any("also occurs" in w for w in pv.warnings)
    assert "WORDING LEFT ELSEWHERE —" not in pv.batch_plan.to_text(per_file=False)
