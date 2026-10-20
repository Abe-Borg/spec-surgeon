"""Character changes, independent re-derivation, and individual user approval."""
import io
from pathlib import Path

import pytest
from docx import Document

from spec_surgeon import compiler
from spec_surgeon.editspec import plan_spec, validate_spec
from spec_surgeon.ops import JobException, get_op, merge_fragments
from spec_surgeon.ops import hygiene
from spec_surgeon.ops.textedit import Edit, RunText
from spec_surgeon.package import DocxPackage
from spec_surgeon.verify import run_suite


def document(*texts):
    d = Document()
    for text in texts:
        p = d.add_paragraph()
        # Force changes to span separate runs with their own formatting.
        for i in range(0, len(text), 3):
            p.add_run(text[i:i+3]).bold = False
    d.sections[0].header.paragraphs[0].text = "Project  Name  "
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


def execute(src, operation, params, name="source.docx", select=None):
    p = DocxPackage(src, name=name)
    op = get_op(operation)
    f = op.plan(p, select or {}, params)
    plan = merge_fragments([f])
    work = p.copy()
    op.apply(work, f)
    result = work.output_bytes()
    rec = run_suite(src, result, plan.to_dict(), name, "out.docx")
    return f, result, rec


def text(data):
    return RunText(next(iter(DocxPackage(data).xml("word/document.xml").iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p")))).text


@pytest.mark.parametrize("tracked", [True, False])
@pytest.mark.parametrize("op,params,before,after", [
    ("normalize_whitespace", {}, "one   two  \nthree  ", "one two\nthree"),
    ("normalize_punctuation", {"quotes": "curly", "dash": "em", "section_separator": "en"},
     '\"Owner\" -- owner\'s Section 21-13—19 and 6\"', '“Owner” — owner’s Section 21–13–19 and 6\"'),
    ("normalize_punctuation", {"quotes": "straight"}, '“Owner” and ‘Work’', '\"Owner\" and \'Work\''),
    ("protect_nonbreaking_spaces", {}, "10 psi; 21 13 19; 10 psig; 5 kPa", "10\u00a0psi; 21\u00a013\u00a019; 10 psig; 5 kPa"),
    ("protect_nonbreaking_spaces", {"units": ["kPa"], "section_numbers": False}, "5 kPa; 21 13 19", "5\u00a0kPa; 21 13 19"),
])
def test_exact_characters_tracked_and_untracked(tracked, op, params, before, after):
    src = document(before)
    frag, out, rec = execute(src, op, {**params, "tracked": tracked})
    assert rec.passed, rec.to_markdown()
    assert frag.transform_data["character_changes"]
    if not tracked:
        assert text(out) == after
    else:
        assert any(d.accepted_view == after.replace("\n", "") for d in frag.diffs)


def test_tabs_and_nbsp_are_not_collapsed_and_headers_are_in_scope():
    src = document("one\t two\u00a0\u00a0three")
    f, out, rec = execute(src, "normalize_whitespace", {"tracked": False})
    assert rec.passed, rec.to_markdown()
    assert text(out) == "one\t two\u00a0\u00a0three"
    assert "word/header1.xml" in f.intended_parts


def test_noop_and_idempotence():
    f, out, rec = execute(document("10 psi"), "protect_nonbreaking_spaces", {"tracked": False})
    again, _, rec2 = execute(out, "protect_nonbreaking_spaces", {"tracked": False})
    assert rec.passed and rec2.passed
    assert again.property_assertions[0]["expected_edits"] == 0


def test_shared_transform_prediction_bug_is_caught_independently(monkeypatch):
    src = document("one  two")
    monkeypatch.setattr(hygiene, "character_edits", lambda text, op, params: [Edit(0, 3, "invented")])
    _, _, rec = execute(src, "normalize_whitespace", {"tracked": False}, select={"scope": ["document"]})
    assert rec.get("V10").status == "FAIL"


def test_same_paragraph_operations_require_a_chain():
    src = DocxPackage(document("one  -- two"))
    with pytest.raises(JobException, match="split into separate specs"):
        plan_spec(src, {"request": "tidy", "summary": "tidy", "mode": "untracked", "ops": [
            {"op": "normalize_whitespace", "params": {"tracked": False}},
            {"op": "normalize_punctuation", "params": {"dash": "em", "tracked": False}}]})


def test_capitalization_is_per_occurrence_and_stale_sources_refuse(tmp_path):
    path = tmp_path / "section.docx"
    src = document("owner and owner and ownership. work")
    path.write_bytes(src)
    records = hygiene.propose_defined_terms(DocxPackage.from_path(path), ["Owner", "Work"])
    assert len(records) == 3
    spec = {"request": "capitalize Owner", "summary": "one occurrence", "mode": "untracked", "ops": [
        {"op": "capitalize_defined_terms", "params": {"tracked": False, "approvals": records[:1]}}]}
    validate_spec(spec)
    _, out, rec = execute(src, "capitalize_defined_terms", spec["ops"][0]["params"], name=str(path))
    assert rec.passed, rec.to_markdown()
    assert text(out) == "Owner and owner and ownership. work"
    _, _, reversed_record = execute(src, "capitalize_defined_terms", {"approvals": list(reversed(records))}, name=str(path))
    assert reversed_record.passed, reversed_record.to_markdown()
    batch = compiler._Batch([path], None)
    assert "own user pick" in " ".join(compiler.check_submission(spec, batch, ["capitalize Owner"]))
    assert compiler.check_submission(spec, batch, [hygiene.approval_source(records[0])]) == []
    spec["ops"][0]["params"]["approvals"] = records[:2]
    assert compiler.check_submission(spec, batch, [hygiene.approval_source(records[0])])
    with pytest.raises(JobException, match="stale"):
        execute(document("owner and new text"), "capitalize_defined_terms", {"approvals": records[:1]}, name=str(path))
    bypass = {**spec, "ops": [{"op": "replace_table", "params": {"tracked": False,
                        "pairs": [{"find": "owner", "replace": "Owner"}]}}]}
    assert "global replacement" in " ".join(compiler.check_submission(bypass, batch, ["capitalize Owner"]))


def test_capitalization_choices_come_from_inspected_occurrences(tmp_path):
    from spec_surgeon.conversation import Conversation
    path = tmp_path / "section.docx"
    path.write_bytes(document("owner and work"))
    conv = Conversation("review defined terms", [path])
    proposal = conv.batch.propose_defined_terms(path.name, ["Owner"])["proposals"][0]
    errors = []
    choices = conv._choices_from([{"kind": "capitalization", "proposal_id": proposal["id"],
                                   "label": "misleading", "wording": "invented"}], errors)
    assert not errors and len(choices) == 1
    choice = choices[0]
    assert "owner" in choice.label and "misleading" not in choice.label
    assert choice.words == ["Owner"] and choice.snapshot["id"] == proposal["id"]
    assert hygiene.approval_source(proposal) not in conv.sources
    assert not conv._choices_from([{"kind": "capitalization", "proposal_id": "uninspected"}], errors)


def test_a_real_conversation_pick_approves_only_that_source_occurrence(tmp_path):
    from fake_claude import FakeClient, resp, tool
    from spec_surgeon.conversation import Conversation
    path = tmp_path / "section.docx"
    path.write_bytes(document("owner and owner"))
    record = hygiene.propose_defined_terms(DocxPackage.from_path(path), ["Owner"])[0]
    spec = {"request": "review", "summary": "one approved occurrence", "mode": "content", "ops": [
        {"op": "capitalize_defined_terms", "params": {"approvals": [record]}}]}
    fake = FakeClient(
        resp(tool("propose_defined_terms", {"file": path.name, "terms": ["Owner"]}, "p1")),
        resp(tool("ask_user", {"understanding": "Review defined terms", "message": "Review this occurrence",
                              "choices": [{"kind": "capitalization", "proposal_id": record["id"]}]}, "a1")),
        resp(tool("submit_editspec", {"spec": spec}, "s1")))
    c = Conversation("review defined terms", [path], client=fake)
    turn = c.start()
    assert turn.status == "questions" and len(turn.choices) == 1
    assert hygiene.approval_source(record) not in c.sources
    assert c.reply(choice_id=turn.choices[0].id).status == "ok"
    assert hygiene.approval_source(record) in c.sources
    assert len([s for s in c.sources if s.startswith("approved-capitalization:")]) == 1


def test_preview_enumerates_every_files_characters_and_nbsp(tmp_path):
    from spec_surgeon.planner import plan_batch
    paths = [tmp_path / "a.docx", tmp_path / "b.docx"]
    for path in paths:
        path.write_bytes(document("10 psi; 21 13 19."))
    spec = {"request": "protect spaces", "summary": "protect spaces", "mode": "content", "ops": [
        {"op": "protect_nonbreaking_spaces", "params": {}}]}
    plan = plan_batch([str(p) for p in paths], spec)
    assert not plan.exceptions
    rendered = plan.to_text(per_file=False)
    assert all(str(path) in rendered for path in paths)
    assert rendered.count('"\\u00a0"') == 6
    assert all(len(file["character_changes"]) == 3 for file in plan.to_dict()["files"])


def test_field_results_are_refused_for_hygiene():
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    d = Document()
    p = d.add_paragraph()
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "DOCPROPERTY ProjectName")
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "two  spaces"
    run.append(t)
    field.append(run)
    p._p.append(field)
    b = io.BytesIO()
    d.save(b)
    with pytest.raises(JobException) as error:
        execute(b.getvalue(), "normalize_whitespace", {})
    assert error.value.code == "field-result"


@pytest.mark.parametrize("op,params", [
    ("normalize_punctuation", {}), ("normalize_punctuation", {"quotes": "guess"}),
    ("protect_nonbreaking_spaces", {"units": [""]}), ("capitalize_defined_terms", {"approvals": []})])
def test_bad_choices_refuse(op, params):
    with pytest.raises(JobException):
        get_op(op).plan(DocxPackage(document("body")), {}, params)


def test_empty_scope_does_not_silently_expand_to_the_document():
    with pytest.raises(JobException, match="scope"):
        get_op("normalize_whitespace").plan(DocxPackage(document("two  spaces")), {"scope": []}, {})


def test_capitalization_preserves_formatting_of_unchanged_letters(tmp_path):
    from spec_surgeon.package import W
    path = tmp_path / "case.docx"
    d = Document()
    p = d.add_paragraph()
    p.add_run("o").bold = True
    p.add_run("wner").italic = True
    d.save(path)
    record = hygiene.propose_defined_terms(DocxPackage.from_path(path), ["Owner"])[0]
    frag, out, rec = execute(path.read_bytes(), "capitalize_defined_terms", {"approvals": [record]}, name=str(path))
    assert rec.passed, rec.to_markdown()
    assert frag.transform_data["character_changes"][0]["before"] == "o"
    assert frag.transform_data["character_changes"][0]["after"] == "O"
    root = DocxPackage(out).xml("word/document.xml")
    remaining = next(t for t in root.iter(W+"t") if t.text == "wner")
    assert remaining.getparent().find(W+"rPr/"+W+"i") is not None


def test_hygiene_chain_verifies_each_stage_and_never_transfers_capitalization_approval(tmp_path):
    from spec_surgeon.chain import ChainError, chain_from_specs, run_chain
    path = tmp_path / "section.docx"
    original = document('owner  said \"Work\" -- 10  psi; Section 21 13 19.')
    path.write_bytes(original)
    record = hygiene.propose_defined_terms(DocxPackage.from_path(path), ["Owner"])[0]
    specs = [("case", "capitalize_defined_terms", {"approvals": [record]}),
             ("spaces", "normalize_whitespace", {}),
             ("punctuation", "normalize_punctuation", {"quotes": "curly", "dash": "em"}),
             ("nonbreaking", "protect_nonbreaking_spaces", {})]
    stages = [(name, {"request": name, "summary": name, "mode": "content", "ops": [{"op": op, "params": params}]})
              for name, op, params in specs]
    chain = chain_from_specs(stages)
    result = run_chain([str(path)], chain, tmp_path / "out", workers=1, do_render=False)
    assert result.files[0].status == "PASS", result.files[0].reason
    assert result.files[0].stage_status == ["PASS"]*4
    assert text(Path(result.files[0].delivered_to).read_bytes()) == (
        'Owner said “Work” — 10\u00a0psi; Section 21\u00a013\u00a019.')
    assert path.read_bytes() == original
    with pytest.raises(ChainError, match="first stage"):
        chain_from_specs(stages[1:2] + stages[:1])
