"""Regressions for paragraph preservation, repeated text and usable render evidence."""
from __future__ import annotations

import io
import zipfile

import pymupdf
import pytest
from docx import Document
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

from spec_surgeon.batch import geometry_check
from spec_surgeon.editspec import apply_spec, plan_spec
from spec_surgeon.package import DocxPackage
from spec_surgeon.verify import run_suite
from spec_surgeon.verify.suite import VerificationRecord
from spec_surgeon.verify.geometry import compare_pdfs
from spec_surgeon.verify.props import sig
from spec_surgeon.verify.reader import VPackage, W
from spec_surgeon.verify.textcheck import reject_view


def saved(doc):
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


def mutate(data, callback, part="word/document.xml"):
    source, destination = zipfile.ZipFile(io.BytesIO(data)), io.BytesIO()
    with zipfile.ZipFile(destination, "w") as output:
        for info in source.infolist():
            content = source.read(info.filename)
            if info.filename == part:
                root = etree.fromstring(content)
                callback(root)
                content = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            output.writestr(info, content)
    return destination.getvalue()


def execute(source, op, params, select=None):
    mode = "content" if params.get("tracked", True) else "untracked"
    spec = {"request": "Approved exact edit", "summary": "Approved exact edit", "mode": mode,
            "ops": [{"op": op, "select": select or {}, "params": {"author": "Audit User", **params}}]}
    package = DocxPackage(source)
    plan, fragments = plan_spec(package, spec)
    output = apply_spec(package, fragments).output_bytes()
    record = run_suite(source, output, plan.to_dict(), "source.docx", "output.docx")
    assert record.passed, record.to_markdown()
    return plan.to_dict(), output


@pytest.mark.parametrize("tracked", [False, True])
@pytest.mark.parametrize("prefix", [False, True])
def test_first_replacement_in_repeated_final_paragraphs(tracked, prefix):
    doc = Document()
    if prefix:
        doc.add_paragraph("Introduction")
    doc.add_paragraph("Owner")
    doc.add_paragraph("Owner")
    source = saved(doc)
    execute(source, "replace_text", {"find": "Owner", "replace": "Client", "occurrences": "first",
                                     "tracked": tracked})


@pytest.mark.parametrize("tracked", [False, True])
def test_repeated_text_does_not_allow_editing_the_wrong_paragraph(tracked):
    doc = Document()
    doc.add_paragraph("Owner")
    doc.add_paragraph("Owner")
    source = saved(doc)
    plan, output = execute(source, "replace_text", {"find": "Owner", "replace": "Client", "tracked": tracked},
                           {"paragraphs": {"index_range": [0, 1]}})

    def wrong_paragraph(root):
        body = root.find(W + "body")
        first, second = list(root.iter(W + "p"))
        body.remove(second)
        body.insert(0, second)

    record = run_suite(source, mutate(output, wrong_paragraph), plan, "source.docx", "output.docx")
    assert record.get("V10").status == "FAIL"


def marked_document():
    doc = Document()
    paragraph = doc.add_paragraph()
    paragraph.add_run("Before ")
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), "4")
    start.set(qn("w:name"), "ImportantClause")
    paragraph._p.append(start)
    paragraph.add_run("BOOKMARKED ")
    proof = OxmlElement("w:proofErr")
    proof.set(qn("w:type"), "spellStart")
    paragraph._p.append(proof)
    paragraph.add_run("TEXT")
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), "4")
    paragraph._p.append(end)
    paragraph.add_run(" After")
    doc.add_paragraph("Retained paragraph")
    return saved(doc)


def paragraph_contents(root):
    return [sig(child) for child in root.find(".//" + W + "p") if child.tag != W + "pPr"]


def test_tracked_deletion_preserves_bookmark_and_proofing_boundaries_on_reject():
    source = marked_document()
    plan, output = execute(source, "delete_paragraphs", {"tracked": True},
                           {"paragraphs": {"index_range": [0, 1]}})
    assert plan["rev_delta"]["del"] == 4
    assert paragraph_contents(VPackage(source).xml("word/document.xml")) == paragraph_contents(
        reject_view(VPackage(output).xml("word/document.xml")))

    def collapse_markers(root):
        paragraph = root.find(".//" + W + "p")
        for child in list(paragraph):
            if child.tag in {W + "bookmarkStart", W + "bookmarkEnd", W + "proofErr"}:
                paragraph.remove(child)
                paragraph.append(child)

    record = run_suite(source, mutate(output, collapse_markers), plan, "source.docx", "output.docx")
    assert record.get("V10").status == "FAIL"


def rich_document():
    doc = Document()
    paragraph = doc.add_paragraph()
    paragraph.add_run("Warning: ").bold = True
    paragraph.add_run("Owner shall retain ")
    paragraph.add_run("\t")
    paragraph.add_run("content; see ")
    link = OxmlElement("w:hyperlink")
    link.set(qn("w:anchor"), "ApprovedClause")
    run, text = OxmlElement("w:r"), OxmlElement("w:t")
    text.text = "reference"
    run.append(text)
    link.append(run)
    paragraph._p.append(link)
    paragraph.add_run(".")
    return saved(doc)


@pytest.mark.parametrize("tracked", [False, True])
@pytest.mark.parametrize("damage", ["bold", "link", "tab", "metadata", "empty_run"])
def test_literal_edit_checks_unchanged_character_format_links_and_structural_positions(tracked, damage):
    source = rich_document()
    plan, output = execute(source, "replace_text", {"find": "Owner", "replace": "Client", "tracked": tracked})

    def change(root):
        paragraph = root.find(".//" + W + "p")
        if damage == "bold":
            first = paragraph.find(W + "r")
            first.remove(first.find(W + "rPr"))
        elif damage == "link":
            paragraph.find(W + "hyperlink").set(W + "anchor", "UnapprovedClause")
        elif damage == "tab":
            tab = paragraph.find(".//" + W + "tab")
            tab.getparent().remove(tab)
            list(paragraph.iter(W + "r"))[-1].append(tab)
        elif damage == "metadata":
            paragraph.set(W + "rsidR", "DEADBEEF")
        else:
            etree.SubElement(paragraph, W + "r")

    record = run_suite(source, mutate(output, change), plan, "source.docx", "output.docx")
    assert record.get("V10").status == "FAIL", record.to_markdown()


def test_predicted_wording_does_not_exempt_a_paragraph_without_an_independent_proof():
    source = rich_document()
    plan, output = execute(source, "replace_text", {"find": "Owner", "replace": "Client", "tracked": False})
    plan["property_assertions"] = []
    record = run_suite(source, output, plan, "source.docx", "output.docx")
    assert record.get("V10").status == "FAIL"
    assert "no independent paragraph preservation proof" in record.get("V10").summary


@pytest.mark.parametrize("tracked", [False, True])
@pytest.mark.parametrize("damage", ["font", "link", "tab"])
def test_custom_edit_primitive_preserves_items_outside_its_declared_range(tracked, damage):
    source = rich_document()
    plan, output = execute(source, "replace_text", {"find": "Owner", "replace": "Client", "tracked": tracked})
    plan["property_assertions"] = [{"kind": "paragraph_edits", "part": "word/document.xml", "block_index": 0,
                                     "edits": [[9, 14, "Client", False]], "tracked": tracked,
                                     "author": "Audit User"}]
    assert run_suite(source, output, plan, "source.docx", "output.docx").passed

    def change(root):
        paragraph = root.find(".//" + W + "p")
        if damage == "font":
            first = paragraph.find(W + "r")
            first.remove(first.find(W + "rPr"))
        elif damage == "link":
            paragraph.find(W + "hyperlink").set(W + "anchor", "WrongClause")
        else:
            tab = paragraph.find(".//" + W + "tab")
            tab.getparent().remove(tab)
            list(paragraph.iter(W + "r"))[-1].append(tab)

    assert run_suite(source, mutate(output, change), plan, "source.docx", "output.docx").get("V10").status == "FAIL"


def test_checked_paragraph_cannot_relocate_into_a_table_without_changing_its_global_index():
    doc = Document()
    doc.add_paragraph("Owner")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Retained cell"
    source = saved(doc)
    plan, output = execute(source, "replace_text", {"find": "Owner", "replace": "Client", "tracked": False})

    def move(root):
        paragraph = root.find(".//" + W + "p")
        cell = root.find(".//" + W + "tc")
        paragraph.getparent().remove(paragraph)
        cell.insert(1, paragraph)

    record = run_suite(source, mutate(output, move), plan, "source.docx", "output.docx")
    assert record.get("V10").status == "FAIL"


def test_pass_record_explicitly_reports_unevaluated_checks_and_automated_scope():
    source = rich_document()
    plan, output = execute(source, "replace_text", {"find": "Owner", "replace": "Client", "tracked": True})
    record = run_suite(source, output, plan, "source.docx", "output.docx")
    data = record.to_dict()
    coverage = data["verification_coverage"]
    assert record.passed
    assert set(coverage["not_applicable_checks"]) == {"V6", "V7", "V9"}
    assert "V10" in coverage["performed_checks"] and "V6" not in coverage["performed_checks"]
    assert coverage["human_visual_acceptance"] == "not_recorded_by_this_suite"
    assert coverage["renderer"] is None
    assert "NA is not evidence of success" in record.to_markdown()
    assert "Human visual acceptance is not recorded" in record.to_markdown()
    assert VerificationRecord.from_dict(data).to_dict() == data


@pytest.mark.parametrize("multiple_paragraphs", [False, True])
def test_cover_numbering_removal_preserves_the_retained_footer_envelope(multiple_paragraphs):
    doc = Document()
    doc.add_paragraph("Cover")
    footer = doc.sections[0].footer
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run, text = OxmlElement("w:r"), OxmlElement("w:t")
    text.text = "1"
    run.append(text)
    field.append(run)
    footer.paragraphs[0]._p.append(field)
    if multiple_paragraphs:
        footer.add_paragraph("Retained footer wording")
    doc.add_section(WD_SECTION.NEW_PAGE).footer.is_linked_to_previous = False
    source = saved(doc)
    package = DocxPackage(source)
    plan, fragments = plan_spec(package, {"request": "Remove cover page number", "summary": "Remove cover number",
        "mode": "formatting", "ops": [{"op": "fix_page_numbering", "params": {"option": "cover-unnumbered"}}]})
    output = apply_spec(package, fragments).output_bytes()
    record = run_suite(source, output, plan.to_dict(), "source.docx", "output.docx")
    assert record.passed, record.to_markdown()
    footer_part = next(a["footer"] for a in plan.property_assertions if a["kind"] == "cover_unnumbered")

    def change(root):
        root.find(".//" + W + "p").set(W + "rsidR", "DEADBEEF")

    record = run_suite(source, mutate(output, change, footer_part), plan.to_dict(), "source.docx", "output.docx")
    assert record.get("V10").status == "FAIL"


def pdf(path, text):
    document = pymupdf.open()
    page = document.new_page()
    if text:
        page.insert_text((72, 72), text)
    document.save(path)
    document.close()
    return path


@pytest.mark.parametrize("output_text", ["", "This approved clause", "This approved clause extra words"])
def test_source_geometry_requires_word_correspondence_and_usable_anchors(tmp_path, output_text):
    source = pdf(tmp_path / "source.pdf", "This approved clause must remain visible.")
    output = pdf(tmp_path / "output.pdf", output_text)
    evidence = compare_pdfs(source, output)
    check = geometry_check({"mode": "formatting", "geometry_reference": "source"}, evidence)
    assert check.status == "FAIL", check.summary


def test_source_geometry_distinguishes_real_anchors_from_two_blank_pages(tmp_path):
    source = pdf(tmp_path / "source.pdf", "Identical clause.")
    output = pdf(tmp_path / "output.pdf", "Identical clause.")
    assert geometry_check({"mode": "formatting"}, compare_pdfs(source, output)).status == "PASS"
    blank = pdf(tmp_path / "blank.pdf", "")
    assert geometry_check({"mode": "formatting"}, compare_pdfs(blank, blank)).status == "NA"
    assert geometry_check({"mode": "formatting"}, compare_pdfs(source, blank, [1])).status != "FAIL"
