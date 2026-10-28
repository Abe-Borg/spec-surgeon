"""Bounded patterns (stress-test finding F8, P2): every pattern a request or the model writes is compiled through
``ops/patterns.compile_pattern`` and searched with a two-second deadline through the ``regex`` module (its default
``V0`` behaviour is ``re``'s); the verifier searches through its own twin, ``verify/patterns``. A search past the
deadline is the refusal ``regex-timeout`` that sets the file aside; in the verifier it is a failed check with that
reason, never an exception out of the suite.

The stress test's own pattern, ``(a+)+$``, is one ``regex`` finishes at once (its engine sees through that
repetition), so the catastrophic case here is ``(a|aa)+$`` on forty letters and a stop, which both engines chase
exponentially (1.7 s on thirty-one letters, 20 s on thirty-six, minutes on forty); the stress pattern is still proved
bounded.

The planted-defect rule of CLAUDE.md, for this change: a pattern that would match differently under ``re`` and
``regex`` is not expected — the parity test below holds ``V0`` to ``re``'s results on the patterns specifications
use — and a flag passed through as a number instead of translated (``re.ASCII``'s value is ``regex.V1``'s) is caught
by the ASCII test on both sides."""
from __future__ import annotations

import re
import time

import pytest
import regex
from docx import Document

from make_fixtures import build_rich, docx_bytes
from spec_surgeon import compiler
from spec_surgeon.catalog import SELECTOR_GRAMMAR
from spec_surgeon.ops import JobException, get_op, merge_fragments
from spec_surgeon.ops import patterns
from spec_surgeon.ops.patterns import MAX_PATTERN_CHARS, REGEX_TIMEOUT_S, compile_pattern
from spec_surgeon.package import DocxPackage
from spec_surgeon.report import FAILURE_EXPLANATIONS
from spec_surgeon.verify import patterns as v_patterns
from spec_surgeon.verify import run_suite
from spec_surgeon.verify.patterns import v_compile, v_search

AUTHOR = "Abraham Borg"
SLOW = r"(a|aa)+$"              # exponential on a run of a's that ends in something else, under re and regex alike
STRESS = r"(a+)+$"              # the stress test's pattern (F8): re never returns on it, regex finishes it at once
TEXT = "a" * 40 + "!"
GRACE = 1.5                     # seconds beyond the deadline the whole plan may take (loaded test workers)


@pytest.fixture(scope="module")
def doc() -> bytes:
    d = Document()
    d.add_paragraph("PART 1 - GENERAL")
    d.add_paragraph(TEXT)
    d.add_paragraph("1.01 SUMMARY")
    return docx_bytes(d)


def plan(src: bytes, select, params=None):
    return get_op("set_paragraph_props").plan(DocxPackage(src), select,
                                             params or {"spacing": {"after": 120}, "author": AUTHOR})


def run(src: bytes, select):
    pkg = DocxPackage(src)
    frag = plan(src, select)
    merged = merge_fragments([frag])
    work = pkg.copy()
    get_op(frag.op).apply(work, frag)
    out = work.output_bytes()
    return merged.to_dict(), out


@pytest.fixture
def quick(monkeypatch):
    """A short deadline on both sides, so the tests that only need *a* deadline do not wait two seconds each."""
    monkeypatch.setattr(patterns, "REGEX_TIMEOUT_S", 0.2)
    monkeypatch.setattr(v_patterns, "V_REGEX_TIMEOUT_S", 0.2)


# ---- the transform side --------------------------------------------------------------------------------

def test_a_catastrophic_text_regex_sets_the_file_aside_within_the_deadline(doc):
    """End to end at the real deadline: the selector's search on the forty-letter paragraph is cut off at two seconds
    and the plan is the refusal, well under the stress test's cutoff."""
    started = time.perf_counter()
    with pytest.raises(JobException) as e:
        plan(doc, {"paragraphs": {"text_regex": SLOW}})
    took = time.perf_counter() - started
    assert e.value.code == "regex-timeout"
    assert SLOW in e.value.reason and "simplify it" in e.value.reason and "paragraph's text" in e.value.reason
    assert e.value.details["timeout_s"] == REGEX_TIMEOUT_S == 2.0
    assert took < REGEX_TIMEOUT_S + GRACE, took


def test_the_stress_test_s_own_pattern_is_bounded_too(doc):
    """``(a+)+$``: regex answers at once (no match) — and whatever a future engine does with it, the deadline holds."""
    started = time.perf_counter()
    try:
        assert compile_pattern(STRESS, where="a paragraph's text").search(TEXT) is None
    except JobException as e:
        assert e.code == "regex-timeout"
    assert time.perf_counter() - started < REGEX_TIMEOUT_S + GRACE


def test_an_ordinary_pattern_plans_and_verifies_as_before():
    rich = build_rich()
    d, out = run(rich, {"paragraphs": {"text_regex": r"^[A-F]\."}})
    assert d["rev_delta"]["pPrChange"] == 5
    rec = run_suite(rich, out, d, "s.docx", "o.docx")
    assert rec.passed, rec.to_markdown()


def test_every_search_method_is_bounded(quick):
    rx = compile_pattern(SLOW, where="a paragraph's text")
    for call in (lambda: rx.search(TEXT), lambda: rx.match(TEXT), lambda: rx.fullmatch(TEXT), lambda: rx.findall(TEXT),
                 lambda: list(rx.finditer(TEXT)), lambda: rx.sub("", TEXT), lambda: rx.subn("", TEXT),
                 lambda: rx.split(TEXT)):
        with pytest.raises(JobException) as e:
            call()
        assert e.value.code == "regex-timeout"


def test_a_long_pattern_a_bad_pattern_and_a_non_string_are_bad_params(doc):
    for bad, words in (("x" * (MAX_PATTERN_CHARS + 1), "characters long"), ("(", "not a regular expression"),
                       (("a" * 10 + "(") * 3, "not a regular expression")):
        with pytest.raises(JobException) as e:
            plan(doc, {"paragraphs": {"text_regex": bad}})
        assert e.value.code == "bad-params" and words in e.value.reason, e.value.reason
    with pytest.raises(JobException) as e:
        compile_pattern(5, where="a paragraph's text")
    assert e.value.code == "bad-params" and "must be text" in e.value.reason
    assert compile_pattern("x" * MAX_PATTERN_CHARS, where="t").search("x" * MAX_PATTERN_CHARS) is not None
    # a pattern the app builds from the user's own literal is as long as the words: no length limit on it
    assert compile_pattern(re.escape("." * 2000), where="t", max_chars=None).search("." * 2000) is not None


def test_flags_are_translated_never_passed_through(doc):
    """re.ASCII's number is regex.V1's: passed through, ``\\w`` would match a letter outside ASCII and the pattern
    would run under another grammar. case_insensitive still means IGNORECASE, and VERBOSE reads back in re's
    numbering (the selector's flexible_whitespace rule relies on it)."""
    assert compile_pattern(r"\w", re.ASCII, where="t").search("é") is None
    assert compile_pattern(r"\w", where="t").search("é") is not None
    assert compile_pattern("(?x) a b", where="t").flags & re.VERBOSE
    assert not compile_pattern("ab", where="t").flags & re.VERBOSE
    assert compile_pattern("ab", re.IGNORECASE, where="t").flags & re.IGNORECASE
    frag = plan(doc, {"paragraphs": {"text_regex": "^part", "case_insensitive": True}})
    assert frag.rev_delta["pPrChange"] == 1
    with pytest.raises(ValueError):
        compile_pattern("a", 1 << 20, where="t")                   # an unknown bit is a programming error


def test_a_match_s_groups_are_the_pattern_s():
    rx = compile_pattern(r"(?P<tag>\[[A-Z]+\])\s*(.*)", where="t")
    assert rx.groups == 2 and dict(rx.groupindex) == {"tag": 1}
    m = rx.search("[WET] Piping")
    assert m.group("tag") == "[WET]" and m.re.groups == 2 and "tag" in m.re.groupindex


# ---- the verifier's twin -------------------------------------------------------------------------------

def test_the_twin_times_out_into_a_value_error_with_the_reason(quick):
    rx = v_compile(SLOW, where="a paragraph's text")
    for call in (lambda: rx.search(TEXT), lambda: rx.match(TEXT), lambda: rx.fullmatch(TEXT), lambda: rx.findall(TEXT),
                 lambda: list(rx.finditer(TEXT)), lambda: rx.sub("", TEXT), lambda: rx.split(TEXT),
                 lambda: v_search(SLOW, TEXT, where="a paragraph's text")):
        with pytest.raises(ValueError) as e:
            call()
        assert "regex-timeout" in str(e.value) and SLOW in str(e.value) and "paragraph's text" in str(e.value)
    for bad in (5, "(", "x" * (v_patterns.V_MAX_PATTERN_CHARS + 1)):
        with pytest.raises(ValueError) as e:
            v_compile(bad)
        assert "the op must refuse" in str(e.value)
    assert v_compile(r"\w", re.ASCII).search("é") is None and v_compile(r"\w").search("é") is not None
    assert v_compile("(?x) a b").flags & re.VERBOSE and not v_compile("ab").flags & re.VERBOSE


def test_a_timeout_in_the_verifier_is_a_failed_check_not_an_exception(doc, quick):
    """A plan whose assertion carries the catastrophic pattern (the selector the verifier re-resolves on its own):
    V10 fails with the reason; the suite completes."""
    d, out = run(doc, {"paragraphs": {"text_regex": "^PART"}})
    assert run_suite(doc, out, d, "s.docx", "o.docx").passed
    d["property_assertions"][0]["select"]["paragraphs"]["text_regex"] = SLOW
    rec = run_suite(doc, out, d, "s.docx", "o.docx")
    v10 = rec.get("V10")
    assert v10.status == "FAIL" and "regex-timeout" in v10.summary and SLOW in v10.summary, v10.summary
    d["property_assertions"][0]["select"]["paragraphs"]["text_regex"] = "("
    v10 = run_suite(doc, out, d, "s.docx", "o.docx").get("V10")
    assert v10.status == "FAIL" and "does not compile" in v10.summary, v10.summary


# ---- V0 is re ------------------------------------------------------------------------------------------

PARITY = [
    (r"^PART\s+\d", "PART 2 - PRODUCTS"),
    (r"(?i)^end\s+of\s+section\b", "END OF SECTION 21 13 13"),
    (r"^\d+\.\d+\s", "2.01 MANUFACTURERS"),
    (r"^[A-F]\.", "A. Comply with NFPA 13."),
    (r"\bNFPA\s+13\b", "A. Comply with NFPA 13."),
    (r"(?x) ^ \s* 2\.01 ", "2.01 MANUFACTURERS"),
    (r"[^\W\d_]+(?:['’][^\W\d_]+)*", "the owner's representative"),
    (r"SECTION\s+(?P<number>\d{2}\s?\d{2}\s?\d{2})", "SECTION 21 13 13 – WET-PIPE SPRINKLER SYSTEMS"),
    (r"\s*[/,;]\s*", "WET / DRY, PREACTION"),
    (r"x*", "ab"),
    (r"\b", "a b"),
    (r"\s", "non breaking"),
    (r"(?<![\w.])21(?!\w)(?!\.\w)", "see 21 13 13 and 21.5"),
]


@pytest.mark.parametrize("pattern,text", PARITY)
def test_v0_matches_as_re_does(pattern, text):
    """The regex module's default behaviour (V0) is re's; the twins return what re returns on the patterns and texts
    specifications are made of, zero-width matches included."""
    assert regex.DEFAULT_VERSION == regex.V0
    ours, theirs = compile_pattern(pattern, where="t"), v_compile(pattern)
    ref = re.compile(pattern)
    for rx in (ours, theirs):
        got, want = rx.search(text), ref.search(text)
        assert (got is None) == (want is None)
        if want is not None:
            assert got.span() == want.span() and got.groups() == want.groups() and got.groupdict() == want.groupdict()
        assert rx.findall(text) == ref.findall(text)
        assert [m.span() for m in rx.finditer(text)] == [m.span() for m in ref.finditer(text)]
        assert rx.sub("-", text) == ref.sub("-", text)
        assert rx.split(text) == ref.split(text)


# ---- the words -----------------------------------------------------------------------------------------

def test_the_refusal_the_catalog_and_the_prompt_say_it():
    title, text = FAILURE_EXPLANATIONS["regex-timeout"]
    assert "two seconds" in text and "set aside" in text
    assert "two seconds" in SELECTOR_GRAMMAR["patterns"] and "1000 characters" in SELECTOR_GRAMMAR["patterns"]
    assert "a pattern that takes more than two seconds on one paragraph sets the file aside" in compiler.SYSTEM
    assert compiler.SYSTEM.count("two seconds") == 1
