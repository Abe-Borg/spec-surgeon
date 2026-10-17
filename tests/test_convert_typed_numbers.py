"""convert_typed_numbers and V10's ``typed_numbers`` kind (built-in batch 5, item 34 ★ — docs/BUILTIN_OPS_BATCH5_DESIGN.md
§2): end to end on the typed spec (build_typed: tabs, spaces, leading whitespace; tracked and plain; numbers split across
runs) converted to the example file's scheme; the guarantee "the document reads the same" (a typo, a list that already
numbers paragraphs, a skipped level, an ambiguous typed number, a level without a suffix each refuse); held schemes;
tables; cleared formatting; the display rule (hidden, capitalised, or a mark that makes the number white, tiny, raised,
struck or symbols — the adversarial review's F1/F4); the review's other fixes (flags for the new style's capitals and
linked levels, held styles unlike the example file, tracked over replaced styles, the messages); wrong OUTPUTS and wrong
IMPLEMENTATIONS shared by plan and apply, each caught by V10;
plan-time conflicts (rule 9); the template as an input (rule 12); the worked example; and the LibreOffice proof — the
converted file rendered to PDF shows every typed number exactly once, in order."""
from __future__ import annotations

import dataclasses
import io
import shutil
import subprocess
import zipfile
from pathlib import Path
from typing import Any, Dict, List, Tuple

import pytest
from docx import Document
from lxml import etree

from make_fixtures import SCHEME_LEVELS, build_scheme_example, build_typed, patch_docx
from spec_surgeon.catalog import COMMON_ASSUMPTIONS, op_catalog
from spec_surgeon.editspec import apply_spec, load_spec, plan_spec, validate_spec
from spec_surgeon.ops import JobException, get_op, merge_fragments
from spec_surgeon.ops import lists as tl
from spec_surgeon.ops.selectors import StyleSheet
from spec_surgeon.package import DocxPackage, W, serialize_xml
from spec_surgeon.report import FAILURE_EXPLANATIONS
from spec_surgeon.verify import run_suite

ROOT = Path(__file__).resolve().parent.parent
MAIN = "word/document.xml"
NUMBERING = "word/numbering.xml"
AUTHOR = "Abraham Borg"
W_DECL = f'xmlns:w="{W[1:-1]}"'

# build_typed(): the typed number of each paragraph, by the text that follows it
TYPED = {"GENERAL": "PART 1 -", "SUMMARY": "1.1", "Section includes wet-pipe sprinkler systems.": "A.",
         "Related requirements:": "B.", 'Division 21 Section "Common Work Results for Fire Suppression".': "1.",
         "Hangers.": "a.", "Supports.": "b.", 'Division 28 Section "Fire Detection and Alarm".': "2.",
         "Submittals.": "C.", "REFERENCES": "1.2", "NFPA 13.": "A.", "PRODUCTS": "PART 2 -", "MANUFACTURERS": "2.1",
         "Tyco.": "A.", "Viking.": "B.", "Reliable.": "C.", "MATERIALS": "2.2", "Steel pipe.": "A.",
         "EXECUTION": "PART 3 -", "INSTALLATION": "3.1", "Install per NFPA 13.": "A.", "Hangers per NFPA 13.": "1."}


def live(p) -> str:
    """What a reader sees of a paragraph's own runs: w:t and a run's w:tab outside w:del / w:moveFrom."""
    out = []
    for el in p.iter(W + "t", W + "tab"):
        q, gone = el.getparent(), False
        while q is not None and q.tag != W + "p":
            gone = gone or q.tag in (W + "del", W + "moveFrom")
            q = q.getparent()
        if q is p and not gone and (el.tag == W + "t" or el.getparent().tag == W + "r"):
            out.append(el.text or "" if el.tag == W + "t" else "\t")
    return "".join(out)


def find(root, text: str):
    return next(p for p in root.iter(W + "p") if live(p) == text)


def _x(xml: str):
    return etree.fromstring(f"<w:x {W_DECL}>{xml}</w:x>")[0]


@pytest.fixture()
def tpl(tmp_path) -> str:
    p = tmp_path / "NUMBERING SCHEME.docx"
    p.write_bytes(build_scheme_example())
    return str(p)


def run_ops(data: bytes, specs: List[Tuple[str, dict, dict]]):
    pkg = DocxPackage(data)
    frags = [get_op(o).plan(pkg, sel, par) for o, sel, par in specs]
    plan = merge_fragments(frags)
    work = pkg.copy()
    for f in frags:
        get_op(f.op).apply(work, f)
    out = work.output_bytes()
    return plan, out, run_suite(data, out, plan.to_dict(), "src.docx", "out.docx")


def converting(data: bytes, tpl: str = None, select: dict = None, **params):
    if tpl is not None:
        params["template"] = tpl
    params.setdefault("levels", list(SCHEME_LEVELS))
    return run_ops(data, [("convert_typed_numbers", select or {}, {"author": AUTHOR, **params})])


def failed(rec) -> List[str]:
    return [c.id for c in rec.checks if c.status == "FAIL"]


def v10(rec) -> str:
    c = next(c for c in rec.checks if c.id == "V10")
    return c.summary + " | " + " | ".join(c.data.get("errors", []))


def numbers(data: bytes) -> Dict[str, Any]:
    """the live text of each numbered paragraph → its rendered number (by the transform's model)."""
    pkg = DocxPackage(data)
    root = pkg.xml(MAIN)
    pars = list(root.iter(W + "p"))
    return {live(pars[x.index]): (x.text if x.evaluated else (x.issue,))
            for x in tl.count(root, StyleSheet(pkg), tl.ListDefs.load(pkg))}


def refused(data: bytes, tpl: str = None, select: dict = None, **params) -> JobException:
    if tpl is not None:
        params["template"] = tpl
    params.setdefault("levels", list(SCHEME_LEVELS))
    with pytest.raises(JobException) as e:
        get_op("convert_typed_numbers").plan(DocxPackage(data), select or {}, params)
    return e.value


def docx_of(doc) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---- end to end -------------------------------------------------------------------------------------------------------

@pytest.mark.parametrize("variant", [{}, {"spaces": True}, {"leading_tab": True}])
@pytest.mark.parametrize("tracked", [True, False])
def test_typed_numbers_become_the_same_list_numbers(tpl, variant, tracked):
    """The starred case: every typed number of the body becomes the scheme's list number — the same number, the typed
    characters (and the space or tab after them, and any whitespace before) removed exactly as predicted; the table's
    typed numbers stay; tracked by default (the typed numbers as deletions)."""
    src = build_typed(**variant)
    plan, out, rec = converting(src, tpl, tracked=tracked)
    assert rec.passed, rec.to_markdown()
    got = numbers(out)
    assert {k: got.get(k) for k in TYPED} == TYPED
    root = DocxPackage(out).xml(MAIN)
    sep = " " if variant.get("spaces") else "\t"
    assert find(root, f"1.{sep}Cell item.") is not None and find(root, f"2.{sep}Cell item two.") is not None   # as typed
    p = find(root, "Section includes wet-pipe sprinkler systems.")
    assert p.find(f"{W}pPr/{W}pStyle").get(W + "val") == "PR1" and p.find(f"{W}pPr/{W}numPr") is None
    a = next(x for x in plan.property_assertions if x["kind"] == "typed_numbers")
    assert len(a["converted"]) == 22
    if tracked:
        assert plan.mode == "content" and plan.rev_delta["pPrChange"] == 22 and plan.rev_delta["del"] >= 22
        dels = "".join(t.text or "" for t in p.iter(W + "delText"))
        assert dels.strip(" ") == "A."                       # split across two runs: two deletions
        assert len(p.findall(W + "del")) == 2
        if not variant.get("spaces"):
            assert any(d.find(f"{W}r/{W}tab") is not None for d in p.findall(W + "del"))   # the tab goes with it
    else:
        assert plan.mode == "untracked" and not plan.rev_delta.get("pPrChange")
        assert p.find(f".//{W}del") is None and not any(True for _ in p.iter(W + "tab"))
        if not variant.get("spaces"):                         # V3: the tab after the number is declared gone
            assert any(d.struct_after is not None for d in plan.diffs)


def test_cleared_formatting_with_a_plain_conversion(tpl):
    def bold(root):
        p = find(root, "B.\tRelated requirements:")
        r = p.find(W + "r")
        r.insert(0, _x("<w:rPr><w:b/><w:i/></w:rPr>"))
        ppr = p.find(W + "pPr")
        if ppr is None:
            ppr = etree.SubElement(p, W + "pPr")
            p.insert(0, ppr)
        ppr.append(_x('<w:jc w:val="both"/>'))
    src = patch_docx(build_typed(), {MAIN: bold})
    plan, out, rec = converting(src, tpl, tracked=False, direct_formatting="clear")
    assert rec.passed, rec.to_markdown()
    p = find(DocxPackage(out).xml(MAIN), "Related requirements:")
    assert [etree.QName(c).localname for c in p.find(W + "pPr")] == ["pStyle"] and p.find(f"{W}r/{W}rPr") is None
    assert next(a for a in plan.property_assertions if a["kind"] == "typed_numbers")["expected_runs"] == 1
    e = refused(build_typed(), tpl, direct_formatting="clear")              # tracked by default: refused
    assert e.code == "bad-params"


def test_tables_and_a_narrowed_selection(tpl):
    plan, out, rec = converting(build_typed(), tpl, {"include_tables": True}, tracked=False)
    assert rec.passed, rec.to_markdown()
    got = numbers(out)
    assert got["Cell item."] == "1." and got["Cell item two."] == "2."
    plan, out, rec = converting(build_typed(), tpl, {"paragraphs": {"text_regex": "^PART"}}, tracked=False)
    assert rec.passed, rec.to_markdown()
    assert [v for v in numbers(out).values()] == ["PART 1 -", "PART 2 -", "PART 3 -"]


def test_own_numbering_off_is_removed(tpl):
    """A typed paragraph whose own numPr turns numbering off (numId 0) is a candidate; the numPr goes."""
    def off(root):
        p = find(root, "B.\tViking.")
        p.insert(0, _x('<w:pPr><w:numPr><w:ilvl w:val="0"/><w:numId w:val="0"/></w:numPr></w:pPr>'))
    src = patch_docx(build_typed(), {MAIN: off})
    plan, out, rec = converting(src, tpl, tracked=False)
    assert rec.passed, rec.to_markdown()
    p = find(DocxPackage(out).xml(MAIN), "Viking.")
    assert p.find(f"{W}pPr/{W}numPr") is None and numbers(out)["Viking."] == "B."


# ---- the guarantee: the document reads the same ----------------------------------------------------------------------

def test_a_typo_refuses_naming_the_paragraph(tpl):
    e = refused(build_typed(typo=True), tpl)
    assert e.code == "numbering-mismatch" and "‘B. Reliable.’ would read ‘C. Reliable.’" in e.reason


def test_a_skipped_level_refuses(tpl):
    """Only the articles selected: "1.1" names a PART level no paragraph reached — not evaluated, so refused."""
    e = refused(build_typed(), tpl, {"paragraphs": {"text_regex": "^[0-9]+\\.[0-9]+"}})
    assert e.code == "numbering-mismatch" and "skipped" in e.reason


def _scheme_doc_with(text: str, before_general: bool) -> bytes:
    doc = Document(io.BytesIO(build_scheme_example()))
    if before_general:
        doc.paragraphs[1].insert_paragraph_before(text)
    else:
        doc.add_paragraph(text)
    return docx_of(doc)


def test_a_list_that_already_numbers_paragraphs_refuses_when_they_would_move():
    """The example file as the document (its own list numbers its paragraphs; no template): a typed PART before its
    first PART would push every PART and article by one — refused; a typed "B." that continues the list converts."""
    e = refused(_scheme_doc_with("PART 1 - INTRODUCTION", True))
    assert e.code == "numbering-mismatch" and "‘PART 1 - GENERAL’ would read ‘PART 2 - GENERAL’" in e.reason
    src = _scheme_doc_with("B.\tSecond paragraph under part 2.", False)
    plan, out, rec = converting(src, tracked=False)
    assert rec.passed, rec.to_markdown()
    assert numbers(out)["Second paragraph under part 2."] == "B."
    assert len(plan.property_assertions) == 1                      # nothing imported


def _patched_template(tmp_path, fn) -> str:
    p = tmp_path / "SCHEME (patched).docx"
    p.write_bytes(patch_docx(build_scheme_example(), {NUMBERING: fn}))
    return str(p)


def _level(root, k: int):
    absn = next(a for a in root.findall(W + "abstractNum") if a.get(W + "abstractNumId") == "10")
    return next(l for l in absn.findall(W + "lvl") if l.get(W + "ilvl") == str(k))


def test_an_ambiguous_typed_number_refuses(tmp_path):
    """Level 4 in upper roman beside level 2 in upper letters: "C." fits both."""
    t = _patched_template(tmp_path, lambda r: _level(r, 4).find(W + "numFmt").set(W + "val", "upperRoman"))
    e = refused(build_typed(), t)
    assert e.code == "ambiguous" and "'C.'" in e.reason


def test_a_level_without_a_suffix_refuses(tmp_path):
    def nothing(root):
        lvl = _level(root, 2)
        lvl.insert(list(lvl).index(lvl.find(W + "lvlText")), _x('<w:suff w:val="nothing"/>'))
    e = refused(build_typed(), _patched_template(tmp_path, nothing))
    assert e.code == "numbering-mismatch" and "w:suff nothing" in e.reason


def test_nothing_typed_is_a_noted_no_op(tpl):
    doc = Document()
    doc.add_paragraph("Nothing typed here.")
    plan, out, rec = converting(docx_of(doc), tpl)
    assert rec.passed and any("no typed number" in n for n in plan.notes)
    assert refused(docx_of(doc), tpl, allow_missing=False).code == "not-found"


@pytest.mark.parametrize("select, params", [
    ({"scope": ["document"]}, {}), ({"include_tables": "yes"}, {}), ({}, {"tracked": "no"}),
    ({}, {"existing": "keep"}), ({}, {"levels": "PRT"}), ({}, {"colour": "red"})])
def test_parameters(tpl, select, params):
    assert refused(build_typed(), None if "existing" in params else tpl, select, **params).code == "bad-params"


# ---- it must READ as typed: hidden and capitalised typed numbers ------------------------------------------------------

STYLES = "word/styles.xml"


def _add_style(xml: str):
    def add(root):
        root.append(etree.fromstring(xml.replace("<w:style ", f"<w:style {W_DECL} ", 1)))
    return add


def _rpr(p_text: str, xml: str, run: int = 0):
    """Add run properties to run ``run`` of the paragraph ``p_text``."""
    def edit(root):
        r = find(root, p_text).findall(W + "r")[run]
        rpr = r.find(W + "rPr")
        if rpr is None:
            rpr = etree.Element(W + "rPr")
            r.insert(0, rpr)
        for el in etree.fromstring(f"<w:x {W_DECL}>{xml}</w:x>"):
            rpr.append(el)
    return edit


def _ppr(p_text: str, xml: str):
    def edit(root):
        p = find(root, p_text)
        ppr = p.find(W + "pPr")
        if ppr is None:
            ppr = etree.Element(W + "pPr")
            p.insert(0, ppr)
        for el in etree.fromstring(f"<w:x {W_DECL}>{xml}</w:x>"):
            ppr.append(el)
    return edit


HIDDEN_CHAR = ('<w:style w:type="character" w:styleId="HiddenChar"><w:name w:val="Hidden Char"/>'
               '<w:rPr><w:vanish/></w:rPr></w:style>')
SMALL_CAPS_PARA = ('<w:style w:type="paragraph" w:styleId="SmallCapsPara"><w:name w:val="Small Caps Para"/>'
                   '<w:basedOn w:val="Normal"/><w:rPr><w:smallCaps/></w:rPr></w:style>')
DISPLAY_CASES = {
    # (1) a hidden typed "A" (its own run, direct) and a "B.⇥…" hidden through a character style
    "hidden prefix": ({MAIN: _rpr("A.\tSection includes wet-pipe sprinkler systems.", "<w:vanish/>")}, "hidden text"),
    "hidden by its character style": ({STYLES: _add_style(HIDDEN_CHAR),
                                       MAIN: _rpr("B.\tRelated requirements:", '<w:rStyle w:val="HiddenChar"/>')},
                                      "hidden text"),
    # (2) a typed "C." on a paragraph whose mark is hidden
    "hidden mark": ({MAIN: _ppr("C.\tReliable.", "<w:rPr><w:vanish/></w:rPr>")}, "paragraph mark is hidden"),
    # (3) typed "a."/"b." the reader sees as "A."/"B.": direct caps, small caps through the paragraph style
    "caps": ({MAIN: _rpr("a.\tHangers.", "<w:caps/>")}, "capitals"),
    "small caps by style": ({STYLES: _add_style(SMALL_CAPS_PARA),
                             MAIN: _ppr("b.\tSupports.", '<w:pStyle w:val="SmallCapsPara"/>')}, "small capitals"),
    # (4) the review's F1: the list number takes the MARK's formatting — white, 1 pt, superscript, struck, symbol glyphs
    "white mark": ({MAIN: _ppr("A.\tTyco.", '<w:rPr><w:color w:val="FFFFFF"/></w:rPr>')}, "in another colour"),
    "1 pt mark": ({MAIN: _ppr("B.\tViking.", '<w:rPr><w:sz w:val="2"/></w:rPr>')}, "in another size"),
    "superscript mark": ({MAIN: _ppr("A.\tSteel pipe.", '<w:rPr><w:vertAlign w:val="superscript"/></w:rPr>')},
                         "raised or lowered"),
    "struck mark": ({MAIN: _ppr("A.\tNFPA 13.", "<w:rPr><w:strike/></w:rPr>")}, "struck through"),
    "Wingdings mark": ({MAIN: _ppr("C.\tReliable.", '<w:rPr><w:rFonts w:ascii="Wingdings" w:hAnsi="Wingdings"/>'
                                                    '</w:rPr>')}, "symbol font"),
    # the review's suspicion: w:specVanish hides like w:vanish
    "spec-vanished mark": ({MAIN: _ppr("A.\tInstall per NFPA 13.", "<w:rPr><w:specVanish/></w:rPr>")},
                           "paragraph mark is hidden"),
    "spec-vanished prefix": ({MAIN: _rpr("1.\tHangers per NFPA 13.", "<w:specVanish/>")}, "hidden text"),
}
F1_CASES = ["white mark", "1 pt mark", "superscript mark", "struck mark", "Wingdings mark", "spec-vanished mark",
            "spec-vanished prefix"]


def _display_case(name: str) -> bytes:
    edits, _ = DISPLAY_CASES[name]
    out = build_typed()
    for part, fn in edits.items():
        out = patch_docx(out, {part: fn})
    return out


@pytest.mark.parametrize("name", list(DISPLAY_CASES))
def test_a_typed_number_that_does_not_read_as_typed_refuses(tpl, name):
    """Run text ignores display properties; the reader does not: a hidden typed number would become a visible list
    number, a typed "a." in capitals reads "A." (and would become a lower-case number), a hidden mark would hide the
    number (Word) or join the paragraph to the next (LibreOffice). The file is refused, never left half typed."""
    e = refused(_display_case(name), tpl)
    assert e.code == "typed-format" and DISPLAY_CASES[name][1] in e.reason, e.reason
    assert "typed-format" in FAILURE_EXPLANATIONS


@pytest.mark.parametrize("name", ["hidden prefix", "hidden mark", "caps", "small caps by style"] + F1_CASES)
def test_a_transform_that_ignores_the_display_rule_is_caught(tpl, name, monkeypatch):
    """A plan and apply that convert those paragraphs anyway: the reading check of the transform (run text) passes,
    V1–V9 pass — the verifier's own display rule says the op had to refuse."""
    from spec_surgeon.ops import convert_typed_numbers as ctn
    monkeypatch.setattr(ctn, "display_problem", lambda *a: (None, []))
    plan, out, rec = converting(_display_case(name), tpl, tracked=False)
    assert "V10" in failed(rec) and "the op must refuse: typed-format" in v10(rec)


def test_the_mark_in_capitals_on_a_lower_case_level_refuses(tmp_path):
    """The example file's PR3 (level 4, "%5." lower letters) in capitals: Word formats the list number with the
    paragraph mark — "a." would show "A.". PRT is in capitals too, harmlessly: "PART %1 -" has no lower-case letter."""
    def caps(root):
        st = next(s for s in root.findall(W + "style") if s.get(W + "styleId") == "PR3")
        etree.SubElement(st, W + "rPr").append(etree.Element(W + "caps"))
    t = tmp_path / "SCHEME (caps).docx"
    t.write_bytes(patch_docx(build_scheme_example(), {STYLES: caps}))
    e = refused(build_typed(), str(t))
    assert e.code == "typed-format" and "PR3" in e.reason and "capitals" in e.reason


def test_a_scheme_level_whose_number_is_hidden_or_in_capitals_refuses(tmp_path, monkeypatch):
    """w:lvl/w:rPr turning on caps (or vanish, smallCaps): refused at scheme resolution — bad-params from the example
    file (both ops), scheme for the document's own list; a transform without the rule is caught by V10."""
    def lvl_caps(root):
        _level(root, 2).append(_x("<w:rPr><w:caps/></w:rPr>"))
    t = _patched_template(tmp_path, lvl_caps)
    assert refused(build_typed(), t).code == "bad-params"
    with pytest.raises(JobException) as e:
        get_op("convert_numbering").plan(DocxPackage(build_typed()), {}, {"levels": SCHEME_LEVELS, "template": t})
    assert e.value.code == "bad-params"
    held = patch_docx(_scheme_doc_with("B.\tSecond paragraph under part 2.", False), {NUMBERING: lambda r: _level(
        r, 2).append(_x("<w:rPr><w:vanish/></w:rPr>"))})
    assert refused(held).code == "scheme"
    from spec_surgeon.ops import numbering_convert
    from spec_surgeon.ops import convert_typed_numbers as ctn
    monkeypatch.setattr(numbering_convert, "level_display_problem", lambda *a: None)
    monkeypatch.setattr(ctn, "display_problem", lambda *a: (None, []))      # its number check sees the level's vanish too
    plan, out, rec = converting(held, tracked=False)
    assert "V10" in failed(rec) and "the op must refuse: scheme" in v10(rec)


# ---- wrong OUTPUTS are caught -----------------------------------------------------------------------------------------

def mutate(docx: bytes, part: str, fn) -> bytes:
    zin = zipfile.ZipFile(io.BytesIO(docx))
    root = etree.fromstring(zin.read(part))
    fn(root)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        for info in zin.infolist():
            z.writestr(info, serialize_xml(root) if info.filename == part else zin.read(info.filename))
    return buf.getvalue()


def _broken(tpl, fn, data=None, **params):
    data = data or build_typed()
    plan, out, rec = converting(data, tpl, **params)
    assert rec.passed, rec.to_markdown()
    rec2 = run_suite(data, mutate(out, MAIN, fn), plan.to_dict(), "src.docx", "out.docx")
    assert "V10" in failed(rec2)
    return rec2


def test_a_tab_left_behind_fails(tpl):
    def fn(root):
        p = find(root, "Section includes wet-pipe sprinkler systems.")
        p.find(W + "r").insert(0, _x("<w:tab/>"))
    rec = _broken(tpl, fn, tracked=False)
    # V3 cannot see it (the output holds as many tabs as the source); the item model does
    assert "typed_numbers" in v10(rec) and "expected text 'S', found run content tab" in v10(rec)


def test_the_old_number_left_in_the_text_fails(tpl):
    def fn(root):
        p = find(root, "Tyco.")
        p.insert(1, _x('<w:r><w:t xml:space="preserve">A. </w:t></w:r>'))
    rec = _broken(tpl, fn, tracked=False)
    assert "does not read like the source" in v10(rec)


def test_a_wrong_levels_style_fails(tpl):
    def fn(root):
        find(root, "Tyco.").find(f"{W}pPr/{W}pStyle").set(W + "val", "PR2")
    rec = _broken(tpl, fn)
    assert "typed_numbers" in v10(rec)


def test_an_own_numbering_left_in_place_fails(tpl):
    def off(root):
        find(root, "B.\tViking.").insert(0, _x('<w:pPr><w:numPr><w:ilvl w:val="0"/><w:numId w:val="0"/></w:numPr>'
                                                '</w:pPr>'))
    src = patch_docx(build_typed(), {MAIN: off})

    def fn(root):
        find(root, "Viking.").find(W + "pPr").insert(1, _x('<w:numPr><w:ilvl w:val="0"/><w:numId w:val="0"/></w:numPr>'))
    rec = _broken(tpl, fn, data=src, tracked=False)
    assert "typed_numbers" in v10(rec)


# ---- wrong IMPLEMENTATIONS shared by plan and apply are caught --------------------------------------------------------

def test_a_matcher_that_keeps_the_separator_is_caught(tpl, monkeypatch):
    """A typed-number matcher whose prefix stops at the number (the tab or spaces after it are left): plan and apply
    agree, the transform's own reading guarantee collapses the whitespace and passes, V2 and V3 pass — only the
    verifier's own prefixes see the separator left behind."""
    real = tl.typed_number

    def keep_separator(text, scheme):
        tn = real(text, scheme)
        return None if tn is None else dataclasses.replace(tn, end=tn.token_start + len(tn.token))
    monkeypatch.setattr(tl, "typed_number", keep_separator)
    plan, out, rec = converting(build_typed(), tpl, tracked=False)
    assert failed(rec) == ["V10"] and "typed numbers found here" in v10(rec)


def test_a_matcher_that_ignores_the_separator_rule_on_a_tracked_conversion_is_caught(tpl, monkeypatch):
    """The same matcher, tracked: the typed numbers become deletions without their tabs; V2/V8 agree with the plan's own
    accepted views, V10 does not."""
    real = tl.typed_number
    monkeypatch.setattr(tl, "typed_number", lambda text, scheme: (
        None if real(text, scheme) is None else
        dataclasses.replace(real(text, scheme), end=real(text, scheme).token_start + len(real(text, scheme).token))))
    plan, out, rec = converting(build_typed(), tpl)
    assert failed(rec) == ["V10"] and "typed_numbers" in v10(rec)


# ---- plan-time conflicts (rule 9) -------------------------------------------------------------------------------------

def _spec(*ops, mode="content"):
    return {"request": "r", "summary": "s", "mode": mode, "ops": list(ops)}


def test_conflicts(tpl):
    typed = {"op": "convert_typed_numbers", "params": {"levels": SCHEME_LEVELS, "template": tpl}}
    for other in ({"op": "replace_text", "params": {"find": "closing", "replace": "final"}},     # an unconverted line
                  {"op": "apply_style", "select": {"paragraphs": {"text_regex": "Tyco"}}, "params": {"style": "Normal"}},
                  {"op": "set_font", "params": {"font": "Arial", "only_where_defined": False}},
                  {"op": "convert_numbering", "params": {"levels": SCHEME_LEVELS, "template": tpl}},
                  # the display rule reads the typed runs' and marks' style chains: a style's caps or font beside it
                  {"op": "set_properties", "select": {"styles": {"ids": ["Normal"]}}, "params": {"run": {"caps": True}}},
                  {"op": "set_style_font", "select": {"styles": {"ids": ["Normal"]}},
                   "params": {"font": "Arial", "only_where_defined": False}}):
        for ops in ((typed, other), (other, typed)):
            with pytest.raises(JobException) as e:
                plan_spec(DocxPackage(build_typed()), _spec(*ops))
            assert e.value.code == "conflict", other["op"]


def test_touches(tpl):
    frag = get_op("convert_typed_numbers").plan(DocxPackage(build_typed()), {},
                                                {"levels": SCHEME_LEVELS, "template": tpl})
    t = set(frag.touches)
    assert {(MAIN, "*", "verify:text"), (MAIN, "scope", "verify"), (MAIN, "*", "verify-props:pStyle"),
            (MAIN, "*", "verify-props:numPr"), (NUMBERING, "*", "verify"), (NUMBERING, "*", "add")} <= t
    edited = [b for m, b, a in t if m == MAIN and a == "text"]
    assert len(edited) == 22 and all((MAIN, b, "pStyle") in t for b in edited)


# ---- the template is an input (rule 12) -------------------------------------------------------------------------------

def test_a_template_changed_after_planning(tpl):
    src = build_typed()
    pkg = DocxPackage(src)
    frag = get_op("convert_typed_numbers").plan(pkg, {}, {"levels": SCHEME_LEVELS, "template": tpl, "tracked": False})
    plan = merge_fragments([frag])
    work = pkg.copy()
    get_op("convert_typed_numbers").apply(work, frag)
    out = work.output_bytes()
    Path(tpl).write_bytes(patch_docx(build_scheme_example(), {NUMBERING: lambda r: _level(r, 3).find(
        W + "lvlText").set(W + "val", "(%4)")}))
    with pytest.raises(JobException):
        get_op("convert_typed_numbers").apply(pkg.copy(), frag)
    rec = run_suite(src, out, plan.to_dict(), "src.docx", "out.docx")
    assert not rec.passed and "template" in v10(rec)


# ---- the worked example, the catalog --------------------------------------------------------------------------------

def test_the_worked_example_plans_and_verifies():
    spec = load_spec(ROOT / "examples" / "specs" / "convert-typed-numbers.json")
    validate_spec(spec)
    src = build_typed()
    plan, frags = plan_spec(DocxPackage(src), spec)
    out = apply_spec(DocxPackage(src), frags).output_bytes()
    rec = run_suite(src, out, plan.to_dict(), "src.docx", "out.docx")
    assert rec.passed, rec.to_markdown()
    assert numbers(out)["SUMMARY"] == "1.1"


def test_catalog_and_explanations():
    assert "convert_typed_numbers" in op_catalog()["ops"] and COMMON_ASSUMPTIONS["convert_typed_numbers"]
    for code in ("numbering-mismatch", "ambiguous", "numbering-unsupported", "scheme"):
        assert code in FAILURE_EXPLANATIONS, code


def test_every_refusal_code_of_the_numbering_conversions_is_explained_and_none_file_specific_is_fatal():
    """The code of every JobException the three modules raise (and of every lists.resolve_scheme call, which raises the
    code its caller passes) has a plain-words explanation; a refusal one file earns never rejects the whole spec at
    compile time (batch 2 rule 5) — only bad-params (the parameters, the example file) is fatal."""
    import ast
    from spec_surgeon import compiler
    codes, computed = set(), []
    for module in ("numbering_convert", "convert_numbering", "convert_typed_numbers"):
        tree = ast.parse((ROOT / "src" / "spec_surgeon" / "ops" / f"{module}.py").read_text(encoding="utf-8"))
        for n in ast.walk(tree):
            if not isinstance(n, ast.Call):
                continue
            name = getattr(n.func, "id", None) or getattr(n.func, "attr", None)
            at = {"JobException": 1, "resolve_scheme": 4}.get(name)
            if at is None:
                continue
            code = n.args[at] if len(n.args) > at else None
            if isinstance(code, ast.Constant):
                codes.add(code.value)
            else:
                computed.append((module, n.lineno))
    assert computed == [] and {"scheme", "scheme-depth", "numbering-mismatch", "numbering-unsupported", "typed-format",
                               "ambiguous", "bad-params"} <= codes
    codes.discard("internal")
    assert sorted(c for c in codes if c not in FAILURE_EXPLANATIONS) == []
    assert (codes & compiler.FATAL_SPEC_CODES) == {"bad-params"}


def test_the_numbering_examples_place_no_words():
    from spec_surgeon import compiler
    for name in ("convert-typed-numbers", "convert-numbering-to-scheme"):
        assert compiler.placed_words(load_spec(ROOT / "examples" / "specs" / f"{name}.json")) == [], name


# ---- the adversarial review's fixes (F1, F3, F4, F6, F9–F14) ------------------------------------------------------------

def _text(old: str, new: str):
    """Replace the w:t text ``old`` (one w:t) by ``new``."""
    def edit(root):
        t = next(t for t in root.iter(W + "t") if t.text == old)
        t.text = new
    return edit


def test_a_bold_mark_is_flagged_not_refused(tpl):
    """F1: bold, italic, underline and a (non-symbol) font of the list number differing from the typed number's are
    flagged; the mark's bold makes 'B.' bold where the typed 'B.' was not."""
    plan, out, rec = converting(patch_docx(build_typed(), {MAIN: _ppr("b.\tSupports.", "<w:rPr><w:b/></w:rPr>")}),
                                tpl, tracked=False)
    assert rec.passed, rec.to_markdown()
    assert any("differ in bold" in f and "‘b. Supports.’" in f for f in plan.flags), plan.flags


def test_a_scheme_style_that_sizes_text_and_number_alike_converts(tmp_path):
    """F1 compares the typed runs and the number over the SAME new style chain: an example file whose PR1 is 9 pt makes
    the text and the number 9 pt alike — no refusal, the document reads as before."""
    def small(root):
        st = next(s for s in root.findall(W + "style") if s.get(W + "styleId") == "PR1")
        etree.SubElement(st, W + "rPr").append(_x('<w:sz w:val="18"/>'))
    t = tmp_path / "SCHEME (9 pt).docx"
    t.write_bytes(patch_docx(build_scheme_example(), {STYLES: small}))
    plan, out, rec = converting(build_typed(), str(t), tracked=False)
    assert rec.passed, rec.to_markdown()


@pytest.mark.parametrize("edit", [_rpr("PART 1 - GENERAL", "<w:b/><w:caps/>"), _rpr("1.1\tSUMMARY", "<w:caps/>"),
                                  _rpr("A.\tTyco.", "<w:caps/>")], ids=["PART", "article", "upper letter"])
def test_capitals_on_a_typed_number_without_lower_case_letters_convert(tpl, edit):
    """F4: "PART 1 -", "1.1" and "A." read the same in capitals — converted and verified (V10's own rule agrees); a
    typed "a." in capitals still refuses (DISPLAY_CASES "caps")."""
    plan, out, rec = converting(patch_docx(build_typed(), {MAIN: edit}), tpl, tracked=False)
    assert rec.passed, rec.to_markdown()
    assert numbers(out)["Tyco."] == "A."


def test_the_new_styles_capitals_are_flagged(tpl):
    """F9: the example file's PRT shows its text in capitals — a typed "PART 1 - General" reads "GENERAL" afterwards."""
    plan, out, rec = converting(patch_docx(build_typed(), {MAIN: _text("PART 1 - GENERAL", "PART 1 - General")}), tpl,
                                tracked=False)
    assert rec.passed, rec.to_markdown()
    assert any("capitals" in f and "‘General’" in f for f in plan.flags), plan.flags


def _linked(root):
    for st in root.findall(W + "style"):
        if st.get(W + "styleId") in ("PR1", "PR2", "PR3"):
            numpr = st.find(f"{W}pPr/{W}numPr")
            numpr.remove(numpr.find(W + "ilvl"))


def test_a_scheme_whose_styles_do_not_name_their_level_is_flagged(tpl, tmp_path):
    """F3: PR1–PR3 name their list without w:ilvl (the level linked through w:lvl/w:pStyle): Word numbers them right,
    LibreOffice at level 0 — flagged, in both ops; the example file as built names every level: no flag."""
    t = tmp_path / "LINKED.docx"
    t.write_bytes(patch_docx(build_scheme_example(), {STYLES: _linked}))
    plan, out, rec = converting(build_typed(), str(t), tracked=False)
    assert rec.passed, rec.to_markdown()
    want = "your scheme's style(s) ['PR1', 'PR2', 'PR3'] do not name their level"
    assert any(want in f and "LibreOffice" in f for f in plan.flags), plan.flags
    from make_fixtures import build_numbering
    frag = get_op("convert_numbering").plan(DocxPackage(build_numbering(False)), {},
                                            {"levels": SCHEME_LEVELS, "template": str(t)})
    assert any(want in f for f in frag.flags)
    frag = get_op("convert_typed_numbers").plan(DocxPackage(build_typed()), {}, {"levels": SCHEME_LEVELS,
                                                                                 "template": tpl})
    assert not any("do not name their level" in f for f in frag.flags)


def _zero_articles(root):
    _level(root, 1).find(W + "numFmt").set(W + "val", "decimalZero")


def test_held_styles_that_number_unlike_the_example_file_refuse_with_keep(tpl, monkeypatch):
    """F6: the example file's own style ids with this file's own list ("1.01" articles): with the default existing
    "keep" the scheme would silently not apply — refused (scheme), naming ART and pointing at "replace", in both ops
    (test_convert_numbering converts it with "replace"). A transform without the rule is caught by V10."""
    arch = patch_docx(build_scheme_example(), {NUMBERING: _zero_articles})
    for op in ("convert_typed_numbers", "convert_numbering"):
        with pytest.raises(JobException) as e:
            get_op(op).plan(DocxPackage(arch), {}, {"levels": SCHEME_LEVELS, "template": tpl})
        assert e.value.code == "scheme" and "ART" in e.value.reason and "'1.01'" in e.value.reason \
            and "existing 'replace'" in e.value.reason, op
    src = patch_docx(_scheme_doc_with("B.\tSecond paragraph under part 2.", False), {NUMBERING: _zero_articles})
    assert refused(src, tpl).code == "scheme"
    from spec_surgeon.ops import numbering_convert
    monkeypatch.setattr(numbering_convert, "level_pattern_sig", lambda *a: 0)
    plan, out, rec = converting(src, tpl, tracked=False)
    assert "V10" in failed(rec) and "the op must refuse: scheme" in v10(rec)


def test_an_unmatched_part_line_is_named_in_the_reason(tpl):
    """F11: "PART 1 – GENERAL" (an en dash) matches no level: the article after it names a level no paragraph reached —
    the reason names the PART line, what it was typed as and what the level needs."""
    e = refused(patch_docx(build_typed(), {MAIN: _text("PART 1 - GENERAL", "PART 1 – GENERAL")}), tpl)
    assert e.code == "numbering-mismatch"
    assert "‘PART 1 – GENERAL’ ('PART 1 –')" in e.reason and "level 0 numbers like 'PART 1 -'" in e.reason, e.reason
    assert "names a level no paragraph reached" in e.reason


def test_a_positional_tab_is_named(tpl):
    """F11: a w:ptab after the typed number is called a positional tab, not a tab."""
    def ptab(root):
        tab = find(root, "A.\tTyco.").find(f"{W}r/{W}tab")
        tab.getparent().replace(tab, _x('<w:ptab w:relativeTo="margin" w:alignment="left" w:leader="none"/>'))
    e = refused(patch_docx(build_typed(), {MAIN: ptab}), tpl)
    assert e.code == "structure" and "a positional tab (w:ptab)" in e.reason, e.reason


REL_NUMBERING = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/numbering"


def _without_numbering(data: bytes) -> bytes:
    zin = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for info in zin.infolist():
            if info.filename == NUMBERING:
                continue
            payload = zin.read(info.filename)
            if info.filename in ("word/_rels/document.xml.rels", "[Content_Types].xml"):
                root = etree.fromstring(payload)
                for el in list(root):
                    if el.get("Type") == REL_NUMBERING or el.get("PartName") == "/" + NUMBERING:
                        root.remove(el)
                payload = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            z.writestr(info, payload)
    return out.getvalue()


def test_a_file_without_list_definitions_is_refused_plainly(tpl):
    """F12 (kept): Word writes word/numbering.xml once a list is used; the engine never creates a part (rule 10) — the
    reason says so in plain words."""
    e = refused(_without_numbering(build_typed()), tpl)
    assert e.code == "style-numbering" and "no list definitions part" in e.reason and "rule 10" in e.reason


def test_a_roman_and_a_letter_level_refuse_with_a_clear_reason(tmp_path):
    """F13 (kept): the classic outline (I. / … / A.) — "II." fits the roman level 0 and the letter level 2."""
    def outline(root):
        lv = _level(root, 0)
        lv.find(W + "numFmt").set(W + "val", "upperRoman")
        lv.find(W + "lvlText").set(W + "val", "%1.")
    e = refused(patch_docx(build_typed(), {MAIN: _text("PART 1 - GENERAL", "II. GENERAL")}),
                _patched_template(tmp_path, outline))
    assert e.code == "ambiguous" and "'II.'" in e.reason and "level 0 (numbered like 'I.')" in e.reason \
        and "a roman numeral is also a run of letters" in e.reason, e.reason


def _scheme_styles_without_numbering(root):
    for sid in SCHEME_LEVELS:
        root.append(etree.fromstring(f'<w:style {W_DECL} w:type="paragraph" w:styleId="{sid}"><w:name w:val="{sid}"/>'
                                     f'<w:basedOn w:val="Normal"/></w:style>'))


def _styled(text: str, sid: str):
    return _ppr(text, f'<w:pStyle w:val="{sid}"/>')


def test_a_tracked_conversion_over_replaced_styles_refuses(tpl, monkeypatch):
    """F14: typed paragraphs that already carry the scheme's style ids (unnumbered here) with existing "replace": the
    definitions are replaced untracked — reject-all would bring the typed numbers back under numbered styles ("A. A.").
    Tracked is refused (tracked-format); plain converts; a transform without the rule is caught by V10."""
    src = patch_docx(build_typed(), {STYLES: _scheme_styles_without_numbering,
                                     MAIN: lambda r: (_styled("PART 1 - GENERAL", "PRT")(r),
                                                      _styled("1.1\tSUMMARY", "ART")(r))})
    e = refused(src, tpl, existing="replace")
    assert e.code == "tracked-format" and "PRT" in e.reason and "reject-all" in e.reason
    plan, out, rec = converting(src, tpl, existing="replace", tracked=False)
    assert rec.passed, rec.to_markdown()
    from spec_surgeon.ops import convert_typed_numbers as ctn
    monkeypatch.setattr(ctn, "replaced_in_chain", lambda *a: [])
    plan, out, rec = converting(src, tpl, existing="replace")
    assert "V10" in failed(rec) and "the op must refuse: tracked-format" in v10(rec)


def test_touches_are_declared_once(tpl):
    """F10: one touch per target and aspect — merge_fragments' pairwise owner loop stays linear."""
    frag = get_op("convert_typed_numbers").plan(DocxPackage(build_typed()), {},
                                                {"levels": SCHEME_LEVELS, "template": tpl})
    assert len(frag.touches) == len(set(frag.touches))


# ---- the LibreOffice proof ---------------------------------------------------------------------------------------------

def _assert_rendered_words(source: bytes, rendered: str) -> None:
    """The fixture's literal typed numbers are the oracle, independent of either numbering model.

    PDF line and page breaks vary with the renderer and installed fonts. Compare every visible token instead:
    a missing, duplicated, changed or reordered number or word still fails, including table and unnumbered text.
    """
    with zipfile.ZipFile(io.BytesIO(source)) as archive:
        root = etree.fromstring(archive.read(MAIN))
    expected = [word for paragraph in root.iter(W + "p") for word in live(paragraph).split()]
    assert rendered.split() == expected


def test_rendered_numbering_oracle_accepts_only_layout_whitespace():
    doc = Document()
    doc.add_paragraph("1. Alpha beta gamma.")
    doc.add_paragraph("2. Delta epsilon.")
    _assert_rendered_words(docx_of(doc), "1.\tAlpha beta\n   gamma.\n\f2. Delta\n epsilon.\n")


@pytest.mark.parametrize("rendered", [
    "Alpha beta gamma. 2. Delta epsilon.",                   # missing number
    "1. 1. Alpha beta gamma. 2. Delta epsilon.",             # duplicate number
    "3. Alpha beta gamma. 2. Delta epsilon.",                # changed number
    "1. Alpha gamma. 2. Delta epsilon.",                    # missing word
    "1. Alpha beta beta gamma. 2. Delta epsilon.",           # duplicate word
    "1. Alpha beta gamma 2. Delta epsilon.",                # changed punctuation
    "2. Delta epsilon. 1. Alpha beta gamma.",                # reordered paragraphs
])
def test_rendered_numbering_oracle_rejects_changed_content(rendered):
    doc = Document()
    doc.add_paragraph("1. Alpha beta gamma.")
    doc.add_paragraph("2. Delta epsilon.")
    with pytest.raises(AssertionError):
        _assert_rendered_words(docx_of(doc), rendered)


@pytest.mark.skipif(not (shutil.which("soffice") and shutil.which("pdftotext")), reason="LibreOffice/poppler missing")
@pytest.mark.parametrize("tracked", [False, True])
@pytest.mark.parametrize("narrow_page", [False, True], ids=["default-width", "wrapped"])
def test_libreoffice_shows_every_typed_number_exactly_once_in_order(tpl, tmp_path, tracked, narrow_page):
    """The converted file rendered by LibreOffice (whose numbering follows the example file's explicit w:ilvl levels,
    ooxml-lessons §6): every converted paragraph shows its number once — never twice, never missing — and all visible
    words and numbers remain in the document's order, even when a paragraph wraps across physical PDF lines."""
    source = build_typed()
    if narrow_page:
        def narrow(root):
            for size in root.iter(W + "pgSz"):
                size.set(W + "w", "10080")  # 7-inch page; the long indented paragraph must wrap.
        source = patch_docx(source, {MAIN: narrow})
    plan, out, rec = converting(source, tpl, tracked=tracked)
    assert rec.passed
    if tracked:                                         # LibreOffice shows deletions struck through: accept them first
        root = DocxPackage(out).xml(MAIN)
        for d in list(root.iter(W + "del")):
            d.getparent().remove(d)
        for c in list(root.iter(W + "pPrChange")):
            c.getparent().remove(c)
        pkg = DocxPackage(out)
        pkg.write_xml(MAIN, root)
        out = pkg.output_bytes()
    doc = tmp_path / "out.docx"
    doc.write_bytes(out)
    # Keep the executable's discovered PATH (including non-system installations) and isolate
    # LibreOffice's profile without replacing the user's HOME or the runtime environment.
    profile = (tmp_path / "libreoffice-profile").as_uri()
    subprocess.run([shutil.which("soffice"), f"-env:UserInstallation={profile}", "--headless", "--convert-to", "pdf",
                    "--outdir", str(tmp_path), str(doc)], check=True, capture_output=True, timeout=240)
    subprocess.run(["pdftotext", "-layout", str(tmp_path / "out.pdf"), str(tmp_path / "out.txt")], check=True,
                   timeout=60)
    rendered = (tmp_path / "out.txt").read_text()
    if narrow_page:
        longest = 'Division 21 Section "Common Work Results for Fire Suppression".'
        assert not any(longest in line for line in rendered.splitlines()), "fixture did not exercise line wrapping"
    _assert_rendered_words(source, rendered)
