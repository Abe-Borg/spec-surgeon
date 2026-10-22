"""Document properties built from each file's own section (docs/PLAN-2026-10-22-batch-consistency.md C5).

From the owner's Division 21 run: File Properties titles from other projects ("SECTION PP-16622 - REMOTE POWER
PANELS", "SECTION 211313 - WET-PIPE SPRINKLER SYSTEMS") sat in the batch, and set_document_properties took typed words
only, so fixing them meant typing every title. ``from_document`` puts each file's own section number and title (the
reference fill rules) where {number} and {title} stand, between the user's own words. V10 reads them again with the
verifier's own rules."""
from __future__ import annotations

import io

import pytest
from docx import Document

from make_fixtures import build
from spec_surgeon import compiler
from spec_surgeon.editspec import apply_spec, plan_spec
from spec_surgeon.ops import JobException
from spec_surgeon.ops import docprops as tdocprops
from spec_surgeon.ops.slots import REFERENCE_FILLS
from spec_surgeon.package import DocxPackage
from spec_surgeon.planner import plan_batch
from spec_surgeon.report import FAILURE_EXPLANATIONS
from spec_surgeon.verify import docprops as vdocprops
from spec_surgeon.verify import run_suite

TITLE = "Identification for Fire Suppression Systems"
WORDS = "SECTION {number} - {title}"


def spec_file(number="21 05 53", title=TITLE, properties_title="SECTION PP-16622 - REMOTE POWER PANELS") -> bytes:
    doc = Document(io.BytesIO(build(number, title)))
    doc.core_properties.title = properties_title
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def spec_of(**params) -> dict:
    return {"request": f"Set every file's title to {WORDS}.", "summary": "titles", "mode": "untracked",
            "ops": [{"op": "set_document_properties", "params": params}]}


def title_of(docx: bytes) -> str:
    return Document(io.BytesIO(docx)).core_properties.title


def run(src: bytes, spec: dict):
    plan, frags = plan_spec(DocxPackage(src), spec)
    out = apply_spec(DocxPackage(src), frags).output_bytes()
    return plan, out, run_suite(src, out, plan.to_dict(), "src.docx", "out.docx")


def failed(rec):
    return [c.id for c in rec.checks if c.status == "FAIL"]


def test_the_stale_title_becomes_the_files_own():
    src = spec_file()
    plan, out, rec = run(src, spec_of(from_document={"title": WORDS}))
    assert rec.passed, failed(rec)
    assert title_of(out) == f"SECTION 21 05 53 - {TITLE}"
    assert any("the file's own section number and title in place" in n for n in plan.notes)
    # already as asked: nothing changes, and the record still passes
    again, out2, rec = run(out, spec_of(from_document={"title": WORDS}))
    assert rec.passed and title_of(out2) == title_of(out)
    # one placeholder alone, beside a typed property
    plan, out, rec = run(src, spec_of(from_document={"subject": "{title}"}, set={"company": "Acme Engineering"}))
    assert rec.passed and Document(io.BytesIO(out)).core_properties.subject == TITLE


def test_a_batch_gets_each_files_own_and_sets_aside_the_file_without(tmp_path):
    files = []
    for n, title in (("21 05 17", "SLEEVES AND SLEEVE SEALS"), ("21 12 00", "FIRE-SUPPRESSION STANDPIPES")):
        path = tmp_path / f"{n}.docx"
        path.write_bytes(spec_file(n, title))
        files.append(path)
    loose = tmp_path / "loose.docx"
    doc = Document()
    doc.add_paragraph("No section line here.")
    doc.save(loose)
    bp = plan_batch([str(f) for f in files + [loose]], spec_of(from_document={"title": WORDS}))
    status = {f.source.name: f.status for f in bp.files}
    assert status == {"21 05 17.docx": "planned", "21 12 00.docx": "planned", "loose.docx": "exception"}
    loose_plan = next(f for f in bp.files if f.source.name == "loose.docx")
    assert loose_plan.exception["code"] == "not-found"
    assert "section number is not found by the reference rules" in loose_plan.exception["reason"]
    for f in files:
        _, out, rec = run(f.read_bytes(), spec_of(from_document={"title": WORDS}))
        assert rec.passed
    assert title_of(run(files[1].read_bytes(), spec_of(from_document={"title": WORDS}))[1]) == \
        "SECTION 21 12 00 - FIRE-SUPPRESSION STANDPIPES"


@pytest.mark.parametrize("params, code, words", [
    ({"from_document": {"title": "SECTION 21 05 53"}}, "bad-params", "{number} and/or {title}"),
    ({"from_document": {"title": "{number} {project}"}}, "bad-params", "no other"),
    ({"from_document": {"title": "{number} - {number}"}}, "bad-params", "at most once"),
    ({"from_document": {"title": WORDS}, "set": {"title": "Typed"}}, "bad-params", "asked for twice"),
    ({"from_document": {"title": WORDS}, "clear": ["title"]}, "bad-params", "asked for twice"),
    ({"from_document": {"title": "\u200b{number}"}}, "bad-params", "invisible"),
    ({"from_document": {"nickname": WORDS}}, "spec", "nickname"),          # the schema refuses it first
])
def test_the_words_are_checked_before_any_file(params, code, words):
    with pytest.raises(JobException) as e:
        plan_spec(DocxPackage(spec_file()), spec_of(**params))
    assert e.value.code == code and words in e.value.reason


def test_a_built_value_that_is_no_property_value_sets_the_file_aside():
    src = spec_file(title="LONG " * 60)
    with pytest.raises(JobException) as e:
        plan_spec(DocxPackage(src), spec_of(from_document={"title": WORDS}))
    assert e.value.code == "property-value" and "255" in e.value.reason
    assert "property-value" in FAILURE_EXPLANATIONS


def test_the_words_around_the_placeholders_are_the_users():
    spec = spec_of(from_document={"title": WORDS})
    assert compiler.placed_words(spec) == ["SECTION", "-"]
    assert compiler.untraceable(spec, [spec["request"]]) == []
    stamped = spec_of(from_document={"title": "ISSUED FOR PERMIT {number}"})
    assert compiler.untraceable(stamped, ["Set the titles to their own numbers."]) == ["ISSUED FOR PERMIT"]


def test_planted_defects_are_caught(monkeypatch):
    src = spec_file()
    plan, frags = plan_spec(DocxPackage(src), spec_of(from_document={"title": WORDS}))
    # a wrong output: the title written without its section title
    work = apply_spec(DocxPackage(src), frags)
    doc = Document(io.BytesIO(work.output_bytes()))
    doc.core_properties.title = "SECTION 21 05 53 - "
    buf = io.BytesIO()
    doc.save(buf)
    rec = run_suite(src, buf.getvalue(), plan.to_dict(), "src.docx", "out.docx")
    assert "V10" in failed(rec)
    # a wrong implementation shared by plan and apply: the title read from the wrong paragraph
    real = tdocprops.extract_fill
    monkeypatch.setattr(tdocprops, "extract_fill", lambda pkg, name, rule: "PROJECT REVISION HISTORY"
                        if name == "SECTION_TITLE" else real(pkg, name, rule))
    plan, out, rec = run(src, spec_of(from_document={"title": WORDS}))
    assert title_of(out) == "SECTION 21 05 53 - PROJECT REVISION HISTORY"     # the defect did what it does
    assert "V10" in failed(rec)


def test_the_verifier_keeps_its_own_copy_of_the_reference_rules():
    assert vdocprops._V_REFERENCE == {"number": REFERENCE_FILLS["SECTION_NUMBER"],
                                      "title": REFERENCE_FILLS["SECTION_TITLE"]}


def test_the_title_it_reads_is_a_read_another_op_cannot_change_unseen():
    spec = spec_of(from_document={"title": WORDS})
    spec["ops"].insert(0, {"op": "replace_text", "select": {"scope": ["document"]},
                           "params": {"find": TITLE, "replace": "Identification for Fire-Suppression Piping",
                                      "tracked": False}})
    with pytest.raises(JobException) as e:
        plan_spec(DocxPackage(spec_file()), spec)
    assert e.value.code == "conflict"


def test_the_prompt_and_the_catalog_name_it():
    assert 'from_document {"title": "SECTION {number} - {title}"}' in compiler.system_prompt(chat=True)
    from spec_surgeon.catalog import COMMON_ASSUMPTIONS, op_sheet
    assert any("from_document" in a for a in COMMON_ASSUMPTIONS["set_document_properties"])
    assert "from_document" in op_sheet()["ops"]["set_document_properties"]["params"]["properties"]
