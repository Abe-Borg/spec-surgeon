"""Build the bundled numbering schemes (docs/SPEC_FORMATTER_MODES_DESIGN.md N5) under
``src/spec_surgeon/examples/schemes/``: one paragraph style per level, numbered by ONE list, and a paragraph of each.

Built with python-docx (a dev dependency the app does not ship) and committed. Every ``.docx`` is written with fixed zip
timestamps, so running this again on the same python-docx gives the same bytes; ``tests/test_schemes.py`` compares the
committed files' document, styles and numbering parts with a fresh build.

    python tests/make_scheme_examples.py              # rewrites src/spec_surgeon/examples/schemes/

* **CSC PageFormat** — the Specification Formatter's built-in list (its ``builtin_scheme.py``): ``PART %1``,
  ``%1.%2``, ``.%3`` … ``.%9``, every level decimal, starting at 1, no restart; PART flush left, every level below it
  one half inch further in with its number hanging; spacing 12 pt before and 6 pt after PART and the article, 6 pt
  after the rest. No character formatting: a conversion keeping each paragraph's own style keeps its look.
* **CSI SectionFormat** — the same layout with ``PART %1``, ``%1.%2``, ``%3.`` (A), ``%4.`` (1), ``%5.`` (a),
  ``%6)`` (1), ``%7)`` (a), ``(%8)`` (1), ``(%9)`` (a).
"""
from __future__ import annotations

import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(HERE.parent / "src"))

from docx import Document  # noqa: E402

from make_fixtures import _abstract_xml, _add_lists, _add_style, _style_name, docx_bytes  # noqa: E402
from make_tutorial_showcase import stable_zip  # noqa: E402
from spec_surgeon.schemes import SCHEMES, SCHEMES_DIR  # noqa: E402

LEVEL_TEXT = {
    "csc": [("decimal", "PART %1"), ("decimal", "%1.%2")] + [("decimal", f".%{k}") for k in range(3, 10)],
    "csi": [("decimal", "PART %1"), ("decimal", "%1.%2"), ("upperLetter", "%3."), ("decimal", "%4."),
            ("lowerLetter", "%5."), ("decimal", "%6)"), ("lowerLetter", "%7)"), ("decimal", "(%8)"),
            ("lowerLetter", "(%9)")],
}
NAMES = {"csc": ["PageFormat Part", "PageFormat Article"] + [f"PageFormat Level {k}" for k in range(3, 10)],
         "csi": ["CSI Part", "CSI Article", "CSI Paragraph", "CSI Subparagraph"] + [f"CSI Level {k}" for k in range(5, 10)]}
SAMPLES = ["GENERAL", "SUMMARY", "A paragraph of the article.", "A subparagraph.", "The fifth level.",
           "The sixth level.", "The seventh level.", "The eighth level.", "The ninth level."]
NUM_ID, ABSTRACT_ID = 40, 40
STEP = 720


def _lvl(ilvl: int, fmt: str, text: str, pstyle: str) -> str:
    left, hanging = (0, 0) if ilvl == 0 else (STEP * ilvl, STEP)
    return (f'<w:lvl w:ilvl="{ilvl}"><w:start w:val="1"/><w:numFmt w:val="{fmt}"/><w:pStyle w:val="{pstyle}"/>'
            f'<w:lvlText w:val="{text}"/><w:lvlJc w:val="left"/>'
            f'<w:pPr><w:ind w:left="{left}" w:hanging="{hanging}"/></w:pPr></w:lvl>')


def build_scheme(scheme: str) -> bytes:
    _, title, ids, _ = SCHEMES[scheme]
    doc = Document()
    for k, (sid, name) in enumerate(zip(ids, NAMES[scheme])):
        spacing = '<w:spacing w:before="240" w:after="120"/>' if k < 2 else '<w:spacing w:before="0" w:after="120"/>'
        nxt = f'<w:next w:val="{ids[k + 1]}"/>' if k + 1 < len(ids) else ""
        _add_style(doc, f'<w:style w:type="paragraph" w:customStyle="1" w:styleId="{sid}"><w:name w:val="{name}"/>'
                        f'<w:basedOn w:val="Normal"/>{nxt}<w:qFormat/><w:pPr><w:numPr><w:ilvl w:val="{k}"/>'
                        f'<w:numId w:val="{NUM_ID}"/></w:numPr>{spacing}</w:pPr></w:style>')
    levels = [_lvl(k, fmt, text, sid) for k, ((fmt, text), sid) in enumerate(zip(LEVEL_TEXT[scheme], ids))]
    _add_lists(doc, [_abstract_xml(ABSTRACT_ID, levels, f"6B6B00{NUM_ID:02d}")],
               [f'<w:num w:numId="{NUM_ID}"><w:abstractNumId w:val="{ABSTRACT_ID}"/></w:num>'])
    doc.add_paragraph(f"{title.upper()} NUMBERING SCHEME (BUILT IN)", style="Title")
    for sid, text in zip(ids, SAMPLES):
        doc.add_paragraph(text, style=doc.styles[_style_name(doc, sid)])
    return stable_zip(docx_bytes(doc))


def schemes() -> dict:
    return {SCHEMES[s][0]: build_scheme(s) for s in SCHEMES}


def main(argv=None) -> int:
    SCHEMES_DIR.mkdir(parents=True, exist_ok=True)
    for name, data in schemes().items():
        (SCHEMES_DIR / name).write_bytes(data)
        print(f"wrote {SCHEMES_DIR / name} ({len(data)} bytes)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
