"""The seed new revision ids count up from (``ooxml.max_revision_id``): every annotation id the file already carries
counts, whichever document part carries it. Word writes a w:tblGridChange with its own w:id when someone resizes table
columns with Track Changes on (common in redlined specifications); a seed read from a short list of revision tags left
it (and w:tblPrExChange) out, handed the same id out again, and V4 set the file aside for a collision the engine made.
The list is now every element the schemas give an annotation id, checked against the vendored XSDs in
tests/test_package_hardening.py."""
from __future__ import annotations

import io
from typing import List

import pytest
from docx import Document
from lxml import etree

from spec_surgeon.editspec import apply_spec, plan_spec
from spec_surgeon.ooxml import max_revision_id
from spec_surgeon.ops import replace_text as replace_module
from spec_surgeon.package import DocxPackage, W
from spec_surgeon.verify import run_suite
from test_text_edits import accepted_text, mutate

AUTHOR = "Abraham Borg"
STYLES = "word/styles.xml"
SENTENCE = "The Owner shall review."
EDITS = 4               # four occurrences: a w:del and a w:ins each, eight new ids; from a seed of 0 they are 1..8


def resized_table(change: str = "tblGridChange", rev_id: str = "7") -> bytes:
    """One table carrying a tracked change with ``rev_id`` (a column resize, or a row's table-property exception),
    then ``EDITS`` paragraphs reading ``SENTENCE``."""
    d = Document()
    tbl = d.add_table(rows=1, cols=2)
    tbl.cell(0, 0).text, tbl.cell(0, 1).text = "Item", "Value"
    for _ in range(EDITS):
        d.add_paragraph(SENTENCE)
    if change == "tblGridChange":
        grid = tbl._tbl.find(W + "tblGrid")
        old = etree.SubElement(etree.SubElement(grid, W + "tblGridChange", {W + "id": rev_id}), W + "tblGrid")
        for _ in range(2):
            etree.SubElement(old, W + "gridCol", {W + "w": "4000"})
    else:
        tr = tbl._tbl.find(W + "tr")
        ex = etree.Element(W + "tblPrEx")
        etree.SubElement(ex, W + "jc", {W + "val": "center"})
        chg = etree.SubElement(ex, W + "tblPrExChange",
                               {W + "id": rev_id, W + "author": "Reviewer", W + "date": "2026-10-01T09:00:00Z"})
        etree.SubElement(chg, W + "tblPrEx")
        tr.insert(0, ex)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def spec() -> dict:
    return {"request": "replace Owner with Client", "summary": "s", "mode": "content",
            "ops": [{"op": "replace_text", "params": {"find": "Owner", "replace": "Client", "author": AUTHOR}}]}


def run(src: bytes):
    pkg = DocxPackage(src)
    plan, frags = plan_spec(pkg, spec())
    out = apply_spec(pkg, frags).output_bytes()
    return plan, out, run_suite(src, out, plan.to_dict(), "src.docx", "out.docx")


def new_ids(out: bytes) -> List[int]:
    root = DocxPackage(out).xml("word/document.xml")
    return sorted(int(e.get(W + "id")) for e in root.iter(W + "ins", W + "del"))


@pytest.mark.parametrize("change", ["tblGridChange", "tblPrExChange"])
def test_a_table_change_raises_the_seed(change):
    assert max_revision_id(DocxPackage(resized_table(change))) == 7


@pytest.mark.parametrize("change", ["tblGridChange", "tblPrExChange"])
def test_a_tracked_replacement_beside_a_resized_table_passes(change):
    src = resized_table(change)
    plan, out, rec = run(src)
    assert plan.rev_delta["ins"] + plan.rev_delta["del"] == 2 * EDITS >= 7     # reaches id 7 from a seed of 0
    assert rec.passed, rec.to_markdown()
    assert rec.get("V4").status == "PASS" and rec.get("V4").data["colliding_new_ids"] == []
    assert new_ids(out) == list(range(8, 8 + 2 * EDITS))
    body = DocxPackage(out).xml("word/document.xml").find(W + "body")
    assert [accepted_text(p) for p in body.findall(W + "p")][-EDITS:] == ["The Client shall review."] * EDITS


def test_every_document_part_and_every_annotation_id_counts():
    """A property change in styles.xml (outside the text parts) and a bookmark's id both raise the seed; ids that are
    not numbers are skipped; a document part that is not well-formed XML and is no text part is skipped, not refused,
    and a part outside the document (customXml) is not read at all."""
    src = resized_table(rev_id="3")

    def style_change(root):
        rpr = root.find(f"{W}style/{W}rPr")
        etree.SubElement(rpr, W + "rPrChange", {W + "id": "40", W + "author": "Reviewer"}).append(etree.Element(W + "rPr"))

    def bookmark(root):
        p = root.find(f"{W}body/{W}p")
        p.insert(0, etree.Element(W + "bookmarkStart", {W + "id": "12", W + "name": "here"}))
        p.append(etree.Element(W + "bookmarkEnd", {W + "id": "12"}))
        p.append(etree.Element(W + "permStart", {W + "id": "not-a-number", W + "edGrp": "everyone"}))

    styled = mutate(src, STYLES, style_change)
    assert max_revision_id(DocxPackage(styled)) == 40
    assert max_revision_id(DocxPackage(mutate(src, "word/document.xml", bookmark))) == 12

    for broken in ("word/item9.xml", "customXml/item9.xml"):
        pkg = DocxPackage(styled)
        pkg.add_member(broken, b"<broken", "application/xml")
        assert max_revision_id(pkg) == 40
        assert (broken in pkg.wordprocessing_xml_parts()) == broken.startswith("word/")   # read and passed over, or not read


def test_shared_defect_the_old_seed_collides_and_v4_catches_it(monkeypatch):
    """The old seed's answer for this file (0: w:tblGridChange was not on its list), shared by plan and apply: the
    verifier's own count of every revision element finds the new id 7 beside the column resize's."""
    monkeypatch.setattr(replace_module, "max_revision_id", lambda pkg: 0)
    _, out, rec = run(resized_table())
    assert not rec.passed and rec.get("V4").status == "FAIL"
    assert rec.get("V4").data["colliding_new_ids"] == ["7"]
    assert new_ids(out) == list(range(1, 1 + 2 * EDITS))
