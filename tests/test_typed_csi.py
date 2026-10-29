"""convert_typed_numbers with ``typed_as: "csi"`` (docs/SPEC_FORMATTER_MODES_DESIGN.md N6): typed CSI numbers ("PART 1 -",
"1.1" / "1.01", "A.", "1.", "a.") become a list's numbers whose look is another's — a Canadian PageFormat list (``PART
%1``, ``%1.%2``, ``.%3``, ``.%4``, ``.%5``) — wherever their values are the list's counters. End to end on the typed spec
(``build_typed``, plain and MasterSpec-like, tracked and plain), the refusals (typed_as without counters, a skipped
letter), a typed Canadian ".1" left as typed, the CSI shapes excluding one another on both sides, and wrong
implementations shared by plan and apply, each caught by the verifier's own reading."""
from __future__ import annotations

import io
from typing import Any, Dict, List, Tuple

import pytest
from docx import Document

from make_fixtures import _abstract_xml, _add_lists, _add_style, _lvl_xml, _style_name, _typed_paragraph, build_typed
from spec_surgeon.catalog import COMMON_ASSUMPTIONS
from spec_surgeon.ops import JobException, get_op, merge_fragments
from spec_surgeon.ops import convert_typed_numbers as ctn
from spec_surgeon.ops import lists as tl
from spec_surgeon.ops.selectors import StyleSheet
from spec_surgeon.package import DocxPackage, W
from spec_surgeon.verify import run_suite
from spec_surgeon.verify import typednumbers as vt

MAIN = "word/document.xml"
PF_LEVELS = ["CSC0", "CSC1", "CSC2", "CSC3", "CSC4"]

# build_typed(): the PageFormat number each typed paragraph takes, by the text that follows its typed number
PAGEFORMAT = {"GENERAL": "PART 1", "SUMMARY": "1.1", "Section includes wet-pipe sprinkler systems.": ".1",
              "Related requirements:": ".2", 'Division 21 Section "Common Work Results for Fire Suppression".': ".1",
              "Hangers.": ".1", "Supports.": ".2", 'Division 28 Section "Fire Detection and Alarm".': ".2",
              "Submittals.": ".3", "REFERENCES": "1.2", "NFPA 13.": ".1", "PRODUCTS": "PART 2", "MANUFACTURERS": "2.1",
              "Tyco.": ".1", "Viking.": ".2", "Reliable.": ".3", "MATERIALS": "2.2", "Steel pipe.": ".1",
              "EXECUTION": "PART 3", "INSTALLATION": "3.1", "Install per NFPA 13.": ".1", "Hangers per NFPA 13.": ".1"}


def build_pageformat_scheme() -> bytes:
    """An example file numbered like CSC PageFormat: paragraph styles CSC0–CSC4 on one list (num 30) — ``PART %1``,
    ``%1.%2``, ``.%3``, ``.%4``, ``.%5``, every level decimal, starting at 1 — each style naming its level, and a
    paragraph of each."""
    doc = Document()
    for k, sid in enumerate(PF_LEVELS):
        _add_style(doc, f'<w:style w:type="paragraph" w:customStyle="1" w:styleId="{sid}"><w:name w:val="{sid}"/>'
                        f'<w:basedOn w:val="Normal"/><w:pPr><w:numPr><w:ilvl w:val="{k}"/><w:numId w:val="30"/>'
                        f'</w:numPr></w:pPr></w:style>')
    texts = ["PART %1", "%1.%2", ".%3", ".%4", ".%5"]
    levels = [_lvl_xml(k, "decimal", texts[k], pstyle=PF_LEVELS[k], left=720 * k, hanging=720 if k else 0)
              for k in range(5)]
    _add_lists(doc, [_abstract_xml(30, levels, "6B6B0030")], ['<w:num w:numId="30"><w:abstractNumId w:val="30"/></w:num>'])
    for sid, text in zip(PF_LEVELS, ("GENERAL", "SUMMARY", "Paragraph.", "Subparagraph.", "Sub-subparagraph.")):
        doc.add_paragraph(text, style=doc.styles[_style_name(doc, sid)])
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_lines(lines: List[str]) -> bytes:
    """A Normal paragraph per line (a tab inside a line becomes a w:tab)."""
    doc = Document()
    for line in lines:
        _typed_paragraph(doc, [line])
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.fixture()
def pf(tmp_path) -> str:
    p = tmp_path / "CSC PageFormat scheme.docx"
    p.write_bytes(build_pageformat_scheme())
    return str(p)


def live(p) -> str:
    out = []
    for el in p.iter(W + "t", W + "tab"):
        q, gone = el.getparent(), False
        while q is not None and q.tag != W + "p":
            gone = gone or q.tag in (W + "del", W + "moveFrom")
            q = q.getparent()
        if q is p and not gone and (el.tag == W + "t" or el.getparent().tag == W + "r"):
            out.append(el.text or "" if el.tag == W + "t" else "\t")
    return "".join(out)


def numbers(data: bytes) -> Dict[str, Any]:
    """The live text of each numbered paragraph → its rendered number (by the transform's model)."""
    pkg = DocxPackage(data)
    root = pkg.xml(MAIN)
    pars = list(root.iter(W + "p"))
    return {live(pars[x.index]): (x.text if x.evaluated else (x.issue,))
            for x in tl.count(root, StyleSheet(pkg), tl.ListDefs.load(pkg))}


def converting(data: bytes, tpl: str, select: dict = None, **params) -> Tuple[Any, bytes, Any]:
    params = {"levels": list(PF_LEVELS), "template": tpl, "match": "counters", "typed_as": "csi",
              "author": "Abraham Borg", **params}
    pkg = DocxPackage(data)
    frag = get_op("convert_typed_numbers").plan(pkg, select or {}, params)
    plan = merge_fragments([frag])
    work = pkg.copy()
    get_op(frag.op).apply(work, frag)
    out = work.output_bytes()
    return plan, out, run_suite(data, out, plan.to_dict(), "src.docx", "out.docx")


def refused(data: bytes, tpl: str, **params) -> JobException:
    params = {"levels": list(PF_LEVELS), "template": tpl, **params}
    with pytest.raises(JobException) as e:
        get_op("convert_typed_numbers").plan(DocxPackage(data), {}, params)
    return e.value


def failed(rec) -> List[str]:
    return [c.id for c in rec.checks if c.status == "FAIL"]


def v10(rec) -> str:
    c = next(c for c in rec.checks if c.id == "V10")
    return c.summary + " | " + " | ".join(c.data.get("errors", []))


# ---- end to end -------------------------------------------------------------------------------------------------------

@pytest.mark.parametrize("masterspec", [False, True])
@pytest.mark.parametrize("tracked", [True, False])
def test_typed_csi_numbers_become_pageformat_numbers(pf, masterspec, tracked):
    """Every typed CSI number of the body becomes the PageFormat list's number with the same values — "A." shows ".1",
    "1.01" shows "1.1", "PART 1 –" shows "PART 1" — the typed characters (and the dash and space after PART) removed,
    nothing written; the table's typed numbers stay as typed; tracked as deletions, or plain."""
    src = build_typed(masterspec=masterspec)
    plan, out, rec = converting(src, pf, tracked=tracked)
    assert rec.passed, rec.to_markdown()
    got = numbers(out)
    assert {k: got.get(k) for k in PAGEFORMAT} == PAGEFORMAT
    root = DocxPackage(out).xml(MAIN)
    assert any(live(p) == "1.\tCell item." for p in root.iter(W + "p"))            # a table cell: left as typed
    general = next(p for p in root.iter(W + "p") if live(p) == "GENERAL")         # "PART 1 - " gone, dash and all
    assert general.find(f"{W}pPr/{W}pStyle").get(W + "val") == "CSC0"
    a = next(x for x in plan.property_assertions if x["kind"] == "typed_numbers")
    assert a["params"]["typed_as"] == "csi" and len(a["converted"]) == 22
    assert plan.mode == ("content" if tracked else "untracked")
    assert any("will show as your list writes them" in f and "‘A.’ → ‘.1’" in f for f in plan.flags), plan.flags
    if masterspec:
        assert any("‘1.01’ → ‘1.1’" in f for f in plan.flags), plan.flags


def test_a_typed_canadian_number_is_left_as_typed(pf):
    """A typed ".1" carries no depth: it is never read, the paragraph stays as typed and reads the same; the CSI
    numbers around it convert."""
    src = build_lines(["PART 1 - GENERAL", "1.1\tSUMMARY", ".1\tSection includes sprinklers.", "A.\tRelated work."])
    plan, out, rec = converting(src, pf, tracked=False)
    assert rec.passed, rec.to_markdown()
    root = DocxPackage(out).xml(MAIN)
    assert any(live(p) == ".1\tSection includes sprinklers." for p in root.iter(W + "p"))
    got = numbers(out)
    assert got["GENERAL"] == "PART 1" and got["SUMMARY"] == "1.1" and got["Related work."] == ".1"
    a = next(x for x in plan.property_assertions if x["kind"] == "typed_numbers")
    assert [c[0] for c in a["converted"]] == [0, 1, 3]


# ---- refusals ---------------------------------------------------------------------------------------------------------

def test_typed_as_needs_counters_and_a_known_value(pf):
    src = build_typed()
    e = refused(src, pf, typed_as="csi")                                   # match exact (the default)
    assert e.code == "bad-params" and "match 'counters'" in e.reason
    e = refused(src, pf, typed_as="csi", match="exact")
    assert e.code == "bad-params"
    e = refused(src, pf, typed_as="canadian", match="counters")
    assert e.code == "bad-params" and "typed_as must be one of" in e.reason


@pytest.mark.parametrize("lines, reader", [
    (["PART 1 - GENERAL", "1.1\tSUMMARY", "A.\tFirst.", "C.\tSkipped a letter."], "Skipped a letter."),
    (["PART 1 - GENERAL", "1.1\tSUMMARY", "A.\tFirst.", "B.\tSecond.", "B.\tTwice."], "Twice."),
    (["PART 1 - GENERAL", "1.2\tSUMMARY", "A.\tFirst."], "SUMMARY"),
])
def test_a_typo_in_the_typed_values_is_refused(pf, lines, reader):
    """A skipped letter, a duplicate or an article typed 1.2 where the list counts 1.1: the values are not the list's
    counters there — refused, naming the paragraph."""
    e = refused(build_lines(lines), pf, match="counters", typed_as="csi")
    assert e.code == "numbering-mismatch" and reader in e.reason, e.reason


def test_the_csi_shapes_exclude_one_another_on_both_sides(pf):
    """A typed CSI number reads at one level at most — no token fits two shapes — and both sides read the same level
    and the same values; letters are letters (no Roman reading): "i." is the ninth letter at level 4."""
    scheme = tl.Scheme(num_id=1, abstract_id=1, levels=[
        tl.SchemeLevel(style=f"S{k}", num_id=1, ilvl=k, level=None, pattern="") for k in range(9)])
    tokens = {"PART 1 - X": (0, {0: 1}), "PART 12 X": (0, {0: 12}), "PART 3: X": (0, {0: 3}),
              "PART 2 – X": (0, {0: 2}), "1.01\tX": (1, {0: 1, 1: 1}), "2.10 X": (1, {0: 2, 1: 10}),
              "A.\tX": (2, {2: 1}), "AA. X": (2, {2: 27}), "I. X": (2, {2: 9}), "12. X": (3, {3: 12}),
              "i.\tX": (4, {4: 9}), "b. X": (4, {4: 2}), "3) X": (5, {5: 3}), "c) X": (6, {6: 3}),
              "(4) X": (7, {7: 4}), "(d) X": (8, {8: 4})}
    for text, (level, values) in tokens.items():
        tn = tl.csi_typed_number(text, scheme)
        assert tn is not None and tn.levels == (level,), (text, tn)
        assert tl.csi_typed_values(tn.token, level) == values, text
        vn = vt._v_csi_typed(text, scheme)
        assert vn is not None and vn.levels == (level,) and (vn.token, vn.end) == (tn.token, tn.end), (text, vn)
        assert vt._v_csi_values(vn.token, level) == values, text
    for text in (".1 X", "1.1.1 X", "PART X", "Part 1 - X", "1.1X", "A.X"):
        assert tl.csi_typed_number(text, scheme) is None and vt._v_csi_typed(text, scheme) is None, text
    # a run of two letters is found at its level but reads as no value (Word repeats one letter: "AA"), on both sides —
    # the conversion then refuses it, as the list's own lenient reading does
    assert tl.csi_typed_number("AB. X", scheme).levels == (2,) and vt._v_csi_typed("AB. X", scheme).levels == (2,)
    assert tl.csi_typed_values("AB.", 2) is None and vt._v_csi_values("AB.", 2) is None


def test_a_scheme_shorter_than_the_typed_levels_leaves_the_deeper_ones_typed(pf, tmp_path):
    """A three-level scheme: typed "1." and "a." find no level of it and stay as typed (reading the same); the
    others convert."""
    src = build_typed()
    plan, out, rec = converting(src, pf, levels=PF_LEVELS[:3], tracked=False)
    assert rec.passed, rec.to_markdown()
    a = next(x for x in plan.property_assertions if x["kind"] == "typed_numbers")
    assert {c[1] for c in a["converted"]} == {0, 1, 2}
    root = DocxPackage(out).xml(MAIN)
    assert any(live(p) == "a.\tHangers." for p in root.iter(W + "p"))


# ---- the words ----------------------------------------------------------------------------------------------------------

def test_the_sentence_and_the_catalog_name_typed_csi_numbers():
    words = get_op("convert_typed_numbers").describe({}, {"levels": PF_LEVELS, "template": "S.docx",
                                                          "match": "counters", "typed_as": "csi"})
    assert "typed CSI numbers such as “A.” become your list's numbers wherever their values are the list's" in words
    assert any("typed_as 'csi'" in x and "never converted" in x for x in COMMON_ASSUMPTIONS["convert_typed_numbers"])
    schema = get_op("convert_typed_numbers").params_schema["properties"]["typed_as"]
    assert schema["enum"] == ["csi"]


# ---- wrong implementations shared by plan and apply are caught ---------------------------------------------------------

def test_an_apply_that_leaves_a_typed_number_is_caught(pf, monkeypatch):
    """The plan is right; the transform skips one paragraph's removal: the verifier sees the typed "A." still there."""
    real = ctn.ConvertTypedNumbers.apply

    def leave_one(self, pkg, frag):
        first = sorted(frag.transform_data["edits"], key=int)[2]
        frag.transform_data = {**frag.transform_data,
                               "edits": {k: v for k, v in frag.transform_data["edits"].items() if k != first}}
        return real(self, pkg, frag)
    monkeypatch.setattr(ctn.ConvertTypedNumbers, "apply", leave_one)
    plan, out, rec = converting(build_typed(), pf, tracked=False)
    assert not rec.passed and "V10" in failed(rec), rec.to_markdown()


def test_a_shared_wrong_reading_of_a_letter_is_caught(pf, monkeypatch):
    """Plan and apply read a typed letter one too low ("C." as 2): a skipped letter then passes the transform's own
    guarantee — the verifier's own reading values "C." at 3 and says the op had to refuse."""
    src = build_lines(["PART 1 - GENERAL", "1.1\tSUMMARY", "A.\tFirst.", "C.\tSkipped a letter."])
    real = tl.csi_typed_values

    def low(token, k):
        values = real(token, k)
        if values and k == 2 and values[2] > 1:
            values = {**values, 2: values[2] - 1}
        return values
    monkeypatch.setattr(tl, "csi_typed_values", low)
    plan, out, rec = converting(src, pf, tracked=False)
    assert "V10" in failed(rec) and "the op must refuse: numbering-mismatch" in v10(rec), v10(rec)


def test_a_shared_shape_that_misses_the_part_dash_is_caught(pf, monkeypatch):
    """Plan and apply read "PART 1" without its dash: "- GENERAL" is left in the paragraph, which the transform's own
    reading accepts — the verifier's own shape takes the dash with the number, and the prefixes differ."""
    shapes = list(tl.CSI_SHAPES)
    shapes[0] = ("PART[  ]+([0-9]+)", shapes[0][1])
    monkeypatch.setattr(tl, "CSI_SHAPES", tuple(shapes))
    plan, out, rec = converting(build_typed(), pf, tracked=False)
    assert "V10" in failed(rec) and "typed numbers found here" in v10(rec), v10(rec)


def test_without_typed_as_the_op_reads_the_list_s_own_shape(pf):
    """The default is untouched: without typed_as, counters mode finds a typed number by the destination list's own
    lenient pattern, so a typed "A." is no ".%3" number and stays as typed (the file reads the same)."""
    pkg = DocxPackage(build_typed())
    frag = get_op("convert_typed_numbers").plan(pkg, {}, {"levels": PF_LEVELS, "template": pf, "match": "counters",
                                                          "tracked": False})
    a = next(x for x in frag.property_assertions if x["kind"] == "typed_numbers")
    assert "typed_as" not in a["params"]
    assert all(c[1] in (0, 1) for c in a["converted"])            # PART and the articles alone, by their own shapes
