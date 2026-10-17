"""Adversarial final-goal checks: a correct formatting assertion cannot hide damage."""
from __future__ import annotations

import copy
import io
import zipfile

import pytest
from docx import Document
from lxml import etree

from spec_surgeon import goals
from spec_surgeon.verify.goals import verify_goal
from spec_surgeon.verify.reader import W
from test_goals import TITLE, execute, goal, write_document


def mutate(data, change, part="word/document.xml"):
    with zipfile.ZipFile(io.BytesIO(data)) as source:
        root = etree.fromstring(source.read(part))
        change(root)
        output = io.BytesIO()
        with zipfile.ZipFile(output, "w") as result:
            for info in source.infolist():
                result.writestr(info, etree.tostring(root) if info.filename == part else source.read(info))
    return output.getvalue()


def paragraphs(root):
    return list(root.iter(W + "p"))


@pytest.fixture
def verified(tmp_path):
    source = write_document(tmp_path / "spec.docx")
    doc = Document(source)
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Table words\tKEEP THEM"
    doc.sections[0].header.paragraphs[0].text = "Original header"
    doc.save(source)
    contract = goal(paragraph={"keepNext": True})
    binding = goals.resolve_goal(source, contract)["binding"]
    output, record, result = execute(source, contract, binding)
    assert record.passed, record.to_markdown()
    assert result["passed"], result
    return source, output, contract, binding


@pytest.mark.parametrize("damage", ["body_text", "table_text", "tab", "title_run", "unrelated_bold",
                                    "run_attribute", "paragraph_attribute", "run_reorder"])
def test_correct_title_formatting_cannot_hide_content_or_structure_damage(verified, damage):
    source, output, contract, binding = verified

    def corrupt(root):
        title, body, cell = paragraphs(root)[1:4]
        if damage == "body_text":
            body.find(".//" + W + "t").text = "Words silently replaced."
        elif damage == "table_text":
            cell.find(".//" + W + "t").text = "LOST TABLE CONTENT"
        elif damage == "tab":
            tab = cell.find(".//" + W + "tab")
            tab.getparent().remove(tab)
        elif damage == "title_run":
            title.remove(title.findall(W + "r")[1])
        elif damage == "unrelated_bold":
            rpr = etree.SubElement(body.find(W + "r"), W + "rPr")
            etree.SubElement(rpr, W + "b")
        elif damage == "run_attribute":
            title.find(W + "r/" + W + "rPr").set(W + "unexpected", "1")
        elif damage == "paragraph_attribute":
            title.find(W + "pPr").set(W + "unexpected", "1")
        elif damage == "run_reorder":
            title.append(title.findall(W + "r")[0])

    result = verify_goal(source, mutate(output, corrupt), contract, binding)
    assert not result["passed"], damage
    assert result["errors"] and any(check["status"] == "FAIL" for check in result["checks"])


@pytest.mark.parametrize("element", ["b", "bCs"])
def test_title_requires_both_effective_bold_properties(verified, element):
    source, output, contract, binding = verified

    def remove(root):
        rpr = paragraphs(root)[1].find(W + "r/" + W + "rPr")
        rpr.remove(rpr.find(W + element))

    result = verify_goal(source, mutate(output, remove), contract, binding)
    assert not result["passed"]
    assert any(element in error for error in result["errors"])


def test_other_package_members_remain_byte_identical(verified):
    source, output, contract, binding = verified
    corrupted = mutate(output, lambda root: setattr(root.find(".//" + W + "t"), "text", "Wrong header"),
                       "word/header1.xml")
    result = verify_goal(source, corrupted, contract, binding)
    assert not result["passed"]
    assert any("header1.xml" in error for error in result["errors"])


def test_final_target_must_retain_its_section_evidence(verified):
    source, output, contract, binding = verified
    corrupted = mutate(output, lambda root: setattr(paragraphs(root)[0].find(".//" + W + "t"),
                                                     "text", "SECTION 22 11 16"))
    result = verify_goal(source, corrupted, contract, binding)
    assert not result["passed"]
    assert any("section evidence" in error for error in result["errors"])


@pytest.mark.parametrize("field", ["source_sha256", "contract_sha256", "candidate_id", "target", "choice"])
def test_forged_binding_never_counts_as_evidence(verified, field):
    source, output, contract, binding = verified
    bad = copy.deepcopy(binding)
    if field == "target":
        bad[field]["in_table"] = 0  # Python equality must not make integers into JSON booleans.
    else:
        bad[field] = "forged"
    result = verify_goal(source, output, contract, bad)
    assert not result["passed"]


def test_changed_contract_or_original_bytes_require_a_new_binding(verified):
    source, output, contract, binding = verified
    revised = copy.deepcopy(contract)
    revised["paragraph"]["keepNext"] = False
    assert not verify_goal(source, output, revised, binding)["passed"]
    assert not verify_goal(output, output, contract, binding)["passed"]


def test_unique_mode_is_independently_checked(tmp_path):
    source = write_document(tmp_path / "duplicate.docx", duplicate=True)
    contract = goal()
    selected = goals.resolve_goal(source, contract)["candidates"][0]
    binding = goals.choose_goal(source, contract, selected["candidate_id"])
    output, _, result = execute(source, contract, binding)
    assert result["passed"]
    assert not verify_goal(source, output, contract, {**binding, "choice": "unique"})["passed"]


def test_supported_title_does_not_silently_beat_an_unsupported_matching_title(tmp_path):
    source = write_document(tmp_path / "textbox.docx")

    def add_textbox(root):
        outer = paragraphs(root)[-1]
        content = etree.SubElement(etree.SubElement(outer, W + "r"), W + "txbxContent")
        title = etree.SubElement(content, W + "p")
        etree.SubElement(etree.SubElement(title, W + "r"), W + "t").text = TITLE

    source.write_bytes(mutate(source.read_bytes(), add_textbox))
    contract = goal()
    candidates = goals.resolve_goal(source, contract)["candidates"]
    binding = goals.choose_goal(source, contract, candidates[0]["candidate_id"])
    # We can verify a deliberate choice without pretending the title is unique.
    output, _, result = execute(source, contract, binding)
    assert result["passed"], result
    assert not verify_goal(source, output, contract, {**binding, "choice": "unique"})["passed"]


def test_verifier_does_not_share_a_broken_transform_target_resolver(tmp_path, monkeypatch):
    source = write_document(tmp_path / "spec.docx")
    original = goals._visible_text
    wrong_title = "THIS TITLE IS NOT IN THE DOCUMENT"
    monkeypatch.setattr(goals, "_visible_text", lambda p: wrong_title if original(p) == TITLE else original(p))
    contract = goal(target={"title": wrong_title})
    binding = goals.resolve_goal(source, contract)["binding"]
    # The buggy resolver and compiler agree about their forged target, and every
    # stage's formatting verification passes. Original-goal verification must fail.
    _, record, final = execute(source, contract, binding)
    assert record.passed, record.to_markdown()
    assert not final["passed"]
    assert any("source candidate" in error for error in final["errors"])


def test_paragraph_only_goal_cannot_authorize_paragraph_mark_formatting(tmp_path):
    source = write_document(tmp_path / "spec.docx")
    contract = {"kind": "format_section_title", "target": {"title": TITLE}, "paragraph": {"keepNext": True}}
    binding = goals.resolve_goal(source, contract)["binding"]
    output, _, result = execute(source, contract, binding)
    assert result["passed"]

    def corrupt(root):
        mark = etree.SubElement(paragraphs(root)[1].find(W + "pPr"), W + "rPr")
        etree.SubElement(mark, W + "b")

    final = verify_goal(source, mutate(output, corrupt), contract, binding)
    assert not final["passed"]
    assert any("paragraph-mark" in error for error in final["errors"])


@pytest.mark.parametrize("damage", ["snapshot", "author", "duplicate_id", "extra_marker", "snapshot_attribute"])
def test_tracked_changes_need_exact_original_snapshots_and_fresh_ids(tmp_path, damage):
    source = write_document(tmp_path / "tracked.docx")
    contract = goal(tracked=True, paragraph={"keepNext": True})
    binding = goals.resolve_goal(source, contract)["binding"]
    output, record, result = execute(source, contract, binding)
    assert record.passed and result["passed"]

    def corrupt(root):
        revisions = list(root.iter(W + "rPrChange"))
        first = revisions[0]
        if damage == "snapshot":
            etree.SubElement(first.find(W + "rPr"), W + "i")
        elif damage == "author":
            first.set(W + "author", "Someone else")
        elif damage == "duplicate_id":
            revisions[1].set(W + "id", first.get(W + "id"))
        elif damage == "extra_marker":
            first.getparent().append(copy.deepcopy(first))
        elif damage == "snapshot_attribute":
            first.find(W + "rPr").set(W + "unexpected", "1")

    final = verify_goal(source, mutate(output, corrupt), contract, binding)
    assert not final["passed"], damage


def test_unknown_or_malformed_contracts_fail_closed(verified):
    source, output, contract, binding = verified
    for malformed in (None, [], {**contract, "ignore_errors": True}, {**contract, "tracked": 1},
                      {**contract, "run": {"b": "true"}}, {**contract, "run": {"highlight": True}}):
        result = verify_goal(source, output, malformed, binding)
        assert result["passed"] is False and result["errors"]
    assert not verify_goal(b"not a ZIP", output, contract, binding)["passed"]


def test_new_revision_id_cannot_collide_with_an_untouched_header_revision(tmp_path):
    source = write_document(tmp_path / "tracked.docx")
    doc = Document(source)
    doc.sections[0].header.paragraphs[0].text = "Existing header insertion"
    doc.save(source)

    def insert_revision(root):
        paragraph = paragraphs(root)[0]
        run = paragraph.find(W + "r")
        paragraph.remove(run)
        insertion = etree.SubElement(paragraph, W + "ins", {W + "id": "321", W + "author": "Original author"})
        insertion.append(run)

    source.write_bytes(mutate(source.read_bytes(), insert_revision, "word/header1.xml"))
    contract = goal(tracked=True)
    binding = goals.resolve_goal(source, contract)["binding"]
    output, record, final = execute(source, contract, binding)
    assert record.passed and final["passed"], final
    corrupted = mutate(output, lambda root: next(root.iter(W + "rPrChange")).set(W + "id", "321"))
    final = verify_goal(source, corrupted, contract, binding)
    assert not final["passed"]
    assert any("reuses id 321" in error for error in final["errors"])


def test_unmodelled_conditional_table_properties_cannot_supply_a_success(tmp_path):
    source = write_document(tmp_path / "table.docx", layout="table")
    doc = Document(source)
    doc.tables[0].style = "Table Grid"
    doc.save(source)

    def conditional(root):
        table_style = next(style for style in root.findall(W + "style") if style.get(W + "styleId") == "TableGrid")
        row = etree.SubElement(table_style, W + "tblStylePr", {W + "type": "firstRow"})
        etree.SubElement(etree.SubElement(row, W + "rPr"), W + "b")

    source.write_bytes(mutate(source.read_bytes(), conditional, "word/styles.xml"))
    contract = goal()
    binding = goals.resolve_goal(source, contract)["binding"]
    _, record, final = execute(source, contract, binding)
    assert record.passed, record.to_markdown()
    assert not final["passed"]
    assert any("conditional table-style" in error for error in final["errors"])
