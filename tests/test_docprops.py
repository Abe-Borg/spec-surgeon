"""set_document_properties and set_document_state (docs/BUILTIN_OPS_BATCH2_DESIGN.md §11, §12) — the two operations the
method used to forbid and now does only on request: document properties set and cleared, custom text properties,
Track Changes turned off or on, Mark as Final. End to end through the whole verification suite; the refusals; planted
defects in the output and in an implementation plan and apply share, each caught; the plan-time conflicts; the two
sides' tables agree."""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any, Dict, Optional

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

from make_fixtures import add_parts, docx_bytes, patch_docx
from spec_surgeon.editspec import SpecError, apply_spec, load_spec, plan_spec
from spec_surgeon.ops import JobException, get_op
from spec_surgeon.ops import docprops as dp
from spec_surgeon.package import DocxPackage
from spec_surgeon.verify import docprops as vdp
from spec_surgeon.verify import run_suite
from test_text_edits import failed, mutate, v10_errors

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W = "{%s}" % W_NS
CP = "{http://schemas.openxmlformats.org/package/2006/metadata/core-properties}"
DC = "{http://purl.org/dc/elements/1.1/}"
EP = "{http://schemas.openxmlformats.org/officeDocument/2006/extended-properties}"
CU_NS = "http://schemas.openxmlformats.org/officeDocument/2006/custom-properties"
VT_NS = "http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes"
CU, VT = "{%s}" % CU_NS, "{%s}" % VT_NS
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
CUSTOM_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/custom-properties"
CUSTOM_CT = "application/vnd.openxmlformats-officedocument.custom-properties+xml"
CORE, APP, SETTINGS, CUSTOM = "docProps/core.xml", "docProps/app.xml", "word/settings.xml", "docProps/custom.xml"
EXAMPLES = Path(__file__).resolve().parent.parent / "examples" / "specs"


# ---- fixtures --------------------------------------------------------------------------------------------------------

def _field(paragraph, instr: str, result: str) -> None:
    for kind, text in (("begin", None), (None, instr), ("separate", None), (None, None), ("end", None)):
        r = paragraph.add_run()
        if kind:
            fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), kind); r._r.append(fc)
        elif text is not None:
            it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = text; r._r.append(it)
        else:
            t = OxmlElement("w:t"); t.text = result; r._r.append(t)


def _custom_xml(props) -> bytes:
    root = etree.Element(CU + "Properties", nsmap={None: CU_NS, "vt": VT_NS})
    for pid, (name, tag, value) in enumerate(props, 2):
        p = etree.SubElement(root, CU + "property")
        p.set("fmtid", "{D5CDD505-2E9C-101B-9397-08002B2CF9AE}"); p.set("pid", str(pid)); p.set("name", name)
        etree.SubElement(p, VT + tag).text = value
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def build(*, title: Optional[str] = None, creator: Optional[str] = None, last_modified_by: Optional[str] = None,
          company: Optional[str] = None, status: Optional[str] = None, track: Optional[str] = None,
          locked: bool = False, custom=None, fields: bool = False, redline: bool = False) -> bytes:
    """A spec-like file with valid settings (in schema order, w:zoom with its percent) and the given properties:
    ``track`` "on" (w:trackRevisions), "off-val" (w:trackRevisions w:val="false") or None; ``locked`` tracking locked
    (w:documentProtection trackedChanges, enforced); ``custom`` [(name, vt tag, value)] in docProps/custom.xml;
    ``fields`` DOCPROPERTY Title and TITLE fields in the body, DOCPROPERTY "Project Number" in the header; ``redline``
    Jeremy's pending insertion and Jack's comment."""
    doc = Document()
    doc.add_paragraph("SECTION 21 13 13 - WET-PIPE SPRINKLER SYSTEMS")
    p = doc.add_paragraph("Body text ")
    if fields:
        _field(p, " DOCPROPERTY  Title  \\* MERGEFORMAT ", "Old Title")
        _field(doc.add_paragraph("By "), " TITLE ", "Old Title")
        _field(doc.sections[0].header.paragraphs[0], ' DOCPROPERTY "Project Number" ', "1001")
    if redline:
        q = doc.add_paragraph("Pipe shall be ")
        r = q.add_run("galvanized")
        ins = OxmlElement("w:ins"); ins.set(qn("w:id"), "31"); ins.set(qn("w:author"), "Jeremy")
        ins.set(qn("w:date"), "2026-09-22T10:15:00Z")
        r._r.addprevious(ins); ins.append(r._r)
        doc.add_comment(q.runs[0], text="Check the coating", author="Jack", initials="J")
    data = docx_bytes(doc)

    def core(root):
        for tag, value in ((DC + "title", title), (DC + "creator", creator), (CP + "lastModifiedBy", last_modified_by),
                           (CP + "contentStatus", status)):
            if value is None:
                continue
            el = root.find(tag)
            if el is None:
                el = etree.SubElement(root, tag)
            el.text = value

    def app(root):
        if company is not None:
            root.find(EP + "Company").text = company

    def settings(root):
        root.find(W + "zoom").set(W + "percent", "100")
        anchor = root.find(W + "defaultTabStop")
        if track:
            el = etree.Element(W + "trackRevisions")
            if track == "off-val":
                el.set(W + "val", "false")
            anchor.addprevious(el)
        if locked:
            prot = etree.Element(W + "documentProtection")
            prot.set(W + "edit", "trackedChanges"); prot.set(W + "enforcement", "1")
            anchor.addprevious(prot)
    data = patch_docx(data, {CORE: core, APP: app, SETTINGS: settings})
    if custom is not None:
        data = add_parts(data, {CUSTOM: _custom_xml(custom)},
                         rels=[("_rels/.rels", "rId90", CUSTOM_REL, CUSTOM)], overrides={CUSTOM: CUSTOM_CT})
    return data


def run_spec(src: bytes, ops, mode: str):
    spec = {"request": "r", "summary": "s", "mode": mode, "ops": ops}
    pkg = DocxPackage(src)
    plan, frags = plan_spec(pkg, spec)
    out = apply_spec(pkg, frags).output_bytes()
    return plan, out, run_suite(src, out, plan.to_dict(), "src.docx", "out.docx")


def props(src, mode="untracked", **params):
    return run_spec(src, [{"op": "set_document_properties", "params": params}], mode)


def state(src, **params):
    return run_spec(src, [{"op": "set_document_state", "params": params}], "formatting")


def verify(src: bytes, out: bytes, plan):
    return run_suite(src, out, plan.to_dict(), "src.docx", "out.docx")


def xml(data: bytes, member: str):
    return DocxPackage(data).xml(member)


def text_of(data: bytes, member: str, tag: str) -> Optional[str]:
    els = xml(data, member).findall(tag)
    assert len(els) <= 1
    return None if not els else (els[0].text or "")


def custom_props(data: bytes) -> Dict[str, Any]:
    pkg = DocxPackage(data)
    if not pkg.has(CUSTOM):
        return {}
    return {p.get("name"): (p.get("pid"), etree.QName(p[0]).localname, p[0].text) for p in xml(data, CUSTOM)}


def plan_only(src: bytes, op: str, **params):
    return get_op(op).plan(DocxPackage(src), {}, params)


# ---- set_document_properties: end to end -----------------------------------------------------------------------------

def test_set_the_title_the_company_and_a_new_custom_property():
    src = build(fields=False)
    plan, out, rec = props(src, set={"title": "Wet-Pipe Sprinkler Systems", "company": "ACME Fire Protection"},
                           custom={"Project Number": "1001"})
    assert rec.passed, rec.to_markdown()
    assert plan.mode == "untracked" and plan.geometry_reference == "source"
    assert text_of(out, CORE, DC + "title") == "Wet-Pipe Sprinkler Systems"
    assert text_of(out, APP, EP + "Company") == "ACME Fire Protection"
    # no custom properties part before: created, with exactly its relationship and content type
    assert plan.members_added == [CUSTOM]
    assert custom_props(out) == {"Project Number": ("2", "lpwstr", "1001")}
    rels = [r for r in xml(out, "_rels/.rels") if r.get("Type") == CUSTOM_REL]
    assert [r.get("Target") for r in rels] == [CUSTOM]
    ct = [o for o in xml(out, "[Content_Types].xml") if o.get("PartName") == "/" + CUSTOM]
    assert [o.get("ContentType") for o in ct] == [CUSTOM_CT]
    assert set(rec.get("V5a").data["baseline"]) >= {CORE, APP, CUSTOM}          # every touched part validated
    assert sorted(rec.get("V10").data["residual_parts"]) == [APP, CORE]     # the rest compared element for element


def test_clear_the_author_names():
    src = build(creator="Jeremy Smith", last_modified_by="Jack Jones", redline=True)
    plan, out, rec = props(src, mode="formatting", clear=["author", "last_modified_by"])
    assert rec.passed, rec.to_markdown()
    assert plan.mode == "formatting"                                               # nothing written: only removed
    assert text_of(out, CORE, DC + "creator") is None and text_of(out, CORE, CP + "lastModifiedBy") is None
    assert any("tracked change(s) keep their authors' names" in f for f in plan.flags)
    assert any("comment(s) keep their authors' names" in f for f in plan.flags)


def test_values_already_as_asked_change_nothing():
    src = build(title="Wet-Pipe Sprinkler Systems")
    # status is absent; category and last-modified-by are empty elements (no value, as Word shows them): all clear
    plan, out, rec = props(src, set={"title": "Wet-Pipe Sprinkler Systems"},
                           clear=["status", "category", "last_modified_by"])
    assert rec.passed and out == src
    assert plan.intended_parts == []
    assert any("already reads" in n for n in plan.notes)
    assert sum("nothing to clear" in n for n in plan.notes) == 3


def test_custom_properties_existing_new_and_cleared():
    src = build(custom=[("Project Number", "lpwstr", "1001"), ("Site", "lpstr", "VA1"), ("Budget", "i4", "5"),
                        ("Old", "lpwstr", "x")])
    plan, out, rec = props(src, custom={"project number": "1003", "Site": "VA3", "Phase": "CD"}, clear_custom=["OLD"])
    assert rec.passed, rec.to_markdown()
    got = custom_props(out)
    assert got["Project Number"] == ("2", "lpwstr", "1003")       # matched without case, its own name kept
    assert got["Site"] == ("3", "lpstr", "VA3")                    # the value element keeps its type
    assert got["Phase"] == ("6", "lpwstr", "CD")                   # new: the next pid after the largest (5)
    assert "Old" not in got and got["Budget"] == ("4", "i4", "5")
    assert plan.members_added == []


def test_a_custom_property_that_is_not_text_is_refused_for_that_file():
    src = build(custom=[("Budget", "i4", "5")])
    with pytest.raises(JobException) as e:
        plan_only(src, "set_document_properties", custom={"budget": "6"})
    assert e.value.code == "property-type"                    # the file's own part decides: never bad-params


def test_fields_that_show_a_changed_property_are_flagged():
    src = build(fields=True, custom=[("Project Number", "lpwstr", "1001")])
    plan, out, rec = props(src, set={"title": "New Title"}, custom={"Project Number": "1003"})
    assert rec.passed, rec.to_markdown()
    flags = "\n".join(plan.flags)
    assert "DOCPROPERTY 'Title' field (shows title)" in flags and "TITLE field (shows title)" in flags
    assert "DOCPROPERTY 'Project Number' field (shows custom property 'Project Number')" in flags
    assert "until Word updates fields" in flags
    # a renderer draws those fields with the new value: the source render is no layout reference for this file
    assert plan.geometry_reference == "none" and "fields show a changed property" in plan.geometry_reason
    # a property that does not change is not flagged
    plan2, _, _ = props(src, set={"title": "Old Title2"}, custom={"Project Number": "1001"})
    assert not any("Project Number" in f for f in plan2.flags)
    plan3, _, _ = props(src, custom={"Project Number": "1001", "Site": "VA3"})      # nothing a field shows changes
    assert plan3.flags == [] and plan3.geometry_reference == "source"


# ---- set_document_properties: refusals -------------------------------------------------------------------------------

@pytest.mark.parametrize("params, code", [
    ({}, "bad-params"),
    ({"set": {"title": ""}}, "bad-params"),
    ({"set": {"title": "   "}}, "bad-params"),
    ({"set": {"title": "Two\nlines"}}, "bad-params"),
    ({"set": {"title": "Zero​width"}}, "bad-params"),
    ({"set": {"title": "x" * 256}}, "bad-params"),
    ({"set": {"owner": "x"}}, "bad-params"),
    ({"set": {"title": "x"}, "clear": ["title"]}, "bad-params"),
    ({"custom": {"Site": "a"}, "clear_custom": ["SITE"]}, "bad-params"),
    ({"custom": {"": "a"}}, "bad-params"),
])
def test_params_that_are_wrong_whatever_the_file(params, code):
    with pytest.raises(JobException) as e:
        plan_only(build(), "set_document_properties", **params)
    assert e.value.code == code


def test_an_unknown_property_is_a_spec_error_before_any_file():
    with pytest.raises(SpecError):
        props(build(), set={"owner": "ACME"})


def _without(data: bytes, member: str, rel_type: str) -> bytes:
    """``data`` without ``member`` and the package relationship naming it (a file another tool wrote)."""
    import io
    import zipfile
    zin = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as z:
        for info in zin.infolist():
            if info.filename == member:
                continue
            payload = zin.read(info.filename)
            if info.filename == "_rels/.rels":
                root = etree.fromstring(payload)
                for r in list(root):
                    if r.get("Type") == rel_type:
                        root.remove(r)
                payload = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            if info.filename == "[Content_Types].xml":
                root = etree.fromstring(payload)
                for o in list(root):
                    if o.get("PartName") == "/" + member:
                        root.remove(o)
                payload = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            z.writestr(info, payload)
    return out.getvalue()


def test_a_missing_core_or_app_part_is_refused():
    no_core = _without(build(), CORE,
                       "http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties")
    with pytest.raises(JobException) as e:
        plan_only(no_core, "set_document_properties", set={"title": "T"})
    assert e.value.code == "no-part"
    no_app = _without(build(), APP,
                      "http://schemas.openxmlformats.org/officeDocument/2006/relationships/extended-properties")
    with pytest.raises(JobException) as e:
        plan_only(no_app, "set_document_properties", set={"company": "C"})
    assert e.value.code == "no-part"
    plan_only(no_app, "set_document_properties", set={"title": "T"})              # core alone is fine


def test_a_custom_part_no_relationship_names_is_refused():
    src = add_parts(build(), {CUSTOM: _custom_xml([("A", "lpwstr", "1")])}, overrides={CUSTOM: CUSTOM_CT})
    with pytest.raises(JobException) as e:
        plan_only(src, "set_document_properties", custom={"B": "2"})
    assert e.value.code == "malformed-properties"


# ---- set_document_state ----------------------------------------------------------------------------------------------

def test_turn_track_changes_off():
    src = build(track="on")
    plan, out, rec = state(src, track_changes="off")
    assert rec.passed, rec.to_markdown()
    assert xml(out, SETTINGS).find(W + "trackRevisions") is None
    assert plan.intended_parts == [SETTINGS]


@pytest.mark.parametrize("track", [None, "off-val"])
def test_turn_track_changes_on(track):
    src = build(track=track)
    plan, out, rec = state(src, track_changes="on")
    assert rec.passed, rec.to_markdown()
    root = xml(out, SETTINGS)
    names = [etree.QName(c).localname for c in root]
    tr = root.find(W + "trackRevisions")
    assert tr is not None and dict(tr.attrib) == {}
    assert names.index("trackRevisions") == names.index("defaultTabStop") - 1    # its place in CT_Settings


def test_already_off_or_on_changes_nothing():
    for track, ask in ((None, "off"), ("off-val", "off"), ("on", "on")):
        src = build(track=track)
        plan, out, rec = state(src, track_changes=ask)
        assert rec.passed and out == src and plan.intended_parts == []
        assert any("already" in n for n in plan.notes)


@pytest.mark.parametrize("track", ["on", None])
def test_locked_tracking_is_refused(track):
    with pytest.raises(JobException) as e:
        plan_only(build(track=track, locked=True), "set_document_state", track_changes="off")
    assert e.value.code == "protected"                    # the lock is read first, whatever the setting says
    plan_only(build(track=track, locked=True), "set_document_state", final=True)      # Mark as Final is not tracking


def test_mark_as_final_creates_the_custom_part():
    src = build()
    plan, out, rec = state(src, final=True)
    assert rec.passed, rec.to_markdown()
    assert custom_props(out) == {"_MarkAsFinal": ("2", "bool", "true")}
    assert text_of(out, CORE, CP + "contentStatus") == "Final"
    assert plan.members_added == [CUSTOM] and plan.mode == "formatting"


def test_mark_as_final_in_an_existing_custom_part_and_unmark():
    src = build(custom=[("Project Number", "lpwstr", "1001")], status="Draft")
    plan, out, rec = state(src, final=True, track_changes="off")
    assert rec.passed, rec.to_markdown()
    assert custom_props(out)["_MarkAsFinal"] == ("3", "bool", "true")
    assert text_of(out, CORE, CP + "contentStatus") == "Final"
    plan2, out2, rec2 = state(out, final=False)
    assert rec2.passed, rec2.to_markdown()
    assert "_MarkAsFinal" not in custom_props(out2) and text_of(out2, CORE, CP + "contentStatus") is None
    # another status is the user's: unmarking leaves it
    src3 = build(status="Issued for Bid")
    plan3, out3, rec3 = state(src3, final=False)
    assert rec3.passed and out3 == src3


def test_a_mark_as_final_property_of_another_type_is_refused():
    with pytest.raises(JobException) as e:
        plan_only(build(custom=[("_MarkAsFinal", "lpwstr", "yes")]), "set_document_state", final=True)
    assert e.value.code == "property-type"


@pytest.mark.parametrize("params", [{}, {"track_changes": "maybe"}, {"final": "yes"}])
def test_state_params_wrong_whatever_the_file(params):
    with pytest.raises(JobException) as e:
        plan_only(build(), "set_document_state", **params)
    assert e.value.code == "bad-params"


def test_no_settings_part_is_refused():
    src = build()

    def drop(root):
        for r in list(root):
            if r.get("Type", "").endswith("/settings"):
                root.remove(r)
    src = patch_docx(src, {"word/_rels/document.xml.rels": drop})
    with pytest.raises(JobException) as e:
        plan_only(src, "set_document_state", track_changes="on")
    assert e.value.code == "no-part"


# ---- planted defects: wrong outputs ----------------------------------------------------------------------------------

def test_another_core_property_changed_too_is_caught():
    src = build(title="Old")
    plan, out, rec = props(src, set={"title": "New"})
    assert rec.passed
    bad = mutate(out, CORE, lambda r: setattr(r.find(DC + "subject"), "text", "sneaked in"))
    rec = verify(src, bad, plan)
    assert failed(rec) == ["V10"] and any(e.startswith(f"{CORE}: unpredicted change") for e in v10_errors(rec))


def test_a_value_written_otherwise_is_caught():
    src = build()
    plan, out, rec = props(src, set={"title": "Wet-Pipe Sprinkler Systems"})
    for wrong in ("Wet-Pipe Sprinkler Systems ", "WET-PIPE SPRINKLER SYSTEMS"):
        bad = mutate(out, CORE, lambda r, w=wrong: setattr(r.find(DC + "title"), "text", w))
        assert any("title reads" in e for e in v10_errors(verify(src, bad, plan)))


def test_track_changes_left_on_or_another_setting_changed_is_caught():
    src = build(track="on")
    plan, out, rec = state(src, track_changes="off")
    rec = verify(src, src, plan)                                     # nothing done
    assert any("still in the output settings" in e for e in v10_errors(rec))
    extra = mutate(out, SETTINGS, lambda r: r.find(W + "defaultTabStop").set(W + "val", "360"))
    rec = verify(src, extra, plan)
    assert any(e.startswith(f"{SETTINGS}: unpredicted change") for e in v10_errors(rec))


def test_an_inserted_track_changes_out_of_schema_order_is_caught_where_v5a_cannot_see_it():
    src = build()
    plan, out, rec = state(src, track_changes="on")

    def misplace(root, first: bool):
        tr = root.find(W + "trackRevisions")
        root.remove(tr)
        if first:
            root.insert(0, tr)                                       # before w:zoom: out of order
        else:
            root.append(tr)                                          # after the w14 extensions: out of order
    for first in (True, False):
        rec = verify(src, mutate(out, SETTINGS, lambda r, f=first: misplace(r, f)), plan)
        assert any("schema order" in e for e in v10_errors(rec))
        # V5a sees the second only: its filter drops the first's message (it lists w:hyphenationZone as expected)
        assert rec.get("V5a").status == ("PASS" if first else "FAIL")


def test_a_created_custom_part_with_more_than_asked_or_a_wrong_package_entry_is_caught():
    src = build()
    plan, out, rec = state(src, final=True)
    assert rec.passed

    def extra(root):
        p = etree.SubElement(root, CU + "property")
        p.set("fmtid", "{D5CDD505-2E9C-101B-9397-08002B2CF9AE}"); p.set("pid", "3"); p.set("name", "Extra")
        etree.SubElement(p, VT + "lpwstr").text = "x"
    assert any("not exactly the asked properties" in e for e in v10_errors(verify(src, mutate(out, CUSTOM, extra),
                                                                                       plan)))

    def second_rel(root):
        r = etree.SubElement(root, "{%s}Relationship" % REL_NS)
        r.set("Id", "rId99"); r.set("Type", CUSTOM_REL); r.set("Target", CUSTOM)
    rec = verify(src, mutate(out, "_rels/.rels", second_rel), plan)
    assert "V10" in failed(rec)

    def no_override(root):
        for o in list(root):
            if o.get("PartName") == "/" + CUSTOM:
                root.remove(o)
    rec = verify(src, mutate(out, "[Content_Types].xml", no_override), plan)
    assert {"V1", "V10"} <= set(failed(rec))


def test_mark_as_final_with_the_wrong_status_is_caught():
    src = build()
    plan, out, rec = state(src, final=True)
    bad = mutate(out, CORE, lambda r: setattr(r.find(CP + "contentStatus"), "text", "final"))
    assert any("status reads 'final'" in e for e in v10_errors(verify(src, bad, plan)))


def test_a_renamed_custom_property_is_caught():
    src = build(custom=[("Project Number", "lpwstr", "1001")])
    plan, out, rec = props(src, custom={"Project Number": "1003"})
    assert rec.passed

    def rename(root):
        root.find(CU + "property").set("name", "project number")
    assert "V10" in failed(verify(src, mutate(out, CUSTOM, rename), plan))


# ---- planted defects: one wrong implementation shared by plan and apply ----------------------------------------------

def test_a_property_table_shared_wrong_by_plan_and_apply_is_caught(monkeypatch):
    src = build()
    monkeypatch.setitem(dp.PROPERTIES, "comments", ("core", DC + "subject"))     # comments are dc:description
    plan, out, rec = props(src, set={"comments": "Issued for permit"})
    assert text_of(out, CORE, DC + "subject") == "Issued for permit"            # plan and apply agreed
    assert failed(rec) == ["V10"]
    assert any("comments reads" in e for e in v10_errors(rec))


def test_a_new_property_shape_shared_wrong_is_caught(monkeypatch):
    src = build()
    monkeypatch.setattr(dp, "FMTID", "{00000000-0000-0000-0000-000000000000}")
    plan, out, rec = state(src, final=True)
    assert failed(rec) == ["V10"]


def test_a_schema_order_shared_wrong_is_caught(monkeypatch):
    src = build()
    order = [W + "trackRevisions"] + [t for t in dp.SETTINGS_ORDER if t != W + "trackRevisions"]
    monkeypatch.setattr(dp, "SETTINGS_ORDER", order)                 # plan and apply put it first, before w:zoom
    plan, out, rec = state(src, track_changes="on")
    assert etree.QName(xml(out, SETTINGS)[0]).localname == "trackRevisions"
    assert failed(rec) == ["V10"] and any("schema order" in e for e in v10_errors(rec))


def test_a_final_marker_shared_wrong_is_caught(monkeypatch):
    monkeypatch.setattr(dp, "FINAL_STATUS", "final")
    plan, out, rec = state(build(), final=True)
    assert failed(rec) == ["V10"]


def test_an_off_reading_shared_wrong_is_caught(monkeypatch):
    src = build(track="off-val")
    monkeypatch.setattr(dp, "FALSE", ())                    # w:val="false" read as on: the plan removes an off setting
    plan, out, rec = state(src, track_changes="off")
    assert xml(out, SETTINGS).find(W + "trackRevisions") is None
    assert failed(rec) == ["V10"]
    errs = v10_errors(rec)
    assert any("the plan predicted" in e for e in errs) and any("unpredicted change" in e for e in errs)


# ---- rule 9: two ops on one property or one new pid ------------------------------------------------------------------

def test_the_status_set_twice_is_a_conflict():
    with pytest.raises(JobException) as e:
        run_spec(build(), [{"op": "set_document_properties", "params": {"set": {"status": "Draft"}}},
                           {"op": "set_document_state", "params": {"final": True}}], "untracked")
    assert e.value.code == "conflict"


def test_two_ops_creating_or_numbering_custom_properties_conflict():
    for custom in (None, [("Project Number", "lpwstr", "1001")]):
        with pytest.raises(JobException) as e:
            run_spec(build(custom=custom),
                     [{"op": "set_document_properties", "params": {"custom": {"Site": "VA3"}}},
                      {"op": "set_document_state", "params": {"final": True}}], "untracked")
        assert e.value.code == "conflict"


def test_the_two_ops_together_on_different_properties_verify():
    src = build(track="on", custom=[("Project Number", "lpwstr", "1001"), ("_MarkAsFinal", "bool", "false")])
    plan, out, rec = run_spec(src, [
        {"op": "set_document_properties", "params": {"set": {"title": "T"}, "custom": {"Project Number": "1003"},
                                                     "clear": ["author"]}},
        {"op": "set_document_state", "params": {"final": True, "track_changes": "off"}}], "untracked")
    assert rec.passed, rec.to_markdown()
    assert custom_props(out)["_MarkAsFinal"] == ("3", "bool", "true")


# ---- the two sides' tables agree; the worked examples run ------------------------------------------------------------

def test_the_property_tables_and_the_settings_order_agree():
    assert dp.PROPERTIES == vdp._V_PROPERTIES
    assert tuple(dp.SETTINGS_ORDER) == vdp._settings_order()      # the transform's table is the bundled schema's


@pytest.mark.parametrize("name, kw", [("set-title-clear-authors.json",
                                       {"creator": "Jeremy", "last_modified_by": "Jack"}),
                                      ("mark-final-tracking-off.json", {"track": "on"})])
def test_the_worked_examples_plan_and_verify(name, kw):
    spec = load_spec(EXAMPLES / name)
    src = build(**kw)
    pkg = DocxPackage(src)
    plan, frags = plan_spec(pkg, spec)
    out = apply_spec(pkg, frags).output_bytes()
    rec = run_suite(src, out, plan.to_dict(), "src.docx", "out.docx")
    assert rec.passed, rec.to_markdown()
    assert json.loads((EXAMPLES / name).read_text(encoding="utf-8"))["request"]


def _two_titles(root):
    etree.SubElement(root, DC + "title").text = "second"


def _two_tracks(root):
    root.find(W + "trackRevisions").addnext(etree.Element(W + "trackRevisions"))


@pytest.mark.parametrize("src, op, params", [
    (lambda: patch_docx(build(), {CORE: _two_titles}), "set_document_properties", {"set": {"title": "T"}}),
    (lambda: build(custom=[("Site", "lpwstr", "a"), ("SITE", "lpwstr", "b")]), "set_document_properties",
     {"custom": {"Site": "c"}}),
    (lambda: patch_docx(build(track="on"), {SETTINGS: _two_tracks}), "set_document_state", {"track_changes": "off"}),
])
def test_a_properties_part_word_would_not_write_is_refused(src, op, params):
    with pytest.raises(JobException) as e:
        plan_only(src(), op, **params)
    assert e.value.code == "malformed-properties"


def test_keywords_with_value_elements_are_replaced_by_the_text():
    def kw(root):
        el = root.find(CP + "keywords")
        el.text = "fire, "
        etree.SubElement(el, CP + "value").text = "sprinkler"
    src = patch_docx(build(), {CORE: kw})
    plan, out, rec = props(src, set={"keywords": "wet-pipe; NFPA 13"})
    assert rec.passed, rec.to_markdown()
    el = xml(out, CORE).find(CP + "keywords")
    assert el.text == "wet-pipe; NFPA 13" and len(el) == 0


def test_a_batch_delivers_every_file_with_its_new_properties(tmp_path):
    from spec_surgeon.batch import run_batch
    folder = tmp_path / "in"
    folder.mkdir()
    (folder / "21 13 13.docx").write_bytes(build(track="on"))
    (folder / "21 13 19.docx").write_bytes(build(custom=[("Project Number", "lpwstr", "1001")]))
    spec = {"request": "mark final, tracking off, title set", "summary": "s", "mode": "untracked", "ops": [
        {"op": "set_document_properties", "params": {"set": {"title": "Fire Protection"}}},
        {"op": "set_document_state", "params": {"track_changes": "off", "final": True}}]}
    batch = run_batch([str(folder)], "", {}, tmp_path / "out", workers=1, do_render=False, spec=spec)
    assert [fr.status for fr in batch.files] == ["PASS", "PASS"], [fr.record.to_markdown() for fr in batch.files]
    for fr in batch.files:
        out = Path(fr.delivered_to).read_bytes()
        assert text_of(out, CORE, DC + "title") == "Fire Protection"
        assert custom_props(out)["_MarkAsFinal"][1:] == ("bool", "true")
        assert xml(out, SETTINGS).find(W + "trackRevisions") is None
