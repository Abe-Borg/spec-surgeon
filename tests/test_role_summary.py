"""Role evidence per layout (``structure.role_summary``): how one file encodes each role, and the concrete selector a
planner can write for it — ``style_in`` when the file's styles make the role unambiguous, a ``text_regex`` from the
typed marker shape otherwise, nothing when neither is reliable. The profile carries it as ``roles``."""
from __future__ import annotations

import io
import json
import re

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from make_fixtures import build, build_scheme_example, build_styles, build_typed
from spec_surgeon.ooxml import block_full_text, paragraphs
from spec_surgeon.ops.selectors import StyleSheet, paragraph_selected
from spec_surgeon.package import DocxPackage
from spec_surgeon.profile import MAX_PROFILE_BYTES, bound, profile
from spec_surgeon.structure import analyze_structure, marker_regex, role_summary


def package(document) -> DocxPackage:
    out = io.BytesIO()
    document.save(out)
    return DocxPackage(out.getvalue())


def selected(pkg: DocxPackage, selector) -> int:
    """How many main-body paragraphs the engine's own selector resolution picks with ``selector``."""
    root = pkg.xml(pkg.main_part())
    sheet = StyleSheet(pkg)
    return sum(1 for i, p in enumerate(paragraphs(root))
               if paragraph_selected(selector, p, i, block_full_text(p), sheet))


def numbering(doc):
    """One multilevel list: ``PART %1 -`` / ``%1.%2`` / ``%3.`` (as test_structure builds it)."""
    root = doc.part.numbering_part.element
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), "77")
    for index, (fmt, text) in enumerate((("decimal", "PART %1 -"), ("decimal", "%1.%2"), ("upperLetter", "%3."))):
        level = OxmlElement("w:lvl")
        level.set(qn("w:ilvl"), str(index))
        for name, value in (("start", "1"), ("numFmt", fmt), ("lvlText", text), ("suff", "tab")):
            child = OxmlElement("w:" + name)
            child.set(qn("w:val"), value)
            level.append(child)
        abstract.append(level)
    root.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), "77")
    reference = OxmlElement("w:abstractNumId")
    reference.set(qn("w:val"), "77")
    num.append(reference)
    root.append(num)


def numbered(p, level):
    num = OxmlElement("w:numPr")
    for name, value in (("ilvl", level), ("numId", 77)):
        child = OxmlElement("w:" + name)
        child.set(qn("w:val"), str(value))
        num.append(child)
    p._p.get_or_add_pPr().append(num)


def test_a_typed_normal_only_spec_reports_typed_markers_and_text_regex_selectors():
    pkg = DocxPackage(build("21 13 13", "WET-PIPE SPRINKLER SYSTEMS"))
    summary = role_summary(pkg)
    article = summary["roles"]["article"]
    assert article["count"] == 30 and article["styles"] == {"Normal": 30} and article["numbering"] == "typed"
    assert article["markers"] == ["1.01", "2.01"]                      # as the file carries them, not canonical "1.1"
    assert article["examples"][0] == "1.01\tSUMMARY" and len(article["examples"]) == 2
    assert all(len(e) <= 80 for role in summary["roles"].values() for e in role["examples"])
    # Normal is every role's style, so no role gets style_in; the typed shape gives the regex
    sel = summary["selectors"]["article"]
    assert set(sel) == {"text_regex"} and sel["text_regex"] == r"^\s*\d+\.\d{2}(?:\s|$)"   # 1.01 … 30.01: \d+ then \d{2}
    rows = analyze_structure(pkg)["paragraphs"]
    rx = re.compile(sel["text_regex"])
    assert all(rx.search(r["text"]) for r in rows if r["role"] == "article")
    assert not any(rx.search(r["text"]) for r in rows if r["role"] not in (None, "article"))
    assert selected(pkg, sel) == 30                                    # through the engine's own selector resolution
    assert summary["selectors"]["paragraph"] == {"text_regex": r"^\s*[A-Z]\.(?:\s|$)"}
    assert summary["selectors"]["section_id"] == {"text_regex": r"^\s*SECTION\s+\d{2}\s+\d{2}\s+\d{2}(?:\s|$)"}
    assert "section_title" not in summary["selectors"]                 # no marker, and a style every role shares
    assert summary["roles"]["section_title"]["numbering"] == "none"
    assert summary["unresolved"] == {"count": 2, "reasons": {"no unambiguous supported specification marker": 2}}
    assert list(summary["roles"]) == ["section_id", "section_title", "article", "paragraph"]   # hierarchy order
    assert not pkg.is_modified()


def test_a_styled_spec_reports_style_selectors():
    summary = role_summary(DocxPackage(build_styles()))
    article = summary["roles"]["article"]
    assert article["styles"] == {"ArchArticle": 2} and article["numbering"] == "typed"
    assert summary["selectors"]["article"] == {"style_in": ["ArchArticle"]}
    # PART headings: two in Heading1, one typed in Normal — Normal belongs to other roles, so the typed shape decides
    assert summary["roles"]["part"]["styles"] == {"Heading1": 2, "Normal": 1}
    part = summary["selectors"]["part"]
    assert set(part) == {"text_regex"} and re.search(part["text_regex"], "PART 2 - PRODUCTS")
    assert not re.search(part["text_regex"], "PART ONE - GENERAL")
    assert summary["unresolved"]["count"] == 7


def test_automatic_numbering_reports_rendered_markers_and_source_text_examples():
    summary = role_summary(DocxPackage(build_scheme_example()))
    assert {r: e["numbering"] for r, e in summary["roles"].items()} == {
        "part": "automatic", "article": "automatic", "paragraph": "automatic", "subparagraph": "automatic",
        "subparagraph_2": "automatic"}
    assert summary["roles"]["part"]["markers"] == ["PART 1 -", "PART 2 -"]
    assert summary["roles"]["part"]["examples"] == ["GENERAL", "PRODUCTS"]     # the text a regex would see
    assert summary["selectors"] == {"part": {"style_in": ["PRT"]}, "article": {"style_in": ["ART"]},
                                    "paragraph": {"style_in": ["PR1"]}, "subparagraph": {"style_in": ["PR2"]},
                                    "subparagraph_2": {"style_in": ["PR3"]}}


def test_table_cells_that_share_a_shape_are_fenced_out_of_the_selector():
    for leading_tab in (False, True):
        pkg = DocxPackage(build_typed(leading_tab=leading_tab))
        summary = role_summary(pkg)
        sub = summary["selectors"]["subparagraph"]
        assert sub == {"text_regex": r"^\s*\d\.(?:\s|$)", "in_table": False}   # "1.⇥Cell item." sits in a table
        assert "in_table" not in summary["selectors"]["paragraph"]
        assert selected(pkg, sub) == 3 and selected(pkg, summary["selectors"]["paragraph"]) == 9
        assert selected(pkg, summary["selectors"]["article"]) == 5                # "1." never selects "1.1"
        assert summary["selectors"]["end_of_section"]["text_regex"].startswith(r"^\s*END\s+OF\s+SECTION")


def test_marker_shapes_that_disagree_or_collide_give_no_regex():
    assert marker_regex(["1.01", "10.01"]) == r"^\s*\d+\.\d{2}(?:\s|$)"
    assert marker_regex(["(1)", "(2)"]) == r"^\s*\(\d\)(?:\s|$)"
    assert marker_regex(["a.", "b."]) == r"^\s*[a-z]\.(?:\s|$)"
    assert marker_regex(["PART 1 -", "PART 2"]) is None                # separators disagree
    assert marker_regex(["A.", "AA."]) is None                         # letter runs disagree
    assert marker_regex(["21 13 13", "21-13-13"]) is None               # punctuation disagrees
    assert marker_regex([]) is None


def test_a_style_shared_with_unmarked_wording_or_mixed_numbering_gives_no_style_selector():
    doc = Document()
    doc.styles.add_style("ART", 1)
    for text, style in (("SECTION 21 13 13", None), ("SPRINKLERS", None), ("1.1 SUMMARY", "ART"),
                        ("Unmarked wording in the article style.", "ART"), ("1.2 REFERENCES", "ART")):
        doc.add_paragraph(text, style=style)
    summary = role_summary(package(doc))
    assert summary["roles"]["article"]["styles"] == {"ART": 2}
    assert summary["selectors"]["article"] == {"text_regex": r"^\s*\d\.\d(?:\s|$)"}    # ART also holds prose

    doc = Document()
    numbering(doc)
    doc.add_paragraph("SECTION 21 13 13")
    doc.add_paragraph("SPRINKLERS")
    numbered(doc.add_paragraph("GENERAL"), 0)
    numbered(doc.add_paragraph("SUMMARY"), 1)                          # the list renders 1.1
    doc.add_paragraph("1.2 REFERENCES")                                 # typed
    summary = role_summary(package(doc))
    assert summary["roles"]["article"]["numbering"] == "mixed"
    assert summary["roles"]["article"]["markers"] == ["1.1", "1.2"]
    assert "article" not in summary["selectors"]


def test_a_document_without_structure_yields_empty_roles():
    doc = Document()
    doc.add_paragraph("Hello world.")
    doc.add_paragraph("A second line of prose.")
    summary = role_summary(package(doc))
    assert summary["roles"] == {} and summary["selectors"] == {}
    assert summary["unresolved"] == {"count": 2, "reasons": {"no unambiguous supported specification marker": 2}}
    assert role_summary(package(Document())) == {"roles": {}, "unresolved": {"count": 0, "reasons": {}},
                                                 "selectors": {}}
    assert profile(package(doc))["roles"]["roles"] == {}


def test_the_profile_carries_roles_within_its_budget_and_keeps_the_selectors_longest():
    prof = profile(DocxPackage(build_styles(), name="styled.docx"))
    assert len(json.dumps(prof, ensure_ascii=False)) <= MAX_PROFILE_BYTES
    json.dumps(prof)                                                    # what the compiler sends the model
    assert prof["roles"]["selectors"]["article"] == {"style_in": ["ArchArticle"]}
    assert prof["roles"]["roles"]["article"]["numbering"] == "typed"
    for limit in (4000, 3000, 2000, 1500, 1000, 600, 300):
        cut = bound(json.loads(json.dumps(prof)), limit=limit)
        assert len(json.dumps(cut, ensure_ascii=False)) <= limit, limit
        if "sections" in cut and ("lists" in cut or "tables" in cut):  # roles go after the list and table summaries
            assert "roles" in cut, limit
    small = bound(json.loads(json.dumps(prof)), limit=1500)             # the last resort keeps the selectors
    assert "sections" not in small and small["roles"] == {"selectors": prof["roles"]["selectors"]}
    assert "role selectors" in small["truncated"][-1]
    tiny = bound(json.loads(json.dumps(prof)), limit=300)
    assert "roles" not in tiny


def test_the_profile_records_a_summary_failure_instead_of_raising(monkeypatch):
    import spec_surgeon.profile as profile_module
    monkeypatch.setattr(profile_module, "role_summary", lambda pkg: (_ for _ in ()).throw(RuntimeError("boom")))
    prof = profile(DocxPackage(build("21 13 13", "WET-PIPE SPRINKLER SYSTEMS"), name="spec.docx"))
    assert prof["roles"] == {"error": "RuntimeError: boom"}
    assert prof["counts"]["paragraphs"] > 0                            # the file still profiles
