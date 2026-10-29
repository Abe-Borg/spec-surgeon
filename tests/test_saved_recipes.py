"""Saved recipes (``saved_recipes``): the runs of What you've done so far, saved under a name and run again on other
documents as ONE plan — one preview, one approval, one checked run, no AI call. Safe mode keeps each button and its
form; Default and Obliterate keep the user's words and the plan they approved. What a recipe cannot repeat (a custom
method, values read from particular documents) is refused with the reason; a recipe runs in its own mode and in any with
fewer guardrails; the window's journal keeps what each run would save."""
from __future__ import annotations

import json
import threading
from pathlib import Path

import pytest
from docx import Document

from fake_claude import FakeClient, resp, tool
from make_fixtures import build
from test_document_task import REQUEST, SPEC, document
from test_recipes import font_fixtures, no_ai
from spec_surgeon import layouts, saved_recipes
from spec_surgeon.document_task import DocumentTask
from spec_surgeon.gui import Session
from spec_surgeon.recovery_state import session_lock
from spec_surgeon.saved_recipes import RecipeRunError, Store

TITLE = {"pattern": "SECTION {number} - {title}"}
WORDS = "Replace “Keep this” with “Retain this”"
REPLACE = {"request": WORDS, "summary": "Replace the words", "mode": "untracked",
           "ops": [{"op": "replace_text", "select": {"scope": ["document"]},
                    "params": {"find": "Keep this", "replace": "Retain this", "tracked": False}}]}


def specs(where: Path, names) -> list:
    where.mkdir(parents=True, exist_ok=True)
    out = []
    for num, title in names:
        path = where / f"{num}.docx"
        path.write_bytes(build(num, title))
        out.append(path)
    return out


@pytest.fixture
def store(tmp_path) -> Store:
    return Store(tmp_path / "config" / "recipes.json")


def safe(tmp_path, store, name="out") -> DocumentTask:
    return DocumentTask(Session(out_root=tmp_path / name, visual_check=False), render=False, mode="safe", recipes=store)


def press(task, recipe, inputs=None):
    view = task.perform("recipe", recipe=recipe, inputs=inputs or {})
    assert view["stage"] == "preview", [m["text"] for m in task.messages][-2:]
    return task.perform("primary", token=view["token"])


def done(task) -> list:
    return task.view()["working"]["done"]


# ---- the store ------------------------------------------------------------------------------------------------------

STEP = saved_recipes.button_step("remove_comments", {}, "Remove all comments")


def test_the_store_keeps_recipes_in_one_plain_json_file_written_atomically(store):
    record = store.add("  Clean   copies ", "for bids", [STEP])
    assert record["name"] == "Clean copies" and record["mode"] == "safe" and record["last_run"] == ""
    data = json.loads(store.path.read_text(encoding="utf-8"))
    assert data["format"] == saved_recipes.FORMAT and [r["id"] for r in data["recipes"]] == [record["id"]]
    assert not store.path.with_name("recipes.json.tmp").exists()
    store.rename(record["id"], "Clean", "for permits")
    copy = store.duplicate(record["id"])
    assert [r["name"] for r in store.all()] == ["Clean", "Clean (copy)"] and copy["note"] == "for permits"
    store.delete(record["id"])
    assert [r["name"] for r in Store(store.path).all()] == ["Clean (copy)"]
    with pytest.raises(RecipeRunError, match="no such recipe"):
        store.delete(record["id"])


@pytest.mark.parametrize("name, why", [("", "Type a name"), ("x" * 121, "at most 120")])
def test_a_recipe_needs_a_name_and_a_step(store, name, why):
    with pytest.raises(RecipeRunError, match=why):
        store.add(name, "", [STEP])
    with pytest.raises(RecipeRunError, match="at least one"):
        store.add("Name", "", [])


def test_two_windows_read_each_others_saves_before_they_change_the_file(store):
    other = Store(store.path)
    store.add("First", "", [STEP])
    other.add("Second", "", [STEP])                   # read again before writing: the first window's save survives
    assert [r["name"] for r in store.all()] == ["First", "Second"]


def test_a_save_another_window_makes_mid_change_waits_for_the_lock_and_neither_is_lost(store):
    other, write, saving = Store(store.path), store._write, []

    def mid_change(recipes):
        if not saving:                            # the other window saves between this one's read and its write
            saving.append(threading.Thread(target=other.add, args=("Second", "", [STEP])))
            saving[0].start()
            saving[0].join(0.3)
            assert saving[0].is_alive()           # it waits for the lock instead of writing first
        write(recipes)
    store._write = mid_change
    store.add("First", "", [STEP])
    saving[0].join(5)
    assert [r["name"] for r in Store(store.path).all()] == ["First", "Second"]


def test_an_unreadable_file_is_kept_aside_and_a_recipe_that_does_not_read_is_skipped(store):
    store.path.parent.mkdir(parents=True)
    store.path.write_text("{not json", encoding="utf-8")
    assert store.all() == [] and store.path.with_name("recipes.json.bad").is_file()
    good = {"id": "r1", "name": "Good", "steps": [STEP]}
    store.path.write_text(json.dumps({"format": 1, "recipes": [good, {"id": "r2", "name": "", "steps": [STEP]},
                                                                {"id": "r3", "name": "No steps", "steps": []}]}),
                          encoding="utf-8")
    assert [r["id"] for r in Store(store.path).all()] == ["r1"]


def test_a_safe_recipe_takes_the_forms_its_run_form_holds(store):
    record = store.add("Stamp", "", [saved_recipes.button_step("title_property", TITLE, "x"), STEP])
    changed = store.update_inputs(record["id"], [{"pattern": "{number}"}, None])
    assert changed["steps"][0]["inputs"] == {"pattern": "{number}"} and changed["steps"][1] == STEP
    with pytest.raises(RecipeRunError, match="changed since"):
        store.update_inputs(record["id"], [None])
    with pytest.raises(RecipeRunError, match="Only a button"):
        plan = saved_recipes.plan_step(WORDS, REPLACE, [], False)
        other = store.add("Words", "", [plan])
        store.update_inputs(other["id"], [{"x": 1}])


# ---- what a recipe can keep -----------------------------------------------------------------------------------------

def test_a_plan_keeps_no_unticked_occurrences_and_refuses_what_is_bound_to_particular_documents():
    table = {"request": "r", "summary": "s", "mode": "untracked",
             "ops": [{"op": "replace_table", "select": {"scope": ["document"]},
                      "params": {"pairs": [{"find": "a", "replace": "b"}], "skip": [{"file": "x.docx"}]}}]}
    kept, why = saved_recipes.portable_plan(table)
    assert why == "" and "skip" not in kept["ops"][0]["params"] and "skip" in table["ops"][0]["params"]
    bound = {**SPEC, "bindings": [{"file": "x.docx", "md5": "0" * 32, "slots": {}}]}
    assert saved_recipes.portable_plan(bound) == (None, saved_recipes.WHY_BOUND)
    typed = {**SPEC, "ops": [{"op": "apply_header_footer_template", "params": {"slot_fills": {
        "NAME": {"rule": "first_match", "rules": [{"rule": "typed", "value": "v", "source_md5": "0" * 32}]}}}}]}
    assert saved_recipes.portable_plan(typed) == (None, saved_recipes.WHY_BOUND)
    capitals = {**SPEC, "ops": [{"op": "capitalize_defined_terms", "params": {"approvals": []}}]}
    assert saved_recipes.portable_plan(capitals) == (None, saved_recipes.WHY_CAPITALS)
    chain = {"stages": [{"spec": SPEC}, {"spec": bound}]}
    assert saved_recipes.portable_plan(chain)[1] == saved_recipes.WHY_BOUND


def test_a_recipe_runs_in_its_own_mode_and_in_any_with_fewer_guardrails():
    button, plan = STEP, saved_recipes.plan_step(WORDS, REPLACE, [], False)
    unguarded = saved_recipes.plan_step(WORDS, REPLACE, [], True)
    assert [saved_recipes.recipe_mode(s) for s in ([button], [button, plan], [plan, unguarded])] == \
        ["safe", "default", "obliterate"]
    assert plan["places"] == ["Retain this"]
    runs = {(m, mode): saved_recipes.runs_in(m, mode) for m in saved_recipes.MODES for mode in saved_recipes.MODES}
    assert [k for k, v in runs.items() if v] == [("safe", "safe"), ("safe", "default"), ("safe", "obliterate"),
                                                 ("default", "default"), ("default", "obliterate"),
                                                 ("obliterate", "obliterate")]
    assert "Safe mode runs only recipes made of its own buttons" in saved_recipes.not_here("default", "safe")
    assert "only there" in saved_recipes.not_here("obliterate", "default")


# ---- Safe mode ------------------------------------------------------------------------------------------------------

def test_safe_buttons_saved_as_a_recipe_run_on_other_documents_as_one_plan(tmp_path, store, monkeypatch):
    no_ai(monkeypatch)
    first = safe(tmp_path, store)
    first.attach(specs(tmp_path / "a", (("21 13 13", "WET-PIPE SPRINKLER SYSTEMS"),)))
    press(first, "remove_comments")
    press(first, "accept_changes", {"action": "accept"})
    press(first, "title_property", TITLE)
    runs = done(first)
    assert [r["buttons"] for r in runs] == [["remove_comments"], ["accept_changes"], ["title_property"]]
    assert all(r["why_not"] == "" and r["steps"][0]["kind"] == "button" for r in runs)
    record = first.save_recipe("Clean and title", "bid set", [r["id"] for r in runs])
    assert record["mode"] == "safe" and [s["button"] for s in record["steps"]] == \
        ["remove_comments", "accept_changes", "title_property"]
    assert record["steps"][2]["inputs"] == TITLE
    assert first.messages[-1]["text"].startswith("Saved the recipe “Clean and title”: 3 steps.")

    others = specs(tmp_path / "b", (("21 05 00", "COMMON WORK RESULTS"), ("21 31 13", "FIRE PUMPS")))
    before = {p: p.read_bytes() for p in others}
    task = safe(tmp_path, store, "out2")
    task.attach(others)
    view = task.perform("saved_recipe", recipe=record["id"], steps=None)
    assert view["stage"] == "preview" and task.session.model_free and task.session.outcome_source == "saved"
    spec = task.session.spec
    assert [st["name"] for st in spec["stages"]] == ["Remove all comments", "Accept or reject all tracked changes",
                                                     "Set the document title property"]
    assert [m["text"] for m in task.messages][-2:] == [
        "Run the recipe “Clean and title”",
        "Its 3 steps, in order, planned without AI: 1. Remove all comments; 2. Accept all tracked changes; "
        "3. Set the document title property to “SECTION {number} - {title}”."]
    view = task.perform("primary", token=view["token"])
    assert view["stage"] == "done" and len(task.results) == 2 and task._budget.calls == 0
    assert {p: p.read_bytes() for p in before} == before                 # the originals are never written
    [run] = done(task)
    assert run["name"] == "Clean and title" and run["request"] == "Run the recipe “Clean and title”"
    assert run["buttons"] == ["remove_comments", "accept_changes", "title_property"] and run["documents"] == 2
    assert store.get(record["id"])["last_run"] and store.get(record["id"])["last_documents"] == 2
    # saving the recipe's run again keeps its steps, not the run as one step
    again = task.save_recipe("Again", "", [run["id"]])
    assert [s["button"] for s in again["steps"]] == ["remove_comments", "accept_changes", "title_property"]
    task.undo()                                                          # one run: one undo takes the whole recipe off
    assert done(task) == [] and task.working == {}


def test_a_step_can_be_skipped_or_run_with_another_form_this_time(tmp_path, store):
    record = store.add("Two", "", [STEP, saved_recipes.button_step("title_property", TITLE, "x")])
    task = safe(tmp_path, store)
    task.attach(specs(tmp_path / "b", (("21 05 00", "COMMON WORK RESULTS"),)))
    view = task.perform("saved_recipe", recipe=record["id"],
                        steps=[{"skip": True}, {"skip": False, "inputs": {"pattern": "{number} {title}"}}])
    assert view["stage"] == "preview"
    assert [o["op"] for o in task.session.spec["ops"]] == ["set_document_properties"]   # one step: no chain
    assert "Skipped this time: 1." in task.messages[-1]["text"]
    task.perform("primary", token=view["token"])
    [run] = done(task)
    assert [s["inputs"] for s in run["steps"]] == [{"pattern": "{number} {title}"}]   # the run keeps what it ran
    assert store.get(record["id"])["steps"][1]["inputs"] == TITLE                       # the recipe as it was saved


@pytest.mark.parametrize("picks, said", [([{"skip": True}], "Every step is skipped"),
                                         ([{}, {}], "changed since its form was opened")])
def test_a_run_form_that_does_not_fit_the_recipe_is_refused(tmp_path, store, picks, said):
    record = store.add("One", "", [STEP])
    task = safe(tmp_path, store)
    task.attach(specs(tmp_path / "b", (("21 05 00", "COMMON WORK RESULTS"),)))
    task.perform("saved_recipe", recipe=record["id"], steps=picks)
    assert task.stage == "empty" and said in task.messages[-1]["text"] and task.session.spec is None


def test_a_step_that_needs_a_template_says_so_and_nothing_is_previewed(tmp_path, store):
    record = store.add("Headers", "", [saved_recipes.button_step("headers_footers_template", {"slots": {}}, "x")])
    task = safe(tmp_path, store)
    task.attach(specs(tmp_path / "b", (("21 05 00", "COMMON WORK RESULTS"),)))
    task.perform("saved_recipe", recipe=record["id"], steps=None)
    assert task.messages[-1]["text"] == "Headers: Step 1, Match headers and footers to template: Attach your template first."
    assert task.stage == "empty" and task.session.spec is None


def test_a_step_that_plans_per_layout_makes_the_recipe_a_plan_per_layout_of_chains(tmp_path, store, monkeypatch):
    no_ai(monkeypatch)
    nine, eleven, _ = font_fixtures(tmp_path)
    record = store.add("Fonts", "", [STEP, saved_recipes.button_step("set_font", {"font": "Arial", "size": "10"}, "x"),
                                     saved_recipes.button_step("remove_highlight", {}, "Remove highlighting")])
    task = safe(tmp_path, store)
    task.attach([nine, eleven])
    view = task.perform("saved_recipe", recipe=record["id"], steps=None)
    plan = task.session.spec
    assert view["stage"] == "preview" and layouts.is_layout_plan(plan) and len(plan["layouts"]) == 2
    for layout in plan["layouts"]:
        assert [st["name"] for st in layout["spec"]["stages"]] == ["Remove all comments", "Change font",
                                                                   "Remove highlighting"]
        assert all("applies_to" not in st["spec"] for st in layout["spec"]["stages"])   # pinned again per layout
    task.perform("primary", token=view["token"])
    assert len(task.results) == 2 and {f.status for f in task.session.last_batch.files} == {"PASS"}


def test_safe_mode_lists_default_recipes_only_to_say_they_run_elsewhere(tmp_path, store):
    record = store.add("Words", "", [saved_recipes.plan_step(WORDS, REPLACE, [], False)])
    task = safe(tmp_path, store)
    [shown] = task.view()["saved_recipes"]
    assert shown["mode"] == "default" and not shown["runs_here"]
    task.attach([document(tmp_path / "Spec.docx")])
    with pytest.raises(ValueError, match="Safe mode runs only recipes made of its own buttons"):
        task.perform("saved_recipe", recipe=record["id"], steps=None)


# ---- Default and Obliterate -----------------------------------------------------------------------------------------

def chat(tmp_path, store, *replies, mode="default", name="out"):
    task = DocumentTask(Session(client=FakeClient(*replies) if replies else None, out_root=tmp_path / name,
                                visual_check=False), render=False, obliterate_allowed=True, recipes=store)
    if mode == "obliterate":
        task.set_mode("obliterate", acknowledged=True)
    return task


def plan(spec, tid="t1", message=None):
    return resp(tool("submit_editspec", {"spec": spec, **({"message": message} if message else {})}, tid))


def test_requests_saved_as_a_recipe_run_again_with_no_ai_call_and_a_document_that_does_not_fit_waits_on_the_list(
        tmp_path, store):
    first = chat(tmp_path, store, plan(SPEC), plan(REPLACE, "t2"))
    first.attach([document(tmp_path / "a" / "Spec.docx")])
    first.perform("primary", token=first.perform("send", text=REQUEST)["token"])
    first.perform("primary", token=first.perform("send", text=WORDS)["token"])
    runs = done(first)
    assert [r["steps"][0]["kind"] for r in runs] == ["plan", "plan"]
    assert runs[1]["steps"][0]["places"] == ["Retain this"] and runs[1]["steps"][0]["plan"]
    record = first.save_recipe("Bold and retain", "", [r["id"] for r in runs])
    assert record["mode"] == "default"
    assert record["steps"][0]["plan"]["ops"] == SPEC["ops"] and record["steps"][0]["request"] == REQUEST

    task = chat(tmp_path, store, name="out2")                          # no client at all: nothing can call a model
    fits = document(tmp_path / "b" / "Other.docx")
    other = tmp_path / "b" / "Untitled.docx"
    other.write_bytes(build("21 31 13", "FIRE PUMPS"))                 # no FIRE SUPPRESSION title, no “Keep this”
    task.attach([fits, other])
    view = task.perform("saved_recipe", recipe=record["id"], steps=None)
    assert view["stage"] == "preview" and task._budget.calls == 0
    assert task.messages[-1]["text"].endswith("Words it places: “Retain this”.")
    view = task.perform("primary", token=view["token"])
    assert view["stage"] == "done" and list(task.results) == [str(fits)] and task._budget.calls == 0
    assert list(task.todo) == [str(other)]                              # Default mode: the to-do list, not beside
    assert task.left_untouched == []
    [run] = done(task)
    assert run["name"] == "Bold and retain" and [s["text"] for s in run["steps"]] == [REQUEST, WORDS]


def test_a_step_that_changes_what_an_earlier_step_made_is_planned_on_the_copies_that_step_makes(tmp_path, store):
    """Runs that built on each other in the window are saved together: the second step finds words only the first one
    places, so it plans on no original and is planned as the step after the first, as the run that made it was."""
    again = "Replace “Retain this” with “Hold this”"
    then = {"request": again, "summary": "Replace the words again", "mode": "untracked",
            "ops": [{"op": "replace_text", "select": {"scope": ["document"]},
                     "params": {"find": "Retain this", "replace": "Hold this", "tracked": False}}]}
    record = store.add("Retain, then hold", "", [saved_recipes.plan_step(WORDS, REPLACE, [], False),
                                                 saved_recipes.plan_step(again, then, [], False)])
    task = chat(tmp_path, store)
    spec = document(tmp_path / "Spec.docx")
    task.attach([spec])
    view = task.perform("saved_recipe", recipe=record["id"], steps=None)
    assert view["stage"] == "preview", task.messages[-1]["text"]
    view = task.perform("primary", token=view["token"])
    assert view["stage"] == "done" and task._budget.calls == 0
    assert Document(task.results[str(spec)]).paragraphs[2].text == "Hold this technical wording intact."

    pairs = [{"find": "Keep this", "replace": "Retain this"}], [{"find": "Retain this", "replace": "Hold this"}]
    buttons = store.add("Two passes", "", [saved_recipes.button_step("find_replace", {"pairs": p}, "") for p in pairs])
    task = safe(tmp_path, store, name="out3")      # Safe mode's find and replace takes a pair it finds nowhere; its
    spec = document(tmp_path / "c" / "Spec.docx")  # passes still run in the recipe's order
    task.attach([spec])
    view = task.perform("saved_recipe", recipe=buttons["id"], steps=None)
    assert view["stage"] == "preview", task.messages[-1]["text"]
    task.perform("primary", token=view["token"])
    assert Document(task.results[str(spec)]).paragraphs[2].text == "Hold this technical wording intact."

    alone = store.add("Hold only", "", [saved_recipes.plan_step(again, then, [], False)])
    task = chat(tmp_path, store, name="out2")
    task.attach([document(tmp_path / "b" / "Spec.docx")])
    task.perform("saved_recipe", recipe=alone["id"], steps=None)          # no earlier step: refused as before
    assert task.stage == "empty" and "Step 1 cannot be used on these documents" in task.messages[-1]["text"]


def test_a_default_recipe_waits_for_an_open_request_and_is_never_handed_to_claude(tmp_path, store):
    record = store.add("Words", "", [saved_recipes.plan_step(WORDS, REPLACE, [], False)])
    task = chat(tmp_path, store, resp(tool("ask_user", {"questions": ["Which title should change?"]})))
    task.attach([document(tmp_path / "Spec.docx")])
    task.perform("send", text=REQUEST)
    assert task.view()["recipe_blocked"] == "Drop the request that is open first, or finish it."
    with pytest.raises(ValueError, match="Drop the request"):
        task.perform("saved_recipe", recipe=record["id"], steps=None)
    task.discard_preview()
    view = task.perform("saved_recipe", recipe=record["id"], steps=None)
    assert view["stage"] == "preview" and view["recipe_blocked"] == "Discard the preview first, or create its copies."
    task.attach([document(tmp_path / "More.docx")])                     # the preview no longer applies
    assert task.decisions == [] and task.stage == "empty"               # nothing left for Claude to plan
    assert task.messages[-1]["text"].endswith("Run the recipe again to preview it on these documents.")


def test_a_custom_method_or_a_lost_plan_cannot_be_saved_and_says_why(tmp_path, store, monkeypatch):
    task = chat(tmp_path, store)
    monkeypatch.setattr(task.session, "generated", object())
    assert task._steps_now()["why_not"] == saved_recipes.WHY_CUSTOM
    monkeypatch.setattr(task.session, "generated", None)
    assert task._steps_now()["why_not"] == saved_recipes.WHY_NOT_KEPT
    task.runs = [{"id": "x", "request": "r", "recipe": "", "documents": [], "not": {}}]
    task.run_steps = {"x": {"steps": [], "name": "", "recipe_id": "", "why_not": saved_recipes.WHY_CUSTOM}}
    [row] = task._done_view()
    assert row["why_not"] == saved_recipes.WHY_CUSTOM and row["steps"] == []
    with pytest.raises(ValueError, match="custom method"):
        task.save_recipe("Name", "", ["x"])


def test_an_obliterate_recipe_carries_claudes_words_and_runs_only_in_obliterate_mode(tmp_path, store):
    shown = "I will put “Retain this” in place of “Keep this”."
    first = chat(tmp_path, store, plan({**REPLACE, "request": "Tidy that sentence"}, message=shown), mode="obliterate")
    first.attach([document(tmp_path / "a" / "Spec.docx")])
    first.perform("primary", token=first.perform("send", text="Tidy that sentence")["token"])
    [run] = done(first)
    assert run["steps"][0]["unguarded"] is True and run["steps"][0]["places"] == ["Retain this"]
    record = first.save_recipe("Tidy", "", [run["id"]])
    assert record["mode"] == "obliterate"

    hidden = DocumentTask(Session(out_root=tmp_path / "o3", visual_check=False), render=False, recipes=store)
    assert hidden.view()["saved_recipes"] == []                          # Obliterate hidden: its recipes are too
    default = chat(tmp_path, store, name="out2")
    [listed] = default.view()["saved_recipes"]
    assert not listed["runs_here"] and listed["why_not_here"] == "It was saved in Obliterate mode and runs only there."
    default.attach([document(tmp_path / "b" / "Other.docx")])
    with pytest.raises(ValueError, match="runs only there"):
        default.perform("saved_recipe", recipe=record["id"], steps=None)

    task = chat(tmp_path, store, mode="obliterate", name="out4")
    task.attach([document(tmp_path / "c" / "Other.docx")])
    view = task.perform("saved_recipe", recipe=record["id"], steps=None)
    assert view["stage"] == "preview"
    assert task.messages[-1]["text"].endswith(
        "Words it places: “Retain this”. It was saved in Obliterate mode, so some of them may be Claude's.")
    assert "Retain this" not in task.word_sources                       # Claude's words never become the user's


# ---- the journal ----------------------------------------------------------------------------------------------------

def test_what_each_run_would_save_survives_a_restart_and_a_waiting_recipe_is_not_resumed_through_claude(tmp_path, store):
    journal = tmp_path / "conversation.json"
    task = DocumentTask(Session(out_root=tmp_path / "out", visual_check=False), render=False, mode="safe",
                        recipes=store, journal=journal)
    task.attach(specs(tmp_path / "a", (("21 13 13", "WET-PIPE SPRINKLER SYSTEMS"),)))
    press(task, "remove_comments")
    record = store.add("One", "", [saved_recipes.button_step("remove_highlight", {}, "Remove highlighting")])
    task.set_mode("default")
    view = task.perform("saved_recipe", recipe=record["id"], steps=None)
    assert view["stage"] == "preview"
    saved = json.loads(journal.read_text(encoding="utf-8"))
    assert saved["recipe_run"] is True and list(saved["run_steps"]) == [r["id"] for r in task.runs]

    again = DocumentTask(Session(out_root=tmp_path / "out", visual_check=False), render=False, recipes=store,
                         journal=journal)
    assert again.restore()
    # pressed again, never compiled by Claude; it spent nothing, so no go-ahead is asked
    assert again.decisions == [] and again.stage == "empty" and again._request_budget.exhausted_reason is None
    [row] = again._done_view()
    assert row["buttons"] == ["remove_comments"] and row["why_not"] == ""


# ---- the window's routes --------------------------------------------------------------------------------------------

def test_the_routes_save_change_and_run_a_recipe(tmp_path, store, monkeypatch):
    from fastapi.testclient import TestClient
    from spec_surgeon.settings import Settings
    from spec_surgeon.web.server import TOKEN_HEADER, WindowState, create_app
    from test_web_server import TOKEN, FakeUpdates, frames
    no_ai(monkeypatch)
    task = safe(tmp_path, store)
    task.attach(specs(tmp_path / "a", (("21 13 13", "WET-PIPE SPRINKLER SYSTEMS"),)))
    press(task, "remove_comments")
    state = WindowState(task, Settings(), token=TOKEN, updates=FakeUpdates(), version="1.6.0")
    client = TestClient(create_app(state), headers={TOKEN_HEADER: TOKEN})
    run_id = state.task.runs[0]["id"]
    assert client.post("/api/saved-recipes/save", json={"name": "", "runs": [run_id]}).json() == {
        "error": "Type a name for the recipe."}
    view = client.post("/api/saved-recipes/save", json={"name": "Clean", "note": "", "runs": [run_id]}).json()
    [recipe] = view["saved_recipes"]
    assert recipe["name"] == "Clean" and recipe["runs_here"] and recipe["steps"][0]["button"] == "remove_comments"
    view = client.post("/api/saved-recipes/change", json={"action": "rename", "id": recipe["id"], "name": "Tidy"}).json()
    assert [r["name"] for r in view["saved_recipes"]] == ["Tidy"]
    view = client.post("/api/saved-recipes/change", json={"action": "duplicate", "id": recipe["id"]}).json()
    assert [r["name"] for r in view["saved_recipes"]] == ["Tidy", "Tidy (copy)"]
    with client.stream("POST", "/api/saved-recipes/run", json={"id": recipe["id"], "steps": [{"skip": False}]}) as r:
        last = frames(r)[-1]
    assert last["type"] == "done" and last["state"]["stage"] == "preview"
    assert client.post("/api/saved-recipes/run", json={"id": recipe["id"], "steps": "all"}).json() == {
        "error": "The recipe's steps could not be read."}
    view = client.post("/api/saved-recipes/change", json={"action": "delete", "id": recipe["id"]}).json()
    assert [r["name"] for r in view["saved_recipes"]] == ["Tidy (copy)"]
    assert client.post("/api/saved-recipes/change", json={"action": "melt", "id": "x"}).json() == {
        "error": "Unknown recipe change."}


def test_a_recipe_file_edited_to_place_words_the_user_never_gave_is_refused_by_the_gate(tmp_path, store):
    """recipes.json is plain JSON the user can edit: a plan step whose words were never traced (``places`` emptied,
    the replacement changed) still passes the one gate, which refuses words outside the step's own sources."""
    step = saved_recipes.plan_step(WORDS, REPLACE, [], False)
    step["plan"]["ops"][0]["params"]["replace"] = "Words nobody approved"
    step["places"] = []
    record = store.add("Tampered", "", [step])
    task = chat(tmp_path, store)
    task.attach([document(tmp_path / "Spec.docx")])
    task.perform("saved_recipe", recipe=record["id"], steps=None)
    assert task.stage == "empty" and task.session.spec is None
    assert "Step 1 cannot be used on these documents" in task.messages[-1]["text"]
    assert "Words nobody approved" in task.messages[-1]["text"]


def test_a_file_held_by_another_program_is_never_set_aside_and_a_failed_write_is_said(store, monkeypatch):
    record = store.add("Kept", "", [STEP])
    held = Store(store.path)

    def locked(*a, **k):
        raise PermissionError(13, "The process cannot access the file")
    monkeypatch.setattr(Path, "read_text", locked)
    assert held.all() == []                                       # nothing read yet: nothing shown, nothing moved
    with pytest.raises(RecipeRunError, match="could not be read"):
        held.add("Another", "", [STEP])                           # a change never writes the list as last read
    monkeypatch.undo()
    assert [r["id"] for r in held.all()] == [record["id"]] and store.path.is_file()
    assert not store.path.with_name("recipes.json.bad").exists()
    monkeypatch.setattr(saved_recipes, "LOCK_WAIT_S", 0.05)
    with session_lock(store.path.parent, filename="recipes.json.lock"):    # another window mid-change
        with pytest.raises(RecipeRunError, match="Another window is changing the recipes"):
            held.add("Another", "", [STEP])
    monkeypatch.setattr("os.replace", locked)
    with pytest.raises(RecipeRunError, match="could not be saved"):
        held.add("Another", "", [STEP])
    assert [r["name"] for r in Store(store.path).all()] == ["Kept"]
