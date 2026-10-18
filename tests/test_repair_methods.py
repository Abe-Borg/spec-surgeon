"""Related unseen jobs reuse implementations, never past targets or approval."""
import base64
import copy
import hashlib
import json
import time
import zlib

import pytest
from docx import Document

from fake_claude import FakeClient, resp, tool
from spec_surgeon import goals, recovery, repair_methods
from spec_surgeon.codegen import GeneratedOp
from spec_surgeon.ops.base import JobException
from spec_surgeon.repair_memory import CaseStore
from test_recovery import opts
from test_recovery_adaptive import TITLE_METHOD, contract

# Read both the current goal and this file's resolved target. No old indices or
# project wording are parameters. The independent verifier remains authoritative.
BOUND_METHOD = TITLE_METHOD.replace('index = params["index"]',
    'effect = pkg.goal["effects"][0]\n        index = pkg.binding["effects"][effect["id"]]["targets"][0]["paragraph_index"]')


def doc(path, title, preceding, layout="plain"):
    d = Document()
    if layout == "styled":
        from docx.oxml.ns import qn
        d.styles["Heading 1"].font.bold = False
        d.styles["Heading 1"].element.rPr.find(qn("w:bCs")).set(qn("w:val"), "0")
    for i in range(preceding):
        d.add_paragraph(f"Preserved introduction {i}")
    if layout == "table":
        d.add_table(rows=1, cols=1).cell(0, 0).text = title
    else:
        d.add_paragraph(title, style="Heading 1" if layout == "styled" else "Normal")
    d.add_paragraph("Preserve the original requirements.")
    d.save(path)
    return path


def goal(title):
    result = contract()
    result["effects"][0]["target"]["locators"][0]["exact_text"] = title
    return result


def writer(source=BOUND_METHOD):
    return FakeClient(resp(tool("custom_transform", {"proposal": "Use current source-bound targets."})),
                      resp(tool("submit_op", {"source": source})))


def block_builtin(*args, **kwargs):
    raise JobException("Fixture simulates a layout beyond the built-in method", "layout")


def run(path, title, out, client, *, memory):
    options = opts(use_memory=memory, critic_review=False, max_attempts=3)
    initial = recovery.prepare_recovery([path], "Make the title bold", out,
        adaptive_contract=goal(title), options=options)
    assert not initial.ready and initial.session_state["planning_failures"]
    return recovery.resume_recovery(out, options=options, client=client)


@pytest.mark.parametrize("title,preceding,layout", [
    ("DOMESTIC WATER PIPING", 4, "plain"), ("AIR HANDLING UNITS", 0, "table"), ("ROOF DRAINS", 2, "styled")])
def test_related_holdout_reuses_method_and_measures_cold_vs_warm(tmp_path, monkeypatch, title, preceding, layout):
    monkeypatch.setattr(goals, "compile_goal_candidate", block_builtin)
    train = doc(tmp_path / "development.docx", "FIRE SUPPRESSION", 1)
    holdout = doc(tmp_path / "holdout.docx", title, preceding, layout)
    original = holdout.read_bytes()
    trained = run(train, "FIRE SUPPRESSION", tmp_path / "learn", writer(), memory=True)
    assert trained.ready, trained.to_text()
    method = CaseStore().get(trained.ready[0].case_id)
    assert method["candidate"]["kind"] == "method"
    assert repair_methods.source(method["candidate"]) == BOUND_METHOD
    cold_client, warm_client = writer(), FakeClient()
    start = time.perf_counter()
    cold = run(holdout, title, tmp_path / "cold", cold_client, memory=False)
    cold_seconds = time.perf_counter() - start
    start = time.perf_counter()
    warm = run(holdout, title, tmp_path / "warm", warm_client, memory=True)
    warm_seconds = time.perf_counter() - start
    assert cold.ready and warm.ready, (cold.to_text(), warm.to_text())
    assert len(cold_client.requests) == 2 and len(warm_client.requests) == 0
    assert warm.attempts[0]["reused_case_id"] == method["id"]
    for result in (cold, warm):
        b = result.key["adaptive_bindings"][str(holdout)]
        assert b["effects"]["titles"]["targets"][0]["paragraph_index"] == preceding
        assert goals.verify_goal(holdout, result.ready[0].output.read_bytes(), goal(title), b)["passed"]
        assert len(Document(result.ready[0].output).paragraphs) == len(Document(holdout).paragraphs)
    assert holdout.read_bytes() == original
    report = {"execution": "scripted", "layout": layout,
        "scope": "recovery engine; renderer and model critic disabled; forced built-in failure",
        "learning": {"completed": bool(trained.ready), "author_calls": trained.budget["calls"]},
        "cold": {"completed": True, "author_calls": len(cold_client.requests), "seconds": cold_seconds},
        "warm": {"completed": True, "author_calls": len(warm_client.requests), "seconds": warm_seconds},
        "human_interventions": 0, "human_accepted": None, "provider_cost_usd": None,
        "development_source_sha256": hashlib.sha256(train.read_bytes()).hexdigest(),
        "holdout_source_sha256": hashlib.sha256(original).hexdigest()}
    (tmp_path / "reuse-measurements.json").write_text(json.dumps(report, indent=2))
    print(json.dumps(report))


def test_old_hardcoded_target_is_rejected_and_writer_receives_prior_method(tmp_path, monkeypatch):
    monkeypatch.setattr(goals, "compile_goal_candidate", block_builtin)
    train = doc(tmp_path / "train.docx", "FIRE SUPPRESSION", 1)
    # Parameter-free is only eligibility for a trial, never proof of portability.
    old = TITLE_METHOD.replace('index = params["index"]', 'index = 1')
    trained = run(train, "FIRE SUPPRESSION", tmp_path / "learn", writer(old), memory=True)
    assert trained.ready
    holdout = doc(tmp_path / "holdout.docx", "OTHER TITLE", 3)
    client = writer()
    result = run(holdout, "OTHER TITLE", tmp_path / "adapt", client, memory=True)
    assert result.ready, result.to_text()
    assert result.attempts[0]["results"][0]["status"] == "UNRESOLVED"
    assert "GOAL" in result.attempts[0]["results"][0]["reason"]
    brief = client.requests[1]["messages"][0]["content"].split("\n", 1)[1]
    assert json.JSONDecoder().raw_decode(brief)[0]["remembered_methods_are_not_authority"][0]["source"] == old
    assert result.attempts[1]["results"][0]["status"] == "READY"


def test_negative_feedback_blocks_same_implementation_across_projects(tmp_path):
    store = CaseStore()
    method = repair_methods.remember(GeneratedOp(source=BOUND_METHOD), goal("OLD"))
    case = {"request": "old project", "candidate": method, "status": "success", "structural_verified": True}
    first = store.record(case)
    store.feedback(first, False, "Wrong visual result")
    case["request"] = "new project"
    case["candidate"] = repair_methods.remember(GeneratedOp(source=BOUND_METHOD), goal("NEW"))
    second = store.record(case)
    assert not repair_methods.eligible(store.get(second), goal("NEW"))
    assert store.get(second)["reuse_blocked"]


def test_bounded_source_integrity_and_parameterized_methods_are_not_auto_reused(tmp_path):
    method = repair_methods.remember(GeneratedOp(source=BOUND_METHOD), goal("TITLE"))
    assert repair_methods.source(method) == BOUND_METHOD
    broken = copy.deepcopy(method)
    broken["source_sha256"] = "0" * 64
    with pytest.raises(ValueError, match="hash changed"):
        repair_methods.source(broken)
    broken["source_zlib"] = [base64.b64encode(zlib.compress(b"x" * (repair_methods.MAX_SOURCE + 1))).decode()]
    with pytest.raises(ValueError, match="oversized"):
        repair_methods.source(broken)
    method["params"] = {"index": 1}
    case = {"candidate": method, "status": "success", "structural_verified": True, "reuse_blocked": False}
    assert not repair_methods.eligible(case, goal("TITLE"))
    assert repair_methods.evidence([case], goal("TITLE"))[0]["source"] == BOUND_METHOD


def test_reused_content_method_reads_current_wording_and_cannot_mutate_authority(tmp_path):
    from test_governed_codegen import CONTENT_OP, contract as content_goal
    from spec_surgeon.codegen import trial
    source = doc(tmp_path / "source.docx", "FIRE SUPPRESSION", 1)
    dynamic = CONTENT_OP.replace('tracked = params.get("tracked", False)',
        'effect = pkg.goal["effects"][0]\n        tracked = effect.get("tracked", False)')
    dynamic = dynamic.replace('params["replacement"]', 'effect["replace"]')
    g = content_goal()
    g["effects"][0]["replace"] = "CURRENT PROJECT"
    case = {"candidate": repair_methods.remember(GeneratedOp(source=dynamic), content_goal()),
            "status": "success", "structural_verified": True, "reuse_blocked": False}
    assert repair_methods.eligible(case, g)
    result = trial(repair_methods.source(case["candidate"]), source, {}, {}, tmp_path / "current",
        goal_contract=g, goal_binding=goals.resolve_goal(source, g)["binding"])
    assert result.passed, result.problems
    assert "CURRENT PROJECT SUPPRESSION" in [p.text for p in Document(result.output).paragraphs]
    # Changing the sandbox's job data cannot change the host's approved contract.
    malicious = dynamic.replace('effect = pkg.goal["effects"][0]',
        'effect = pkg.goal["effects"][0]\n        effect["replace"] = "OLD PROJECT"')
    rejected = trial(malicious, source, {}, {}, tmp_path / "wrong",
        goal_contract=g, goal_binding=goals.resolve_goal(source, g)["binding"])
    assert not rejected.passed and rejected.record.get("GOAL").status == "FAIL"
    assert g["effects"][0]["replace"] == "CURRENT PROJECT"
