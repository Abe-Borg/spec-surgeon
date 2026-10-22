"""Every file of a batch side by side, the odd one named (docs/PLAN-2026-10-22-batch-consistency.md C2).

From the owner's Division 21 run: 21 05 53's footer showed its title in mixed case while the other files showed
capitals, and File Properties titles from other projects sat in 21 12 00 and 21 31 13 — nothing compared the files,
and only four structure patterns were profiled up front. The overview reads every file without a model call; its flags
describe and decide nothing."""
from __future__ import annotations

import json
from pathlib import Path

import pytest
from docx import Document

from fake_claude import FakeClient, resp, tool
from spec_surgeon import compiler, recovery
from spec_surgeon.document_task import TOOL_ACTIVITY, DocumentTask
from spec_surgeon.gui import Session
from spec_surgeon.ops.effective import PropResolver
from spec_surgeon.ops.selectors import StyleSheet
from spec_surgeon.package import DocxPackage, W
from spec_surgeon.profile import (COMPARE_PAGE, MAX_OVERVIEW_BYTES, OVERVIEW_FLAGS, OVERVIEW_LINES_CUT, OWN_NUMBER,
                                  OWN_TITLE, SAME_AS_MOST, batch_overview, letters_shown)

FOOTER = ["{title}", "2026-10-01", "{number} - 1"]


def spec(path: Path, number: str, title: str, *, footer=FOOTER, footer_caps=False, body_title=None,
         properties_title="", section_line=True, first_footer=None, title_page=False, even_footer=None,
         even_and_odd=False) -> Path:
    """A specification: "SECTION <number>", its title (``body_title`` as typed, else in capitals), a footer of
    ``footer`` lines ({title} and {number} the file's own; the first line under w:caps when ``footer_caps``)."""
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    sec = doc.sections[0]

    def fill(hf, lines, caps_first=False):
        hf.is_linked_to_previous = False
        for i, line in enumerate(lines):
            para = hf.paragraphs[0] if i == 0 else hf.add_paragraph()
            para.add_run(line.format(title=title, number=number)).font.all_caps = caps_first and i == 0

    if footer is not None:
        fill(sec.footer, footer, footer_caps)
    if first_footer is not None:
        sec.different_first_page_header_footer = True
        fill(sec.first_page_footer, first_footer)
        sec.different_first_page_header_footer = title_page
    if even_footer is not None:
        doc.settings.odd_and_even_pages_header_footer = True
        fill(sec.even_page_footer, even_footer)
        doc.settings.odd_and_even_pages_header_footer = even_and_odd
    if section_line:
        doc.add_paragraph(f"SECTION {number}")
    doc.add_paragraph(body_title if body_title is not None else title.upper())
    doc.add_paragraph("PART 1 - GENERAL")
    doc.core_properties.title = properties_title
    doc.save(path)
    return path


DIVISION_21 = [("21 05 17", "SLEEVES AND SLEEVE SEALS FOR FIRE-SUPPRESSION PIPING", "SECTION 21 05 17 - SLEEVES"),
               ("21 05 29", "HANGERS AND SUPPORTS FOR FIRE-SUPPRESSION PIPING", ""),
               ("21 05 53", "Identification for Fire Suppression Systems", ""),
               ("21 12 00", "FIRE-SUPPRESSION STANDPIPES", "SECTION 211313 - WET-PIPE SPRINKLER SYSTEMS"),
               ("21 31 13", "ELECTRIC-DRIVE, CENTRIFUGAL FIRE PUMPS", "SECTION PP-16622 - REMOTE POWER PANELS")]


def division_21(folder: Path):
    return [spec(folder / f"{n}.docx", n, t, body_title=t, properties_title=p) for n, t, p in DIVISION_21]


def flags_of(ov, file):
    return [f["flag"] for f in ov.flags if f["file"] == file]


def test_the_odd_footer_title_and_the_stale_properties_are_named(tmp_path):
    ov = batch_overview(division_21(tmp_path))
    assert flags_of(ov, "21 05 53.docx") == [
        "its title “Identification for Fire Suppression Systems” shows in mixed case; most files (4 of 5) show their "
        "title in capitals",
        "footer line “Identification for Fire Suppression Systems” shows in mixed case; most files (4 of 5) show that "
        "line in capitals"]
    assert flags_of(ov, "21 12 00.docx") == [
        "its File Properties title “SECTION 211313 - WET-PIPE SPRINKLER SYSTEMS” names section 211313, while the file "
        "is 21 12 00, and has none of the words of its title “FIRE-SUPPRESSION STANDPIPES”"]
    assert flags_of(ov, "21 31 13.docx") == [
        "its File Properties title “SECTION PP-16622 - REMOTE POWER PANELS” names section 16622, while the file is "
        "21 31 13, and has none of the words of its title “ELECTRIC-DRIVE, CENTRIFUGAL FIRE PUMPS”"]
    assert not flags_of(ov, "21 05 17.docx") and not flags_of(ov, "21 05 29.docx")   # "SLEEVES" is a title word
    out = ov.summary()
    assert out["most_files_show"] == {"footer.default": {
        "lines": [OWN_TITLE, "2026-10-01", f"{OWN_NUMBER} - 1"], "files_showing_every_line": "5 of 5"}}
    rows = {r["file"]: r for r in out["rows"]}
    assert rows["21 05 17.docx"]["headers_footers"] == {"footer.default": SAME_AS_MOST}
    odd = rows["21 05 53.docx"]                    # a flagged header/footer is shown in full
    assert odd["headers_footers"]["footer.default"]["lines"] == [
        "Identification for Fire Suppression Systems", "2026-10-01", "21 05 53 - 1"]
    assert (odd["section_number"], odd["title_shown_in"]) == ("21 05 53", "mixed case")
    assert rows["21 31 13.docx"]["properties_title"] == "SECTION PP-16622 - REMOTE POWER PANELS"
    assert "properties_title" not in rows["21 05 29.docx"]


def test_capitals_by_formatting_are_capitals(tmp_path):
    """Typed capitals and capitals through w:caps look the same to a reader: neither is flagged against the other, and
    the row shows the text as typed with shown_in_capitals (as the profiles do)."""
    files = [spec(tmp_path / "a.docx", "21 05 17", "SLEEVES FOR PIPING"),
             spec(tmp_path / "b.docx", "21 05 29", "HANGERS FOR PIPING"),
             spec(tmp_path / "c.docx", "21 05 53", "Identification for Piping", footer_caps=True)]
    ov = batch_overview(files)
    assert not ov.flags
    row = ov.compare(["c.docx"])["rows"][0]["headers_footers"]["footer.default"]
    assert row["lines"][0] == "Identification for Piping" and row["shown_in_capitals"] == {"0": "all capitals"}
    assert batch_overview(files[:2] + [spec(tmp_path / "d.docx", "21 05 53", "Identification for Piping")]).flags


def test_a_line_that_did_not_take_a_line_of_its_own_and_a_renamed_title(tmp_path):
    files = [spec(tmp_path / f"{n}.docx", n, f"TITLE {n[-2:]} OF PIPING") for n in ("21 05 17", "21 05 29", "21 05 48")]
    files.append(spec(tmp_path / "no-date.docx", "21 13 13", "WET-PIPE SPRINKLER SYSTEMS",
                      footer=["{title}", "{number} - 1"]))
    files.append(spec(tmp_path / "draft.docx", "21 13 16", "DRY-PIPE SPRINKLER SYSTEMS", footer=FOOTER + ["DRAFT"]))
    renamed = ["Identification for Fire-Suppression Piping and Equipment", "2026-10-01", "{number} - 1"]
    files.append(spec(tmp_path / "renamed.docx", "21 05 53", "IDENTIFICATION FOR FIRE SUPPRESSION SYSTEMS",
                      footer=renamed))
    ov = batch_overview(files)
    assert flags_of(ov, "no-date.docx") == ["footer has no line like “2026-10-01”; most files (5 of 6) show one"]
    assert flags_of(ov, "draft.docx") == ["footer shows “DRAFT”, a line most files do not"]
    assert flags_of(ov, "renamed.docx") == [
        "footer reads “Identification for Fire-Suppression Piping and Equipment” (in mixed case, where most show "
        "capitals) where most files (5 of 6) show “«SECTION TITLE»”"]
    assert len(ov.flags) == 3
    rows = {r["file"]: r for r in ov.summary()["rows"]}
    assert all(isinstance(rows[f]["headers_footers"]["footer.default"], dict)
               for f in ("no-date.docx", "draft.docx", "renamed.docx"))
    assert ov.summary()["most_files_show"]["footer.default"]["files_showing_every_line"] == "4 of 6"


def test_own_values_page_numbers_and_dates_never_make_lines_differ(tmp_path):
    files = [spec(tmp_path / "a.docx", "21 05 17", "FIRE SUPPRESSION PIPING"),
             spec(tmp_path / "b.docx", "21 05 29", "FIRE-SUPPRESSION HANGERS", footer=["{title}", "2026-11-14",
                                                                                     "210529 – 3"]),
             spec(tmp_path / "c.docx", "21 05 53", "FIRE SUPPRESSION IDENTIFICATION",
                  footer=["FIRE-SUPPRESSION   IDENTIFICATION", "2026-10-01", "21-05-53 - 12"])]
    assert batch_overview(files).flags == []


def test_a_missing_and_an_extra_header_footer(tmp_path):
    files = [spec(tmp_path / f"{n}.docx", n, f"TITLE {n[-2:]}") for n in ("21 05 17", "21 05 29", "21 05 48")]
    files.append(spec(tmp_path / "bare.docx", "21 13 13", "WET-PIPE", footer=None))
    files.append(spec(tmp_path / "cover.docx", "21 13 16", "DRY-PIPE", first_footer=["COVER"], title_page=True))
    # a first-page footer Word does not show (no titlePg) and an even-page one without evenAndOddHeaders: not listed
    files.append(spec(tmp_path / "hidden.docx", "21 13 19", "PREACTION", first_footer=["COVER"],
                      even_footer=["EVEN"]))
    ov = batch_overview(files)
    assert flags_of(ov, "bare.docx") == ["shows no footer; most files (5 of 6) show one"]
    assert flags_of(ov, "cover.docx") == ["shows a footer (first page) most files do not (1 of 6)"]
    assert not flags_of(ov, "hidden.docx")
    rows = {r["file"]: r for r in ov.compare()["rows"]}
    assert set(rows["hidden.docx"]["headers_footers"]) == {"footer.default"}
    assert set(rows["cover.docx"]["headers_footers"]) == {"footer.default", "footer.first"}
    assert rows["bare.docx"]["headers_footers"] == {}
    shown = batch_overview([spec(tmp_path / "even.docx", "21 13 19", "PREACTION", even_footer=["EVEN"],
                                 even_and_odd=True)]).compare()["rows"][0]["headers_footers"]
    assert set(shown) == {"footer.default", "footer.even"}


def test_two_files_have_no_majority(tmp_path):
    files = [spec(tmp_path / "a.docx", "21 05 17", "SLEEVES"),
             spec(tmp_path / "b.docx", "21 05 29", "Hangers", body_title="Hangers", footer=["{title}", "DRAFT"])]
    assert batch_overview(files).flags == []


def test_a_file_the_reference_rules_cannot_read_is_compared_as_typed(tmp_path):
    files = [spec(tmp_path / f"{n}.docx", n, f"TITLE {n[-2:]}") for n in ("21 05 17", "21 05 29", "21 05 48")]
    files.append(spec(tmp_path / "loose.docx", "21 13 13", "WET-PIPE SPRINKLER", section_line=False))
    ov = batch_overview(files)
    assert flags_of(ov, "loose.docx")[0] == ("the reference rules find no section number or section title in it "
                                             "(most files have theirs); its headers/footers are compared as typed")
    assert flags_of(ov, "loose.docx")[1:] == [
        "footer reads “WET-PIPE SPRINKLER” where most files (3 of 4) show “«SECTION TITLE»”",
        "footer reads “21 13 13 - 1” where most files (3 of 4) show “«SECTION NUMBER» - 1”"]
    row = ov.compare(["loose.docx"])["rows"][0]
    assert row["section_number"] == row["section_title"] == "(not found by the reference rules)"


def test_properties_titles_read_as_section_numbers_and_title_words(tmp_path):
    def flagged(props, number="21 05 53", title="IDENTIFICATION FOR FIRE SUPPRESSION"):
        ov = batch_overview([spec(tmp_path / "p.docx", number, title, properties_title=props)])
        return [f["flag"] for f in ov.flags]
    assert flagged("SECTION 21 05 53 - IDENTIFICATION") == []
    assert flagged("210553") == []                                    # the own number, no words: nothing to say
    assert flagged("Identification, issued 2026-10-01") == []         # a date is no section number
    assert flagged("") == []
    assert flagged("Fire suppression identification 21 13 13") == [
        "its File Properties title “Fire suppression identification 21 13 13” names section 21 13 13, while the file "
        "is 21 05 53"]
    assert flagged("Remote power panels") == [
        "its File Properties title “Remote power panels” has none of the words of its title “IDENTIFICATION FOR FIRE "
        "SUPPRESSION”"]
    assert flagged("For the Owner") == [                             # "for" and "the" are not title words
        "its File Properties title “For the Owner” has none of the words of its title “IDENTIFICATION FOR FIRE "
        "SUPPRESSION”"]
    assert flagged("SECTION 21 05 53.13 - IDENTIFICATION", number="21 05 53") == []


def test_letters_shown_reads_the_look_of_a_line():
    doc = Document()
    rows = [("SECTION 21 05 53", {}), ("Fire Suppression", {}), ("fire suppression", {}),
            ("Fire Suppression", {"caps": True}), ("Fire Suppression", {"small_caps": True}), ("21 05 53 - 1", {})]
    for text, look in rows:
        run = doc.add_paragraph().add_run(text)
        run.font.all_caps = look.get("caps")
        run.font.small_caps = look.get("small_caps")
    import io
    buf = io.BytesIO()
    doc.save(buf)
    pkg = DocxPackage(buf.getvalue(), name="x.docx")
    props = PropResolver(StyleSheet(pkg))
    paras = list(pkg.xml(pkg.main_part()).iter(W + "p"))
    assert [letters_shown(p, props) for p in paras] == ["capitals", "mixed case", "lower case", "capitals",
                                                         "small capitals", None]


def test_the_overview_is_bounded_and_compare_files_pages_through_every_file(tmp_path):
    files = [spec(tmp_path / f"{i:03d}.docx", f"21 {i // 100:02d} {i % 100:02d}", f"TITLE NUMBER {i}",
                  footer=["{title}", "2026-10-01", f"SHEET {chr(65 + i % 26)}{chr(65 + i // 26)}"])
             for i in range(70)]
    ov = batch_overview(files)
    out = ov.summary()
    assert len(json.dumps(out, ensure_ascii=False)) <= MAX_OVERVIEW_BYTES
    assert out["truncated"][0] == OVERVIEW_LINES_CUT
    cut = int(out["truncated"][-1].split("last ")[1].split(" ")[0])
    assert len(out["rows"]) + cut == 70 and cut > 0
    assert out["flags"] == ov.flags[:OVERVIEW_FLAGS] and len(ov.flags) == 70     # every file's SHEET line is its own
    assert out["flags_more"].startswith(f"{70 - OVERVIEW_FLAGS} more")
    seen, start = [], 0
    while start is not None:
        page = ov.compare(start=start, count=COMPARE_PAGE)
        seen += [r["file"] for r in page["rows"]]
        assert all(len(r["headers_footers"]["footer.default"]["lines"]) == 3 for r in page["rows"])
        start = page["next_start"]
    assert seen == [f.name for f in files]
    assert ov.compare(["005.docx"])["flags"] == [f for f in ov.flags if f["file"] == "005.docx"]
    with pytest.raises(ValueError):
        ov.compare(count=COMPARE_PAGE + 1)
    with pytest.raises(KeyError):
        ov.compare(["nope.docx"])


def test_an_unreadable_file_is_listed_and_the_others_compare(tmp_path):
    import zipfile
    files = division_21(tmp_path)
    bad = tmp_path / "broken.docx"
    bad.write_bytes(b"not a zip")
    ov = batch_overview(files + [bad])
    assert [u["file"] for u in ov.summary()["unreadable"]] == ["broken.docx"]
    assert flags_of(ov, "21 05 53.docx")
    # File Properties that do not parse leave out only the Properties title
    odd = tmp_path / "props.docx"
    with zipfile.ZipFile(files[3]) as src, zipfile.ZipFile(odd, "w") as dst:
        for item in src.infolist():
            dst.writestr(item, b"<cp:coreProperties" if item.filename == "docProps/core.xml" else src.read(item))
    row = batch_overview([odd]).compare()["rows"][0]
    assert row["section_number"] == "21 12 00" and "properties_title" not in row


def test_the_batch_opens_with_the_overview_and_offers_compare_files(tmp_path):
    a = division_21(tmp_path / "a")
    b = spec(tmp_path / "b" / "21 05 53.docx", "21 05 53", "IDENTIFICATION FOR FIRE SUPPRESSION SYSTEMS")
    batch = compiler._Batch(a + [b], None)
    summary = batch.summary()
    assert list(summary)[:2] == ["files", "overview"]
    labels = [r["file"] for r in summary["overview"]["rows"]]
    assert "a/21 05 53.docx" in labels and "b/21 05 53.docx" in labels     # the labels the tools take
    assert batch.compare_files(["a/21 05 53.docx"])["flags"][0]["file"] == "a/21 05 53.docx"
    message = compiler.build_messages("Make the footers match.", batch, [])[0]["content"]
    assert message.index('"overview"') < message.index('"patterns"')
    for tools in (compiler.TOOLS, compiler.CHAT_TOOLS, recovery.TOOLS):
        assert "compare_files" in {t["name"] for t in tools}
    assert TOOL_ACTIVITY["compare_files"]
    assert "compare_files" in compiler.system_prompt(chat=True) and "most_files_show" in compiler.system_prompt()


def _batch_of(request) -> dict:
    content = request["messages"][0]["content"]
    if isinstance(content, list):
        content = "".join(block.get("text", "") for block in content if isinstance(block, dict))
    return json.loads(content.split("THE BATCH (", 1)[1].split(":\n", 1)[1])


def test_a_request_after_a_run_reads_the_copies_it_made(tmp_path):
    """The overview is computed whenever a batch is: the next request's batch is the working set, so it reads the
    copies the run made, not the originals."""
    files = division_21(tmp_path / "in")
    caps = {"request": "Show the first footer line in capitals.", "summary": "Footer title in capitals",
            "mode": "formatting",
            "ops": [{"op": "set_properties", "select": {"scope": ["footers"], "paragraphs": {"index_range": [0, 1]}},
                     "params": {"run": {"caps": True}, "tracked": False}}]}
    bold = {"request": "Make PART 1 bold.", "summary": "Bold PART 1", "mode": "formatting",
            "ops": [{"op": "set_properties", "select": {"paragraphs": {"text_regex": "^PART 1"}},
                     "params": {"run": {"b": True}, "tracked": False}}]}
    fake = FakeClient(resp(tool("submit_editspec", {"spec": caps})), resp(tool("submit_editspec", {"spec": bold})))
    task = DocumentTask(Session(client=fake, out_root=tmp_path / "out", visual_check=False), render=False)
    task.attach(files)
    view = task.perform("send", text=caps["request"])
    first = _batch_of(fake.requests[0])["overview"]
    assert any("footer line" in f["flag"] for f in first["flags"] if f["file"] == "21 05 53.docx")
    task.perform("primary", token=view["token"])
    assert task.stage == "done"
    task.perform("send", text=bold["request"])
    second = _batch_of(fake.requests[1])["overview"]
    odd = [f["flag"] for f in second["flags"] if f["file"] == "21 05 53.docx"]
    assert odd and not any("footer line" in f for f in odd)       # the copy shows its footer title in capitals
    assert any(f.startswith("its title") for f in odd)            # its body title is still typed in mixed case
