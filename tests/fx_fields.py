"""Fixtures for insert_field (built-in batch 6, docs/BUILTIN_OPS_BATCH6_DESIGN.md §9).

``build_fields()`` — a specification section with the placeholders where its fields go:

* body 0 ``SECTION 21 13 13 - WET-PIPE SPRINKLER SYSTEMS`` in Heading 1 (what a STYLEREF shows);
* body 1 ``Issued {{DATE}} for review.`` — the placeholder inside one run;
* body 2 ``Printed {{DATE}} by the office.`` — the placeholder over two bold runs, a proofing mark between them;
* body 3 ``See {{FILENAME}}`` — the placeholder in a hyperlink (untracked only: a tracked change stays out of links);
* body 4 ``Rev {{DATE}} draft`` — ``{{DATE}} draft`` in Jack's pending insertion (the core splits it, tracked);
* body 5 ``Project {{PROJECT}}`` — for DOCPROPERTY "Project Number" (``docProps/custom.xml``);
* body 6 ``End of section.`` — italic, for a field at the start or end;
* the footer ``Page {{PAGE}} of {{NUMPAGES}}`` (two placeholders in one paragraph, small caps) and ``21 13 13``;
* the header ``{{TITLE}}`` (STYLEREF Heading 1) and an empty second paragraph whose mark is bold.

Variants: ``in_field`` body 7 a REF field showing ``{{PAGE}}``; ``deleted_inside`` body 7 ``{{PA`` + Jack's deleted
``x`` + ``GE}}``; ``tracked_format`` body 1's run carries a pending tracked formatting change; ``textbox`` body 7 a text
box holding ``{{DATE}}``; ``toc`` body 7–8 a table of contents over two paragraphs (its begin in 7, its end in 8);
``footnote`` a footnotes part with ``Footnote dated {{DATE}}.``; ``custom=False`` no custom properties.

``build_placeholders()`` — "Page X of Y" in the footer and body words holding the placeholders' letters (whole-word
matching; its own docstring).
"""
from __future__ import annotations

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

from make_fixtures import W, W_NS, XML_SPACE, _textbox_run, add_parts, docx_bytes, patch_docx

DOC = "word/document.xml"
FOOTER = "word/footer1.xml"
HEADER = "word/header1.xml"
FOOTNOTES = "word/footnotes.xml"
CUSTOM = "docProps/custom.xml"
CU_NS = "http://schemas.openxmlformats.org/officeDocument/2006/custom-properties"
VT_NS = "http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes"
CUSTOM_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/custom-properties"
CUSTOM_CT = "application/vnd.openxmlformats-officedocument.custom-properties+xml"
FOOTNOTES_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes"
FOOTNOTES_CT = "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"
HYPERLINK_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink"
REVISION = {W + "author": "Jack", W + "date": "2026-10-02T09:00:00Z"}


def _el(tag: str, **attrs):
    e = etree.Element(W + tag)
    for k, v in attrs.items():
        e.set(W + k, v)
    return e


def _run(text: str, *props: str):
    r = _el("r")
    if props:
        rpr = etree.SubElement(r, W + "rPr")
        for name in props:
            etree.SubElement(rpr, W + name)
    t = etree.SubElement(r, W + "t")
    t.text = text
    if text != text.strip():
        t.set(XML_SPACE, "preserve")
    return r


def _revision(tag: str, rid: str, *children):
    e = _el(tag, id=rid)
    for k, v in REVISION.items():
        e.set(k, v)
    for c in children:
        e.append(c)
    return e


def _char_run(kind: str):
    r = _el("r")
    r.append(_el("fldChar", fldCharType=kind))
    return r


def _field_runs(instruction: str, result: str, end: bool = True):
    """A complex field as Word writes it, showing ``result``; ``end=False`` leaves it open (it ends later)."""
    code = _el("instrText")
    code.set(XML_SPACE, "preserve")
    code.text = instruction
    holder = _el("r")
    holder.append(code)
    return [_char_run("begin"), holder, _char_run("separate"), _run(result)] + ([_char_run("end")] if end else [])


def _custom_xml() -> bytes:
    root = etree.Element("{%s}Properties" % CU_NS, nsmap={None: CU_NS, "vt": VT_NS})
    p = etree.SubElement(root, "{%s}property" % CU_NS)
    p.set("fmtid", "{D5CDD505-2E9C-101B-9397-08002B2CF9AE}")
    p.set("pid", "2")
    p.set("name", "Project Number")
    etree.SubElement(p, "{%s}lpwstr" % VT_NS).text = "P-1003"
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _footnotes_xml() -> bytes:
    return (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<w:footnotes xmlns:w="{W_NS}">'
            '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
            '<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:r><w:continuationSeparator/></w:r></w:p>'
            '</w:footnote><w:footnote w:id="1"><w:p><w:r><w:t>Footnote dated {{DATE}}.</w:t></w:r></w:p>'
            '</w:footnote></w:footnotes>').encode("utf-8")


def build_fields(*, in_field: bool = False, deleted_inside: bool = False, tracked_format: bool = False,
                 textbox: bool = False, toc: bool = False, footnote: bool = False, custom: bool = True) -> bytes:
    doc = Document()
    doc.add_paragraph("SECTION 21 13 13 - WET-PIPE SPRINKLER SYSTEMS", style="Heading 1")
    doc.add_paragraph("Issued {{DATE}} for review.")
    for _ in range(5):
        doc.add_paragraph()
    if in_field or deleted_inside or textbox or toc:
        doc.add_paragraph()
    if toc:
        doc.add_paragraph()
    footer = doc.sections[0].footer
    footer.add_paragraph("21 13 13")
    header = doc.sections[0].header
    header.paragraphs[0].text = "{{TITLE}}"
    header.add_paragraph()
    if textbox:
        _textbox_run(doc.paragraphs[7], "Box dated {{DATE}}")
    data = docx_bytes(doc)

    def fix_document(root):
        pars = [p for p in root.find(W + "body") if p.tag == W + "p"]
        p = pars[2]
        p.append(_run("Printed "))
        p.append(_run("{{DA", "b"))
        p.append(_el("proofErr", type="spellStart"))
        p.append(_run("TE}}", "b"))
        p.append(_run(" by the office."))
        p = pars[3]
        p.append(_run("See "))
        link = _el("hyperlink")
        link.set("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id", "rId900")
        link.append(_run("{{FILENAME}}"))
        p.append(link)
        p = pars[4]
        p.append(_run("Rev "))
        p.append(_revision("ins", "800", _run("{{DATE}} draft")))
        pars[5].append(_run("Project {{PROJECT}}"))
        pars[6].append(_run("End of section.", "i"))
        if tracked_format:
            rpr = etree.Element(W + "rPr")
            change = _revision("rPrChange", "801", etree.Element(W + "rPr"))
            rpr.append(etree.Element(W + "b"))
            rpr.append(change)
            pars[1].find(W + "r").insert(0, rpr)
        if in_field:
            pars[7].append(_run("See page "))
            for r in _field_runs(" REF _Ref1 \\h ", "{{PAGE}}"):
                pars[7].append(r)
        if deleted_inside:
            pars[7].append(_run("{{PA"))
            gone = _run("x")
            gone.find(W + "t").tag = W + "delText"
            pars[7].append(_revision("del", "802", gone))
            pars[7].append(_run("GE}}"))
        if toc:
            for r in _field_runs(" TOC \\o \"1-3\" ", "Part 1", end=False):
                pars[7].append(r)
            pars[8].append(_run("Part 2"))
            pars[8].append(_char_run("end"))                     # the table of contents ends here

    def fix_footer(root):
        p = root.find(W + "p")
        for text in ("Page ", "{{PAGE}}", " of ", "{{NUMPAGES}}"):
            p.append(_run(text, "smallCaps"))

    def fix_header(root):
        empty = root.findall(W + "p")[1]
        ppr = etree.SubElement(empty, W + "pPr")
        etree.SubElement(etree.SubElement(ppr, W + "rPr"), W + "b")

    data = patch_docx(data, {DOC: fix_document, FOOTER: fix_footer, HEADER: fix_header})
    rels = [("word/_rels/document.xml.rels", "rId900", HYPERLINK_REL, "https://example.com/specs")]
    parts, overrides = {}, {}
    if custom:
        parts[CUSTOM] = _custom_xml()
        rels.append(("_rels/.rels", "rId90", CUSTOM_REL, CUSTOM))
        overrides[CUSTOM] = CUSTOM_CT
    if footnote:
        parts[FOOTNOTES] = _footnotes_xml()
        rels.append(("word/_rels/document.xml.rels", "rId901", FOOTNOTES_REL, "footnotes.xml"))
        overrides[FOOTNOTES] = FOOTNOTES_CT
    data = add_parts(data, parts, rels=rels, overrides=overrides)
    return patch_docx(data, {"word/_rels/document.xml.rels": _external_link})


def build_placeholders() -> bytes:
    """"Page X of Y" in the footer — the user's one-letter placeholders — and body words holding the same letters:

    * body 0 ``WET-PIPE SPRINKLER SYSTEMS`` (a Y inside a word);
    * body 1 ``A.<tab>Provide FLEXIBLE couplings at EXTERIOR walls; MAXIMUM spacing per NFPA 13.`` — "EXTERIOR" over two
      runs ("E" | "XTERIOR …"): an X inside a word across a run boundary;
    * body 2 ``B.<tab>See Detail X on sheet X2 for the X_RAY room.`` — a whole X, then an X before a digit and before an
      underscore (word characters both);
    * body 3 ``C.<tab>UPDATE the schedule and VALIDATE the CANDIDATE products by DATE.`` — one whole DATE, three inside
      words.
    """
    doc = Document()
    doc.add_paragraph("WET-PIPE SPRINKLER SYSTEMS")
    p = doc.add_paragraph()
    p.add_run("A.\tProvide FLEXIBLE couplings at E")
    p.add_run("XTERIOR walls; MAXIMUM spacing per NFPA 13.")
    doc.add_paragraph("B.\tSee Detail X on sheet X2 for the X_RAY room.")
    doc.add_paragraph("C.\tUPDATE the schedule and VALIDATE the CANDIDATE products by DATE.")
    doc.sections[0].footer.paragraphs[0].add_run("Page X of Y")
    return docx_bytes(doc)


def _external_link(root):
    for rel in root:
        if rel.get("Id") == "rId900":
            rel.set("TargetMode", "External")
