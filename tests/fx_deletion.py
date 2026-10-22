"""Fixtures for built-in batch 6's deletion — delete_paragraphs over "all", tables, fields, section breaks, comments,
notes and bookmarks, and remove_footnotes (docs/BUILTIN_OPS_BATCH6_DESIGN.md §1, §2). Every paragraph is found by its
text in the tests, never by index.

``build_tables()`` — tables to delete whole, a paragraph between two tables, cells with several paragraphs, a nested
table; ``build_fields()`` — a local PAGE field, a simple field, a contents field over three paragraphs, a REF field to a
bookmark; ``build_sections()`` — three sections, the cover's header holding a picture, the body's footer and the
appendix's header inherited; ``build_marks()`` — comments (one inside a paragraph with a reply, one across two
paragraphs), bookmarks (inside, across), footnotes (one plain, one holding a link and a picture, one with a custom mark),
an endnote, a picture and a link in the body, a REF field to a bookmark; ``build_wipe()`` — everything at once over three
sections, for deleting the whole body; ``build_pictures()`` — paragraphs without text that hold a picture, a text box, a
field or a shape, and a table of pictures, between empty ones; ``build_kept_marks(shape)`` — a paragraph the deletion
keeps and empties (a cell's last, one between two tables, the body's last, the last section's) holding one end of a
bookmark whose other end stays. ``with_printer_settings`` names a printer settings part from section breaks;
``with_last_comment_paragraph`` gives a threaded comment a last paragraph of its own.
"""
from __future__ import annotations

import io

from docx import Document
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from lxml import etree

from make_fixtures import W, W_NS, _hyperlink, _textbox_run, add_parts, docx_bytes, patch_docx, png

R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
W14 = "http://schemas.microsoft.com/office/word/2010/wordml"
W15 = "http://schemas.microsoft.com/office/word/2012/wordml"
DOC = "word/document.xml"
FOOTNOTES = "word/footnotes.xml"
ENDNOTES = "word/endnotes.xml"
COMMENTS = "word/comments.xml"
_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"
_CT = "application/vnd.openxmlformats-officedocument.wordprocessingml.{}+xml"
EXTENDED_REL = "http://schemas.microsoft.com/office/2011/relationships/commentsExtended"


def text_of(p) -> str:
    return "".join(t.text or "" for t in p.iter(W + "t", W + "delText"))


def by_text(root, start: str):
    return next(p for p in root.iter(W + "p") if text_of(p).startswith(start))


def field(paragraph, instr: str, result: str) -> None:
    """A complex field in one paragraph: begin, instruction, separate, result, end — each its own run."""
    for kind, value in (("begin", None), ("instr", instr), ("separate", None), ("result", result), ("end", None)):
        r = paragraph.add_run()
        if kind == "instr":
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = value
        elif kind == "result":
            el = OxmlElement("w:t"); el.text = value
        else:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), kind)
        r._r.append(el)


def char(paragraph, kind: str, instr: str = None) -> None:
    r = paragraph.add_run()
    if instr is not None:
        el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = instr
    else:
        el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), kind)
    r._r.append(el)


def bookmark(paragraph, name: str, bid: int, where: str = "around") -> None:
    """A bookmark around the paragraph's runs (``around``), or only its start (``start``) or end (``end``)."""
    p = paragraph._p
    if where in ("around", "start"):
        s = OxmlElement("w:bookmarkStart"); s.set(qn("w:id"), str(bid)); s.set(qn("w:name"), name)
        ppr = p.find(qn("w:pPr"))
        (ppr.addnext(s) if ppr is not None else p.insert(0, s))
    if where in ("around", "end"):
        e = OxmlElement("w:bookmarkEnd"); e.set(qn("w:id"), str(bid))
        p.append(e)


def note_reference(paragraph, kind: str, nid: int, custom: bool = False) -> None:
    r = paragraph.add_run()
    rpr = OxmlElement("w:rPr")
    rs = OxmlElement("w:rStyle"); rs.set(qn("w:val"), "FootnoteReference" if kind == "footnote" else "EndnoteReference")
    rpr.append(rs)
    r._r.append(rpr)
    ref = OxmlElement(f"w:{kind}Reference"); ref.set(qn("w:id"), str(nid))
    if custom:
        ref.set(qn("w:customMarkFollows"), "1")
    r._r.append(ref)
    if custom:
        t = OxmlElement("w:t"); t.text = "*"
        r2 = paragraph.add_run()
        r2._r.append(t)


def _drawing(rid: str, docpr: int) -> str:
    return (f'<w:drawing xmlns:w="{W_NS}" xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
            'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
            'xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture" '
            f'xmlns:r="{R_NS}"><wp:inline><wp:extent cx="228600" cy="228600"/><wp:docPr id="{docpr}" name="Note picture"/>'
            '<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture"><pic:pic>'
            f'<pic:nvPicPr><pic:cNvPr id="{docpr}" name="note.png"/><pic:cNvPicPr/></pic:nvPicPr><pic:blipFill>'
            f'<a:blip r:embed="{rid}"/><a:stretch><a:fillRect/></a:stretch></pic:blipFill><pic:spPr><a:xfrm>'
            '<a:off x="0" y="0"/><a:ext cx="228600" cy="228600"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/>'
            '</a:prstGeom></pic:spPr></pic:pic></a:graphicData></a:graphic></wp:inline></w:drawing>')


def _note(kind: str, nid: int, body: str) -> str:
    style = "FootnoteText" if kind == "footnote" else "EndnoteText"
    ref = "footnoteRef" if kind == "footnote" else "endnoteRef"
    return (f'<w:{kind} w:id="{nid}"><w:p><w:pPr><w:pStyle w:val="{style}"/></w:pPr><w:r><w:rPr><w:rStyle '
            f'w:val="{"FootnoteReference" if kind == "footnote" else "EndnoteReference"}"/></w:rPr><w:{ref}/></w:r>'
            f'{body}</w:p></w:{kind}>')


def _notes_part(kind: str, notes: str) -> bytes:
    return (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<w:{kind}s xmlns:w="{W_NS}" xmlns:r="{R_NS}">'
            f'<w:{kind} w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:{kind}>'
            f'<w:{kind} w:type="continuationSeparator" w:id="0"><w:p><w:r><w:continuationSeparator/></w:r></w:p>'
            f'</w:{kind}>{notes}</w:{kind}s>').encode()


def with_notes(data: bytes, footnotes: str = "", endnotes: str = "", footnote_rels: str = "",
               media: dict = None) -> bytes:
    """``data`` with a footnotes part (and an endnotes part) holding the given note entries, their relationships and
    media."""
    parts, rels, overrides = {}, [], {}
    if footnotes:
        parts[FOOTNOTES] = _notes_part("footnote", footnotes)
        rels.append(("word/_rels/document.xml.rels", "rId801", _REL + "footnotes", "footnotes.xml"))
        overrides[FOOTNOTES] = _CT.format("footnotes")
    if endnotes:
        parts[ENDNOTES] = _notes_part("endnote", endnotes)
        rels.append(("word/_rels/document.xml.rels", "rId802", _REL + "endnotes", "endnotes.xml"))
        overrides[ENDNOTES] = _CT.format("endnotes")
    if footnote_rels:
        parts["word/_rels/footnotes.xml.rels"] = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<Relationships xmlns="http://schemas.'
            f'openxmlformats.org/package/2006/relationships">{footnote_rels}</Relationships>').encode()
    parts.update(media or {})
    overrides.update({name: "image/png" for name in media or {}})
    return add_parts(data, parts, rels=rels, overrides=overrides)


def _comment_ids(root, para_ids: dict) -> None:
    for c in root.findall(W + "comment"):
        for k, cp in enumerate(c.iter(W + "p")):
            cp.set("{%s}paraId" % W14, para_ids[c.get(W + "id")] if k == 0 else f"{0x3C000000 + k:08X}")


def _extended(threads) -> bytes:
    ex = etree.Element("{%s}commentsEx" % W15, nsmap={"w15": W15})
    for pid, parent in threads:
        e = etree.SubElement(ex, "{%s}commentEx" % W15)
        e.set("{%s}paraId" % W15, pid)
        if parent:
            e.set("{%s}paraIdParent" % W15, parent)
        e.set("{%s}done" % W15, "0")
    return etree.tostring(ex, xml_declaration=True, encoding="UTF-8", standalone=True)


# ---- tables -------------------------------------------------------------------------------------------------------

def _fill(table, rows):
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            table.cell(r, c).text = text


def build_tables() -> bytes:
    doc = Document()
    doc.add_paragraph("SECTION 21 13 13")
    doc.add_paragraph("PART 1 - GENERAL")
    _fill(doc.add_table(rows=2, cols=2), [["Tag", "Size"], ["P-1", "2 in"]])
    doc.add_paragraph("Between the tables.")
    t2 = doc.add_table(rows=2, cols=2)
    _fill(t2, [["Valve", "Rating"], ["V-1", "175 psi"]])
    t2.cell(1, 1).add_paragraph("Rated per UL.")
    doc.add_paragraph("After the tables.")
    t3 = doc.add_table(rows=1, cols=2)
    t3.cell(0, 0).text = "Outer cell."
    t3.cell(0, 1).text = "Above the inner table."
    inner = t3.cell(0, 1).add_table(rows=1, cols=1)
    inner.cell(0, 0).text = "Inner cell."
    doc.add_paragraph("END OF SECTION")
    return docx_bytes(doc)


# ---- fields -------------------------------------------------------------------------------------------------------

def build_fields() -> bytes:
    doc = Document()
    doc.add_paragraph("SECTION 21 13 13")
    p = doc.add_paragraph("A. Page ")
    field(p, " PAGE ", "3")
    p.add_run(" of the specification.")
    p = doc.add_paragraph("B. See ")
    fs = OxmlElement("w:fldSimple"); fs.set(qn("w:instr"), " REF Clause1 \\h ")
    r = OxmlElement("w:r"); t = OxmlElement("w:t"); t.text = "2.01"; r.append(t); fs.append(r)
    p._p.append(fs)
    p.add_run(".")
    p = doc.add_paragraph()
    char(p, "begin"); char(p, None, " TOC \\o \"1-3\" "); char(p, "separate")
    p.add_run("1.01 SUMMARY")
    doc.add_paragraph("1.02 SUBMITTALS")
    p = doc.add_paragraph("1.03 QUALITY ASSURANCE")
    char(p, "end")
    p = doc.add_paragraph("C. The clause a reference points to.")
    bookmark(p, "Clause1", 11)
    p = doc.add_paragraph("D. Refers to ")
    field(p, " REF Clause1 \\h ", "C.")
    doc.add_paragraph("END OF SECTION")
    return docx_bytes(doc)


# ---- sections -----------------------------------------------------------------------------------------------------

def build_sections() -> bytes:
    doc = Document()
    cover = doc.sections[0]
    cover.header.is_linked_to_previous = False
    hp = cover.header.paragraphs[0]
    hp.add_run("COVER HEADER ")
    hp.add_run().add_picture(io.BytesIO(png((200, 30, 30), size=5)), width=Pt(12))
    cover.footer.is_linked_to_previous = False
    cover.footer.paragraphs[0].add_run("Cover footer")
    doc.add_paragraph("PROJECT MANUAL")
    doc.add_paragraph("Cover note.")
    body = doc.add_section(WD_SECTION.NEW_PAGE)
    body.header.is_linked_to_previous = False
    body.header.paragraphs[0].add_run("BODY HEADER")
    body.footer.is_linked_to_previous = True                  # inherits the cover's footer
    doc.add_paragraph("PART 1 - GENERAL")
    doc.add_paragraph("A. Body text.")
    appendix = doc.add_section(WD_SECTION.NEW_PAGE)
    appendix.header.is_linked_to_previous = True              # inherits the body's header
    appendix.footer.is_linked_to_previous = False
    appendix.footer.paragraphs[0].add_run("Appendix footer")
    doc.add_paragraph("APPENDIX A")
    doc.add_paragraph("Attached.")
    data = docx_bytes(doc)

    def name_breaks(root):                                    # python-docx's break paragraphs are empty: name them
        breaks = [p for p in root.iter(W + "p") if p.find(f"{W}pPr/{W}sectPr") is not None]
        for p, text in zip(breaks, ("End of the cover.", "End of the body.")):
            r = etree.SubElement(p, W + "r")
            etree.SubElement(r, W + "t").text = text
    return patch_docx(data, {DOC: name_breaks})


# ---- comments, bookmarks, notes, pictures, links -------------------------------------------------------------------

FOOTNOTE_IMAGE = "word/media/note1.png"


def build_marks(textbox: bool = False) -> bytes:
    doc = Document()
    doc.add_paragraph("SECTION 21 13 13")
    p = doc.add_paragraph("A. Comment ")
    r = p.add_run("wholly inside")
    p.add_run(" this paragraph.")
    doc.add_comment(r, text="Check this.", author="Jeremy", initials="JS")          # 0
    doc.add_comment(r, text="Checked.", author="Jack", initials="JK")               # 1, a reply to 0
    p = doc.add_paragraph("B. A comment starts ")
    r = p.add_run("here")
    doc.add_comment(r, text="Spans B and C.", author="Abe", initials="AB")          # 2, its end moved into C
    doc.add_paragraph("C. and ends in this paragraph.")
    p = doc.add_paragraph("D. Bookmark wholly inside.")
    bookmark(p, "Inside", 21)
    p = doc.add_paragraph("E. A bookmark starts here")
    bookmark(p, "Across", 22, "start")
    p = doc.add_paragraph("F. and ends here.")
    bookmark(p, "Across", 22, "end")
    p = doc.add_paragraph("G. A footnote here.")
    note_reference(p, "footnote", 1)
    p = doc.add_paragraph("H. A footnote with a link and a picture.")
    note_reference(p, "footnote", 2)
    p = doc.add_paragraph("I. An endnote.")
    note_reference(p, "endnote", 1)
    p = doc.add_paragraph("J. A custom mark.")
    note_reference(p, "footnote", 3, custom=True)
    p = doc.add_paragraph("K. Picture: ")
    p.add_run().add_picture(io.BytesIO(png((30, 30, 200), size=4)), width=Pt(14))
    p = doc.add_paragraph("L. Link: ")
    _hyperlink(p, "nfpa.org", "https://nfpa.org/")
    p = doc.add_paragraph("M. Refers to the bookmark ")
    field(p, " REF Inside \\h ", "D.")
    if textbox:
        _textbox_run(doc.add_paragraph("N. Text box: "), "Boxed words")
    doc.add_paragraph("END OF SECTION")
    data = docx_bytes(doc)

    def fix_document(root):
        b, c = by_text(root, "B."), by_text(root, "C.")
        for el in list(b):
            if el.tag == W + "commentRangeEnd" or (el.tag == W + "r" and el.find(W + "commentReference") is not None):
                b.remove(el)
                c.append(el)
    data = patch_docx(data, {DOC: fix_document, COMMENTS: lambda root: _comment_ids(
        root, {"0": "1A000001", "1": "1A000002", "2": "1A000003"})})
    data = add_parts(data, {"word/commentsExtended.xml": _extended([("1A000001", None), ("1A000002", "1A000001"),
                                                                     ("1A000003", None)])},
                     rels=[("word/_rels/document.xml.rels", "rId890", EXTENDED_REL, "commentsExtended.xml")],
                     overrides={"word/commentsExtended.xml": _CT.format("commentsExtended")})
    footnotes = (_note("footnote", 1, '<w:r><w:t xml:space="preserve"> The first note.</w:t></w:r>')
                 + _note("footnote", 2, '<w:r><w:t xml:space="preserve"> See </w:t></w:r><w:hyperlink r:id="rId1">'
                                        '<w:r><w:t>nfpa.org</w:t></w:r></w:hyperlink><w:r>'
                                        + _drawing("rId2", 951) + '</w:r>')
                 + _note("footnote", 3, '<w:r><w:t xml:space="preserve"> A custom-marked note.</w:t></w:r>'))
    endnotes = _note("endnote", 1, '<w:r><w:t xml:space="preserve"> The endnote.</w:t></w:r>')
    rels = ('<Relationship Id="rId1" Type="' + _REL + 'hyperlink" Target="https://nfpa.org/" TargetMode="External"/>'
            '<Relationship Id="rId2" Type="' + _REL + 'image" Target="media/note1.png"/>')
    return with_notes(data, footnotes, endnotes, rels, {FOOTNOTE_IMAGE: png((20, 160, 20), size=3)})


# ---- notes ----------------------------------------------------------------------------------------------------------

def build_notes(custom: bool = False, twice: bool = False, shared_run: bool = False) -> bytes:
    """Footnote 1 (plain, a bookmark inside), footnote 2 (a link and a picture), endnotes 1 and 2, each referenced from
    its own body paragraph; ``custom`` adds footnote 3 with a custom mark, ``twice`` a second reference to footnote 1,
    ``shared_run`` puts footnote 1's reference in one run with text."""
    doc = Document()
    doc.add_paragraph("SECTION 21 13 13")
    note_reference(doc.add_paragraph("A. First reference."), "footnote", 1)
    note_reference(doc.add_paragraph("B. Second reference."), "footnote", 2)
    note_reference(doc.add_paragraph("C. Endnote reference."), "endnote", 1)
    note_reference(doc.add_paragraph("D. Another endnote."), "endnote", 2)
    doc.add_paragraph("E. Plain paragraph.")
    if custom:
        note_reference(doc.add_paragraph("F. Custom."), "footnote", 3, custom=True)
    if twice:
        note_reference(doc.add_paragraph("G. Again."), "footnote", 1)
    data = docx_bytes(doc)
    if shared_run:
        def share(root):
            ref = next(root.iter(W + "footnoteReference"))
            t = etree.SubElement(ref.getparent(), W + "t")
            t.text = "x"
        data = patch_docx(data, {DOC: share})
    footnotes = (_note("footnote", 1, '<w:bookmarkStart w:id="41" w:name="InNote"/><w:r><w:t xml:space="preserve"> '
                                      'See NFPA 13.</w:t></w:r><w:bookmarkEnd w:id="41"/>')
                 + _note("footnote", 2, '<w:r><w:t xml:space="preserve"> See </w:t></w:r><w:hyperlink r:id="rId1">'
                                        '<w:r><w:t>nfpa.org</w:t></w:r></w:hyperlink><w:r>'
                                        + _drawing("rId2", 952) + '</w:r>')
                 + (_note("footnote", 3, '<w:r><w:t xml:space="preserve"> Custom note.</w:t></w:r>') if custom else ""))
    endnotes = (_note("endnote", 1, '<w:r><w:t xml:space="preserve"> The first endnote.</w:t></w:r>')
                + _note("endnote", 2, '<w:r><w:t xml:space="preserve"> The second endnote.</w:t></w:r>'))
    rels = ('<Relationship Id="rId1" Type="' + _REL + 'hyperlink" Target="https://nfpa.org/" TargetMode="External"/>'
            '<Relationship Id="rId2" Type="' + _REL + 'image" Target="media/note1.png"/>')
    return with_notes(data, footnotes, endnotes, rels, {FOOTNOTE_IMAGE: png((20, 160, 20), size=3)})


# ---- the whole body ------------------------------------------------------------------------------------------------

def build_wipe() -> bytes:
    """Three sections (the cover's header with a picture), a table, a field, a bookmark, a comment and a footnote."""
    def more(root):
        general = by_text(root, "PART 1 - GENERAL")
        tbl = etree.fromstring(
            f'<w:tbl xmlns:w="{W_NS}"><w:tblPr><w:tblW w:w="0" w:type="auto"/></w:tblPr><w:tblGrid><w:gridCol '
            'w:w="4000"/></w:tblGrid><w:tr><w:tc><w:tcPr><w:tcW w:w="4000" w:type="dxa"/></w:tcPr><w:p><w:r><w:t>'
            'Schedule</w:t></w:r></w:p></w:tc></w:tr></w:tbl>')
        general.addnext(tbl)
        p = etree.fromstring(
            f'<w:p xmlns:w="{W_NS}"><w:bookmarkStart w:id="31" w:name="Wiped"/><w:r><w:t xml:space="preserve">B. Page '
            '</w:t></w:r><w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText xml:space="preserve"> PAGE '
            '</w:instrText></w:r><w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:t>2</w:t></w:r><w:r>'
            '<w:fldChar w:fldCharType="end"/></w:r><w:bookmarkEnd w:id="31"/><w:r><w:rPr><w:rStyle '
            'w:val="FootnoteReference"/></w:rPr><w:footnoteReference w:id="1"/></w:r></w:p>')
        tbl.addnext(p)
    data = with_notes(patch_docx(build_sections(), {DOC: more}),
                      _note("footnote", 1, '<w:r><w:t xml:space="preserve"> Wiped note.</w:t></w:r>'))
    doc = Document(io.BytesIO(data))
    target = next(p for p in doc.paragraphs if p.text.startswith("A. Body text"))
    doc.add_comment(target.runs[0], text="Gone with the body.", author="Jeremy", initials="JS")
    return docx_bytes(doc)


# ---- pictures in paragraphs without text (the built-in batch 6 review) ---------------------------------------------

def build_pictures() -> bytes:
    """Paragraphs without text that are not empty, between truly empty ones: "SECTION 21 13 13", an empty paragraph, a
    logo picture alone in its paragraph, "Figure 1 - Riser diagram", a signature table (a picture in its first cell,
    its second cell blank), "Signed.", a text box alone in its paragraph, a PAGE field with no result shown, a VML
    shape alone in its paragraph, another empty paragraph, "END OF SECTION"."""
    doc = Document()
    doc.add_paragraph("SECTION 21 13 13")
    doc.add_paragraph("")
    doc.add_paragraph().add_run().add_picture(io.BytesIO(png((200, 30, 30), size=8)), width=Pt(72))
    doc.add_paragraph("Figure 1 - Riser diagram")
    t = doc.add_table(rows=1, cols=2)
    t.cell(0, 0).paragraphs[0].add_run().add_picture(io.BytesIO(png((30, 30, 200), size=8)), width=Pt(72))
    doc.add_paragraph("Signed.")
    _textbox_run(doc.add_paragraph(), "Boxed words")
    fs = OxmlElement("w:fldSimple"); fs.set(qn("w:instr"), " PAGE ")
    doc.add_paragraph()._p.append(fs)
    doc.add_paragraph()._p.append(etree.fromstring(
        f'<w:r xmlns:w="{W_NS}" xmlns:v="urn:schemas-microsoft-com:vml"><w:pict><v:rect style="width:36pt;height:18pt"/>'
        '</w:pict></w:r>'))
    doc.add_paragraph("")
    doc.add_paragraph("END OF SECTION")
    return docx_bytes(doc)


# ---- range marks in a paragraph the deletion keeps -------------------------------------------------------------------

KEPT_SHAPES = {"cell": "^Q\\. ", "between": "^R\\. ", "body": "^S\\. ", "section": "^T\\. "}


def _range_mark(paragraph, tag: str, mid: int, first: bool = False) -> None:
    """A ``w:permStart`` / ``w:permEnd`` (``tag``) in the paragraph: after its properties (``first``) or at its end."""
    el = OxmlElement(f"w:{tag}"); el.set(qn("w:id"), str(mid))
    if tag == "permStart":
        el.set(qn("w:edGrp"), "everyone")
    p = paragraph._p
    ppr = p.find(qn("w:pPr"))
    if first:
        (ppr.addnext(el) if ppr is not None else p.insert(0, el))
    else:
        p.append(el)


def build_kept_marks(shape: str) -> bytes:
    """One paragraph to delete (its text starts with the letter ``KEPT_SHAPES`` selects) that the deletion must keep,
    emptied, holding one end of a bookmark whose other end is in text that stays:

    * ``cell`` — "Q. Gone …" and "Q. Kept …", the two paragraphs of a cell: the second is the cell's last. It holds
      the start of "KeptCell" (its end in the next cell, a REF field to it below the table) and of a permission range
      (its end in the next cell too), all of "Inner", and the end of "Wholly", which starts in "Q. Gone": both ends of
      those two lie in the deleted text;
    * ``between`` — "R. Between the tables.", alone between two tables, holds the end of "Span" (its start in
      "Lead.");
    * ``body`` — "S. Last of the body.", after a table, holds the end of "ToEnd" (its start in "Lead.");
    * ``section`` — "T. Appendix text.", all of the last section, holds the end of "IntoAppendix" (its start in
      "Lead.", before the section break)."""
    doc = Document()
    doc.add_paragraph("SECTION 21 13 13")
    if shape == "cell":
        t = doc.add_table(rows=1, cols=2)
        gone = t.cell(0, 0).paragraphs[0]
        gone.add_run("Q. Gone with its cell's text.")
        kept = t.cell(0, 0).add_paragraph("Q. Kept, the cell's last.")
        nxt = t.cell(0, 1).paragraphs[0]
        nxt.add_run("Next cell.")
        bookmark(gone, "Wholly", 30, "start")
        bookmark(kept, "Wholly", 30, "end")
        bookmark(kept, "Inner", 31)
        _range_mark(kept, "permStart", 33, first=True)
        bookmark(kept, "KeptCell", 32, "start")
        bookmark(nxt, "KeptCell", 32, "end")
        _range_mark(nxt, "permEnd", 33)
        p = doc.add_paragraph("A. See ")
        field(p, " REF KeptCell \\h ", "Q.")
        doc.add_paragraph("END OF SECTION")
    elif shape == "between":
        lead = doc.add_paragraph("Lead.")
        doc.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0].add_run("Alpha")
        between = doc.add_paragraph("R. Between the tables.")
        doc.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0].add_run("Beta")
        doc.add_paragraph("END OF SECTION")
        bookmark(lead, "Span", 34, "start")
        bookmark(between, "Span", 34, "end")
    elif shape == "body":
        lead = doc.add_paragraph("Lead.")
        doc.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0].add_run("Gamma")
        last = doc.add_paragraph("S. Last of the body.")
        bookmark(lead, "ToEnd", 35, "start")
        bookmark(last, "ToEnd", 35, "end")
    elif shape == "section":
        lead = doc.add_paragraph("Lead.")
        doc.add_section(WD_SECTION.NEW_PAGE)
        appendix = doc.add_paragraph("T. Appendix text.")
        bookmark(lead, "IntoAppendix", 36, "start")
        bookmark(appendix, "IntoAppendix", 36, "end")
    else:
        raise ValueError(shape)
    return docx_bytes(doc)


# ---- printer settings and comment threads (the built-in batch 6 review) ---------------------------------------------

PRINTER = "word/printerSettings/printerSettings1.bin"
PRINTER_REL = _REL + "printerSettings"


def with_printer_settings(data: bytes, sections=(0,), rid: str = "rIdPS1") -> bytes:
    """``data`` with a printer settings part (as Word on Windows writes one) named by the sectPr of each of ``sections``
    (indices among the paragraph-level sectPrs, then the body's own)."""
    data = add_parts(data, {PRINTER: b"\x00printer\x00" * 8},
                     rels=[("word/_rels/document.xml.rels", rid, PRINTER_REL, "printerSettings/printerSettings1.bin")],
                     overrides={PRINTER: "application/vnd.openxmlformats-officedocument.wordprocessingml."
                                         "printerSettings"})

    def name(root):
        body = root.find(W + "body")
        secs = [p.find(f"{W}pPr/{W}sectPr") for p in body.iter(W + "p") if p.find(f"{W}pPr/{W}sectPr") is not None]
        secs.append(body.find(W + "sectPr"))
        for i in sections:
            etree.SubElement(secs[i], W + "printerSettings").set("{%s}id" % R_NS, rid)
    return patch_docx(data, {DOC: name})


def with_last_comment_paragraph(data: bytes) -> bytes:
    """``build_marks()`` with comment 0 holding a second, last paragraph "Second line TBD." — the paragraph its
    commentsExtended entry names (paraId 1A000001: Word names a thread by its comment's last paragraph); its first
    paragraph takes another id."""
    def two_paragraphs(root):
        c0 = next(c for c in root.findall(W + "comment") if c.get(W + "id") == "0")
        c0.find(W + "p").set("{%s}paraId" % W14, "3C000009")
        p = etree.SubElement(c0, W + "p")
        p.set("{%s}paraId" % W14, "1A000001")
        p.set("{%s}textId" % W14, "77777777")
        etree.SubElement(etree.SubElement(p, W + "r"), W + "t").text = "Second line TBD."
    return patch_docx(data, {COMMENTS: two_paragraphs})
