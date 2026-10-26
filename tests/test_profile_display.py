"""What the profiles and the paragraph tools say about how text LOOKS, beside the text as typed (live-run finding
2026-10-22): lines Word shows in capitals through formatting are named,
and a header/footer view the size bound thins says how many lines it left out — a footer read as two lines when it
has four looks like a template that did not fully take."""
from __future__ import annotations

import io
import json

from docx import Document
from docx.shared import Pt

from make_fixtures import build_rich
from spec_surgeon.compiler import _Batch
from spec_surgeon.ops.effective import PropResolver
from spec_surgeon.ops.selectors import StyleSheet
from spec_surgeon.package import DocxPackage, W
from spec_surgeon.profile import PART_LINES_CUT, bound, profile, shown_case

FOOTER = ["Identification for Fire-Suppression Piping and Equipment", "2026-10-01", "21 05 53 - 1",
          "CONFIDENTIAL - DO NOT DISTRIBUTE"]


def _doc(footer_lines=FOOTER, *, footer_caps=False, title_style_caps=False, styles=0) -> bytes:
    """A spec whose footer shows ``footer_lines`` (the first one in capitals by formatting when ``footer_caps``) and
    whose title paragraph is typed in mixed case, in a style showing capitals when ``title_style_caps``; ``styles``
    extra styled paragraphs make the profile big enough for the size bound to thin it."""
    doc = Document()
    footer = doc.sections[0].footer
    footer.is_linked_to_previous = False
    run = footer.paragraphs[0].add_run(footer_lines[0])
    run.font.all_caps = footer_caps
    for line in footer_lines[1:]:
        footer.add_paragraph(line)
    doc.add_paragraph("SECTION 21 05 53")
    title = doc.styles.add_style("Spec Title", 1)
    title.font.all_caps = title_style_caps
    doc.add_paragraph("Identification for Fire Suppression Systems", style=title)
    doc.add_paragraph("PART 1 - GENERAL")                                     # typed capitals: never marked
    for i in range(styles):
        st = doc.styles.add_style(f"Custom Style Number {i}", 1)
        st.font.name = "Arial"
        st.font.size = Pt(10)
        doc.add_paragraph(f"{i}.01\tArticle text that is a fairly long sample sentence for style {i}.", style=st)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _pkg(data: bytes, name="21 05 53.docx") -> DocxPackage:
    return DocxPackage(data, name=name)


def _footer_view(prof):
    return next(v for k, v in prof["header_footer_parts"].items() if "footer" in k)


def test_typed_and_displayed_capitals_are_told_apart():
    plain = profile(_pkg(_doc()))
    assert _footer_view(plain)["text"] == FOOTER and "shown_in_capitals" not in _footer_view(plain)
    assert "first_paragraphs_shown_in_capitals" not in plain
    styled = profile(_pkg(_doc(footer_caps=True, title_style_caps=True)))
    assert _footer_view(styled)["text"][0] == FOOTER[0]                        # the text is still as typed
    assert _footer_view(styled)["shown_in_capitals"] == {"0": "all capitals"}
    # the title (its style shows capitals) is named; "SECTION 21 05 53" and "PART 1 - GENERAL" are typed capitals
    assert styled["first_paragraphs"][1] == "Identification for Fire Suppression Systems"
    assert styled["first_paragraphs_shown_in_capitals"] == {"1": "all capitals"}


def test_shown_case_names_partly_and_small_capitals():
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("Fire ").font.all_caps = True
    p.add_run("Suppression")
    q = doc.add_paragraph()
    q.add_run("Small Capitals").font.small_caps = True
    q.add_run(" 21 05 53")                                                    # no lowercase letter: no say
    buf = io.BytesIO()
    doc.save(buf)
    pkg = _pkg(buf.getvalue())
    sheet = StyleSheet(pkg)
    props = PropResolver(sheet)
    body = [x for x in pkg.xml(pkg.main_part()).iter(W + "p")]
    assert [shown_case(x, props) for x in body] == ["partly capitals", "small capitals"]


def test_a_thinned_footer_says_how_many_lines_it_left_out():
    full = profile(_pkg(_doc(styles=30)))
    assert _footer_view(full)["text"] == FOOTER and not full.get("truncated")
    unbounded = json.loads(json.dumps(full))
    # a limit that reaches the header/footer thinning but not the dropping of the text
    for limit in range(len(json.dumps(full, ensure_ascii=False)), 1000, -100):
        thinned = bound(json.loads(json.dumps(unbounded)), limit=limit)
        view = _footer_view(thinned)
        if view.get("text") and len(view["text"]) < len(FOOTER):
            break
    else:
        raise AssertionError("the bound never thinned the footer")
    assert view["text"] == FOOTER[:2] and view["more_lines"] == 2
    assert PART_LINES_CUT in thinned["truncated"]
    # once the text goes altogether, "header/footer text" says it, and the line count goes with the text
    dropped = bound(json.loads(json.dumps(unbounded)), limit=2500)
    assert "header/footer text" in dropped["truncated"] and PART_LINES_CUT not in dropped["truncated"]
    assert not {"text", "more_lines", "shown_in_capitals"} & set(_footer_view(dropped))


def test_a_long_footer_view_counts_the_lines_past_six():
    lines = [f"Footer line {i}" for i in range(9)]
    view = _footer_view(profile(_pkg(_doc(lines))))
    assert view["text"] == lines[:6] and view["more_lines"] == 3


def test_the_paragraph_tools_say_what_is_shown_in_capitals(tmp_path):
    path = tmp_path / "21 05 53.docx"
    path.write_bytes(_doc(footer_caps=True, title_style_caps=True))
    batch = _Batch([path], None)
    rows = batch.read_paragraphs("21 05 53.docx")["paragraphs"]
    assert [r.get("shown_in_capitals") for r in rows] == [None, "all capitals", None]
    footer = next(p for p in DocxPackage.from_path(path).text_parts() if "footer" in p)
    rows = batch.read_paragraphs("21 05 53.docx", part=footer)["paragraphs"]
    assert rows[0]["shown_in_capitals"] == "all capitals" and "shown_in_capitals" not in rows[1]
    sample = batch.sample_paragraphs("21 05 53.docx", regex="Identification")
    assert sample == [{"index": 1, "style": "SpecTitle", "text": "Identification for Fire Suppression Systems",
                       "shown_in_capitals": "all capitals"}]


def test_the_rich_fixture_still_fits_without_cuts():
    prof = profile(_pkg(build_rich(), name="rich.docx"))
    assert not prof.get("truncated")
