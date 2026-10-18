"""Approved effects stay constant while source evidence and concrete recipes vary."""
from __future__ import annotations

import copy
import hashlib

import pytest
from docx import Document
from lxml import etree

from spec_surgeon.editspec import apply_spec, plan_spec
from spec_surgeon.goal_contracts import GoalMethodRequired
from spec_surgeon.goals import (GoalError, NeedsInput, approved_goal_words, choose_goal, compile_goal_candidate,
                                goal_catalog, goal_templates, resolve_goal, validate_contract, verify_goal)
from spec_surgeon.package import DocxPackage, W
from spec_surgeon.verify import run_suite


def contract(*effects):
    return {"kind": "document_effects", "effects": list(effects)}


def exact(text, **extra):
    return {"locators": [{"exact_text": text, **extra}]}


def document(path, layout="manual", title="FIRE SUPPRESSION"):
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "Original project header"
    if layout == "table":
        cells = doc.add_table(rows=1, cols=2).rows[0].cells
        cells[0].text = "SECTION 21 13 13"
        paragraph = cells[1].paragraphs[0]
    elif layout == "inline":
        paragraph = doc.add_paragraph()
        paragraph.add_run("SECTION 21 13 13 — ")
    else:
        doc.add_paragraph("SECTION 21 13 13")
        paragraph = doc.add_paragraph(style="Heading 1" if layout == "styled" else "Normal")
    paragraph.add_run(title[:5]).italic = True
    paragraph.add_run(title[5:]).bold = False
    body = doc.add_paragraph()
    body.add_run("Ow")
    body.add_run("ner shall retain ").italic = True
    body.add_run("all content.")
    doc.add_paragraph("Untouched lexical  spacing and NFPA 13.")
    doc.save(path)
    return path


def execute(source, goal, binding=None):
    binding = binding or resolve_goal(source, goal)["binding"]
    candidate = compile_goal_candidate(source, goal, binding)
    stages = ([{"spec": candidate["spec"]}] if candidate["kind"] == "spec" else candidate["chain"]["stages"])
    data = source.read_bytes()
    for stage in stages:
        pkg = DocxPackage(data)
        plan, fragments = plan_spec(pkg, stage["spec"])
        output = apply_spec(pkg, fragments).output_bytes()
        record = run_suite(data, output, plan.to_dict(), str(source), "candidate.docx")
        assert record.passed, record.to_markdown()
        data = output
    verdict = verify_goal(source, data, goal, binding)
    assert verdict["passed"], verdict
    return data


@pytest.mark.parametrize("layout", ["manual", "styled", "table", "inline"])
def test_semantic_role_varies_with_source_and_composes_with_font_text_pages(tmp_path, layout):
    source = document(tmp_path / "spec.docx", layout, "ELECTRICAL REQUIREMENTS")
    goal = contract(
        {"id": "body-font", "type": "format", "target": {"all": True}, "font": {"name": "Arial", "size": 11}},
        {"id": "title", "type": "format", "target": {"locators": [{"role": "section_title", "style_hint": "WrongStyle",
             "section_number": "21 13 13"}]}, "run": {"b": True}, "paragraph": {"keepNext": True}},
        {"id": "wording", "type": "replace_text", "target": {"all": True}, "find": "Owner", "replace": "Client",
         "expected_matches": 1, "tracked": True},
        {"id": "page", "type": "page_setup", "margins": {"left": 1440}, "orientation": "landscape"})
    resolution = resolve_goal(source, goal)
    assert resolution["status"] == "resolved"
    title = resolution["binding"]["effects"]["title"]["targets"][0]
    assert "ELECTRICAL REQUIREMENTS" in title["text"]
    assert title["in_table"] == (layout == "table")
    result = execute(source, goal)
    assert result != source.read_bytes()


def test_multiple_approved_locators_require_user_choice_and_keep_cardinality(tmp_path):
    source = document(tmp_path / "spec.docx")
    goal = contract({"id": "title", "type": "format", "target": {"locators": [
        {"role": "section_title"}, {"contains_text": "Untouched"}], "cardinality": 1}, "run": {"b": True}})
    result = resolve_goal(source, goal)
    assert result["binding"] is None and len(result["candidates"]) == 2
    with pytest.raises(NeedsInput):
        compile_goal_candidate(source, goal)
    chosen = result["effects"]["title"]["candidates"][1]
    binding = choose_goal(source, goal, [chosen["candidate_id"]], effect_id="title")
    execute(source, goal, binding)
    with pytest.raises(GoalError, match="cardinality"):
        choose_goal(source, goal, [c["candidate_id"] for c in result["candidates"]], effect_id="title")
    tampered = copy.deepcopy(binding)
    tampered["effects"]["title"]["targets"][0]["in_table"] = 0
    with pytest.raises(GoalError, match="evidence"):
        compile_goal_candidate(source, goal, tampered)


def test_scoped_anchors_are_exact_and_explicit_style_is_a_predicate(tmp_path):
    source = document(tmp_path / "spec.docx")
    good = contract({"id": "title", "type": "format", "target": exact("FIRE SUPPRESSION", after_text="SECTION 21 13 13",
                    before_text="Owner shall retain all content."), "run": {"b": True}})
    execute(source, good)
    bad = copy.deepcopy(good)
    bad["effects"][0]["target"]["locators"][0]["after_text"] = "SECTION 211313"
    assert resolve_goal(source, bad)["status"] == "needs_input"
    bad["effects"][0]["target"] = {"locators": [{"style_id": "Heading1"}]}
    assert resolve_goal(source, bad)["status"] == "needs_input"


def test_global_format_of_absent_optional_parts_is_a_verified_noop(tmp_path):
    source = document(tmp_path / "spec.docx")
    goal = contract({"id": "footnotes", "type": "format", "target": {"scope": ["footnotes"], "all": True},
                     "font": {"name": "Arial", "size": 11}})
    assert resolve_goal(source, goal)["status"] == "resolved"
    assert execute(source, goal) == source.read_bytes()
    goal["effects"][0]["target"] = {"scope": ["footnotes"], "locators": [{"exact_text": "Missing wording"}],
                                     "cardinality": "all"}
    assert resolve_goal(source, goal)["status"] == "needs_input"


@pytest.mark.parametrize("tracked", [False, True])
@pytest.mark.parametrize("kind", ["insert_paragraph", "delete_paragraphs"])
def test_exact_structural_effects_preserve_other_paragraphs(tmp_path, tracked, kind):
    source = document(tmp_path / "spec.docx")
    effect = {"id": "content", "type": kind, "target": exact("Owner shall retain all content."), "tracked": tracked}
    if kind == "insert_paragraph":
        effect.update(text="Approved new obligation.", position="after", formatting="inherit")
    execute(source, contract(effect))


@pytest.mark.parametrize("tracked", [False, True])
def test_exact_replacement_supports_split_runs_and_empty_literal(tmp_path, tracked):
    source = document(tmp_path / "spec.docx")
    goal = contract({"id": "wording", "type": "replace_text", "target": {"all": True}, "find": "Owner",
                     "replace": "", "expected_matches": 1, "tracked": tracked})
    execute(source, goal)
    changed = copy.deepcopy(goal)
    changed["effects"][0]["expected_matches"] = 2
    with pytest.raises(NeedsInput, match="found 1"):
        compile_goal_candidate(source, changed)


def test_bound_goal_fails_closed_for_stale_source_and_collateral_mutation(tmp_path):
    source = document(tmp_path / "spec.docx")
    goal = contract({"id": "title", "type": "format", "target": {"locators": [{"role": "section_title"}]}, "run": {"b": True}})
    bound = resolve_goal(source, goal)["binding"]
    data = execute(source, goal, bound)
    bad = DocxPackage(data)
    root = bad.xml(bad.main_part())
    texts = list(root.iter(W + "t"))
    texts[-1].text = "Unauthorized changed words."
    bad.write_xml(bad.main_part(), root)
    assert not verify_goal(source, bad.output_bytes(), goal, bound)["passed"]
    source.write_bytes(data)
    with pytest.raises(GoalError, match="changed"):
        compile_goal_candidate(source, goal, bound)


def test_a_self_consistent_wrong_target_method_does_not_satisfy_the_goal(tmp_path, monkeypatch):
    from spec_surgeon import goal_contracts
    source = document(tmp_path / "spec.docx")
    goal = contract({"id": "title", "type": "format", "target": {"locators": [{"role": "section_title"}]},
                     "run": {"b": True}})
    binding = resolve_goal(source, goal)["binding"]
    # Both plan and apply agree on this defective method. The original goal must still catch it.
    monkeypatch.setattr(goal_contracts, "_selection", lambda pkg, candidate:
                        {"scope": ["document"], "paragraphs": {"index_range": [2, 3]}})
    method = compile_goal_candidate(source, goal, binding)["spec"]
    pkg = DocxPackage(source.read_bytes())
    plan, fragments = plan_spec(pkg, method)
    output = apply_spec(pkg, fragments).output_bytes()
    assert run_suite(source.read_bytes(), output, plan.to_dict(), "source", "candidate").passed
    verdict = verify_goal(source, output, goal, binding)
    assert not verdict["passed"]
    assert any("format" in error or "effective" in error or "residual" in error for error in verdict["errors"])


def test_missing_inherited_size_is_caught_even_when_plan_and_apply_share_the_bug(tmp_path, monkeypatch):
    from spec_surgeon.ops import set_font
    source = document(tmp_path / "spec.docx")
    goal = contract({"id": "font", "type": "format", "target": {"all": True}, "font": {"size": 13.5}})
    binding = resolve_goal(source, goal)["binding"]
    expected = set_font.size_applies
    monkeypatch.setattr(set_font, "size_applies", lambda own, size, missing=False: expected(own, size, False))
    method = compile_goal_candidate(source, goal, binding)["spec"]
    pkg = DocxPackage(source.read_bytes())
    plan, fragments = plan_spec(pkg, method)
    output = apply_spec(pkg, fragments).output_bytes()
    record = run_suite(source.read_bytes(), output, plan.to_dict(), "source", "candidate")
    assert record.get("V10").status == "FAIL"
    assert not verify_goal(source, output, goal, binding)["passed"]


def template_effect(path):
    return {"id": "branding", "type": "header_footer_template", "template": str(path),
            "template_sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
            "params": {"template_sections": {"all": 0}, "slot_fills": {"TITLE": {
                "rule": "paragraph_after", "anchor_regex": "^SECTION "}}}}


def test_template_body_font_title_and_page_compose_from_pinned_inputs(tmp_path):
    source = document(tmp_path / "spec.docx")
    template = tmp_path / "template.docx"
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "Project template — {{TITLE}}"
    doc.sections[0].footer.paragraphs[0].text = "Approved project footer"
    doc.save(template)
    goal = contract(template_effect(template),
        {"id": "body", "type": "format", "target": {"all": True}, "font": {"name": "Arial", "size": 11}},
        {"id": "title", "type": "format", "target": {"locators": [{"role": "section_title"}]}, "run": {"b": True}},
        {"id": "page", "type": "page_setup", "margins": {"left": 1440, "right": 1440}})
    assert goal_templates(goal) == [template]
    execute(source, goal)
    doc.add_paragraph("Template changed after approval")
    doc.save(template)
    with pytest.raises(GoalError, match="template bytes"):
        resolve_goal(source, goal)


def test_template_style_import_cannot_redefine_an_unresolved_original_body_reference(tmp_path):
    from docx.enum.style import WD_STYLE_TYPE
    source = document(tmp_path / "spec.docx")
    package = DocxPackage(source.read_bytes())
    root = package.xml(package.main_part())
    body = list(root.iter(W + "p"))[-1]
    properties = etree.Element(W + "pPr")
    etree.SubElement(properties, W + "pStyle").set(W + "val", "ImportedHeader")
    body.insert(0, properties)
    package.write_xml(package.main_part(), root)
    source.write_bytes(package.output_bytes())
    template = tmp_path / "template.docx"
    doc = Document()
    style = doc.styles.add_style("ImportedHeader", WD_STYLE_TYPE.PARAGRAPH)
    style.font.bold = True
    paragraph = doc.sections[0].header.paragraphs[0]
    paragraph.style = style
    paragraph.text = "Project template"
    doc.save(template)
    effect = template_effect(template)
    effect["params"]["slot_fills"] = {}
    goal = contract(effect)
    binding = resolve_goal(source, goal)["binding"]
    method = compile_goal_candidate(source, goal, binding)["spec"]
    package = DocxPackage(source.read_bytes())
    _, fragments = plan_spec(package, method)
    output = apply_spec(package, fragments).output_bytes()
    verdict = verify_goal(source, output, goal, binding)
    assert verdict["passed"], verdict
    safe = DocxPackage(output)
    style_root = safe.xml("word/styles.xml")
    assert not any(s.get(W + "styleId") == "ImportedHeader" for s in style_root)
    imported = next(s for s in style_root if s.get(W + "styleId") == "TemplateImportedHeader")
    original_body = list(package.xml(package.main_part()).iter(W + "p"))[-1]
    final_body = list(safe.xml(safe.main_part()).iter(W + "p"))[-1]
    assert etree.tostring(original_body) == etree.tostring(final_body)
    # A faulty method can still recreate the old bug. Independent verification
    # must reject that output even though the correct method can now recover.
    imported.set(W + "styleId", "ImportedHeader")
    safe.write_xml("word/styles.xml", style_root)
    for part in safe.parts_by_kind()["headers"]:
        root = safe.xml(part)
        for ref in root.iter(W + "pStyle"):
            if ref.get(W + "val") == "TemplateImportedHeader":
                ref.set(W + "val", "ImportedHeader")
        safe.write_xml(part, root)
    verdict = verify_goal(source, safe.output_bytes(), goal, binding)
    assert not verdict["passed"]
    assert any("original" in error and "style" in error for error in verdict["errors"]), verdict


def test_header_footer_removal_keeps_body_and_validates_whole_graph(tmp_path):
    source = document(tmp_path / "spec.docx")
    execute(source, contract({"id": "remove", "type": "remove_header_footer", "kinds": ["header"]},
        {"id": "title", "type": "format", "target": {"locators": [{"role": "section_title"}]}, "run": {"b": True}}))


def test_approved_words_do_not_include_notes_find_text_or_extraction_patterns(tmp_path):
    template = tmp_path / "template.docx"
    Document().save(template)
    effect = template_effect(template)
    effect["params"]["slot_fills"] = {"TITLE": {"rule": "first_match", "rules": [
        {"rule": "regex", "pattern": "Original source phrase"}, {"rule": "literal", "value": "Approved footer"}]}}
    goal = contract(effect, {"id": "replace", "type": "replace_text", "target": {"all": True},
        "find": "Original old wording", "replace": "Approved wording", "expected_matches": 1},
        {"id": "insert", "type": "insert_paragraph", "target": exact("anchor"), "text": "Approved insertion"})
    goal["request"] = "Notes are not literal approval."
    assert set(approved_goal_words(goal)) == {"Approved footer", "Approved wording", "Approved insertion"}
    assert set(goal_catalog()["effects"]) >= {"format", "replace_text", "page_setup", "insert_paragraph", "delete_paragraphs",
                                              "header_footer_template", "remove_header_footer"}


@pytest.mark.parametrize("effect", [
    {"id": "a", "type": "format", "target": {"all": True}, "font": {"size": 10.25}},
    {"id": "a", "type": "replace_text", "target": {"all": True}, "find": "old", "replace": "new"},
    {"id": "a", "type": "page_setup", "margins": {"left": -1}},
    {"id": "a", "type": "format", "target": {"locators": [{"role": "guess_any_title"}]}, "run": {"b": True}},
    {"id": "a", "type": "header_footer_template", "template": "/x.docx", "params": {"template_sections": {"all": 0}, "slot_fills": {}}},
])
def test_goal_grammar_refuses_unapproved_or_underspecified_effects(effect):
    with pytest.raises(GoalError):
        validate_contract(contract(effect))
