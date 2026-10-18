"""Real persisted project workflows; only model answers are scripted."""
import json
import os
from pathlib import Path

import pytest

from fake_claude import FakeClient, resp, tool
from make_fixtures import build
from spec_surgeon.campaign import create_campaign, load_campaign
from spec_surgeon.recovery import RecoveryOptions


def opts():
    return RecoveryOptions(render=False, use_memory=False)


def effect_goal(left=1440):
    return {"kind": "document_effects", "effects": [
        {"id": "page", "type": "page_setup", "sections": "all", "margins": {"left": left}}]}


def proposal(contract=None):
    return resp(tool("propose_plan", {"contract": contract or effect_goal(),
                "rationale": "Set the requested left margin in every selected section.",
                "risks": ["Pagination may change."], "evidence": ["The request specifies one inch."]}))


def adaptive_submit(source, contract=None):
    from spec_surgeon.goals import compile_goal_candidate
    return resp(tool("submit_repair", {"candidate": compile_goal_candidate(source, contract or effect_goal())}))


def critic_accept():
    return resp(tool("review_repair", {"verdict": "accept", "reason": "The checked source-to-output margin matches the approved goal."}))


def doc(folder, name="21 13 13.docx"):
    folder.mkdir(parents=True, exist_ok=True)
    result = folder / name
    result.write_bytes(build("21 13 13", "SPRINKLER SYSTEMS"))
    return result


def goal():
    return {"request": "Set left margin to one inch", "summary": "Margins", "mode": "formatting",
            "ops": [{"op": "set_page_setup", "params": {"margins": {"left": 1440}}}]}


def goal_needing_repair():
    """A saved recipe's mistaken count needs repair without changing its approved effect."""
    spec = goal()
    spec["ops"][0]["expect"] = {"text_diffs": 1}
    return spec


def submit():
    return resp(tool("submit_repair", {"candidate": {"kind": "spec", "spec": goal()}}))


def make(tmp_path, request=""):
    root = tmp_path / "jobs"
    for name in ("north", "south"):
        folder = root / name
        doc(folder)
        (folder / "PROJECT.json").write_text(json.dumps({"name": name, "stamp": name.upper()}))
    return create_campaign([str(root)], tmp_path / "workspace", request)


def test_restart_and_context_changes_invalidate_only_affected_project(tmp_path):
    campaign = make(tmp_path)
    north, south = campaign.projects
    campaign.set_goal(north, goal()["request"], spec=goal())
    loaded = load_campaign(campaign.out_dir)
    assert loaded.projects[north]["goal"]["spec"] == goal()
    assert not loaded.projects[south]["goal"]
    loaded.set_project_context(north, {"name": "North", "stamp": "NORTH REVISED"})
    assert loaded.projects[south]["context"]["stamp"] == "SOUTH"
    assert [e["action"] for e in map(json.loads, (campaign.out_dir / "events.jsonl").read_text().splitlines())] == [
        "created", "set_goal", "set_project_context"]
    with pytest.raises(ValueError, match="another window"):
        campaign.set_project_context(south, {"name": "Stale writer"})


def test_project_assignment_and_scoped_answers_survive_reload(tmp_path):
    campaign = make(tmp_path)
    north, south = campaign.projects
    doc_id = campaign.projects[north]["document_ids"][0]
    campaign.set_project([doc_id], "new-project")
    campaign.set_goal("new-project", "Format the title")
    campaign.answer("new-project", "Keep the project name as NORTH", files=[doc_id])
    loaded = load_campaign(campaign.out_dir)
    assert loaded.documents[doc_id]["project_id"] == "new-project"
    assert loaded.projects["new-project"]["decisions"][0]["document_ids"] == [doc_id]
    assert not loaded.projects[south]["decisions"]
    with pytest.raises(ValueError, match="only to documents"):
        loaded.answer(south, "Wrong scope", files=[doc_id])


def test_real_project_trials_keep_context_and_word_sources_separate_then_export_after_restart(tmp_path):
    campaign = make(tmp_path)
    for pid in campaign.projects:
        campaign.set_goal(pid, goal()["request"], spec=goal_needing_repair())
    client = FakeClient(submit(), submit())
    campaign.run(options=opts(), client=client)
    assert all(d["status"] == "READY" for d in campaign.documents.values()), campaign.to_text()
    requests = [json.dumps(r["messages"]) for r in client.requests]
    assert len(requests) == 2
    assert "NORTH" in requests[0] and "SOUTH" not in requests[0]
    assert "SOUTH" in requests[1] and "NORTH" not in requests[1]
    loaded = load_campaign(campaign.out_dir)
    loaded.run(options=opts(), client=FakeClient())  # Already verified: no model work on restart.
    exported = loaded.export(tmp_path / "delivery")
    outputs = list(exported.rglob("*.docx"))
    assert len(outputs) == 2
    assert (exported / "export.json").is_file()
    assert all(d["status"] == "EXPORTED" for d in loaded.documents.values())


def test_same_stat_source_drift_blocks_only_that_project(tmp_path):
    campaign = make(tmp_path)
    north, south = campaign.projects
    for pid in campaign.projects:
        campaign.set_goal(pid, goal()["request"], spec=goal())
    path = next(d["source"] for d in campaign.documents.values() if d["project_id"] == north)
    from pathlib import Path
    path = Path(path)
    stat = path.stat()
    data = bytearray(path.read_bytes())
    data[-1] ^= 1
    path.write_bytes(data)
    os.utime(path, ns=(stat.st_atime_ns, stat.st_mtime_ns))
    campaign.run(options=opts(), client=FakeClient(submit()))
    assert campaign.projects[north]["status"] == "DRIFTED"
    assert campaign.projects[south]["status"] == "READY", campaign.to_text()


def test_ready_project_is_retained_while_another_waits_for_answer(tmp_path):
    campaign = make(tmp_path)
    for pid in campaign.projects:
        campaign.set_goal(pid, goal()["request"], spec=goal_needing_repair())
    campaign.run(options=opts(), client=FakeClient(submit(), resp(tool("need_user", {"question": "Same margin here?"}))))
    north, south = campaign.projects
    assert campaign.projects[north]["status"] == "READY"
    assert campaign.projects[south]["status"] == "NEEDS_INPUT"
    output = next(d["output"] for d in campaign.documents.values() if d["project_id"] == north)
    loaded = load_campaign(campaign.out_dir)
    loaded.answer(south, "Yes, the same one-inch left margin.")
    loaded.run(options=opts(), client=FakeClient(submit()))
    assert all(d["status"] == "READY" for d in loaded.documents.values()), loaded.to_text()
    assert next(d["output"] for d in loaded.documents.values() if d["project_id"] == north) == output


def test_tampered_output_is_never_exported(tmp_path):
    campaign = make(tmp_path)
    pid = next(iter(campaign.projects))
    campaign.set_goal(pid, goal()["request"], spec=goal())
    campaign.run(options=opts(), client=FakeClient(submit()))
    from pathlib import Path
    path = Path(next(d["output"] for d in campaign.documents.values() if d["project_id"] == pid))
    path.write_bytes(path.read_bytes() + b"changed")
    destination = tmp_path / "delivery"
    with pytest.raises(ValueError):
        load_campaign(campaign.out_dir).export(destination)
    assert not destination.exists()


def test_campaign_load_refuses_duplicate_membership_or_external_artifacts(tmp_path):
    campaign = make(tmp_path)
    state = campaign.to_dict()
    north, south = state["projects"]
    state["projects"][south]["document_ids"].append(state["projects"][north]["document_ids"][0])
    (campaign.out_dir / "campaign.json").write_text(json.dumps(state))
    with pytest.raises(ValueError, match="identity or project membership"):
        load_campaign(campaign.out_dir)


def test_adaptive_goal_reaches_engine_and_survives_reload_without_model_calls(tmp_path):
    source = doc(tmp_path / "source")
    campaign = create_campaign([str(source)], tmp_path / "workspace")
    pid = next(iter(campaign.projects))
    contract = {"kind": "format_section_title", "target": {"title": "SPRINKLER SYSTEMS"}, "run": {"b": True}}
    campaign.set_goal(pid, "Make SPRINKLER SYSTEMS bold", adaptive_contract=contract)
    campaign.run(options=opts(), client=FakeClient())
    assert campaign.projects[pid]["status"] == "READY", campaign.to_text()
    loaded = load_campaign(campaign.out_dir)
    destination = loaded.export(tmp_path / "delivery")
    assert len(list(destination.rglob("*.docx"))) == 1
    assert next(iter(loaded.documents.values()))["adaptive_verification"]["passed"]


def test_question_answer_defaults_to_only_the_files_the_question_names(tmp_path):
    folder = tmp_path / "source"
    first = doc(folder, "a.docx")
    second = doc(folder, "b.docx")
    second.write_bytes(build("21 13 13", "SPRINKLER SYSTEMS", single_section=True))
    campaign = create_campaign([str(folder)], tmp_path / "workspace")
    pid = next(iter(campaign.projects))
    campaign.set_goal(pid, goal()["request"], spec=goal_needing_repair())
    campaign.run(options=opts(), client=FakeClient(resp(tool("need_user", {"question": "One inch for this layout?"})), submit()))
    project = campaign.projects[pid]
    question = project["questions"][0]
    first_id = next(d["id"] for d in campaign.documents.values() if d["source"] == str(first))
    assert question["id"] == "q0001" and question["document_ids"] == [first_id]
    with pytest.raises(ValueError, match="affected document scope"):
        campaign.answer(pid, "Yes", question_id=question["id"], files=project["document_ids"])
    campaign.answer(pid, "Yes, one inch.", question_id=question["id"])
    assert project["decisions"][-1]["document_ids"] == [first_id]
    loaded = load_campaign(campaign.out_dir)
    loaded.run(options=opts(), client=FakeClient(submit()))
    assert all(d["status"] == "READY" for d in loaded.documents.values()), loaded.to_text()


def test_changed_project_metadata_requires_explicit_context_update(tmp_path):
    campaign = make(tmp_path)
    north, south = campaign.projects
    for pid in campaign.projects:
        campaign.set_goal(pid, goal()["request"], spec=goal())
    from pathlib import Path
    Path(campaign.projects[north]["context_source"]).write_text('{"stamp":"UNREVIEWED CHANGE"}')
    campaign.run(options=opts(), client=FakeClient(submit()))
    assert campaign.projects[north]["status"] == "DRIFTED"
    assert "PROJECT.json changed" in campaign.projects[north]["error"]
    assert campaign.projects[south]["status"] == "READY"
    campaign.set_project_context(north, {"stamp": "APPROVED CHANGE"})
    campaign.run(options=opts(), client=FakeClient(submit()))
    assert campaign.projects[north]["status"] == "READY", campaign.to_text()


def test_ambiguous_adaptive_target_waits_while_other_file_finishes_then_resumes_user_choice(tmp_path):
    from docx import Document
    folder = tmp_path / "source"
    first = doc(folder, "a.docx")
    document = Document(first)
    document.add_paragraph("SPRINKLER SYSTEMS")
    document.save(first)
    doc(folder, "b.docx")
    campaign = create_campaign([str(folder)], tmp_path / "workspace")
    pid = next(iter(campaign.projects))
    contract = {"kind": "format_section_title", "target": {"title": "SPRINKLER SYSTEMS"}, "run": {"b": True}}
    campaign.set_goal(pid, "Make the section title SPRINKLER SYSTEMS bold", adaptive_contract=contract)
    campaign.run(options=opts(), client=FakeClient())
    records = list(campaign.documents.values())
    assert [d["status"] for d in records] == ["NEEDS_INPUT", "READY"], campaign.to_text()
    good_output = records[1]["output"]
    question = campaign.projects[pid]["questions"][0]
    campaign.bind(question["document_id"], question["candidates"][0]["candidate_id"])
    loaded = load_campaign(campaign.out_dir)
    loaded.run(options=opts(), client=FakeClient())
    assert all(d["status"] == "READY" for d in loaded.documents.values()), loaded.to_text()
    assert loaded.documents[records[1]["id"]]["output"] == good_output
    loaded.export(tmp_path / "delivery")
    assert len(list((tmp_path / "delivery").rglob("*.docx"))) == 2


def test_adaptive_compile_error_isolated_per_document(tmp_path, monkeypatch):
    from spec_surgeon import goals
    folder = tmp_path / "source"
    first = doc(folder, "a.docx")
    doc(folder, "b.docx")
    campaign = create_campaign([str(folder)], tmp_path / "workspace")
    pid = next(iter(campaign.projects))
    contract = {"kind": "format_section_title", "target": {"title": "SPRINKLER SYSTEMS"}, "run": {"b": True}}
    campaign.set_goal(pid, "Make the title bold", adaptive_contract=contract)
    real_compile = goals.compile_goal

    def compile_one(source, contract, binding):
        if source == first:
            raise ValueError("This source needs an unsupported structural repair")
        return real_compile(source, contract, binding)

    monkeypatch.setattr(goals, "compile_goal", compile_one)
    campaign.run(options=opts(), client=FakeClient())
    assert [d["status"] for d in campaign.documents.values()] == ["ERROR", "READY"], campaign.to_text()


def test_interrupted_adaptive_preparation_does_not_poison_later_export(tmp_path, monkeypatch):
    from spec_surgeon import recovery
    source = doc(tmp_path / "source")
    campaign = create_campaign([str(source)], tmp_path / "workspace")
    pid = next(iter(campaign.projects))
    contract = {"kind": "format_section_title", "target": {"title": "SPRINKLER SYSTEMS"}, "run": {"b": True}}
    campaign.set_goal(pid, "Make the title bold", adaptive_contract=contract)
    real_prepare = recovery.prepare_recovery

    def interrupted(*args, **kwargs):
        raise OSError("Interrupted before checkpoint")

    monkeypatch.setattr(recovery, "prepare_recovery", interrupted)
    campaign.run(options=opts(), client=FakeClient())
    assert next(iter(campaign.documents.values()))["status"] == "ERROR"
    monkeypatch.setattr(recovery, "prepare_recovery", real_prepare)
    loaded = load_campaign(campaign.out_dir)
    loaded.run(options=opts(), client=FakeClient())
    assert len(loaded.projects[pid]["runs"]) == 2
    assert next(iter(loaded.documents.values()))["status"] == "READY", loaded.to_text()
    assert len(list(loaded.export(tmp_path / "delivery").rglob("*.docx"))) == 1


def test_template_inside_input_folder_is_a_reference_not_unfinished_work(tmp_path):
    folder = tmp_path / "source"
    doc(folder)
    template = folder / "template.docx"
    template.write_bytes(build("21 05 00", "TEMPLATE REFERENCE"))
    campaign = create_campaign([str(folder)], tmp_path / "workspace")
    pid = next(iter(campaign.projects))
    campaign.set_goal(pid, goal()["request"], spec=goal(), template=template)
    campaign.run(options=opts(), client=FakeClient(submit()))
    assert campaign.projects[pid]["status"] == "READY", campaign.to_text()
    assert sorted(d["status"] for d in campaign.documents.values()) == ["READY", "REFERENCE"]


def test_advisory_campaign_lock_rejects_simultaneous_handles_and_stays_reusable(tmp_path):
    campaign = make(tmp_path)
    second = load_campaign(campaign.out_dir)
    pid = next(iter(campaign.projects))
    with campaign._lock():
        with pytest.raises(ValueError, match="another worker"):
            second.set_project_context(pid, {"stamp": "BLOCKED"})
        with campaign._lock():  # The owning operation may checkpoint while holding its lock.
            assert (campaign.out_dir / ".campaign.lock").exists()
    assert (campaign.out_dir / ".campaign.lock").exists()  # Never unlink an advisory lock inode.
    fresh = load_campaign(campaign.out_dir)
    assert fresh.projects[pid]["context"]["stamp"] != "BLOCKED"
    fresh.set_project_context(pid, {"stamp": "AFTER RELEASE"})
    assert load_campaign(campaign.out_dir).projects[pid]["context"]["stamp"] == "AFTER RELEASE"


def test_campaign_advisory_lock_releases_after_abrupt_subprocess_exit(tmp_path):
    import subprocess
    import sys
    from pathlib import Path
    campaign = make(tmp_path)
    code = ("import os, sys\nfrom spec_surgeon.campaign import load_campaign\n"
            "campaign = load_campaign(sys.argv[1])\nwith campaign._lock():\n    os._exit(23)\n")
    environment = dict(os.environ, PYTHONPATH=str(Path(__file__).resolve().parents[1] / "src"))
    stopped = subprocess.run([sys.executable, "-c", code, str(campaign.out_dir)], env=environment,
                             capture_output=True, text=True, timeout=20)
    assert stopped.returncode == 23, stopped.stderr
    loaded = load_campaign(campaign.out_dir)
    pid = next(iter(loaded.projects))
    loaded.set_project_context(pid, {"stamp": "RECOVERED AFTER EXIT"})
    assert load_campaign(campaign.out_dir).projects[pid]["context"]["stamp"] == "RECOVERED AFTER EXIT"


def test_each_adaptive_question_is_checkpointed_and_retained_across_other_files(tmp_path, monkeypatch):
    from docx import Document
    folder = tmp_path / "source"
    for name in ("a.docx", "b.docx", "c.docx"):
        source = doc(folder, name)
        if name != "b.docx":
            document = Document(source)
            document.add_paragraph("SPRINKLER SYSTEMS")
            document.save(source)
    campaign = create_campaign([str(folder)], tmp_path / "workspace")
    pid = next(iter(campaign.projects))
    ids = campaign.projects[pid]["document_ids"]
    contract = {"kind": "format_section_title", "target": {"title": "SPRINKLER SYSTEMS"}, "run": {"b": True}}
    campaign.set_goal(pid, "Make the title bold", adaptive_contract=contract)
    saved = campaign._save
    snapshots = []

    def interrupt_after_second_file(action, details=None):
        saved(action, details)
        if action == "initial_document_prepared":
            snapshots.append(load_campaign(campaign.out_dir).projects[pid]["questions"])
            if details["document_id"] == ids[1]:
                raise KeyboardInterrupt("Process stopped after the second document")

    monkeypatch.setattr(campaign, "_save", interrupt_after_second_file)
    with pytest.raises(KeyboardInterrupt):
        campaign.run(options=opts(), client=FakeClient())
    assert len(snapshots) == 2
    assert all([q["document_ids"] for q in questions] == [[ids[0]]] for questions in snapshots)
    loaded = load_campaign(campaign.out_dir)
    good_output = loaded.documents[ids[1]]["output"]
    loaded.run(options=opts(), client=FakeClient())
    assert [q["document_ids"] for q in loaded.projects[pid]["questions"]] == [[ids[0]], [ids[2]]]
    assert loaded.documents[ids[1]]["output"] == good_output
    assert loaded.documents[ids[1]]["status"] == "READY"
    loaded.export(tmp_path / "ready-delivery")
    reopened = load_campaign(loaded.out_dir)
    assert [q["document_ids"] for q in reopened.projects[pid]["questions"]] == [[ids[0]], [ids[2]]]
    assert reopened.documents[ids[1]]["status"] == "EXPORTED"


def test_campaign_resume_reuses_saved_resource_policy_and_validates_before_saving(tmp_path, monkeypatch):
    from dataclasses import asdict
    from spec_surgeon import recovery
    source = doc(tmp_path / "source")
    campaign = create_campaign([str(source)], tmp_path / "workspace")
    pid = next(iter(campaign.projects))
    campaign.set_goal(pid, goal()["request"], spec=goal_needing_repair())
    # Leave room for the full operation catalog; the one-call limit is what
    # intentionally pauses this test before the successful continuation.
    policy = RecoveryOptions(max_attempts=2, max_calls=1, max_cost_usd=12.5, max_output_tokens=8000,
                             max_total_tokens=500000, effort="high", render=False, use_memory=False)
    inspection_client = FakeClient(resp(tool("list_styles", {"file": source.name})))
    campaign.run(options=policy, client=inspection_client)
    assert len(inspection_client.requests) == 1
    run_dir = Path(campaign.projects[pid]["runs"][-1]["path"])
    paused = recovery.load_recovery(run_dir)
    assert paused.budget["calls"] == 1 and "call limit" in paused.paused_reason
    loaded = load_campaign(campaign.out_dir)
    assert loaded.to_dict()["resource_options"] == asdict(policy)
    real_resume = recovery.resume_recovery
    observed = []

    def resume(*args, **kwargs):
        observed.append(asdict(kwargs["options"]))
        return real_resume(*args, **kwargs)

    monkeypatch.setattr(recovery, "resume_recovery", resume)
    repair_client = FakeClient(submit())
    loaded.run(client=repair_client)
    assert len(repair_client.requests) == 1
    assert observed == [asdict(policy)]
    assert next(iter(loaded.documents.values()))["status"] == "READY", loaded.to_text()
    before = (loaded.out_dir / "campaign.json").read_bytes()
    with pytest.raises(ValueError):
        loaded.run(options=RecoveryOptions(max_cost_usd=-1), client=FakeClient())
    assert (loaded.out_dir / "campaign.json").read_bytes() == before


def test_planner_revises_more_than_three_invalid_proposals_then_requires_review(tmp_path):
    source = doc(tmp_path / "source")
    campaign = create_campaign([str(source)], tmp_path / "workspace")
    pid = next(iter(campaign.projects))
    client = FakeClient(*[proposal({"kind": "unknown", "attempt": i}) for i in range(4)], proposal())
    campaign.plan(pid, "Set every left margin to one inch.", options=opts(), client=client)
    plan = campaign.projects[pid]["proposed_plan"]
    assert plan["status"] == "PROPOSED", campaign.to_text()
    assert len(client.requests) == 5
    assert len(campaign.projects[pid]["planning_state"]["attempts"]) == 4
    assert not campaign.projects[pid]["runs"]
    assert '"left": 1440' in campaign.to_text()
    campaign.run(client=FakeClient())
    assert not campaign.projects[pid]["runs"]
    campaign.approve_plan(pid, plan["id"])
    campaign.run(options=opts(), client=FakeClient(adaptive_submit(source), critic_accept()))
    assert next(iter(campaign.documents.values()))["status"] == "READY", campaign.to_text()
    restored = load_campaign(campaign.out_dir)
    restored.run(client=FakeClient())
    assert restored.to_dict()["projects"][pid]["resource_totals"]["calls"] == 5
    assert len(list(restored.export(tmp_path / "delivery").rglob("*.docx"))) == 1


def test_planner_transport_pause_and_scoped_question_resume_preserve_evidence(tmp_path):
    from fake_claude import api_error, raise_on_final, thinking
    campaign = make(tmp_path)
    north, south = campaign.projects
    did = campaign.projects[north]["document_ids"][0]
    question = resp(thinking("signed-thought"), tool("need_user", {"question": "Which margin?", "document_ids": [did]}))
    campaign.plan(north, "Fix the margins", options=opts(), client=FakeClient(question))
    first = campaign.projects[north]["proposed_plan"]
    assert first["status"] == "NEEDS_INPUT"
    campaign.answer(north, "One inch on the left", first["questions"][0]["id"])
    campaign.plan(north, options=opts(), client=FakeClient(raise_on_final(api_error())))
    assert campaign.projects[north]["proposed_plan"]["status"] == "PAUSED"
    restored = load_campaign(campaign.out_dir)
    client = FakeClient(proposal())
    restored.plan(north, options=opts(), client=client)
    messages = json.dumps(client.requests[0]["messages"])
    assert "One inch on the left" in messages and "signed-thought" in messages
    assert "SOUTH" not in messages
    assert restored.projects[north]["proposed_plan"]["status"] == "PROPOSED"
    assert not restored.projects[south].get("proposed_plan")
    assert restored.to_dict()["projects"][north]["resource_totals"]["calls"] == 3


def test_planner_first_terminal_tool_is_fixed_and_approval_checks_context_file(tmp_path):
    campaign = make(tmp_path)
    north = next(iter(campaign.projects))
    response = resp(tool("propose_plan", {"contract": effect_goal(), "rationale": "One inch left."}, "first"),
                    tool("propose_plan", {"contract": effect_goal(2880), "rationale": "Two inches instead."}, "later"))
    campaign.plan(north, "Set margin", options=opts(), client=FakeClient(response))
    plan = campaign.projects[north]["proposed_plan"]
    assert plan["contract"] == effect_goal()
    assert len(campaign.projects[north]["planning_state"]["messages"][-1]["content"]) == 2
    from pathlib import Path
    Path(campaign.projects[north]["context_source"]).write_text('{"name":"Changed"}')
    with pytest.raises(ValueError, match="PROJECT.json changed"):
        campaign.approve_plan(north, plan["id"])


def test_scoped_replanning_preserves_other_ready_outputs_and_refresh_repairs_bad_source(tmp_path):
    source = doc(tmp_path / "source", "a.docx")
    second = doc(source.parent, "b.docx")
    bad = source.parent / "broken.docx"
    bad.write_bytes(b"broken package")
    campaign = create_campaign([str(source.parent)], tmp_path / "workspace")
    pid = next(iter(campaign.projects))
    ids = {d["source"]: i for i, d in campaign.documents.items()}
    campaign.plan(pid, "Set left margin", doc_ids=[ids[str(source)], ids[str(second)]], options=opts(), client=FakeClient(proposal()))
    campaign.approve_plan(pid, campaign.projects[pid]["proposed_plan"]["id"])
    campaign.run(options=opts(), client=FakeClient(adaptive_submit(source), critic_accept(), adaptive_submit(second), critic_accept()))
    original_second = campaign.documents[ids[str(second)]]["output"]
    assert campaign.documents[ids[str(source)]]["status"] == "READY", campaign.to_text()
    campaign.save_plan_draft(pid, "Two-inch margin on a only", doc_ids=[ids[str(source)]])
    assert campaign.documents[ids[str(source)]]["status"] == "READY"
    campaign.plan(pid, doc_ids=[ids[str(source)]], options=opts(), client=FakeClient(proposal(effect_goal(2880))))
    assert campaign.documents[ids[str(second)]]["output"] == original_second
    campaign.approve_plan(pid, campaign.projects[pid]["proposed_plan"]["id"])
    assert campaign.documents[ids[str(second)]]["status"] == "READY"
    campaign.run(options=opts(), client=FakeClient(adaptive_submit(source, effect_goal(2880)), critic_accept()))
    assert campaign.documents[ids[str(second)]]["output"] == original_second
    bad.write_bytes(build("23 05 00", "HVAC GENERAL"))
    campaign.refresh([ids[str(bad)]])
    assert campaign.documents[ids[str(bad)]]["status"] == "PENDING"
    assert campaign.documents[ids[str(second)]]["output"] == original_second
    campaign.plan(pid, "One-inch margin", doc_ids=[ids[str(bad)]], options=opts(), client=FakeClient(proposal()))
    campaign.approve_plan(pid, campaign.projects[pid]["proposed_plan"]["id"])
    campaign.run(options=opts(), client=FakeClient(adaptive_submit(bad), critic_accept()))
    assert all(d["status"] == "READY" for d in campaign.documents.values()), campaign.to_text()
    assert len(list(campaign.export(tmp_path / "delivery").rglob("*.docx"))) == 3


def test_project_allowance_is_shared_by_document_sessions_and_continuation_keeps_ready(tmp_path, monkeypatch):
    # Force the custom-method route, preserving real sessions, calls, critic, independent
    # verification and continuation. Simple built-in goals normally need no model calls.
    from spec_surgeon.campaign import Campaign
    original = Campaign._adaptive_document
    def custom_method(self, project, document, options, **kwargs):
        if kwargs.get("prepare_only"):
            return []
        return original(self, project, document, options, **kwargs)
    monkeypatch.setattr(Campaign, "_adaptive_document", custom_method)
    source = doc(tmp_path / "source", "a.docx")
    second, third = doc(source.parent, "b.docx"), doc(source.parent, "c.docx")
    campaign = create_campaign([str(source.parent)], tmp_path / "workspace")
    pid = next(iter(campaign.projects))
    campaign.set_goal(pid, "Set left margin", adaptive_contract=effect_goal())
    policy = RecoveryOptions(render=False, use_memory=False, max_calls=2)
    client = FakeClient(adaptive_submit(source), critic_accept())
    campaign.run(options=policy, client=client)
    docs = list(campaign.documents.values())
    assert docs[0]["status"] == "READY", campaign.to_text()
    assert all(d["status"] != "READY" for d in docs[1:])
    assert len(client.requests) == 2
    assert campaign.projects[pid]["resource_history"][0]["calls"] == 2
    output = docs[0]["output"]
    loaded = load_campaign(campaign.out_dir)
    loaded.run(client=FakeClient(adaptive_submit(second), critic_accept()))
    assert loaded.documents[docs[0]["id"]]["output"] == output
    assert loaded.documents[docs[1]["id"]]["status"] == "READY", loaded.to_text()
    assert loaded.documents[docs[2]["id"]]["status"] != "READY"
    loaded.run(client=FakeClient(adaptive_submit(third), critic_accept()))
    assert all(d["status"] == "READY" for d in loaded.documents.values()), loaded.to_text()
    assert loaded.to_dict()["projects"][pid]["resource_totals"]["calls"] == 6


def test_generalized_ambiguous_choice_is_scoped_to_effect_and_resumes_same_session(tmp_path):
    from docx import Document
    from spec_surgeon.goals import choose_goal, compile_goal_candidate
    source = tmp_path / "ambiguous.docx"
    document = Document()
    document.add_paragraph("TITLE")
    document.add_paragraph("TITLE")
    document.save(source)
    contract = {"kind": "document_effects", "effects": [
        {"id": "title", "type": "format", "target": {"locators": [{"exact_text": "TITLE"}], "cardinality": 1}, "run": {"b": True}},
        effect_goal()["effects"][0]]}
    campaign = create_campaign([str(source)], tmp_path / "workspace")
    pid, did = next(iter(campaign.projects)), next(iter(campaign.documents))
    campaign.set_goal(pid, "Bold the intended title and set left margin", adaptive_contract=contract)
    campaign.run(options=opts(), client=FakeClient())
    assert campaign.documents[did]["status"] == "NEEDS_INPUT", campaign.to_text()
    question = campaign.projects[pid]["questions"][0]
    assert question["document_ids"] == [did] and question["kind"] == "goal_binding"
    selected = question["resolution"]["effects"]["title"]["candidates"][1]["candidate_id"]
    run = campaign.projects[pid]["runs"][0]["path"]
    loaded = load_campaign(campaign.out_dir)
    loaded.bind(did, effect_id="title", candidate_ids=[selected])
    binding = choose_goal(source, contract, [selected], effect_id="title")
    candidate = compile_goal_candidate(source, contract, binding)
    loaded.run(client=FakeClient(resp(tool("submit_repair", {"candidate": candidate})), critic_accept()))
    assert loaded.documents[did]["status"] == "READY", loaded.to_text()
    assert loaded.projects[pid]["runs"][0]["path"] == run
    assert not loaded.projects[pid]["questions"]
    assert len(list(loaded.export(tmp_path / "delivery").rglob("*.docx"))) == 1


def test_scoped_new_plan_migrates_shared_legacy_proof_without_silent_export_loss(tmp_path):
    source = doc(tmp_path / "source", "a.docx")
    second = doc(source.parent, "b.docx")
    campaign = create_campaign([str(source.parent)], tmp_path / "workspace")
    pid = next(iter(campaign.projects))
    ids = {d["source"]: i for i, d in campaign.documents.items()}
    campaign.set_goal(pid, goal()["request"], spec=goal())
    campaign.run(options=opts(), client=FakeClient(submit()))
    assert all(d["status"] == "READY" for d in campaign.documents.values()), campaign.to_text()
    campaign.plan(pid, "Two-inch left margin for a", doc_ids=[ids[str(source)]], options=opts(),
                  client=FakeClient(proposal(effect_goal(2880))))
    campaign.approve_plan(pid, campaign.projects[pid]["proposed_plan"]["id"])
    assert campaign.documents[ids[str(second)]]["requires_revalidation"]
    with pytest.raises(ValueError, match="per-document revalidation"):
        campaign.export(tmp_path / "too-soon")
    campaign.run(client=FakeClient())
    assert all(d["status"] == "READY" for d in campaign.documents.values()), campaign.to_text()
    assert not campaign.documents[ids[str(second)]].get("requires_revalidation")
    assert len(list(campaign.export(tmp_path / "delivery").rglob("*.docx"))) == 2


def test_global_preparation_finishes_later_clean_project_before_first_model_call(tmp_path):
    campaign = make(tmp_path)
    first, second = list(campaign.projects)
    difficult = {"kind": "document_effects", "effects": [
        {"id": "title", "type": "format", "target": {"locators": [{"exact_text": "SPRINKLER SYSTEMS"}], "cardinality": 1},
         "font": {"name": "Arial", "size": 12}, "run": {"b": True}, "tracked": True}]}
    campaign.set_goal(first, "Tracked title formatting needs a combined method", adaptive_contract=difficult)
    campaign.set_goal(second, "Set left margin", adaptive_contract=effect_goal())
    class ObserveClient(FakeClient):
        def _stream(self, **kwargs):
            assert all(campaign.documents[d]["status"] == "READY" for d in campaign.projects[second]["document_ids"])
            assert "NORTH" in json.dumps(kwargs["messages"])
            assert "SOUTH" not in json.dumps(kwargs["messages"])
            return super()._stream(**kwargs)
    client = ObserveClient(resp(tool("need_user", {"question": "Confirm the tracked title formatting scope?"})))
    campaign.run(options=opts(), client=client)
    assert len(client.requests) == 1
    assert campaign.projects[second]["status"] == "READY", campaign.to_text()
    assert campaign.projects[first]["status"] == "NEEDS_INPUT"


def test_unknown_spend_remains_unknown_in_combined_planning_and_execution_totals(tmp_path):
    campaign = make(tmp_path)
    pid = next(iter(campaign.projects))
    campaign.projects[pid]["planning_resource_history"] = [{"calls": 1, "accounted_tokens": 30, "estimated_cost_usd": 0.25}]
    campaign.projects[pid]["resource_history"] = [{"calls": 2, "accounted_tokens": 50, "estimated_cost_usd": None}]
    campaign._save("record_usage")
    result = load_campaign(campaign.out_dir).to_dict()["projects"][pid]["resource_totals"]
    assert result == {"calls": 3, "accounted_tokens": 80, "estimated_cost_usd": None}


def test_approval_refuses_reassigned_document_and_planner_rejects_cross_project_scope(tmp_path):
    campaign = make(tmp_path)
    first, second = campaign.projects
    first_doc = campaign.projects[first]["document_ids"][0]
    second_doc = campaign.projects[second]["document_ids"][0]
    with pytest.raises(ValueError, match="within this project"):
        campaign.plan(first, "Margins", doc_ids=[second_doc], options=opts(), client=FakeClient())
    campaign.plan(first, "Margins", options=opts(), client=FakeClient(proposal()))
    plan_id = campaign.projects[first]["proposed_plan"]["id"]
    campaign.set_project([first_doc], second)
    with pytest.raises(ValueError, match="within this project"):
        campaign.approve_plan(first, plan_id)
    assert not campaign.projects[second]["goal"]


def _hard_campaign(tmp_path, names=("a.docx", "b.docx")):
    folder = tmp_path / "source"
    for name in names:
        doc(folder, name)
    campaign = create_campaign([str(folder)], tmp_path / "workspace")
    pid = next(iter(campaign.projects))
    contract = {"kind": "document_effects", "effects": [
        {"id": "title", "type": "format", "target": {"locators": [{"exact_text": "SPRINKLER SYSTEMS"}], "cardinality": 1},
         "font": {"name": "Arial", "size": 12}, "run": {"b": True}, "tracked": True}]}
    campaign.set_goal(pid, "Track the combined title font and bold change", adaptive_contract=contract)
    return campaign, pid


def test_reassigned_documents_need_new_scoped_approval_even_when_moved_back(tmp_path):
    campaign = make(tmp_path)
    north, south = campaign.projects
    original = campaign.projects[north]["document_ids"][0]
    moved = campaign.projects[south]["document_ids"][0]
    campaign.plan(north, "One-inch left margin", options=opts(), client=FakeClient(proposal()))
    campaign.approve_plan(north, campaign.projects[north]["proposed_plan"]["id"])
    for _ in range(2):
        campaign.set_project([moved], north)
        campaign = load_campaign(campaign.out_dir)
        campaign.run(options=opts(), client=FakeClient())
        assert campaign.documents[original]["status"] == "READY", campaign.to_text()
        assert campaign.documents[moved]["status"] != "READY", campaign.to_text()
        assert not campaign.documents[moved].get("output")
        assert moved in campaign.projects[north]["review_required_docs"]
        campaign.plan(north, "One-inch left margin", doc_ids=[moved], options=opts(), client=FakeClient(proposal()))
        campaign.approve_plan(north, campaign.projects[north]["proposed_plan"]["id"])
        campaign.run(options=opts(), client=FakeClient())
        assert campaign.documents[moved]["status"] == "READY", campaign.to_text()
        campaign.set_project([moved], south)


def test_missing_document_goal_cannot_inherit_a_scoped_plan_after_reload(tmp_path):
    source = doc(tmp_path / "source", "a.docx")
    doc(source.parent, "b.docx")
    campaign = create_campaign([str(source.parent)], tmp_path / "workspace")
    pid = next(iter(campaign.projects))
    approved, unapproved = campaign.projects[pid]["document_ids"]
    campaign.plan(pid, "One-inch left margin", doc_ids=[approved], options=opts(), client=FakeClient(proposal()))
    campaign.approve_plan(pid, campaign.projects[pid]["proposed_plan"]["id"])
    # Older workspaces could contain newly assigned documents without map entries.
    del campaign.projects[pid]["document_goals"][unapproved]
    campaign._save("legacy_missing_scoped_goal")
    loaded = load_campaign(campaign.out_dir)
    loaded.run(options=opts(), client=FakeClient())
    assert loaded.documents[approved]["status"] == "READY", loaded.to_text()
    assert loaded.documents[unapproved]["status"] != "READY", loaded.to_text()
    assert not loaded.documents[unapproved].get("output")


def test_ready_output_without_a_scoped_goal_cannot_be_exported_after_reload(tmp_path):
    source = doc(tmp_path / "source")
    campaign = create_campaign([str(source)], tmp_path / "workspace")
    pid, did = next(iter(campaign.projects)), next(iter(campaign.documents))
    campaign.plan(pid, "One-inch left margin", options=opts(), client=FakeClient(proposal()))
    campaign.approve_plan(pid, campaign.projects[pid]["proposed_plan"]["id"])
    campaign.run(options=opts(), client=FakeClient())
    assert campaign.documents[did]["status"] == "READY", campaign.to_text()
    # A pre-fix checkpoint may contain a verified candidate with no scoped approval.
    del campaign.projects[pid]["document_goals"][did]
    campaign._save("legacy_ready_output_without_approval")
    loaded = load_campaign(campaign.out_dir)
    with pytest.raises(ValueError, match="no approved scoped goal"):
        loaded.export(tmp_path / "delivery")
    assert not (tmp_path / "delivery").exists()


@pytest.mark.parametrize("changed", ["request", "scope", "template"])
def test_approval_rejects_changed_saved_draft_after_reload(tmp_path, changed):
    source = doc(tmp_path / "source", "a.docx")
    doc(source.parent, "b.docx")
    campaign = create_campaign([str(source.parent)], tmp_path / "workspace")
    pid = next(iter(campaign.projects))
    request = "One-inch left margin"
    campaign.plan(pid, request, options=opts(), client=FakeClient(proposal()))
    plan_id = campaign.projects[pid]["proposed_plan"]["id"]
    ids = campaign.projects[pid]["document_ids"]
    template = str(doc(tmp_path / "templates")) if changed == "template" else None
    campaign.save_plan_draft(pid, "Two-inch left margin" if changed == "request" else request,
                             doc_ids=ids[:1] if changed == "scope" else ids, template=template)
    before = (campaign.out_dir / "campaign.json").read_bytes()
    loaded = load_campaign(campaign.out_dir)
    with pytest.raises(ValueError, match="draft changed"):
        loaded.approve_plan(pid, plan_id)
    assert (campaign.out_dir / "campaign.json").read_bytes() == before
    assert not loaded.projects[pid].get("approvals")


def test_approval_accepts_unchanged_draft_with_reordered_scope(tmp_path):
    source = doc(tmp_path / "source", "a.docx")
    doc(source.parent, "b.docx")
    campaign = create_campaign([str(source.parent)], tmp_path / "workspace")
    pid = next(iter(campaign.projects))
    campaign.plan(pid, "One-inch left margin", options=opts(), client=FakeClient(proposal()))
    plan_id = campaign.projects[pid]["proposed_plan"]["id"]
    campaign.save_plan_draft(pid, "One-inch left margin", doc_ids=campaign.projects[pid]["document_ids"][::-1])
    loaded = load_campaign(campaign.out_dir)
    loaded.approve_plan(pid, plan_id)
    loaded.run(options=opts(), client=FakeClient())
    assert all(d["status"] == "READY" for d in loaded.documents.values()), loaded.to_text()


def test_refinement_rotates_hard_documents_without_resetting_attempt_allowance(tmp_path):
    campaign, pid = _hard_campaign(tmp_path)
    wrong_second = goal()
    wrong_second["ops"][0]["params"]["margins"]["left"] = 2880
    client = FakeClient(submit(), resp(tool("need_user", {"question": "Which title in b?"})),
                        resp(tool("submit_repair", {"candidate": {"kind": "spec", "spec": wrong_second}})))
    policy = RecoveryOptions(render=False, use_memory=False, group_call_slice=1, max_calls=10, max_attempts=2)
    campaign.run(options=policy, client=client)
    assert len(client.requests) == 3, campaign.to_text()
    assert "a.docx" in json.dumps(client.requests[0]["messages"])
    assert "b.docx" in json.dumps(client.requests[1]["messages"])
    assert "a.docx" in json.dumps(client.requests[2]["messages"])
    ids = {Path(d["source"]).name: did for did, d in campaign.documents.items()}
    assert campaign.documents[ids["b.docx"]]["status"] == "NEEDS_INPUT"
    first_run = next(r for r in campaign.projects[pid]["runs"] if r["document_id"] == ids["a.docx"])
    from spec_surgeon.recovery import load_recovery
    result = load_recovery(first_run["path"])
    state = next(iter(result.session_state["groups"].values()))
    assert state["used"] == state["attempt_limit"] == 2
    assert state["status"] == "paused"
    assert result.budget["calls"] == 2  # One live allowance survived the two internal slices.
    assert campaign.projects[pid]["resource_history"][-1]["calls"] == 3


def test_short_project_allowance_persists_fair_next_document_across_restart(tmp_path):
    campaign, pid = _hard_campaign(tmp_path, ("a.docx", "b.docx", "c.docx"))
    policy = RecoveryOptions(render=False, use_memory=False, group_call_slice=1, max_calls=2)
    first = FakeClient(resp(tool("list_styles", {"file": "a.docx"})),
                       resp(tool("list_styles", {"file": "b.docx"})))
    campaign.run(options=policy, client=first)
    assert len(first.requests) == 2
    ids = {Path(d["source"]).name: did for did, d in campaign.documents.items()}
    assert campaign.projects[pid]["scheduler"]["next_document_id"] == ids["c.docx"]
    reopened = load_campaign(campaign.out_dir)
    second = FakeClient(resp(tool("list_styles", {"file": "c.docx"})),
                        resp(tool("read_paragraphs", {"file": "a.docx", "count": 1})))
    reopened.run(client=second)
    assert len(second.requests) == 2
    assert "c.docx" in json.dumps(second.requests[0]["messages"])
    assert "a.docx" in json.dumps(second.requests[1]["messages"])
    assert reopened.projects[pid]["scheduler"]["next_document_id"] == ids["b.docx"]
    assert reopened.to_dict()["projects"][pid]["resource_totals"]["calls"] == 4
    assert all(d["status"] == "UNRESOLVED" for d in reopened.documents.values()), reopened.to_text()


def test_answer_recovery_checkpoint_is_not_applied_twice_after_campaign_interruption(tmp_path, monkeypatch):
    from docx import Document
    from spec_surgeon.goals import choose_goal, compile_goal_candidate
    source = tmp_path / "ambiguous.docx"
    document = Document()
    document.add_paragraph("TITLE")
    document.add_paragraph("TITLE")
    document.save(source)
    contract = {"kind": "document_effects", "effects": [{"id": "title", "type": "format",
        "target": {"locators": [{"exact_text": "TITLE"}], "cardinality": 1}, "run": {"b": True}}]}
    campaign = create_campaign([str(source)], tmp_path / "workspace")
    pid, did = next(iter(campaign.projects)), next(iter(campaign.documents))
    campaign.set_goal(pid, "Bold the intended title", adaptive_contract=contract)
    campaign.run(options=opts(), client=FakeClient())
    question = campaign.projects[pid]["questions"][0]
    selected = question["resolution"]["effects"]["title"]["candidates"][0]["candidate_id"]
    campaign.bind(did, effect_id="title", candidate_ids=[selected])
    binding = choose_goal(source, contract, [selected], effect_id="title")
    candidate = compile_goal_candidate(source, contract, binding)
    saved = campaign._save
    interrupted = False
    def interrupt(action, details=None):
        nonlocal interrupted
        if action == "decision_applied" or interrupted:
            interrupted = True  # Simulate abrupt exit: finalizers cannot checkpoint memory either.
            raise KeyboardInterrupt("Recovery checkpoint committed before campaign's decision checkpoint")
        saved(action, details)
    monkeypatch.setattr(campaign, "_save", interrupt)
    with pytest.raises(KeyboardInterrupt):
        campaign.run(client=FakeClient(resp(tool("submit_repair", {"candidate": candidate})), critic_accept()))
    restored = load_campaign(campaign.out_dir)
    assert not restored.projects[pid]["decisions"][0]["applied"]
    # Model work already completed, and replaying the same stable answer must not add
    # another recovery decision or require another model response.
    restored.run(client=FakeClient())
    assert restored.documents[did]["status"] == "READY", restored.to_text()
    from spec_surgeon.recovery import load_recovery
    result = load_recovery(restored.projects[pid]["runs"][0]["path"])
    assert len(result.decisions) == 1
    assert result.decisions[0]["external_decision_id"] == restored.projects[pid]["decisions"][0]["id"]
