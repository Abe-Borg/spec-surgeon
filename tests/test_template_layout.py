"""V7's template reference (verify/template_layout.py, live QA finding F08): the template as each document's plan
makes it show — its fills in the slots, in capitals where caps_fills asks, each template section the output uses at
the output section's text width with tab stops rescaled as the plan declares, the template's vertical layout kept —
checked here without a renderer against the copy the op makes: the same words per header/footer paragraph, the same
effective tab stops, every other member the template's own; and, where the plan keeps the document's header/footer
distances, the source's distances (the deep-dive stress test's D4). The rendered comparisons are in tests/test_template.py
(test_v7_template_reference_*)."""
from __future__ import annotations

import copy
import hashlib
import io
import shutil
import zipfile
from pathlib import Path

import pytest
from lxml import etree

from make_fixtures import REFERENCE_FILLS, TEMPLATE_FILLS, build, build_reference_spec, build_template
from spec_surgeon.ops import get_op, merge_fragments
from spec_surgeon.package import DocxPackage, W
from spec_surgeon.verify import run_suite
from spec_surgeon.verify.geometry import template_section_map
from spec_surgeon.verify.props import sect_prs
from spec_surgeon.verify.reader import VPackage
from spec_surgeon.verify.template import Eff, _refs_shown, _width, display_of
from spec_surgeon.verify.template_layout import template_reference

OP = "apply_header_footer_template"
KINDS = ["header", "footer"]


@pytest.fixture(scope="module")
def template(tmp_path_factory) -> Path:
    p = tmp_path_factory.mktemp("layout") / "00_00_00_TEMPLATE.docx"
    p.write_bytes(build_template())
    return p


def params(tpl: Path, **kw) -> dict:
    d = {"template": str(tpl), "template_sections": {"cover": 0, "body": 1}, "slot_fills": copy.deepcopy(TEMPLATE_FILLS)}
    d.update(kw)
    return d


def made(src: bytes, prm: dict):
    """The plan, the copy (verified: V1–V10 pass) and its reference with the section map. The plan carries the
    approved op and its params, as a spec's plan does (``editspec.plan_spec``)."""
    pkg = DocxPackage(src)
    frag = get_op(OP).plan(pkg, {"sections": "all"}, prm)
    plan = merge_fragments([frag]).to_dict()
    plan["ops"] = [{"op": OP, "select": {"sections": "all"}, "params": prm}]
    work = pkg.copy()
    get_op(OP).apply(work, frag)
    out = work.output_bytes()
    rec = run_suite(src, out, plan, "src.docx", "out.docx")
    assert rec.passed, rec.to_markdown()
    n = len(sect_prs(VPackage(out).xml("word/document.xml")))
    ref, smap = template_reference(plan, out, template_section_map(prm["template_sections"], n),
                                   prm.get("kinds", KINDS), source=src)
    return plan, out, ref, smap


def shown_paragraphs(pkg: VPackage, k: int):
    """{(kind, type): [top-level paragraph elements]} of the header/footer parts section k shows."""
    return {key: [p for p in pkg.xml(target).iter(W + "p")
                  if not any(q.tag == W + "p" for q in p.iterancestors())]
            for key, target in _refs_shown(pkg)[k].items() if target}


def test_the_reference_shows_the_documents_fills_at_its_text_width(template):
    """The template's sections at the copy's text width (8640; the template's is 9360), its header/footer distances,
    page height and titlePg kept; each header/footer the copy's text, paragraph by paragraph, with the copy's
    effective tab stops; the fills' highlights gone; every template member byte for byte but the main part, its
    relationships and the content types; no dangling relationship or untyped member the template lacks."""
    src = build("21 05 00", "COMMON WORK RESULTS FOR FIRE SUPPRESSION")
    plan, out, ref, smap = made(src, params(template))
    tpl, opkg, rpkg = VPackage(template.read_bytes()), VPackage(out), VPackage(ref)
    assert smap == [0, 1]
    t_secs, o_secs, r_secs = (sect_prs(p.xml("word/document.xml")) for p in (tpl, opkg, rpkg))
    assert len(r_secs) == 2
    for g, (t_sp, o_sp, r_sp) in enumerate(zip(t_secs, o_secs, r_secs)):
        assert _width(t_sp) == 9360 and _width(r_sp) == _width(o_sp) == 8640
        t_mar, r_mar = t_sp.find(W + "pgMar"), r_sp.find(W + "pgMar")
        assert [r_mar.get(W + k) for k in ("top", "bottom", "header", "footer")] == \
               [t_mar.get(W + k) for k in ("top", "bottom", "header", "footer")]
        assert r_sp.find(W + "pgSz").get(W + "h") == t_sp.find(W + "pgSz").get(W + "h")
        assert (r_sp.find(W + "titlePg") is None) == (t_sp.find(W + "titlePg") is None)
        ref_pars, out_pars = shown_paragraphs(rpkg, g), shown_paragraphs(opkg, g)
        assert set(ref_pars) == set(out_pars) and ref_pars
        r_eff, o_eff = Eff(rpkg), Eff(opkg)
        for key in ref_pars:
            assert [display_of(p) for p in ref_pars[key]] == [display_of(p) for p in out_pars[key]], key
            for rp, op_ in zip(ref_pars[key], out_pars[key]):
                if not any(q.tag == W + "tbl" for q in rp.iterancestors()):
                    assert r_eff.para(rp)["tabs"] == o_eff.para(op_)["tabs"], key
    texts = [display_of(p) for key, pars in shown_paragraphs(rpkg, 1).items() if key[0] == "header" for p in pars]
    assert "21 05 00 – COMMON WORK RESULTS FOR FIRE SUPPRESSIONISSUED FOR CONSTRUCTION  2026-02-15" in texts  # w:t
    added = [n for n in rpkg.names if n not in tpl.names]
    assert added and all("v7-reference-" in n for n in added)
    assert not [h for n in added if not n.endswith(".rels") for h in rpkg.xml(n).iter(W + "highlight")]
    assert [h for n in tpl.names if "header" in n and not n.endswith(".rels") for h in tpl.xml(n).iter(W + "highlight")]
    assert {n for n in tpl.names if rpkg.read(n) != tpl.read(n)} == \
           {"[Content_Types].xml", "word/_rels/document.xml.rels", "word/document.xml"}
    assert rpkg.integrity_problems() == tpl.integrity_problems()       # none of its own (the reader's .rels quirk)


def margins(pkg: VPackage):
    return [[sp.find(W + "pgMar").get(W + k) for k in ("top", "bottom", "header", "footer")]
            for sp in sect_prs(pkg.xml("word/document.xml"))]


def test_a_plan_that_keeps_the_distances_holds_the_copy_to_the_sources(template):
    """align_header_footer_distance false: the copy keeps the document's header and footer distances, and so does its
    reference — read from the source, section by section; the template's top and bottom margins, page height and
    titlePg stay. Aligned, the reference keeps the template's distances (above)."""
    tpl = VPackage(template.read_bytes())
    src = build("21 22 00", "CLEAN-AGENT", header_distance_twips=1440, footer_distance_twips=1080)
    plan, out, ref, smap = made(src, params(template, align_header_footer_distance=False))
    assert smap == [0, 1]
    t_m, s_m, o_m, r_m = (margins(p) for p in (tpl, VPackage(src), VPackage(out), VPackage(ref)))
    assert [m[2:] for m in s_m] == [m[2:] for m in o_m] == [m[2:] for m in r_m] == [["1440", "1080"]] * 2
    assert [m[:2] for m in r_m] == [m[:2] for m in t_m] and [m[2:] for m in t_m] != [m[2:] for m in r_m]
    _, _, aligned, _ = made(src, params(template))
    assert [m[2:] for m in margins(VPackage(aligned))] == [m[2:] for m in t_m]


def test_sections_with_their_own_distances_get_reference_sections_of_their_own(template):
    """Two body sections mapped to one template section at one width, their distances differing in the source: one
    reference section each, so neither is held to the other's. A section that states no distance has none in its
    reference section either (Word's default for both)."""
    pkg = DocxPackage(build("21 22 00", "CLEAN-AGENT"))           # a cover section and a body section
    root = pkg.xml(pkg.main_part())
    body = root.find(W + "body")
    final = body.find(W + "sectPr")
    middle = copy.deepcopy(final)                                  # a body section before the last, header at 1000
    for el in middle.findall(W + "headerReference") + middle.findall(W + "footerReference"):
        middle.remove(el)
    middle.find(W + "pgMar").set(W + "header", "1000")
    p = etree.Element(W + "p")
    etree.SubElement(p, W + "pPr").append(middle)
    body.insert(list(body).index(final), p)
    del final.find(W + "pgMar").attrib[W + "footer"]               # the last states no footer distance
    pkg.write_xml(pkg.main_part(), root)
    src = pkg.output_bytes()
    _, out, ref, smap = made(src, params(template, align_header_footer_distance=False))
    s_secs = sect_prs(VPackage(src).xml("word/document.xml"))
    assert len(s_secs) == 3 and smap == [0, 1, 2]
    r_secs = sect_prs(VPackage(ref).xml("word/document.xml"))
    assert [(sp.find(W + "pgMar").get(W + "header"), sp.find(W + "pgMar").get(W + "footer")) for sp in r_secs] == \
           [(sp.find(W + "pgMar").get(W + "header"), sp.find(W + "pgMar").get(W + "footer")) for sp in s_secs]
    assert r_secs[2].find(W + "pgMar").get(W + "footer") is None


def test_a_plan_that_keeps_the_distances_needs_the_source_to_read_them_from(template):
    src = build("21 22 00", "CLEAN-AGENT", header_distance_twips=1440)
    plan, out, _, _ = made(src, params(template, align_header_footer_distance=False))
    with pytest.raises(ValueError, match="no source was given"):
        template_reference(plan, out, [0, 1], KINDS)
    with pytest.raises(ValueError, match=r"the source has 1 section\(s\), the output 2"):
        template_reference(plan, out, [0, 1], KINDS, source=build("21 22 00", "CLEAN-AGENT", single_section=True))


def test_a_caps_fill_shows_in_capitals_by_formatting(template):
    """caps_fills: the fill's run shows in capitals (w:caps on), its letters as the document has them."""
    src = build("21 05 00", "Common Work Results for Fire Suppression")
    _, _, ref, _ = made(src, params(template, caps_fills=["SECTION_TITLE"]))
    rpkg = VPackage(ref)
    eff = Eff(rpkg)
    runs = [(r, p) for pars in shown_paragraphs(rpkg, 1).values() for p in pars for r in p.iter(W + "r")
            if "".join(t.text or "" for t in r.findall(W + "t")) == "Common Work Results for Fire Suppression"]
    assert runs and all(eff.run(r.find(W + "rPr"), p)["caps"] for r, p in runs)
    plain = [r for pars in shown_paragraphs(rpkg, 1).values() for p in pars for r in p.iter(W + "r")
             if "21 05 00" in "".join(t.text or "" for t in r.findall(W + "t"))]
    assert plain and not any(r.find(f"{W}rPr/{W}caps") is not None for r in plain)


def _landscape_last_section(src: bytes) -> bytes:
    zin, buf = zipfile.ZipFile(io.BytesIO(src)), io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for info in zin.infolist():
            data = zin.read(info.filename)
            if info.filename == "word/document.xml":
                root = etree.fromstring(data)
                sz = root.find(f"{W}body/{W}sectPr/{W}pgSz")
                sz.set(W + "w", "15840"); sz.set(W + "h", "12240"); sz.set(W + "orient", "landscape")
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            z.writestr(info, data)
    return buf.getvalue()


def test_sections_of_another_width_get_a_reference_section_of_their_own(template):
    """A landscape appendix the body template section serves: the op makes it copies of its own, and so does the
    reference — three sections (cover, body, body at 12240), the appendix's header with its right stop at its own
    width; the template's sections stay in the template's order."""
    src = _landscape_last_section(build("21 22 00", "CLEAN-AGENT", appendix_restart=True))
    _, out, ref, smap = made(src, params(template))
    rpkg, opkg = VPackage(ref), VPackage(out)
    assert smap == [0, 1, 2]
    assert [_width(sp) for sp in sect_prs(rpkg.xml("word/document.xml"))] == [8640, 8640, 12240]
    r_eff, o_eff = Eff(rpkg), Eff(opkg)
    for k in (1, 2):
        r_head = shown_paragraphs(rpkg, k)[("header", "default")][0]
        o_head = shown_paragraphs(opkg, k)[("header", "default")][0]
        assert r_eff.para(r_head)["tabs"] == o_eff.para(o_head)["tabs"]
    tabs = [r_eff.para(shown_paragraphs(rpkg, k)[("header", "default")][0])["tabs"] for k in (1, 2)]
    assert tabs[0] != tabs[1]


def test_reference_slots_show_the_documents_own_values(tmp_path):
    """A template spec (slot_markers ["reference"]): where its header and footer show its own section number and
    title, the reference shows the document's."""
    tpl = tmp_path / "08 71 00 - DOOR HARDWARE (ACME).docx"
    tpl.write_bytes(build_reference_spec())
    src = build("21 05 00", "COMMON WORK RESULTS FOR FIRE SUPPRESSION")
    prm = {"template": str(tpl), "template_sections": {"all": 0}, "slot_fills": copy.deepcopy(REFERENCE_FILLS),
           "slot_markers": ["reference"]}
    _, out, ref, smap = made(src, prm)
    rpkg, opkg = VPackage(ref), VPackage(out)
    assert smap == [0, 0]
    for key, pars in shown_paragraphs(rpkg, 0).items():
        assert [display_of(p) for p in pars] == [display_of(p) for p in shown_paragraphs(opkg, 0)[key]], key
    words = " ".join(display_of(p) for pars in shown_paragraphs(rpkg, 0).values() for p in pars)
    assert "21 05 00" in words and "08 71 00" not in words and "DOOR HARDWARE" not in words


def test_a_reference_that_cannot_be_made_raises(template, tmp_path):
    """No template in the plan, a template changed since planning, a map that does not cover the output's
    sections: ValueError (V7 then fails for want of its reference)."""
    src = build("21 05 00", "COMMON WORK RESULTS FOR FIRE SUPPRESSION")
    plan, out, _, _ = made(src, params(template))
    with pytest.raises(ValueError, match="names no template"):
        template_reference({"property_assertions": []}, out, [0, 1], KINDS)
    with pytest.raises(ValueError, match=r"2 section\(s\), the map 1"):
        template_reference(plan, out, [0], KINDS)
    changed = tmp_path / "changed.docx"
    changed.write_bytes(template.read_bytes() + b"\0")
    moved = copy.deepcopy(plan)
    for a in moved["property_assertions"]:
        if a.get("kind") == "template_parts":
            a["template"] = str(changed)
    with pytest.raises(ValueError, match="changed since the plan was made"):
        template_reference(moved, out, [0, 1], KINDS)
    assert hashlib.md5(template.read_bytes()).hexdigest() == plan["property_assertions"][1]["template_md5"]


def footer_only_template(path: Path) -> Path:
    """A template with a footer and no header at all."""
    from docx import Document
    from make_fixtures import _page_field
    doc = Document()
    f = doc.sections[0].footer
    f.is_linked_to_previous = False
    p = f.paragraphs[0]
    p.add_run("{{SECTION_NUMBER}} - ")
    _page_field(p)
    doc.add_paragraph("PART 1 - GENERAL")
    buf = io.BytesIO()
    doc.save(buf)
    path.write_bytes(buf.getvalue())
    return path


def test_a_template_that_shows_none_of_the_kinds_is_a_blank_reference(tmp_path):
    """Codex review on #144: headers from a template that has none — the op copies no part (the plan has no
    template_parts) and removes the document's headers. The reference is the template the plan names, its sections
    referencing no header: blank, as the copy's are."""
    tpl = footer_only_template(tmp_path / "footer only.docx")
    plan, out, ref, smap = made(build("21 13 13", "SPRINKLERS"),
                                {"template": str(tpl), "template_sections": {"all": 0}, "kinds": ["header"],
                                 "slot_fills": {}})
    assert not [a for a in plan["property_assertions"] if a["kind"] == "template_parts"]
    rpkg = VPackage(ref)
    assert smap == [0, 0]
    for sp in sect_prs(rpkg.xml("word/document.xml")):
        assert sp.find(W + "headerReference") is None and sp.find(W + "footerReference") is None
    assert not [k for k in _refs_shown(VPackage(out))[1] if k[0] == "header"]


@pytest.mark.render
@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")
def test_a_blank_reference_passes_v7_rendered(tmp_path):
    """The same request run with LibreOffice: the copy's headers are gone, as the template's are, and V7 passes on
    every page (blank against blank) instead of failing for want of a reference."""
    from spec_surgeon.batch import run_batch
    tpl = footer_only_template(tmp_path / "footer only.docx")
    src = tmp_path / "in"
    src.mkdir()
    (src / "a.docx").write_bytes(build("21 13 13", "SPRINKLERS"))
    spec = {"request": "the template's headers", "summary": "s", "assumptions": [], "questions": [],
            "applies_to": {"or": "all"}, "mode": "formatting",
            "ops": [{"op": OP, "select": {"sections": "all"},
                     "params": {"template": str(tpl), "template_sections": {"all": 0}, "kinds": ["header"],
                                "slot_fills": {}}}],
            "review": {"sample_files": 1, "geometry_reference": "template"}}
    batch = run_batch([str(src)], "edit-spec", {}, tmp_path / "run", workers=1, renderer="libreoffice", spec=spec)
    rec = batch.files[0].record
    assert rec.passed and rec.get("V7").status == "PASS", rec.to_markdown()
    assert all(p["anchors"] == 0 and not p["problems"] for p in rec.get("V7").data["pages"])
