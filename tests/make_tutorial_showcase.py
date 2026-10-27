"""Build the guided tour's bundled showcase: the practice files under ``src/spec_surgeon/examples/tutorial/``.

The tour (``spec_surgeon.tutorial``) runs on these files alone — no model call, no file of the user's — so they are
built once here, with python-docx (a dev dependency the app does not ship), and committed. Every ``.docx`` is written
with fixed zip timestamps, so running this again on the same python-docx gives the same bytes; ``tests/test_tutorial.py``
compares the committed files' document, header and footer parts with a fresh build, so a stale showcase fails CI.

    python tests/make_tutorial_showcase.py            # rewrites src/spec_surgeon/examples/tutorial/

What the showcase holds, and why (docs/CONVERSATIONAL_DESKTOP.md, "The guided tour"):

* two specifications built like the engine's own fixtures (a cover section with its footer, a body section that
  restarts its page numbers, an issue stamp in the header, a pending insertion and a pending deletion), so the pages
  and the notes look like a real set;
* one well-styled specification (Heading styles, automatic numbering), so the set is a mixed batch;
* one specification from "another firm" with no SECTION line, which the template button cannot fill: the tour's
  honest example of a document left untouched with its reason;
* a file that is not a Word package at all, named ``.docx``: triage refuses it, and the list shows why;
* a prepared header/footer template with one placeholder the form asks for (``ISSUED FOR PERMIT``);
* a small logo (a picture for the edit to use) and a table of find-and-replace pairs (a CSV).
"""
from __future__ import annotations

import io
import sys
import zipfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

from docx import Document  # noqa: E402
from docx.shared import Pt  # noqa: E402

from make_fixtures import build, build_dirty, build_template, png  # noqa: E402

TUTORIAL_DIR = HERE.parent / "src" / "spec_surgeon" / "examples" / "tutorial"

SPECS = (("21 13 13", "WET-PIPE SPRINKLER SYSTEMS"), ("21 13 16", "DRY-PIPE SPRINKLER SYSTEMS"))
STYLED = ("21 05 00", "COMMON WORK RESULTS FOR FIRE SUPPRESSION")
OTHER_FIRM = "Clean Agent Systems (other firm).docx"
REFUSED = "21 00 00 - TABLE OF CONTENTS.docx"
TEMPLATE = "Office template.docx"
LOGO = "logo.png"
PAIRS = "project values.csv"

# What the window reads from a bundled file's name
DOCUMENTS = tuple(f"{num} - {title}.docx" for num, title in SPECS) + (f"{STYLED[0]} - {STYLED[1]}.docx", OTHER_FIRM)
FIXED_DATE = (1980, 1, 1, 0, 0, 0)


def stable_zip(data: bytes) -> bytes:
    """The same members in the same order, every entry dated 1980-01-01: python-docx stamps the time it saved."""
    zin = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as zout:
        for info in zin.infolist():
            zi = zipfile.ZipInfo(info.filename, date_time=FIXED_DATE)
            zi.compress_type = zipfile.ZIP_DEFLATED
            zout.writestr(zi, zin.read(info.filename))
    return out.getvalue()


def build_other_firm() -> bytes:
    """A specification laid out the way another firm writes them: a DIVISION line and a title, never a ``SECTION nn nn
    nn`` line. The template button's fill rules read nothing from it, so it is left untouched with its reason."""
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    doc.add_paragraph("DIVISION 21 — FIRE SUPPRESSION")
    doc.add_paragraph("Clean Agent Fire Extinguishing Systems")
    doc.add_paragraph("PART 1 - GENERAL")
    doc.add_paragraph("1.1 SUMMARY")
    doc.add_paragraph("A. This Section includes clean agent fire extinguishing systems for the electrical rooms.")
    doc.add_paragraph("B. The Owner's insurer is FM Global and the systems shall be FM Approved.")
    doc.add_paragraph("1.2 REFERENCES")
    doc.add_paragraph("A. NFPA 2001, Standard on Clean Agent Fire Extinguishing Systems.")
    doc.add_paragraph("B. NFPA 72, National Fire Alarm and Signaling Code.")
    for i in range(3, 15):
        doc.add_paragraph(f"1.{i} Article body text for the clean agent systems. " * 5)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def showcase() -> dict:
    """Every bundled file, name → bytes."""
    files = {}
    for num, title in SPECS:
        files[f"{num} - {title}.docx"] = stable_zip(build(num, title))
    files[f"{STYLED[0]} - {STYLED[1]}.docx"] = stable_zip(build_dirty("styled", *STYLED))
    files[OTHER_FIRM] = stable_zip(build_other_firm())
    # Not a Word package: a text file saved under a .docx name (the list shows the refusal and what to do)
    files[REFUSED] = ("TABLE OF CONTENTS\n\nThis file is plain text saved with a .docx name: Word did not write it, so "
                      "Spec Surgeon refuses it on arrival and says so.\n").encode("utf-8")
    files[TEMPLATE] = stable_zip(build_template())
    files[LOGO] = png((28, 92, 158), 16)
    files[PAIRS] = ("find,replace\nISSUED FOR CONSTRUCTION,ISSUED FOR PERMIT\nSchedule 40,Schedule 10\n").encode("utf-8")
    return files


def main(argv=None) -> int:
    TUTORIAL_DIR.mkdir(parents=True, exist_ok=True)
    for name, data in showcase().items():
        (TUTORIAL_DIR / name).write_bytes(data)
        print(f"wrote {TUTORIAL_DIR / name} ({len(data)} bytes)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
