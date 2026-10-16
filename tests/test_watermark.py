"""add_watermark / remove_watermark (docs/BUILTIN_OPS_BATCH2_DESIGN.md §7): a text watermark in every header the
document shows, as Word writes it, sized to the page; Word's own watermarks (VML in a "Watermarks" content control, a
DrawingML one with a text box, a picture) taken out; the refusals; the two readings of the headers, the anchor, the
watermarks and the size compared; planted defects — a wrong OUTPUT and wrong IMPLEMENTATIONS shared by plan and apply —
each caught by V10; and the rendered pages."""
from __future__ import annotations

import io
import math
import shutil
from pathlib import Path
from typing import Any, Dict, List

import pytest
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.shared import Inches, Pt
from lxml import etree

from make_fixtures import docx_bytes, patch_docx, png
from spec_surgeon.editspec import apply_spec, plan_spec
from spec_surgeon.ops import JobException, get_op, merge_fragments
from spec_surgeon.ops import watermark as wm
from spec_surgeon.package import DocxPackage
from spec_surgeon.verify import run_suite
from spec_surgeon.verify import watermark as vwm
from spec_surgeon.verify.reader import VPackage

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W = "{%s}" % W_NS
V = "{urn:schemas-microsoft-com:vml}"
O = "{urn:schemas-microsoft-com:office:office}"
MC = "{http://schemas.openxmlformats.org/markup-compatibility/2006}"
NS_DECL = ('xmlns:w="%s" xmlns:v="urn:schemas-microsoft-com:vml" xmlns:o="urn:schemas-microsoft-com:office:office" '
           'xmlns:w10="urn:schemas-microsoft-com:office:word" '
           'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
           'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
           'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
           'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
           'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"' % W_NS)


# ---- fixtures -----------------------------------------------------------------------------------------------------

def document(*, title_page: bool = False, second: str = "none", header: bool = True, body: int = 3) -> bytes:
    """A spec-like file: a default header "ACME DATA CENTERS" (``header``), a first-page header (``title_page``), a
    second section (``second``: "none" | "inherit" — it shows the first section's headers — | "landscape")."""
    doc = Document()
    sec = doc.sections[0]
    if header:
        sec.header.paragraphs[0].add_run("ACME DATA CENTERS\tISSUED FOR BID")
    if title_page:
        sec.different_first_page_header_footer = True
        sec.first_page_header.paragraphs[0].add_run("COVER — SECTION 21 13 13")
    for i in range(body):
        doc.add_paragraph(f"1.0{i}\tArticle {i} of the specification.")
    if second != "none":
        s2 = doc.add_section(WD_SECTION.NEW_PAGE)
        if second == "landscape":
            s2.orientation = WD_ORIENT.LANDSCAPE
            s2.page_width, s2.page_height = Inches(11), Inches(8.5)
            s2.left_margin = s2.right_margin = Inches(2)
        doc.add_paragraph("Second section text.")
    return docx_bytes(doc)


WORD_SHAPETYPE = ('<v:shapetype id="_x0000_t136" coordsize="21600,21600" o:spt="136" adj="10800" '
                  'path="m@7,l@8,m@5,21600l@6,21600e"><v:formulas><v:f eqn="sum #0 0 10800"/></v:formulas>'
                  '<v:textpath on="t" fitshape="t"/><o:lock v:ext="edit" text="t" shapetype="t"/></v:shapetype>')


def word_vml(text: str = "DRAFT", sid: str = "PowerPlusWaterMarkObject357476642", spid: str = "_x0000_s2049",
             shapetype: bool = True) -> str:
    """A text watermark run as Word writes it (VML, silver, half transparent)."""
    return ('<w:r %s><w:rPr><w:noProof/></w:rPr><w:pict>%s<v:shape id="%s" o:spid="%s" type="#_x0000_t136" '
            'style="position:absolute;margin-left:0;margin-top:0;width:412.4pt;height:247.45pt;rotation:315;'
            'z-index:-251657216;mso-position-horizontal:center;mso-position-horizontal-relative:margin;'
            'mso-position-vertical:center;mso-position-vertical-relative:margin" o:allowincell="f" fillcolor="silver" '
            'stroked="f"><v:fill opacity=".5"/><v:textpath style="font-family:&quot;Calibri&quot;;font-size:1pt" '
            'string="%s"/><w10:wrap anchorx="margin" anchory="margin"/></v:shape></w:pict></w:r>'
            % (NS_DECL, WORD_SHAPETYPE if shapetype else "", sid, spid, text))


def in_sdt(run_xml: str) -> str:
    """Word's "Watermarks" content control around a paragraph holding the watermark run."""
    return ('<w:sdt %s><w:sdtPr><w:id w:val="-1588376018"/><w:docPartObj><w:docPartGallery w:val="Watermarks"/>'
            '<w:docPartUnique/></w:docPartObj></w:sdtPr><w:sdtContent><w:p><w:pPr><w:pStyle w:val="Header"/></w:pPr>'
            '%s</w:p></w:sdtContent></w:sdt>' % (NS_DECL, run_xml))


def drawingml(text: str = "CONFIDENTIAL") -> str:
    """A DrawingML watermark with a VML fallback, its words in a text box (the text exists twice)."""
    box = '<w:txbxContent><w:p><w:r><w:t>%s</w:t></w:r></w:p></w:txbxContent>' % text
    return ('<w:r %s><w:rPr><w:noProof/></w:rPr><mc:AlternateContent><mc:Choice Requires="wps"><w:drawing>'
            '<wp:inline><wp:extent cx="5000000" cy="1000000"/><wp:docPr id="77" name="PowerPlusWaterMarkObject9"/>'
            '<a:graphic><a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
            '<wps:wsp><wps:txbx>%s</wps:txbx><wps:bodyPr/></wps:wsp></a:graphicData></a:graphic></wp:inline>'
            '</w:drawing>'
            '</mc:Choice><mc:Fallback><w:pict><v:shape id="PowerPlusWaterMarkObject9" o:spid="_x0000_s2050" '
            'style="position:absolute"><v:textbox>%s</v:textbox></v:shape></w:pict></mc:Fallback>'
            '</mc:AlternateContent></w:r>' % (NS_DECL, box, box))


def picture_vml(rid: str) -> str:
    return ('<w:r %s><w:rPr><w:noProof/></w:rPr><w:pict><v:shapetype id="_x0000_t75" coordsize="21600,21600" '
            'o:spt="75" path="m@4@5l@4@11@9@11@9@5xe"><v:stroke joinstyle="miter"/></v:shapetype>'
            '<v:shape id="WordPictureWatermark123456" o:spid="_x0000_s2051" type="#_x0000_t75" '
            'style="position:absolute;margin-left:0;margin-top:0;width:200pt;height:100pt;z-index:-251656192" '
            'o:allowincell="f"><v:imagedata r:id="%s" o:title="logo"/></v:shape></w:pict></w:r>' % (NS_DECL, rid))


HEADER = "word/header1.xml"


def with_header(data: bytes, part: str, fn) -> bytes:
    return patch_docx(data, {part: fn})


def prepend_block(xml: str):
    def fn(root):
        root.insert(0, etree.fromstring(xml))
    return fn


def first_para_run(xml: str, at: int = 1):
    """Put a run into the header's first paragraph (after its pPr)."""
    def fn(root):
        p = root.find(W + "p")
        p.insert(at, etree.fromstring(xml))
    return fn


def word_watermarked(text: str = "DRAFT", sdt: bool = True) -> bytes:
    data = document()
    if sdt:
        return with_header(data, HEADER, prepend_block(in_sdt(word_vml(text))))
    return with_header(data, HEADER, first_para_run(word_vml(text)))


def picture_watermarked() -> bytes:
    doc = Document()
    run = doc.sections[0].header.paragraphs[0].add_run()
    run.add_picture(io.BytesIO(png((40, 40, 40), size=6)), width=Pt(20))
    doc.add_paragraph("Body.")
    data = docx_bytes(doc)
    rid = [None]

    def fn(root):
        p = root.find(W + "p")
        r = next(x for x in p.iter(W + "r") if x.find(W + "drawing") is not None)
        blip = next(r.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}blip"))
        rid[0] = blip.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
        p.replace(r, etree.fromstring(picture_vml(rid[0])))
    return with_header(data, HEADER, fn)


# ---- running --------------------------------------------------------------------------------------------------------

def spec(op: str, mode: str, **params) -> Dict[str, Any]:
    return {"request": "r", "summary": "s", "mode": mode, "ops": [{"op": op, "params": params}]}


def run(src: bytes, op: str = "add_watermark", **params):
    mode = "untracked" if op == "add_watermark" else "formatting"
    pkg = DocxPackage(src)
    plan, frags = plan_spec(pkg, spec(op, mode, **params))
    out = apply_spec(pkg, frags).output_bytes()
    return plan, out, run_suite(src, out, plan.to_dict(), "src.docx", "out.docx")


def suite(src: bytes, out: bytes, plan) -> Any:
    return run_suite(src, out, plan.to_dict(), "src.docx", "out.docx")


def failed(rec) -> List[str]:
    return [c.id for c in rec.checks if c.status == "FAIL"]


def v10(rec) -> List[str]:
    return rec.get("V10").data.get("errors", [])


def refusal(src: bytes, op: str = "add_watermark", **params) -> JobException:
    with pytest.raises(JobException) as e:
        get_op(op).plan(DocxPackage(src), {}, params)
    return e.value


def shapes(docx: bytes, part: str = HEADER) -> List[Any]:
    return [s for s in DocxPackage(docx).xml(part).iter(V + "shape")]


def style_of(shape) -> Dict[str, str]:
    return dict(x.split(":", 1) for x in shape.get("style").split(";"))


# ---- add: what it does -----------------------------------------------------------------------------------------------

def test_draft_goes_into_every_header_the_document_shows():
    src = document(title_page=True, second="inherit")
    plan, out, rec = run(src, text="DRAFT")
    assert not failed(rec), [(c.id, c.summary) for c in rec.checks if c.status == "FAIL"]
    pkg = DocxPackage(out)
    parts = sorted(plan.property_assertions[0]["parts"])
    assert len(parts) == 2 and plan.intended_parts == list(plan.property_assertions[0]["parts"])
    for part in parts:
        root = pkg.xml(part)
        p = root.find(W + "p")
        first = [c for c in p if c.tag != W + "pPr"][0]
        assert first.find(W + "rPr/" + W + "noProof") is not None
        shape = first.find(W + "pict/" + V + "shape")
        assert shape.get("id").startswith("PowerPlusWaterMarkObject") and shape.get(O + "spid").startswith("_x0000_s")
        assert shape.find(V + "textpath").get("string") == "DRAFT"
        assert shape.get("fillcolor") == "#C0C0C0" and shape.find(V + "fill").get("opacity") == "0.5"
        assert style_of(shape)["rotation"] == "315"
        assert first.find(W + "pict/" + V + "shapetype").get("id") == "_x0000_t136"
    ids = [s.get("id") for part in parts for s in pkg.xml(part).iter(V + "shape")]
    spids = [s.get(O + "spid") for part in parts for s in pkg.xml(part).iter(V + "shape")]
    assert len(set(ids)) == len(set(spids)) == 2
    assert plan.mode == "untracked" and plan.geometry_reference == "none"
    assert any("places the words 'DRAFT' on every page" in f for f in plan.flags)
    assert "DRAFT" not in "".join(t.text or "" for t in DocxPackage(out).xml(parts[0]).iter(W + "t"))


def test_the_box_takes_the_words_proportions_and_fills_the_page_diagonally():
    plan, out, rec = run(document(), text="NOT FOR CONSTRUCTION")
    assert not failed(rec)
    w, h = plan.property_assertions[0]["parts"][HEADER]["box"]
    em, _ = wm.advance("NOT FOR CONSTRUCTION", "Calibri")
    assert h / w == pytest.approx(1.55 / em, rel=1e-3)
    assert (w + h) / math.sqrt(2) == pytest.approx(0.98 * 432, abs=0.02)        # Letter, 1.25" sides: 432 × 648 pt
    s = style_of(shapes(out)[0])
    assert s["width"] == f"{wm.number(w)}pt" and s["height"] == f"{wm.number(h)}pt"


def test_horizontal_custom_look_and_a_size_in_points():
    plan, out, rec = run(document(), text="PRELIMINARY", layout="horizontal", font="Arial", color="#0000ff",
                         opacity=0.3, size=60)
    assert not failed(rec)
    shape = shapes(out)[0]
    s = style_of(shape)
    assert "rotation" not in s
    assert shape.get("fillcolor") == "#0000FF" and shape.find(V + "fill").get("opacity") == "0.3"
    assert shape.find(V + "textpath").get("style") == 'font-family:"Arial";font-size:1pt'
    w, h = plan.property_assertions[0]["parts"][HEADER]["box"]
    assert h == pytest.approx(1.55 * 60, abs=0.01)
    assert w == pytest.approx(round(wm.advance("PRELIMINARY", "Arial")[0] * 60, 2), abs=0.01)


def test_the_smallest_text_area_of_the_sections_showing_a_header_decides():
    src = document(second="landscape")             # (11 − 2 × 2) × (8.5 − 2) in: 504 × 468 pt; Letter: 432 × 648 pt
    plan, out, rec = run(src, text="DRAFT")
    assert not failed(rec)
    w, h = plan.property_assertions[0]["parts"][HEADER]["box"]
    assert (w + h) / math.sqrt(2) == pytest.approx(0.98 * 432, abs=0.02)        # the smallest side of either area
    touched = {(m, t, a) for m, t, a in DocxPackage(src).sections() and plan_spec(DocxPackage(src),
               spec("add_watermark", "untracked", text="DRAFT"))[1][0].touches}
    assert {("word/document.xml", "sectPr:0", "verify:pgSz"), ("word/document.xml", "sectPr:1", "verify:pgMar"),
            ("word/document.xml", "sectPr:1", "verify:hfref")} <= touched


def test_the_ids_are_new_and_unique_in_the_package():
    src = with_header(document(), HEADER, first_para_run(word_vml("OLD", sid="PowerPlusWaterMarkObject5",
                                                                  spid="_x0000_s3000"), at=1))
    plan, out, rec = run(src, text="DRAFT", replace=True)
    assert not failed(rec)
    shape = shapes(out)[0]
    assert shape.get(O + "spid") == "_x0000_s3001"
    assert shape.get("id") != "PowerPlusWaterMarkObject5"


def test_an_unknown_font_and_characters_outside_ascii_are_estimated_and_flagged():
    plan, out, rec = run(document(), text="ÉBAUCHE", font="Garamond")
    assert not failed(rec)
    flag = next(f for f in plan.flags if "the size is estimated" in f)
    assert "Garamond" in flag and "'É'" in flag


def test_a_page_type_without_a_header_is_flagged():
    data = document(title_page=True)

    def drop_first(root):
        sp = root.find(W + "body/" + W + "sectPr")
        for ref in list(sp):
            if ref.tag == W + "headerReference" and ref.get(W + "type") == "first":
                sp.remove(ref)
    src = patch_docx(data, {"word/document.xml": drop_first})
    plan, out, rec = run(src, text="DRAFT")
    assert not failed(rec)
    assert any("section 1's first page shows no header" in f for f in plan.flags)


def test_an_empty_header_made_by_word_for_its_watermark_is_replaced_in_place():
    """Word puts a watermark into an empty header as a content control around its own paragraph: that paragraph is
    the anchor, the old run goes and the new one takes its place."""
    data = Document()
    data.add_paragraph("Body.")
    sec = data.sections[0]
    sec.header.is_linked_to_previous = False
    src = docx_bytes(data)

    def only_sdt(root):
        for p in list(root):
            root.remove(p)
        root.append(etree.fromstring(in_sdt(word_vml("DRAFT"))))
    src = patch_docx(src, {HEADER: only_sdt})
    plan, out, rec = run(src, text="NOT FOR CONSTRUCTION", replace=True)
    assert not failed(rec), v10(rec)
    root = DocxPackage(out).xml(HEADER)
    assert [s.find(V + "textpath").get("string") for s in root.iter(V + "shape")] == ["NOT FOR CONSTRUCTION"]
    assert root.find(W + "sdt") is not None


def test_the_worked_example_replaces_words_draft_watermark():
    from spec_surgeon.editspec import load_spec
    example = load_spec(Path(__file__).resolve().parent.parent / "examples" / "specs" / "add-draft-watermark.json")
    src = word_watermarked("DRAFT")
    pkg = DocxPackage(src)
    plan, frags = plan_spec(pkg, example)
    out = apply_spec(pkg, frags).output_bytes()
    rec = suite(src, out, plan)
    assert not failed(rec), v10(rec)
    assert [s.find(V + "textpath").get("string") for s in shapes(out)] == ["NOT FOR CONSTRUCTION"]


# ---- add: refusals ---------------------------------------------------------------------------------------------------

def test_a_file_that_already_has_a_watermark_is_refused_unless_replace():
    src = word_watermarked("DRAFT")
    e = refusal(src, text="NOT FOR CONSTRUCTION")
    assert e.code == "watermark-exists" and "'DRAFT'" in e.reason
    plan, out, rec = run(src, text="NOT FOR CONSTRUCTION", replace=True)
    assert not failed(rec), v10(rec)
    texts = [s.find(V + "textpath").get("string") for s in shapes(out)]
    assert texts == ["NOT FOR CONSTRUCTION"]
    # the old run carried the shape definition: the new run brings its own
    assert len(list(DocxPackage(out).xml(HEADER).iter(V + "shapetype"))) == 1
    assert any("1 old watermark(s) replaced" in n for n in plan.notes)


def test_a_file_without_a_header_is_refused():
    e = refusal(document(header=False), text="DRAFT")
    assert e.code == "no-header" and "never created" in e.reason


@pytest.mark.parametrize("params,words", [
    ({}, "needs 'text'"),
    ({"text": "   "}, "needs 'text'"),
    ({"text": "DRA​FT"}, "invisible"),
    ({"text": "DRAFT\nCOPY"}, "control"),
    ({"text": "X" * 101}, "at most 100"),
    ({"text": "DRAFT", "color": "silver"}, "6 hex digits"),
    ({"text": "DRAFT", "opacity": 0}, "0.05 to 1"),
    ({"text": "DRAFT", "opacity": True}, "0.05 to 1"),
    ({"text": "DRAFT", "layout": "vertical"}, "'layout'"),
    ({"text": "DRAFT", "size": 0}, "'size'"),
    ({"text": "DRAFT", "size": "big"}, "'size'"),
    ({"text": "DRAFT", "font": 'Calibri";color:red'}, "'font'"),
    ({"text": "DRAFT", "replace": "yes"}, "'replace'"),
])
def test_bad_params_are_refused_for_every_file(params, words):
    e = refusal(document(), **params)
    assert e.code == "bad-params" and words in e.reason


def test_a_size_that_does_not_fit_this_files_page_is_the_files_refusal():
    e = refusal(document(), text="NOT FOR CONSTRUCTION", size=200)
    assert e.code == "watermark-size" and "432 × 648 pt" in e.reason
    assert get_op("add_watermark").plan(DocxPackage(document()), {}, {"text": "NOT FOR CONSTRUCTION", "size": 40})


def test_a_selector_is_refused():
    with pytest.raises(JobException) as e:
        get_op("add_watermark").plan(DocxPackage(document()), {"scope": ["headers"]}, {"text": "DRAFT"})
    assert e.value.code == "bad-params"


def test_the_spec_mode_is_untracked_for_add_and_formatting_for_remove():
    with pytest.raises(JobException) as e:
        plan_spec(DocxPackage(document()), spec("add_watermark", "formatting", text="DRAFT"))
    assert e.value.code == "spec-mode"
    with pytest.raises(JobException) as e:
        plan_spec(DocxPackage(word_watermarked()), spec("remove_watermark", "untracked"))
    assert e.value.code == "spec-mode"


# ---- remove ----------------------------------------------------------------------------------------------------------

def test_words_vml_watermark_comes_out_and_its_paragraph_stays():
    src = word_watermarked("DRAFT")
    plan, out, rec = run(src, "remove_watermark")
    assert not failed(rec), v10(rec)
    root = DocxPackage(out).xml(HEADER)
    assert not list(root.iter(V + "shape")) and root.find(W + "sdt/" + W + "sdtContent/" + W + "p") is not None
    assert "".join(t.text for t in root.iter(W + "t")) == "ACME DATA CENTERS\tISSUED FOR BID".replace("\t", "")
    assert plan.mode == "formatting" and not plan.diffs
    assert any("'DRAFT'" in n for n in plan.notes)


def test_a_drawingml_watermark_with_a_text_box_comes_out_whole():
    src = with_header(document(), HEADER, first_para_run(drawingml("CONFIDENTIAL")))
    plan, out, rec = run(src, "remove_watermark")
    assert not failed(rec), [(c.id, c.summary) for c in rec.checks if c.status == "FAIL"]
    deleted = [d for d in plan.diffs if d.after is None]
    assert [d.before for d in deleted] == ["CONFIDENTIAL", "CONFIDENTIAL"]          # the text box exists twice
    struct = [d for d in plan.diffs if d.structure_only]
    assert len(struct) == 1 and struct[0].struct_after == {"tab": 1}               # its w:drawing went
    assert ("word/header1.xml", "*", "text") in plan_spec(DocxPackage(src),
                                                          spec("remove_watermark", "formatting"))[1][0].touches
    assert not list(DocxPackage(out).xml(HEADER).iter(MC + "AlternateContent"))


def test_a_picture_watermark_comes_out_and_its_picture_file_stays():
    src = picture_watermarked()
    plan, out, rec = run(src, "remove_watermark")
    assert not failed(rec), v10(rec)
    assert not plan.members_removed and any("a picture" in n for n in plan.notes)


def test_text_takes_out_only_the_watermarks_showing_those_words():
    data = document(title_page=True)
    src = patch_docx(data, {HEADER: first_para_run(word_vml("DRAFT")),
                            "word/header2.xml": first_para_run(word_vml("CONFIDENTIAL", sid="PowerPlusWaterMarkObject8",
                                                                        spid="_x0000_s2060"))})
    plan, out, rec = run(src, "remove_watermark", text="  draft ")
    assert not failed(rec), v10(rec)
    left = [s.find(V + "textpath").get("string") for part in (HEADER, "word/header2.xml")
            for s in DocxPackage(out).xml(part).iter(V + "shape")]
    assert left == ["CONFIDENTIAL"]


def test_no_watermark_is_a_no_op_by_default_and_a_refusal_when_asked():
    src = document()
    plan, out, rec = run(src, "remove_watermark")
    assert not failed(rec) and out == src
    assert any("no watermark: nothing to do" in n for n in plan.notes)
    e = refusal(src, "remove_watermark", allow_missing=False)
    assert e.code == "not-found"
    e = refusal(word_watermarked("DRAFT"), "remove_watermark", text="FINAL", allow_missing=False)
    assert e.code == "not-found" and "'FINAL'" in e.reason


def test_watermarks_that_cannot_be_taken_out_are_refused():
    mixed = word_vml().replace("<w:pict>", "<w:t>x</w:t><w:pict>", 1)
    e = refusal(with_header(document(), HEADER, first_para_run(mixed)), "remove_watermark")
    assert e.code == "structure" and "shares its run" in e.reason
    tracked = ('<w:ins %s w:id="9" w:author="J" w:date="2026-01-01T00:00:00Z">%s</w:ins>' % (NS_DECL, word_vml()))
    e = refusal(with_header(document(), HEADER, first_para_run(tracked)), "remove_watermark")
    assert e.code == "existing-revisions"
    # the watermark's run carries the shape definition another shape of the header still uses
    other = ('<w:r %s><w:pict><v:shape id="_x0000_s1030" type="#_x0000_t136" style="width:10pt;height:5pt">'
             '<v:textpath string="WordArt"/></v:shape></w:pict></w:r>' % NS_DECL)
    src = with_header(document(), HEADER, lambda r: (first_para_run(word_vml())(r), first_para_run(other, at=2)(r)))
    e = refusal(src, "remove_watermark")
    assert e.code == "structure" and "_x0000_t136" in e.reason
    # … which replacing it does not break: the new watermark brings the definition back
    plan, out, rec = run(src, text="NEW", replace=True)
    assert not failed(rec), v10(rec)


# ---- the two readings agree ------------------------------------------------------------------------------------------

@pytest.mark.parametrize("build", [lambda: document(title_page=True, second="inherit"), lambda: word_watermarked(),
                                   picture_watermarked, lambda: document(second="landscape")])
def test_the_two_readings_of_headers_anchors_and_watermarks_agree(build):
    data = build()
    pkg, vp = DocxPackage(data), VPackage(data)
    order, showing, _ = wm.headers_shown(pkg)
    v_order, v_showing = vwm._v_shows(vp)
    assert order == v_order and showing == v_showing
    from spec_surgeon.ops.hf_common import KINDS  # noqa: F401 — the transform's header model is the one compared
    secs = DocxPackage.sect_prs(pkg.xml(pkg.main_part()))
    for sp in secs:
        assert wm.text_area(sp) == vwm._v_area(sp)
    for part in pkg.parts_by_kind()["headers"]:
        root = pkg.xml(part)
        assert wm.first_paragraph(root) is vwm._v_anchor(root)
        mine = wm.watermark_runs(root, part)
        theirs = vwm._v_watermarks(root)
        assert [(f["run"], f["text"]) for f in mine] == [(f["run"], f["text"]) for f in theirs]


@pytest.mark.parametrize("x", [0, 0.5, 0.3, 412.4, 247.45, 100, 100.004, 99.995, 1e-7, 379.4949, 1000, -0.001])
def test_the_two_number_writers_agree(x):
    assert wm.number(x) == vwm._v_num(x)


def test_the_two_number_writers_agree_on_many_values():
    import random
    rnd = random.Random(16)
    for _ in range(5000):
        x = round(rnd.uniform(0, 1500), rnd.choice([0, 1, 2, 3, 6]))
        assert wm.number(x) == vwm._v_num(x), x


def test_adding_then_removing_gives_the_header_back():
    src = document(title_page=True)
    _, added_, rec = run(src, text="DRAFT")
    assert not failed(rec)
    _, back, rec = run(added_, "remove_watermark")
    assert not failed(rec), v10(rec)
    from spec_surgeon.verify.edits import v_items
    from spec_surgeon.verify.edits import compare_items
    for part in (HEADER, "word/header2.xml"):
        a, b = DocxPackage(src).xml(part), DocxPackage(back).xml(part)
        assert compare_items(v_items(a.find(W + "p")), v_items(b.find(W + "p")), set()) is None


# ---- planted defects: a wrong output ---------------------------------------------------------------------------------

@pytest.fixture()
def added():
    src = document()
    plan, out, rec = run(src, text="DRAFT")
    assert not failed(rec)
    return src, plan, out


def mutate_header(out: bytes, fn) -> bytes:
    return patch_docx(out, {HEADER: fn})


def first_shape(root):
    return next(root.iter(V + "shape"))


@pytest.mark.parametrize("damage,words", [
    (lambda r: first_shape(r).find(V + "textpath").set("string", "DRAFT COPY"), "not exactly the predicted change"),
    (lambda r: first_shape(r).set("fillcolor", "#FF0000"), "not exactly the predicted change"),
    (lambda r: first_shape(r).find(V + "fill").set("opacity", "0.9"), "not exactly the predicted change"),
    (lambda r: first_shape(r).set("style", first_shape(r).get("style").replace("rotation:315;", "")),
     "not exactly the predicted change"),
])
def test_another_look_than_asked_is_caught(added, damage, words):
    src, plan, out = added
    rec = suite(src, mutate_header(out, damage), plan)
    assert failed(rec) == ["V10"] and any(words in e for e in v10(rec)), v10(rec)


def test_a_second_watermark_is_caught(added):
    src, plan, out = added

    def twice(root):
        p = root.find(W + "p")
        run = next(x for x in p if x.tag == W + "r")
        p.append(etree.fromstring(etree.tostring(run)))
    rec = suite(src, mutate_header(out, twice), plan)
    assert "V10" in failed(rec)


def test_a_moved_watermark_run_is_caught(added):
    src, plan, out = added

    def move(root):
        p = root.find(W + "p")
        run = next(x for x in p if x.tag == W + "r")
        p.remove(run)
        p.append(run)
    rec = suite(src, mutate_header(out, move), plan)
    assert failed(rec) == ["V10"] and any("does not start with a watermark run" in e for e in v10(rec))


def test_a_reused_id_is_caught(added):
    src, plan, out = added

    def reuse(root):
        first_shape(root).set(O + "spid", "_x0000_s1")
        p = root.find(W + "p")
        p.append(etree.fromstring('<w:r %s><w:pict><v:shape o:spid="_x0000_s1" style="width:1pt"/></w:pict></w:r>'
                                  % NS_DECL))
    rec = suite(src, mutate_header(out, reuse), plan)
    assert "V10" in failed(rec) and any("not new and unique" in e for e in v10(rec))


def test_a_new_number_that_repeats_another_shapes_id_is_caught(added):
    src, plan, out = added

    def clash(root):
        spid = first_shape(root).get(O + "spid")
        p = root.find(W + "p")
        p.append(etree.fromstring('<w:r %s><w:pict><v:shape id="%s" style="width:1pt"/></w:pict></w:r>'
                                  % (NS_DECL, spid)))
    rec = suite(src, mutate_header(out, clash), plan)
    assert "V10" in failed(rec) and any("not new and unique" in e for e in v10(rec))


def test_a_box_of_another_size_is_caught(added):
    src, plan, out = added
    rec = suite(src, mutate_header(out, lambda r: first_shape(r).set(
        "style", first_shape(r).get("style").replace("width:", "width:1"))), plan)
    assert failed(rec) == ["V10"]


def test_a_header_text_changed_beside_the_watermark_is_caught(added):
    src, plan, out = added

    def edit(root):
        t = next(root.iter(W + "t"))
        t.text = t.text.replace("ACME", "ACNE")
    rec = suite(src, mutate_header(out, edit), plan)
    assert {"V2", "V10"} <= set(failed(rec))


def test_a_watermark_left_behind_by_remove_is_caught():
    src = word_watermarked("DRAFT")
    plan, out, rec = run(src, "remove_watermark")
    assert not failed(rec)
    rec = suite(src, src, plan)                                     # nothing taken out
    assert failed(rec) == ["V10"] and any("not exactly the predicted change" in e for e in v10(rec))


# ---- planted defects: wrong implementations shared by plan and apply ---------------------------------------------

def test_a_shared_template_defect_is_caught(monkeypatch):
    monkeypatch.setattr(wm, "Z_INDEX", "251654144")                   # in front of the text: plan and apply agree
    _, _, rec = run(document(), text="DRAFT")
    assert failed(rec) == ["V10"] and any("not exactly the predicted change" in e for e in v10(rec))


def test_a_shared_formula_defect_is_caught(monkeypatch):
    monkeypatch.setattr(wm, "_FORMULAS", wm._FORMULAS[:-1])
    _, _, rec = run(document(), text="DRAFT")
    assert failed(rec) == ["V10"]


def test_a_shared_sizing_defect_is_caught(monkeypatch):
    monkeypatch.setattr(wm, "DIAGONAL_FILL", 1.3)                     # off the page: plan and apply agree
    _, _, rec = run(document(), text="DRAFT")
    assert failed(rec) == ["V10"] and any("breaks the sizing rule" in e for e in v10(rec))


def test_a_shared_header_set_defect_is_caught(monkeypatch):
    real = wm.headers_shown

    def default_only(pkg):
        order, showing, gaps = real(pkg)
        return order[:1], {k: v for k, v in showing.items() if k == order[0]}, gaps
    monkeypatch.setattr(wm, "headers_shown", default_only)            # the cover's first-page header forgotten
    _, _, rec = run(document(title_page=True), text="DRAFT")
    assert failed(rec) == ["V10"] and any("the document shows the headers" in e for e in v10(rec))


def test_a_shared_anchor_defect_is_caught(monkeypatch):
    monkeypatch.setattr(wm, "first_paragraph", lambda root: list(root.iter(W + "p"))[-1])
    src = with_header(document(), HEADER, lambda r: r.append(etree.fromstring('<w:p %s/>' % NS_DECL)))
    _, _, rec = run(src, text="DRAFT")
    assert failed(rec) == ["V10"] and any("belongs in paragraph 0" in e for e in v10(rec))


def test_a_shared_detection_defect_is_caught(monkeypatch):
    real = wm.watermark_runs
    monkeypatch.setattr(wm, "watermark_runs",
                        lambda root, part: [f for f in real(root, part) if f["text"] is not None])   # pictures missed
    src = patch_docx(picture_watermarked(), {HEADER: first_para_run(word_vml("DRAFT"), at=1)})
    _, _, rec = run(src, "remove_watermark")
    assert failed(rec) == ["V10"] and any("the watermarks here are" in e for e in v10(rec))


def test_a_shared_id_defect_is_caught(monkeypatch):
    """The ids picked without looking at the package: the new shape number repeats a shape's of the body."""
    body_shape = ('<w:r %s><w:pict><v:shape id="_x0000_s2049" o:spid="_x0000_s2049" style="width:20pt;height:20pt"/>'
                  '</w:pict></w:r>' % NS_DECL)
    src = patch_docx(document(), {"word/document.xml": lambda r: r.find(W + "body/" + W + "p").append(
        etree.fromstring(body_shape))})
    monkeypatch.setattr(wm, "package_ids", lambda pkg: (set(), wm.FIRST_SPID - 1))
    _, out, rec = run(src, text="DRAFT")
    assert shapes(out)[0].get(O + "spid") == "_x0000_s2049"
    assert failed(rec) == ["V10"] and any("not new and unique" in e for e in v10(rec))


def test_a_shared_detection_defect_that_takes_another_shape_is_caught(monkeypatch):
    """Every VML shape read as a watermark: plan and apply take the WordArt out with the watermark."""
    other = ('<w:r %s><w:pict>%s<v:shape id="_x0000_s1030" type="#_x0000_t136" style="width:10pt;height:5pt">'
             '<v:textpath string="WordArt"/></v:shape></w:pict></w:r>' % (NS_DECL, WORD_SHAPETYPE))
    src = with_header(document(), HEADER, lambda r: (first_para_run(word_vml(shapetype=False))(r),
                                                     first_para_run(other, at=2)(r)))
    monkeypatch.setattr(wm, "PREFIXES", wm.PREFIXES + ("_x0000_s",))
    _, out, rec = run(src, "remove_watermark")
    assert not shapes(out)
    assert failed(rec) == ["V10"] and any("the watermarks here are" in e for e in v10(rec))


def test_a_shared_defect_that_ignores_a_still_used_definition_is_caught(monkeypatch):
    monkeypatch.setattr(wm, "check_shapetype_users", lambda *a, **k: None)
    other = ('<w:r %s><w:pict><v:shape id="_x0000_s1030" type="#_x0000_t136" style="width:10pt;height:5pt">'
             '<v:textpath string="WordArt"/></v:shape></w:pict></w:r>' % NS_DECL)
    src = with_header(document(), HEADER, lambda r: (first_para_run(word_vml())(r), first_para_run(other, at=2)(r)))
    _, _, rec = run(src, "remove_watermark")
    assert "V10" in failed(rec) and any("_x0000_t136" in e for e in v10(rec))


# ---- conflicts -----------------------------------------------------------------------------------------------------

def test_what_the_watermark_reads_and_changes_is_declared():
    src = document()
    pkg = DocxPackage(src)
    add = get_op("add_watermark").plan(pkg, {}, {"text": "DRAFT"})
    header_edit = get_op("replace_table").plan(pkg, {"scope": ["headers"]},
                                               {"pairs": [{"find": "ACME", "replace": "ACNE"}], "author": "A"})
    with pytest.raises(JobException) as e:
        merge_fragments([add, header_edit])
    assert e.value.code == "conflict"
    margins = get_op("set_page_setup").plan(pkg, {"sections": "all"}, {"margins": {"left": 1800}})
    with pytest.raises(JobException) as e:
        merge_fragments([add, margins])
    assert e.value.code == "conflict" and "pgMar" in e.value.reason
    body_edit = get_op("replace_table").plan(pkg, {"scope": ["document"]},
                                             {"pairs": [{"find": "Article", "replace": "Section"}], "author": "A"})
    with pytest.raises(JobException):
        merge_fragments([add, body_edit])                       # untracked + tracked content: two specs (the mode rule)
    body_edit_clean = get_op("replace_table").plan(pkg, {"scope": ["document"]},
                                                   {"pairs": [{"find": "Article", "replace": "Section"}],
                                                    "tracked": False})
    assert merge_fragments([add, body_edit_clean]).mode == "untracked"


# ---- the rendered pages ----------------------------------------------------------------------------------------------

def _grey(pdf_path: Path) -> List[int]:
    import pymupdf
    out = []
    with pymupdf.open(str(pdf_path)) as pdf:
        for pg in pdf:
            pix = pg.get_pixmap(dpi=40)
            s, n = pix.samples, pix.n
            out.append(sum(1 for i in range(0, len(s), n)
                           if 205 <= s[i] <= 240 and abs(s[i] - s[i + 1]) <= 3 and abs(s[i + 1] - s[i + 2]) <= 3))
    return out


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")
def test_the_watermark_is_drawn_behind_every_page(tmp_path):
    from spec_surgeon.render import libreoffice
    src = document(second="inherit", body=60)
    _, out, rec = run(src, text="DRAFT")
    assert not failed(rec)
    (tmp_path / "src.docx").write_bytes(src)
    (tmp_path / "out.docx").write_bytes(out)
    pdfs = libreoffice.render_pdfs([tmp_path / "src.docx", tmp_path / "out.docx"], tmp_path / "pdf", workers=1)
    if len(pdfs) != 2:
        pytest.skip("LibreOffice did not render here")
    before, after = _grey(pdfs[tmp_path / "src.docx"]), _grey(pdfs[tmp_path / "out.docx"])
    assert len(before) == len(after) >= 3
    assert all(a > b + 1000 for a, b in zip(after, before)), (before, after)
