"""EditSpec font defaults must not turn preserved glyph fonts into universal final requirements."""
import copy
import io
import json
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from fake_claude import FakeClient, resp, tool
from make_fixtures import build_rich
from spec_surgeon.chain import run_chain
from spec_surgeon.ops.selectors import DEFAULT_PRESERVE_FONTS
from spec_surgeon.recovery import RecoveryOptions, recover
from test_chain_outcomes import _document, _font


def _recover(source, goal, out, candidate=None):
    client = FakeClient(resp(tool("submit_repair", {"candidate": {
        "kind": "spec", "spec": candidate if candidate is not None else goal}})))
    return recover([source], goal["request"], out, spec=goal, client=client,
                   options=RecoveryOptions(render=False, use_memory=False, max_attempts=1))


@pytest.mark.parametrize("explicit_defaults", [False, True])
def test_implicit_and_explicit_symbol_exceptions_never_infer_universal_font(tmp_path, explicit_defaults):
    source = tmp_path / "symbols.docx"
    source.write_bytes(build_rich())
    original = source.read_bytes()
    goal = json.loads((Path(__file__).resolve().parents[1] / "examples/specs/change-all-font-arial.json").read_text())
    if explicit_defaults:
        goal["assertions"][0]["preserve_fonts"] = list(DEFAULT_PRESERVE_FONTS)
    frozen_goal = copy.deepcopy(goal)

    ordinary = run_chain([str(source)], {"stages": [{"spec": goal}]}, tmp_path / "ordinary",
                         workers=1, do_render=False)
    assert ordinary.files[0].status == "PASS", ordinary.files[0].reason
    result = _recover(source, goal, tmp_path / "recovery")
    assert result.key["assertions"] == [], "Symbol exceptions were incorrectly strengthened to a universal font"
    assert not result.ready
    assert "source-relative" in result.files[0].reason
    assert not (result.out_dir / "attempts/001/chain.json").exists()
    assert source.read_bytes() == original and goal == frozen_goal


def test_explicit_empty_preservation_list_allows_supported_universal_final_check(tmp_path):
    source = tmp_path / "plain.docx"
    source.write_bytes(_document())
    goal = _font("Courier New")
    goal["assertions"] = [{"kind": "effective_font", "value": "Courier New", "preserve_fonts": []}]
    result = _recover(source, goal, tmp_path / "recovery")
    assert result.ready, result.to_text()
    final = result.key["assertions"][0]
    assert "preserve_fonts" not in final
    assert final["slots"] == ["ascii", "hAnsi", "cs"]
    assert final["scope"] == ["document"] and final["include_paragraph_marks"] is True
    assert result.ready[0].outcomes[0]["status"] == "PASS"


@pytest.mark.parametrize("replacement", [[], [{"kind": "effective_font", "value": "Courier New", "preserve_fonts": []}]])
def test_candidate_cannot_erase_an_original_implicit_preservation_rule(tmp_path, replacement):
    source = tmp_path / "plain.docx"
    source.write_bytes(_document())
    goal = _font("Courier New")
    goal["assertions"] = [{"kind": "effective_font", "value": "Courier New"}]
    candidate = copy.deepcopy(goal)
    candidate["assertions"] = replacement
    result = _recover(source, goal, tmp_path / "recovery", candidate)
    assert not result.ready
    assert "source-relative" in result.files[0].reason


def test_explicit_empty_font_preservation_does_not_remove_symbol_run_size_exception(tmp_path):
    doc = Document()
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_paragraph("Ordinary specification text")
    run = doc.add_paragraph().add_run("Symbol annotation")
    run.font.size = Pt(20)
    symbol = OxmlElement("w:sym")
    symbol.set(qn("w:font"), "Wingdings")
    symbol.set(qn("w:char"), "F0B7")
    run._r.append(symbol)
    saved = io.BytesIO()
    doc.save(saved)
    source = tmp_path / "symbol-size.docx"
    source.write_bytes(saved.getvalue())
    original = source.read_bytes()
    goal = {"request": "Set the header distance while retaining the specified font contract",
            "summary": "Header distance", "mode": "formatting", "ops": [
                {"op": "set_page_setup", "params": {"margins": {"header": 900}}}],
            "assertions": [{"kind": "effective_font", "size": 11, "preserve_fonts": []}]}

    # A text-bearing w:sym run is exempt from the EditSpec size check even with preserve_fonts: [].
    # The ordinary checker accepts its 20-point size; a universal 11-point final assertion would strengthen it.
    ordinary = run_chain([str(source)], {"stages": [{"spec": goal}]}, tmp_path / "ordinary",
                         workers=1, do_render=False)
    assert ordinary.files[0].status == "PASS", ordinary.files[0].reason
    result = _recover(source, goal, tmp_path / "recovery")
    assert all(a["kind"] != "effective_font" for a in result.key["assertions"])
    assert not result.ready
    assert "source-relative" in result.files[0].reason
    assert not (result.out_dir / "attempts/001/chain.json").exists()
    assert source.read_bytes() == original
