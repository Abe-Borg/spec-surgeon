"""Original-to-final operation goals reject corruption without stage permissions."""
from __future__ import annotations

import copy
import hashlib
import io

import pytest
from docx import Document
from lxml import etree

from spec_surgeon.ops import get_op
from spec_surgeon.package import DocxPackage, W
from spec_surgeon.verify.effects import verify_effects
from spec_surgeon.verify.goals import _hash_json


def source_bytes():
    d = Document()
    d.add_paragraph("FIRE SUPPRESSION")
    d.add_paragraph("Retain every original word.")
    d.add_table(rows=2, cols=2).cell(0, 0).text = "SCHEDULE"
    stream = io.BytesIO()
    d.save(stream)
    return stream.getvalue()


def apply(data, name, select, params):
    pkg = DocxPackage(data)
    op = get_op(name)
    plan = op.plan(pkg, select, params)
    out = pkg.copy()
    op.apply(out, plan)
    return out.output_bytes()


def goal(effect, data, count):
    c = {"kind": "document_effects", "effects": [{"id": "approved", **effect}]}
    b = {"source_sha256": hashlib.sha256(data).hexdigest(), "contract_sha256": _hash_json(c),
         "effects": {"approved": {"targets": [], "choice": "unique", "source_count": count}}}
    return c, b


def mutate(data, fn, part="word/document.xml"):
    pkg = DocxPackage(data)
    root = pkg.xml(part)
    fn(root)
    pkg.write_xml(part, root)
    return pkg.output_bytes()


@pytest.mark.parametrize("kind,op,select,params,count", [
    ("paragraph_style", "apply_style", {"paragraphs": {"text_regex": "^FIRE"}},
     {"style": "Heading 1", "tracked": False}, 1),
    ("paragraph_style", "apply_style", {"paragraphs": {"text_regex": "^FIRE"}},
     {"style": "Heading 1", "tracked": True, "author": "Abraham Borg", "direct_formatting": "clear"}, 1),
    ("style_mapping", "map_styles", {"scope": ["document"]},
     {"map": [{"from": "Normal", "to": "Body Text"}]}, 6),
    ("proofing_language", "set_language", {}, {"language": "en-CA"}, 1),
    ("table_format", "set_table_props", {"tables": {"first_row_regex": "SCHEDULE"}},
     {"width": {"pct": 100}, "repeat_header_rows": 1, "cant_split": True}, 1),
    ("table_style", "apply_table_style", {}, {"style": "Light Shading Accent 1"}, 1),
])
def test_source_derived_proof_and_corruption(kind, op, select, params, count):
    source = source_bytes()
    output = apply(source, op, select, params)
    c, b = goal({"type": kind, "select": select, "params": params, "expected_count": count}, source, count)
    result = verify_effects(source, output, c, b)
    assert result["passed"], result
    corrupted = mutate(output, lambda root: setattr(next(root.iter(W + "t")), "text", "UNAPPROVED"))
    assert not verify_effects(source, corrupted, c, b)["passed"]
    b["effects"]["approved"]["source_count"] += 1
    assert not verify_effects(source, output, c, b)["passed"]


def test_all_count_is_pinned_to_source_and_selector_cannot_expand():
    source = source_bytes()
    params = {"style": "Heading 1", "tracked": False}
    sel = {"paragraphs": {"text_regex": "^FIRE"}}
    c, b = goal({"type": "paragraph_style", "select": sel, "params": params, "expected_count": "all"}, source, 1)
    out = apply(source, "apply_style", sel, params)
    assert verify_effects(source, out, c, b)["passed"]
    expanded = apply(source, "apply_style", {}, params)
    assert not verify_effects(source, expanded, c, b)["passed"]


def test_proofing_language_cannot_change_other_style_properties():
    source = source_bytes()
    params = {"language": "en-CA"}
    output = apply(source, "set_language", {}, params)
    c, b = goal({"type": "proofing_language", "params": params, "expected_count": 1}, source, 1)
    corrupted = mutate(output, lambda root: root.find(W + "style").set(W + "customStyle", "1"), "word/styles.xml")
    assert not verify_effects(source, corrupted, c, b)["passed"]


def test_pinned_style_template_import_and_drift(tmp_path):
    from docx.enum.style import WD_STYLE_TYPE
    source = source_bytes()
    d = Document()
    d.styles.add_style("Office Clause", WD_STYLE_TYPE.PARAGRAPH)
    template = tmp_path / "office.docx"
    d.save(template)
    params = {"style": "Office Clause", "tracked": False}
    sel = {"paragraphs": {"text_regex": "^FIRE"}}
    output = apply(source, "apply_style", sel, {**params, "template": str(template)})
    c, b = goal({"type": "paragraph_style", "select": sel, "params": params, "expected_count": 1,
                 "template": str(template), "template_sha256": hashlib.sha256(template.read_bytes()).hexdigest()}, source, 1)
    result = verify_effects(source, output, c, b)
    assert result["passed"], result
    template.write_bytes(source)
    assert not verify_effects(source, output, c, b)["passed"]


def test_page_content_selection_and_first_page_false_are_verified():
    source = source_bytes()
    output = apply(source, "set_page_setup", {"sections": {"with_tables": True}}, {"different_first_page": True})
    c = {"kind": "document_effects", "effects": [{"id": "cover", "type": "page_setup",
          "sections": {"with_tables": True}, "different_first_page": True}]}
    b = {"source_sha256": hashlib.sha256(source).hexdigest(), "contract_sha256": _hash_json(c),
         "effects": {"cover": {"targets": [], "choice": "unique"}}}
    result = verify_effects(source, output, c, b)
    assert result["passed"], result
    assert not verify_effects(source, source, c, b)["passed"]


@pytest.mark.parametrize("kind,params", [("numbering", {}), ("typed_numbering", {}),
                                         ("typed_numbering", {"tracked": False, "direct_formatting": "clear"})])
def test_numbering_effect_complete_residual(tmp_path, kind, params):
    from make_fixtures import SCHEME_LEVELS, build_numbering, build_scheme_example, build_typed
    from spec_surgeon.verify.goal_numbering import derive
    from spec_surgeon.verify.props import Ctx
    from spec_surgeon.verify.reader import VPackage
    source = build_typed() if kind == "typed_numbering" else build_numbering(shared_restart=False)
    template = tmp_path / "scheme.docx"
    template.write_bytes(build_scheme_example())
    params = {"levels": list(SCHEME_LEVELS), "author": "Abraham Borg", **params}
    effect = {"id": "approved", "type": kind, "params": params, "select": {}, "expected_count": "all",
              "template": str(template), "template_sha256": hashlib.sha256(template.read_bytes()).hexdigest()}
    pkg = VPackage(source)
    try:
        _, count = derive(Ctx(pkg, pkg, {}), effect)
    finally:
        pkg.zip.close()
    output = apply(source, "convert_typed_numbers" if kind == "typed_numbering" else "convert_numbering", {},
                   {**params, "template": str(template)})
    c, b = goal(effect, source, count)
    result = verify_effects(source, output, c, b)
    assert result["passed"], result
    assert not verify_effects(source, source, c, b)["passed"]
    # Numbering permission cannot conceal a mutation in an unrelated part.
    corrupted = mutate(output, lambda root: root.set("unapproved", "changed"), "word/settings.xml")
    assert not verify_effects(source, corrupted, c, b)["passed"]


def test_property_approval_cannot_move_untouched_paragraph():
    source = source_bytes()
    params = {"style": "Heading 1", "tracked": False}
    sel = {"paragraphs": {"text_regex": "^FIRE"}}
    output = apply(source, "apply_style", sel, params)
    c, b = goal({"type": "paragraph_style", "select": sel, "params": params, "expected_count": 1}, source, 1)

    def move(root):
        body = root.find(W + "body")
        paragraph = body.findall(W + "p")[1]
        body.remove(paragraph)
        next(root.iter(W + "tc")).append(paragraph)

    assert not verify_effects(source, mutate(output, move), c, b)["passed"]


def test_tables_and_page_content_selector_compose():
    source = source_bytes()
    table = {"id": "tables", "type": "table_format", "params": {"cant_split": True}, "expected_count": "all"}
    page = {"id": "pages", "type": "page_setup", "sections": {"containing": {"text_regex": "^SCHEDULE"}},
            "different_first_page": False, "orientation": "landscape"}
    output = apply(source, "set_table_props", {}, table["params"])
    output = apply(output, "set_page_setup", {"sections": page["sections"]},
                   {"different_first_page": False, "orientation": "landscape"})
    c = {"kind": "document_effects", "effects": [table, page]}
    b = {"source_sha256": hashlib.sha256(source).hexdigest(), "contract_sha256": _hash_json(c), "effects": {
         "tables": {"targets": [], "choice": "unique", "source_count": 1},
         "pages": {"targets": [], "choice": "unique"}}}
    result = verify_effects(source, output, c, b)
    assert result["passed"], result


def test_first_page_boolean_cannot_hide_xml_payload():
    source = source_bytes()
    output = apply(source, "set_page_setup", {}, {"different_first_page": True})
    c = {"kind": "document_effects", "effects": [{"id": "cover", "type": "page_setup", "different_first_page": True}]}
    b = {"source_sha256": hashlib.sha256(source).hexdigest(), "contract_sha256": _hash_json(c),
         "effects": {"cover": {"targets": [], "choice": "unique"}}}
    bad = mutate(output, lambda root: setattr(next(root.iter(W + "titlePg")), "text", "UNAPPROVED"))
    assert not verify_effects(source, bad, c, b)["passed"]


@pytest.mark.parametrize("payload", ["duplicate", "text"])
def test_language_property_mask_preserves_unchecked_containers(payload):
    source = source_bytes()

    def original_language(root):
        run = next(root.iter(W + "r"))
        rpr = etree.Element(W + "rPr")
        etree.SubElement(rpr, W + "lang").set(W + "val", "en-US")
        run.insert(0, rpr)

    source = mutate(source, original_language)
    params = {"language": "en-CA"}
    output = apply(source, "set_language", {}, params)
    c, b = goal({"type": "proofing_language", "params": params, "expected_count": 1}, source, 1)

    def corrupt(root):
        run = next(root.iter(W + "r"))
        if payload == "duplicate":
            extra = etree.SubElement(run, W + "rPr")
            etree.SubElement(extra, W + "vanish")
        else:
            run.find(W + "rPr").text = "UNAPPROVED"

    assert not verify_effects(source, mutate(output, corrupt), c, b)["passed"]


def test_table_style_and_row_format_composition():
    source = source_bytes()
    style = {"id": "style", "type": "table_style", "params": {"style": "Light Shading Accent 1"}, "expected_count": 1}
    rows = {"id": "rows", "type": "table_format", "params": {"repeat_header_rows": 1, "cant_split": True}, "expected_count": 1}
    output = apply(source, "apply_table_style", {}, style["params"])
    output = apply(output, "set_table_props", {}, rows["params"])
    c = {"kind": "document_effects", "effects": [style, rows]}
    b = {"source_sha256": hashlib.sha256(source).hexdigest(), "contract_sha256": _hash_json(c), "effects": {
         e["id"]: {"targets": [], "choice": "unique", "source_count": 1} for e in c["effects"]}}
    result = verify_effects(source, output, c, b)
    assert result["passed"], result
    rows["params"]["width"] = {"pct": 100}
    b["contract_sha256"] = _hash_json(c)
    result = verify_effects(source, output, c, b)
    assert not result["passed"] and "separate approved goals" in str(result["errors"])


def test_paragraph_mapping_and_run_format_share_original_boundaries(tmp_path):
    from spec_surgeon.goals import resolve_goal, verify_goal
    from test_goal_contracts import execute
    source = tmp_path / "spec.docx"
    source.write_bytes(source_bytes())
    c = {"kind": "document_effects", "effects": [
        {"id": "styles", "type": "style_mapping", "select": {"scope": ["document"]},
         "params": {"map": [{"from": "Normal", "to": "Body Text"}]}, "expected_count": "all"},
        {"id": "font", "type": "format", "target": {"all": True}, "run": {"b": True}, "tracked": False},
    ]}
    binding = resolve_goal(source, c)["binding"]
    output = execute(source, c, binding)
    corrupted = mutate(output, lambda root: setattr(next(root.iter(W + "t")), "text", "UNAPPROVED"))
    assert not verify_goal(source, corrupted, c, binding)["passed"]


def test_typed_conversion_envelope_cannot_hide_empty_run_or_text_metadata(tmp_path):
    from make_fixtures import SCHEME_LEVELS, build_scheme_example, build_typed
    from spec_surgeon.goals import resolve_goal, verify_goal
    from test_goal_contracts import execute
    source = tmp_path / "spec.docx"
    source.write_bytes(build_typed())
    template = tmp_path / "scheme.docx"
    template.write_bytes(build_scheme_example())
    c = {"kind": "document_effects", "effects": [{"id": "numbers", "type": "typed_numbering", "expected_count": "all",
         "params": {"levels": list(SCHEME_LEVELS)}, "template": str(template),
         "template_sha256": hashlib.sha256(template.read_bytes()).hexdigest()}]}
    binding = resolve_goal(source, c)["binding"]
    output = execute(source, c, binding)

    def corrupt(root):
        p = next(p for p in root.iter(W + "p") if any(t.text == "GENERAL" for t in p.iter(W + "t")))
        r = etree.SubElement(p, W + "r")
        etree.SubElement(etree.SubElement(r, W + "rPr"), W + "vanish")

    assert not verify_goal(source, mutate(output, corrupt), c, binding)["passed"]


@pytest.mark.parametrize("other", [
    {"id": "bold", "type": "format", "target": {"all": True}, "run": {"b": True}},
    {"id": "font", "type": "format", "target": {"all": True}, "font": {"name": "Arial"}},
    {"id": "style", "type": "paragraph_style", "params": {"style": "Heading 1"}, "expected_count": "all"},
    {"id": "table", "type": "table_style", "params": {"style": "Table Grid"}, "expected_count": "all"},
])
def test_language_full_property_proof_conflicts_are_refused_before_execution(other):
    from spec_surgeon.verify.effects import _validate
    c = {"kind": "document_effects", "effects": [
        {"id": "language", "type": "proofing_language", "params": {"language": "en-US"}, "expected_count": 1}, other]}
    with pytest.raises(ValueError, match="separate approved goals"):
        _validate(c)


def test_cleared_formatting_and_overrides_need_separate_approved_goals():
    from spec_surgeon.verify.effects import _validate
    clear = {"id": "clear", "type": "paragraph_style", "params": {"style": "Heading 1", "direct_formatting": "clear"},
             "expected_count": "all"}
    bold = {"id": "bold", "type": "format", "target": {"all": True}, "run": {"b": True}}
    with pytest.raises(ValueError, match="clearing direct formatting needs a separate"):
        _validate({"kind": "document_effects", "effects": [clear, bold]})
    _validate({"kind": "document_effects", "effects": [clear,
              {"id": "page", "type": "page_setup", "different_first_page": True}]})


def test_language_and_paragraph_only_format_compose(tmp_path):
    from test_goal_contracts import execute
    source = tmp_path / "spec.docx"
    source.write_bytes(source_bytes())
    c = {"kind": "document_effects", "effects": [
        {"id": "language", "type": "proofing_language", "params": {"language": "en-CA"}, "expected_count": 1},
        {"id": "paragraphs", "type": "format", "target": {"all": True}, "paragraph": {"keepLines": True}},
    ]}
    execute(source, c)
