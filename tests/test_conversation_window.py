"""The installed entry point uses the conversation UI and preserves approval/lifetime boundaries."""
import threading
from pathlib import Path

import pytest

from spec_surgeon.gui import Session, build_window
from spec_surgeon.settings import Settings
from test_document_task import REQUEST, client, document


@pytest.fixture
def qt(monkeypatch):
    from qt_support import close_windows, qt_app_or_skip
    app = qt_app_or_skip()
    monkeypatch.setenv("SPEC_SURGEON_NO_UPDATE_CHECK", "1")
    windows = []
    yield app, windows
    close_windows(app, windows)


def window(qt, tmp_path, **kwargs):
    from spec_surgeon.ui.app import RecordingPrompts
    w = build_window(Session(client=client(), out_root=tmp_path / "out", visual_check=False),
                     settings=Settings(), prompts=RecordingPrompts(), **kwargs)
    w.task.render = False
    qt[1].append(w)
    w.show()
    qt[0].processEvents()
    return w


def test_default_window_has_only_documents_template_and_chat(qt, tmp_path):
    from PySide6.QtWidgets import QComboBox, QSlider, QSpinBox
    from spec_surgeon.ui.conversation_window import ConversationWindow
    w = window(qt, tmp_path)
    assert isinstance(w, ConversationWindow)
    assert not w.findChildren(QComboBox)
    assert not w.findChildren(QSlider)
    assert not w.findChildren(QSpinBox)
    assert not hasattr(w, "files_page")
    assert not w.primary.isVisible()
    assert not w.send.isEnabled()


def test_chat_to_verified_copies_and_unsent_revision_blocks_approval(qt, tmp_path):
    source = document(tmp_path / "Spec.docx")
    w = window(qt, tmp_path)
    w.add_documents([str(source)])
    w.composer.setPlainText(REQUEST)
    w.do_send()
    w._spec_surgeon["wait"]()
    assert w.primary.isEnabled()
    assert w.primary.text() == "Create edited copies"
    w.composer.setPlainText("Actually, only the cover")
    assert not w.primary.isEnabled()
    w.composer.clear()
    w.do_primary()
    w._spec_surgeon["wait"]()
    assert w.task.stage == "done"
    assert len(w.task.results) == 1
    assert Path(next(iter(w.task.results.values()))).is_file()


def test_multiple_file_drop_does_not_choose_containing_folder(qt, tmp_path):
    from PySide6.QtCore import QMimeData, QPointF, Qt, QUrl
    from PySide6.QtGui import QDropEvent
    first = document(tmp_path / "a" / "One.docx")
    second = document(tmp_path / "b" / "Two.docx")
    document(tmp_path / "a" / "Unselected.docx")
    w = window(qt, tmp_path)
    mime = QMimeData()
    mime.setUrls([QUrl.fromLocalFile(str(p)) for p in (first, second)])
    event = QDropEvent(QPointF(10, 10), Qt.DropAction.CopyAction, mime,
                       Qt.MouseButton.LeftButton, Qt.KeyboardModifier.NoModifier)
    w.centralWidget().dropEvent(event)
    assert w.session.files() == [first, second]


def test_attachment_mutation_and_close_are_blocked_during_work(qt, tmp_path, monkeypatch):
    source = document(tmp_path / "Spec.docx")
    w = window(qt, tmp_path)
    w.add_documents([str(source)])
    started, finish = threading.Event(), threading.Event()
    def work(*args, **kwargs):
        started.set()
        assert finish.wait(10)
        return w.task.view()
    monkeypatch.setattr(w.task, "perform", work)
    w.composer.setPlainText(REQUEST)
    w.do_send()
    assert started.wait(5)
    try:
        assert w.busy and not w.attachments.isEnabled()
        w.remove_document(source)
        assert w.session.documents == [source]
        w.close()
        assert w.isVisible()
    finally:
        finish.set()
        w._spec_surgeon["wait"]()
    assert not w.busy


def test_no_key_preserves_composer_and_does_not_offer_recipe_modes(qt, tmp_path, monkeypatch):
    from PySide6.QtWidgets import QPushButton
    w = window(qt, tmp_path)
    monkeypatch.setattr(w.session, "claude_ready", lambda: False)
    w.add_documents([str(document(tmp_path / "Spec.docx"))])
    w.composer.setPlainText("My request stays here")
    assert not w.send.isEnabled()
    assert w.composer.isVisible()
    assert w.composer.toPlainText() == "My request stays here"
    visible = [button.text() for button in w.findChildren(QPushButton) if button.isVisible()]
    assert not any("recipe" in text.lower() or "operation" in text.lower() for text in visible)


def test_document_text_is_plain_text_in_chat(qt, tmp_path):
    from PySide6.QtCore import Qt
    from PySide6.QtWidgets import QLabel
    w = window(qt, tmp_path)
    w.task.say('<img src="https://example.com/tracker">', "assistant")
    w._render()
    text = next(label for label in w.findChildren(QLabel) if "tracker" in label.text())
    assert text.textFormat() == Qt.TextFormat.PlainText


def test_stop_action_prevents_the_next_model_call(qt, tmp_path, monkeypatch):
    from spec_surgeon import llm
    from fake_claude import FakeClient, resp, text
    w = window(qt, tmp_path)
    w.add_documents([str(document(tmp_path / "Spec.docx"))])
    fake = FakeClient(resp(text("first response")), resp(text("must not run")))
    reached, release = threading.Event(), threading.Event()
    def work(*args, **kwargs):
        arguments = dict(model=llm.DEFAULT_MODEL, system="s", tools=[], messages=[])
        llm.call(fake, **arguments)
        reached.set()
        assert release.wait(10)
        llm.call(fake, **arguments)
    monkeypatch.setattr(w.task, "perform", work)
    w.composer.setPlainText(REQUEST)
    w.do_send()
    assert reached.wait(5)
    try:
        assert w.primary.text() == "Stop after current step"
        w.do_primary()
        assert not w.primary.isEnabled()
    finally:
        release.set()
        w._spec_surgeon["wait"]()
    assert len(fake.requests) == 1
    assert "Stopped before the next AI step" in w.notice.text()


def test_rendered_page_growth_is_visible_in_the_conversation(qt, tmp_path):
    from PySide6.QtWidgets import QLabel
    w = window(qt, tmp_path)
    w.add_documents([str(document(tmp_path / "Spec.docx"))])
    w.task.perform("send", text=REQUEST)
    sample = w.session.preview.samples[0]
    sample.pages_before, sample.pages_after = 2, 3
    w._render()
    assert any(label.text() == "Pages: 2 → 3 (+1)" for label in w.findChildren(QLabel))


def test_installed_smoke_uses_default_chat_and_delivers_verified_copies(qt, tmp_path):
    import importlib.util
    import json
    from spec_surgeon import catalog
    from spec_surgeon.ui.app import RecordingPrompts
    path = Path(__file__).resolve().parents[1] / "packaging/windows/installed_smoke.py"
    module = importlib.util.spec_from_file_location("_conversation_smoke", path)
    smoke = importlib.util.module_from_spec(module)
    module.loader.exec_module(smoke)
    sources = [document(tmp_path / name) for name in ("One.docx", "Two.docx")]
    for path in sources:
        from docx import Document
        doc = Document(path)
        doc.sections[0].header.paragraphs[0].text = "Remove this header"
        doc.save(path)
    original = [p.read_bytes() for p in sources]
    spec = json.loads((catalog.EXAMPLES_DIR / "remove-all-headers-footers.json").read_text())
    session = Session(client=smoke._offline_client(spec), out_root=tmp_path / "out", visual_check=False)
    w = build_window(session, settings=Settings(), prompts=RecordingPrompts())
    qt[1].append(w)
    w.show()
    smoke._conversation_run(w, sources)
    assert original == [p.read_bytes() for p in sources]
    for output in w.task.results.values():
        assert not Document(output).sections[0].header.paragraphs[0].text


def test_choices_are_buttons_and_a_click_answers(qt, tmp_path):
    from fake_claude import FakeClient, resp, tool
    from test_document_task import SPEC
    fake = FakeClient(resp(tool("ask_user", {"message": "Two titles match.", "choices": [
        {"kind": "clarify", "label": "Only the first title"}, {"kind": "clarify", "label": "Both titles"}]})),
        resp(tool("submit_editspec", {"spec": SPEC, "message": "I’ll bold the first title."})))
    w = window(qt, tmp_path)
    w.session.client = fake
    w.add_documents([str(document(tmp_path / "Spec.docx"))])
    w.composer.setPlainText(REQUEST)
    w.do_send()
    w._spec_surgeon["wait"]()
    assert w.task.stage == "question" and not w.primary.isVisible()
    assert [b.text() for b in w._choice_buttons] == ["Only the first title", "Both titles"]
    assert all(b.isEnabled() for b in w._choice_buttons)
    w._choice_buttons[1].click()
    w._spec_surgeon["wait"]()
    assert w.task.stage == "preview" and w.primary.text() == "Create edited copies"
    assert "Both titles" in w.task.decisions
    assert [m["text"] for m in w.task.messages if m["role"] == "assistant"][-1] == "I’ll bold the first title."
    assert not any("•" in m["text"] for m in w.task.messages)


def test_a_mixed_batch_shows_one_card_per_layout_with_its_assumptions(qt, tmp_path):
    from PySide6.QtWidgets import QLabel
    from fake_claude import FakeClient, resp, tool
    from make_fixtures import build, build_styles
    from spec_surgeon.fingerprint import fingerprint
    from spec_surgeon.package import DocxPackage
    typed, styled = tmp_path / "a" / "21 13 13.docx", tmp_path / "b" / "21 13 16.docx"
    typed.parent.mkdir(); styled.parent.mkdir()
    typed.write_bytes(build("21 13 13", "WET-PIPE SPRINKLER SYSTEMS"))
    styled.write_bytes(build_styles("21 13 16", "DRY-PIPE SPRINKLER SYSTEMS"))
    pid = lambda p: fingerprint(DocxPackage.from_path(p))["pattern_id"]  # noqa: E731
    bold = {"run": {"b": True}, "tracked": False}
    submission = {"summary": "two layouts", "layouts": [
        {"id": "A", "label": "typed numbers, no styles", "pattern_ids": [pid(typed)],
         "spec": {"summary": "by text", "mode": "formatting", "assumptions": ["Body paragraphs start with A. to E."],
                  "ops": [{"op": "set_properties", "select": {"paragraphs": {"text_regex": r"^[A-E]\.\t"}}, "params": bold}]}},
        {"id": "B", "label": "styled body", "pattern_ids": [pid(styled)],
         "spec": {"summary": "by style", "mode": "formatting",
                  "ops": [{"op": "set_properties", "select": {"paragraphs": {"style_in": ["ArchPara"]}}, "params": bold}]}}]}
    w = window(qt, tmp_path)
    w.session.client = FakeClient(resp(tool("submit_editspec", {"spec": submission, "message": "Two layouts, one plan."})))
    w.add_documents([str(typed), str(styled)])
    w.composer.setPlainText("Make the body paragraphs bold.")
    w.do_send()
    w._spec_surgeon["wait"]()
    assert w.task.stage == "preview" and w.primary.text() == "Create edited copies"
    texts = [label.text() for label in w.findChildren(QLabel)]
    assert "Your documents use 2 layouts" in texts
    assert "typed numbers, no styles · 1 document" in texts
    assert "styled body · 1 document" in texts
    assert "Assuming: Body paragraphs start with A. to E." in texts
    assert not any(t.startswith("• Assumptions") for t in texts)
    assert all(not t.startswith("• ") or t.endswith(".") for t in texts)


def test_set_aside_documents_are_named_with_their_reason(qt, tmp_path):
    from PySide6.QtWidgets import QLabel
    from spec_surgeon.planner import FilePlan
    w = window(qt, tmp_path)
    w.add_documents([str(document(tmp_path / "Spec.docx"))])
    w.task.perform("send", text=REQUEST)
    w.session.preview.batch_plan.files.append(FilePlan(tmp_path / "Other.docx", "x", "exception",
        exception={"code": "layout", "reason": "Other.docx is built differently; it was left untouched."}))
    w._render()
    texts = [label.text() for label in w.findChildren(QLabel)]
    assert "Left untouched" in texts
    assert "Other.docx: Other.docx is built differently; it was left untouched." in texts
