"""remove_highlight (built-in batch 2, docs/BUILTIN_OPS_BATCH2_DESIGN.md §3): Word's highlighter comes off — every colour or the
listed ones — in the body, headers, footers and notes, plain or tracked; the text and every other formatting stay; styles
are flagged, never changed — and the verifier (verify/highlight.py) catches a wrong output and a wrong selection shared by
plan and apply."""
from __future__ import annotations

import io
import zipfile
from pathlib import Path

import pytest
from lxml import etree

from fx_comments import DOC, build_highlighted, build_plain
from make_fixtures import W, patch_docx
from spec_surgeon.editspec import apply_spec, load_spec, plan_spec, validate_spec
from spec_surgeon.ops import JobException
from spec_surgeon.ops import highlight as hop
from spec_surgeon.package import DocxPackage
from spec_surgeon.verify.suite import run_suite

EXAMPLES = Path(__file__).resolve().parent.parent / "examples" / "specs"


def spec(*ops, mode="formatting"):
    return {"request": "remove the highlighting", "summary": "s", "mode": mode, "ops": list(ops)}


def unhighlight(select=None, **params):
    o = {"op": "remove_highlight", "params": params}
    if select:
        o["select"] = select
    return o


def run(data: bytes, sp):
    pkg = DocxPackage(data, "s.docx")
    plan, frags = plan_spec(pkg, sp)
    out = apply_spec(pkg, frags).output_bytes()
    return plan, out, run_suite(data, out, plan.to_dict(), "s.docx", "o.docx")


def failed(rec):
    return {c.id: c.summary for c in rec.checks if c.status == "FAIL"}


def member(data: bytes, name: str):
    return etree.fromstring(zipfile.ZipFile(io.BytesIO(data)).read(name))


def header_part(data: bytes) -> str:
    return next(n for n in zipfile.ZipFile(io.BytesIO(data)).namelist() if n.startswith("word/header"))


def highlights(data: bytes, part=DOC):
    """Colour of every w:highlight in the part, rPrChange snapshots and styles left out."""
    return [h.get(W + "val") for h in member(data, part).iter(W + "highlight")
            if not any(a.tag == W + "rPrChange" for a in h.iterancestors())]


def para(root, start: str):
    return next(p for p in root.iter(W + "p") if "".join(t.text or "" for t in p.iter(W + "t", W + "delText")).startswith(start))


# ---- happy paths ---------------------------------------------------------------------------------------------------

def test_every_colour_comes_off_and_everything_else_stays():
    data = build_highlighted()
    plan, out, rec = run(data, spec(unhighlight()))
    assert rec.passed, failed(rec)
    assert sorted(highlights(data)) == ["cyan", "darkYellow", "green", "magenta", "magenta", "none", "yellow", "yellow",
                                        "yellow", "yellow"]
    assert highlights(out) == ["none"]                                    # "none" is no highlight: never touched
    assert highlights(out, header_part(out)) == []
    root = member(out, DOC)
    approved = [r for r in para(root, "A.").iter(W + "r") if "".join(t.text or "" for t in r.iter(W + "t")) == "approved"][0]
    assert approved.find(W + "rPr/" + W + "b") is not None               # bold stays on the highlighted word
    assert para(root, "C.").find(W + "pPr") is None                       # the emptied mark rPr and pPr went
    assert member(out, "word/styles.xml") is not None and zipfile.ZipFile(io.BytesIO(out)).read("word/styles.xml") == \
        zipfile.ZipFile(io.BytesIO(data)).read("word/styles.xml")         # a style is never changed …
    assert any("style 'Marker' highlights" in f for f in plan.flags)      # … it is flagged
    a = plan.property_assertions[0]
    assert (a["expected_runs"], a["expected_marks"]) == (9, 1)


def test_only_the_listed_colours_come_off():
    data = build_highlighted()
    plan, out, rec = run(data, spec(unhighlight(colors=["yellow", "darkYellow"])))
    assert rec.passed, failed(rec)
    assert sorted(highlights(out)) == ["cyan", "green", "magenta", "magenta", "none"]
    assert plan.property_assertions[0]["colors"] == ["yellow", "darkYellow"]
    # a spec names Word's colours exactly (the schema's list); the op itself reads any case
    assert hop.canonical_colors(["YELLOW", "darkyellow", "Yellow"]) == ["yellow", "darkYellow"]


def test_deleted_text_and_text_boxes_lose_their_highlight_too():
    data = build_highlighted()
    plan, out, rec = run(data, spec(unhighlight(colors=["yellow", "magenta"])))
    assert rec.passed, failed(rec)
    root = member(out, DOC)
    deleted = [r for r in root.iter(W + "r") if r.getparent().tag == W + "del"]
    assert deleted and all(r.find(W + "rPr/" + W + "highlight") is None for r in deleted)
    assert all(h.get(W + "val") != "magenta" for h in root.iter(W + "highlight"))


def test_tracked_removal_records_a_formatting_change_per_run():
    data = build_highlighted()
    plan, out, rec = run(data, spec(unhighlight(colors=["green", "cyan"], tracked=True, author="Abe")))
    assert rec.passed, failed(rec)
    assert plan.rev_delta == {"ins": 0, "del": 0, "rPrChange": 2}
    root = member(out, DOC)
    changes = [c for c in root.iter(W + "rPrChange") if c.get(W + "author") == "Abe"]
    assert len(changes) == 2
    assert all(c.find(W + "rPr/" + W + "highlight") is not None for c in changes)   # reject-all restores it
    assert rec.get("V4").status == "PASS" and rec.get("V5b").status == "PASS"


def test_tracked_removal_refuses_a_run_that_already_carries_a_tracked_formatting_change():
    with pytest.raises(JobException) as e:
        plan_spec(DocxPackage(build_highlighted(), "s.docx"), spec(unhighlight(tracked=True)))
    assert e.value.code == "tracked-format"
    # plain removal leaves the earlier reviewer's change as it was
    plan, out, rec = run(build_highlighted(), spec(unhighlight()))
    assert rec.passed
    assert [c.get(W + "author") for c in member(out, DOC).iter(W + "rPrChange")] == ["Jeremy"]


def test_tracked_removal_refuses_a_paragraph_mark_that_is_itself_a_pending_insertion():
    """The rPrChange snapshot would copy the mark's own w:ins: a second revision with the same id (V4 caught it as an
    unpredicted ins_mark and a duplicate id) — the op refuses that mark instead; plain removal is fine."""
    from docx import Document
    from docx.enum.text import WD_COLOR_INDEX
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from make_fixtures import docx_bytes
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("inserted and highlighted").font.highlight_color = WD_COLOR_INDEX.YELLOW
    mark = OxmlElement("w:rPr")
    ins = OxmlElement("w:ins")
    for k, v in (("w:id", "5"), ("w:author", "Jeremy"), ("w:date", "2026-01-01T00:00:00Z")):
        ins.set(qn(k), v)
    hl = OxmlElement("w:highlight"); hl.set(qn("w:val"), "yellow")
    mark.append(ins); mark.append(hl)
    p._p.get_or_add_pPr().append(mark)
    data = docx_bytes(doc)
    with pytest.raises(JobException) as e:
        plan_spec(DocxPackage(data, "s.docx"), spec(unhighlight(tracked=True)))
    assert e.value.code == "tracked-format" and "pending tracked change" in e.value.reason
    plan, out, rec = run(data, spec(unhighlight()))
    assert rec.passed, failed(rec)
    assert highlights(out) == []


def test_the_scope_limits_where_it_comes_off():
    data = build_highlighted()
    plan, out, rec = run(data, spec(unhighlight({"scope": ["headers"]})))
    assert rec.passed, failed(rec)
    assert highlights(out, header_part(out)) == []
    assert sorted(highlights(out)) == sorted(highlights(data))
    assert plan.intended_parts == [header_part(data)]


def test_nothing_to_remove_is_a_no_op_unless_asked_to_insist():
    data = build_plain()
    plan, frags = plan_spec(DocxPackage(data, "s.docx"), spec(unhighlight()))
    assert plan.intended_parts == [] and any("nothing to do" in n for n in plan.notes)
    assert apply_spec(DocxPackage(data, "s.docx"), frags).output_bytes() == data
    with pytest.raises(JobException) as e:
        plan_spec(DocxPackage(data, "s.docx"), spec(unhighlight(allow_missing=False)))
    assert e.value.code == "not-found"
    with pytest.raises(JobException) as e:
        plan_spec(DocxPackage(build_highlighted(), "s.docx"), spec(unhighlight(colors=["red"], allow_missing=False)))
    assert e.value.code == "not-found" and "red" in e.value.reason


@pytest.mark.parametrize("colors", [["pink"], [], ["none"], "yellow"])
def test_a_colour_word_does_not_know_is_refused(colors):
    with pytest.raises(JobException) as e:
        hop.canonical_colors(colors)
    assert e.value.code == "bad-params"


# ---- wrong outputs ------------------------------------------------------------------------------------------------

def _good(**params):
    data = build_highlighted()
    plan, out, rec = run(data, spec(unhighlight(**params)))
    assert rec.passed
    return data, plan, out


def test_another_property_taken_off_with_the_highlight_fails():
    data, plan, out = _good()

    def unbold(root):
        for b in list(para(root, "A.").iter(W + "b")):
            b.getparent().remove(b)
    rec = run_suite(data, patch_docx(out, {DOC: unbold}), plan.to_dict(), "s.docx", "o.docx")
    assert "V10" in failed(rec) and "other than the highlight" in failed(rec)["V10"]


def test_a_highlight_left_on_fails():
    data, plan, out = _good()

    def relight(root):
        r = [r for r in para(root, "B.").iter(W + "r")][1]
        rpr = r.find(W + "rPr")
        if rpr is None:
            rpr = etree.Element(W + "rPr"); r.insert(0, rpr)
        etree.SubElement(rpr, W + "highlight").set(W + "val", "cyan")
    rec = run_suite(data, patch_docx(out, {DOC: relight}), plan.to_dict(), "s.docx", "o.docx")
    assert "V10" in failed(rec) and "still highlighted" in failed(rec)["V10"]


def test_an_explicit_no_highlight_taken_off_fails():
    data, plan, out = _good()

    def drop_none(root):
        for h in list(root.iter(W + "highlight")):
            if h.get(W + "val") == "none":
                h.getparent().remove(h)
    rec = run_suite(data, patch_docx(out, {DOC: drop_none}), plan.to_dict(), "s.docx", "o.docx")
    assert "V10" in failed(rec)


def test_a_tracked_change_without_its_snapshot_fails():
    data, plan, out = _good(colors=["green", "cyan"], tracked=True, author="Abe")

    def empty_snapshot(root):
        for c in root.iter(W + "rPrChange"):
            if c.get(W + "author") == "Abe":
                for x in list(c.find(W + "rPr")):
                    c.find(W + "rPr").remove(x)
    rec = run_suite(data, patch_docx(out, {DOC: empty_snapshot}), plan.to_dict(), "s.docx", "o.docx")
    assert "V10" in failed(rec) and "snapshot" in failed(rec)["V10"]


# ---- wrong implementations shared by plan and apply -----------------------------------------------------------------

def test_a_selection_that_counts_none_as_a_highlight_is_caught(monkeypatch):
    real = hop.highlighted
    monkeypatch.setattr(hop, "highlighted", lambda rpr, colors: real(rpr, list(colors) + ["none"]))
    plan, out, rec = run(build_highlighted(), spec(unhighlight()))
    assert "V10" in failed(rec)


def test_a_colour_reading_that_takes_dark_yellow_for_yellow_is_caught(monkeypatch):
    real = hop.canonical_colors
    monkeypatch.setattr(hop, "canonical_colors",
                        lambda raw: real(raw) + (["darkYellow"] if raw and "yellow" in raw else []))
    plan, out, rec = run(build_highlighted(), spec(unhighlight(colors=["yellow"])))
    assert "darkYellow" in plan.property_assertions[0]["colors"]          # plan and apply both read it so …
    assert "V10" in failed(rec)                                           # … the verifier reads the user's own words
    assert "the asked colours are ['yellow'] here" in failed(rec)["V10"]


def test_a_default_scope_that_forgets_the_headers_is_caught(monkeypatch):
    monkeypatch.setattr(hop, "DEFAULT_SCOPE", ["document"])
    plan, out, rec = run(build_highlighted(), spec(unhighlight()))
    assert highlights(out, header_part(out)) == ["yellow"]                # plan and apply left the header alone …
    assert "V10" in failed(rec)                                           # … the verifier's own default has headers


def test_a_strip_that_also_drops_other_properties_is_caught(monkeypatch):
    real = hop.strip_highlight

    def strip_more(rpr, stamp, rev_id):
        for b in rpr.findall(W + "b"):
            rpr.remove(b)
        return real(rpr, stamp, rev_id)
    monkeypatch.setattr(hop, "strip_highlight", strip_more)
    plan, out, rec = run(build_highlighted(), spec(unhighlight()))
    assert "V10" in failed(rec)


# ---- rule 9 and the rest of the app -----------------------------------------------------------------------------------

def test_formatting_the_same_paragraphs_in_one_spec_is_a_plan_time_conflict():
    sp = spec(unhighlight(), {"op": "set_properties", "select": {"runs": "all"},
                              "params": {"run": {"b": True}, "tracked": False}})
    with pytest.raises(JobException) as e:
        plan_spec(DocxPackage(build_highlighted(), "s.docx"), sp)
    assert e.value.code == "conflict"


def test_the_worked_example_validates_and_plans():
    sp = load_spec(EXAMPLES / "remove-highlighting.json")
    validate_spec(sp)
    plan, out, rec = run(build_highlighted(), sp)
    assert rec.passed and sorted(highlights(out)) == ["cyan", "darkYellow", "green", "magenta", "magenta", "none"]
