"""Mixed input discovery must preserve project boundaries, evidence and broken files."""
import json

from make_fixtures import build
from spec_surgeon.inventory import inventory_inputs


def write_doc(folder, name="21 13 13.docx"):
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / name
    path.write_bytes(build("21 13 13", "SPRINKLER SYSTEMS"))
    return path


def test_nearest_metadata_scopes_same_structure_and_never_reads_above_root(tmp_path):
    (tmp_path / "PROJECT.json").write_text('{"stamp":"WRONG OUTSIDE ROOT"}')
    root = tmp_path / "jobs"
    a, b = root / "a", root / "b"
    write_doc(a / "mechanical")
    write_doc(b / "mechanical")
    (a / "PROJECT.json").write_text(json.dumps({"name": "North", "stamp": "NORTH"}))
    (b / "PROJECT.json").write_text(json.dumps({"name": "South", "stamp": "SOUTH"}))
    result = inventory_inputs([str(root)])
    assert len(result["projects"]) == 2
    assert {p["context"]["stamp"] for p in result["projects"].values()} == {"NORTH", "SOUTH"}
    assert len({d["pattern_id"] for d in result["documents"].values()}) == 1
    assert all(d["section"] == "21 13 13" and d["discipline"] == "Fire suppression"
               for d in result["documents"].values())
    narrow = inventory_inputs([str(a / "mechanical")])
    project = next(iter(narrow["projects"].values()))
    assert project["context"] == {} and project["provisional"]


def test_bad_documents_locks_previous_results_and_overlapping_inputs(tmp_path):
    root = tmp_path / "input"
    good = write_doc(root)
    (root / "bad.docx").write_bytes(b"not a zip")
    (root / "~$21 13 13.docx").write_bytes(b"Word lock")
    write_doc(root / "old-run")
    (root / "old-run" / "recovery.json").write_text("{}")
    write_doc(root / "new-workspace")
    result = inventory_inputs([str(root), str(good)], exclude=[root / "new-workspace"])
    assert len(result["documents"]) == 2
    assert {d["status"] for d in result["documents"].values()} == {"PENDING", "UNREADABLE"}
    assert all(d["source_sha256"] for d in result["documents"].values())
    assert result == inventory_inputs([str(root), str(good)], exclude=[root / "new-workspace"])


def test_no_metadata_folders_are_separate_and_conflicting_section_evidence_is_unresolved(tmp_path):
    write_doc(tmp_path / "a", "22 11 16.docx")
    write_doc(tmp_path / "b")
    result = inventory_inputs([str(tmp_path)])
    assert len(result["projects"]) == 2
    conflict = next(d for d in result["documents"].values() if "22 11 16" in d["source"])
    assert conflict["section"] is None and conflict["unresolved"]
    assert all(p["provisional"] for p in result["projects"].values())


def test_nonrecursive_intake_and_explicit_result_file_are_excluded(tmp_path):
    source = write_doc(tmp_path / "root")
    write_doc(source.parent / "nested")
    assert len(inventory_inputs([str(source.parent)], recursive=False)["documents"]) == 1
    result = write_doc(tmp_path / "result")
    (result.parent / "campaign.json").write_text("{}")
    assert not inventory_inputs([str(result)])["documents"]


def test_full_document_classification_covers_late_architectural_section_and_quality(tmp_path):
    from docx import Document
    document = Document()
    for index in range(100):
        document.add_paragraph(f"Front matter {index}")
    document.add_paragraph("SECTION 08 71 00")
    document.add_paragraph("DOOR HARDWARE")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Hardware schedule"
    source = tmp_path / "spec.docx"
    document.save(source)
    record = next(iter(inventory_inputs([str(source)])["documents"].values()))
    assert record["section"] == "08 71 00"
    assert record["discipline"] == "Openings"
    assert record["section_status"] == "identified"
    assert "tables" in record["quality_flags"]
    assert "missing_styles_part" not in record["quality_flags"]
    assert record["status"] == "PENDING"


def test_unidentified_section_is_explicit_and_empty_revision_nodes_are_detected(tmp_path):
    import io
    import zipfile
    from docx import Document
    document = Document()
    document.add_paragraph("Unclassified document")
    source = tmp_path / "misc.docx"
    document.save(source)
    raw = source.read_bytes()
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(raw)) as original, zipfile.ZipFile(output, "w") as result:
        for item in original.infolist():
            data = original.read(item.filename)
            if item.filename == "word/document.xml":
                data = data.replace(b"</w:p>", b"<w:ins w:id=\"1\"/></w:p>", 1)
            result.writestr(item, data)
    source.write_bytes(output.getvalue())
    record = next(iter(inventory_inputs([str(source)])["documents"].values()))
    assert record["section_status"] == "not_found"
    assert "pending_revisions" in record["quality_flags"]
