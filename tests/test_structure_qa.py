"""Semantic QA is advisory, scoped to provable hierarchy and explicit citations."""
import hashlib

from docx import Document

from spec_surgeon.spec_qa import audit_batch, audit_markdown


def source(tmp_path, *lines, table=None):
    doc = Document()
    for line in ("SECTION 21 13 13", "SPRINKLER SYSTEMS", "PART 1 GENERAL", "1.1 SUMMARY", *lines):
        doc.add_paragraph(line)
    if table:
        doc.add_table(rows=1, cols=1).cell(0, 0).text = table
    path = tmp_path / "input.docx"
    doc.save(path)
    return path


def test_duplicates_are_siblings_and_numeric_tokens_must_agree(tmp_path):
    path = source(tmp_path, "A. Provide 4-inch pipe with the specified finish.",
                  "B. Provide 4-inch pipe with the specified finish.",
                  "C. Provide 6-inch pipe with the specified finish.",
                  "D. Provide 4-inch pipe with specified finish.",
                  "1.2 MATERIALS", "A. Provide 4-inch pipe with the specified finish.")
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    report = audit_batch([path], audits=["duplicate_provisions"])
    audit = report["files"][0]["duplicate_provisions"]
    assert [(f["paragraph"], f["other"]["paragraph"], f["status"]) for f in audit["findings"]] == [
        (5, 4, "duplicate"), (7, 4, "near_duplicate")]
    assert audit["status"] == "evaluated" and audit["unassessed"] == []
    assert "compare Section 21 13 13; Article 1.1" in audit_markdown(report)
    assert hashlib.sha256(path.read_bytes()).hexdigest() == digest


def test_internal_references_match_full_and_relative_and_flag_missing(tmp_path):
    path = source(tmp_path, "A. See Paragraph 1.1.B.1 and Paragraph B.1.",
                  "B. Scope includes the following.", "1. Provide documentation.",
                  "C. See Article 1.2 and Paragraph 1.1.B.2.", "1.2 SUBMITTALS", "A. Submit product data.")
    report = audit_batch([path], audits=["internal_references"])
    audit = report["files"][0]["internal_references"]
    assert [f["status"] for f in audit["findings"]] == ["matched", "matched", "matched", "missing_target"]
    assert audit["findings"][1]["reference_id"] == "1.1.B.1"
    assert audit["findings"][0]["targets"][0]["paragraph"] == 6
    assert audit["status"] == "evaluated"
    assert "missing_target" in audit_markdown(report)


def test_unknown_hierarchy_cannot_prove_missing_target(tmp_path):
    path = source(tmp_path, "A. See Paragraph 1.1.D.", "I. Ambiguous marker", "Unmarked explanation")
    audit = audit_batch([path], audits=["internal_references"])["files"][0]["internal_references"]
    assert audit["findings"][0]["status"] == "unassessed"
    assert audit["status"] == "partial" and len(audit["unassessed"]) == 2


def test_excluded_tables_are_visible_coverage_and_cannot_supply_targets(tmp_path):
    path = source(tmp_path, "A. See Paragraph 1.1.B.", table="B. Only in a table. See Article 9.9.")
    report = audit_batch([path], audits=["internal_references", "duplicate_provisions"])
    audit = report["files"][0]["internal_references"]
    assert [f["status"] for f in audit["findings"]] == ["unassessed", "unassessed"]
    assert audit["coverage"]["excluded"] == 1
    assert "Unassessed" in audit_markdown(report) and "table paragraph" in audit_markdown(report)


def test_duplicate_article_labels_produce_ambiguous_targets(tmp_path):
    path = source(tmp_path, "A. See Article 1.1.", "1.1 REPEATED", "A. Repeated provision.")
    audit = audit_batch([path], audits=["internal_references"])["files"][0]["internal_references"]
    assert audit["findings"][0]["status"] == "ambiguous_target"
    assert len(audit["findings"][0]["targets"]) == 2


def test_standard_and_other_section_qualifiers_do_not_become_broken_internal_refs(tmp_path):
    path = source(tmp_path, "A. See Paragraph 9.2 of NFPA 13; Article 2.3 of Section 21 05 00.",
                  "B. Per NFPA 13, Paragraph 9.3 governs installation.",
                  "C. Comply with NFPA 13. See Article 1.1.",
                  "D. See Article 1.1; NFPA 13 also applies.")
    audit = audit_batch([path], audits=["internal_references"])["files"][0]["internal_references"]
    assert [f["status"] for f in audit["findings"]] == ["external_reference", "external_reference",
                                                      "external_reference", "matched", "matched"]


def test_article_definitions_are_not_reported_as_self_references(tmp_path):
    path = source(tmp_path, "ARTICLE 1.2 PRODUCTS", "A. Provide listed products.")
    audit = audit_batch([path], audits=["internal_references"])["files"][0]["internal_references"]
    assert audit["findings"] == []


def test_requested_register_comparison_retains_source_and_basis(tmp_path):
    path = source(tmp_path, "A. NFPA 13, 2022 Edition; UL 199.")
    context = {"project_evidence": {"version": 1, "facts": [], "standards": [
        {"id": "nfpa13", "identifier": "NFPA 13", "edition": "2019", "basis": "Owner-approved contract register",
         "scope": {"kind": "project"}, "source": {"kind": "user", "reference": "Approved register"},
         "confirmed_by": "Reviewer"}]}}
    report = audit_batch([path], audits=["standards"], project_context=context)
    comparison = report["files"][0]["standards_register"]
    assert [x["status"] for x in comparison] == ["mismatch", "unassessed"]
    assert comparison[0]["location"]["paragraph"] == 4
    assert comparison[0]["register_entries"][0]["basis"] == "Owner-approved contract register"
    assert "Approved standards register" in audit_markdown(report)


def test_register_comparison_honors_both_document_and_section_scope(tmp_path):
    path = source(tmp_path, "A. NFPA 13, 2022 Edition; UL 199, 2023 Edition.")
    context = {"project_evidence": {"version": 1, "facts": [], "standards": [
        {"id": "nfpa13", "identifier": "NFPA 13", "edition": "2022", "basis": "Selected section edition",
         "scope": {"kind": "section", "value": "21 13 13"}, "source": {"kind": "user", "reference": "Register"},
         "confirmed_by": "Reviewer"},
        {"id": "ul199", "identifier": "UL 199", "edition": "2023", "basis": "Selected document edition",
         "scope": {"kind": "document", "value": str(path)}, "source": {"kind": "user", "reference": "Register"},
         "confirmed_by": "Reviewer"}]}}
    report = audit_batch([path], audits=["standards"], project_context=context)
    assert [x["status"] for x in report["files"][0]["standards_register"]] == ["match", "match"]


def test_filename_only_section_identity_does_not_activate_section_register(tmp_path):
    doc = Document()
    doc.add_paragraph("NFPA 13, 2022 Edition.")
    path = tmp_path / "21 13 13.docx"
    doc.save(path)
    context = {"project_evidence": {"version": 1, "facts": [], "standards": [
        {"id": "nfpa13", "identifier": "NFPA 13", "edition": "2022", "basis": "Selected section edition",
         "scope": {"kind": "section", "value": "21 13 13"}, "source": {"kind": "user", "reference": "Register"},
         "confirmed_by": "Reviewer"}]}}
    report = audit_batch([path], audits=["standards"], project_context=context)
    assert report["files"][0]["standards_register"][0]["status"] == "unassessed"
