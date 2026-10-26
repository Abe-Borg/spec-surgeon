"""What a custom method may format: more kinds of formatting, never more kinds of change.

A1 — the limits say what the checker proves. ``CUSTOM_TRANSFORM_LIMITS`` (2026-10-17) listed highlight under "cannot",
though V10's ``properties`` kind has verified highlight, superscript and subscript, hidden text, character spacing and
double strikethrough since set_properties grew them (2026-10-21). A custom method picks its targets by its own logic
and names them (an ``index_range`` per paragraph); the kind checks those got exactly the asked property and nothing
else changed. Each scripted method here runs through the real trial: the whitelist, the sandbox, ``hard_limits`` and
the suite."""
from __future__ import annotations

import io
from pathlib import Path

import pytest
from docx import Document

from lxml import etree

from make_fixtures import build, build_numbering
from spec_surgeon import codegen
from spec_surgeon.catalog import CUSTOM_TRANSFORM_LIMITS, op_catalog
from spec_surgeon.codegen import check_source, hard_limits, trial
from spec_surgeon.ops import get_op
from spec_surgeon.package import W, DocxPackage
from spec_surgeon.verify import props as vprops

HIGHLIGHT_OP = '''
from __future__ import annotations
from lxml import etree
from spec_surgeon.ops.base import Op, JobException
from spec_surgeon.package import W
from spec_surgeon.ooxml import RPR_ORDER, insert_in_order, paragraphs, block_full_text, run_rpr, mark_rpr
from spec_surgeon.ops.selectors import DEFAULT_PRESERVE_FONTS


class HighlightInsurerParagraphs(Op):
    name = "highlight_insurer_paragraphs"
    mode = "formatting"
    description = "Highlight the paragraphs that name the insurer, but not one holding a field or a tracked change."
    params_schema = {"type": "object", "properties": {"color": {"type": "string"}, "words": {"type": "string"}}}
    select_schema = {}

    def plan(self, pkg, select, params):
        color, words = params.get("color", "yellow"), params.get("words", "FM Global")
        main = pkg.main_part()
        frag = self.fragment(select, params)
        chosen = []
        for bi, p in enumerate(paragraphs(pkg.xml(main))):
            if words not in block_full_text(p):
                continue
            if p.find(".//" + W + "fldChar") is not None or p.find(".//" + W + "ins") is not None \\
                    or p.find(".//" + W + "del") is not None:
                continue
            runs = [r for r in p.iter(W + "r") if any(c.tag != W + "rPr" for c in r)]
            frag.property_assertions.append({
                "kind": "properties", "target": "direct", "scope": ["document"], "parts": [main],
                "select": {"runs": "all", "paragraphs": {"index_range": [bi, bi + 1]}},
                "set": {"run": {"highlight": color}, "paragraph": {}}, "preserve_fonts": list(DEFAULT_PRESERVE_FONTS),
                "skip_sym_runs": True, "in_revisions": "include", "include_paragraph_marks": True, "tracked": False,
                "author": None, "expected_runs": len(runs), "expected_marks": 1, "expected_paragraphs": 0,
                "expected_shown_marks": 0})
            frag.touches.append((main, str(bi), "rPr"))
            chosen.append(bi)
        if not chosen:
            raise JobException("no paragraph names the insurer", "not-found")
        frag.intended_parts = [main]
        frag.transform_data = {"main": main, "paragraphs": chosen, "color": color}
        frag.geometry_reference = "source"
        return frag

    def apply(self, pkg, frag):
        td = frag.transform_data
        root = pkg.xml(td["main"])
        pars = paragraphs(root)
        for bi in td["paragraphs"]:
            p = pars[bi]
            targets = [run_rpr(r) for r in p.iter(W + "r") if any(c.tag != W + "rPr" for c in r)]
            targets.append(mark_rpr(p, create=True))
            for rpr in targets:
                h = rpr.find(W + "highlight")
                if h is None:
                    h = insert_in_order(rpr, etree.Element(W + "highlight"), RPR_ORDER)
                h.set(W + "val", td["color"])
        pkg.write_xml(td["main"], root)
'''

APPLY_LOOP = 'for bi in td["paragraphs"]:'
SET_VALUE = 'h.set(W + "val", td["color"])'
PLANTED = {
    # the colour written is not the one asserted
    "a wrong colour": HIGHLIGHT_OP.replace(SET_VALUE, 'h.set(W + "val", "green")'),
    # a paragraph the plan never named highlighted too
    "an unlisted paragraph": HIGHLIGHT_OP.replace(APPLY_LOOP, 'for bi in td["paragraphs"] + [td["paragraphs"][-1] + 1]:'),
    # a word changed beside the formatting
    "a changed word": HIGHLIGHT_OP.replace(
        SET_VALUE, SET_VALUE + '\n            for t in p.iter(W + "t"):\n'
                               '                t.text = (t.text or "").replace("FM Global", "FM Approvals")'),
}


@pytest.fixture
def spec(tmp_path) -> Path:
    f = tmp_path / "21 05 00 - COMMON WORK RESULTS.docx"
    f.write_bytes(build("21 05 00", "COMMON WORK RESULTS"))
    return f


def highlighted(docx: bytes):
    out = []
    for p in Document(io.BytesIO(docx)).paragraphs:
        colours = {r.font.highlight_color for r in p.runs if r.text}
        if colours and None not in colours:
            out.append(p.text[:2])
    return out


def test_a_custom_method_highlights_the_paragraphs_it_picks(spec, tmp_path):
    assert check_source(HIGHLIGHT_OP) == []
    before = spec.read_bytes()
    t = trial(HIGHLIGHT_OP, spec, {}, {"color": "yellow"}, tmp_path / "sb")
    assert t.passed, t.problems
    assert "properties" in t.record.get("V10").summary
    # A. and B. name the insurer; C. holds Jeremy's insertion and E. a field: the method's own logic left them
    assert highlighted(Path(t.output).read_bytes()) == ["A.", "B."]
    assert spec.read_bytes() == before


@pytest.mark.parametrize("case", sorted(PLANTED))
def test_a_wrong_custom_method_fails_the_trial(case, spec, tmp_path):
    t = trial(PLANTED[case], spec, {}, {"color": "yellow"}, tmp_path / "sb")
    assert not t.passed, case
    assert any(p.startswith(("V10 FAIL", "V2 FAIL")) for p in t.problems), t.problems


# what the `properties` family's words name, per property the kind takes (set_properties' own parameters)
WORDS = {"b": "bold", "i": "italic", "caps": "caps", "smallCaps": "small caps", "strike": "strikethrough",
         "dstrike": "double strikethrough", "u": "underline", "color": "text colour", "vertAlign": "superscript and "
         "subscript", "vanish": "hidden text", "spacing": "character spacing", "highlight": "highlight",
         "keepNext": "keep-with-next", "keepLines": "keep lines together", "widowControl": "widow control",
         "pageBreakBefore": "page break before"}
SAMPLES = {"u": "single", "color": "FF0000", "vertAlign": "superscript", "highlight": "yellow", "spacing": 1}


def test_the_properties_family_names_every_property_the_kind_verifies():
    family = next(words for words, kind in CUSTOM_TRANSFORM_LIMITS["families"].items() if kind == "properties")
    schema = get_op("set_properties").params_schema["properties"]
    run_keys, para_keys = list(schema["run"]["properties"]), list(schema["paragraph"]["properties"])
    assert set(WORDS) == set(run_keys) | set(para_keys), "a property was added: name it in the family and here"
    for k in run_keys:
        assert vprops._v_run_elements({k: SAMPLES.get(k, True)}), k          # the kind verifies it
        assert WORDS[k] in family, k
    for k in para_keys:
        assert vprops._v_para_elements({k: True}), k
        assert WORDS[k] in family, k
    assert set(CUSTOM_TRANSFORM_LIMITS["families"].values()) <= set(vprops.KINDS)


def test_highlight_is_no_longer_out_of_reach_and_the_writer_is_told():
    assert not any("highlight" in c for c in CUSTOM_TRANSFORM_LIMITS["cannot"])
    assert "highlight of selected text" in codegen.system_prompt()
    assert "highlight of selected text" in str(op_catalog(examples=False)["custom_transform_limits"])


# ---- A2 — table appearance: table_props and table_cells, the targets named by position -----------------------------

SHADE_OP = '''
from __future__ import annotations
from lxml import etree
from spec_surgeon.ops.base import Op, JobException
from spec_surgeon.package import W
from spec_surgeon.ooxml import TCPR_ORDER, insert_in_order, block_full_text


def own(el, tag):
    return [c for c in el if c.tag == W + tag]


def first_cell(tbl):
    rows = own(tbl, "tr")
    cells = own(rows[0], "tc") if rows else []
    return " ".join(block_full_text(p) for p in own(cells[0], "p")).strip() if cells else None


class ShadeEveryOtherRow(Op):
    name = "shade_every_other_row"
    mode = "formatting"
    description = "Shade every other row of the revision tables (the tables whose first cell reads Rev)."
    params_schema = {"type": "object", "properties": {"fill": {"type": "string"}, "first_cell": {"type": "string"}}}
    select_schema = {}

    def plan(self, pkg, select, params):
        fill, first = params.get("fill", "D9D9D9").upper(), params.get("first_cell", "Rev")
        main = pkg.main_part()
        frag = self.fragment(select, params)
        targets = []
        want = {W + "val": "clear", W + "color": "auto", W + "fill": fill}
        for k, tbl in enumerate(pkg.xml(main).iter(W + "tbl")):
            if first_cell(tbl) != first:
                continue
            rows = own(tbl, "tr")
            for r in range(1, len(rows), 2):
                cells = []
                for c, tc in enumerate(own(rows[r], "tc")):
                    pr = tc.find(W + "tcPr")
                    shd = pr.find(W + "shd") if pr is not None else None
                    if shd is None or dict(shd.attrib) != want:
                        cells.append(c)
                frag.property_assertions.append({
                    "kind": "table_cells", "parts": [main],
                    "select": {"scope": ["document"], "tables": {"index_range": [k, k]},
                               "rows": {"index_range": [r, r]}},
                    "params": {"shading": fill}, "expected_tables": 1 if cells else 0,
                    "expected_cells": len(cells), "expected_columns": 0,
                    "changed": {main: {"cells": [[k, r, c, ["shd"]] for c in cells]}} if cells else {}})
                targets += [[k, r, c] for c in cells]
                frag.touches.append((main, "tr:" + str(k) + ":" + str(r), "tcPr"))
        if not frag.property_assertions:
            raise JobException("no revision table in this file", "not-found")
        frag.intended_parts = [main] if targets else []
        frag.transform_data = {"main": main, "targets": targets, "fill": fill}
        frag.geometry_reference = "source"
        return frag

    def apply(self, pkg, frag):
        td = frag.transform_data
        if not td["targets"]:
            return
        root = pkg.xml(td["main"])
        tables = list(root.iter(W + "tbl"))
        for k, r, c in td["targets"]:
            tc = own(own(tables[k], "tr")[r], "tc")[c]
            pr = tc.find(W + "tcPr")
            if pr is None:
                pr = etree.Element(W + "tcPr")
                tc.insert(0, pr)
            old = pr.find(W + "shd")
            if old is not None:
                pr.remove(old)
            insert_in_order(pr, etree.Element(W + "shd", {W + "val": "clear", W + "color": "auto",
                                                          W + "fill": td["fill"]}), TCPR_ORDER)
        pkg.write_xml(td["main"], root)
'''

HEADER_OP = '''
from __future__ import annotations
from lxml import etree
from spec_surgeon.ops.base import Op, JobException
from spec_surgeon.package import W
from spec_surgeon.ooxml import TRPR_ORDER, insert_in_order, block_full_text


def own(el, tag):
    return [c for c in el if c.tag == W + tag]


def plain_on(pr, name):
    els = pr.findall(W + name) if pr is not None else []
    return len(els) == 1 and not els[0].attrib


class RepeatRevisionHeader(Op):
    name = "repeat_revision_header"
    mode = "formatting"
    description = "Repeat the first row of the revision tables on every page and keep their rows whole."
    params_schema = {"type": "object", "properties": {}}
    select_schema = {}

    def plan(self, pkg, select, params):
        main = pkg.main_part()
        frag = self.fragment(select, params)
        rows_changed = []
        for k, tbl in enumerate(pkg.xml(main).iter(W + "tbl")):
            rows = own(tbl, "tr")
            cells = own(rows[0], "tc") if rows else []
            if not cells or " ".join(block_full_text(p) for p in own(cells[0], "p")).strip() != "Rev":
                continue
            changed = []
            for r, tr in enumerate(rows):
                pr = tr.find(W + "trPr")
                names = [n for n in (["tblHeader"] if r == 0 else []) + ["cantSplit"] if not plain_on(pr, n)]
                if r > 0 and pr is not None and pr.find(W + "tblHeader") is not None:
                    raise JobException("a body row is marked as a header row", "structure")
                if names:
                    changed.append([k, r, sorted(names)])
            frag.property_assertions.append({
                "kind": "table_props", "parts": [main],
                "select": {"scope": ["document"], "tables": {"index_range": [k, k]}},
                "params": {"repeat_header_rows": 1, "cant_split": True},
                "expected_tables": 1 if changed else 0, "expected_rows": len(changed), "expected_cells": 0,
                "changed": {main: {"rows": changed}} if changed else {}})
            rows_changed += changed
            frag.touches.append((main, "tbl:" + str(k), "trPr"))
        if not frag.property_assertions:
            raise JobException("no revision table in this file", "not-found")
        frag.intended_parts = [main] if rows_changed else []
        frag.transform_data = {"main": main, "rows": rows_changed}
        frag.geometry_reference = "none"
        frag.geometry_reason = "rows kept whole may move to the next page"
        return frag

    def apply(self, pkg, frag):
        td = frag.transform_data
        if not td["rows"]:
            return
        root = pkg.xml(td["main"])
        tables = list(root.iter(W + "tbl"))
        for k, r, names in td["rows"]:
            tr = own(tables[k], "tr")[r]
            pr = tr.find(W + "trPr")
            if pr is None:
                pr = etree.Element(W + "trPr")
                ex = tr.find(W + "tblPrEx")
                tr.insert(tr.index(ex) + 1 if ex is not None else 0, pr)
            for n in names:
                for old in pr.findall(W + n):
                    pr.remove(old)
                insert_in_order(pr, etree.Element(W + n), TRPR_ORDER)
        pkg.write_xml(td["main"], root)
'''

SHADE_LOOP = 'for k, r, c in td["targets"]:'
SHADE_PLANTED = {
    # a row the plan never named shaded too
    "an unlisted row": SHADE_OP.replace(SHADE_LOOP, 'for k, r, c in td["targets"] + [[td["targets"][0][0], 2, 0]]:'),
    # a cell's words changed beside its shading
    "a changed cell text": SHADE_OP.replace(
        "        pkg.write_xml(td[\"main\"], root)",
        "        for t in tables[td[\"targets\"][0][0]].iter(W + \"t\"):\n"
        "            t.text = (t.text or \"\").replace(\"Issued\", \"Withdrawn\")\n"
        "        pkg.write_xml(td[\"main\"], root)"),
    # the merged cell of row 3 split back into two cells
    "a merged cell split": SHADE_OP.replace(
        "        pkg.write_xml(td[\"main\"], root)",
        "        for tc in own(own(tables[td[\"targets\"][0][0]], \"tr\")[3], \"tc\"):\n"
        "            span = tc.find(W + \"tcPr/\" + W + \"gridSpan\")\n"
        "            if span is not None:\n"
        "                span.getparent().remove(span)\n"
        "                tc.addnext(etree.fromstring(etree.tostring(tc)))\n"
        "                break\n"
        "        pkg.write_xml(td[\"main\"], root)"),
    # a grid column added to the table
    "a grid column added": SHADE_OP.replace(
        "        pkg.write_xml(td[\"main\"], root)",
        "        grid = tables[td[\"targets\"][0][0]].find(W + \"tblGrid\")\n"
        "        grid.append(etree.Element(W + \"gridCol\", {W + \"w\": \"720\"}))\n"
        "        pkg.write_xml(td[\"main\"], root)"),
}


@pytest.fixture
def schedule(tmp_path) -> Path:
    """The revision table with five rows (Rev, then 0–3), row 3's last two cells merged."""
    doc = Document(io.BytesIO(build("21 05 00", "COMMON WORK RESULTS")))
    tbl = doc.tables[0]
    for n, (date, what) in enumerate((("2026-02-01", "Addendum 1"), ("2026-03-01", "Addendum 2"),
                                      ("2026-04-01", "Issued for Permit")), start=1):
        cells = tbl.add_row().cells
        cells[0].text, cells[1].text, cells[2].text = str(n), date, what
    merged = tbl.rows[3].cells[1].merge(tbl.rows[3].cells[2])
    merged.text = "2026-03-01 Addendum 2"
    f = tmp_path / "21 05 00 - COMMON WORK RESULTS.docx"
    doc.save(f)
    return f


def shaded_rows(docx: bytes):
    tbl = Document(io.BytesIO(docx)).tables[0]
    out = []
    for r, row in enumerate(tbl.rows):
        fills = {tc._tc.tcPr.find(W_ + "shd").get(W_ + "fill") if tc._tc.tcPr is not None
                 and tc._tc.tcPr.find(W_ + "shd") is not None else None for tc in row.cells}
        if fills == {"D9D9D9"}:
            out.append(r)
    return out


W_ = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def test_a_custom_method_shades_every_other_row_it_picks(schedule, tmp_path):
    assert check_source(SHADE_OP) == []
    t = trial(SHADE_OP, schedule, {}, {"fill": "d9d9d9"}, tmp_path / "sb")
    assert t.passed, t.problems
    assert "table_cells" in t.record.get("V10").summary
    assert shaded_rows(Path(t.output).read_bytes()) == [1, 3]      # the merged cell of row 3 shaded with its row


def test_a_custom_method_repeats_the_header_row_and_keeps_rows_whole(schedule, tmp_path):
    assert check_source(HEADER_OP) == []
    t = trial(HEADER_OP, schedule, {}, {}, tmp_path / "sb")
    assert t.passed, t.problems
    assert "table_props" in t.record.get("V10").summary
    rows = Document(io.BytesIO(Path(t.output).read_bytes())).tables[0].rows
    flags = [(tr._tr.trPr.find(W_ + "tblHeader") is not None, tr._tr.trPr.find(W_ + "cantSplit") is not None)
             for tr in rows]
    assert flags == [(True, True)] + [(False, True)] * 4


@pytest.mark.parametrize("case", sorted(SHADE_PLANTED))
def test_a_custom_method_that_touches_table_content_fails_the_trial(case, schedule, tmp_path):
    t = trial(SHADE_PLANTED[case], schedule, {}, {"fill": "D9D9D9"}, tmp_path / "sb")
    assert not t.passed, case
    assert any(p.startswith(("V10 FAIL", "V2 FAIL", "V3 FAIL")) for p in t.problems), t.problems


def test_table_appearance_is_a_custom_method_family_and_structure_is_not():
    families = CUSTOM_TRANSFORM_LIMITS["families"]
    assert {"table_props", "table_cells"} <= set(families.values())
    assert not any(c == "tables" for c in CUSTOM_TRANSFORM_LIMITS["cannot"])
    assert any("table structure" in c for c in CUSTOM_TRANSFORM_LIMITS["cannot"])
    assert {"TBLPR_ORDER", "TRPR_ORDER", "TCPR_ORDER"} <= codegen.ALLOWED_FROM["spec_surgeon.ooxml"]


def test_a_semantic_selector_is_re_derived_by_the_kind_not_trusted(schedule, tmp_path):
    """Naming targets by position is guidance, not the boundary (Codex on #87): the kinds re-derive any selector they
    take with their own code, so a method selecting by text is checked the same way, and one whose selector does not
    pick what it changed fails."""
    by_text = HEADER_OP.replace('"tables": {"index_range": [k, k]}', '"tables": {"first_row_regex": "^Rev"}')
    t = trial(by_text, schedule, {}, {}, tmp_path / "sb1")
    assert t.passed, t.problems
    # rows named by a text the shaded row does not hold: the kind picks row 4 ("Issued for Permit"), the method shaded 1
    wrong = SHADE_OP.replace('"rows": {"index_range": [r, r]}', '"rows": {"text_regex": "Permit"}')
    t = trial(wrong, schedule, {}, {"fill": "D9D9D9"}, tmp_path / "sb2")
    assert not t.passed and any(p.startswith("V10 FAIL") for p in t.problems), t.problems


# ---- A3 — list label formatting: numbering_labels, the levels named -------------------------------------------------

LABEL_OP = '''
from __future__ import annotations
from lxml import etree
from spec_surgeon.ops.base import Op, JobException
from spec_surgeon.package import W

NUMBERING = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/numbering"


class OutlineLabelsInArial(Op):
    name = "outline_labels_in_arial"
    mode = "formatting"
    description = "Show the outline list's numbers (the list whose first level reads PART) in Arial at 10 points."
    params_schema = {"type": "object", "properties": {"font": {"type": "string"}, "size": {"type": "number"}}}
    select_schema = {}

    def plan(self, pkg, select, params):
        font, size = params.get("font", "Arial"), params.get("size", 10)
        main = pkg.main_part()
        part = next((r.resolved for r in pkg.rels(main).values() if r.type == NUMBERING), None)
        if part is None:
            raise JobException("the document has no lists", "not-found")
        frag = self.fragment(select, params)
        levels = []
        for an in pkg.xml(part).findall(W + "abstractNum"):
            lvls = an.findall(W + "lvl")
            first = next((l for l in lvls if l.get(W + "ilvl") == "0"), None)
            text = first.find(W + "lvlText") if first is not None else None
            if text is None or not (text.get(W + "val") or "").startswith("PART"):
                continue
            for l in lvls:
                if l.find(W + "lvlPicBulletId") is not None:
                    raise JobException("a picture bullet", "structure")
                levels.append({"abstractNum": int(an.get(W + "abstractNumId")), "ilvl": int(l.get(W + "ilvl")),
                               "set": {"fonts": {"ascii": font, "hAnsi": font}, "size": size}})
        if not levels:
            raise JobException("no outline list", "not-found")
        frag.property_assertions.append({"kind": "numbering_labels", "numbering_part": part, "levels": levels,
                                         "expected_levels": len(levels)})
        frag.intended_parts = [part]
        frag.touches.append((part, "*", "lvl"))
        frag.transform_data = {"part": part, "levels": levels, "font": font, "half": str(int(size * 2))}
        frag.geometry_reference = "none"
        frag.geometry_reason = "a label in another font and size can move its text"
        return frag

    def apply(self, pkg, frag):
        td = frag.transform_data
        root = pkg.xml(td["part"])
        for lv in td["levels"]:
            an = next(a for a in root.findall(W + "abstractNum") if a.get(W + "abstractNumId") == str(lv["abstractNum"]))
            lvl = next(l for l in an.findall(W + "lvl") if l.get(W + "ilvl") == str(lv["ilvl"]))
            rpr = lvl.find(W + "rPr")
            if rpr is None:
                rpr = etree.SubElement(lvl, W + "rPr")
            rf = rpr.find(W + "rFonts")
            if rf is None:
                rf = etree.SubElement(rpr, W + "rFonts")
            for slot, theme in (("ascii", "asciiTheme"), ("hAnsi", "hAnsiTheme")):
                rf.set(W + slot, td["font"])
                rf.attrib.pop(W + theme, None)
            for n in ("sz", "szCs"):
                for old in rpr.findall(W + n):
                    rpr.remove(old)
                etree.SubElement(rpr, W + n).set(W + "val", td["half"])
        pkg.write_xml(td["part"], root)
'''

LABEL_WRITE = '        pkg.write_xml(td["part"], root)'


def _also(lines: str) -> str:
    return LABEL_OP.replace(LABEL_WRITE, lines + "\n" + LABEL_WRITE)


OUTLINE = '        an = next(a for a in root.findall(W + "abstractNum") if a.get(W + "abstractNumId") == "10")\n'
LABEL_PLANTED = {
    "a changed level text": _also(OUTLINE + '        an.findall(W + "lvl")[1].find(W + "lvlText").set(W + "val", "%1-%2")'),
    "a changed start": _also(OUTLINE + '        an.findall(W + "lvl")[2].find(W + "start").set(W + "val", "2")'),
    "a changed indent": _also(OUTLINE + '        an.findall(W + "lvl")[2].find(W + "pPr/" + W + "ind").set(W + "left", "1800")'),
    "an unlisted level": _also(
        '        bullet = next(a for a in root.findall(W + "abstractNum") if a.get(W + "abstractNumId") == "21")\n'
        '        etree.SubElement(bullet.find(W + "lvl"), W + "rPr").append(etree.Element(W + "b"))'),
    # the theme font left beside the name: Word still shows the theme font
    "a theme font kept": LABEL_OP.replace("                rf.attrib.pop(W + theme, None)\n", ""),
}


def themed(docx: bytes) -> bytes:
    """The outline list's first level with its own label rPr naming a theme font (Word shows that one)."""
    pkg = DocxPackage(docx)
    root = pkg.xml("word/numbering.xml")
    an = next(a for a in root.findall(W + "abstractNum") if a.get(W + "abstractNumId") == "10")
    rpr = etree.SubElement(an.findall(W + "lvl")[0], W + "rPr")
    etree.SubElement(rpr, W + "rFonts", {W + "asciiTheme": "minorHAnsi", W + "hAnsiTheme": "minorHAnsi"})
    pkg.write_xml("word/numbering.xml", root)
    return pkg.output_bytes()


@pytest.fixture
def outline(tmp_path) -> Path:
    f = tmp_path / "21 13 13 - WET-PIPE SPRINKLER SYSTEMS.docx"
    f.write_bytes(themed(build_numbering()))
    return f


def label_props(docx: bytes, abstract: str = "10"):
    root = DocxPackage(docx).xml("word/numbering.xml")
    an = next(a for a in root.findall(W + "abstractNum") if a.get(W + "abstractNumId") == abstract)
    out = []
    for lvl in an.findall(W + "lvl"):
        rf, sz = lvl.find(f"{W}rPr/{W}rFonts"), lvl.find(f"{W}rPr/{W}sz")
        out.append((rf.get(W + "ascii") if rf is not None else None, rf.get(W + "asciiTheme") if rf is not None else None,
                    sz.get(W + "val") if sz is not None else None))
    return out


def test_a_custom_method_formats_the_list_labels_it_picks(outline, tmp_path):
    assert check_source(LABEL_OP) == []
    before = outline.read_bytes()
    t = trial(LABEL_OP, outline, {}, {"font": "Arial", "size": 10}, tmp_path / "sb")
    assert t.passed, t.problems
    assert "numbering_labels" in t.record.get("V10").summary
    out = Path(t.output).read_bytes()
    assert label_props(out) == [("Arial", None, "20")] * 5            # the theme font cleared on level 0
    assert label_props(out, "21") == label_props(before, "21")       # the bullet list untouched
    assert outline.read_bytes() == before


@pytest.mark.parametrize("case", sorted(LABEL_PLANTED))
def test_a_custom_method_that_changes_the_numbers_fails_the_trial(case, outline, tmp_path):
    t = trial(LABEL_PLANTED[case], outline, {}, {"font": "Arial", "size": 10}, tmp_path / "sb")
    assert not t.passed, case
    assert any(p.startswith("V10 FAIL") for p in t.problems), t.problems


def test_the_numbering_part_may_change_only_behind_a_numbering_labels_assertion(outline):
    from spec_surgeon.codegen import allowed_parts_of
    allowed = allowed_parts_of(outline)
    assert "word/numbering.xml" not in allowed
    plan = {"diffs": [], "rev_delta": {}, "intended_parts": ["word/numbering.xml"], "property_assertions": []}
    assert any("numbering.xml" in p for p in hard_limits(plan, allowed))
    plan["property_assertions"] = [{"kind": "numbering_labels", "numbering_part": "word/numbering.xml"}]
    assert hard_limits(plan, allowed) == []
    plan["intended_parts"].append("docProps/core.xml")                 # the assertion opens no other part
    assert any("docProps/core.xml" in p for p in hard_limits(plan, allowed))


def _labels_plan(levels, part="word/numbering.xml"):
    return {"property_assertions": [{"kind": "numbering_labels", "numbering_part": part, "levels": levels,
                                     "expected_levels": len(levels)}]}


@pytest.mark.parametrize("levels, words", [
    ([{"abstractNum": 10, "ilvl": 0, "set": {"size": 10}}] * 2, "named twice"),
    ([{"abstractNum": 99, "ilvl": 0, "set": {"size": 10}}], "not uniquely defined"),
    ([{"abstractNum": 10, "ilvl": 0, "set": {"lvlText": "x"}}], "bad set"),
    ([{"abstractNum": 10, "ilvl": 0, "set": {"size": 10.3}}], "bad size"),
    ([{"abstractNum": 10, "ilvl": 0, "num": 10, "set": {"size": 10}}], "bad level"),
    ([{"abstractNum": 21, "ilvl": 0, "set": {"fonts": {"ascii": "Arial"}}}], "symbol font"),
])
def test_the_kind_says_what_the_method_must_refuse(levels, words, outline):
    from spec_surgeon.verify import run_suite
    src = outline.read_bytes()
    pkg = DocxPackage(src)                                    # the bullet in a symbol font, for the last case
    root = pkg.xml("word/numbering.xml")
    bullet = next(a for a in root.findall(W + "abstractNum") if a.get(W + "abstractNumId") == "21")
    etree.SubElement(etree.SubElement(bullet.find(W + "lvl"), W + "rPr"), W + "rFonts", {W + "ascii": "Symbol"})
    pkg.write_xml("word/numbering.xml", root)
    src = pkg.output_bytes()
    plan = {"diffs": [], "rev_delta": {}, "intended_parts": [], "members_removed": [], "members_added": [],
            **_labels_plan(levels)}
    rec = run_suite(src, src, plan, "src.docx", "out.docx")
    v10 = next(c for c in rec.checks if c.id == "V10")
    assert v10.status == "FAIL" and "must refuse" in v10.summary + " ".join(v10.data.get("errors", [])), v10.summary
    assert words in v10.summary + " ".join(v10.data.get("errors", []))


def test_list_labels_are_a_custom_method_family_and_list_structure_is_not():
    assert "numbering_labels" in CUSTOM_TRANSFORM_LIMITS["families"].values()
    assert any("list structure" in c for c in CUSTOM_TRANSFORM_LIMITS["cannot"])
    assert any("numbering part" in r for r in CUSTOM_TRANSFORM_LIMITS["rules"])
