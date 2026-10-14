"""package.py working overlay: remove_member, add_member, edit_rels, unreferenced_media, repack."""
from __future__ import annotations

import io
import zipfile

import pytest
from docx import Document
from docx.shared import Pt

from make_fixtures import build, png
from spec_surgeon.package import DocxPackage, HEADER_REL, CT_PART


def _doc_with_media() -> bytes:
    d = Document()
    d.add_paragraph("body")
    d.add_picture(io.BytesIO(png((0, 0, 200))))                       # shared: body + header
    hdr = d.sections[0].header
    hdr.is_linked_to_previous = False
    hdr.paragraphs[0].add_run().add_picture(io.BytesIO(png((0, 0, 200))), width=Pt(10))
    hdr.paragraphs[0].add_run().add_picture(io.BytesIO(png((200, 0, 0))), width=Pt(10))  # header only
    b = io.BytesIO(); d.save(b)
    return b.getvalue()


def test_overlay_reads_writes_and_repacks_by_substitution():
    src = DocxPackage(build("21 05 00", "A"))
    work = src.copy()
    assert not work.is_modified() and work.output_bytes() == src.data
    work.write("word/settings.xml", work.read("word/settings.xml"))
    assert work.substitutions() == {"word/settings.xml": src.read("word/settings.xml")}
    assert src.substitutions() == {}                        # copies do not share the overlay
    out = DocxPackage(work.output_bytes())
    assert out.member_names == src.member_names
    with pytest.raises(KeyError):
        work.write("word/new.xml", b"<x/>")                 # new members go through add_member


def test_remove_and_add_members_with_content_types():
    src = DocxPackage(build("21 05 00", "A"))
    work = src.copy()
    hdr = src.sections()[0].header_refs["default"]
    rid = [r.rid for r in src.rels().values() if r.resolved == hdr][0]
    work.edit_rels("word/document.xml", remove=[rid])
    work.remove_member(hdr)
    assert not work.has(hdr) and hdr not in work.content_types()[1]
    work.add_member("word/extra.xml", b"<x/>", "application/vnd.example+xml")
    new = work.edit_rels("word/document.xml", add=[{"type": "urn:example", "target": "extra.xml"}])
    assert new and new[0].startswith("rId")
    out = DocxPackage(work.output_bytes())
    # order: source order minus removed, then added members in add order
    assert out.member_names == [n for n in src.member_names if n != hdr] + ["word/extra.xml"]
    assert out.content_type_of("word/extra.xml") == "application/vnd.example+xml"
    assert hdr not in [r.resolved for r in out.rels().values()]
    assert set(src.member_names) - set(out.member_names) == {hdr}
    # unchanged members are byte-identical
    zs, zo = zipfile.ZipFile(io.BytesIO(src.data)), zipfile.ZipFile(io.BytesIO(out.data))
    for n in src.member_names:
        if n not in (hdr, CT_PART, "word/_rels/document.xml.rels"):
            assert zs.read(n) == zo.read(n), n
    with pytest.raises(ValueError):
        work.add_member("word/document.xml", b"", None)
    with pytest.raises(KeyError):
        work.edit_rels("word/document.xml", remove=["rId999"])


def test_unreferenced_media_reports_only_media_orphaned_by_this_edit():
    src = DocxPackage(_doc_with_media())
    hdr = src.sections()[0].header_refs["default"]
    hdr_media = sorted(r.resolved for r in src.rels(hdr).values())
    body_media = sorted(r.resolved for r in src.rels().values() if r.type.endswith("/image"))
    assert len(hdr_media) == 2 and len(body_media) == 1 and body_media[0] in hdr_media
    work = src.copy()
    assert work.unreferenced_media() == []
    rid = [r.rid for r in src.rels().values() if r.resolved == hdr][0]
    work.edit_rels("word/document.xml", remove=[rid])
    work.remove_member(hdr)
    work.remove_member(DocxPackage.rels_path_for(hdr))
    orphaned = work.unreferenced_media()
    assert orphaned == [m for m in hdr_media if m not in body_media]   # the shared image stays referenced
