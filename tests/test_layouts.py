"""One intent, one EditSpec per layout, one verification per file (Workstream L)."""
from __future__ import annotations

import copy
from pathlib import Path

import pytest
from docx import Document

from make_fixtures import build, build_rich, build_styles
from spec_surgeon import layouts
from spec_surgeon.editspec import SpecError
from spec_surgeon.fingerprint import fingerprint
from spec_surgeon.package import DocxPackage


def _write(path: Path, data: bytes) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(data)
    return path


def pattern(path: Path) -> str:
    return fingerprint(DocxPackage.from_path(path))["pattern_id"]


@pytest.fixture
def batch(tmp_path):
    typed = _write(tmp_path / "typed" / "21 13 13 - WET-PIPE.docx", build("21 13 13", "WET-PIPE SPRINKLER SYSTEMS"))
    styled = _write(tmp_path / "styled" / "21 13 16 - DRY-PIPE.docx", build_styles("21 13 16", "DRY-PIPE SPRINKLER SYSTEMS"))
    other = _write(tmp_path / "other" / "21 13 19 - PREACTION.docx", build_rich("21 13 19", "PREACTION SPRINKLER SYSTEMS"))
    return typed, styled, other


def plan_for(typed: Path, styled: Path) -> dict:
    bold = {"run": {"b": True}, "tracked": False}
    return {"kind": "layouts", "request": "Make the first article title bold.", "summary": "Bold article titles",
            "layouts": [
                {"id": "A", "label": "typed numbers, no styles", "pattern_ids": [pattern(typed)],
                 "spec": {"summary": "bold by text", "mode": "formatting",
                          "ops": [{"op": "set_properties", "select": {"paragraphs": {"text_regex": r"^1\.01\t"}},
                                   "params": bold}]}},
                {"id": "B", "label": "styled articles", "pattern_ids": [pattern(styled)],
                 "spec": {"summary": "bold by style", "mode": "formatting",
                          "ops": [{"op": "set_properties", "select": {"paragraphs": {"style_in": ["ArchPara"]}},
                                   "params": bold}]}}]}


def test_plan_validation_refuses_overlaps_and_missing_specs(batch):
    typed, styled, _ = batch
    plan = plan_for(typed, styled)
    layouts.validate_layout_plan(plan)
    overlap = copy.deepcopy(plan)
    overlap["layouts"][1]["pattern_ids"] = overlap["layouts"][0]["pattern_ids"]
    with pytest.raises(SpecError, match="only one plan"):
        layouts.validate_layout_plan(overlap)
    missing = copy.deepcopy(plan)
    del missing["layouts"][0]["spec"]
    with pytest.raises(SpecError, match="needs a spec"):
        layouts.validate_layout_plan(missing)
    with pytest.raises(SpecError):
        layouts.validate_layout_plan({"kind": "layouts", "request": "x", "layouts": []})
    one = {"kind": "layouts", "request": plan["request"], "layouts": [plan["layouts"][0]]}
    assert layouts.single_spec(one)["ops"] == plan["layouts"][0]["spec"]["ops"]
    assert layouts.single_spec(one)["request"] == plan["request"]
    assert layouts.single_spec(plan) is None


def test_routing_sends_each_file_to_its_layout_and_names_the_rest(batch, tmp_path):
    typed, styled, other = batch
    broken = _write(tmp_path / "broken.docx", b"not a package")
    routing = layouts.route([typed, styled, other, broken], plan_for(typed, styled))
    assert routing.files == {"A": [typed], "B": [styled]}
    assert routing.unassigned == [(other, pattern(other))]
    assert [p for p, _ in routing.unreadable] == [broken]
    assert routing.layout_of(styled) == "B" and routing.layout_of(other) is None
    reason = layouts.unassigned_exception(other, pattern(other), plan_for(typed, styled))["reason"]
    assert other.name in reason and "typed numbers, no styles" in reason and "left untouched" in reason


def test_preview_covers_every_layout_and_sets_aside_the_rest(batch, tmp_path):
    typed, styled, other = batch
    plan = plan_for(typed, styled)
    pv = layouts.make_layout_preview(plan, [str(typed), str(styled), str(other)], tmp_path / "preview", do_render=False)
    assert [layout["id"] for layout, _ in pv.previews] == ["A", "B"]
    assert len(pv.samples) == 2 and all(s.passed for s in pv.samples)
    assert pv.layout_for_sample(pv.samples[1])["id"] == "B"
    statuses = {Path(f.source).name: f.status for f in pv.batch_plan.files}
    assert statuses == {typed.name: "planned", styled.name: "planned", other.name: "exception"}
    aside = next(f for f in pv.batch_plan.exceptions)
    assert aside.exception["code"] == layouts.LAYOUT_CODE and "left untouched" in aside.exception["reason"]
    assert (tmp_path / "preview" / "layouts.json").is_file()
    assert "Layout A" in pv.to_text() and "SET ASIDE" in pv.to_text()
    # the token binds the plan, the files and the routing
    again = layouts.make_layout_preview(plan, [str(typed), str(styled), str(other)], tmp_path / "preview2", do_render=False)
    assert again.token == pv.token
    assert layouts.layout_token(plan, [typed, styled, other]) == pv.token
    changed = copy.deepcopy(plan)
    changed["layouts"][0]["spec"]["ops"][0]["params"]["run"] = {"b": True, "i": True}
    assert layouts.layout_token(changed, [typed, styled, other]) != pv.token
    assert layouts.layout_token(plan, [typed, styled]) != pv.token


def test_a_layout_with_no_files_warns_instead_of_failing(batch, tmp_path):
    typed, styled, _ = batch
    plan = plan_for(typed, styled)
    pv = layouts.make_layout_preview(plan, [str(typed)], tmp_path / "preview", do_render=False)
    assert [layout["id"] for layout, _ in pv.previews] == ["A"]
    assert any("styled articles" in w for w in pv.warnings)


def test_run_delivers_each_layout_and_refuses_moved_inputs(batch, tmp_path):
    typed, styled, other = batch
    plan = plan_for(typed, styled)
    before = {p: p.read_bytes() for p in (typed, styled, other)}
    pv = layouts.make_layout_preview(plan, [str(typed), str(styled), str(other)], tmp_path / "preview", do_render=False)
    result = layouts.run_layouts([str(typed), str(styled), str(other)], plan, tmp_path / "run", do_render=False,
                                 approved_inputs=pv.token_inputs)
    assert result.counts() == {"PASS": 2, "FAIL": 0, "EXCEPTION": 1, "PENDING": 0}
    by_name = {Path(f.source).name: f for f in result.files}
    assert by_name[other.name].exception["code"] == layouts.LAYOUT_CODE
    assert Path(by_name[typed.name].delivered_to).is_file() and Path(by_name[styled.name].delivered_to).is_file()
    assert Path(by_name[typed.name].delivered_to).parent == tmp_path / "run" / "layout-A" / "outputs"
    assert {p: p.read_bytes() for p in (typed, styled, other)} == before
    assert (tmp_path / "run" / "LAYOUT_REPORT.md").is_file()
    typed_out = Document(by_name[typed.name].delivered_to)
    assert any(p.text.startswith("1.01") and any(r.bold for r in p.runs if r.text.strip()) for p in typed_out.paragraphs)
    styled_out = Document(by_name[styled.name].delivered_to)
    body = [p for p in styled_out.paragraphs if p.style.style_id == "ArchPara"]
    assert body and all(all(r.bold for r in p.runs if r.text.strip()) for p in body)
    assert not any(r.bold for p in styled_out.paragraphs if p.style.style_id == "Normal" for r in p.runs
                   if r.text.strip() and not r.text.startswith(("SECTION", "DRY-PIPE", "PART", "Earlier")))
    # a source that moved since the preview is refused before anything runs
    typed.write_bytes(before[typed] + b" ")
    with pytest.raises(ValueError, match="changed since the preview"):
        layouts.run_layouts([str(typed), str(styled), str(other)], plan, tmp_path / "run2", do_render=False,
                            approved_inputs=pv.token_inputs)
    assert not (tmp_path / "run2" / "layout-A").exists()
