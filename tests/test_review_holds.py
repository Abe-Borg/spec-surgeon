"""The per-preview review (2026-10-21): a green checkmark says a document's preview looks right; a red X sets it
aside. Create edited copies runs the documents that look right, the set-aside ones wait on a to-do list, and the
pass over them begins when the user says so — planned alone, from the request they were set aside from, with the
model told what they missed and what the user saw. Every file that runs is still held to the identity the user
approved; a file outside the approval never runs."""
from __future__ import annotations

import copy
from pathlib import Path

import pytest
from docx import Document

from fake_claude import FakeClient, resp, tool
from spec_surgeon.document_task import START_OVER, WORK_TODO, DocumentTask
from spec_surgeon.gui import Session
from test_document_task import REQUEST, SPEC, document

ITALIC = copy.deepcopy(SPEC)
ITALIC["ops"][0]["params"]["run"] = {"i": True}
ITALIC["summary"] = "Make the title italic"


def task_for(tmp_path, fake, journal=None):
    return DocumentTask(Session(client=fake, out_root=tmp_path / "out", visual_check=False), render=False,
                        journal=journal)


def previewed(tmp_path, fake, journal=None, names=("A.docx", "B.docx")):
    task = task_for(tmp_path, fake, journal=journal)
    sources = [document(tmp_path / name) for name in names]
    task.attach(sources)
    view = task.perform("send", text=REQUEST)
    assert view["stage"] == "preview" and view["can_hold"] is True
    assert {str(s) for s in sources} == {str(sv.source) for sv in view["samples"]}   # every document is previewed
    assert view["copies_to_create"] == len(sources) and view["unpreviewed"] == 0
    return task, sources, view


def test_a_document_set_aside_at_approval_waits_on_the_to_do_list_while_the_others_get_copies(tmp_path):
    fake = FakeClient(resp(tool("submit_editspec", {"spec": SPEC})))
    task, (first, second), view = previewed(tmp_path, fake)
    before = second.read_bytes()
    task.perform("primary", token=view["token"], hold=[str(second)])
    assert task.stage == "done"
    assert list(task.results) == [str(first)]                      # one copy, for the document that looked right
    assert Path(task.results[str(first)]).is_file()
    assert second.read_bytes() == before
    assert list(task.todo) == [str(second)] and task.todo[str(second)]["name"] == "B.docx"
    assert len(task.todo[str(second)]["seen"]) == 1 and "bold" in task.todo[str(second)]["seen"][0]
    assert list(task.applied) == [str(first)] and task.applied[str(first)] == task.history   # only A carries it
    assert task.documents() == [Path(task.results[str(first)])]   # the set-aside document is out of scope meanwhile
    view = task.view()
    assert view["todo"] == [{"path": str(second), "name": "B.docx", "why": "held",
                             "reason": "You marked its preview as needing more work."}] and view["focus"] == []
    assert view["chips"][0] == WORK_TODO and START_OVER in view["chips"]
    assert view["can_hold"] is False
    assert view["copies_to_create"] == 0 and view["unpreviewed"] == 0          # nothing waits to be created
    assert any("to-do list" in m["text"] and "B.docx" in m["text"] for m in task.messages if m["role"] == "app")


def test_the_pass_over_the_set_aside_documents_plans_them_alone_from_the_request_and_the_note(tmp_path):
    fake = FakeClient(resp(tool("submit_editspec", {"spec": SPEC})), resp(tool("submit_editspec", {"spec": ITALIC})),
                      resp(tool("submit_editspec", {"spec": SPEC})))
    task, (first, second), view = previewed(tmp_path, fake)
    task.perform("primary", token=view["token"], hold=[str(second)])
    first_copy = Path(task.results[str(first)])
    view = task.perform("send", text=WORK_TODO)
    assert task.focus == {"documents": [str(second)], "seen": task.history,
                          "why": {str(second): {"why": "held", "reason": "You marked its preview as needing more work."}},
                          "decisions": [REQUEST]}                  # what a Discard puts back on the list with it
    assert task.todo == {} and task.documents() == [second]        # the set-aside document alone, as it is
    assert task.decisions == [REQUEST]                              # the request it was set aside from stands
    assert view["stage"] == "resume" and view["focus"] == [str(second)] and WORK_TODO not in view["chips"]
    assert any("Working on B.docx" in m["text"] for m in task.messages if m["role"] == "app")
    view = task.perform("send", text="Italic, not bold, for this one.")
    sent = fake.requests[1]["messages"][0]["content"]
    assert REQUEST in sent and "Italic, not bold, for this one." in sent
    assert "SET ASIDE AFTER A PREVIEW" in sent and "THE PLAN THEY SAW" in sent
    assert sent.index("THE PLAN THEY SAW") < sent.index("Apply bold in paragraphs")   # the plan the user saw
    assert "CHANGES ALREADY MADE" not in sent                       # it carries none of the other copies' changes
    assert task.session.files() == [second]                         # planned alone
    assert view["stage"] == "preview"
    task.perform("primary", token=view["token"])
    assert task.stage == "done" and task.focus is None and task.todo == {}
    second_copy = Path(task.results[str(second)])
    assert second_copy.is_file() and Document(second_copy).paragraphs[1].runs[0].italic
    assert not Document(second_copy).paragraphs[1].runs[0].bold
    assert Document(first_copy).paragraphs[1].runs[0].bold
    assert task.documents() == [first_copy, second_copy]            # both copies are the working set again
    assert task.generation == 2 and len(task.history) == 2
    assert WORK_TODO not in task.view()["chips"]
    # each copy carries its own changes, and the next request is told per document, never a shared story
    assert [s[:10] for s in task.applied[str(first)]] == ["Apply bold"] and [s[:12] for s in task.applied[str(second)]] == ["Apply italic"]
    task.perform("send", text="Now make the title red.")
    sent = fake.requests[2]["messages"][0]["content"]
    block = sent.split("CHANGES ALREADY MADE")[1].split("\n\n")[0]
    assert "(not in B.docx)" in block and "(not in A.docx)" in block
    assert block.index("Apply bold") < block.index("Apply italic")


def test_setting_every_document_aside_creates_nothing_and_begins_the_pass_at_once(tmp_path):
    fake = FakeClient(resp(tool("submit_editspec", {"spec": SPEC})))
    task, (first, second), view = previewed(tmp_path, fake)
    task.perform("primary", token=view["token"], hold=[str(first), str(second)])
    assert task.results == {} and task.generation == 0
    assert task.focus["documents"] == [str(first), str(second)] and len(task.focus["seen"]) == 1
    assert task.documents() == [first, second] and task.decisions == [REQUEST]
    assert any("every document aside" in m["text"] for m in task.messages if m["role"] == "app")
    assert task.view()["stage"] == "resume"


def test_a_hold_outside_the_plan_is_refused_and_nothing_runs(tmp_path):
    fake = FakeClient(resp(tool("submit_editspec", {"spec": SPEC})))
    task, (first, second), view = previewed(tmp_path, fake)
    stranger = document(tmp_path / "elsewhere" / "C.docx")
    with pytest.raises(ValueError, match="documents or plan changed"):
        task.perform("primary", token=view["token"], hold=[str(stranger)])
    assert task.results == {} and task.todo == {}
    # the session itself refuses a file outside the approval, whatever the caller says; the preview is still there
    assert task.session.preview is not None
    with pytest.raises(ValueError, match="documents or plan changed"):
        task.session.run(False, False, lambda _: None, only=[stranger])
    with pytest.raises(ValueError, match="nothing to create"):
        task.session.run(False, False, lambda _: None, only=[])
    assert task.session.last_batch is None


def test_a_partial_run_keeps_every_approved_pin_and_refuses_a_file_that_changed(tmp_path):
    fake = FakeClient(resp(tool("submit_editspec", {"spec": SPEC})))
    task, (first, second), view = previewed(tmp_path, fake)
    document(first)                                                 # the file that would run changed since the preview
    with pytest.raises(ValueError):
        task.perform("primary", token=view["token"], hold=[str(second)])
    assert task.results == {}


def test_the_journal_keeps_the_to_do_list_and_the_pass(tmp_path):
    fake = FakeClient(resp(tool("submit_editspec", {"spec": SPEC})))
    journal = tmp_path / "conversation.json"
    task, (first, second), view = previewed(tmp_path, fake, journal=journal)
    task.perform("primary", token=view["token"], hold=[str(second)])
    again = task_for(tmp_path, FakeClient(), journal=journal)
    assert again.restore()
    assert list(again.todo) == [str(second)] and again.todo[str(second)]["decisions"] == [REQUEST]
    assert again.documents() == [Path(task.results[str(first)])]
    assert again.view()["chips"][0] == WORK_TODO
    again.perform("send", text=WORK_TODO)
    once_more = task_for(tmp_path, FakeClient(), journal=journal)
    assert once_more.restore()
    assert once_more.focus == {"documents": [str(second)], "seen": task.history,
                               "why": {str(second): {"why": "held", "reason": "You marked its preview as needing more work."}},
                               "decisions": [REQUEST]}
    assert once_more.documents() == [second] and once_more.decisions == [REQUEST] and once_more.stage == "resume"
    assert once_more.applied == {str(first): task.history}
    # a journal from before per-copy histories: every copy carried every change
    import json
    data = json.loads(journal.read_text(encoding="utf-8"))
    del data["applied"]
    journal.write_text(json.dumps(data), encoding="utf-8")
    older = task_for(tmp_path, FakeClient(), journal=journal)
    assert older.restore() and older.applied == {str(first): task.history}
    # a malformed record is refused, never guessed at
    data = json.loads(journal.read_text(encoding="utf-8"))
    data["todo"] = {"x": "not a record"}
    journal.write_text(json.dumps(data), encoding="utf-8")
    with pytest.raises(ValueError):
        task_for(tmp_path, FakeClient(), journal=journal).restore()


def test_start_over_and_remove_drop_the_to_do_list(tmp_path):
    fake = FakeClient(resp(tool("submit_editspec", {"spec": SPEC})), resp(tool("submit_editspec", {"spec": SPEC})))
    task, (first, second), view = previewed(tmp_path, fake)
    task.perform("primary", token=view["token"], hold=[str(second)])
    task.remove(second)
    assert task.todo == {} and task.documents() == [Path(task.results[str(first)])]
    task.attach([second])
    task.perform("send", text=START_OVER)
    assert task.todo == {} and task.focus is None and task.documents() == [first, second]
    view = task.perform("send", text=REQUEST)
    task.perform("primary", token=view["token"], hold=[str(first), str(second)])
    assert task.focus and task.view()["working"]["can_start_over"]
    task.perform("send", text=START_OVER)
    assert task.focus is None and task.decisions == [] and task.view()["stage"] == "empty"


def test_the_to_do_reply_with_nothing_set_aside_is_answered_not_acted_on(tmp_path):
    task = task_for(tmp_path, FakeClient())
    task.attach([document(tmp_path / "A.docx")])
    task.perform("send", text=WORK_TODO)
    assert task.messages[-1]["role"] == "app" and "Nothing is set aside" in task.messages[-1]["text"]
    assert task.focus is None


def test_the_narrowed_approval_inputs_are_a_fresh_approval_of_the_subset(tmp_path):
    """A run on part of an approved batch is held to the approved pins of exactly those files: the narrowed inputs
    equal what approving the subset alone would give (while nothing changed), so every run-time scope check holds."""
    from spec_surgeon.editspec import approval_key
    from spec_surgeon.preview import approval_inputs, narrow_approval_inputs
    first, second, third = (document(tmp_path / name) for name in ("A.docx", "B.docx", "C.docx"))
    key = approval_key(SPEC)
    whole = approval_inputs(key, [first, second, third])
    assert narrow_approval_inputs(whole, [first, third]) == approval_inputs(key, [first, third])
    assert narrow_approval_inputs(whole, [second])["files"] == {str(second.resolve()): whole["files"][str(second.resolve())]}
    assert narrow_approval_inputs(whole, [second])["spec"] == whole["spec"]


def test_documents_set_aside_from_different_requests_are_worked_on_one_request_at_a_time(tmp_path):
    """B set aside from the first request, A from the second (made on the others): the pass takes B alone with the
    first request and the plan B saw; A waits and gets its own pass with the second request and the changes its copy
    does carry."""
    fake = FakeClient(resp(tool("submit_editspec", {"spec": SPEC})), resp(tool("submit_editspec", {"spec": ITALIC})),
                      resp(tool("submit_editspec", {"spec": SPEC})), resp(tool("submit_editspec", {"spec": ITALIC})))
    task, (first, second, third), view = previewed(tmp_path, fake, names=("A.docx", "B.docx", "C.docx"))
    task.perform("primary", token=view["token"], hold=[str(second)])                 # round 1: B waits
    view = task.perform("send", text="Also make it italic.")                          # round 2 on A and C
    assert task.session.files() == [Path(task.results[str(first)]), Path(task.results[str(third)])]
    task.perform("primary", token=view["token"], hold=[str(task.current_of(first))])  # A waits too
    assert [e["name"] for e in task.todo.values()] == ["B.docx", "A.docx"]
    assert task.todo[str(second)]["decisions"] == [REQUEST] and task.todo[str(first)]["decisions"] == ["Also make it italic."]
    task.perform("send", text=WORK_TODO)
    assert task.focus["documents"] == [str(second)] and task.decisions == [REQUEST]    # B first, with its own request
    assert list(task.todo) == [str(first)]                                             # A waits for the next pass
    assert any("A.docx: set aside from a different request" in m["text"] for m in task.messages)
    view = task.perform("send", text="Keep it as it is, just bold.")
    sent = fake.requests[2]["messages"][0]["content"]
    assert "CHANGES ALREADY MADE" not in sent and "Apply bold" in sent.split("THE PLAN THEY SAW")[1][:200]
    assert "italic" not in sent.split("THE PLAN THEY SAW")[1].split("\n\n")[0].lower()
    task.perform("primary", token=view["token"])
    assert task.focus is None and list(task.todo) == [str(first)] and WORK_TODO in task.view()["chips"]
    task.perform("send", text=WORK_TODO)
    assert task.focus["documents"] == [str(first)] and task.decisions == ["Also make it italic."]
    view = task.perform("send", text="Italic it is.")
    sent = fake.requests[3]["messages"][0]["content"]
    block = sent.split("CHANGES ALREADY MADE")[1].split("\n\n")[0]
    assert "Apply bold" in block and "(" not in block.split("\n")[1]                   # A's copy carries the bold change
    assert "Apply italic" in sent.split("THE PLAN THEY SAW")[1][:200]
    task.perform("primary", token=view["token"])
    assert task.todo == {} and task.focus is None
    assert [s[:12] for s in task.applied[str(first)]] == ["Apply bold i", "Apply italic"]
    assert [s[:10] for s in task.applied[str(second)]] == ["Apply bold"]
    assert [s[:12] for s in task.applied[str(third)]] == ["Apply bold i", "Apply italic"]


def test_removing_the_last_document_of_a_pass_ends_it_with_its_request(tmp_path):
    fake = FakeClient(resp(tool("submit_editspec", {"spec": SPEC})))
    task, (first, second), view = previewed(tmp_path, fake)
    task.perform("primary", token=view["token"], hold=[str(second)])
    task.perform("send", text=WORK_TODO)
    assert task.decisions == [REQUEST]
    task.remove(second)
    assert task.focus is None and task.decisions == []                 # the request left with the document
    assert task.documents() == [Path(task.results[str(first)])] and task.view()["stage"] == "empty"
    assert any("that pass is over" in m["text"] for m in task.messages if m["role"] == "app")
