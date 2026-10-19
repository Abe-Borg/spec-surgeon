"""Ordinary delivery depends on the approved recipe and the independent engine suite."""
from __future__ import annotations

import copy
import json
from types import SimpleNamespace

import pytest
from docx import Document

from fake_claude import FakeClient, resp, tool
from spec_surgeon import compiler, ordinary_outcomes
from spec_surgeon.cli import main
from spec_surgeon.gui import Session
from spec_surgeon.batch import run_batch
from spec_surgeon.chain import chain_from_specs, run_chain

REQUEST = "Make the FIRE SUPPRESSION section title bold in every document."
CONTRACT = {"kind": "document_effects", "effects": [{"id": "title", "type": "format",
    "target": {"locators": [{"exact_text": "FIRE SUPPRESSION"}], "cardinality": 1}, "run": {"b": True}}]}


def document(path, *, title_style="Normal", body_style="Normal", duplicate=False):
    doc = Document()
    doc.add_paragraph("SECTION 21 13 13")
    title = doc.add_paragraph("FIRE SUPPRESSION", style=title_style)
    title.runs[0].bold = False
    if duplicate:
        doc.add_paragraph("FIRE SUPPRESSION")
    body = doc.add_paragraph("Owner shall retain all content.", style=body_style)
    body.runs[0].bold = False
    doc.save(path)
    return path


def recipe(predicate):
    return {"request": REQUEST, "summary": "Bold the requested section titles", "mode": "formatting",
        "ops": [{"op": "set_properties", "select": {"paragraphs": predicate},
                 "params": {"run": {"b": True}, "tracked": False}}]}


def test_review_approval_refuses_stale_scope_source_contract_and_preview(tmp_path):
    source = document(tmp_path / "spec.docx")
    review = ordinary_outcomes.build_review(REQUEST, [source], CONTRACT)
    ordinary_outcomes.validate(review, [source], approval=review["digest"])
    changed = copy.deepcopy(review)
    changed["contract"]["effects"][0]["run"]["b"] = False
    with pytest.raises(ValueError, match="review changed"):
        ordinary_outcomes.validate(changed, [source], approval=review["digest"])
    with pytest.raises(ValueError, match="scope changed"):
        ordinary_outcomes.validate(review, [], approval=review["digest"])
    source.write_bytes(source.read_bytes() + b"changed")
    with pytest.raises(ValueError, match="document changed"):
        ordinary_outcomes.validate(review, [source], approval=review["digest"])


def test_ambiguous_targets_need_deliberate_choice_and_new_review_digest(tmp_path):
    source = document(tmp_path / "spec.docx", duplicate=True)
    review = ordinary_outcomes.build_review(REQUEST, [source], CONTRACT)
    assert review["status"] == "NEEDS_INPUT"
    candidates = review["documents"][0]["resolution"]["effects"]["title"]["candidates"]
    with pytest.raises(ValueError, match="not independently covered"):
        ordinary_outcomes.validate(review, [source], approval=review["digest"])
    selected = ordinary_outcomes.choose_targets(review, source, [candidates[1]["candidate_id"]], effect_id="title")
    assert selected["status"] == "PROPOSED" and selected["digest"] != review["digest"]
    ordinary_outcomes.validate(selected, [source], approval=selected["digest"])



@pytest.mark.parametrize("in_place", [False, True])
def test_preview_is_the_only_ordinary_approval_and_calls_no_second_planner(tmp_path, in_place, monkeypatch):
    folder = tmp_path / "sources"
    folder.mkdir()
    source = document(folder / "spec.docx")
    before = source.read_bytes()
    # The preview establishes the desired scope, including when it differs from an old request contract.
    spec = recipe({"text_regex": "^Owner"})
    fake = FakeClient(resp(tool("submit_editspec", {"spec": spec})))
    def unexpected(*args, **kwargs):
        pytest.fail("ordinary preview must not call the outcome planner")
    monkeypatch.setattr(ordinary_outcomes, "propose", unexpected)
    session = Session(folder=folder, out_root=tmp_path / "results", client=fake, visual_check=False)
    session.compile(REQUEST, lambda _: None)
    session.make_preview(False, lambda _: None)
    assert session.outcome_review is None and session.run_blocked_reason() == ""
    assert len(fake.requests) == 1
    session.run(in_place, False, lambda _: None)
    assert session.last_batch.counts() == {"PASS": 1, "FAIL": 0, "EXCEPTION": 0, "PENDING": 0}
    delivered = source if in_place else next((session.last_run_dir / "outputs").glob("*.docx"))
    assert Document(delivered).paragraphs[2].runs[0].bold
    assert session.last_batch.files[0].record.get("GOAL") is None
    if not in_place:
        assert source.read_bytes() == before


@pytest.mark.parametrize("chain,in_place", [(False, False), (False, True), (True, False), (True, True)])
def test_legacy_outcome_metadata_cannot_withhold_verified_delivery(tmp_path, chain, in_place):
    source = document(tmp_path / "spec.docx")
    spec = recipe({"text_regex": "^Owner"})
    out = tmp_path / "run"
    kwargs = dict(workers=1, do_render=False, in_place=in_place,
                  goal_contract=CONTRACT, goal_bindings={})  # even missing old bindings cannot create a gate
    result = (run_chain([str(source)], chain_from_specs([("bold", spec)]), out, **kwargs) if chain else
              run_batch([str(source)], "edit-spec", {}, out, spec=spec, **kwargs))
    assert result.counts() == {"PASS": 1, "FAIL": 0, "EXCEPTION": 0, "PENDING": 0}
    assert result.goal_contract == CONTRACT
    assert result.files[0].delivered_to
    record = result.files[0].results[-1].record if chain else result.files[0].record
    assert record.get("GOAL") is None


def test_cli_yes_runs_without_a_planner_and_legacy_flags_are_advisory(tmp_path, monkeypatch, capsys):
    source = document(tmp_path / "spec.docx")
    spec = recipe({"text_regex": "^Owner"})
    fake = FakeClient(resp(tool("submit_editspec", {"spec": spec})))
    monkeypatch.setattr(compiler, "_client", lambda: fake)
    monkeypatch.setattr("spec_surgeon.llm.credentials_found", lambda: True)
    monkeypatch.setattr("spec_surgeon.cli._interactive", lambda: False)
    result_dir = tmp_path / "ordinary"
    assert main(["do", REQUEST, str(source), "--out", str(result_dir), "--yes", "--no-render",
                 "--workers", "1"]) == 0
    assert len(fake.requests) == 1
    assert not (result_dir / "outcome-review.json").exists()
    assert (result_dir / "run" / "outputs" / source.name).is_file()
    review = ordinary_outcomes.build_review(REQUEST, [source], CONTRACT)
    review_path = tmp_path / "review.json"
    ordinary_outcomes.save(review, review_path)
    assert main(["run-spec", str(result_dir / "preview" / "spec.json"), str(source),
                 "--out", str(tmp_path / "legacy"), "--outcome-review", str(review_path),
                 "--outcome-approval", "obsolete-digest", "--no-render", "--workers", "1"]) == 0
    assert "does not approve or withhold delivery" in capsys.readouterr().out
    assert (tmp_path / "legacy" / "legacy-outcome-review.json").is_file()


def test_ordinary_preview_still_refuses_changed_source(tmp_path):
    folder = tmp_path / "sources"
    folder.mkdir()
    source = document(folder / "spec.docx")
    session = Session(folder=folder, out_root=tmp_path / "results", visual_check=False)
    session._set_outcome(spec=recipe({"text_regex": "^Owner"}), source="compiled")
    session.make_preview(False, lambda _: None)
    doc = Document(source)
    doc.add_paragraph("A new user edit")
    doc.save(source)
    with pytest.raises(ValueError, match="changed since the preview"):
        session.run(False, False, lambda _: None)
    assert session.last_batch is None


def test_resolve_with_no_exceptions_returns_without_model_calls(tmp_path, monkeypatch):
    folder = tmp_path / "sources"
    folder.mkdir()
    document(folder / "spec.docx")
    session = Session(folder=folder, out_root=tmp_path / "results", visual_check=False)
    session._set_outcome(spec=recipe({"text_regex": "^Owner"}), source="compiled")
    session.make_preview(False, lambda _: None)
    def unexpected(*args, **kwargs):
        pytest.fail("nothing to resolve must not invoke recovery")
    monkeypatch.setattr("spec_surgeon.recovery.recover", unexpected)
    monkeypatch.setattr("spec_surgeon.recovery.recover_from_run", unexpected)
    assert "Nothing to resolve: 0" in session.resolve_remaining(lambda _: None)
    session.run(False, False, lambda _: None)
    assert "Nothing to resolve: 0" in session.resolve_remaining(lambda _: None)


def test_legacy_set_aside_contract_does_not_force_ordinary_retry_into_recovery(tmp_path, monkeypatch):
    from spec_surgeon import cli
    source = document(tmp_path / "spec.docx")
    run = tmp_path / "old-run"
    run.mkdir()
    (run / "set_aside.json").write_text(json.dumps({"request": REQUEST, "goal_contract": CONTRACT,
        "goal_bindings": {}, "files": [{"source": str(source), "name": source.name, "code": "verify-fail",
        "category": "Verification", "reason": "old GOAL mismatch"}]}))
    session = Session(out_root=tmp_path / "results")
    session.load_set_aside(run)
    assert session.pending_recovery_run is None and session.selected_files == [source]
    calls = []
    monkeypatch.setattr(cli, "cmd_do", lambda args: calls.append(args) or 0)
    assert main(["retry", str(run), "--yes", "--no-render"]) == 0
    assert calls[0].from_set_aside == str(run)
