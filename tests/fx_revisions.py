"""Fixtures for resolve_tracked_changes (tests/test_revisions.py): documents whose tracked changes cover every kind the op
resolves — and the ones it leaves to Word. ``tracked_column_resize`` (tests/test_ops.py): a table whose column widths
carry a pending change, the file's only w:id — new revision ids must start above it."""
from __future__ import annotations

import copy
from typing import Optional

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

from make_fixtures import docx_bytes, patch_docx

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W = "{%s}" % W_NS
W16DU = "http://schemas.microsoft.com/office/word/2023/wordml/word16du"
M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
JEREMY, JACK = ("Jeremy", "2026-09-22T10:15:00Z"), ("Jack", "2026-09-23T08:00:00Z")
_ids = [500]


def _rid() -> str:
    _ids[0] += 1
    return str(_ids[0])


def stamp(el, who, utc: Optional[str] = None):
    el.set(qn("w:id"), _rid())
    el.set(qn("w:author"), who[0])
    if who[1] is not None:
        el.set(qn("w:date"), who[1])
    if utc:
        el.set("{%s}dateUtc" % W16DU, utc)
    return el


def wrap(run, tag: str, who, utc: Optional[str] = None):
    """Put a python-docx run (or a w:r element) in a new w:ins / w:del / w:moveFrom / w:moveTo by ``who``."""
    r = getattr(run, "_r", run)
    el = stamp(OxmlElement(f"w:{tag}"), who, utc)
    if tag == "del":
        for t in r.findall(qn("w:t")):
            t.tag = qn("w:delText")
            t.set(qn("xml:space"), "preserve")
        for t in r.findall(qn("w:instrText")):
            t.tag = qn("w:delInstrText")
    r.addprevious(el)
    el.append(r)
    return el


def mark(paragraph, tag: str, who):
    """A paragraph-mark revision (w:pPr/w:rPr/w:ins | w:del …)."""
    p = getattr(paragraph, "_p", paragraph)
    ppr = p.get_or_add_pPr() if hasattr(p, "get_or_add_pPr") else p.find(qn("w:pPr"))
    rpr = ppr.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        sect = ppr.find(qn("w:sectPr"))
        (sect.addprevious(rpr) if sect is not None else ppr.append(rpr))
    m = stamp(OxmlElement(f"w:{tag}"), who)
    rpr.insert(0, m)
    return m


def _field(paragraph, instr: str, result: str):
    runs = []
    for kind in ("begin",):
        r = paragraph.add_run()
        fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), kind); r._r.append(fc); runs.append(r)
    r = paragraph.add_run()
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr; r._r.append(it); runs.append(r)
    r = paragraph.add_run()
    fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "separate"); r._r.append(fc); runs.append(r)
    runs.append(paragraph.add_run(result))
    r = paragraph.add_run()
    fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "end"); r._r.append(fc); runs.append(r)
    return runs


def _goback(paragraph):
    bs = OxmlElement("w:bookmarkStart"); bs.set(qn("w:id"), "0"); bs.set(qn("w:name"), "_GoBack")
    be = OxmlElement("w:bookmarkEnd"); be.set(qn("w:id"), "0")
    paragraph._p.append(bs)
    paragraph._p.append(be)


def redline(moves: bool = True, extras: bool = True) -> bytes:
    """A spec redlined by Jeremy (2026-09-22) and Jack (2026-09-23) — every paragraph found by its text:

    * A — Jeremy inserted "hot-dip galvanized " (w16du:dateUtc too), Jack deleted "black steel ";
    * B — Jack deleted "nested " inside Jeremy's insertion;
    * C — a paragraph Jeremy inserted whole (content and mark); D — one Jack deleted whole, Word's _GoBack bookmark in it;
    * E — Jeremy made a run bold (rPrChange), Jack centred the paragraph (pPrChange);
    * F/G — Jeremy moved "Moved sentence. " from F to G (named move ranges), when ``moves``;
    * H — Jack deleted a REF field (its code as w:delInstrText);
    * the header: Jeremy's insertion; the comments: a comment whose text holds Jack's deletion (``extras``).
    """
    _ids[0] = 500
    doc = Document()
    p = doc.add_paragraph("A. Pipe shall be ")
    wrap(p.add_run("hot-dip galvanized "), "ins", JEREMY, utc="2026-09-22T17:15:00Z")
    wrap(p.add_run("black steel "), "del", JACK)
    p.add_run("Schedule 40.")
    p = doc.add_paragraph("B. Jeremy's insertion ")
    ins = wrap(p.add_run("with Jack's "), "ins", JEREMY)
    inner = p.add_run("nested ")
    ins.append(inner._r)
    wrap(inner, "del", JACK)
    tail = p.add_run("deletion")
    ins.append(tail._r)
    p = doc.add_paragraph()
    wrap(p.add_run("C. Inserted paragraph."), "ins", JEREMY)
    mark(p, "ins", JEREMY)
    p = doc.add_paragraph()
    wrap(p.add_run("D. Deleted paragraph."), "del", JACK)
    mark(p, "del", JACK)
    _goback(p)
    p = doc.add_paragraph("E. Formatting: ")
    r = p.add_run("bold now")
    r.bold = True
    chg = stamp(OxmlElement("w:rPrChange"), JEREMY)
    chg.append(OxmlElement("w:rPr"))
    r._r.get_or_add_rPr().append(chg)
    ppr = p._p.get_or_add_pPr()
    jc = OxmlElement("w:jc"); jc.set(qn("w:val"), "center"); ppr.append(jc)
    pchg = stamp(OxmlElement("w:pPrChange"), JACK)
    old = OxmlElement("w:pPr"); ojc = OxmlElement("w:jc"); ojc.set(qn("w:val"), "left"); old.append(ojc)
    pchg.append(old)
    ppr.append(pchg)
    if moves:
        p = doc.add_paragraph("F. Text before. ")
        rs = stamp(OxmlElement("w:moveFromRangeStart"), JEREMY); rs.set(qn("w:name"), "move1")
        p._p.append(rs)
        moved = p.add_run("Moved sentence. ")
        wrap(moved, "moveFrom", JEREMY)
        re_ = OxmlElement("w:moveFromRangeEnd"); re_.set(qn("w:id"), rs.get(qn("w:id"))); p._p.append(re_)
        p.add_run("Text after.")
        p = doc.add_paragraph("G. Other text. ")
        ts = stamp(OxmlElement("w:moveToRangeStart"), JEREMY); ts.set(qn("w:name"), "move1")
        p._p.append(ts)
        wrap(p.add_run("Moved sentence. "), "moveTo", JEREMY)
        te = OxmlElement("w:moveToRangeEnd"); te.set(qn("w:id"), ts.get(qn("w:id"))); p._p.append(te)
        p.add_run("End of G.")
    p = doc.add_paragraph("H. See ")
    fruns = _field(p, " REF _Ref1 \\h ", "21 10 00")
    d = stamp(OxmlElement("w:del"), JACK)
    fruns[0]._r.addprevious(d)
    for fr in fruns:
        for t in fr._r.findall(qn("w:t")):
            t.tag = qn("w:delText")
        for t in fr._r.findall(qn("w:instrText")):
            t.tag = qn("w:delInstrText")
        d.append(fr._r)
    p.add_run(" for details.")
    doc.add_paragraph("Z. Final paragraph.")
    hp = doc.sections[0].header.paragraphs[0]
    hp.add_run("ACME SPEC ")
    wrap(hp.add_run("REV B"), "ins", JEREMY)
    if extras:
        doc.add_comment(doc.paragraphs[0].runs[0], text="Check the finish", author="Jack", initials="J")
    data = docx_bytes(doc)
    if extras:
        def in_comment(root):
            for t in root.iter(W + "t"):
                if t.text == "Check the finish":
                    t.text = "Check "
                    r = t.getparent()
                    r2 = copy.deepcopy(r)
                    r2.find(W + "t").text = "the finish"
                    r.addnext(r2)
                    d = etree.Element(W + "del")
                    d.set(W + "id", "990"); d.set(W + "author", "Jack"); d.set(W + "date", JACK[1])
                    r2.addprevious(d); d.append(r2)
                    r2.find(W + "t").tag = W + "delText"
        data = patch_docx(data, {"word/comments.xml": in_comment})
    return data


def awkward(kind: str) -> bytes:
    """One case the op leaves to Word (or a boundary case), on a small document."""
    _ids[0] = 700
    doc = Document()
    doc.add_paragraph("Intro paragraph.")
    if kind == "join":                                   # a deleted mark with live text on both sides
        p = doc.add_paragraph("First half ")
        wrap(p.add_run("and the end"), "del", JACK)
        mark(p, "del", JACK)
        doc.add_paragraph("second half.")
    elif kind == "last":                                 # the body's last paragraph deleted whole
        p = doc.add_paragraph()
        wrap(p.add_run("Last words."), "del", JACK)
        mark(p, "del", JACK)
    elif kind == "before-table":
        p = doc.add_paragraph()
        wrap(p.add_run("Deleted before a table."), "del", JACK)
        mark(p, "del", JACK)
        doc.add_table(rows=1, cols=1).cell(0, 0).text = "cell"
        doc.add_paragraph("After the table.")
    elif kind == "bookmark":                             # a real bookmark in a paragraph deleted whole
        p = doc.add_paragraph()
        bs = OxmlElement("w:bookmarkStart"); bs.set(qn("w:id"), "3"); bs.set(qn("w:name"), "Spec_Ref")
        p._p.append(bs)
        wrap(p.add_run("Deleted with a bookmark."), "del", JACK)
        be = OxmlElement("w:bookmarkEnd"); be.set(qn("w:id"), "3"); p._p.append(be)
        mark(p, "del", JACK)
        doc.add_paragraph("Next.")
    elif kind == "section":
        p = doc.add_paragraph()
        wrap(p.add_run("Deleted section end."), "del", JACK)
        p._p.get_or_add_pPr().append(copy.deepcopy(doc.element.body.find(qn("w:sectPr"))))
        mark(p, "del", JACK)
        doc.add_paragraph("Body.")
    elif kind == "row":
        t = doc.add_table(rows=2, cols=1)
        t.cell(1, 0).text = "inserted row"
        tr = t.rows[1]._tr
        trpr = tr.get_or_add_trPr()
        trpr.append(stamp(OxmlElement("w:ins"), JEREMY))
        doc.add_paragraph("After.")
    elif kind == "math":
        p = doc.add_paragraph("Equation: ")
        om = etree.SubElement(p._p, "{%s}oMath" % M_NS)
        mr = etree.SubElement(om, "{%s}r" % M_NS)
        rpr = etree.SubElement(mr, W + "rPr")
        rpr.append(stamp(OxmlElement("w:ins"), JEREMY))
        etree.SubElement(mr, "{%s}t" % M_NS).text = "x"
    elif kind == "sectchange":
        sect = doc.element.body.find(qn("w:sectPr"))
        chg = stamp(OxmlElement("w:sectPrChange"), JACK)
        old = OxmlElement("w:sectPr")
        chg.append(old)
        sect.append(chg)
        doc.add_paragraph("Body.")
    elif kind == "partial-move":                         # Jeremy's move whose 'to' end Jack wrote
        p = doc.add_paragraph("From here: ")
        rs = stamp(OxmlElement("w:moveFromRangeStart"), JEREMY); rs.set(qn("w:name"), "move9"); p._p.append(rs)
        wrap(p.add_run("moving words "), "moveFrom", JEREMY)
        re_ = OxmlElement("w:moveFromRangeEnd"); re_.set(qn("w:id"), rs.get(qn("w:id"))); p._p.append(re_)
        p = doc.add_paragraph("To here: ")
        ts = stamp(OxmlElement("w:moveToRangeStart"), JACK); ts.set(qn("w:name"), "move9"); p._p.append(ts)
        wrap(p.add_run("moving words "), "moveTo", JACK)
        te = OxmlElement("w:moveToRangeEnd"); te.set(qn("w:id"), ts.get(qn("w:id"))); p._p.append(te)
    elif kind == "undated":
        p = doc.add_paragraph("Undated: ")
        wrap(p.add_run("no date here "), "ins", ("Jeremy", None))
        wrap(p.add_run("dated "), "ins", JEREMY)
    elif kind == "clean":
        doc.add_paragraph("No tracked changes at all.")
    elif kind == "cell-last":                            # a cell's only paragraph, deleted whole
        t = doc.add_table(rows=1, cols=1)
        cp = t.cell(0, 0).paragraphs[0]
        wrap(cp.add_run("cell text"), "del", JACK)
        mark(cp, "del", JACK)
        doc.add_paragraph("After.")
    elif kind == "chain":                                # three paragraphs deleted whole in a row
        for k in range(3):
            p = doc.add_paragraph()
            wrap(p.add_run(f"Deleted {k}."), "del", JACK)
            mark(p, "del", JACK)
        doc.add_paragraph("Kept.")
    elif kind == "mark-rprchange":                       # a paragraph mark's formatting change, with its own marker
        p = doc.add_paragraph()
        wrap(p.add_run("Inserted and formatted."), "ins", JEREMY)
        m = mark(p, "ins", JEREMY)
        rpr = m.getparent()
        b = OxmlElement("w:b"); rpr.append(b)
        chg = stamp(OxmlElement("w:rPrChange"), JACK)
        chg.append(OxmlElement("w:rPr"))
        rpr.append(chg)
        doc.add_paragraph("Next.")
    elif kind == "comment-in-del":                       # a comment anchored on text Jack deleted
        p = doc.add_paragraph("Keep this. ")
        gone = p.add_run("Commented and deleted.")
        doc.add_comment(gone, text="Why?", author="Jeremy", initials="JR")
        d = wrap(gone, "del", JACK)
        ref = next(r for r in p._p.iter(qn("w:r")) if r.find(qn("w:commentReference")) is not None)
        d.append(ref)                                    # Word deletes the comment's mark with the text
    elif kind in ("bookmark-straddle", "goback-straddle"):   # a bookmark that starts in Jeremy's insertion, ends outside
        p = doc.add_paragraph("Before ")
        r = p.add_run("inserted words ")
        bs = OxmlElement("w:bookmarkStart"); bs.set(qn("w:id"), "4")
        bs.set(qn("w:name"), "_GoBack" if kind == "goback-straddle" else "Spec_Ref")
        r._r.addprevious(bs)
        ins = wrap(r, "ins", JEREMY)
        ins.insert(0, bs)
        be = OxmlElement("w:bookmarkEnd"); be.set(qn("w:id"), "4")
        p.add_run("after.")._r.addnext(be)
    elif kind == "field-part":                           # a field whose start Jack deleted, its end left in place
        p = doc.add_paragraph("See ")
        fruns = _field(p, " REF _Ref2 \\h ", "21 30 00")
        d = stamp(OxmlElement("w:del"), JACK)
        fruns[0]._r.addprevious(d)
        for fr in fruns[:2]:
            for t in fr._r.findall(qn("w:instrText")):
                t.tag = qn("w:delInstrText")
            d.append(fr._r)
    elif kind == "tab-only":                             # Jack deleted a tab and nothing else: structure, no text
        p = doc.add_paragraph("Before the tab")
        r = p.add_run()
        r._r.append(OxmlElement("w:tab"))
        wrap(r, "del", JACK)
        p.add_run("after it.")
    elif kind == "loose-move":                           # move pieces Word left without their ranges
        p = doc.add_paragraph("Loose: ")
        wrap(p.add_run("moved away "), "moveFrom", JEREMY)
        p.add_run("stays ")
        wrap(p.add_run("moved here"), "moveTo", JACK)
    if kind != "last":
        doc.add_paragraph("Closing paragraph.")
    return docx_bytes(doc)


def tracked_column_resize(grid_change_id: int) -> bytes:
    """A table whose column widths carry a pending tracked change (Word writes ``w:tblGridChange`` when columns are
    resized with Track Changes on) and a paragraph with "Owner" twice; that change's id is the file's only w:id."""
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table._tbl.tblGrid.append(etree.fromstring(
        f'<w:tblGridChange xmlns:w="{W_NS}" w:id="{grid_change_id}"><w:tblGrid><w:gridCol w:w="2000"/>'
        f'<w:gridCol w:w="2000"/></w:tblGrid></w:tblGridChange>'))
    doc.add_paragraph("The Owner shall review. The Owner shall approve.")
    return docx_bytes(doc)
