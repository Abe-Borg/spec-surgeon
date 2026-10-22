"""filter_tagged_paragraphs (docs/BUILTIN_OPS_BATCH1_DESIGN.md §4): keep or delete the paragraphs tagged for a system,
the tag taken off the ones that stay — end to end tracked and untracked (V2 pairs a stripped paragraph beside a deleted
one), keep vs delete lists, the refusals, a bookmark running out of a deleted paragraph (it stays where the paragraph
was), and planted defects caught by the verifier."""
from __future__ import annotations

import io
from typing import Any, Dict, List

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

from make_fixtures import _ref_field, build_batch1
from spec_surgeon.ops import JobException, get_op, merge_fragments
from spec_surgeon.ops import filter_tagged as tagged_module
from spec_surgeon.package import DocxPackage, W
from spec_surgeon.verify import run_suite
from test_text_edits import accepted_text, failed, mutate, v10_errors

AUTHOR = "Abraham Borg"
TAG = r"^\[(?P<tag>[A-Z /]+)\]\s*"


@pytest.fixture(scope="module")
def doc() -> bytes:
    return build_batch1()


def run_op(src: bytes, params: Dict[str, Any]):
    pkg = DocxPackage(src)
    frag = get_op("filter_tagged_paragraphs").plan(pkg, {}, {"author": AUTHOR, "tag_regex": TAG, **params})
    plan = merge_fragments([frag])
    work = pkg.copy()
    get_op("filter_tagged_paragraphs").apply(work, frag)
    out = work.output_bytes()
    return plan, out, run_suite(src, out, plan.to_dict(), "src.docx", "out.docx")


def visible(data: bytes) -> List[str]:
    """Paragraph texts with every change accepted (tracked-deleted paragraphs gone)."""
    root = DocxPackage(data).xml("word/document.xml")
    out = []
    for p in root.iter(W + "p"):
        mark = p.find(f"{W}pPr/{W}rPr/{W}del")
        text = accepted_text(p)
        if mark is None or text:
            out.append(text)
    return out


SYSTEMS_AFTER = ["Wet system paragraph.", "Preaction system paragraph.", "Wet/Dry system paragraph."]


@pytest.mark.parametrize("tracked", [True, False])
def test_keep_preaction_and_wet(doc, tracked):
    plan, out, rec = run_op(doc, {"keep": ["PREACTION", "wet"], "tracked": tracked})
    assert rec.passed, rec.to_markdown()
    texts = visible(out)
    assert [t for t in texts if "system paragraph" in t] == SYSTEMS_AFTER
    assert any("dry (kept 1, deleted 1)" in n for n in plan.notes)
    if not tracked:
        assert rec.get("V2").status == "PASS"                  # a stripped paragraph right after a deleted one
        assert len(list(DocxPackage(out).xml("word/document.xml").iter(W + "p"))) == \
            len(list(DocxPackage(doc).xml("word/document.xml").iter(W + "p"))) - 2
    else:
        assert plan.rev_delta.get("del_mark") == 2


def test_delete_list_is_the_mirror_of_keep(doc):
    plan, out, rec = run_op(doc, {"delete": ["DRY", "Clean  Agent"], "tracked": False})
    assert rec.passed
    assert [t for t in visible(out) if "system paragraph" in t] == SYSTEMS_AFTER


def test_tags_stay_when_asked(doc):
    plan, out, rec = run_op(doc, {"keep": ["preaction"], "strip_tags": False, "tracked": False})
    assert rec.passed
    assert [t for t in visible(out) if "system paragraph" in t] == ["[PREACTION] Preaction system paragraph."]


def test_params(doc):
    for params in ({"keep": ["wet"], "delete": ["dry"]}, {}, {"keep": ["wet"], "tag_regex": "("}):
        with pytest.raises(JobException) as e:
            run_op(doc, params)
        assert e.value.code == "bad-params"
    with pytest.raises(JobException) as e:
        run_op(doc, {"keep": ["wet"], "tag_regex": r"^<<(\w+)>>"})
    assert e.value.code == "not-found"
    plan, out, rec = run_op(doc, {"keep": ["wet"], "tag_regex": r"^<<(\w+)>>", "allow_missing": True})
    assert rec.passed and out == doc


def _awkward(kind: str) -> bytes:
    d = Document()
    d.add_paragraph("Intro.")
    if kind == "field":
        p = d.add_paragraph("[DRY] See ")
        _ref_field(p)
    elif kind == "revision":
        p = d.add_paragraph("[DRY] Dry text ")
        r = p.add_run("added")
        ins = OxmlElement("w:ins"); ins.set(qn("w:id"), "5"); ins.set(qn("w:author"), "J"); ins.set(qn("w:date"), "2026-01-01T00:00:00Z")
        r._r.addprevious(ins); ins.append(r._r)
    elif kind == "cell":
        t = d.add_table(rows=1, cols=1)
        t.cell(0, 0).paragraphs[0].add_run("[DRY] the only paragraph of a cell")
    elif kind == "section":
        import copy as _copy
        p = d.add_paragraph("[DRY] last paragraph of the first section")       # its mark carries the section break
        p._p.get_or_add_pPr().append(_copy.deepcopy(d.element.body.find(qn("w:sectPr"))))
        d.add_paragraph("Body.")
    buf = io.BytesIO(); d.save(buf)
    return buf.getvalue()


@pytest.mark.parametrize("kind,tracked,code", [("field", True, "field-result"), ("revision", True, "existing-revisions"),
                                               ("cell", False, "table-cell"), ("section", True, "section-break")])
def test_refusals(kind, tracked, code):
    with pytest.raises(JobException) as e:
        run_op(_awkward(kind), {"keep": ["wet"], "tracked": tracked})
    assert e.value.code == code


def test_untracked_deletion_drops_pending_revisions_and_flags_them():
    plan, out, rec = run_op(_awkward("revision"), {"keep": ["wet"], "tracked": False})
    assert rec.passed, rec.to_markdown()
    assert plan.rev_delta.get("ins") == -1 and any("pending revisions" in f for f in plan.flags)


# ---- planted defects ---------------------------------------------------------------------------------------------

def test_defect_a_run_left_out_of_the_deletion(doc):
    plan, out, rec = run_op(doc, {"keep": ["wet"], "tracked": True})

    def unwrap(root):
        for d in root.iter(W + "del"):
            if d.getparent().tag == W + "p" and d.get(W + "author") == AUTHOR:
                r = d.find(W + "r")
                for t in r.findall(W + "delText"):
                    t.tag = W + "t"
                d.addprevious(r)
                d.getparent().remove(d)
                break
    rec = run_suite(doc, mutate(out, "word/document.xml", unwrap), plan.to_dict(), "s", "o")
    assert "V10" in failed(rec)


def test_shared_defect_multi_system_tags_need_every_system(doc, monkeypatch):
    monkeypatch.setattr(tagged_module, "stays", lambda systems, keep, delete: all(s in keep for s in systems))
    plan, out, rec = run_op(doc, {"keep": ["wet", "preaction"], "tracked": False})
    assert "Wet/Dry system paragraph." not in visible(out)          # deleted by plan and apply alike …
    assert failed(rec) == ["V10"]                                   # … the verifier keeps it


def test_shared_defect_separator_ignored(doc, monkeypatch):
    monkeypatch.setattr(tagged_module, "systems_of", lambda tag, sep: [tagged_module.norm_system(tag)])
    plan, out, rec = run_op(doc, {"keep": ["wet"], "tracked": True})
    assert failed(rec) == ["V10"]
    assert any("tagged_paragraphs" in e for e in v10_errors(rec))


def test_defect_an_extra_paragraph_removed(doc):
    plan, out, rec = run_op(doc, {"keep": ["wet"], "tracked": False})

    def drop(root):
        for p in root.iter(W + "p"):
            if "".join(t.text or "" for t in p.iter(W + "t")).startswith("N."):
                p.getparent().remove(p)
                break
    rec = run_suite(doc, mutate(out, "word/document.xml", drop), plan.to_dict(), "s", "o")
    assert "V2" in failed(rec) and "V10" in failed(rec)


# ---- range marks running out of a deleted paragraph (2026-10-21 review fix) ----------------------------------------

KEPT = "bookmark or range mark(s) in the deleted paragraph run on into text that stays"


def _tagged_marks() -> bytes:
    """"[DRY] Dry text." holds all of "Gone", then the start of "DryRef", whose end is in "Shared text." (untagged:
    it stays); "See ⟨REF DryRef⟩." cross-refers to it."""
    d = Document()
    d.add_paragraph("Intro.")
    dry = d.add_paragraph("[DRY] Dry text.")
    shared = d.add_paragraph("Shared text.")
    d.add_paragraph("[WET] Wet text.")
    see = d.add_paragraph("See ")
    r = see.add_run()
    for tag, val in (("fldChar", "begin"), ("instrText", " REF DryRef \\h "), ("fldChar", "separate")):
        el = OxmlElement(f"w:{tag}")
        el.set(qn("w:fldCharType"), val) if tag == "fldChar" else setattr(el, "text", val)
        r._r.append(el)
    see.add_run("Shared text.")
    end = see.add_run()._r
    el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), "end"); end.append(el)
    for p, tag, mid, name in ((dry, "bookmarkStart", "8", "Gone"), (dry, "bookmarkEnd", "8", ""),
                              (dry, "bookmarkStart", "9", "DryRef"), (shared, "bookmarkEnd", "9", "")):
        el = OxmlElement(f"w:{tag}"); el.set(qn("w:id"), mid)
        if name:
            el.set(qn("w:name"), name)
        p._p.append(el)
    buf = io.BytesIO(); d.save(buf)
    return buf.getvalue()


def _tagged(plan) -> Dict[str, Any]:
    return next(a for a in plan.property_assertions if a["kind"] == "tagged_paragraphs")


def _bookmarks(data: bytes) -> List[str]:
    """The body's children: a paragraph by its text, a bookmark mark between paragraphs by its kind and id."""
    body = DocxPackage(data).xml("word/document.xml").find(W + "body")
    return [accepted_text(c) if c.tag == W + "p" else f"{etree.QName(c).localname} {c.get(W + 'id')}"
            for c in body if c.tag in (W + "p", W + "bookmarkStart", W + "bookmarkEnd")]


@pytest.mark.parametrize("tracked", [True, False])
def test_a_bookmark_running_out_of_a_deleted_paragraph_stays_where_it_was(tracked):
    """Untracked, the start of "DryRef" stays where its deleted paragraph was, so the REF field still finds it; "Gone"
    (both ends in the deleted paragraph) goes. Tracked, the deleted paragraph keeps its bookmarks until accepted."""
    src = _tagged_marks()
    plan, out, rec = run_op(src, {"keep": ["wet"], "tracked": tracked})
    assert rec.passed, rec.to_markdown()
    root = DocxPackage(out).xml("word/document.xml")
    assert sorted(b.get(W + "id") for b in root.iter(W + "bookmarkEnd")) == \
        sorted(b.get(W + "id") for b in root.iter(W + "bookmarkStart"))
    if tracked:
        assert _tagged(plan)["expected_marks"] == {}
        assert [b.get(W + "name") for b in root.iter(W + "bookmarkStart")] == ["Gone", "DryRef"]
        assert not any(KEPT in f for f in plan.flags)
        return
    assert _tagged(plan)["expected_marks"] == {"word/document.xml": {"1": [2]}}
    assert _bookmarks(out) == ["Intro.", "bookmarkStart 9", "Shared text.", "Wet text.", "See Shared text."]
    assert [f for f in plan.flags if KEPT in f] == [
        f"word/document.xml paragraph 1: 1 {KEPT}; each stays where the paragraph was, so its bookmark keeps both ends"]


def test_shared_defect_the_old_rule_drops_a_crossing_bookmark(monkeypatch):
    monkeypatch.setattr(tagged_module, "kept_marks", lambda root, removed: {})
    src = _tagged_marks()
    plan, out, rec = run_op(src, {"keep": ["wet"], "tracked": False})
    assert [b.get(W + "id") for b in DocxPackage(out).xml("word/document.xml").iter(W + "bookmarkStart")] == []
    assert failed(rec) == ["V10"], rec.to_markdown()                  # plan and apply agree; the verifier does not
    assert any("tagged_paragraphs" in e and "the plan predicted" in e for e in v10_errors(rec))


def test_defect_the_kept_bookmark_dropped_from_the_output():
    src = _tagged_marks()
    plan, out, rec = run_op(src, {"keep": ["wet"], "tracked": False})
    assert rec.passed, rec.to_markdown()

    def orphan(root):
        body = root.find(W + "body")
        body.remove(body.find(W + "bookmarkStart"))
    rec = run_suite(src, mutate(out, "word/document.xml", orphan), plan.to_dict(), "s", "o")
    assert "V10" in failed(rec) and any("unpredicted change" in e for e in v10_errors(rec))
