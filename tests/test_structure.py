"""Structure is source evidence: mixed markers, explicit gaps, no document edits."""
import hashlib
import io

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from spec_surgeon.package import DocxPackage
from spec_surgeon.structure import analyze_structure


def package(document):
    out = io.BytesIO()
    document.save(out)
    return DocxPackage(out.getvalue())


def document(*lines):
    doc = Document()
    for line in lines:
        doc.add_paragraph(line)
    return doc


def numbering(doc, formats=None):
    formats = formats or [("decimal", "PART %1 -"), ("decimal", "%1.%2"), ("upperLetter", "%3."),
                          ("decimal", "%4."), ("lowerLetter", "%5.")]
    root = doc.part.numbering_part.element
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), "77")
    for index, (fmt, text) in enumerate(formats):
        level = OxmlElement("w:lvl")
        level.set(qn("w:ilvl"), str(index))
        for name, value in (("start", "1"), ("numFmt", fmt), ("lvlText", text), ("suff", "tab")):
            child = OxmlElement("w:" + name)
            child.set(qn("w:val"), value)
            level.append(child)
        abstract.append(level)
    root.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), "77")
    reference = OxmlElement("w:abstractNumId")
    reference.set(qn("w:val"), "77")
    num.append(reference)
    root.append(num)


def numbered(p, level, num_id=77):
    num = OxmlElement("w:numPr")
    for name, value in (("ilvl", level), ("numId", num_id)):
        child = OxmlElement("w:" + name)
        child.set(qn("w:val"), str(value))
        num.append(child)
    p._p.get_or_add_pPr().append(num)


def test_complete_typed_hierarchy_has_exact_source_identity_and_full_text():
    long = "A. " + "Supply approved products and retain instructions. " * 20
    pkg = package(document("SECTION 21 13 13", "WET PIPE SPRINKLER SYSTEMS", "PART 1 - GENERAL",
                           "1.01 SUMMARY", long, "1. Performance criteria", "a. First detail",
                           "(1) Nested detail", "(a) Deeper detail", "END OF SECTION"))
    original = pkg.data
    result = analyze_structure(pkg)
    rows = result["paragraphs"]
    assert [r["role"] for r in rows] == ["section_id", "section_title", "part", "article", "paragraph",
                                           "subparagraph", "subparagraph_2", "subparagraph_3", "subparagraph_4",
                                           "end_of_section"]
    assert [r["paragraph"] for r in rows] == list(range(10))
    assert rows[4]["text"] == long
    assert rows[8]["parent"] == 7 and rows[8]["reference_id"] == "1.1.A.1.a.1.a"
    assert rows[4]["section_id"] == "21 13 13" and rows[4]["article_label"] == "1.1"
    assert "Section 21 13 13; Article 1.1; beginning" in rows[4]["location"]
    assert result["source_sha256"] == hashlib.sha256(original).hexdigest()
    assert result["coverage"] == {"total": 10, "classified": 10, "unresolved": 0, "excluded": 0}
    assert pkg.data == original and not pkg.is_modified()


def test_automatic_numbering_precedes_typed_prefix_and_reads_inherited_style():
    doc = document("SECTION 21 13 13", "SPRINKLERS")
    numbering(doc)
    for level, text in ((0, "GENERAL"), (1, "SUMMARY"), (2, "1.1 This is provision wording"),
                        (3, "Subordinate requirement"), (4, "Detail")):
        p = doc.add_paragraph(text)
        numbered(p, level)
    # Effective list numbering inherited through a paragraph style is still seen.
    style = doc.styles.add_style("OfficeClause", 1)
    style._element.get_or_add_pPr().append(doc.paragraphs[4]._p.pPr.numPr)
    doc.paragraphs[4].style = style
    result = analyze_structure(package(doc))
    rows = result["paragraphs"]
    assert [r["role"] for r in rows[2:]] == ["part", "article", "paragraph", "subparagraph", "subparagraph_2"]
    assert rows[4]["content"] == "1.1 This is provision wording"
    assert rows[4]["numbering"]["source"] == "style"
    assert rows[4]["numbering"]["marker"] == "A."


def test_ambiguous_roman_dangling_numbering_and_parent_gaps_remain_unresolved():
    doc = document("SECTION 21 13 13", "SPRINKLERS", "PART 1 GENERAL", "1.1 SUMMARY",
                   "I. Could be Roman", "1. Missing parent", "A. Explicit provision",
                   "1.1 Looks typed but is in a broken list", "Unnumbered explanation")
    numbered(doc.paragraphs[7], 0, num_id=999)
    rows = analyze_structure(package(doc))["paragraphs"]
    assert rows[4]["disposition"] == "unresolved" and "Roman" in rows[4]["reason"]
    assert rows[5]["disposition"] == "unresolved" and "parent" in rows[5]["reason"]
    assert rows[6]["role"] == "paragraph"
    assert rows[7]["disposition"] == "unresolved" and "numbering" in rows[7]["reason"]
    assert rows[8]["disposition"] == "unresolved"


def test_sequential_typed_letters_disambiguate_c_but_not_isolated_i():
    doc = document("SECTION 21 13 13", "SPRINKLERS", "1.1 SUMMARY", "A. First", "B. Second", "C. Third",
                   "I. Ambiguous after C")
    rows = analyze_structure(package(doc))["paragraphs"]
    assert rows[5]["role"] == "paragraph"
    assert rows[6]["disposition"] == "unresolved"


def test_automatic_roman_pattern_remains_unresolved_even_with_article_context():
    doc = document("SECTION 21 13 13", "SPRINKLERS", "1.1 SUMMARY")
    numbering(doc, [("upperRoman", "%1.")])
    p = doc.add_paragraph("Possibly a different hierarchy")
    numbered(p, 0)
    assert analyze_structure(package(doc))["paragraphs"][-1]["disposition"] == "unresolved"


def test_table_textbox_drawing_and_revision_dispositions_keep_global_ordinals():
    doc = document("SECTION 21 13 13", "SPRINKLERS", "1.1 SUMMARY")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "A. Table requirement"
    host = doc.add_paragraph("A. Drawing host")
    pict = OxmlElement("w:pict")
    textbox = OxmlElement("w:txbxContent")
    p = OxmlElement("w:p")
    run, text = OxmlElement("w:r"), OxmlElement("w:t")
    text.text = "1. Textbox detail"
    run.append(text)
    p.append(run)
    textbox.append(p)
    pict.append(textbox)
    host._p.append(pict)
    changed = doc.add_paragraph("A. Changed provision")
    ins = OxmlElement("w:ins")
    ins.set(qn("w:id"), "8")
    ins.append(changed.runs[0]._r)
    changed._p.append(ins)
    doc.add_paragraph("A. Body requirement")
    result = analyze_structure(package(doc))
    assert [r["paragraph"] for r in result["paragraphs"]] == list(range(8))
    assert all(r["disposition"] == "excluded" for r in result["paragraphs"][3:7])
    assert result["paragraphs"][4]["text"] == "A. Drawing host"
    assert result["paragraphs"][5]["text"] == "1. Textbox detail"
    # The excluded heading/revision (the wording inserted whole: with the change rejected the line is empty, another
    # reading) may have changed parent context: do not inherit the old article until a new explicit article
    # establishes it.
    assert result["paragraphs"][7]["disposition"] == "unresolved"


def test_repeated_logical_addresses_invalidate_both_parents_and_descendants():
    doc = document("SECTION 21 13 13", "SPRINKLERS", "1.1 SUMMARY", "A. First", "1.1 DUPLICATE", "A. Second")
    rows = analyze_structure(package(doc))["paragraphs"]
    assert not any(r["hierarchy_reliable"] for r in rows[2:])
    assert rows[3]["parent"] == 2 and rows[5]["parent"] == 4


def test_labels_repeat_safely_in_distinct_section_occurrences():
    doc = document("SECTION 21 13 13", "SPRINKLERS", "1.1 SUMMARY", "A. First",
                   "SECTION 21 05 00", "COMMON", "1.1 SUMMARY", "A. Second")
    rows = analyze_structure(package(doc))["paragraphs"]
    assert all(r["hierarchy_reliable"] for r in rows)
    assert rows[3]["section_paragraph"] == 0 and rows[7]["section_paragraph"] == 4


def test_modified_overlay_refused_instead_of_mislabeling_source_pin():
    pkg = package(document("SECTION 21 13 13"))
    pkg.write(pkg.main_part(), pkg.read(pkg.main_part()))
    with pytest.raises(ValueError, match="unmodified"):
        analyze_structure(pkg)


def test_tab_stop_definitions_are_not_text_and_hidden_headings_are_excluded():
    from docx.shared import Inches
    doc = document("SECTION 21 13 13", "SPRINKLERS", "1.1 SUMMARY", "A. Visible provision", "1.2 Hidden article")
    doc.paragraphs[3].paragraph_format.tab_stops.add_tab_stop(Inches(1))
    doc.paragraphs[4].runs[0].font.hidden = True
    rows = analyze_structure(package(doc))["paragraphs"]
    assert rows[3]["text"] == "A. Visible provision"
    assert rows[4]["disposition"] == "excluded" and "hidden" in rows[4]["reason"]


def test_unmarked_heading_interrupts_inherited_parent_context():
    doc = document("SECTION 21 13 13", "SPRINKLERS", "1.1 SUMMARY", "A. First", "Unmarked potential heading",
                   "B. Parent is uncertain", "1.2 REQUIREMENTS", "A. Parent is established")
    rows = analyze_structure(package(doc))["paragraphs"]
    assert rows[5]["disposition"] == "unresolved"
    assert rows[7]["reference_id"] == "1.2.A"


def test_mixed_case_section_reference_does_not_create_a_new_section():
    doc = document("SECTION 21 13 13", "SPRINKLERS", "1.1 REFERENCES", "Section 21 05 00 - Common work results")
    rows = analyze_structure(package(doc))["paragraphs"]
    assert rows[-1]["disposition"] == "unresolved"
    assert rows[-1]["section_id"] == "21 13 13"


def test_numbering_after_a_pending_revision_cannot_claim_stable_addresses():
    doc = document("SECTION 21 13 13", "SPRINKLERS")
    numbering(doc)
    for level, text in ((0, "GENERAL"), (1, "SUMMARY"), (2, "Inserted clause"), (2, "Later clause")):
        numbered(doc.add_paragraph(text), level)
    changed = doc.paragraphs[4]
    mark, ins = OxmlElement("w:rPr"), OxmlElement("w:ins")      # the paragraph mark itself is inserted: with the
    ins.set(qn("w:id"), "8")                                     # change rejected the list is one item shorter
    mark.append(ins)
    changed._p.get_or_add_pPr().append(mark)
    rows = analyze_structure(package(doc))["paragraphs"]
    assert rows[4]["disposition"] == "excluded"
    assert rows[5]["disposition"] == "unresolved" and "earlier pending revisions" in rows[5]["reason"]
