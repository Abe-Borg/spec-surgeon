"""Desktop scope, offline execution and shutdown use the real document pipeline."""
from __future__ import annotations

import json
import threading
from pathlib import Path

import pytest
from docx import Document

from spec_surgeon.campaign import create_campaign, load_campaign
from spec_surgeon.gui import Session


@pytest.fixture
def desktop(tmp_path, monkeypatch):
    from qt_support import close_windows, qt_app_or_skip
    from spec_surgeon import llm
    from spec_surgeon.ui.app import RecordingPrompts, build_legacy_window as build_window

    app = qt_app_or_skip()
    monkeypatch.setenv("SPEC_SURGEON_CONFIG_DIR", str(tmp_path / "config"))
    monkeypatch.setenv("SPEC_SURGEON_NO_UPDATE_CHECK", "1")
    monkeypatch.setattr(llm, "credentials_found", lambda: False)

    def no_model():
        raise AssertionError("An offline desktop action must not construct a model client")

    monkeypatch.setattr(llm, "client", no_model)
    windows = []

    def create(session, prompts=None):
        window = build_window(session, prompts=prompts or RecordingPrompts())
        windows.append(window)
        return window

    yield app, create
    close_windows(app, windows)


def project(tmp_path):
    inputs = tmp_path / "inputs"
    inputs.mkdir()
    (inputs / "PROJECT.json").write_text('{"name":"Selected project"}', encoding="utf-8")
    for name in ("First.docx", "Second.docx"):
        document = Document()
        document.add_paragraph("ALPHA")
        document.add_paragraph("Keep this wording.")
        document.sections[0].header.paragraphs[0].text = "HEADER TO REMOVE"
        document.save(inputs / name)
    return create_campaign([str(inputs)], tmp_path / "workspace")


def margin_spec():
    return {"request": "Set left margin to one inch", "summary": "Left margin", "mode": "formatting",
            "ops": [{"op": "set_page_setup", "params": {"margins": {"left": 1440}}}]}


@pytest.mark.parametrize("kind", ["title", "spec", "chain"])
@pytest.mark.parametrize("with_model", [False, True])
def test_direct_goal_changes_only_selected_document_and_reopens_with_that_scope(desktop, tmp_path, kind, with_model):
    from fake_claude import FakeClient
    from spec_surgeon.ui.app import RecordingPrompts

    campaign = project(tmp_path)
    originals = {d["id"]: Path(d["source"]).read_bytes() for d in campaign.documents.values()}
    path = tmp_path / "saved.json"
    saved = margin_spec()
    if kind == "chain":
        saved = {"name": "Margin chain", "stages": [{"name": "Left margin", "spec": saved}]}
    path.write_text(json.dumps(saved), encoding="utf-8")
    fake = FakeClient() if with_model else None
    window = desktop[1](Session(campaign=campaign, out_root=tmp_path / "out", client=fake),
                        RecordingPrompts(files=[str(path)]))
    window.do_campaign()
    dialog = window.campaign_dialog
    dialog.documents.selectRow(0)
    selected = dialog._selected_documents()[0]
    other = next(d for d in campaign.documents if d != selected)
    dialog.scope.setCurrentIndex(1)
    dialog.request.setPlainText("Make ALPHA bold" if kind == "title" else margin_spec()["request"])
    if kind == "title":
        dialog.goal_mode.setCurrentIndex(1)
        dialog.title_text.setText("ALPHA")
        dialog.properties["run", "b"].setCurrentIndex(1)
    else:
        dialog.pick_spec()
    dialog.render.setChecked(False)
    dialog.memory.setChecked(False)
    dialog.run_workspace()
    window._spec_surgeon["wait"]()
    assert campaign.documents[selected]["status"] == "READY", dialog.evidence.toPlainText()
    assert campaign.documents[other]["status"] != "READY"
    assert not campaign.documents[other].get("output")
    output = Document(campaign.documents[selected]["output"])
    if kind == "title":
        assert output.paragraphs[0].runs[0].bold is True
    else:
        assert output.sections[0].left_margin.twips == 1440
    assert all(Path(campaign.documents[d]["source"]).read_bytes() == data for d, data in originals.items())
    assert dialog.scope.currentIndex() == 1
    assert dialog._selected_documents() == [selected]
    restored = load_campaign(campaign.out_dir)
    assert restored.projects[dialog._selected_project]["goal_scope"] == [selected]
    destination = tmp_path / "delivery"
    restored.export(destination)
    assert [p.name for p in destination.rglob("*.docx")] == [Path(campaign.documents[selected]["source"]).name]
    if fake is not None:
        assert fake.requests == []


def test_selected_scope_without_a_selection_cannot_fall_back_to_all_documents(desktop, tmp_path):
    campaign = project(tmp_path)
    window = desktop[1](Session(campaign=campaign, out_root=tmp_path / "out"))
    window.do_campaign()
    dialog = window.campaign_dialog
    dialog.scope.setCurrentIndex(1)
    dialog.request.setPlainText("Make ALPHA bold")
    dialog.goal_mode.setCurrentIndex(1)
    dialog.title_text.setText("ALPHA")
    dialog.properties["run", "b"].setCurrentIndex(1)
    dialog.run_workspace()
    assert not window.busy
    assert "Select the documents" in dialog.status.text()
    assert not campaign.projects[dialog._selected_project]["goal"]
    assert not campaign.projects[dialog._selected_project]["runs"]
    assert all(not d.get("output") for d in campaign.documents.values())
    dialog.scope.setCurrentIndex(0)  # Leave valid instructions for the close-time save.


@pytest.mark.parametrize("failure", ["conflict", "wrong-pattern"])
def test_keyless_saved_spec_is_tried_before_model_and_pauses_cleanly_if_it_needs_repair(desktop, tmp_path, failure):
    from spec_surgeon.ui.app import RecordingPrompts

    campaign = project(tmp_path)
    saved = margin_spec()
    if failure == "conflict":
        saved["ops"].append({"op": "set_page_setup", "params": {"margins": {"right": 2880}}})
    else:
        saved["applies_to"] = {"pattern_ids": ["different-source-pattern"]}
    path = tmp_path / "conflicting.json"
    path.write_text(json.dumps(saved), encoding="utf-8")
    window = desktop[1](Session(campaign=campaign, out_root=tmp_path / "out"),
                        RecordingPrompts(files=[str(path)]))
    window.do_campaign()
    dialog = window.campaign_dialog
    dialog.request.setPlainText(saved["request"])
    dialog.pick_spec()
    dialog.render.setChecked(False)
    dialog.memory.setChecked(False)
    dialog.run_workspace()
    window._spec_surgeon["wait"]()
    assert not window.claude_ok
    current = campaign.projects[dialog._selected_project]
    assert "needs an Anthropic API key" in current["error"]
    assert all(d["status"] == "UNRESOLVED" and not d.get("output") for d in campaign.documents.values())
    recovery = json.loads((Path(current["runs"][0]["path"]) / "recovery.json").read_text(encoding="utf-8"))
    assert len(recovery["attempts"]) == 1
    assert recovery["budget"]["calls"] == 0
    assert ("split into separate specs" if failure == "conflict" else "pattern") in json.dumps(recovery["attempts"])


def test_scoped_goal_revision_retains_other_verified_output(desktop, tmp_path):
    from spec_surgeon.recovery import RecoveryOptions

    campaign = project(tmp_path)
    pid = next(iter(campaign.projects))
    ids = campaign.projects[pid]["document_ids"]
    contract = {"kind": "format_section_title", "target": {"title": "ALPHA"}, "run": {"b": True}}
    campaign.set_goal(pid, "Make ALPHA bold", adaptive_contract=contract)
    options = RecoveryOptions(render=False, use_memory=False)
    campaign.run(options=options)
    other_output = Path(campaign.documents[ids[1]]["output"])
    other_bytes = other_output.read_bytes()
    revised = {"kind": "format_section_title", "target": {"title": "ALPHA"}, "run": {"i": True}}
    campaign.set_goal(pid, "Make ALPHA italic", adaptive_contract=revised, doc_ids=[ids[0]])
    campaign.run(options=options)
    assert campaign.documents[ids[1]]["output"] == str(other_output)
    assert other_output.read_bytes() == other_bytes
    assert all(d["status"] == "READY" for d in campaign.documents.values())
    assert Document(campaign.documents[ids[0]]["output"]).paragraphs[0].runs[0].italic is True
    assert Document(other_output).paragraphs[0].runs[0].bold is True
    campaign.export(tmp_path / "delivery")


def test_close_during_real_run_keeps_window_and_worker_until_delivery_and_reports_finish(desktop, tmp_path, monkeypatch):
    from spec_surgeon import batch
    from spec_surgeon.ui.app import RecordingPrompts

    campaign = project(tmp_path)
    source = Path(next(iter(campaign.documents.values()))["source"])
    source_dir = tmp_path / "single"
    source_dir.mkdir()
    only_source = source_dir / source.name
    only_source.write_bytes(source.read_bytes())
    session = Session(folder=source_dir, out_root=tmp_path / "out")
    session.load_spec(Path(__file__).resolve().parent.parent / "examples/specs/remove-all-headers-footers.json")
    prompts = RecordingPrompts(answers=[True])
    window = desktop[1](session, prompts)
    window.show()
    window.request_page.render_cb.setChecked(False)
    window.do_preview()
    window._spec_surgeon["wait"]()
    entered, release = threading.Event(), threading.Event()
    render = batch._render_and_compare

    def pause_before_delivery(*args, **kwargs):
        entered.set()
        assert release.wait(10)
        return render(*args, **kwargs)

    monkeypatch.setattr(batch, "_render_and_compare", pause_before_delivery)
    window.run_page.inplace_rb.setChecked(True)
    window.do_run()
    try:
        assert entered.wait(10)
        window.close()
        assert window.isVisible() and window.busy
        assert "finish" in window.status_lbl.text().lower()
        worker = next(t for t in threading.enumerate() if t.name == "spec-surgeon-job")
        assert not worker.daemon
        assert len([p for p in prompts.shown if p[0] == "ask"]) == 1
    finally:
        release.set()
        window._spec_surgeon["wait"]()
    assert Document(only_source).sections[0].header.paragraphs[0].text == ""
    assert (session.last_run_dir / "BATCH_REPORT.md").is_file()
    assert session.last_batch.files[0].delivered_to == str(only_source)
    window.close()
    assert not window.isVisible()


def test_natural_language_preview_is_the_only_approval_and_pins_current_options(desktop, tmp_path):
    from fake_claude import FakeClient, resp, tool

    campaign = project(tmp_path)
    inputs = Path(next(iter(campaign.documents.values()))["source"]).parent
    saved = margin_spec()
    fake = FakeClient(resp(tool("submit_editspec", {"spec": saved})))
    session = Session(folder=inputs, out_root=tmp_path / "out", client=fake)
    window = desktop[1](session)
    window.request_page.request.setPlainText(saved["request"])
    window.request_page.render_cb.setChecked(False)
    assert not window.do_run()
    window.do_preview()
    window._spec_surgeon["wait"]()
    assert session.outcome_review is None, window.log.toPlainText()
    assert window.preview_page.outcome_card.isHidden()
    assert {Path(row.source).name for row in session.preview.batch_plan.files} == {"First.docx", "Second.docx"}
    assert window.run_page.run_btn.isEnabled()
    window.request_page.render_cb.setChecked(True)
    assert not window.run_page.run_btn.isEnabled()  # changing preview options invalidates consent
    assert not window.do_run()
    window.request_page.render_cb.setChecked(False)
    assert window.run_page.run_btn.isEnabled()
    window.do_run()
    window._spec_surgeon["wait"]()
    assert all(f.status == "PASS" for f in session.last_batch.files), window.log.toPlainText()
    assert all(Document(f.output).sections[0].left_margin.twips == 1440 for f in session.last_batch.files)
    assert all(Document(d["source"]).sections[0].left_margin.twips == 1800 for d in campaign.documents.values())
    assert len(fake.requests) == 1  # compilation, with no second intent planner


def test_preview_exposes_the_recipe_target_without_a_second_target_selection(desktop, tmp_path):
    from fake_claude import FakeClient, resp, tool

    inputs = tmp_path / "ambiguous"
    inputs.mkdir()
    source = inputs / "Repeated title.docx"
    document = Document()
    for text in ("ALPHA", "Keep this wording.", "ALPHA"):
        document.add_paragraph(text)
    document.save(source)
    request = "Make the intended ALPHA title bold"
    spec = {"request": request, "summary": "Bold the first title", "mode": "formatting",
            "ops": [{"op": "set_properties", "select": {"scope": ["document"], "paragraphs": {"index_range": [0, 1]}},
                     "params": {"run": {"b": True}, "tracked": False}}]}
    fake = FakeClient(resp(tool("submit_editspec", {"spec": spec})))
    session = Session(folder=inputs, out_root=tmp_path / "out", client=fake)
    window = desktop[1](session)
    window.request_page.request.setPlainText(request)
    window.request_page.render_cb.setChecked(False)
    window.do_preview()
    window._spec_surgeon["wait"]()
    page = window.preview_page
    assert session.outcome_review is None, window.log.toPlainText()
    assert page.outcome_card.isHidden()
    assert session.preview.samples[0].passed
    assert window.run_page.run_btn.isEnabled()
    window.do_run()
    window._spec_surgeon["wait"]()
    assert session.last_batch.files[0].status == "PASS", window.log.toPlainText()
    output = Document(session.last_batch.files[0].output)
    assert output.paragraphs[0].runs[0].bold is True
    assert output.paragraphs[2].runs[0].bold is not True
    assert all(p.runs[0].bold is not True for p in Document(source).paragraphs)
    assert len(fake.requests) == 1


def test_preview_of_retained_retry_uses_current_recovery_options_and_keeps_approved_outcomes(desktop, tmp_path):
    from fake_claude import FakeClient, resp, tool
    from spec_surgeon.goals import compile_goal_candidate

    campaign = project(tmp_path)
    source = Path(next(iter(campaign.documents.values()))["source"])
    contract = {"kind": "document_effects", "effects": [
        {"id": "margin", "type": "page_setup", "sections": "all", "margins": {"left": 1440}}]}
    wrong = margin_spec()
    wrong["ops"][0]["params"]["margins"]["left"] = 2880
    candidate = compile_goal_candidate(source, contract)
    # This separate, explicitly enforced workflow retains its contract across retry.
    from spec_surgeon.batch import run_batch
    from spec_surgeon.goals import resolve_goal
    sources = [Path(document["source"]) for document in campaign.documents.values()]
    bindings = {str(path): resolve_goal(path, contract)["binding"] for path in sources}
    original_run = tmp_path / "enforced-run"
    original = run_batch([str(path) for path in sources], "edit-spec", {}, original_run,
                         spec=wrong, workers=1, do_render=False, goal_contract=contract,
                         goal_bindings=bindings, enforce_goals=True)
    assert all(file.status == "FAIL" for file in original.files)
    fake = FakeClient(resp(tool("submit_repair", {"candidate": candidate})),
                      resp(tool("review_repair", {"verdict": "accept", "reason": "The verified margin matches the approved outcome."})))
    session = Session(folder=source.parent, out_root=tmp_path / "out", client=fake)
    window = desktop[1](session)
    window.request_page.request.setPlainText(margin_spec()["request"])
    window.request_page.render_cb.setChecked(False)
    window.run_page.recovery_memory.setChecked(False)
    session.last_run_dir, session.last_batch = original_run, original
    window.do_set_aside()
    window._spec_surgeon["wait"]()
    assert session.pending_recovery_run == original_run
    window.do_preview()
    window._spec_surgeon["wait"]()
    assert session.recovery is not None, window.log.toPlainText()
    assert session.pending_recovery_run is None
    assert session.recovery.key["adaptive_contract"] == contract
    assert not session.recovery.options_snapshot["render"]
    assert not session.recovery.options_snapshot["use_memory"]
    assert len(session.recovery.ready) == 2, window.log.toPlainText()
    assert window.stack.currentWidget() is window.preview_page
    assert window._preview_current()
    assert window.run_page.run_btn.isEnabled()
    assert window.run_page.run_btn.text() == "Export recovery copies"
    window.do_run()
    window._spec_surgeon["wait"]()
    assert session.last_recovery is not None
    assert all(Document(f.delivered_to).sections[0].left_margin.twips == 1440 for f in session.last_recovery.ready)
    assert len(fake.requests) == 2


def test_custom_generation_prepares_a_verified_preview_for_the_single_run_decision(desktop, tmp_path):
    from fake_claude import FakeClient
    from test_phase8_fixes_window import OFFER_TURN, WRITER

    campaign = project(tmp_path)
    source = Path(next(iter(campaign.documents.values()))["source"])
    before = {Path(d["source"]): Path(d["source"]).read_bytes() for d in campaign.documents.values()}
    fake = FakeClient(OFFER_TURN, WRITER)
    session = Session(folder=source.parent, out_root=tmp_path / "out", client=fake)
    window = desktop[1](session)
    window.request_page.request.setPlainText("put the header 0.625 in from the top")
    window.request_page.render_cb.setChecked(False)
    window.do_preview()
    window._spec_surgeon["wait"]()
    (_, choose), = window._spec_surgeon["choices"]()
    choose.click()
    window._spec_surgeon["wait"]()
    assert session.pending_custom_transform is None
    assert session.generated is not None, window.log.toPlainText()
    assert session.outcome_review is None
    assert window.preview_page.outcome_card.isHidden()
    assert window.stack.currentWidget() is window.preview_page
    assert window._preview_current()
    assert window.run_page.run_btn.isEnabled(), window.run_page.run_hint.text()
    assert window.preview_page.continue_btn.isEnabled()
    assert window.preview_page.save_btn.isEnabled()
    assert len(fake.requests) == 2
    assert all(path.read_bytes() == data for path, data in before.items())
    window.do_run()
    window._spec_surgeon["wait"]()
    assert all(trial.passed for trial in session.last_trials), window.log.toPlainText()
    assert all(Document(trial.output).sections[0].header_distance.twips == 900 for trial in session.last_trials)
    assert all(path.read_bytes() == data for path, data in before.items())
    assert len(fake.requests) == 2
