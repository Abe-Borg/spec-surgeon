"""The simple experience must preserve the engine's authority and explicit document scope."""
from __future__ import annotations

import copy
import json
from pathlib import Path

import pytest
from docx import Document

from fake_claude import FakeClient, api_error, raise_on_enter, resp, text, tool
from spec_surgeon.document_task import DocumentTask
from spec_surgeon.gui import Session
from spec_surgeon import llm, review_display
from spec_surgeon.resources import DEFAULT_MAX_CALLS


REQUEST = "Make the FIRE SUPPRESSION title bold."
CONTRACT = {"kind": "document_effects", "effects": [{"id": "title", "type": "format",
    "target": {"locators": [{"exact_text": "FIRE SUPPRESSION"}], "cardinality": 1}, "run": {"b": True}}]}
SPEC = {"request": REQUEST, "summary": "Make the title bold", "mode": "formatting",
        "ops": [{"op": "set_properties", "select": {"paragraphs": {"text_regex": "^FIRE SUPPRESSION$"}},
                 "params": {"run": {"b": True}, "tracked": False}}]}


def document(path, duplicate=False):
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_paragraph("SECTION 21 13 13")
    doc.add_paragraph("FIRE SUPPRESSION")
    if duplicate:
        doc.add_paragraph("FIRE SUPPRESSION")
    doc.add_paragraph("Keep this technical wording intact.")
    doc.save(path)
    return path


def client(spec=None):
    return FakeClient(resp(tool("submit_editspec", {"spec": spec or SPEC})))


def check_after(task, monkeypatch, calls):
    """Set the spending check so that a request makes ``calls`` scripted model calls and stops before the next: a call
    reserves its largest answer (128,000 output tokens at $20 per million, $2.56) and, here, no input; a scripted call
    costs $0.0008 (100 input and 20 output tokens)."""
    from spec_surgeon import llm
    monkeypatch.setattr(llm, "estimate_input_tokens", lambda **_: 0)
    task.set_spending_limit(2.56 + 0.0008 * calls - 0.0004)
    return task.spending_limit


def task_for(tmp_path, fake=None, **kwargs):
    return DocumentTask(Session(client=fake or client(), out_root=tmp_path / "out", visual_check=False),
                        render=False, **kwargs)


def test_explicit_attachments_never_expand_to_a_folder_and_support_same_names(tmp_path):
    first = document(tmp_path / "a" / "Spec.docx")
    second = document(tmp_path / "b" / "Spec.docx")
    document(tmp_path / "a" / "Unselected.docx")
    task = task_for(tmp_path)
    task.attach([first, second, first])
    assert task.session.files() == [first, second]
    before = [p.read_bytes() for p in (first, second)]
    view = task.perform("send", text=REQUEST)
    assert view["stage"] == "preview"
    assert len(task.session.client.requests) == 1  # one plan, no second intent model
    assert not task.results
    task.perform("primary", token=view["token"])
    assert task.stage == "done"
    assert set(task.results) == {str(first), str(second)}
    assert len(set(task.results.values())) == 2
    assert [p.read_bytes() for p in (first, second)] == before
    for output in task.results.values():
        doc = Document(output)
        assert doc.paragraphs[1].runs[0].bold
        assert doc.paragraphs[2].text == "Keep this technical wording intact."


def test_missing_attachment_is_not_silently_removed_from_scope(tmp_path):
    source = document(tmp_path / "Spec.docx")
    task = task_for(tmp_path)
    task.attach([source])
    source.unlink()
    with pytest.raises(ValueError, match="missing"):
        task.session.files()


def test_template_is_separate_and_invalid_drop_is_atomic(tmp_path):
    source = document(tmp_path / "Spec.docx")
    template = document(tmp_path / "Template.docx")
    task = task_for(tmp_path)
    with pytest.raises(ValueError, match="Word"):
        task.attach([source, tmp_path])
    assert task.session.documents == []
    task.attach([source, template])
    task.attach([template], template=True)
    assert task.session.documents == [source]
    assert task.session.template == template


def test_changed_source_cannot_use_the_displayed_approval(tmp_path):
    source = document(tmp_path / "Spec.docx")
    task = task_for(tmp_path)
    task.attach([source])
    view = task.perform("send", text=REQUEST)
    source.write_bytes(source.read_bytes() + b"changed")
    with pytest.raises(ValueError, match="changed"):
        task.perform("primary", token=view["token"])
    assert not task.results


def test_attachment_change_keeps_request_but_invalidates_plan(tmp_path):
    source = document(tmp_path / "a" / "Spec.docx")
    extra = document(tmp_path / "b" / "Other.docx")
    task = task_for(tmp_path)
    task.attach([source])
    view = task.perform("send", text=REQUEST)
    task.attach([extra])
    assert task.decisions == [REQUEST]
    assert task.session.preview is None
    assert task.session.conversation is None
    with pytest.raises(ValueError, match="changed"):
        task.perform("primary", token=view["token"])


def test_target_question_uses_chat_numbers_without_approving_plan(tmp_path):
    source = document(tmp_path / "Spec.docx", duplicate=True)
    selected = copy.deepcopy(SPEC)
    selected["ops"][0]["select"]["paragraphs"]["index_range"] = [2, 3]
    fake = FakeClient(resp(tool("ask_user", {"message": "Which title should change?", "choices": [
        {"kind": "clarify", "label": "Paragraph 2: FIRE SUPPRESSION"},
        {"kind": "clarify", "label": "Paragraph 3: FIRE SUPPRESSION"}]})),
        resp(tool("submit_editspec", {"spec": selected})))
    task = task_for(tmp_path, fake)
    task.attach([source])
    view = task.perform("send", text=REQUEST)
    assert view["stage"] == "question" and not view["action"]
    labels = "\n".join(c["label"] for c in view["choices"])
    assert "Paragraph 2" in labels and "Paragraph 3" in labels
    view = task.perform("send", text="2")
    assert view["stage"] == "preview"
    assert len(task.session.client.requests) == 2
    assert not task.results
    task.perform("primary", token=view["token"])
    result = Document(task.results[str(source)])
    assert not result.paragraphs[1].runs[0].bold
    assert result.paragraphs[2].runs[0].bold


def test_clarification_continues_same_conversation(tmp_path):
    fake = FakeClient(resp(tool("ask_user", {"questions": ["Which title should change?"]})),
        resp(tool("submit_editspec", {"spec": SPEC})))
    task = task_for(tmp_path, fake)
    task.attach([document(tmp_path / "Spec.docx")])
    task.perform("send", text="Make the title bold")
    original = task.session.conversation
    view = task.perform("send", text="FIRE SUPPRESSION")
    assert task.session.conversation is original
    assert view["stage"] == "preview"
    assert len(task.messages) >= 2


# ---- each of Claude's words shows once: in the chat, or in the clarification card ---------------------------------

REPLY = "You're right, the logo matches.\n\nThese differences remain:\n\n• 21 05 00 uses 11 pt\n• 28 46 21 has an extra tab"


def said(view):
    return [m["text"] for m in view["messages"] if m["role"] == "assistant"]


def test_a_plain_reply_shows_in_the_chat_and_never_again_in_the_card(tmp_path):
    for answer in (resp(text(REPLY), stop="end_turn"),                                   # words, no tool
                   resp(tool("ask_user", {"message": REPLY, "suggested_replies": ["Fix the sizes"]}))):
        task = task_for(tmp_path, FakeClient(answer))
        task.attach([document(tmp_path / "Spec.docx")])
        view = task.perform("send", text=REQUEST)
        assert view["stage"] == "question"
        assert said(view) == [REPLY]                    # Claude's own "• " lines stay in its bubble
        assert view["questions"] == [] and view["choices"] == []
    assert view["chips"] == ["Fix the sizes"]           # the card keeps only what the user can tap


def test_a_question_of_several_lines_shows_whole_in_the_card_and_never_in_the_chat(tmp_path):
    question = "Which one should change?\n\n- the header\n- the footer"
    task = task_for(tmp_path, FakeClient(resp(tool("ask_user", {"message": "I checked both.", "questions": [question]}))))
    task.attach([document(tmp_path / "Spec.docx")])
    view = task.perform("send", text=REQUEST)
    assert said(view) == ["I checked both."]
    assert view["questions"] == [question]


def test_a_failure_is_said_once_in_the_chat(tmp_path):
    task = task_for(tmp_path, FakeClient(resp(stop="end_turn")))                        # an empty answer
    task.attach([document(tmp_path / "Spec.docx")])
    view = task.perform("send", text=REQUEST)
    assert view["stage"] == "question"
    assert said(view) == ["Claude returned an empty answer"] and view["questions"] == []


def test_a_reply_the_chat_did_not_say_still_reaches_the_card(tmp_path):
    fake = FakeClient(resp(tool("ask_user", {"questions": ["Which title?"]})),
                      raise_on_enter(api_error("Overloaded", 529)))
    task = task_for(tmp_path, fake)
    task.attach([document(tmp_path / "Spec.docx")])
    task.perform("send", text=REQUEST)
    task.session.conversation.recovery_active = True    # a recovery turn's words are not the chat's
    view = task.perform("send", text="FIRE SUPPRESSION")
    message = task.session.turn.message
    assert view["stage"] == "question" and message and message not in said(view)
    assert view["questions"] == [message]


def test_an_older_bubble_with_the_same_words_does_not_hide_a_reply_the_chat_did_not_say(tmp_path):
    words = "Planning paused after APIStatusError: Overloaded"     # what the recovery turn below will say
    fake = FakeClient(resp(text(words), stop="end_turn"), raise_on_enter(api_error("Overloaded", 529)))
    task = task_for(tmp_path, fake)
    task.attach([document(tmp_path / "Spec.docx")])
    task.perform("send", text=REQUEST)
    task.session.conversation.recovery_active = True
    view = task.perform("send", text="FIRE SUPPRESSION")
    assert task.session.turn.message == words and said(view) == [words]   # the bubble answers the first request
    assert view["questions"] == [words]                                    # the card answers the second


def test_readable_review_preserves_exact_wording_and_every_effect(tmp_path):
    contract = copy.deepcopy(CONTRACT)
    contract["effects"].append({"id": "words", "type": "replace_text", "target": {"all": True},
        "find": "Old <name>", "replace": "New & Co.", "expected_matches": 1})
    sections = review_display.effects(contract)
    assert len(sections) == 2
    text = "\n".join(line for section in sections for line in section["details"])
    assert "Old <name>" in text and "New & Co." in text
    assert "Bold: yes" in text
    assert "Required text replacements: 1" in text
    assert '"effects"' not in text
    assert "Definition · Kind: bullet" in review_display.fields({"definition": {"kind": "bullet"}})


def test_journal_restores_work_without_restoring_approval_or_resetting_spend(tmp_path):
    journal = tmp_path / "conversation.json"
    task = task_for(tmp_path, journal=journal)
    task.attach([document(tmp_path / "Spec.docx")])
    view = task.perform("send", text=REQUEST)
    task.perform("primary", token=view["token"])
    task.save()
    restored = task_for(tmp_path, journal=journal)
    assert restored.restore()
    assert restored.stage == "done"                   # the run finished: the copies are the working set now
    assert restored.decisions == [] and restored.generation == 1 and len(restored.history) == 1
    assert restored.session.preview is None
    assert restored._budget.calls == task._budget.calls
    assert restored._budget.estimated_cost_usd == task._budget.estimated_cost_usd
    assert "outcome_approval" not in json.loads(journal.read_text(encoding="utf-8"))


def test_interrupted_call_pauses_before_new_spending(tmp_path):
    journal = tmp_path / "conversation.json"
    task = task_for(tmp_path, journal=journal)
    task.attach([document(tmp_path / "Spec.docx")])
    task.decisions = [REQUEST]
    task.save(in_progress=True)
    restored = task_for(tmp_path, journal=journal)
    restored.restore()
    assert restored.stage == "budget" and restored.pause_reason() == "interrupted"
    view = restored.view()
    assert view["action"] == "Continue" and view["budget_paused"] and view["pause_reason"] == "interrupted"
    restored.perform("primary", token=view["token"])
    assert len(restored.session.client.requests) == 1
    assert restored._budget.max_cost_usd is None and restored._budget.max_total_tokens is None
    assert restored.stage == "preview" and restored.messages[-2]["text"] == "Continuing where the work stopped."


def test_every_desktop_call_gets_the_models_full_output_at_medium_effort_with_no_spending_ceiling(tmp_path):
    """The conversation runs at the model's own default effort (fewer tool rounds per turn); recovery, the hard path,
    keeps high."""
    journal = tmp_path / "conversation.json"
    task = task_for(tmp_path, journal=journal)
    task.attach([document(tmp_path / "Spec.docx")])
    task.perform("send", text=REQUEST)
    request = task.session.client.requests[0]
    assert request["max_tokens"] == 128000 and request["output_config"] == {"effort": "medium"}
    # no call-count pause (the owner's decision of 2026-10-27): the conversation's ledger has no limit of its own,
    # and the spending check is the user's (off here: the window applies Settings')
    assert (task._budget.max_cost_usd, task._budget.max_total_tokens, task._budget.max_calls) == (
        None, None, llm.UNLIMITED_CALLS)
    assert task.spending_limit is None and task._request_budget.max_cost_usd is None
    options = task._recovery_options()
    options.validate()
    assert (options.max_output_tokens, options.effort, options.max_calls) == (128000, "high", DEFAULT_MAX_CALLS)
    assert (options.max_cost_usd, options.max_total_tokens) == (None, None)
    task.save()
    restored = task_for(tmp_path, journal=journal)
    assert restored.restore()
    assert (restored._budget.max_output_tokens, restored._budget.effort) == (128000, "medium")
    assert (restored._budget.max_cost_usd, restored._budget.max_total_tokens) == (None, None)


def test_a_journal_from_the_allowance_version_keeps_its_spending_but_not_its_ceilings(tmp_path):
    journal = tmp_path / "conversation.json"
    task = task_for(tmp_path, journal=journal)
    task.attach([document(tmp_path / "Spec.docx")])
    task.decisions = [REQUEST]
    task.save()
    data = json.loads(journal.read_text(encoding="utf-8"))
    data["budget"].update(max_cost_usd=10.0, max_total_tokens=1_000_000, calls=30, estimated_cost_usd=10.2,
                          exhausted_reason="estimated recovery spending ceiling reached ($10.2000 accounted, "
                                           "$10.0000 ceiling); provider billing may differ")
    journal.write_text(json.dumps(data))
    restored = task_for(tmp_path, journal=journal)
    assert restored.restore()
    assert restored.stage == "resume" and restored._budget.exhausted_reason is None
    assert (restored._budget.max_cost_usd, restored._budget.max_total_tokens) == (None, None)
    assert restored._budget.calls == 30 and restored._budget.spent_usd == pytest.approx(10.2)
    assert restored.view()["cost"] == pytest.approx(10.2)


def test_each_request_gets_the_spending_check_afresh(tmp_path, monkeypatch):
    """The check is per request: a message, a pick, an approval. Two one-call requests under a one-call check both
    run; the conversation's ledger counts them both."""
    fake = FakeClient(resp(tool("ask_user", {"questions": ["Which title?"]})),
                      resp(tool("submit_editspec", {"spec": SPEC})))
    task = task_for(tmp_path, fake)
    task.attach([document(tmp_path / "Spec.docx")])
    check_after(task, monkeypatch, 1)
    assert task.perform("send", text=REQUEST)["stage"] == "question"
    view = task.perform("send", text="FIRE SUPPRESSION")
    assert view["stage"] == "preview" and len(fake.requests) == task._budget.calls == 2
    assert task._request_budget.calls == 1 and view["cost"] == pytest.approx(0.0016)


def test_the_spending_check_stops_a_request_before_a_call_that_could_pass_it(tmp_path, monkeypatch):
    fake = FakeClient(resp(tool("list_styles", {"file": "Spec.docx"}, "t1")),
                      resp(tool("submit_editspec", {"spec": SPEC}, "t2")))
    task = task_for(tmp_path, fake)
    task.attach([document(tmp_path / "Spec.docx")])
    limit = check_after(task, monkeypatch, 1)
    view = task.perform("send", text=REQUEST)
    assert view["stage"] == "budget" and view["pause_reason"] == "limit" and view["action"] == "Continue"
    assert len(fake.requests) == 1 and not task.results                       # the second call never started
    assert task._cost()["request"] == pytest.approx(0.0008) and task._cost()["limit"] == limit
    # the user's go-ahead allows the limit again on top of what the request spent, and the work goes on
    view = task.perform("primary", token=view["token"])
    assert view["stage"] == "preview" and len(fake.requests) == 2
    assert task._request_budget.max_cost_usd == pytest.approx(0.0008 + limit)
    assert "may spend up to $" in [m["text"] for m in task.messages if m["role"] == "app"][-1]


def test_with_the_check_off_nothing_pauses_and_settings_apply_from_the_next_request(tmp_path, monkeypatch):
    fake = FakeClient(*[resp(tool("list_styles", {"file": "Spec.docx"}, f"t{i}")) for i in range(60)],
                      resp(tool("submit_editspec", {"spec": SPEC}, "done")))
    task = task_for(tmp_path, fake)
    task.attach([document(tmp_path / "Spec.docx")])
    assert task.perform("send", text=REQUEST)["stage"] == "preview" and len(fake.requests) == 61   # no 48-call pause
    task.set_spending_limit(0)
    assert task.spending_limit is None
    with pytest.raises(ValueError, match="US dollars"):
        task.set_spending_limit(-1)
    check_after(task, monkeypatch, 0)
    assert task._request_budget.max_cost_usd is None                 # the open request keeps its check; the next gets this one


def test_a_model_the_price_table_cannot_estimate_pauses_the_check_and_continue_goes_on_without_it(tmp_path):
    from spec_surgeon import llm
    fake = FakeClient(resp(tool("list_styles", {"file": "Spec.docx"}, "t1"), model="claude-unpriced-9"),
                      resp(tool("submit_editspec", {"spec": SPEC}, "t2")))
    task = task_for(tmp_path, fake)
    task.attach([document(tmp_path / "Spec.docx")])
    task.set_spending_limit(50)
    view = task.perform("send", text=REQUEST)
    assert view["stage"] == "budget" and view["pause_reason"] == "unpriced" and len(fake.requests) == 1
    view = task.perform("primary", token=view["token"])
    assert view["stage"] == "preview" and task._request_budget.max_cost_usd is None
    assert "price table cannot estimate" in [m["text"] for m in task.messages if m["role"] == "app"][-1]
    assert llm.UNLIMITED_CALLS == task._request_budget.max_calls


def test_a_request_paused_at_the_check_reopens_paused_and_goes_on(tmp_path, monkeypatch):
    journal = tmp_path / "conversation.json"
    fake = FakeClient(resp(tool("list_styles", {"file": "Spec.docx"}, "t1")),
                      resp(tool("submit_editspec", {"spec": SPEC}, "t2")))
    task = task_for(tmp_path, fake, journal=journal)
    task.attach([document(tmp_path / "Spec.docx")])
    limit = check_after(task, monkeypatch, 1)
    assert task.perform("send", text=REQUEST)["stage"] == "budget"
    restored = task_for(tmp_path, fake, journal=journal, spending_limit=limit)
    assert restored.restore()
    assert restored.stage == "budget" and restored.pause_reason() == "limit"
    assert restored._request_budget.spent_usd == pytest.approx(0.0008)
    restored.perform("primary", token=restored.action_token())
    assert restored._request_budget.max_cost_usd == pytest.approx(0.0008 + limit)
    # the check turned off in the meantime: nothing to ask on reopening
    off = task_for(tmp_path, fake, journal=journal)
    task.save()
    assert off.restore() and off.stage == "resume"


def test_simple_prompt_keeps_drafting_behind_exact_wording_approval(tmp_path):
    task = task_for(tmp_path)
    task.attach([document(tmp_path / "Spec.docx")])
    task.perform("send", text=REQUEST)
    prompt = task.session.client.requests[0]["system"]
    assert "it is applied only if the user picks it" in prompt
    assert "Never invent project facts" in prompt
    assert "never both" in prompt and "At most three questions per turn" in prompt


def test_proposed_rewrite_requires_wording_choice_then_preview_approval(tmp_path):
    original = "Keep this technical wording intact."
    draft = "Preserve the technical wording."
    request = "Rewrite the final sentence more concisely."
    spec = {"request": request, "summary": "Shorten the final sentence", "mode": "content",
            "ops": [{"op": "replace_text", "select": {"scope": ["document"]},
                     "params": {"find": original, "replace": draft, "tracked": True, "author": "Abraham Borg"}}]}
    fake = FakeClient(resp(tool("ask_user", {"message": "Here is a proposed shorter sentence.",
        "choices": [{"kind": "rephrase", "label": "Use this sentence", "wording": draft}]})),
        resp(tool("submit_editspec", {"spec": spec})))
    task = task_for(tmp_path, fake)
    source = document(tmp_path / "Spec.docx")
    before = source.read_bytes()
    task.attach([source])
    view = task.perform("send", text=request)
    assert view["stage"] == "question" and not view["action"]
    assert draft in "\n".join(c["label"] + " " + " ".join(c["words"]) for c in view["choices"])
    assert draft not in task.word_sources and not task.results
    view = task.perform("send", text="1")
    assert view["stage"] == "preview"
    assert draft in task.word_sources and not task.results
    task.perform("primary", token=view["token"])
    assert task.stage == "done"
    from zipfile import ZipFile
    from lxml import etree
    with ZipFile(task.results[str(source)]) as package:
        root = etree.fromstring(package.read("word/document.xml"))
    inserted = root.xpath("//w:ins//w:t/text()", namespaces={"w":
        "http://schemas.openxmlformats.org/wordprocessingml/2006/main"})
    assert "".join(inserted) == draft
    assert source.read_bytes() == before


def test_an_edited_draft_is_the_users_wording_end_to_end(tmp_path):
    """The choice offers Claude's draft pre-filled and editable; the user's edit is
    the approved wording, the copy carries exactly it, and the draft is never approved. A bad edit is refused before
    anything is said or kept."""
    original = "Keep this technical wording intact."
    draft = "Preserve the technical wording."
    mine = "Preserve the technical wording exactly."
    request = "Rewrite the final sentence more concisely."
    spec = {"request": request, "summary": "Shorten the final sentence", "mode": "content",
            "ops": [{"op": "replace_text", "select": {"scope": ["document"]},
                     "params": {"find": original, "replace": mine, "tracked": True, "author": "Abraham Borg"}}]}
    fake = FakeClient(resp(tool("ask_user", {"message": "Here is a proposed shorter sentence.",
        "choices": [{"kind": "rephrase", "label": "Use this sentence", "wording": draft}]})),
        resp(tool("submit_editspec", {"spec": spec})))
    task = task_for(tmp_path, fake)
    source = document(tmp_path / "Spec.docx")
    task.attach([source])
    view = task.perform("send", text=request)
    choice = view["choices"][0]
    assert choice["editable"] is True and choice["words"] == [draft]
    with pytest.raises(ValueError, match="one line of visible characters"):
        task.perform("choose", choice_id=choice["id"], wording="Two\nlines")
    assert task.stage == "question" and mine not in task.word_sources and len(fake.requests) == 1
    view = task.perform("choose", choice_id=choice["id"], wording=mine)
    assert view["stage"] == "preview", [m["text"] for m in view["messages"]]
    assert mine in task.word_sources and draft not in task.word_sources and mine in task.decisions
    assert any(m["role"] == "user" and m["text"] == f"Chose: Use this sentence — “{mine}”" for m in view["messages"])
    task.perform("primary", token=view["token"])
    from zipfile import ZipFile
    from lxml import etree
    with ZipFile(task.results[str(source)]) as package:
        root = etree.fromstring(package.read("word/document.xml"))
    inserted = root.xpath("//w:ins//w:t/text()", namespaces={"w":
        "http://schemas.openxmlformats.org/wordprocessingml/2006/main"})
    assert "".join(inserted) == mine


def test_recovery_plan_and_preview_stay_in_same_task_without_json(tmp_path, monkeypatch):
    from spec_surgeon import continuous_recovery
    from spec_surgeon.recovery import RecoveryOptions
    refusal = resp(tool("submit_editspec", {"spec": {"request": REQUEST, "summary": "Title",
        "unsupported": "Try another method", "proposal": "Investigate"},
        "custom_transform": "outside_limits", "custom_transform_reason": "Use outcome checks."}))
    fake = FakeClient(refusal, resp(tool("propose_plan", {"contract": CONTRACT,
                      "rationale": "Format the title."})))
    monkeypatch.setattr(continuous_recovery, "options_for", lambda *a, **kw:
                        RecoveryOptions(render=False, critic_review=False, visual_review=False))
    task = task_for(tmp_path, fake)
    source = document(tmp_path / "Spec.docx")
    task.attach([source])
    before = source.read_bytes()
    view = task.perform("send", text=REQUEST)
    # the outcome plan is the user's to pick (F1): the card shows the outcomes in full and the choice that carries
    # them out, with every word they would place beneath it; nothing is tried until the pick
    assert view["stage"] == "question", view
    [choice] = view["choices"]
    assert choice["label"] == "Use these outcomes and continue" and '"exact_text": "FIRE SUPPRESSION"' in choice["detail"]
    assert view["questions"] and '"b": true' in view["questions"][0] and "Format the title." in view["questions"][0]
    assert len(fake.requests) == 2 and not task.results and view["action"] == ""
    view = task.perform("choose", choice_id=choice["id"])
    assert view["stage"] == "preview", view   # one decision remains
    assert any(m["role"] == "user" and m["text"].startswith("Chose: Use these outcomes") for m in task.messages)
    assert view["effects"] and not task.results
    assert not any('"effects"' in message["text"] for message in view["messages"])
    assert view["action"] == "Create edited copies"
    task.perform("primary", token=view["token"])
    assert task.stage == "done" and Path(task.results[str(source)]).is_file()
    assert source.read_bytes() == before


def test_visual_findings_are_visible_and_bound_to_current_action(tmp_path):
    from spec_surgeon.vision import VisualReview
    task = task_for(tmp_path)
    task.attach([document(tmp_path / "Spec.docx")])
    old_view = task.perform("send", text=REQUEST)
    task.session.preview.review = VisualReview(findings=[{"severity": "block", "file": "Spec.docx",
                                                        "page": 2, "text": "The title overlaps the header."}])
    view = task.view()
    assert "reviewed the flagged pages" in view["action"]
    assert "Spec.docx · page 2: The title overlaps the header." in view["warnings"]
    assert task.session.preview.review_ack == ""
    with pytest.raises(ValueError, match="changed"):
        task.perform("primary", token=old_view["token"])
    assert task.session.preview.review_ack == ""


def test_replayed_choice_labels_never_become_approved_document_wording(tmp_path):
    task = task_for(tmp_path)
    task.attach([document(tmp_path / "Spec.docx")])
    task.decisions = [REQUEST, "MODEL LABEL CONTAINING UNAPPROVED WORDS"]
    task.word_sources = [REQUEST]
    task.stage = "resume"
    task.perform("primary", token=task.action_token())
    assert "MODEL LABEL CONTAINING UNAPPROVED WORDS" in task.session.conversation.request
    assert "MODEL LABEL CONTAINING UNAPPROVED WORDS" not in task.session.conversation.sources


def test_later_revision_retains_earlier_copies_and_originals(tmp_path):
    second = copy.deepcopy(SPEC)
    second["ops"][0]["params"]["run"] = {"b": True, "i": True}
    fake = FakeClient(
        resp(tool("submit_editspec", {"spec": SPEC})),
        resp(tool("submit_editspec", {"spec": second})))
    task = task_for(tmp_path, fake)
    source = document(tmp_path / "Spec.docx")
    before = source.read_bytes()
    task.attach([source])
    view = task.perform("send", text=REQUEST)
    task.perform("primary", token=view["token"])
    old_output = Path(task.results[str(source)])
    old_bytes = old_output.read_bytes()
    view = task.perform("send", text="Also make that title italic.")
    task.perform("primary", token=view["token"])
    assert Path(task.results[str(source)]) != old_output
    assert old_output.read_bytes() == old_bytes
    assert source.read_bytes() == before
    run = Document(task.results[str(source)]).paragraphs[1].runs[0]
    assert run.bold and run.italic


def test_custom_method_is_prepared_internally_but_only_delivered_after_review(tmp_path, monkeypatch):
    from test_codegen import GOOD_OP
    request = "Set the header distance to 900 twips."
    fake = FakeClient(resp(tool("submit_editspec", {"spec": {"request": request, "summary": "Move header",
        "unsupported": "Use a custom method", "proposal": "Set the header distance to 900 twips"},
        "custom_transform": "possible", "custom_transform_reason": "Formatting only"})),
        resp(tool("submit_op", {"source": GOOD_OP, "params": {"twips": 900}})),
        resp(tool("ask_user", {"message": "Anything else?"}, "t3")))
    task = task_for(tmp_path, fake)
    source = document(tmp_path / "Spec.docx")
    task.attach([source])
    view = task.perform("send", text=request)
    assert len(fake.requests) == 2
    assert view["stage"] == "preview"
    assert task.session.generated.passed and not task.results
    assert task.session.conversation.batch.inputs == []   # nothing beside the attachments is ever read
    task.perform("primary", token=view["token"])
    assert task.stage == "done" and Path(task.results[str(source)]).is_file()
    assert Document(task.results[str(source)]).sections[0].header_distance.twips == 900
    # the chat thinks at medium, the custom-method writer at xhigh (resources.CODEGEN_EFFORT), and the chat's next
    # call is back at medium; every one of them counted in the window's one ledger
    task.perform("send", text="Thanks.")
    assert [r["output_config"]["effort"] for r in fake.requests] == ["medium", "xhigh", "medium"]
    assert task._budget.calls == 3 and task._budget.effort == "medium"


def test_partial_failure_recovers_automatically_and_can_reopen(tmp_path, monkeypatch):
    from spec_surgeon.recovery import RecoveryOptions
    from spec_surgeon import batch
    from spec_surgeon.verify.suite import CheckResult
    source = document(tmp_path / "First.docx")
    second = document(tmp_path / "Second.docx")
    original = batch._render_and_compare
    injected = False
    def fail_one_render(result, *args, **kwargs):
        nonlocal injected
        original(result, *args, **kwargs)
        if not injected:
            failed = next(f for f in result.files if f.source == second)
            failed.record.set(CheckResult("V6", "Page rendering", "FAIL", "Renderer stopped unexpectedly"))
            injected = True
    monkeypatch.setattr(batch, "_render_and_compare", fail_one_render)
    fake = FakeClient(resp(tool("submit_editspec", {"spec": SPEC})),
        resp(tool("submit_repair", {"candidate": {"kind": "spec", "spec": SPEC}})))
    task = task_for(tmp_path, fake, journal=tmp_path / "journal.json")
    monkeypatch.setattr(task, "_recovery_options", lambda: RecoveryOptions(render=False, critic_review=False,
                                                visual_review=False))
    task.attach([source, second])
    before = [p.read_bytes() for p in (source, second)]
    view = task.perform("send", text=REQUEST)
    view = task.perform("primary", token=view["token"])
    assert len(task.results) == 1
    assert task.session.recovery and view["stage"] == "preview"
    assert view["document_count"] == 1
    restored = task_for(tmp_path, journal=task.journal)
    restored.restore()
    view = restored.perform("primary", token=restored.action_token())
    assert view["stage"] == "preview" and len(restored.results) == 1
    restored.perform("primary", token=view["token"])
    assert len(restored.results) == 2
    assert [p.read_bytes() for p in (source, second)] == before


def test_only_attached_files_are_inspected_not_neighboring_project_data(tmp_path):
    source = document(tmp_path / "Spec.docx")
    (tmp_path / "PROJECT.json").write_text('{"secret": "unrelated project wording"}')
    (tmp_path / "values.csv").write_text("name,value\nsecret,unrelated table wording\n")
    task = task_for(tmp_path)
    task.attach([source])
    task.perform("send", text=REQUEST)
    assert task.session.conversation.batch.inputs == []
    assert "unrelated project wording" not in str(task.session.client.requests)
    assert "unrelated table wording" not in str(task.session.client.requests)
    tools = task.session.client.requests[0]["tools"]
    assert not {"list_library", "read_library_item"}.intersection(t["name"] for t in tools)
    choices = next(t for t in tools if t["name"] == "ask_user")["input_schema"]["properties"]["choices"]["items"]
    assert "recipe" not in choices["properties"]["kind"]["enum"]
    assert "YOUR LIBRARY" not in task.session.client.requests[0]["system"]
    assert "PROJECT MEMORY" not in str(task.session.client.requests[0]["messages"])


@pytest.mark.parametrize("change_source", [False, True])
def test_more_budget_continues_pending_custom_preparation(tmp_path, change_source, monkeypatch):
    from test_codegen import GOOD_OP
    from spec_surgeon.gui import TransformFailed
    request = "Set the header distance to 900 twips."
    fake = FakeClient(resp(tool("submit_editspec", {"spec": {"request": request, "summary": "Move header",
        "unsupported": "Use a custom method", "proposal": request},
        "custom_transform": "possible", "custom_transform_reason": "Formatting only"})),
        resp(tool("submit_op", {"source": GOOD_OP, "params": {"twips": 900}})))
    task = task_for(tmp_path, fake)
    source = document(tmp_path / "Spec.docx")
    task.attach([source])
    check_after(task, monkeypatch, 1)
    with pytest.raises(TransformFailed):
        task.perform("send", text=request)
    assert task.stage == "budget" and task.session.pending_custom_transform
    conversation = task.session.conversation
    if change_source:
        source.write_bytes(source.read_bytes() + b"changed")
        with pytest.raises(ValueError, match="original files changed"):
            task.perform("primary", token=task.action_token())
        assert len(fake.requests) == 1 and not task.results
        return
    view = task.perform("primary", token=task.action_token())
    assert view["stage"] == "preview" and not task.results
    assert task.session.conversation is conversation
    assert task.session.generated.passed and task.session.pending_custom_transform is None
    assert len(fake.requests) == task._budget.calls == 2
    task.perform("primary", token=view["token"])
    assert Document(task.results[str(source)]).sections[0].header_distance.twips == 900


@pytest.mark.parametrize("reopen", [False, True])
def test_more_budget_resumes_unadopted_recovery_checkpoint(tmp_path, reopen, monkeypatch):
    from spec_surgeon import llm, recovery
    from test_recovery import candidate, margin, opts, spec, submit
    source = document(tmp_path / "Spec.docx")
    goal = spec([margin(), margin("right", 2880)])
    chain = {"kind": "chain", "chain": {"stages": [{"spec": spec([op])} for op in goal["ops"]]}}
    fake = FakeClient(submit(candidate(goal)), submit(chain))
    task = task_for(tmp_path, fake, journal=tmp_path / "journal.json")
    task.attach([source])
    task.decisions = task.word_sources = [goal["request"]]
    # Model execution pauses after a failed candidate has been recorded. The coordinator
    # has the checkpoint path, as it does when interrupted before adopting the result.
    checkpoint = tmp_path / "recovery"
    task.session.recovery_checkpoint = checkpoint
    one_call = llm.CallBudget(max_calls=1, max_cost_usd=None, max_total_tokens=None)
    with llm.budget_scope(task._budget), llm.budget_scope(one_call):
        paused = recovery.recover([source], goal["request"], checkpoint, spec=goal,
            options=opts(critic_review=False, visual_review=False), client=fake)
    assert paused.paused_reason and len(paused.attempts) == 1
    task.stage = "budget"
    task.save()
    old_cost = task._budget.estimated_cost_usd
    if reopen:
        task = task_for(tmp_path, fake, journal=task.journal)
        task.restore()
    monkeypatch.setattr(task, "_recovery_options", lambda:
                        opts(critic_review=False, visual_review=False))
    view = task.perform("primary", token=task.action_token())
    assert view["stage"] == "preview" and not task.results
    assert task.session.recovery.out_dir == checkpoint
    assert len(task.session.recovery.attempts) == 2
    assert len(fake.requests) == task._budget.calls == 2
    assert task._budget.estimated_cost_usd >= old_cost
    assert task.session.recovery.resource_history[0]["calls"] == 1
    task.perform("primary", token=view["token"])
    assert Path(task.results[str(source)]).is_file()


def test_recovery_resume_refuses_changed_checkpoint_sources(tmp_path):
    from spec_surgeon import recovery
    from test_recovery import margin, opts, spec
    source = document(tmp_path / "Spec.docx")
    task = task_for(tmp_path, FakeClient())
    task.attach([source])
    goal = spec([margin()])
    ready = recovery.prepare_recovery([source], goal["request"], tmp_path / "recovery", spec=goal, options=opts())
    task.session.recovery_checkpoint = ready.out_dir
    task.stage = "budget"
    source.write_bytes(source.read_bytes() + b"changed")
    with pytest.raises(ValueError, match="changed"):
        task.perform("primary", token=task.action_token())
    assert not task.session.client.requests and not task.results


@pytest.mark.parametrize("reopen", [False, True])
def test_partial_recovery_can_export_again_after_an_answer(tmp_path, reopen, monkeypatch):
    from spec_surgeon import recovery
    from test_recovery import candidate, margin, opts, spec, submit
    from test_recovery_resume import distinct_failures
    paths = [document(tmp_path / name) for name in ("First.docx", "Second.docx")]
    original = [p.read_bytes() for p in paths]
    goal = spec([margin()])
    fake = FakeClient(submit(candidate(goal)),
                      resp(tool("need_user", {"question": "Use these margins on the second document?"})),
                      submit(candidate(goal)))
    task = task_for(tmp_path, fake, journal=tmp_path / "journal.json")
    task.attach(paths)
    task.decisions = task.word_sources = [goal["request"]]
    result = recovery.recover(paths, goal["request"], tmp_path / "recovery", spec=goal,
                              options=opts(),
                              prior_failures=distinct_failures(paths), client=fake)
    assert len(result.ready) == 1 and result.questions
    task.session._adopt_recovery(result)
    task._settle(lambda _: None)
    task.perform("primary", token=task.action_token())
    first_output = Path(task.results[str(paths[0])])
    first_bytes = first_output.read_bytes()
    assert task.stage == "question" and len(task.results) == 1
    if reopen:
        task = task_for(tmp_path, fake, journal=task.journal)
        task.restore()
        task.perform("primary", token=task.action_token())
    monkeypatch.setattr(task, "_recovery_options", lambda:
                        opts(critic_review=False, visual_review=False))
    view = task.perform("send", text="Yes, use those margins on the second document.")
    assert view["stage"] == "preview" and len(task.results) == 1
    task.perform("primary", token=view["token"])
    assert task.stage == "done" and len(task.results) == 2
    assert first_output.read_bytes() == first_bytes
    assert Path(task.results[str(paths[1])]).parent != first_output.parent
    assert all(Path(output).is_file() for output in task.results.values())
    assert [p.read_bytes() for p in paths] == original
    assert len(fake.requests) == 3


def test_a_saved_copys_preview_leaves_while_the_rest_are_still_discussed(tmp_path):
    """A partial save delivers the ready copies; their pages and the visual check's findings about them leave the
    window, and the conversation about the other documents goes on without them. Neither comes back for a document
    whose copy is saved, across a resume; a finding about another document, or about none, stays."""
    from spec_surgeon import recovery, vision
    from test_recovery import candidate, margin, opts, spec, submit
    from test_recovery_resume import distinct_failures
    paths = [document(tmp_path / name) for name in ("First.docx", "Second.docx")]
    goal = spec([margin()])
    fake = FakeClient(submit(candidate(goal)),
                      resp(tool("need_user", {"question": "Use these margins on the second document?"})),
                      submit(candidate(goal)))
    task = task_for(tmp_path, fake)
    task.attach(paths)
    task.decisions = task.word_sources = [goal["request"]]
    result = recovery.recover(paths, goal["request"], tmp_path / "recovery", spec=goal, options=opts(),
                              prior_failures=distinct_failures(paths), client=fake)
    result.review = vision.VisualReview(findings=[
        {"severity": "block", "file": "First.docx", "page": 1, "text": "The header overlaps the title."},
        {"severity": "warn", "file": "Second.docx", "page": None, "text": "The footer is close to the edge."},
        {"severity": "warn", "file": "", "page": None, "text": "Fonts were substituted while rendering."}])
    result._reviews = {str(paths[0]): result.review}        # per file, as recovery keeps the check of each trial
    task.session._adopt_recovery(result)
    task._settle(lambda _: None)
    view = task.view()
    assert view["action"] == "Save 1 ready copy" and view["blocking"]
    assert [Path(sv.source).name for sv in view["samples"]] == ["First.docx"]
    assert any("header overlaps" in w for w in view["warnings"])
    view = task.perform("primary", token=view["token"])
    assert view["stage"] == "question" and list(task.results) == [str(paths[0])]
    assert any(sv.source == paths[0] for sv in task.session.recovery.samples)    # the evidence is kept
    assert view["samples"] == []                                                  # the window shows none of it
    assert not view["blocking"] and not any("header overlaps" in w for w in view["warnings"])
    assert any("footer is close" in w for w in view["warnings"])
    assert any("substituted" in w for w in view["warnings"])
    task._recovery_options = lambda: opts(critic_review=False, visual_review=False)
    view = task.perform("send", text="Yes, use those margins on the second document.")
    assert view["stage"] == "preview"
    assert [Path(sv.source).name for sv in task.session.recovery.samples] == ["First.docx", "Second.docx"]
    assert [Path(sv.source).name for sv in view["samples"]] == ["Second.docx"]
    assert task.session.outcome().blocking                                        # the finding is still on record
    assert not view["blocking"] and view["action"] == "Create edited copies"
    assert not any("header overlaps" in w for w in view["warnings"])


def test_a_mixed_batch_gets_one_plan_per_layout_and_one_decision(tmp_path):
    from make_fixtures import build, build_styles, build_rich
    from spec_surgeon.fingerprint import fingerprint
    from spec_surgeon.package import DocxPackage
    typed = tmp_path / "a" / "21 13 13.docx"
    styled = tmp_path / "b" / "21 13 16.docx"
    other = tmp_path / "c" / "21 13 19.docx"
    for path, data in ((typed, build("21 13 13", "WET-PIPE SPRINKLER SYSTEMS")),
                       (styled, build_styles("21 13 16", "DRY-PIPE SPRINKLER SYSTEMS")),
                       (other, build_rich("21 13 19", "PREACTION SPRINKLER SYSTEMS"))):
        path.parent.mkdir(parents=True)
        path.write_bytes(data)
    pid = lambda p: fingerprint(DocxPackage.from_path(p))["pattern_id"]  # noqa: E731
    bold = {"run": {"b": True}, "tracked": False}
    submission = {"summary": "Bold body paragraphs in two layouts", "layouts": [
        {"id": "A", "label": "typed numbers, no styles", "pattern_ids": [pid(typed)],
         "spec": {"summary": "by text", "mode": "formatting", "assumptions": ["Body paragraphs start with A. to E."],
                  "ops": [{"op": "set_properties", "select": {"paragraphs": {"text_regex": r"^[A-E]\.\t"}}, "params": bold}]}},
        {"id": "B", "label": "styled body", "pattern_ids": [pid(styled), pid(other)],
         "spec": {"summary": "by style", "mode": "formatting",
                  "ops": [{"op": "set_properties", "select": {"paragraphs": {"style_in": ["ArchPara", "Normal"]}}, "params": bold}]}}]}
    fake = FakeClient(resp(tool("submit_editspec", {"spec": submission, "message": "I’ll bold the body paragraphs; your documents use two layouts."})))
    task = task_for(tmp_path, fake)
    task.attach([typed, styled, other])
    before = [p.read_bytes() for p in (typed, styled, other)]
    view = task.perform("send", text="Make the body paragraphs bold.")
    assert view["stage"] == "preview" and view["action"] == "Create edited copies"
    assert [m["text"] for m in view["messages"] if m["role"] == "assistant"] == [
        "I’ll bold the body paragraphs; your documents use two layouts."]
    groups = view["groups"]
    assert [(g["id"], g["label"], g["count"]) for g in groups] == [("A", "typed numbers, no styles", 1), ("B", "styled body", 2)]
    assert groups[0]["assumptions"] == ["Body paragraphs start with A. to E."]
    assert all(change["summary"].endswith(".") for g in groups for change in g["changes"])
    assert len(view["samples"]) == 3 and not view["set_aside"]   # one sample per structure pattern
    task.perform("primary", token=view["token"])
    assert task.stage == "done" and len(task.results) == 3
    assert [p.read_bytes() for p in (typed, styled, other)] == before
    assert len(fake.requests) == 1


def _layout_submission(tmp_path):
    from make_fixtures import build, build_rich, build_styles
    from spec_surgeon.fingerprint import fingerprint
    from spec_surgeon.package import DocxPackage
    typed, styled, other = tmp_path / "a" / "21 13 13.docx", tmp_path / "b" / "21 13 16.docx", tmp_path / "c" / "21 13 19.docx"
    for path, data in ((typed, build("21 13 13", "WET-PIPE SPRINKLER SYSTEMS")),
                       (styled, build_styles("21 13 16", "DRY-PIPE SPRINKLER SYSTEMS")),
                       (other, build_rich("21 13 19", "PREACTION SPRINKLER SYSTEMS"))):
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(data)
    pid = lambda p: fingerprint(DocxPackage.from_path(p))["pattern_id"]  # noqa: E731
    bold = {"run": {"b": True}, "tracked": False}
    submission = {"summary": "Bold body paragraphs in two layouts", "layouts": [
        {"id": "A", "label": "typed numbers", "pattern_ids": [pid(typed)],
         "spec": {"summary": "by text", "mode": "formatting",
                  "ops": [{"op": "set_properties", "select": {"paragraphs": {"text_regex": r"^[A-E]\.\t"}}, "params": bold}]}},
        {"id": "B", "label": "styled body", "pattern_ids": [pid(styled)],
         "spec": {"summary": "by style", "mode": "formatting",
                  "ops": [{"op": "set_properties", "select": {"paragraphs": {"style_in": ["ArchPara"]}}, "params": bold}]}}]}
    return submission, [typed, styled, other]


def test_after_one_layouts_recovery_is_saved_the_next_layouts_documents_are_checked(tmp_path, monkeypatch):
    """Both layouts set their document aside at run time: recovery takes layout A first; once its copies are saved
    the coordinator moves on to layout B by itself, and the queue survives a reopened conversation. (The compiler
    refuses a plan that fails on its own documents, so the plan is set as the outcome directly.)"""
    from spec_surgeon import recovery
    submission, docs = _layout_submission(tmp_path)
    plan = {"kind": "layouts", "request": "Make the body paragraphs bold.", "summary": submission["summary"],
            "layouts": submission["layouts"]}
    for layout in plan["layouts"]:
        layout["spec"]["request"] = plan["request"]
        layout["spec"]["ops"][0]["expect"] = {"text_diffs": 5}
    calls = []

    def fake_recover(run_dir, out_dir, *, request=None, options=None, client=None, log=None):
        calls.append(Path(run_dir).name)
        out = Path(out_dir)
        (out / "outputs").mkdir(parents=True, exist_ok=True)   # a delivery is not part of this test
        return recovery.RecoveryResult(out_dir=out, request=str(request or ""))
    monkeypatch.setattr(recovery, "recover_from_run", fake_recover)
    journal = tmp_path / "journal.json"
    task = task_for(tmp_path, journal=journal)
    task.attach(docs[:2])
    task.session._set_outcome(spec=plan, source="test")
    task._settle(lambda m: None)
    assert task.stage == "preview"
    view = task.view()
    assert [item["name"] for item in view["set_aside"]] == ["21 13 13.docx", "21 13 16.docx"]
    view = task.perform("primary", token=view["token"])
    run = task.session.last_run_dir
    assert calls == ["layout-A"] and task.session.pending_layout_recovery == [run / "layout-B"]
    assert view["stage"] == "preview"                                # layout A's recovery awaits its review
    assert json.loads(journal.read_text(encoding="utf-8"))["pending_layouts"] == [str(run / "layout-B")]
    reopened = task_for(tmp_path, journal=journal)                   # a reopened conversation keeps the queue
    assert reopened.restore() and reopened._resume_layouts == [str(run / "layout-B")]
    # the empty stub recovery has nothing to export; the coordinator's own step to the next layout is the subject
    task.session.last_recovery, task.session.recovery = task.session.recovery, None
    task.session.resolve_next_layout(lambda m: None)
    assert calls == ["layout-A", "layout-B"] and task.session.pending_layout_recovery == []


# ---- no allowance: the pause, the stop and Continue ---------------------------------------------------------------

def _tool_results(messages):
    return [b["tool_use_id"] for m in messages if m["role"] == "user" and isinstance(m["content"], list)
            for b in m["content"] if isinstance(b, dict) and b.get("type") == "tool_result"]


def test_continue_after_the_pause_picks_the_conversation_up_where_it_stopped(tmp_path, monkeypatch):
    fake = FakeClient(resp(tool("list_styles", {"file": "Spec.docx"}, "t1")),
                      resp(tool("list_styles", {"file": "Spec.docx"}, "t2")),
                      resp(tool("submit_editspec", {"spec": SPEC}, "t3")))
    task = task_for(tmp_path, fake)
    task.attach([document(tmp_path / "Spec.docx")])
    limit = check_after(task, monkeypatch, 2)
    view = task.perform("send", text=REQUEST)
    assert view["stage"] == "budget" and view["action"] == "Continue" and len(fake.requests) == 2
    conversation = task.session.conversation
    view = task.perform("primary", token=view["token"])
    assert view["stage"] == "preview" and len(fake.requests) == 3
    assert task.session.conversation is conversation                     # the same investigation, not a new one
    third = fake.requests[2]["messages"]
    assert third[0] == fake.requests[0]["messages"][0] and _tool_results(third) == ["t1", "t2"]
    assert task._request_budget.max_cost_usd == pytest.approx(0.0016 + limit)


def test_continue_after_a_stop_sends_the_stopped_turn_again_as_it_was(tmp_path):
    from fake_claude import raise_on_enter
    from spec_surgeon import llm
    fake = FakeClient(resp(tool("list_styles", {"file": "Spec.docx"}, "t1")),
                      raise_on_enter(llm.UserCancelled(llm.STOPPED)),
                      resp(tool("submit_editspec", {"spec": SPEC}, "t2")))
    task = task_for(tmp_path, fake)
    task.attach([document(tmp_path / "Spec.docx")])
    with pytest.raises(llm.UserCancelled):
        task.perform("send", text=REQUEST)
    assert task.stage == "resume" and task.view()["action"] == "Continue"
    conversation = task.session.conversation
    assert conversation.resumable()
    view = task.perform("primary", token=task.view()["token"])
    assert view["stage"] == "preview" and task.session.conversation is conversation
    assert fake.requests[2]["messages"] == fake.requests[1]["messages"]
    assert not conversation.resumable()                                   # a spec now waits for the user


def test_a_conversation_waiting_for_the_users_answer_is_not_resumed(tmp_path):
    fake = FakeClient(resp(tool("ask_user", {"questions": ["Which title?"]})))
    task = task_for(tmp_path, fake)
    task.attach([document(tmp_path / "Spec.docx")])
    task.perform("send", text=REQUEST)
    assert task.stage == "question" and not task.session.conversation.resumable()
    with pytest.raises(ValueError, match="waiting for the user"):
        task.session.conversation.resume()


# ---- pictures and tables attached for the edit to use -------------------------------------------------------------

def test_pictures_and_tables_attach_for_the_edit_to_use_and_reach_the_conversation(tmp_path):
    from make_fixtures import png
    from test_text_edits import _xlsx
    logo, sheet, broken = tmp_path / "logo.png", tmp_path / "PROJECT.xlsx", tmp_path / "broken.png"
    logo.write_bytes(png(size=6))
    sheet.write_bytes(_xlsx([["Variable", "VA3"], ["SITE_NAME", "Leesburg"]]))
    broken.write_bytes(b"\x89PNG\r\n\x1a\n" + b"\x00" * 10)
    journal = tmp_path / "conversation.json"
    task = task_for(tmp_path, journal=journal)
    spec_file = document(tmp_path / "Spec.docx")
    task.attach([spec_file, logo, sheet, broken])
    refs = [a for a in task.view()["attachments"] if a["reference"]]
    assert [Path(a["path"]).name for a in refs] == ["logo.png", "PROJECT.xlsx", "broken.png"]
    assert refs[0]["state"] == "note" and "6 × 6 pixels" in refs[0]["notes"][0]
    assert refs[1]["state"] == "note" and "1 sheet" in refs[1]["notes"][0]
    assert refs[2]["state"] == "refused" and refs[2]["reason"].startswith("broken.png cannot be read")
    assert task.documents() == [spec_file.resolve()]
    assert task.session.inputs == [logo.resolve(), sheet.resolve()]
    (tmp_path / "notes.pdf").write_bytes(b"%PDF-1.7")
    with pytest.raises(ValueError, match="pictures"):
        task.attach([tmp_path / "notes.pdf"])
    with pytest.raises(ValueError, match="Word .docx"):
        task.attach([logo], template=True)
    task.perform("send", text=REQUEST)
    first = task.session.client.requests[0]["messages"][0]["content"]
    first = first if isinstance(first, str) else json.dumps(first, ensure_ascii=False)
    assert "INPUT FILES" in first and "logo.png" in first and "PROJECT.xlsx" in first and "broken.png" not in first
    conversation = task.session.conversation
    assert conversation.batch.inputs == [logo.resolve(), sheet.resolve()]
    assert task.session.conversation_matches(conversation.request)
    restored = task_for(tmp_path, journal=journal)
    assert restored.restore()
    assert restored.session.inputs == [logo.resolve(), sheet.resolve()]
    task.remove(logo.resolve())
    assert task.session.inputs == [sheet.resolve()] and task.session.conversation is None


def test_a_table_attached_in_the_window_fills_the_documents_end_to_end(tmp_path):
    import zipfile
    from make_fixtures import build_batch1
    from test_text_edits import _xlsx
    doc, sheet = tmp_path / "211319 - Preaction.docx", tmp_path / "PROJECT.xlsx"
    doc.write_bytes(build_batch1())
    sheet.write_bytes(_xlsx([["Variable", "VA1", "VA3"], ["SITE_NAME", "Ashburn", "Leesburg"],
                             ["OWNER_NAME", "Contoso", "Contoso"], ["PROJECT NO", "1001", "1003"]]))
    request = "Fill in the VA3 site values from PROJECT.xlsx."
    spec = {"request": request, "summary": "Fill the project values", "mode": "content", "ops": [
        {"op": "fill_variables", "params": {"table": {"file": str(sheet.resolve()), "name_column": "Variable",
                                                      "value_column": "VA3"}, "author": "A"}}]}
    fake = FakeClient(resp(tool("read_input_file", {"path": str(sheet.resolve())}, "r1")),
                      resp(tool("submit_editspec", {"spec": spec}, "s1")))
    task = task_for(tmp_path, fake)
    task.attach([doc, sheet])
    view = task.perform("send", text=request)
    assert view["stage"] == "preview", (view["messages"], view["warnings"])
    read = fake.requests[1]["messages"][-1]["content"][0]
    assert read["tool_use_id"] == "r1" and not read.get("is_error") and "Leesburg" in read["content"]
    view = task.perform("primary", token=view["token"])
    assert view["stage"] == "done" and len(task.results) == 1
    with zipfile.ZipFile(next(iter(task.results.values()))) as z:
        headers = "".join(z.read(n).decode("utf-8") for n in z.namelist() if n.startswith("word/header"))
    # filled as a tracked change: the placeholder deleted, the sheet's value inserted
    assert "<w:t>Leesburg</w:t>" in headers and "<w:delText>{{SITE_NAME}}</w:delText>" in headers
    assert "<w:t>{{SITE_NAME}}</w:t>" not in headers


def test_a_reply_claude_declines_leaves_the_window_asking_in_plain_words(tmp_path):
    """The 1.8.0 live run showed the provider's policy text for a declined reply. The window says what happened and
    waits for the next reply, which goes out without the declined one."""
    from types import SimpleNamespace
    task = task_for(tmp_path)
    fake = FakeClient(resp(tool("ask_user", {"questions": ["Which title?"]}, "a1")),
                      resp(stop="refusal", details=SimpleNamespace(category="reasoning_extraction",
                                                                   explanation="This request was blocked")),
                      resp(tool("ask_user", {"questions": ["Anything else?"]}, "a2")))
    task.session.client = fake
    task.attach([document(tmp_path / "Spec.docx")])
    task.perform("send", text=REQUEST)
    task.perform("send", text="what were you thinking")
    assert task.stage == "question"
    said = task.messages[-1]["text"]
    assert "taken back out of the conversation" in said and "blocked" not in said
    task.perform("send", text="FIRE SUPPRESSION")
    sent = json.dumps(fake.requests[-1]["messages"], default=str)
    assert task.stage == "question" and "FIRE SUPPRESSION" in sent and "what were you thinking" not in sent


def test_two_attached_documents_of_one_name_are_named_with_their_folders(tmp_path):
    """The 1.8.0 live run attached a template's folder copy of 21 13 13 as a 13th document, beside the batch's own
    21 13 13: nothing said so, and it was planned and edited like the others."""
    first = document(tmp_path / "test-set-1" / "21 13 13 Wet-Pipe Sprinkler Systems.docx")
    task = task_for(tmp_path)
    task.attach([first, document(tmp_path / "test-set-1" / "21 12 00 Standpipes.docx")])
    assert not any("are named" in m["text"] for m in task.messages)
    task.attach([document(tmp_path / "template" / "21 13 13 Wet-Pipe Sprinkler Systems.docx")])
    [said] = [m["text"] for m in task.messages if "are named" in m["text"]]
    assert "“test-set-1”" in said and "“template”" in said and "attach it as the template" in said
    task.attach([document(tmp_path / "other.docx")], template=True)       # the template itself is not a document
    assert sum("are named" in m["text"] for m in task.messages) == 1


def test_an_attach_that_raises_on_one_file_attaches_none_of_them(tmp_path, monkeypatch):
    """The stress test's F6: a valid document attached together with a malformed one in one request was added while
    the request failed, leaving the session's documents out of step with the task's attachments. Every file is
    triaged before any is attached, so a triage that raises leaves the task exactly as it was: the attachments, the
    documents, the fresh list (here, over a working set), the pass and the view. A refusal is not an exception: that
    file is attached, refused, with its reason."""
    from spec_surgeon import document_task
    from spec_surgeon.document_task import Attachment
    task = task_for(tmp_path)
    task.attach([document(tmp_path / "First.docx")])
    view = task.perform("send", text=REQUEST)
    task.perform("primary", token=view["token"])
    assert task.working and task.fresh == []
    good, bad = document(tmp_path / "Good.docx"), document(tmp_path / "Bad.docx")
    real = document_task.triage

    def triage(path, **kwargs):
        if Path(path).name == "Bad.docx":
            raise RuntimeError("the stress test's malformed zip")
        if Path(path).name == "Refused.docx":
            return Attachment(Path(path), state="refused", reason="Refused.docx is not a Word document.")
        return real(path, **kwargs)

    monkeypatch.setattr(document_task, "triage", triage)

    def state():
        return ([a.to_dict() for a in task.attachments], task.documents(), list(task.fresh), task.focus,
                task.session.documents, task.session.inputs, task.view())

    before, before_attachments = state(), list(task.attachments)
    with pytest.raises(RuntimeError, match="malformed zip"):
        task.attach([good, bad])
    assert state() == before
    assert good not in task.documents() and str(good) not in task.fresh and str(good) not in task.working
    refused = document(tmp_path / "Refused.docx")
    task.attach([good, refused])
    assert task.documents() == [*before[1], good] and task.session.documents == task.documents()
    assert [a.path for a in task.attachments] == [*(a.path for a in before_attachments), good, refused]
    assert task.fresh == [str(good), str(refused)]
    [found] = [a for a in task.attachments if a.path == refused]
    assert found.refused and found.reason == "Refused.docx is not a Word document."


def test_continue_after_the_spending_check_takes_the_custom_method_up_where_it_stopped(tmp_path, monkeypatch):
    """Live QA finding F09: a writer paused at the spending check went on from its brief alone in a new folder and
    repeated its inspections. Now it continues from its checkpoint, in its own folder, with what it has read."""
    from test_codegen import GOOD_OP, UNASSERTED_OP
    from spec_surgeon.gui import TransformFailed
    request = "Set the header distance to 900 twips."
    source = document(tmp_path / "Spec.docx")
    fake = FakeClient(resp(tool("submit_editspec", {"spec": {"request": request, "summary": "Move header",
        "unsupported": "Use a custom method", "proposal": request},
        "custom_transform": "possible", "custom_transform_reason": "Formatting only"})),
        resp(tool("show_xml", {"file": source.name, "start": 0, "count": 2}, "x1")),
        resp(tool("submit_op", {"source": UNASSERTED_OP, "params": {"twips": 900}}, "a1")),
        resp(tool("submit_op", {"source": GOOD_OP, "params": {"twips": 900}}, "a2")))
    task = task_for(tmp_path, fake)
    task.attach([source])
    check_after(task, monkeypatch, 3)       # the plan, one inspection, one failed attempt, then the check stops it
    with pytest.raises(TransformFailed):    # a pause, never a failure for the outcome investigation to take up
        task.perform("send", text=request)
    pending = task.session.pending_custom_transform
    assert task.stage == "budget" and task.pause_reason() == "limit" and pending["attempts"] == 1
    writer = Path(pending["paused"])
    assert writer == Path(pending["out"]) / "codegen" and writer.is_dir()
    paused = json.loads((writer / "history.json").read_text(encoding="utf-8"))
    assert paused["budget_stop"] is True and paused["calls"] == 2 and len(paused["attempts"]) == 1
    view = task.perform("primary", token=task.action_token())
    assert view["stage"] == "preview" and task.session.generated.passed and task.session.pending_custom_transform is None
    said = [m["text"] for m in task.messages]
    assert any(t.startswith("Continuing: this request may spend up to $") for t in said)
    assert "The custom method continues where it stopped, with everything it had looked at." in said
    # one writer folder, its history whole; the request after the pause carries the inspection, its answer and the
    # failed attempt's verdict
    assert sorted(p.name for p in Path(pending["out"]).glob("codegen*")) == ["codegen"]
    resumed = fake.requests[3]["messages"]
    assert len(resumed) == 5 and resumed[1]["content"][0]["name"] == "show_xml"
    assert resumed[2]["content"][0]["tool_use_id"] == "x1" and "<w:p" in resumed[2]["content"][0]["content"]
    assert resumed[4]["content"][0]["tool_use_id"] == "a1" and "V10 FAIL" in resumed[4]["content"][0]["content"]
    assert fake.requests[3]["system"] == fake.requests[2]["system"]
    history = json.loads((writer / "history.json").read_text(encoding="utf-8"))
    assert history["passed"] is True and history["calls"] == 3 and history["budget_stop"] is False
    assert [a["attempt"] for a in history["attempts"]] == [1, 2] and len(task.session.generated.attempts) == 2
    assert len(fake.requests) == task._budget.calls == 4
    task.perform("primary", token=view["token"])
    assert Document(task.results[str(source)]).sections[0].header_distance.twips == 900


def test_a_failed_custom_method_attempt_starts_the_next_afresh(tmp_path, monkeypatch):
    """A writer that failed (not a pause) is a spent attempt: the next one starts from the brief in codegen-02."""
    from test_codegen import GOOD_OP
    from spec_surgeon.gui import TransformFailed
    request = "Set the header distance to 900 twips."
    source = document(tmp_path / "Spec.docx")
    fake = FakeClient(resp(tool("submit_editspec", {"spec": {"request": request, "summary": "Move header",
        "unsupported": "Use a custom method", "proposal": request},
        "custom_transform": "possible", "custom_transform_reason": "Formatting only"})),
        resp(tool("show_xml", {"file": source.name, "start": 0, "count": 2}, "x1")),
        raise_on_enter(api_error("Overloaded", 529, "req_1")),
        resp(tool("submit_op", {"source": GOOD_OP, "params": {"twips": 900}}, "a1")))
    task = task_for(tmp_path, fake)
    task.attach([source])
    with pytest.raises(TransformFailed, match="Overloaded"):
        task.perform("send", text=request)
    pending = task.session.pending_custom_transform
    assert pending["attempts"] == 1 and pending["paused"] is None and task.stage != "budget"
    view = task.perform("primary", token=task.action_token())
    assert view["stage"] == "preview" and task.session.generated.passed and pending["attempts"] == 2
    assert sorted(p.name for p in Path(pending["out"]).glob("codegen*")) == ["codegen", "codegen-02"]
    assert len(fake.requests[3]["messages"]) == 1                       # the brief alone: nothing carried over
    assert len(fake.requests) == 4
