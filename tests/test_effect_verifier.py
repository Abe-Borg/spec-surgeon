"""Adversarial original-goal checks independent of the candidate method's plan."""
from __future__ import annotations

import copy

import pytest
from docx import Document
from lxml import etree

from spec_surgeon.goals import resolve_goal
from spec_surgeon.verify.effects import verify_effects
from spec_surgeon.verify.reader import REL, W
from test_goal_contracts import contract, document, exact, execute, template_effect
from test_goal_verifier import mutate


def bound_execute(source, goal):
    binding = resolve_goal(source, goal)["binding"]
    return execute(source, goal, binding), binding


def paragraphs(root):
    return list(root.iter(W + "p"))


def text(paragraph):
    return "".join(t.text or "" for t in paragraph.iter(W + "t"))


def paragraph(root, prefix):
    return next(p for p in paragraphs(root) if text(p).startswith(prefix))


def rejects(source, candidate, goal, binding):
    result = verify_effects(source, candidate, goal, binding)
    assert not result["passed"], result
    assert result["errors"] and any(check["status"] == "FAIL" for check in result["checks"])
    return result


@pytest.fixture
def compound(tmp_path):
    source = document(tmp_path / "spec.docx", "table")
    doc = Document(source)
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Preserve table\tALL WORDS"
    doc.save(source)
    goal = contract(
        {"id": "font", "type": "format", "target": {"all": True}, "font": {"name": "Arial", "size": 11}},
        {"id": "title", "type": "format", "target": {"locators": [{"role": "section_title"}]},
         "run": {"b": True}, "paragraph": {"keepNext": True}},
        {"id": "wording", "type": "replace_text", "target": exact("Owner shall retain all content."),
         "find": "Owner", "replace": "Client", "expected_matches": 1, "tracked": True},
        {"id": "page", "type": "page_setup", "margins": {"left": 1440}})
    output, binding = bound_execute(source, goal)
    return source, output, goal, binding


@pytest.mark.parametrize("damage", ["font", "size", "italic", "title_bold", "title_cs_bold", "mark_bold",
                                    "table_tab", "table_words", "section", "page", "paragraph_attr"])
def test_compound_effects_do_not_hide_collateral_damage(compound, damage):
    source, output, goal, binding = compound

    def change(root):
        title = paragraph(root, "FIRE SUPPRESSION")
        rpr = title.find(W + "r/" + W + "rPr")
        if damage == "font":
            rpr.find(W + "rFonts").set(W + "eastAsia", "Wrong font")
        elif damage == "size":
            rpr.find(W + "sz").set(W + "val", "24")
        elif damage == "italic":
            rpr.find(W + "i").set(W + "val", "0")
        elif damage in {"title_bold", "title_cs_bold"}:
            rpr.find(W + ("b" if damage == "title_bold" else "bCs")).set(W + "val", "0")
        elif damage == "mark_bold":
            title.find(W + "pPr/" + W + "rPr/" + W + "b").set(W + "val", "0")
        elif damage == "table_tab":
            tab = paragraph(root, "Preserve table").find(".//" + W + "tab")
            tab.getparent().remove(tab)
        elif damage == "table_words":
            paragraph(root, "Preserve table").find(".//" + W + "t").text = "Lost words"
        elif damage == "section":
            paragraph(root, "SECTION ").find(".//" + W + "t").text = "SECTION 22 11 16"
        elif damage == "page":
            root.find(".//" + W + "pgMar").set(W + "right", "33")
        else:
            paragraph(root, "Client").set(W + "rsidR", "DEADBEEF")

    rejects(source, mutate(output, change), goal, binding)


def test_effective_format_can_use_inherited_size_instead_of_builtin_direct_size(compound):
    source, output, goal, binding = compound

    def alternate(root):
        rpr = paragraph(root, "FIRE SUPPRESSION").find(W + "r/" + W + "rPr")
        rpr.remove(rpr.find(W + "sz"))  # Normal already inherits the requested eleven points.

    result = verify_effects(source, mutate(output, alternate), goal, binding)
    assert result["passed"], result


@pytest.mark.parametrize("damage", ["deleted_word", "author", "id", "duplicate_id", "new_words", "italic",
                                    "empty_revision", "space", "revision_attr", "second_rpr", "second_ppr", "text_attr"])
def test_final_visible_wording_does_not_excuse_broken_redline_provenance(compound, damage):
    source, output, goal, binding = compound

    def change(root):
        changed = paragraph(root, "Client")
        if damage == "deleted_word":
            changed.find(".//" + W + "delText").text = "Original wording lost"
        elif damage == "author":
            changed.find(W + "ins").set(W + "author", "Unapproved author")
        elif damage == "id":
            changed.find(W + "ins").set(W + "id", "nonnumeric")
        elif damage == "duplicate_id":
            changed.find(W + "ins").set(W + "id", changed.find(W + "del").get(W + "id"))
        elif damage == "new_words":
            changed.find(W + "ins/" + W + "r/" + W + "t").text = "Invented obligation"
        elif damage == "italic":
            unchanged = next(r for r in changed.iter(W + "r") if "shall" in text(r))
            unchanged.find(W + "rPr/" + W + "i").set(W + "val", "0")
        elif damage == "empty_revision":
            extra = etree.SubElement(changed, W + "ins")
            for key, value in {"id": "999", "author": "Abraham Borg", "date": "2026-01-01T00:00:00Z"}.items():
                extra.set(W + key, value)
        elif damage == "space":
            node = next(node for node in changed.iter(W + "t") if (node.text or "").startswith(" shall"))
            node.attrib.pop("{http://www.w3.org/XML/1998/namespace}space", None)
        elif damage == "revision_attr":
            changed.find(W + "ins").set(W + "unapproved", "true")
        elif damage == "second_rpr":
            etree.SubElement(next(changed.iter(W + "r")), W + "rPr")
        elif damage == "second_ppr":
            etree.SubElement(changed, W + "pPr")
        else:
            changed.find(".//" + W + "t").set(W + "unapproved", "true")

    rejects(source, mutate(output, change), goal, binding)


@pytest.mark.parametrize("kind", ["replace_text", "delete_paragraphs"])
@pytest.mark.parametrize("damage", ["empty_run", "empty_container", "duplicate_properties", "extra_revision"])
def test_content_effects_preserve_nontext_envelope_and_reject_empty_revision_noise(tmp_path, kind, damage):
    source = document(tmp_path / "spec.docx")
    doc = Document(source)
    body = doc.paragraphs[2]
    body.add_run("")
    if kind == "replace_text":
        # Empty original wrappers also must survive character edits.
        body._p.append(etree.Element(W + "hyperlink"))
    doc.save(source)
    effect = {"id": "edit", "type": kind, "target": exact("Owner shall retain all content."), "tracked": True}
    if kind == "replace_text":
        effect.update(find="Owner", replace="Client", expected_matches=1)
    goal = contract(effect)
    output, binding = bound_execute(source, goal)

    def change(root):
        body = paragraphs(root)[2]
        if damage == "empty_run":
            run = next(run for run in body.iter(W + "r") if not list(run.iter(W + "t", W + "delText")))
            run.getparent().remove(run)
        elif damage == "empty_container":
            if kind == "replace_text":
                body.remove(body.find(W + "hyperlink"))
            else:
                etree.SubElement(body, W + "hyperlink")
        elif damage == "duplicate_properties":
            etree.SubElement(next(body.iter(W + "r")), W + "rPr")
            etree.SubElement(next(body.iter(W + "r")), W + "rPr")
        else:
            extra = etree.SubElement(body, W + "del")
            for key, value in {"id": "999", "author": "Abraham Borg", "date": "2026-01-01T00:00:00Z"}.items():
                extra.set(W + key, value)

    rejects(source, mutate(output, change), goal, binding)


@pytest.mark.parametrize("damage", ["all_count", "all_subset", "all_listed", "bool_is_int", "matched_locators",
                                    "unique_subset", "source", "contract"])
def test_target_evidence_cannot_be_forged_or_narrow_a_global_request(compound, damage):
    source, output, goal, binding = compound
    bad = copy.deepcopy(binding)
    assert binding["effects"]["font"] == {"targets": "all", "choice": "unique",
                                         "source_count": binding["effects"]["font"]["source_count"]}   # once, by count
    if damage == "all_count":
        bad["effects"]["font"]["source_count"] -= 1
    elif damage == "all_subset":
        bad["effects"]["font"] = {"targets": [], "choice": "user"}
    elif damage == "all_listed":          # a paragraph list is no global binding, however many it names
        bad["effects"]["font"] = {"targets": [copy.deepcopy(binding["effects"]["title"]["targets"][0])],
                                  "choice": "unique"}
    elif damage == "bool_is_int":
        bad["effects"]["title"]["targets"][0]["in_table"] = 1
    elif damage == "matched_locators":
        bad["effects"]["title"]["targets"][0]["matched_locators"] = [99]
    elif damage == "unique_subset":
        bad["effects"]["title"]["targets"] = []
    elif damage == "source":
        bad["source_sha256"] = "0" * 64
    else:
        bad["contract_sha256"] = "0" * 64
    rejects(source, output, goal, bad)


def test_replaced_paragraph_cannot_move_from_body_into_preceding_table(compound):
    source, output, goal, binding = compound

    def move(root):
        changed = paragraph(root, "Client")
        title_cell = paragraph(root, "FIRE SUPPRESSION").getparent()
        title_cell.append(changed)

    result = rejects(source, mutate(output, move), goal, binding)
    assert any("residual" in error for error in result["errors"]), result


@pytest.mark.parametrize("tracked", [False, True])
def test_inserted_paragraph_must_stay_in_its_approved_anchor_container(tmp_path, tracked):
    source = document(tmp_path / "spec.docx", "table")
    goal = contract({"id": "insert", "type": "insert_paragraph", "target": exact("FIRE SUPPRESSION"),
                     "position": "after", "text": "Approved addition.", "tracked": tracked})
    output, binding = bound_execute(source, goal)

    def move(root):
        added = paragraph(root, "Approved addition.")
        body = root.find(W + "body")
        body.insert(body.index(body.find(W + "tbl")) + 1, added)

    result = rejects(source, mutate(output, move), goal, binding)
    assert any("siblings" in error for error in result["errors"]), result


@pytest.mark.parametrize("damage", ["snapshot", "author", "same_id", "extra_attribute"])
def test_tracked_formatting_preserves_exact_original_snapshots(tmp_path, damage):
    source = document(tmp_path / "spec.docx")
    goal = contract({"id": "title", "type": "format", "target": exact("FIRE SUPPRESSION"),
                     "run": {"b": True}, "paragraph": {"keepNext": True}, "tracked": True})
    output, binding = bound_execute(source, goal)

    def change(root):
        revisions = list(paragraph(root, "FIRE SUPPRESSION").iter(W + "rPrChange"))
        if damage == "snapshot":
            etree.SubElement(revisions[0][0], W + "strike")
        elif damage == "author":
            revisions[0].set(W + "author", "Wrong author")
        elif damage == "same_id":
            revisions[1].set(W + "id", revisions[0].get(W + "id"))
        else:
            revisions[0].set(W + "invented", "true")

    rejects(source, mutate(output, change), goal, binding)


def test_symbol_runs_keep_their_symbol_font_while_normal_runs_get_the_standard_font(tmp_path):
    source = document(tmp_path / "spec.docx")
    doc = Document(source)
    symbol = doc.add_paragraph().add_run("Symbol text")
    symbol.font.name = "Wingdings"
    doc.save(source)
    goal = contract({"id": "font", "type": "format", "target": {"all": True},
                     "font": {"name": "Arial", "size": 11}})
    output, binding = bound_execute(source, goal)

    def change(root):
        fonts = paragraph(root, "Symbol text").find(W + "r/" + W + "rPr/" + W + "rFonts")
        fonts.set(W + "ascii", "Arial")
        fonts.set(W + "hAnsi", "Arial")

    rejects(source, mutate(output, change), goal, binding)


def test_removal_cannot_change_an_unrelated_package_relationship(tmp_path):
    source = document(tmp_path / "spec.docx")
    goal = contract({"id": "remove", "type": "remove_header_footer", "kinds": ["header"]})
    output, binding = bound_execute(source, goal)

    def change(root):
        root.find(REL + "Relationship").set("Target", "unapproved.xml")

    rejects(source, mutate(output, change, "word/_rels/document.xml.rels"), goal, binding)


@pytest.mark.parametrize("damage", ["sections", "unknown_param", "nonbool", "unknown_fill", "bad_nested_fill",
                                    "missing_sections", "missing_fills", "literal_trim", "duplicate_kind",
                                    "mixed_markers", "tracked", "page_conflict", "scope_conflict"])
def test_independent_template_grammar_is_closed_even_if_a_method_accepts_it(tmp_path, damage):
    source = document(tmp_path / "spec.docx")
    template = tmp_path / "template.docx"
    Document().save(template)
    effect = template_effect(template)
    goal = contract(effect)
    if damage == "sections":
        effect["sections"] = [0]
    elif damage == "unknown_param":
        effect["params"]["invented"] = True
    elif damage == "nonbool":
        effect["params"]["rescale_tabs"] = 1
    elif damage == "unknown_fill":
        effect["params"]["slot_fills"]["TITLE"]["invented"] = True
    elif damage == "bad_nested_fill":
        effect["params"]["slot_fills"]["TITLE"] = {"rule": "first_match", "rules": [
            {"rule": "literal", "value": "Approved text", "invented": True}]}
    elif damage == "missing_sections":
        del effect["params"]["template_sections"]
    elif damage == "missing_fills":
        del effect["params"]["slot_fills"]
    elif damage == "literal_trim":
        effect["params"]["slot_fills"]["TITLE"] = {"rule": "literal", "value": " Approved text "}
    elif damage == "duplicate_kind":
        effect["params"]["kinds"] = ["header", "header"]
    elif damage == "mixed_markers":
        effect["params"]["slot_markers"] = ["reference", "braces"]
    elif damage == "tracked":
        effect["params"]["tracked"] = True
    elif damage == "page_conflict":
        goal["effects"].append({"id": "page", "type": "page_setup", "margins": {"header": 444}})
    else:
        goal["effects"].append({"id": "header-font", "type": "format", "target": {"scope": ["headers"], "all": True},
                                "font": {"name": "Arial"}})
    result = rejects(source, source, goal, {})
    assert result["checks"][0]["id"] == "E1" and result["checks"][0]["status"] == "FAIL", result
