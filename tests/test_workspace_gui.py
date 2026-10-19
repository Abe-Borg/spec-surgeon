"""The desktop workspace keeps project scope, user answers and saved work across runs."""
from __future__ import annotations

import json
import threading
from pathlib import Path

import pytest

from make_fixtures import build
from spec_surgeon.campaign import create_campaign, load_campaign
from spec_surgeon.gui import Session


@pytest.fixture
def qt_env(monkeypatch):
    from qt_support import close_windows, qt_app_or_skip
    app = qt_app_or_skip()
    monkeypatch.setenv("SPEC_SURGEON_NO_UPDATE_CHECK", "1")
    windows = []
    yield app, windows
    close_windows(app, windows)


@pytest.fixture
def campaign(tmp_path):
    roots = []
    for name in ("Alpha", "Beta"):
        root = tmp_path / name
        root.mkdir()
        (root / "PROJECT.json").write_text(json.dumps({"name": name}), encoding="utf-8")
        (root / "Spec.docx").write_bytes(build("21 00 00", name.upper()))
        roots.append(str(root))
    return create_campaign(roots, tmp_path / "workspace")


def workspace(qt_env, campaign, tmp_path, **kwargs):
    from spec_surgeon.ui.app import RecordingPrompts, build_legacy_window as build_window
    session = Session(campaign=campaign, out_root=tmp_path / "out")
    window = build_window(session, prompts=kwargs.pop("prompts", RecordingPrompts()), **kwargs)
    qt_env[1].append(window)
    assert window.do_campaign()
    return window, window.campaign_dialog


def test_project_editor_preserves_separate_goals_and_template(qt_env, campaign, tmp_path):
    from spec_surgeon import resources
    window, dialog = workspace(qt_env, campaign, tmp_path)
    assert dialog.resources.calls.value() == resources.DEFAULT_MAX_CALLS
    assert dialog.resources.cost.value() == resources.DEFAULT_MAX_COST_USD
    assert dialog.resources.total_tokens.value() == resources.DEFAULT_MAX_TOTAL_TOKENS
    assert dialog.resources.effort.currentText() == resources.DEFAULT_EFFORT
    assert "$500.00" in dialog.allowance.text()
    first = dialog._selected_project
    template = tmp_path / "template.docx"
    template.write_bytes(build("21 00 00", "TEMPLATE"))
    dialog.request.setPlainText("Use Alpha wording")
    dialog.template.setText(str(template))
    dialog.context.setPlainText('{"name":"Alpha","approved":"ALPHA PROJECT"}')
    dialog.projects.setCurrentRow(1)  # leaving an edited project saves it
    second = dialog._selected_project
    assert first != second
    dialog.request.setPlainText("Use Beta wording")
    assert dialog.save_project()
    restored = load_campaign(campaign.out_dir)
    assert restored.projects[first]["planning_draft"]["request"] == "Use Alpha wording"
    assert restored.projects[first]["planning_draft"]["template"] == str(template)
    assert restored.projects[first]["context"]["approved"] == "ALPHA PROJECT"
    assert restored.projects[second]["planning_draft"]["request"] == "Use Beta wording"
    assert not restored.projects[second]["planning_draft"]["template"]


def test_bad_context_does_not_discard_editor_on_project_switch(qt_env, campaign, tmp_path):
    window, dialog = workspace(qt_env, campaign, tmp_path)
    first = dialog._selected_project
    dialog.request.setPlainText("My unsaved goal")
    dialog.context.setPlainText("invalid JSON")
    dialog.projects.setCurrentRow(1)
    assert dialog._selected_project == first
    assert dialog.request.toPlainText() == "My unsaved goal"
    assert dialog._dirty
    dialog.context.setPlainText("{}")  # valid before teardown saves the editor


def test_assignment_and_answer_are_scoped_to_selected_project(qt_env, campaign, tmp_path):
    window, dialog = workspace(qt_env, campaign, tmp_path)
    doc_id = dialog.documents.item(0, 0).data(256)
    dialog.documents.selectRow(0)
    dialog.project_id.setEditText("Reviewed project")
    dialog.assign_documents()
    assert campaign.documents[doc_id]["project_id"] == "Reviewed project"
    assert dialog._selected_project == "Reviewed project"
    dialog.request.setPlainText("Format this project")
    assert dialog.save_project()
    project = campaign.projects["Reviewed project"]
    project["questions"] = [{"id": "which-font", "text": "Which font?", "document_ids": [doc_id]}]
    dialog.refresh()
    dialog.questions.setCurrentIndex(1)
    dialog.answer.setPlainText("Arial")
    dialog.save_answer()
    assert project["decisions"][-1]["question_id"] == "which-font"
    assert project["decisions"][-1]["document_ids"] == [doc_id]
    assert all(not p["decisions"] for pid, p in campaign.projects.items() if pid != "Reviewed project")


def test_workspace_run_uses_background_worker_and_clears_visual_ack(qt_env, campaign, tmp_path, monkeypatch):
    window, dialog = workspace(qt_env, campaign, tmp_path)
    dialog.request.setPlainText("Change font")
    dialog.review_ack.setChecked(True)
    calls = []
    main_thread = threading.get_ident()

    def run(**kwargs):
        calls.append((threading.get_ident(), kwargs))
        kwargs["log"]("Retained 1 ready document; investigating the remaining files")
        return campaign

    monkeypatch.setattr(campaign, "run", run)
    dialog.resources.total_tokens.setValue(25000000)
    dialog.run_workspace()
    window._spec_surgeon["wait"]()
    assert calls and calls[0][0] != main_thread
    assert calls[0][1]["options"].max_total_tokens == 25000000
    assert not dialog.review_ack.isChecked()
    assert not dialog.busy and not window.busy
    assert "Retained 1 ready document" in dialog.log.toPlainText()


def test_workspace_can_create_and_reopen_recursively_without_a_key(qt_env, tmp_path, monkeypatch):
    from spec_surgeon import llm
    from spec_surgeon.ui.app import RecordingPrompts
    monkeypatch.setattr(llm, "credentials_found", lambda: False)
    root = tmp_path / "nested"
    (root / "discipline").mkdir(parents=True)
    (root / "discipline" / "Spec.docx").write_bytes(build("21 00 00", "NESTED"))
    destination = tmp_path / "saved"
    prompts = RecordingPrompts(files=[str(root), str(destination)], answers=[False])
    window, dialog = workspace(qt_env, None, tmp_path, prompts=prompts)
    dialog.create_workspace()
    window._spec_surgeon["wait"]()
    assert len(window.session.campaign.documents) == 1
    assert (destination / "campaign.json").is_file()
    prompts.files.append(str(destination))
    dialog.review_ack.setChecked(True)
    dialog.open_workspace()
    window._spec_surgeon["wait"]()
    assert len(dialog.campaign.documents) == 1
    assert not dialog.review_ack.isChecked()


def test_title_goal_editor_records_explicit_target_and_properties(qt_env, campaign, tmp_path):
    window, dialog = workspace(qt_env, campaign, tmp_path)
    dialog.request.setPlainText("Make ALPHA bold")
    dialog.goal_mode.setCurrentIndex(1)
    dialog.title_text.setText("ALPHA")
    dialog.section_number.setText("21 00 00")
    dialog.properties["run", "b"].setCurrentIndex(1)
    assert dialog.save_project()
    contract = campaign.projects[dialog._selected_project]["goal"]["adaptive_contract"]
    assert contract["target"] == {"title": "ALPHA", "section_number": "21 00 00"}
    assert contract["run"] == {"b": True}
    assert "paragraph" not in contract


def test_title_workspace_runs_reopens_and_exports_verified_copies_keylessly(qt_env, campaign, tmp_path, monkeypatch):
    from spec_surgeon import llm
    monkeypatch.setattr(llm, "credentials_found", lambda: False)
    window, dialog = workspace(qt_env, campaign, tmp_path)
    project_id = dialog._selected_project
    project = campaign.projects[project_id]
    source = Path(campaign.documents[project["document_ids"][0]]["source"])
    before = source.read_bytes()
    dialog.request.setPlainText("Make ALPHA bold")
    dialog.goal_mode.setCurrentIndex(1)
    dialog.title_text.setText("ALPHA")
    dialog.properties["run", "b"].setCurrentIndex(1)
    dialog.render.setChecked(False)
    dialog.memory.setChecked(False)
    dialog.run_workspace()
    window._spec_surgeon["wait"]()
    assert project["status"] == "READY", dialog.evidence.toPlainText()
    assert "READY" in dialog.evidence.toPlainText()
    candidate = campaign.documents[project["document_ids"][0]]["output"]
    assert Path(candidate).is_file()
    assert source.read_bytes() == before
    window.prompts.files.append(str(campaign.out_dir))
    dialog.open_workspace()
    window._spec_surgeon["wait"]()
    assert dialog.campaign.projects[project_id]["status"] == "READY"
    destination = tmp_path / "delivery"
    window.prompts.files.append(str(destination))
    dialog.export_workspace()
    window._spec_surgeon["wait"]()
    assert (destination / "export.json").is_file(), dialog.log.toPlainText()
    assert list(destination.rglob("*.docx"))
    assert source.read_bytes() == before


def test_duplicate_title_requires_explicit_gui_choice_before_resume(qt_env, tmp_path):
    from docx import Document
    root = tmp_path / "inputs"
    root.mkdir()
    source = root / "ambiguous.docx"
    document = Document()
    document.add_paragraph("FIRE SUPPRESSION")
    document.add_paragraph("Other content stays unchanged.")
    document.add_paragraph("FIRE SUPPRESSION")
    document.save(source)
    campaign = create_campaign([str(root)], tmp_path / "workspace")
    window, dialog = workspace(qt_env, campaign, tmp_path)
    dialog.request.setPlainText("Make the intended FIRE SUPPRESSION title bold")
    dialog.goal_mode.setCurrentIndex(1)
    dialog.title_text.setText("FIRE SUPPRESSION")
    dialog.properties["run", "b"].setCurrentIndex(1)
    dialog.render.setChecked(False)
    dialog.memory.setChecked(False)
    dialog.run_workspace()
    window._spec_surgeon["wait"]()
    project = campaign.projects[dialog._selected_project]
    assert project["status"] == "NEEDS_INPUT", dialog.evidence.toPlainText()
    assert dialog.questions.count() == 2
    dialog.questions.setCurrentIndex(1)
    assert dialog.bindings.count() == 2
    dialog.bindings.setCurrentIndex(1)
    doc_id, candidate_id = dialog.bindings.currentData()
    dialog.bind_target()
    assert campaign.documents[doc_id]["binding"]["candidate_id"] == candidate_id
    dialog.run_workspace()
    window._spec_surgeon["wait"]()
    assert project["status"] == "READY", dialog.evidence.toPlainText()


def test_workspace_restores_saved_allowance_and_preserves_hidden_options(qt_env, campaign, tmp_path, monkeypatch):
    from spec_surgeon.recovery import RecoveryOptions
    options = RecoveryOptions(max_cost_usd=2.5, max_calls=7, max_attempts=3, effort="medium",
                              render=False, use_memory=False, allow_codegen=True,
                              model="saved-model", renderer="none", group_call_slice=8,
                              max_context_tokens=65000, critic_review=False)
    campaign.run(options=options)  # no project goals: persist settings without invoking a model
    restored = load_campaign(campaign.out_dir)
    window, dialog = workspace(qt_env, restored, tmp_path)
    assert dialog.resources.cost.value() == 2.5
    assert dialog.resources.calls.value() == 7
    assert dialog.resources.effort.currentText() == "medium"
    assert dialog.codegen.isChecked() and not dialog.memory.isChecked() and not dialog.render.isChecked()
    received = []
    monkeypatch.setattr(restored, "run", lambda **kw: received.append(kw["options"]) or restored)
    dialog.run_workspace()
    window._spec_surgeon["wait"]()
    assert received[0] == options


def test_proposal_is_reviewed_and_approved_before_execution_and_survives_restart(qt_env, campaign, tmp_path):
    from fake_claude import FakeClient, resp, tool
    contract = {"kind": "document_effects", "effects": [
        {"id": "title", "type": "format", "target": {"scope": ["document"],
         "locators": [{"exact_text": "ALPHA"}], "cardinality": 1}, "run": {"b": True}},
        {"id": "wording", "type": "replace_text", "target": {"scope": ["document"],
         "locators": [{"exact_text": "ALPHA"}], "cardinality": 1},
         "find": "ALPHA", "replace": "ALPHA PROJECT", "expected_matches": 1, "tracked": True}]}
    fake = FakeClient(resp(tool("propose_plan", {"contract": contract, "rationale": "Format and update the title",
                      "risks": ["The title text changes; execution must preserve the requested formatting."],
                      "evidence": ["The selected source has one exact ALPHA title."]})))
    window, dialog = workspace(qt_env, campaign, tmp_path)
    window.session.client = fake
    project_id = dialog._selected_project
    project = campaign.projects[project_id]
    originals = {doc_id: Path(doc["source"]).read_bytes() for doc_id, doc in campaign.documents.items()}
    dialog.request.setPlainText('Make the ALPHA title bold and replace it with "ALPHA PROJECT".')
    dialog.render.setChecked(False)
    dialog.memory.setChecked(False)
    dialog.generate_plan()
    window._spec_surgeon["wait"]()
    proposed = project["proposed_plan"]
    assert proposed["status"] == "PROPOSED", dialog.log.toPlainText()
    assert dialog.tabs.currentWidget() is dialog.plan_page
    assert 'Exact text after:  "ALPHA PROJECT"' in dialog.plan_review.toPlainText()
    assert dialog.approve_btn.isEnabled()
    assert not project["runs"]
    dialog.run_workspace()
    window._spec_surgeon["wait"]()
    assert not project["runs"]  # a model proposal is not execution authorization
    window.prompts.files.append(str(campaign.out_dir))
    dialog.open_workspace()
    window._spec_surgeon["wait"]()
    assert dialog._reviewed_plan_id == proposed["id"]
    dialog.approve_plan()
    window._spec_surgeon["wait"]()
    assert dialog.campaign.projects[project_id]["proposed_plan"]["status"] == "APPROVED"
    assert all(Path(doc["source"]).read_bytes() == originals[doc_id]
               for doc_id, doc in dialog.campaign.documents.items())
    # The approved unambiguous builtin effects compile and verify without another
    # model call. Planning, source pins, execution, persistence and export are real.
    source = dialog.campaign.documents[proposed["document_ids"][0]]["source"]
    dialog.run_workspace()
    window._spec_surgeon["wait"]()
    current = dialog.campaign.projects[project_id]
    assert current["status"] == "READY", dialog.evidence.toPlainText() + dialog.log.toPlainText()
    assert len(fake.requests) == 1  # the proposal; clean builtin execution is deterministic
    ready_path = dialog.campaign.documents[proposed["document_ids"][0]]["output"]
    window.prompts.files.append(str(dialog.campaign.out_dir))
    dialog.open_workspace()
    window._spec_surgeon["wait"]()
    assert dialog.campaign.documents[proposed["document_ids"][0]]["output"] == ready_path
    destination = tmp_path / "planned-delivery"
    window.prompts.files.append(str(destination))
    dialog.export_workspace()
    window._spec_surgeon["wait"]()
    outputs = list(destination.rglob("*.docx"))
    assert len(outputs) == 1, dialog.log.toPlainText()
    from spec_surgeon.goals import verify_goal
    binding = dialog.campaign.documents[proposed["document_ids"][0]]["binding"]
    assert verify_goal(Path(source), outputs[0], contract, binding)["passed"]
    assert all(Path(doc["source"]).read_bytes() == originals[doc_id]
               for doc_id, doc in dialog.campaign.documents.items())


def test_saved_draft_scope_changes_disable_old_proposal_approval(qt_env, campaign, tmp_path):
    from fake_claude import FakeClient, resp, tool
    contract = {"kind": "format_section_title", "target": {"title": "ALPHA"}, "run": {"b": True}}
    window, dialog = workspace(qt_env, campaign, tmp_path)
    window.session.client = FakeClient(resp(tool("propose_plan", {"contract": contract,
                                     "rationale": "One exact title", "evidence": [], "risks": []})))
    dialog.request.setPlainText("Make ALPHA bold")
    dialog.generate_plan()
    window._spec_surgeon["wait"]()
    assert dialog.approve_btn.isEnabled(), dialog.log.toPlainText()
    dialog.request.setPlainText("Actually make ALPHA italic")
    assert dialog.save_project()
    dialog.refresh()
    assert not dialog.approve_btn.isEnabled()
    dialog.approve_plan()
    assert "revised plan" in dialog.status.text()


def test_selected_document_plan_retains_other_ready_work_in_same_project(qt_env, tmp_path):
    from fake_claude import FakeClient, resp, tool
    from spec_surgeon.recovery import RecoveryOptions
    root = tmp_path / "inputs"
    root.mkdir()
    for name in ("first", "second"):
        (root / (name + ".docx")).write_bytes(build("21 00 00", "ALPHA"))
    campaign = create_campaign([str(root)], tmp_path / "workspace")
    pid = next(iter(campaign.projects))
    campaign.set_goal(pid, "Make ALPHA bold", adaptive_contract={"kind": "format_section_title",
                      "target": {"title": "ALPHA"}, "run": {"b": True}})
    campaign.run(options=RecoveryOptions(render=False, use_memory=False))
    assert all(doc["status"] == "READY" for doc in campaign.documents.values())
    window, dialog = workspace(qt_env, campaign, tmp_path)
    selected, retained = campaign.projects[pid]["document_ids"]
    retained_path = campaign.documents[retained]["output"]
    retained_bytes = Path(retained_path).read_bytes()
    contract = {"kind": "format_section_title", "target": {"title": "ALPHA"}, "run": {"i": True}}
    window.session.client = FakeClient(resp(tool("propose_plan", {"contract": contract,
                                      "rationale": "The selected document has this title", "risks": [], "evidence": []})))
    dialog.request.setPlainText("Make ALPHA italic in the selected document")
    dialog.documents.selectRow(0)
    dialog.scope.setCurrentIndex(1)
    dialog.generate_plan()
    window._spec_surgeon["wait"]()
    proposed = campaign.projects[pid]["proposed_plan"]
    assert proposed["document_ids"] == [selected], dialog.log.toPlainText()
    assert all(doc["status"] == "READY" for doc in campaign.documents.values())
    dialog.approve_plan()
    window._spec_surgeon["wait"]()
    assert campaign.documents[selected]["status"] == "PENDING"
    assert campaign.documents[retained]["status"] == "READY"
    assert campaign.documents[retained]["output"] == retained_path
    assert Path(retained_path).read_bytes() == retained_bytes
    window.prompts.files.append(str(campaign.out_dir))
    dialog.open_workspace()
    window._spec_surgeon["wait"]()
    assert dialog.scope.currentIndex() == 1
    assert dialog._selected_documents() == [selected]


def test_refresh_repaired_source_preserves_other_ready_copies_and_continues(qt_env, tmp_path):
    from spec_surgeon.recovery import RecoveryOptions
    root = tmp_path / "inputs"
    root.mkdir()
    source, broken = root / "ready.docx", root / "broken.docx"
    source.write_bytes(build("21 00 00", "ALPHA"))
    broken.write_bytes(b"malformed docx")
    campaign = create_campaign([str(root)], tmp_path / "workspace")
    pid = next(iter(campaign.projects))
    campaign.set_goal(pid, "Make ALPHA bold", adaptive_contract={"kind": "format_section_title",
                      "target": {"title": "ALPHA"}, "run": {"b": True}})
    campaign.run(options=RecoveryOptions(render=False, use_memory=False))
    ready_id = next(i for i, doc in campaign.documents.items() if doc["source"] == str(source))
    broken_id = next(i for i, doc in campaign.documents.items() if doc["source"] == str(broken))
    retained_path = campaign.documents[ready_id]["output"]
    assert campaign.documents[broken_id]["status"] == "UNREADABLE"
    window, dialog = workspace(qt_env, campaign, tmp_path)
    broken.write_bytes(build("21 00 00", "ALPHA"))
    for row in range(dialog.documents.rowCount()):
        if dialog.documents.item(row, 0).data(256) == broken_id:
            dialog.documents.selectRow(row)
    dialog.refresh_sources()
    window._spec_surgeon["wait"]()
    assert campaign.documents[broken_id]["status"] == "PENDING", dialog.log.toPlainText()
    assert campaign.documents[ready_id]["status"] == "READY"
    assert campaign.documents[ready_id]["output"] == retained_path
    dialog.run_workspace()
    window._spec_surgeon["wait"]()
    assert all(doc["status"] == "READY" for doc in campaign.documents.values()), dialog.evidence.toPlainText()


def test_general_target_choice_resumes_saved_investigation_with_writer_and_critic(qt_env, tmp_path):
    from docx import Document
    from fake_claude import FakeClient, resp, tool
    root = tmp_path / "inputs"
    root.mkdir()
    source = root / "duplicates.docx"
    document = Document()
    document.add_paragraph("FIRE SUPPRESSION")
    document.add_paragraph("This is the unaffected intervening paragraph.")
    document.add_paragraph("FIRE SUPPRESSION")
    document.save(source)
    original = source.read_bytes()
    campaign = create_campaign([str(root)], tmp_path / "workspace")
    window, dialog = workspace(qt_env, campaign, tmp_path)
    pid = dialog._selected_project
    contract = {"kind": "document_effects", "effects": [{"id": "title", "type": "format",
                "target": {"scope": ["document"], "locators": [{"exact_text": "FIRE SUPPRESSION"}],
                           "cardinality": 1}, "run": {"b": True}}]}
    fake = FakeClient(resp(tool("propose_plan", {"contract": contract, "rationale": "Change the intended title",
                      "risks": ["The title occurs twice; execution needs an explicit target choice."], "evidence": []})))
    window.session.client = fake
    dialog.request.setPlainText("Make the intended FIRE SUPPRESSION title bold")
    dialog.render.setChecked(False)
    dialog.memory.setChecked(False)
    dialog.generate_plan()
    window._spec_surgeon["wait"]()
    dialog.approve_plan()
    window._spec_surgeon["wait"]()
    dialog.run_workspace()
    window._spec_surgeon["wait"]()
    project = dialog.campaign.projects[pid]
    assert project["status"] == "NEEDS_INPUT", dialog.evidence.toPlainText()
    run_path = project["runs"][-1]["path"]
    window.prompts.files.append(str(campaign.out_dir))
    dialog.open_workspace()
    window._spec_surgeon["wait"]()
    dialog.questions.setCurrentIndex(1)
    assert dialog.binding_effect.count() == 1
    assert dialog.binding_candidates.count() == 2
    dialog.binding_candidates.item(1).setSelected(True)
    selected_id = dialog.binding_candidates.item(1).data(256)
    dialog.bind_target()
    project = dialog.campaign.projects[pid]
    assert project["decisions"][-1]["candidate_ids"] == [selected_id]
    assert project["decisions"][-1]["effect_id"] == "title"
    original_stream, writer_turns = fake.beta.messages.stream, []

    def execute_stream(**kwargs):
        names = {definition["name"] for definition in kwargs["tools"]}
        if "review_repair" in names:
            response = resp(tool("review_repair", {"verdict": "accept",
                            "reason": "The explicit selected target is bold and surrounding content is preserved."}, "critic"))
        elif not writer_turns:
            response = resp(tool("replan_goal", {"file": str(source)}, "resolve-targets"))
            writer_turns.append("replan")
        else:
            assert len(writer_turns) == 1, "Unexpected additional writer turn"
            plans = []
            for message in kwargs["messages"]:
                if isinstance(message.get("content"), list):
                    for block in message["content"]:
                        if isinstance(block, dict) and block.get("type") == "tool_result" and isinstance(block.get("content"), str):
                            try:
                                value = json.loads(block["content"])
                            except ValueError:
                                continue
                            if isinstance(value, dict) and value.get("plan_id"):
                                plans.append(value["plan_id"])
            assert plans, "The writer needs a source-bound plan before submitting"
            response = resp(tool("submit_repair", {"plan_id": plans[-1]}, "execute-plan"))
            writer_turns.append("submit")
        fake.responses.append(response)
        return original_stream(**kwargs)

    fake.beta.messages.stream = execute_stream
    dialog.run_workspace()
    window._spec_surgeon["wait"]()
    assert project["status"] == "READY", dialog.evidence.toPlainText() + dialog.log.toPlainText()
    assert project["runs"][-1]["path"] == run_path
    assert len(fake.requests) == 4
    assert source.read_bytes() == original
    doc_id = project["document_ids"][0]
    output = Document(dialog.campaign.documents[doc_id]["output"])
    assert output.paragraphs[0].runs[0].bold is not True
    assert output.paragraphs[2].runs[0].bold is True


def test_planning_question_and_answer_continue_to_reviewable_proposal(qt_env, campaign, tmp_path):
    from fake_claude import FakeClient, resp, tool
    window, dialog = workspace(qt_env, campaign, tmp_path)
    pid = dialog._selected_project
    doc_id = campaign.projects[pid]["document_ids"][0]
    fake = FakeClient(resp(tool("need_user", {"question": "Which exact title wording should I use?", "document_ids": [doc_id]})),
                      resp(tool("propose_plan", {"contract": {"kind": "document_effects", "effects": [
                          {"id": "wording", "type": "replace_text", "target": {"scope": ["document"],
                           "locators": [{"exact_text": "ALPHA"}], "cardinality": 1}, "find": "ALPHA",
                           "replace": "ALPHA PROJECT", "expected_matches": 1}]},
                           "rationale": "Use the user's exact answer for this project."})))
    window.session.client = fake
    dialog.request.setPlainText("Update the project title wording")
    dialog.generate_plan()
    window._spec_surgeon["wait"]()
    assert campaign.projects[pid]["proposed_plan"]["status"] == "NEEDS_INPUT"
    assert "Planning needs your answer" in dialog.status.text()
    assert not dialog.approve_btn.isEnabled()
    dialog.questions.setCurrentIndex(1)
    dialog.answer.setPlainText("Use exactly ALPHA PROJECT")
    dialog.save_answer()
    assert campaign.projects[pid]["decisions"][-1]["document_ids"] == [doc_id]
    assert not dialog.answer.toPlainText()
    dialog.generate_plan()
    window._spec_surgeon["wait"]()
    assert campaign.projects[pid]["proposed_plan"]["status"] == "PROPOSED", dialog.log.toPlainText()
    assert dialog.approve_btn.isEnabled()
    assert 'Exact text after:  "ALPHA PROJECT"' in dialog.plan_review.toPlainText()
    assert not campaign.projects[pid]["runs"]
    assert "Use exactly ALPHA PROJECT" in json.dumps(fake.requests[-1]["messages"], default=str)
