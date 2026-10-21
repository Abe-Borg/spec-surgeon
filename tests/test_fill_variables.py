"""fill_variables (docs/BUILTIN_OPS_BATCH1_DESIGN.md §1): placeholders and highlighted text in the body, headers and
footers filled from typed values or a table the user keeps — end to end, a site column per project, the refusals, and
planted defects (a wrong OUTPUT; wrong IMPLEMENTATIONS shared by plan and apply), each caught by V10 alone."""
from __future__ import annotations

import csv
import json
from typing import Any, Dict, List
from pathlib import Path

import pytest

from make_fixtures import build_batch1, png
from spec_surgeon.ops import JobException, get_op, merge_fragments
from spec_surgeon.ops import fill_variables as fill_module
from spec_surgeon.ops import tables as tables_module
from spec_surgeon.package import DocxPackage, W
from spec_surgeon.verify import run_suite
from test_text_edits import _xlsx, accepted_text, failed, mutate, para, v10_errors
from spec_surgeon.compiler import _Batch, check_submission

AUTHOR = "Abraham Borg"
VALUES = {"SITE_NAME": "Ashburn VA3", "OWNER_NAME": "Acme Cloud LLC", "PROJECT NO": "2026-117"}


@pytest.fixture(scope="module")
def doc() -> bytes:
    return build_batch1()


def run_op(src: bytes, params: Dict[str, Any], select: Dict[str, Any] = None):
    pkg = DocxPackage(src)
    frag = get_op("fill_variables").plan(pkg, select or {}, {"author": AUTHOR, **params})
    plan = merge_fragments([frag])
    work = pkg.copy()
    get_op("fill_variables").apply(work, frag)
    out = work.output_bytes()
    return plan, out, run_suite(src, out, plan.to_dict(), "src.docx", "out.docx")


def header_text(data: bytes) -> str:
    return accepted_text(next(DocxPackage(data).xml("word/header1.xml").iter(W + "p")))


@pytest.mark.parametrize("tracked", [True, False])
def test_fills_body_header_split_placeholders_and_highlights(doc, tracked):
    plan, out, rec = run_op(doc, {"values": VALUES, "tracked": tracked})
    assert rec.passed, rec.to_markdown()
    assert accepted_text(para(out, "A.")) == "A.\tThe Owner is Acme Cloud LLC for the Ashburn VA3 campus."
    assert accepted_text(para(out, "B.")) == "B.\tThe project number is 2026-117."
    assert header_text(out).startswith("ACME DATA CENTERS — Ashburn VA3\t")
    # the slot's highlight is gone from the value (a tracked deletion keeps the old text's formatting)
    assert all(any(q.tag == W + "del" for q in h.iterancestors()) for h in para(out, "B.").iter(W + "highlight"))
    assert plan.mode == ("content" if tracked else "untracked")
    assert any("SITE_NAME → 'Ashburn VA3' ×2" in n for n in plan.notes)


@pytest.mark.parametrize("kind", ["xlsx", "csv", "json"])
def test_one_column_per_site(doc, tmp_path, kind):
    rows = [["Variable", "VA3", "TX1"], ["SITE_NAME", "Ashburn VA3", "Dallas TX1"],
            ["OWNER_NAME", "Acme Cloud LLC", "Acme Cloud LLC"], ["PROJECT NO", "2026-117", "2026-121"]]
    if kind == "xlsx":
        path = tmp_path / "Sites.xlsx"
        path.write_bytes(_xlsx(rows))
        table = {"file": str(path), "value_column": "TX1"}
    elif kind == "csv":
        path = tmp_path / "Sites.csv"
        with open(path, "w", newline="", encoding="utf-8") as f:
            csv.writer(f).writerows(rows)
        table = {"file": str(path), "value_column": "TX1"}
    else:
        path = tmp_path / "PROJECT.json"
        path.write_text(json.dumps({"stamp": "x", "sites": {"TX1": {r[0]: r[2] for r in rows[1:]}}}), encoding="utf-8")
        table = {"file": str(path), "key": "sites.TX1"}
    plan, out, rec = run_op(doc, {"table": table, "tracked": False})
    assert rec.passed, rec.to_markdown()
    assert accepted_text(para(out, "B.")) == "B.\tThe project number is 2026-121."
    assert header_text(out).startswith("ACME DATA CENTERS — Dallas TX1\t")


def test_typed_values_win_over_the_table(doc, tmp_path):
    path = tmp_path / "v.json"
    path.write_text(json.dumps({"SITE_NAME": "From the table", "OWNER_NAME": "O", "PROJECT NO": "1"}), encoding="utf-8")
    plan, out, rec = run_op(doc, {"table": {"file": str(path)}, "values": {"SITE_NAME": "Typed"}, "tracked": False})
    assert rec.passed and header_text(out).startswith("ACME DATA CENTERS — Typed\t")


def test_a_placeholder_without_a_value(doc):
    with pytest.raises(JobException) as e:
        run_op(doc, {"values": {"SITE_NAME": "Ashburn VA3"}})
    assert e.value.code == "unfilled-slot"
    plan, out, rec = run_op(doc, {"values": {"SITE_NAME": "Ashburn VA3"}, "unknown": "leave", "tracked": False})
    assert rec.passed, rec.to_markdown()
    assert "{{OWNER_NAME}}" in accepted_text(para(out, "A."))
    assert any("left as they are" in f for f in plan.flags)


def test_unusable_values_are_refused(doc, tmp_path):
    path = tmp_path / "v.xlsx"
    path.write_bytes(_xlsx([["Name", "Value"], ["SITE_NAME", 46068], ["OWNER_NAME", "O"], ["PROJECT NO", "P"]],
                           number_format=14))                                  # a date where text belongs
    with pytest.raises(JobException) as e:
        run_op(doc, {"table": {"file": str(path)}})
    assert e.value.code == "table-value"
    with pytest.raises(JobException) as e:
        run_op(doc, {"values": {**VALUES, "SITE_NAME": "Ashburn​VA3"}})   # a zero-width space
    assert e.value.code == "table-value"


def test_textbox_placeholder_is_refused(tmp_path):
    from make_fixtures import _textbox_run
    from docx import Document
    import io
    d = Document()
    p = d.add_paragraph("Site: ")
    _textbox_run(p, "{{SITE_NAME}}")
    buf = io.BytesIO(); d.save(buf)
    with pytest.raises(JobException) as e:
        run_op(buf.getvalue(), {"values": VALUES})
    assert e.value.code == "textbox"


def test_needs_values_or_a_table(doc):
    with pytest.raises(JobException) as e:
        run_op(doc, {})
    assert e.value.code == "bad-params"


# ---- planted defects ---------------------------------------------------------------------------------------------

def test_defect_highlight_left_on_the_filled_slot(doc):
    plan, out, rec = run_op(doc, {"values": VALUES, "tracked": False})

    def highlight_again(root):
        from lxml import etree
        for r in root.iter(W + "r"):
            if any((t.text or "") == "2026-117" for t in r.iter(W + "t")):
                rpr = r.find(W + "rPr")
                if rpr is None:
                    rpr = etree.Element(W + "rPr")
                    r.insert(0, rpr)
                h = etree.SubElement(rpr, W + "highlight")
                h.set(W + "val", "yellow")
    rec = run_suite(doc, mutate(out, "word/document.xml", highlight_again), plan.to_dict(), "s", "o")
    assert failed(rec) == ["V10"]


def test_shared_defect_slot_range_off_by_one(doc, monkeypatch):
    real = fill_module.slot_range
    monkeypatch.setattr(fill_module, "slot_range",
                        lambda idx, s: (lambda a, b: (a, b - 1) if s.marker == "braces" else (a, b))(*real(idx, s)))
    plan, out, rec = run_op(doc, {"values": VALUES, "tracked": False})
    assert header_text(out).startswith("ACME DATA CENTERS — Ashburn VA3}\t")     # the last brace stayed
    assert rec.get("V2").status == "PASS"                  # plan and apply agree with each other …
    assert failed(rec) == ["V10"]                          # … not with the verifier's own slots


def test_shared_defect_highlight_slots_missed(doc, monkeypatch):
    real = fill_module.find_slots
    monkeypatch.setattr(fill_module, "find_slots",
                        lambda p, markers, *a, **k: real(p, [m for m in markers if m != "highlight"], *a, **k))
    plan, out, rec = run_op(doc, {"values": VALUES, "tracked": False})
    assert "PROJECT NO" in accepted_text(para(out, "B."))
    assert failed(rec) == ["V10"]
    assert any("predicted" in e for e in v10_errors(rec))


def test_shared_defect_table_column_ignored(doc, tmp_path, monkeypatch):
    path = tmp_path / "Sites.csv"
    with open(path, "w", newline="", encoding="utf-8") as f:
        csv.writer(f).writerows([["Variable", "VA3", "TX1"], ["SITE_NAME", "Ashburn VA3", "Dallas TX1"],
                                 ["OWNER_NAME", "Acme", "Acme"], ["PROJECT NO", "1", "2"]])
    real = tables_module.variables
    monkeypatch.setattr(tables_module, "variables",
                        lambda spec, where="table": real({k: v for k, v in spec.items() if k != "value_column"}, where))
    plan, out, rec = run_op(doc, {"table": {"file": str(path), "value_column": "TX1"}, "tracked": False})
    assert header_text(out).startswith("ACME DATA CENTERS — Ashburn VA3\t")   # the wrong site, consistently
    assert failed(rec) == ["V10"]


def test_conflicts_with_another_op_on_the_same_paragraph(doc):
    pkg = DocxPackage(doc)
    a = get_op("fill_variables").plan(pkg, {}, {"values": VALUES, "author": AUTHOR})
    b = get_op("replace_table").plan(pkg, {}, {"pairs": [{"find": "campus", "replace": "site"}], "author": AUTHOR})
    with pytest.raises(JobException) as e:
        merge_fragments([a, b])
    assert e.value.code == "conflict"
    c = get_op("replace_table").plan(pkg, {}, {"pairs": [{"find": "data halls", "replace": "halls"}],
                                               "author": AUTHOR})
    plan = merge_fragments([a, c])                       # different paragraphs: one spec, both verified
    work = pkg.copy()
    for f in (a, c):
        get_op(f.op).apply(work, f)
    assert run_suite(doc, work.output_bytes(), plan.to_dict(), "s", "o").passed


@pytest.fixture
def inputs_folder(tmp_path) -> Path:
    d = tmp_path / "specs"
    d.mkdir()
    (d / "211319 - Preaction.docx").write_bytes(build_batch1())
    (d / "PROJECT.xlsx").write_bytes(_xlsx([["Variable", "VA1", "VA3"], ["SITE_NAME", "Ashburn", "Leesburg"],
                                            ["OWNER_NAME", "Contoso", "Contoso"], ["PROJECT NO", "1001", "1003"]]))
    (d / "pairs.csv").write_text("find,replace\nOwner,Client\n", encoding="utf-8")
    (d / "logo.png").write_bytes(png((1, 2, 3), size=6))
    return d


def fill_spec(inputs_folder, **values):
    params = {"table": {"file": str(inputs_folder / "PROJECT.xlsx"), "name_column": "Variable", "value_column": "VA3"},
              "author": "A"}
    if values:
        params["values"] = values
    return {"request": "fill the VA3 site", "summary": "s", "mode": "content", "ops": [
        {"op": "fill_variables", "params": params}]}


def test_a_typed_fill_value_must_be_the_users_and_a_table_value_needs_no_trace(inputs_folder):
    batch = _Batch(sorted(inputs_folder.glob("*.docx")), None)
    assert check_submission(fill_spec(inputs_folder), batch, ["fill the VA3 site from PROJECT.xlsx"]) == []
    errs = check_submission(fill_spec(inputs_folder, SITE_NAME="Loudoun"), batch, ["fill the VA3 site"])
    assert errs and "never gave them verbatim" in errs[0] and "Loudoun" in errs[0]
