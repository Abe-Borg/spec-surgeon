"""Tests for the fixes taken from the 2026-10-14 external audit
(docs/audits/2026-10-15_v0.4.0_external_audit_response.md).

One test per finding taken — F1, F2, F4, F5, F6, F8a, F9, F10, F12–F17, the two smaller items and the Strict Open XML
message: the wrong behaviour the audit reproduced, now refused or contained, and where a wrong document used to PASS,
the planted case the verifier catches on its own.
"""
from __future__ import annotations

import contextlib
import hashlib
import io
import os
import sys
import types
import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

from make_fixtures import _styleref_field, build
from spec_surgeon import batch as batch_mod
from spec_surgeon import guard
from spec_surgeon import render as render_mod
from spec_surgeon import advisor as advisor_module
from spec_surgeon.advisor import advise, advise_batch
from spec_surgeon.batch import FileResult, check_inputs_outside, output_names, run_batch, verified_bytes
from spec_surgeon.codegen import run_in_sandbox, trial
from spec_surgeon.compiler import _Batch, check_submission
from spec_surgeon.editspec import apply_spec, plan_spec
from spec_surgeon.ooxml import block_full_text, paragraphs
from spec_surgeon.ops import JobException
from spec_surgeon.ops.hf_common import effective_refs
from spec_surgeon.ops.replace_text import field_start_depths
from spec_surgeon.package import DocxPackage, UnsupportedFormat, W
from spec_surgeon.planner import plan_file
from spec_surgeon.profile import profile_batch
from spec_surgeon.render import libreoffice, word_com
from spec_surgeon.report import explain_failure
from spec_surgeon.verify import run_suite
from spec_surgeon.verify.reader import VPackage
from spec_surgeon.verify.template import _v_field_runs, slots_of
from test_codegen import GOOD_OP
from test_gui import qt, window  # noqa: F401 — the Qt fixture and the window helper
from test_compiler import FakeClient, resp, tool


def _spec(mode, ops):
    return {"request": "x", "summary": "x", "assumptions": [], "questions": [], "applies_to": {"or": "all"},
            "mode": mode, "ops": ops}


REPLACE = _spec("content", [{"op": "replace_text", "params": {"find": "FM Global", "replace": "FM Approved"}}])
NUMBER = lambda option: _spec("formatting", [{"op": "fix_page_numbering", "params": {"option": option}}])


def _rewrite(data: bytes, part: str, fn) -> bytes:
    """``data`` with ``fn(root)`` applied to the XML of ``part``."""
    zin = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            b = zin.read(info.filename)
            if info.filename == part:
                root = etree.fromstring(b)
                fn(root)
                b = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            zout.writestr(info, b)
    return out.getvalue()


def _title(path: Path) -> str:
    return block_full_text(paragraphs(DocxPackage.from_path(path).xml("word/document.xml"))[1])


def _fld(paragraph, kind):
    r = paragraph.add_run()
    fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), kind); r._r.append(fc)


def _instr(paragraph, text):
    r = paragraph.add_run()
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = text; r._r.append(it)


# ---- F1. inputs that share a name ---------------------------------------------------------------------------------

def test_outputs_and_records_stay_distinct_when_inputs_share_a_name(tmp_path):
    files = []
    for folder, num, title in (("a", "21 00 01", "ALPHA"), ("b", "21 00 02", "BRAVO")):
        d = tmp_path / folder
        d.mkdir()
        (d / "Spec.docx").write_bytes(build(num, title))
        files.append(d / "Spec.docx")
    (tmp_path / "c").mkdir()
    files.append(tmp_path / "c" / "Spec (1).docx")
    files[-1].write_bytes(build("21 00 03", "CHARLIE"))
    # a name an input really has is never handed to another input; the first keeps its own, the next gets a free one
    assert [output_names(files)[str(f)] for f in files] == ["Spec.docx", "Spec (2).docx", "Spec (1).docx"]
    assert list(output_names([Path("x/SPEC.docx"), Path("y/spec.docx")]).values()) == ["SPEC.docx", "spec (1).docx"]
    before = {f: _title(f) for f in files}
    b = run_batch([str(f) for f in files], "edit-spec", {}, tmp_path / "out", workers=1, renderer="none",
                  do_render=False, in_place=True, spec=REPLACE)
    assert [f.status for f in b.files] == ["PASS"] * 3
    assert [f.out_name for f in b.files] == ["Spec.docx", "Spec (2).docx", "Spec (1).docx"]
    assert [f.delivered_to for f in b.files] == [str(f) for f in files]
    assert {f: _title(f) for f in files} == before                 # every original is still its own document
    assert sorted(p.name for p in (tmp_path / "out" / "records").glob("*.record.json")) == \
        ["Spec (1).record.json", "Spec (2).record.json", "Spec.record.json"]
    assert len(list((tmp_path / "out" / "plans").glob("*.plan.json"))) == 3


# ---- F2. an input inside the results folder -------------------------------------------------------------------

def test_an_input_inside_the_results_folder_is_refused_before_anything_is_written(tmp_path):
    results = tmp_path / "results"
    src = results / "outputs" / "Spec.docx"
    src.parent.mkdir(parents=True)
    src.write_bytes(build("21 00 04", "DELTA"))
    md5 = hashlib.md5(src.read_bytes()).hexdigest()
    with pytest.raises(ValueError, match="inside the results folder"):
        run_batch([str(src)], "edit-spec", {}, results, workers=1, renderer="none", do_render=False, spec=REPLACE)
    assert hashlib.md5(src.read_bytes()).hexdigest() == md5
    assert sorted(p.name for p in results.iterdir()) == ["outputs"]       # nothing created
    check_inputs_outside([tmp_path / "in" / "x.docx"], results)             # outside: fine
    with pytest.raises(ValueError):
        check_inputs_outside([results / "deeper" / "x.docx"], results)


# ---- F8a. verify before writing; delivery bound to the verified bytes ------------------------------------------

def test_a_verifier_crash_leaves_nothing_in_outputs(tmp_path, monkeypatch):
    f = tmp_path / "A.docx"
    f.write_bytes(build("21 00 05", "ECHO"))

    def boom(*a, **k):
        raise RuntimeError("injected verifier crash")
    monkeypatch.setattr(batch_mod, "run_suite", boom)
    b = run_batch([str(f)], "edit-spec", {}, tmp_path / "out", workers=1, renderer="none", do_render=False, spec=REPLACE)
    fr = b.files[0]
    assert fr.status == "EXCEPTION" and fr.exception["code"] == "crash" and "injected" in fr.exception["reason"]
    assert not list((tmp_path / "out" / "outputs").iterdir()) and not (tmp_path / "out" / "failed").exists()


def test_delivery_is_bound_to_the_bytes_the_suite_verified(tmp_path, monkeypatch):
    f = tmp_path / "A.docx"
    f.write_bytes(build("21 00 06", "FOXTROT"))
    md5 = hashlib.md5(f.read_bytes()).hexdigest()
    real = batch_mod._process_one

    def tampered(*a, **k):
        r = real(*a, **k)
        Path(r["output"]).write_bytes(build("99 99 99", "SOMETHING ELSE"))  # replaced after verification
        return r
    monkeypatch.setattr(batch_mod, "_process_one", tampered)
    for in_place in (True, False):
        out = tmp_path / ("in_place" if in_place else "copies")
        b = run_batch([str(f)], "edit-spec", {}, out, workers=1, renderer="none", do_render=False,
                      in_place=in_place, spec=REPLACE)
        fr = b.files[0]
        v9 = fr.record.get("V9")
        assert fr.status == "FAIL" and v9.status == "FAIL" and "not the one the suite verified" in v9.summary
        assert fr.delivered_to is None and (out / "failed" / "A.docx").is_file() and not (out / "outputs" / "A.docx").exists()
        assert hashlib.md5(f.read_bytes()).hexdigest() == md5                 # the original untouched
    fr2 = FileResult(source=f, output=f, output_md5="0" * 32)
    with pytest.raises(guard.GuardError, match="not the one the suite verified"):
        verified_bytes(fr2)


# ---- F5. a cover footer the body inherits -----------------------------------------------------------------------

def test_cover_unnumbered_refuses_a_cover_footer_the_body_inherits():
    pkg = DocxPackage(build("21 13 13", "WET-PIPE SPRINKLER SYSTEMS", body_footer_linked=True))
    assert pkg.sections()[1].footer_refs == {}                                 # "Link to Previous": no reference
    assert effective_refs(pkg)[1][("footer", "default")] == "word/footer1.xml"   # … so Word shows the cover's
    with pytest.raises(JobException) as e:
        plan_spec(pkg, NUMBER("cover-unnumbered"))
    assert e.value.code == "shared-part" and "inherited" in e.value.reason and "Link to Previous" in e.value.reason
    # a body with a footer of its own still plans, applies and verifies
    pkg = DocxPackage(build("21 13 13", "WET-PIPE SPRINKLER SYSTEMS"))
    plan, frags = plan_spec(pkg, NUMBER("cover-unnumbered"))
    out = apply_spec(pkg, frags).output_bytes()
    assert run_suite(pkg.data, out, plan.to_dict(), "a.docx", "b.docx").passed


# ---- F10. every restart after the cover --------------------------------------------------------------------------

def test_continuous_removes_every_restart_after_the_cover():
    pkg = DocxPackage(build("21 13 13", "WET-PIPE SPRINKLER SYSTEMS", appendix_restart=True))
    assert [s.pg_num_start for s in pkg.sections()] == [None, 1, 1]
    plan, frags = plan_spec(pkg, NUMBER("continuous"))
    a = plan.property_assertions[0]
    assert a["kind"] == "section_props" and a["sections"] == [1, 2] and a["expected_sections"] == 2
    assert any("restarted" in fl for fl in plan.flags)
    out = apply_spec(pkg, frags).output_bytes()
    assert run_suite(pkg.data, out, plan.to_dict(), "a.docx", "b.docx").passed
    assert [s.pg_num_start for s in DocxPackage(out).sections()] == [None, None, None]

    def keep_appendix_restart(root):                       # the output the old op produced
        sp = root.find(W + "body").find(W + "sectPr")
        pg = etree.SubElement(sp, W + "pgNumType")
        pg.set(W + "start", "1")
    planted = _rewrite(out, "word/document.xml", keep_appendix_restart)
    rec = run_suite(pkg.data, planted, plan.to_dict(), "a.docx", "b.docx")
    assert not rec.passed and "section 2" in rec.get("V10").summary
    # two sections: as before
    pkg = DocxPackage(build("21 13 13", "WET-PIPE SPRINKLER SYSTEMS"))
    plan, frags = plan_spec(pkg, NUMBER("continuous"))
    assert plan.property_assertions[0]["sections"] == [1] and not plan.flags
    assert run_suite(pkg.data, apply_spec(pkg, frags).output_bytes(), plan.to_dict(), "a.docx", "b.docx").passed


# ---- F9. a field that spans paragraphs ----------------------------------------------------------------------------

def test_replace_text_refuses_a_toc_entry_in_a_later_paragraph_of_the_field():
    doc = Document()
    doc.add_paragraph("CONTENTS")
    p = doc.add_paragraph()
    _fld(p, "begin"); _instr(p, ' TOC \\o "1-3" \\h '); _fld(p, "separate"); p.add_run("PART 1 GENERAL\t1")
    doc.add_paragraph().add_run("NFPA 14 STANDPIPES\t3")              # a TOC entry: cached field result
    p = doc.add_paragraph()
    p.add_run("PART 3 EXECUTION\t5"); _fld(p, "end")
    doc.add_paragraph("Install standpipes per NFPA 14.")
    buf = io.BytesIO()
    doc.save(buf)
    pkg = DocxPackage(buf.getvalue())
    assert field_start_depths(paragraphs(pkg.xml("word/document.xml"))) == [0, 0, 1, 1, 0]
    with pytest.raises(JobException) as e:
        plan_spec(pkg, _spec("content", [{"op": "replace_text", "params": {"find": "NFPA 14", "replace": "NFPA 13"}}]))
    assert e.value.code == "field-result" and "earlier paragraph" in e.value.reason
    plan, frags = plan_spec(pkg, _spec("content", [{"op": "replace_text", "params": {
        "find": "standpipes per NFPA 14", "replace": "standpipes per NFPA 13"}}]))
    assert [(d.part, d.block_index) for d in plan.diffs] == [("word/document.xml", 4)]
    assert run_suite(pkg.data, apply_spec(pkg, frags).output_bytes(), plan.to_dict(), "a.docx", "b.docx").passed


# ---- F4. a slot a field shows -------------------------------------------------------------------------------------

def _template_with_field_slots(tmp_path: Path) -> Path:
    doc = Document()
    s0 = doc.sections[0]
    h = s0.header
    h.is_linked_to_previous = False
    p = h.paragraphs[0]
    p.add_run("SECTION {{SECTION_NUMBER}} – ")
    _styleref_field(p, "{{SECTION_TITLE}}")                 # the title typed as a STYLEREF field's cached result
    p = h.add_paragraph()
    fs = OxmlElement("w:fldSimple"); fs.set(qn("w:instr"), " DOCPROPERTY Title ")
    r = OxmlElement("w:r"); rpr = OxmlElement("w:rPr"); hl = OxmlElement("w:highlight"); hl.set(qn("w:val"), "yellow")
    rpr.append(hl); r.append(rpr); t = OxmlElement("w:t"); t.text = "SECTION_TITLE"; r.append(t); fs.append(r)
    p._p.append(fs)                                         # a highlighted slot inside a simple field
    f = s0.footer
    f.is_linked_to_previous = False
    f.paragraphs[0].add_run("{{SECTION_NUMBER}} - page")
    doc.add_paragraph("TEMPLATE BODY")
    tpl = tmp_path / "field_slots_template.docx"
    buf = io.BytesIO()
    doc.save(buf)
    tpl.write_bytes(buf.getvalue())
    return tpl


def test_a_slot_shown_by_a_field_is_refused_on_both_sides(tmp_path):
    tpl = _template_with_field_slots(tmp_path)
    fills = {"SECTION_NUMBER": {"rule": "regex", "source": "footer.default", "pattern": "^(\\d{2} \\d{2} \\d{2})"},
             "SECTION_TITLE": {"rule": "paragraph_after", "source": "document",
                               "anchor_regex": "^SECTION \\d{2} \\d{2} \\d{2}$"}}
    pkg = DocxPackage(build("21 13 13", "WET-PIPE SPRINKLER SYSTEMS"))
    for markers in (["braces"], ["highlight"]):
        spec = _spec("untracked", [{"op": "apply_header_footer_template", "select": {"sections": "all"},
                                    "params": {"template": str(tpl), "template_sections": {"cover": 0, "body": 0},
                                               "slot_markers": markers, "slot_fills": fills}}])
        with pytest.raises(JobException) as e:
            plan_spec(pkg, spec)
        assert e.value.code == "slot-in-field" and "SECTION_TITLE" in e.value.reason
    root = VPackage(tpl.read_bytes()).xml("word/header1.xml")
    pars = list(root.iter(W + "p"))
    for par, markers in ((pars[0], ["braces"]), (pars[1], ["highlight"])):
        with pytest.raises(ValueError, match="inside a field"):
            slots_of(par, markers, ["yellow"], None, _v_field_runs(root))
        assert "SECTION_TITLE" in [s["name"] for s in slots_of(par, markers, ["yellow"], None, None)]


# ---- F12. spacing in lines ----------------------------------------------------------------------------------------

def test_spacing_given_in_lines_refuses_a_point_value():
    def after_lines(root):
        p = list(root.iter(W + "p"))[6]
        ppr = p.find(W + "pPr")
        if ppr is None:
            ppr = etree.Element(W + "pPr"); p.insert(0, ppr)
        etree.SubElement(ppr, W + "spacing").set(W + "afterLines", "200")
    pkg = DocxPackage(_rewrite(build("21 00 12", "GOLF"), "word/document.xml", after_lines))
    op = lambda spacing: _spec("formatting", [{"op": "set_paragraph_props", "select": {"scope": ["document"]},
                                               "params": {"spacing": spacing, "tracked": False}}])
    with pytest.raises(JobException) as e:
        plan_spec(pkg, op({"after": 120}))
    assert e.value.code == "conflict-property" and "afterLines=200" in e.value.reason
    plan, frags = plan_spec(pkg, op({"after": 120, "afterLines": 0}))        # the spec says what the lines become
    assert run_suite(pkg.data, apply_spec(pkg, frags).output_bytes(), plan.to_dict(), "a.docx", "b.docx").passed

    def style_lines(root):
        for st in root.iter(W + "style"):
            if st.get(W + "styleId") == "Normal":
                ppr = st.find(W + "pPr")
                if ppr is None:
                    ppr = etree.SubElement(st, W + "pPr")
                etree.SubElement(ppr, W + "spacing").set(W + "beforeLines", "100")
    pkg = DocxPackage(_rewrite(build("21 00 12", "GOLF"), "word/styles.xml", style_lines))
    with pytest.raises(JobException) as e:
        plan_spec(pkg, op({"before": 120}))
    assert e.value.code == "conflict-property" and "style 'Normal'" in e.value.reason


# ---- F13. the compile gate and the pattern's other files -----------------------------------------------------

def test_the_compile_gate_tries_every_file_of_a_pattern_before_rejecting(tmp_path):
    def nfpa14(root):
        for t in root.iter(W + "t"):
            if t.text and "NFPA 13" in t.text:
                t.text = t.text.replace("NFPA 13", "NFPA 14")
    (tmp_path / "A.docx").write_bytes(_rewrite(build("21 00 13", "HOTEL"), "word/document.xml", nfpa14))
    (tmp_path / "B.docx").write_bytes(build("21 00 14", "INDIA"))
    spec = _spec("content", [{"op": "replace_text", "params": {"find": "NFPA 13", "replace": "NFPA 13 (2025)"}}])
    for order in (["A.docx", "B.docx"], ["B.docx", "A.docx"]):
        b = _Batch([tmp_path / n for n in order], None)
        assert len(b.prof["patterns"]) == 1
        assert check_submission(spec, b, ["NFPA 13 (2025)"]) == []
    nowhere = _spec("content", [{"op": "replace_text", "params": {"find": "NFPA 99", "replace": "NFPA 13 (2025)"}}])
    errs = check_submission(nowhere, _Batch([tmp_path / "A.docx", tmp_path / "B.docx"], None), ["NFPA 13 (2025)"])
    assert errs and "none of this pattern's 2 files plans" in errs[1]
    # two folders, one file name: the gate goes by path, so the file the name map shadows is still tried
    for d in ("a", "b"):
        (tmp_path / d).mkdir()
    (tmp_path / "a" / "Spec.docx").write_bytes(build("21 00 14", "INDIA"))                        # has NFPA 13
    (tmp_path / "b" / "Spec.docx").write_bytes(_rewrite(build("21 00 13", "HOTEL"), "word/document.xml", nfpa14))
    b = _Batch([tmp_path / "a" / "Spec.docx", tmp_path / "b" / "Spec.docx"], None)
    assert b.prof["patterns"][0]["paths"] == [str(tmp_path / "a" / "Spec.docx"), str(tmp_path / "b" / "Spec.docx")]
    assert check_submission(spec, b, ["NFPA 13 (2025)"]) == []


# ---- F14. a file the profiler cannot read -------------------------------------------------------------------------

def test_profile_batch_takes_the_next_file_when_the_first_cannot_be_profiled(tmp_path):
    def bad_margin(root):
        for m in root.iter(W + "pgMar"):
            m.set(W + "left", "abc")
    (tmp_path / "A.docx").write_bytes(_rewrite(build("21 00 15", "JULIET"), "word/document.xml", bad_margin))
    (tmp_path / "B.docx").write_bytes(build("21 00 16", "KILO"))
    pb = profile_batch([tmp_path / "A.docx", tmp_path / "B.docx"])
    [pat] = pb["patterns"]
    assert pat["files"] == ["B.docx", "A.docx"] and "error" not in pat["profile"]
    assert pat["unprofiled"] and pat["unprofiled"][0].startswith("A.docx")
    assert _Batch([tmp_path / "A.docx", tmp_path / "B.docx"], None).representatives() == [tmp_path / "B.docx"]
    only = profile_batch([tmp_path / "A.docx"])                                # no file profiles: said, not raised
    assert only["patterns"][0]["profile"]["error"]


# ---- F15. advice and an unreadable file ---------------------------------------------------------------------------

def test_advice_survives_an_unreadable_file(tmp_path, monkeypatch):
    folder = tmp_path / "in"
    folder.mkdir()
    (folder / "A text.docx").write_text("not a docx", encoding="utf-8")
    (folder / "B single.docx").write_bytes(build("21 00 17", "LIMA", single_section=True))
    out = tmp_path / "out"
    b = run_batch([str(folder)], "edit-spec", {}, out, workers=1, renderer="none", do_render=False,
                  spec=NUMBER("continuous"))
    assert [f.exception["code"] for f in b.files] == ["crash", "structure"]
    rec = {"source": str(folder / "A text.docx"), "exception": b.files[0].exception, "spec": NUMBER("continuous")}
    adv = advise(rec, client=object())                                         # no model call for an unreadable file
    assert adv.kind == "recommendation" and "cannot be opened" in adv.text and not adv.errors

    def no_client():                                                           # no API key, no anthropic package
        raise RuntimeError("no API key")
    monkeypatch.setattr(advisor_module, "_client", no_client)
    adv = advise(rec)                                                          # … and none is built for it either
    assert adv.kind == "recommendation" and "cannot be opened" in adv.text
    fake = FakeClient(resp(tool("submit_recommendation", {"text": "One section: nothing to make continuous."})))
    advs = advise_batch(out, client=fake)
    assert [a.kind for a in advs] == ["recommendation", "recommendation"]       # the second one still got advice
    assert "One section" in advs[1].text and len(fake.requests) == 1


# ---- F16, F17. the generated-op sandbox ---------------------------------------------------------------------------

def test_the_sandbox_takes_a_relative_work_folder(tmp_path, monkeypatch):
    f = tmp_path / "A.docx"
    f.write_bytes(build("21 00 18", "MIKE"))
    monkeypatch.chdir(tmp_path)
    sb = run_in_sandbox(GOOD_OP, Path("A.docx"), {}, {"twips": 900}, Path("rel") / "sb")
    assert sb.ok and sb.output and sb.output.is_file(), (sb.code, sb.error)


def test_a_generated_op_that_breaks_the_package_is_a_failed_trial_not_a_crash(tmp_path):
    f = tmp_path / "A.docx"
    f.write_bytes(build("21 00 19", "NOVEMBER"))
    broken = GOOD_OP.replace("def apply(self, pkg, frag):",
                             "def apply(self, pkg, frag):\n        pkg.remove_member('word/document.xml')\n"
                             "        return\n    def _old_apply(self, pkg, frag):")
    assert broken != GOOD_OP
    t = trial(broken, f, {}, {"twips": 900}, tmp_path / "sb")
    assert t.passed is False and "could not check the output" in t.problems[0] and t.record is None


# ---- F6. Word detection and renderer failures ---------------------------------------------------------------------

def test_word_is_picked_only_when_it_is_installed(monkeypatch):
    monkeypatch.setattr(sys, "platform", "win32")
    for name in ("win32com", "win32com.client", "pythoncom"):
        monkeypatch.setitem(sys.modules, name, types.ModuleType(name))

    class NoWord:                                   # pywin32 present (the installer bundles it), Word absent
        HKEY_CLASSES_ROOT = 0

        @staticmethod
        def OpenKey(root, key):
            raise OSError("key not found")
    monkeypatch.setitem(sys.modules, "winreg", NoWord)
    monkeypatch.setattr(libreoffice, "find_soffice", lambda: None)
    assert word_com.available() is False
    assert render_mod.pick("auto") is None and render_mod.pick("word") is None

    class Word(NoWord):
        @staticmethod
        def OpenKey(root, key):
            assert key == word_com.WORD_PROGID_KEY
            return contextlib.nullcontext()
    monkeypatch.setitem(sys.modules, "winreg", Word)
    assert word_com.available() is True and render_mod.pick("auto") == "word"


def test_a_renderer_failure_is_this_files_failed_render_never_a_crashed_batch(tmp_path, monkeypatch):
    msgs = []
    assert render_mod.render_pdfs([tmp_path / "x.docx"], tmp_path / "pdf", "word", log=msgs.append) == {}
    assert msgs and "rendering with word failed" in msgs[0]

    def convert(soffice, files, out_dir, profile_dir, timeout):      # soffice dies on one file of a chunk
        if any(f.name == "bad.docx" for f in files):
            raise RuntimeError("soffice crashed")
        for f in files:
            (out_dir / (f.stem + ".pdf")).write_bytes(b"%PDF")
    monkeypatch.setattr(libreoffice, "find_soffice", lambda: "soffice")
    monkeypatch.setattr(libreoffice, "_convert_batch", convert)
    good, bad = tmp_path / "good.docx", tmp_path / "bad.docx"
    good.write_bytes(b"x")
    bad.write_bytes(b"x")
    msgs = []
    out = render_mod.render_pdfs([good, bad], tmp_path / "pdf2", "libreoffice", workers=1, log=msgs.append)
    assert set(out) == {good} and any("bad.docx" in m for m in msgs)       # one bad file costs its own PDF only

    f = tmp_path / "A.docx"
    f.write_bytes(build("21 00 20", "OSCAR"))

    def always_fails(*a, **k):
        raise RuntimeError("soffice crashed")
    monkeypatch.setattr(libreoffice, "_convert_batch", always_fails)
    b = run_batch([str(f)], "edit-spec", {}, tmp_path / "out", workers=1, renderer="libreoffice", do_render=True,
                  spec=REPLACE)
    fr = b.files[0]
    assert b.renderer == "libreoffice" and fr.status == "FAIL"
    assert fr.record.get("V6").status == "FAIL" and "render failed" in fr.record.get("V6").summary
    assert (tmp_path / "out" / "failed" / "A.docx").is_file() and (tmp_path / "out" / "BATCH_REPORT.md").is_file()


# ---- the two smaller items ------------------------------------------------------------------------------------------

def test_the_guard_sees_a_content_change_that_keeps_size_and_mtime(tmp_path):
    f = tmp_path / "A.docx"
    f.write_bytes(b"0123456789")
    ident = guard.acquire(f)
    st = f.stat()
    f.write_bytes(b"0123456780")
    os.utime(f, ns=(st.st_atime_ns, st.st_mtime_ns))
    assert f.stat().st_size == ident.size and f.stat().st_mtime_ns == ident.mtime_ns
    assert "content changed" in (guard.drifted(ident) or "")
    with pytest.raises(guard.GuardError, match="content changed"):
        guard.write_back(ident, b"new")
    assert f.read_bytes() == b"0123456780"


def test_changing_the_results_folder_after_a_preview_disables_run(qt, tmp_path):
    from spec_surgeon.gui import Session
    folder = tmp_path / "Specs"
    folder.mkdir()
    (folder / "21 00 21 – PAPA.docx").write_bytes(build("21 00 21", "PAPA"))
    s = Session(out_root=tmp_path / "out")
    s.load_spec(Path(__file__).resolve().parent.parent / "examples" / "specs" / "remove-all-headers-footers.json")
    w, ui = window(qt, s)
    ui["vars"]["folder"].set(str(folder))
    ui["vars"]["out"].set(str(tmp_path / "out"))
    ui["vars"]["render"].set(False)
    ui["buttons"]["preview"].click()
    ui["wait"]()
    assert ui["buttons"]["run"].isEnabled()
    ui["vars"]["out"].set(str(tmp_path / "elsewhere"))
    assert not ui["buttons"]["run"].isEnabled() and w.do_run() is False
    ui["vars"]["out"].set(str(tmp_path / "out"))
    assert ui["buttons"]["run"].isEnabled()


# ---- Strict Open XML ------------------------------------------------------------------------------------------------

def _strict(data: bytes) -> bytes:
    """The fixture re-saved as ISO Strict: every transitional namespace moved to purl.oclc.org."""
    ns = {"http://schemas.openxmlformats.org/wordprocessingml/2006/main": "http://purl.oclc.org/ooxml/wordprocessingml/main",
          "http://schemas.openxmlformats.org/officeDocument/2006/relationships":
              "http://purl.oclc.org/ooxml/officeDocument/relationships"}
    zin = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            b = zin.read(info.filename)
            if info.filename.endswith((".xml", ".rels")):
                s = b.decode("utf-8")
                for a, c in ns.items():
                    s = s.replace(a, c)
                b = s.encode("utf-8")
            zout.writestr(info, b)
    return out.getvalue()


def test_a_strict_open_xml_file_is_set_aside_with_a_plain_message(tmp_path):
    f = tmp_path / "STRICT.docx"
    f.write_bytes(_strict(build("21 00 22", "QUEBEC")))
    with pytest.raises(UnsupportedFormat, match="Strict Open XML"):
        DocxPackage.from_path(f).main_part()
    fp = plan_file(f, REPLACE)
    assert fp.status == "exception" and fp.exception["code"] == "unsupported-format" and "Save As" in fp.exception["reason"]
    b = run_batch([str(f)], "edit-spec", {}, tmp_path / "out", workers=1, renderer="none", do_render=False, spec=REPLACE)
    assert b.files[0].status == "EXCEPTION" and b.files[0].exception["code"] == "unsupported-format"
    assert explain_failure("unsupported-format", "x")[0] == "Strict Open XML Document"
    assert DocxPackage(build("21 00 22", "QUEBEC")).main_part() == "word/document.xml"      # transitional: as before
