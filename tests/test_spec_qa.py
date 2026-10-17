"""Batch-wide read-only audits retain every occurrence and never declare partial scans clean."""
import hashlib
import io
import json

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from spec_surgeon.cli import main
from spec_surgeon.spec_qa import AUDITS, audit_batch, audit_markdown


def write_doc(path, number, title, *body, header=""):
    d = Document()
    d.add_paragraph("SECTION " + number)
    d.add_paragraph(title)
    for line in body:
        p = d.add_paragraph()
        for piece in (line[:len(line)//2], line[len(line)//2:]):
            p.add_run(piece)
    d.sections[0].header.paragraphs[0].text = header
    d.save(path)
    return path


def test_all_occurrences_section_titles_and_no_incoming_references(tmp_path):
    a = write_doc(tmp_path / "a.docx", "21 13 19", "SPRINKLERS", "[ ] [wet] <Insert name> XXXX TBD TBD {{SITE}}",
                  'See Section 21 05 00 “COMMON WORK RESULTS”; Section 99 99 99.',
                  'See Section 21 06 00 - WRONG TITLE.', "NFPA 13, 2022 Edition; ASTM A795/A795M-21; UL 199; ASME B31.9-2020; FM 1635, 2023.")
    b = write_doc(tmp_path / "b.docx", "21 05 00", "COMMON WORK RESULTS")
    c = write_doc(tmp_path / "c.docx", "21 06 00", "SCHEDULES")
    before = {p: hashlib.sha256(p.read_bytes()).hexdigest() for p in (a, b, c)}
    report = audit_batch([a, b, c])
    assert report["report_only"] and report["complete"]
    f = report["files"][0]
    assert len(f["placeholders"]) == 7
    assert {r["section"] for r in f["placeholders"]} == {"21 13 19"}
    assert [r["status"] for r in f["cross_references"]] == ["matched", "missing_section", "title_mismatch"]
    assert [s["number"] for s in report["unreferenced_sections"]] == ["21 13 19"]
    standards = {r["family"]: r for r in f["standards"]}
    assert standards["NFPA"]["edition"] == "2022"
    assert standards["ASTM"]["edition"] == "21"
    assert standards["ASME"]["edition"] == "2020"
    assert standards["UL"]["status"] == "edition_missing"
    assert standards["FM"]["edition"] == "2023"
    assert f["readiness"]["status"] == "not_ready"
    assert "title_mismatch" in audit_markdown(report)
    assert before == {p: hashlib.sha256(p.read_bytes()).hexdigest() for p in (a, b, c)}


def test_issue_readiness_includes_style_hidden_highlight_and_revisions(tmp_path):
    path = tmp_path / "markup.docx"
    d = Document()
    d.add_paragraph("SECTION 21 13 19")
    d.add_paragraph("SPRINKLERS")
    p = d.add_paragraph("hidden and highlighted")
    p.style.font.hidden = True
    h = OxmlElement("w:highlight")
    h.set(qn("w:val"), "yellow")
    p.runs[0]._r.get_or_add_rPr().append(h)
    ins = OxmlElement("w:ins")
    ins.set(qn("w:id"), "7")
    ins.set(qn("w:author"), "Reviewer")
    ins.append(p.runs[0]._r)
    p._p.append(ins)
    d.save(path)
    f = audit_batch([path])["files"][0]
    assert f["readiness"]["status"] == "not_ready"
    assert {"hidden_text", "highlights", "tracked_changes"} <= set(f["readiness"]["counts"])


def test_identity_variants_keep_values_and_locations_across_files(tmp_path):
    a = write_doc(tmp_path / "a.docx", "21 05 00", "COMMON", "Project Name: Harbor Tower", "Project No.: H-123",
                  header="Harbor Tower\t10/17/2026")
    b = write_doc(tmp_path / "b.docx", "21 13 19", "SPRINKLERS", "Project Name: Harbour Tower", "Project No.: H-124",
                  header="Harbour Tower\tOctober 17, 2026")
    report = audit_batch([a, b], identity={"project_name": "Harbor Tower", "project_number": "H-123"})
    variants = report["identity_variants"]
    assert variants["project_name"]["has_variants"]
    assert "Harbour Tower" in variants["project_name"]["different_from_expected"]
    assert {"H-123", "H-124"} <= set(variants["project_number"]["variants"])
    assert variants["date"]["has_variants"]
    assert all(r["part"] and r["file"] for values in variants["project_name"]["variants"].values() for r in values)


def test_unreadable_file_and_duplicate_sections_are_explicit(tmp_path):
    a = write_doc(tmp_path / "a.docx", "21 05 00", "COMMON", "See Section 21 13 19.")
    b = write_doc(tmp_path / "b.docx", "21 13 19", "SPRINKLERS")
    c = write_doc(tmp_path / "c.docx", "21 13 19", "OTHER TITLE")
    corrupt = tmp_path / "bad.docx"
    corrupt.write_bytes(b"not a ZIP")
    report = audit_batch([a, b, c, corrupt])
    assert not report["complete"] and len(report["errors"]) == 1
    assert report["files"][0]["cross_references"][0]["status"] == "ambiguous_target"
    assert "21 13 19" in report["duplicate_sections"]
    assert "batch contains unreadable files" in report["files"][1]["readiness"]["limitations"]


def test_cli_selects_report_and_never_overwrites_a_docx(tmp_path, capsys):
    src = write_doc(tmp_path / "a.docx", "21 05 00", "COMMON", "TBD")
    original = src.read_bytes()
    output = tmp_path / "report.json"
    assert main(["audit", str(src), "--audit", "placeholders", "--out", str(output)]) == 0
    report = json.loads(output.read_text())
    assert report["audits"] == ["placeholders"] and "readiness" not in report["files"][0]
    assert main(["audit", str(src), "--out", str(src)]) == 2
    assert src.read_bytes() == original
    assert main(["audit", str(tmp_path / "missing")]) == 2


def test_filename_only_and_missing_titles_never_claim_a_title_match(tmp_path):
    a = write_doc(tmp_path / "a.docx", "21 05 00", "COMMON", "See Section 21 13 19.")
    b = tmp_path / "21 13 19 - SPRINKLERS.docx"
    Document().save(b)
    report = audit_batch([a, b])
    assert report["files"][0]["cross_references"][0]["status"] == "target_heading_unconfirmed"


def test_unknown_identity_and_bad_options_are_explicit(tmp_path):
    p = write_doc(tmp_path / "a.docx", "21 05 00", "COMMON")
    assert audit_batch([p])["identity_variants"]["project_name"]["missing_in_files"] == [str(p)]
    with pytest.raises(ValueError):
        audit_batch([p], audits=["fix_editions"])
    with pytest.raises(ValueError):
        audit_batch([p], identity={"project_name": 123})


def test_readiness_only_retains_the_actual_placeholder_occurrences(tmp_path):
    p = write_doc(tmp_path / "a.docx", "21 05 00", "COMMON", "<Insert name> TBD")
    report = audit_batch([p], audits=["readiness"])
    findings = report["files"][0]["readiness"]["findings"]
    assert [f["text"] for f in findings if f["type"] == "placeholders"] == ["<Insert name>", "TBD"]
    assert "<Insert name>" in audit_markdown(report)


def test_a_shared_header_does_not_make_a_section_reference_itself(tmp_path):
    path = write_doc(tmp_path / "combined.docx", "21 05 00", "COMMON", "SECTION 21 13 19", "SPRINKLERS",
                     header="SECTION 21 05 00")
    report = audit_batch([path])
    assert {s["number"] for s in report["unreferenced_sections"]} == {"21 05 00", "21 13 19"}


def test_compact_and_spaced_section_identities_match_without_editing(tmp_path):
    a = write_doc(tmp_path / "a.docx", "210500", "COMMON", 'See Section 21 13 19 “SPRINKLERS”.')
    b = write_doc(tmp_path / "b.docx", "211319", "SPRINKLERS", 'See Section 210500 “COMMON”.')
    report = audit_batch([a, b])
    assert all(f["cross_references"][0]["status"] == "matched" for f in report["files"])
    assert report["unreferenced_sections"] == []


@pytest.mark.parametrize("date", ["Oct. 2026", "17 Oct 2026", "2026.10.17"])
def test_date_candidates_include_common_issue_stamp_spellings(tmp_path, date):
    p = write_doc(tmp_path / "a.docx", "21 05 00", "COMMON", header=date)
    assert date in audit_batch([p])["identity_variants"]["date"]["variants"]


def test_fm_data_sheets_notes_textboxes_and_comments_are_audited(tmp_path):
    from fx_comments import build_commented
    path = tmp_path / "comments.docx"
    path.write_bytes(build_commented(textbox=True))
    report = audit_batch([path])
    assert report["files"][0]["readiness"]["counts"]["comments"] > 0
    fm = write_doc(tmp_path / "fm.docx", "21 05 00", "COMMON", "FM Global Data Sheet 2-0, October 2021.")
    standard = audit_batch([fm])["files"][0]["standards"][0]
    assert standard["designation"] == "2-0" and standard["edition"] == "2021"


def test_desktop_audit_and_exports_are_read_only(tmp_path, monkeypatch):
    from qt_support import close_windows, qt_app_or_skip
    from spec_surgeon.gui import Session
    from spec_surgeon.ui.app import RecordingPrompts, build_window
    app = qt_app_or_skip()
    monkeypatch.setenv("SPEC_SURGEON_CONFIG_DIR", str(tmp_path / "config"))
    monkeypatch.setenv("SPEC_SURGEON_NO_UPDATE_CHECK", "1")
    folder = tmp_path / "specs"
    folder.mkdir()
    path = write_doc(folder / "a.docx", "21 05 00", "COMMON", "TBD")
    original = path.read_bytes()
    output = tmp_path / "report.json"
    prompts = RecordingPrompts(files=[str(path), str(output)])
    w = build_window(Session(folder=folder, out_root=tmp_path / "out"), prompts=prompts)
    try:
        w.act_audit.trigger()
        w._spec_surgeon["wait"]()
        app.processEvents()
        dialog = w.audit_dialog
        assert "TBD" in dialog.text.toPlainText()
        assert not dialog.save("json")
        assert dialog.save("json")
        assert json.loads(output.read_text())["report_only"]
        assert path.read_bytes() == original
        dialog.identity_fields["project_name"].setText("Harbor Tower")
        dialog.recheck_button.click()
        w._spec_surgeon["wait"]()
        assert w.audit_dialog.report["identity_variants"]["project_name"]["expected"] == "Harbor Tower"
    finally:
        close_windows(app, [w])
