"""Final requirements must catch a valid later stage undoing an earlier requirement."""
from __future__ import annotations

import io
import json

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from spec_surgeon import batch
from spec_surgeon import library
from spec_surgeon.chain import ChainError, chain_from_specs, chain_key, run_chain, validate_chain
from spec_surgeon.preview import approval_token, make_chain_preview
from spec_surgeon.verify.outcomes import assert_outcomes


def _document(font="Arial", header_font=None):
    doc = Document()
    doc.styles["Normal"].font.name = font
    doc.add_paragraph("SECTION 21 13 13")
    doc.add_paragraph("Keep every character of this specification.")
    if header_font:
        doc.sections[0].header.paragraphs[0].add_run("Architect header").font.name = header_font
    data = io.BytesIO()
    doc.save(data)
    return data.getvalue()


def _font(name):
    return {"request": f"Use {name}", "summary": f"Use {name}", "mode": "formatting",
            "applies_to": {"or": "all"}, "ops": [
                {"op": "set_font", "select": {"scope": ["document"], "runs": "all"}, "params": {"font": name}},
                {"op": "set_style_font", "select": {"styles": "all"}, "params": {"font": name}},
                {"op": "set_doc_defaults_font", "params": {"font": name}}]}


def _require_font(name="Arial", **kw):
    return {"kind": "effective_font", "value": name, "slots": ["ascii", "hAnsi"], **kw}


def _run(tmp_path, chain):
    source = tmp_path / "source.docx"
    before = _document()
    source.write_bytes(before)
    result = run_chain([str(source)], chain, tmp_path / "run", workers=1, renderer="none", do_render=False)
    assert source.read_bytes() == before
    return source, result


def test_valid_stages_can_fail_final_requirement_without_delivering_or_touching_original(tmp_path):
    chain = chain_from_specs([("Arial", _font("Arial")), ("Courier", _font("Courier New"))],
                             assertions=[_require_font()])
    source, result = _run(tmp_path, chain)
    f = result.files[0]
    assert f.stage_status == ["PASS", "PASS"]
    assert f.status == "FAIL" and f.code == "chain-outcome" and f.delivered_to is None
    assert not (tmp_path / "run" / "outputs").exists()
    assert assert_outcomes(source.read_bytes(), [_require_font()])[0]["status"] == "PASS"
    assert f.outcomes[0]["status"] == "FAIL" and f.outcomes[0]["checked"] > 0
    report = json.loads((tmp_path / "run" / "chain.json").read_text())
    assert report["files"][0]["outcomes"] == f.outcomes
    assert "Final requirements" in (tmp_path / "run" / "CHAIN_REPORT.md").read_text()
    diagnostics = (tmp_path / "run" / "DIAGNOSTICS.md").read_text()
    assert "final requirements failed" in diagnostics and "Courier New" in diagnostics
    manifest = json.loads((tmp_path / "run" / "set_aside.json").read_text())
    assert manifest["files"][0]["stage"] == "chain final requirements"
    assert manifest["files"][0]["source"] == str(source.resolve())


def test_backward_compatible_user_order_can_intentionally_replace_the_earlier_font(tmp_path):
    _, result = _run(tmp_path, chain_from_specs([("Arial", _font("Arial")), ("Courier", _font("Courier New"))]))
    f = result.files[0]
    assert f.status == "PASS" and f.outcomes == []
    assert assert_outcomes(f.output.read_bytes(), [_require_font("Courier New")])[0]["status"] == "PASS"


def test_all_final_requirements_checked_after_later_page_setup_and_persisted(tmp_path):
    margins = {"request": "One-inch side margins", "summary": "One-inch side margins",
               "mode": "formatting", "applies_to": {"or": "all"},
               "ops": [{"op": "set_page_setup", "params": {"margins": {"left": 1440, "right": 1440}}}]}
    assertions = [_require_font(), {"kind": "page_setup", "sections": [0],
                                   "margins": {"left": 1440, "right": 1440}, "orientation": "portrait"}]
    source, result = _run(tmp_path, chain_from_specs([("font", _font("Arial")), ("margins", margins)], assertions=assertions))
    f = result.files[0]
    assert f.status == "PASS" and len(f.outcomes) == 2 and all(o["status"] == "PASS" for o in f.outcomes)
    assert f.output.read_bytes() == f.results[-1].output.read_bytes()


def test_changed_final_requirement_changes_approval_token(tmp_path):
    source = tmp_path / "source.docx"
    source.write_bytes(_document())
    chain = chain_from_specs([("font", _font("Arial"))], assertions=[_require_font()])
    frozen_key = chain_key(chain)
    before = approval_token(frozen_key, [source])
    chain["assertions"][0]["value"] = "Courier New"
    assert frozen_key["assertions"][0]["value"] == "Arial"
    assert approval_token(chain_key(chain), [source]) != before


def test_library_round_trip_retains_requirements_and_preview_cannot_call_failed_outcome_passed(tmp_path):
    source = tmp_path / "source.docx"
    source.write_bytes(_document())
    chain = chain_from_specs([("courier", _font("Courier New"))], assertions=[_require_font()])
    library.save_chain("Keep Arial", chain)
    saved = library.load_chain("Keep Arial")
    assert saved["assertions"] == chain["assertions"]
    pv = make_chain_preview(saved, [str(source)], tmp_path / "preview", renderer="none", do_render=False)
    sample = pv.samples[0]
    assert sample.stages[0]["passed"] is True and sample.passed is False
    assert sample.checks["FINAL"] == "FAIL" and sample.outcomes[0]["status"] == "FAIL"
    assert "final requirement 1 FAIL" in pv.to_text()
    assert pv.to_dict()["samples"][0]["outcomes"] == sample.outcomes


def test_caller_cannot_change_requirements_while_a_stage_runs(tmp_path, monkeypatch):
    chain = chain_from_specs([("courier", _font("Courier New"))], assertions=[_require_font()])
    original = batch.run_batch

    def change_requirements(*args, **kw):
        chain["assertions"][0]["value"] = "Courier New"
        return original(*args, **kw)

    monkeypatch.setattr(batch, "run_batch", change_requirements)
    _, result = _run(tmp_path, chain)
    assert result.files[0].status == "FAIL"
    assert result.chain["assertions"][0]["value"] == "Arial"


def test_final_assertions_refuse_in_place_before_any_file_or_folder_changes(tmp_path):
    source = tmp_path / "source.docx"
    before = _document()
    source.write_bytes(before)
    chain = chain_from_specs([("font", _font("Courier New"))], assertions=[_require_font()])
    with pytest.raises(ChainError, match="copies mode"):
        run_chain([str(source)], chain, tmp_path / "run", in_place=True)
    assert source.read_bytes() == before and not (tmp_path / "run").exists()


@pytest.mark.parametrize("assertions", [
    None, {}, [None], [{"kind": []}], [{"kind": "unknown"}],
    [_require_font(scope=[])], [_require_font(scope=["document", "bogus"])],
    [_require_font(slots=[])], [_require_font(include_paragraph_marks="false")],
    [_require_font(allow_absent_scopes="true")],
    [_require_font(preserve_fonts=["Symbol"])], [_require_font(size_from=[10])],
    [{"kind": "effective_font", "size": True}], [{"kind": "effective_font", "size": float("nan")}],
    [{"kind": "effective_font", "size": 10.1}], [{"kind": "effective_font", "value": ""}],
    [{"kind": "page_setup", "sections": [True], "orientation": "portrait"}],
    [{"kind": "page_setup", "sections": [], "orientation": "portrait"}],
    [{"kind": "page_setup", "margins": {"left": True}}],
    [{"kind": "page_setup", "size": {"w": 12240}}],
    [{"kind": "effective_properties", "run": {"b": "yes"}}],
    [{"kind": "effective_properties", "paragraph": {"spacing": 20}}],
    [{"kind": "effective_properties", "run": {"caps": True, "smallCaps": True}}],
])
def test_malformed_or_unsupported_requirements_refuse_the_chain(assertions):
    chain = chain_from_specs([("font", _font("Arial"))])
    chain["assertions"] = assertions
    with pytest.raises(ChainError):
        validate_chain(chain)


def test_font_requirements_cover_new_headers_and_direct_formatting_overrides():
    records = assert_outcomes(_document(header_font="Courier New"),
                             [_require_font(scope=["document", "headers"])])
    assert records[0]["status"] == "FAIL"
    assert any("header" in e and "Courier New" in e for e in records[0]["errors"])


def test_missing_scope_missing_section_and_corrupt_output_never_pass_vacuously():
    assert assert_outcomes(_document(), [_require_font(scope=["comments"])])[0]["status"] == "FAIL"
    assert assert_outcomes(_document(), [{"kind": "page_setup", "sections": [3], "margins": {"left": 1440}}])[0]["status"] == "FAIL"
    assert assert_outcomes(b"not a zip", [_require_font()])[0]["status"] == "FAIL"


def test_optional_absent_scopes_still_check_present_and_new_parts_and_never_pass_empty_selection():
    assertion = _require_font(scope=["document", "headers", "footnotes"], allow_absent_scopes=True)
    record = assert_outcomes(_document(), [assertion])[0]
    assert record["status"] == "PASS" and record["checked"] > 0
    assert assert_outcomes(_document(header_font="Courier New"), [assertion])[0]["status"] == "FAIL"
    assertion["scope"] = ["footnotes", "comments"]
    assert assert_outcomes(_document(), [assertion])[0]["status"] == "FAIL"


def test_orientation_requires_the_dimensions_as_well_as_the_flag():
    doc = Document(io.BytesIO(_document()))
    pg = doc.sections[0]._sectPr.find(qn("w:pgSz"))
    pg.set(qn("w:orient"), "landscape")  # deliberately leave the portrait dimensions
    data = io.BytesIO()
    doc.save(data)
    record = assert_outcomes(data.getvalue(), [{"kind": "page_setup", "orientation": "landscape"}])[0]
    assert record["status"] == "FAIL" and any("dimensions" in e for e in record["errors"])


def test_size_is_checked_independently_of_font_name():
    doc = Document(io.BytesIO(_document()))
    for p in doc.paragraphs:
        for r in p.runs:
            size = OxmlElement("w:sz")
            size.set(qn("w:val"), "24")
            r._r.get_or_add_rPr().append(size)
    data = io.BytesIO()
    doc.save(data)
    assertion = _require_font(size=12, include_paragraph_marks=False)
    assert assert_outcomes(data.getvalue(), [assertion])[0]["status"] == "PASS"
    assertion["size"] = 10
    assert assert_outcomes(data.getvalue(), [assertion])[0]["status"] == "FAIL"


def test_properties_require_both_regular_and_complex_script_bold_and_paragraph_flags():
    doc = Document(io.BytesIO(_document()))
    for p in doc.paragraphs:
        p.paragraph_format.keep_with_next = True
        for r in p.runs:
            r.bold = True
    data = io.BytesIO()
    doc.save(data)
    assertion = {"kind": "effective_properties", "run": {"b": True}, "paragraph": {"keepNext": True},
                 "include_paragraph_marks": False}
    failure = assert_outcomes(data.getvalue(), [assertion])[0]
    assert failure["status"] == "FAIL" and any("bCs" in e for e in failure["errors"])
    for p in doc.paragraphs:
        for r in p.runs:
            r._r.get_or_add_rPr().append(OxmlElement("w:bCs"))
    data = io.BytesIO()
    doc.save(data)
    assert assert_outcomes(data.getvalue(), [assertion])[0]["status"] == "PASS"
    assertion["paragraph"]["keepNext"] = False
    assert assert_outcomes(data.getvalue(), [assertion])[0]["status"] == "FAIL"


def test_conditional_table_font_cannot_hide_behind_the_base_style_resolver():
    doc = Document(io.BytesIO(_document()))
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Conditional cell"
    table.style = "Table Grid"
    conditional = OxmlElement("w:tblStylePr")
    conditional.set(qn("w:type"), "firstRow")
    rpr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Courier New")
    rpr.append(fonts)
    conditional.append(rpr)
    doc.styles["Table Grid"].element.append(conditional)
    data = io.BytesIO()
    doc.save(data)
    result = assert_outcomes(data.getvalue(), [_require_font()])[0]
    assert result["status"] == "FAIL" and any("conditional table" in e for e in result["errors"])
