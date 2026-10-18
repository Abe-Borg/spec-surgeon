"""An approved header removal never exempts body words or whole rendered pages."""
from __future__ import annotations

import pymupdf
import pytest
from docx import Document
from docx.enum.section import WD_SECTION

from spec_surgeon.batch import geometry_check
from spec_surgeon.verify.geometry import compare_pdfs, hf_removal_copy
from spec_surgeon.verify.props import sect_prs, sig
from spec_surgeon.verify.reader import VPackage, W


def source_document(path):
    doc = Document()
    doc.add_paragraph("Repeated retained body words")
    for i in range(2):
        section = doc.sections[0] if not i else doc.add_section(WD_SECTION.NEW_PAGE)
        section.header.is_linked_to_previous = False
        section.header.paragraphs[0].text = f"Header {i}"
        section.first_page_footer.is_linked_to_previous = False
        section.first_page_footer.paragraphs[0].text = f"First footer {i}"
    doc.save(path)
    return path


def rendered(path, *, header=True, body=True, footer=True, moved=False, missing=False, extra=False):
    doc = pymupdf.open()
    page = doc.new_page(width=612, height=792)
    if header:
        page.insert_text((72, 30), "Repeated")
    if body:
        page.insert_text((72, 200 if not moved else 220), "Repeated retained body" if missing else "Repeated retained body words")
    if footer:
        page.insert_text((72, 760), "Retained footer")
    if extra:
        page.insert_text((72, 40), "Invented header")
    doc.save(path)
    doc.close()
    return path


def test_removal_reference_uses_raw_selector_and_preserves_every_other_source_item(tmp_path):
    source = source_document(tmp_path / "source.docx")
    plan = {"ops": [{"op": "remove_header_footer", "select": {
        "sections": [1], "kinds": ["footer"], "types": ["first"]}, "params": {},
        "transform_data": {"sections": [0]}}],
        "property_assertions": [{"kind": "section_hf_refs", "sections": [0]}]}
    reference = hf_removal_copy(source, tmp_path / "reference.docx", plan)
    before, after = VPackage(source.read_bytes()), VPackage(reference.read_bytes())
    original, expected = before.xml(before.main_part()), after.xml(after.main_part())
    selected = sect_prs(original)[1]
    reference_to_remove = selected.find(W + "footerReference")
    assert reference_to_remove.get(W + "type") == "first"
    selected.remove(reference_to_remove)
    assert sig(original) == sig(expected)
    assert before.names == after.names
    assert all(before.read(name) == after.read(name) for name in before.names if name != before.main_part())
    assert hf_removal_copy(source, tmp_path / "unused.docx", {"ops": []}) is None
    assert not (tmp_path / "unused.docx").exists()


@pytest.mark.parametrize("damage", ["blank", "missing", "moved", "extra"])
def test_header_removal_keeps_all_body_word_and_position_requirements(tmp_path, damage):
    source = source_document(tmp_path / "source.docx")
    original = rendered(tmp_path / "source.pdf")
    reference = rendered(tmp_path / "reference.pdf", header=False)
    output = rendered(tmp_path / "output.pdf", header=False, body=damage != "blank", footer=damage != "blank",
                      missing=damage == "missing", moved=damage == "moved", extra=damage == "extra")
    evidence = compare_pdfs(original, output, removed_hf_pdf=reference, removed_hf_source=source)
    assert not evidence.evidence_problems
    assert evidence.pages[0].expected_removed_words == 1
    assert geometry_check({"mode": "formatting"}, evidence).status == "FAIL"


def test_header_removal_matches_repeated_words_by_position_and_retains_footer(tmp_path):
    source = source_document(tmp_path / "source.docx")
    original = rendered(tmp_path / "source.pdf")
    output = rendered(tmp_path / "output.pdf", header=False)
    evidence = compare_pdfs(original, output, removed_hf_pdf=output, removed_hf_source=source)
    assert not evidence.evidence_problems
    assert evidence.pages[0].matched_words == 6
    assert evidence.pages[0].expected_removed_words == 1
    assert geometry_check({"mode": "formatting"}, evidence).status == "PASS"


@pytest.mark.parametrize("damage", ["missing", "moved", "blank", "bounds", "pages"])
def test_independent_removal_reference_cannot_hide_body_damage(tmp_path, damage):
    source = source_document(tmp_path / "source.docx")
    original = rendered(tmp_path / "source.pdf")
    reference = rendered(tmp_path / "reference.pdf", header=False, body=damage != "blank",
                         missing=damage == "missing", moved=damage == "moved")
    if damage == "pages":
        with pymupdf.open(reference) as document:
            document.new_page()
            document.save(tmp_path / "two-pages.pdf")
        reference = tmp_path / "two-pages.pdf"
    evidence = compare_pdfs(original, reference, removed_hf_pdf=reference,
                            removed_hf_source=None if damage == "bounds" else source)
    assert evidence.evidence_problems
    assert geometry_check({"mode": "formatting"}, evidence).status == "FAIL"
