"""Phase 3: document profiles, the batch planner (dry run) and the op catalog / EditSpec schema."""
from __future__ import annotations

import io
import json
import shutil
from pathlib import Path
import zipfile

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import RGBColor
import pytest
from lxml import etree

from make_fixtures import (W, build, build_numbering, build_rich, build_scheme_example, build_style_template,
                           build_styles, build_tables, build_template, build_typed, docx_bytes, patch_docx)
from spec_surgeon import compiler
from spec_surgeon.catalog import COMMON_ASSUMPTIONS, CUSTOM_TRANSFORM_LIMITS, editspec_json_schema, op_catalog
from spec_surgeon.compiler import _Batch
from spec_surgeon.editspec import load_spec, plan_spec, validate_spec
from spec_surgeon.ops import OPS
from spec_surgeon.package import DocxPackage
from spec_surgeon.planner import plan_batch, summarize_plan
from spec_surgeon.profile import MAX_PROFILE_BYTES, bound, markup, profile, profile_batch, template_spec_summary
from spec_surgeon.verify import props as vprops
from test_adaptive_messy import _build_messy_inline_spec

SPECS = Path(__file__).resolve().parent.parent / "examples" / "specs"


@pytest.fixture(scope="module")
def folder(tmp_path_factory) -> Path:
    d = tmp_path_factory.mktemp("batch")
    for i in range(3):
        (d / f"21 0{i} 00 – SPEC {i}.docx").write_bytes(build(f"21 0{i} 00", f"SPEC {i}"))
    (d / "21 13 13 – RICH.docx").write_bytes(build_rich())
    return d


def test_profile_describes_sections_parts_styles_and_fits_the_budget():
    prof = profile(DocxPackage(build_rich(), name="rich.docx"))
    assert len(json.dumps(prof, ensure_ascii=False)) <= MAX_PROFILE_BYTES and not prof.get("truncated")
    s0, s1 = prof["sections"][0], prof["sections"][1]
    assert s0["titlePg"] and s0["headers"]["first"]["part"] == "word/header1.xml"
    s2 = prof["sections"][2]
    assert s1["headers"]["even"]["own"] and s1["footers"]["default"]["part"] == s0["footers"]["default"]["part"]
    assert not any(v["own"] for k in ("headers", "footers") for v in s2[k].values())     # section 2 inherits all
    assert prof["settings"]["evenAndOddHeaders"]
    assert "EVEN PAGE HEADER" in prof["header_footer_parts"]["word/header4.xml"]["text"][0]
    styles = {s["id"]: s for s in prof["styles_in_use"]}
    assert styles["SpecNote"]["font"] == "Courier New"
    assert "Wingdings" in prof["fonts"]["direct_in_runs"]


def test_profile_of_a_template_lists_its_slots():
    prof = profile(DocxPackage(build_template(), name="T.docx"))
    parts = prof["header_footer_parts"]
    assert parts["word/header2.xml"]["slots"] == ["SECTION_NUMBER", "SECTION_TITLE", "ISSUED FOR PERMIT"]
    assert parts["word/footer1.xml"]["slots"] == ["XX XX XX"]
    assert "⇥" in parts["word/header1.xml"]["text"][0]                                 # tabs stay visible


def test_bound_always_fits_and_cuts_samples_before_structure():
    """Codex review on PR #3: the bound must hold even when structure alone exceeds it; every structural cut is
    reported in ``truncated``."""
    prof = profile(DocxPackage(build_rich(), name="rich.docx"))
    # just room for every section once the descriptive parts are cut (3500 before batch 5's table and list summaries,
    # whose omission records take 20 bytes of it)
    roomy = bound(json.loads(json.dumps(prof)), limit=3520)
    assert roomy["sections"] == prof["sections"] and not roomy["first_paragraphs"]          # samples went first
    # batch 4's character/table styles in use go before any section or part (they matter less than the page setup), and
    # so do batch 5's table and list summaries (the table count stays under counts)
    assert roomy["truncated"] == ["header/footer text", "character styles in use", "tables", "lists"]
    for limit in (3000, 1500, 800, 300):
        small = bound(json.loads(json.dumps(prof)), limit=limit)
        assert len(json.dumps(small, ensure_ascii=False)) <= limit, limit
    tiny = bound(json.loads(json.dumps(prof)), limit=300)
    assert set(tiny) <= {"file", "pattern_id", "settings", "counts", "truncated"} and tiny["truncated"]


def test_an_unreadable_file_is_a_record_not_an_abort(folder, tmp_path):
    """Codex review on PR #3: one truncated .docx must not stop the dry run or the profiler."""
    d = tmp_path / "with-broken"
    shutil.copytree(folder, d)
    (d / "BROKEN.docx").write_bytes(b"PK\x03\x04 not really a zip")
    bp = plan_batch([str(d)], load_spec(SPECS / "remove-all-headers-footers.json"))
    broken = [f for f in bp.files if f.source.name == "BROKEN.docx"][0]
    assert broken.status == "exception" and broken.exception["code"] == "unreadable"
    assert broken.exception["kind"] == "unreadable" and broken.exception["stage"] == "open"
    assert "damaged or cut short" in broken.exception["reason"]           # it starts like a zip, but is not one
    assert sum(f.status == "planned" for f in bp.files) == 4
    pb = profile_batch(sorted(d.glob("*.docx")))
    assert [u["file"] for u in pb["unreadable"]] == ["BROKEN.docx"] and sum(p["count"] for p in pb["patterns"]) == 4


def test_profile_batch_groups_by_pattern(folder):
    pb = profile_batch(sorted(folder.glob("*.docx")))
    assert pb["files"] == 4 and [p["count"] for p in pb["patterns"]] == [3, 1]
    assert pb["patterns"][0]["files"][0].startswith("21 00 00")


def test_plan_batch_dry_run(folder):
    before = {f: f.read_bytes() for f in folder.glob("*.docx")}
    bp = plan_batch([str(folder)], load_spec(SPECS / "apply-header-footer-template.json"))
    assert len(bp.files) == 4 and len(bp.patterns) == 2
    exc = bp.exceptions
    assert [f.source.name for f in exc] == ["21 13 13 – RICH.docx"] and exc[0].exception["code"] == "even-odd"
    ok = [f for f in bp.files if f.status == "planned"][0]
    assert any(line.startswith("members added: header2.xml") for line in ok.summary)
    assert any("literal fill" in line for line in ok.summary)
    assert {f: f.read_bytes() for f in folder.glob("*.docx")} == before                  # nothing written
    text = bp.to_text(per_file=False)
    assert "4 file(s), 3 plan cleanly, 1 exception(s)" in text and "… and 2 more file(s)" in text
    json.dumps(bp.to_dict())


def test_plan_batch_skips_the_spec_template(folder, tmp_path):
    spec = load_spec(SPECS / "apply-header-footer-template.json")
    d = tmp_path / "with-template"
    shutil.copytree(folder, d)
    shutil.copy(spec["ops"][0]["params"]["template"], d / "00 00 00 TEMPLATE.docx")
    bp = plan_batch([str(d)], spec)
    assert [Path(s).name for s in bp.skipped] == ["00 00 00 TEMPLATE.docx"] and len(bp.files) == 4


def test_catalog_and_schema_cover_every_op():
    cat = op_catalog()
    assert set(cat["ops"]) == set(OPS)
    for name, entry in cat["ops"].items():
        assert entry["description"] and entry["common_assumptions"], name
    consts = [item["properties"]["op"]["const"] for item in cat["editspec_schema"]["properties"]["ops"]["items"]["oneOf"]]
    assert sorted(consts) == sorted(OPS)
    assert {e["request"] for e in cat["worked_examples"]} >= {"change all font to Arial", "remove all headers and footers"}
    json.dumps(cat)


def test_example_specs_conform_to_the_published_schema():
    jsonschema = pytest.importorskip("jsonschema")
    schema = editspec_json_schema()
    for f in SPECS.glob("*.json"):
        jsonschema.validate(json.loads(f.read_text(encoding="utf-8")), schema)
    bad = json.loads((SPECS / "remove-all-headers-footers.json").read_text(encoding="utf-8"))
    bad["ops"][0]["params"]["colour"] = "red"
    with pytest.raises(jsonschema.ValidationError):
        jsonschema.validate(bad, schema)


# ---- the profile: borders, shading, styles in use, languages (built-in batch 4) ------------------------------------------

def _prof(data: bytes):
    return profile(DocxPackage(data, name="spec.docx"))


def _style(root, sid):
    return next(s for s in root.findall(W + "style") if s.get(W + "styleId") == sid)


def test_styles_show_their_border_box_and_shading():
    styles = {s["id"]: s for s in _prof(build_styles())["styles_in_use"]}
    assert styles["ArchNote"]["borders"] == ["top", "left", "bottom", "right"]
    assert styles["ArchNote"]["shading"] == "D9D9D9"
    assert styles["ArchNoteChild"]["borders"] == ["top", "left", "right"]      # its own bottom nil wins


def test_a_theme_fill_is_reported_as_the_theme_colour():
    """Codex on #31: a w:shd with w:themeFill and no cached w:fill is shading Word renders."""
    def edit(root):
        shd = _style(root, "ArchNote").find(f"{W}pPr/{W}shd")
        del shd.attrib[W + "fill"]
        shd.set(W + "themeFill", "accent1")
        shd.set(W + "themeFillTint", "33")
    styles = {s["id"]: s for s in _prof(patch_docx(build_styles(), {"word/styles.xml": edit}))["styles_in_use"]}
    assert styles["ArchNote"]["shading"] == "accent1 tint 33"


def test_character_and_table_styles_in_use_cover_every_part_and_marks_but_no_snapshot():
    """Codex on #31: map_styles works over every text part and on paragraph marks; a tracked change's snapshot is what
    reject-all would restore, not what the file shows."""
    def edit_header(root):
        p = root.find(f".//{W}p")
        r = etree.SubElement(p, W + "r")
        rpr = etree.SubElement(r, W + "rPr")
        etree.SubElement(rpr, W + "rStyle").set(W + "val", "ArchUnusedChar")
        etree.SubElement(r, W + "t").text = "header words"

    def edit_doc(root):
        p = root.findall(f".//{W}body/{W}p")[3]
        ppr = p.find(W + "pPr")
        mark = etree.SubElement(ppr, W + "rPr")
        etree.SubElement(mark, W + "rStyle").set(W + "val", "ArchLinkedChar")
        r = p.find(W + "r")
        rpr = r.find(W + "rPr") if r.find(W + "rPr") is not None else etree.SubElement(r, W + "rPr")
        chg = etree.SubElement(rpr, W + "rPrChange")
        chg.set(W + "id", "900"); chg.set(W + "author", "Jeremy")
        snap = etree.SubElement(chg, W + "rPr")
        etree.SubElement(snap, W + "rStyle").set(W + "val", "SubtleEmphasis")
    data = patch_docx(build_styles(), {"word/header1.xml": edit_header, "word/document.xml": edit_doc})
    prof = _prof(data)
    chars = {e["id"] for e in prof["character_styles_in_use"]}
    assert {"ArchUnusedChar", "ArchLinkedChar", "ArchStrong", "Heading1Char"} <= chars
    assert "SubtleEmphasis" not in chars
    assert [e["id"] for e in prof["table_styles_in_use"]] == ["ArchTable"]


def test_languages_leave_out_tracked_snapshots():
    """Codex on #31: a w:pPrChange snapshot's paragraph-mark language is not the file's."""
    def edit(root):
        chg = root.find(f".//{W}pPrChange/{W}pPr")
        rpr = etree.SubElement(chg, W + "rPr")
        etree.SubElement(rpr, W + "lang").set(W + "val", "de-DE")
    langs = _prof(patch_docx(build_styles(), {"word/document.xml": edit}))["markup"]["languages"]
    assert langs["default"] == "en-GB" and langs["styles"] == {"en-GB": 1}
    assert langs["direct"] == {"en-GB": 3, "fr-FR": 1}


def test_the_bound_thins_the_new_style_lists_before_dropping_everything():
    """Codex on #31: long character/table style lists must be trimmed by the staged truncation, not push the profile
    into the last-resort branch."""
    prof = _prof(build_styles())
    prof["character_styles_in_use"] = [{"id": f"CharStyle{i:03d}", "count": i, "name": "A long character style name " * 3}
                                       for i in range(80)]
    prof["table_styles_in_use"] = [{"id": f"TableStyle{i:03d}", "count": i, "name": "A long table style name " * 3}
                                   for i in range(80)]
    out = bound(prof, limit=6000)
    assert "character_styles_in_use" in out and len(out["character_styles_in_use"]) == 4
    assert "name" not in out["character_styles_in_use"][0]
    assert any("character style(s) in use" in t for t in out["truncated"])


def test_a_template_lists_the_styles_it_offers():
    summary = template_spec_summary(DocxPackage(build_style_template(), name="OFFICE STYLES.docx"))
    rows = {s["id"]: s for s in summary["styles"]}
    assert rows["PRT"]["numbered"] is True and rows["PR2"]["basedOn"] == "PR1"
    assert rows["MyBody"]["name"] == "Body Text" and rows["TblMine"]["type"] == "table"


# ---- the profile: lists and tables (built-in batch 5) --------------------------------------------------------------------

def _profile(data: bytes):
    return profile(DocxPackage(data, name="x.docx"))


def test_the_example_file_shows_its_scheme_as_numbered_styles():
    lists = _profile(build_scheme_example())["lists"]
    shown = [(e["style"], e["level"], e.get("shows")) for e in lists["numbered_styles"][:5]]
    assert shown == [("PRT", 0, "PART 1 -"), ("ART", 1, "1.1"), ("PR1", 2, "A."), ("PR2", 3, "1."), ("PR3", 4, "a.")]
    assert len({e["list"] for e in lists["numbered_styles"][:5]}) == 1            # one list: a scheme
    outline = lists["lists"][0]
    assert [lv.get("first") for lv in outline["levels"]] == ["PART 1 -", "1.1", "A.", "1.", "a."]
    assert "not_evaluated" not in outline


def test_unused_single_level_list_styles_are_left_out():
    """python-docx's List Number styles number nothing in these files: noise for a scheme."""
    lists = _profile(build_rich()).get("lists") or {}
    assert not lists.get("numbered_styles")


def test_lists_say_which_paragraphs_the_engine_does_not_evaluate():
    outline = _profile(build_numbering(shared_restart=True))["lists"]["lists"][0]
    assert outline["lists"] == [10, 11] and outline["not_evaluated"] == {"shared": outline["paragraphs"]}
    clean = _profile(build_numbering(shared_restart=False))["lists"]["lists"][0]
    # "Product data." skips level 2: the three level-2 paragraphs after it are not evaluated (lists.skipped_below)
    assert clean["not_evaluated"] == {"skipped": 3} and clean["levels"][1]["first"] == "1.01"
    assert clean["levels"][2]["styles"][0] == "ArchPara"


def test_typed_numbers_are_counted_with_tab_separators():
    lists = _profile(build_typed())["lists"]
    assert lists["typed_numbers"] >= 20 and "lists" not in lists


def test_tables_are_summarised_by_the_selector_s_ordinal():
    tables = _profile(build_tables())["tables"]
    assert tables["count"] == len(tables["first"]) + tables.get("more", 0)
    assert [t["index"] for t in tables["first"]] == list(range(len(tables["first"])))
    assert any(t.get("nested") for t in tables["first"]) and any(t["header_rows"] for t in tables["first"])
    assert "tables" not in _profile(build_scheme_example())              # a body without a table says nothing


def test_the_summaries_go_before_any_section_when_the_profile_is_too_big():
    prof = _profile(build_numbering(shared_restart=True))
    assert "lists" in prof and "tables" in prof
    for limit in (6000, 4500, 3500):
        cut = bound(json.loads(json.dumps(prof)), limit=limit)
        if len(cut.get("sections", [])) < len(prof["sections"]):
            assert "lists" not in cut and "tables" not in cut, limit


def _marked() -> bytes:
    doc = Document()
    p = doc.add_paragraph("Plain ")
    r = p.add_run("hidden note"); r.font.hidden = True
    r = p.add_run("yellow"); r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    r = p.add_run("blue"); r.font.color.rgb = RGBColor(0, 0, 0xFF)
    doc.core_properties.title = "Spec title"
    doc.core_properties.author = "Jeremy"
    return docx_bytes(doc)


def test_the_profile_shows_what_the_clean_up_ops_would_find():
    m = markup(DocxPackage(_marked()))
    assert m["hidden_runs"] == 1 and m["highlight_colors"] == {"yellow": 1} and m["run_colors"] == {"0000FF": 1}
    assert m["properties"]["title"] == "Spec title" and m["properties"]["author"] == "Jeremy"
    rich = profile(DocxPackage(build_rich()))
    assert rich["markup"]["tracked_changes"]["Jeremy"]["ins"] == 2
    assert rich["markup"]["tracked_change_dates"] == ["2026-09-22", "2026-09-23"]


NEW_OPS = ("set_borders_shading", "clear_direct_formatting", "set_language", "apply_style", "map_styles",
           "remove_unused_styles")


def test_catalog_prompt_and_custom_transform_limits_know_the_batch():
    cat = op_catalog()
    for op in NEW_OPS:
        assert op in cat["ops"] and COMMON_ASSUMPTIONS[op], op
        assert op in compiler.SYSTEM, op
    assert "styles" in cat["ops"]["set_paragraph_props"]["select"]["properties"]
    assert CUSTOM_TRANSFORM_LIMITS["families"]["paragraph borders and shading of paragraphs or of paragraph styles"] \
        == "borders_shading"
    assert not any("borders" in c or "shading" in c for c in CUSTOM_TRANSFORM_LIMITS["cannot"])
    assert set(CUSTOM_TRANSFORM_LIMITS["families"].values()) <= set(vprops.KINDS)
    for kind in ("borders_shading", "cleared_formatting", "language", "applied_style", "mapped_styles",
                 "styles_imported", "removed_styles"):
        assert kind in vprops.KINDS, kind
    assert "ATTACHED TEMPLATE" in compiler.SYSTEM and "read_template lists its styles" in compiler.SYSTEM


@pytest.fixture(scope="module")
def src() -> bytes:
    return build_styles()


def _with_part(data: bytes, name: str, edit) -> bytes:
    zin = zipfile.ZipFile(io.BytesIO(data))
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            raw = zin.read(item.filename)
            if item.filename == name:
                raw = edit(raw)
            zout.writestr(item, raw)
    return buf.getvalue()


def test_the_profile_survives_a_malformed_numbering_part(src):
    """Review D6: the unused-style counts parse numbering and settings; a part that will not parse leaves them out, the
    file still profiles."""
    data = _with_part(src, "word/numbering.xml", lambda raw: raw[: len(raw) // 2])
    prof = profile(DocxPackage(data, name="spec.docx"))
    assert prof.get("styles_unused") is None and prof["styles_in_use"]


def _spec(request, mode, ops, **extra):
    return {"request": request, "summary": request, "assumptions": [], "questions": [], "applies_to": {"or": "all"},
            "mode": mode, "ops": ops, **extra}


def test_a_pure_tracked_deletion_is_named_in_the_plan_summary():
    pkg = DocxPackage(build("21 13 13", "WET-PIPE SPRINKLER SYSTEMS"))
    plan, _ = plan_spec(pkg, _spec("delete the words 'FM Global'", "content",
                                   [{"op": "replace_text", "params": {"find": "FM Global", "replace": ""}}]))
    lines = summarize_plan(plan)
    assert any("tracked deletion" in line for line in lines)
    assert not any(line == "text: no block changes text" for line in lines)


EXAMPLES = Path(__file__).resolve().parent.parent / "examples" / "specs"


def test_catalog_documents_fill_rules_and_expect_ranges_and_the_fallback_example_plans_on_both_layouts():
    cat = op_catalog()
    assert set(cat["fill_rules"]) == {"literal", "regex", "paragraph_after", "first_match", "whitespace", "in_template"}
    assert "'min'" in cat["expect_syntax"]["value"] and "rPrChange" in cat["expect_syntax"]["rev_delta"]
    assert "first_match" in cat["ops"]["apply_header_footer_template"]["params"]["properties"]["slot_fills"]["description"]
    assert any("first_match" in json.dumps(e) for e in cat["worked_examples"])
    spec = load_spec(EXAMPLES / "apply-header-footer-template-fallback.json")
    validate_spec(spec)
    for raw, title in ((build("21 05 00", "COMMON WORK RESULTS"), "COMMON WORK RESULTS"),
                       (_build_messy_inline_spec("21 13 13", "WET-PIPE SPRINKLER SYSTEMS"), "WET-PIPE SPRINKLER SYSTEMS")):
        plan, frags = plan_spec(DocxPackage(raw), spec)
        assert frags[0].transform_data["fills"]["SECTION_TITLE"] == title


def _rewrite(data: bytes, part: str, fn) -> bytes:
    """``data`` with ``fn(root)`` applied to the XML of ``part``."""
    zin = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            b = zin.read(info.filename)
            if info.filename == part:
                root = etree.fromstring(b)
                fn(root)
                b = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            zout.writestr(info, b)
    return out.getvalue()


def test_profile_batch_takes_the_next_file_when_the_first_cannot_be_profiled(tmp_path):
    def bad_margin(root):
        for m in root.iter(W + "pgMar"):
            m.set(W + "left", "abc")
    (tmp_path / "A.docx").write_bytes(_rewrite(build("21 00 15", "JULIET"), "word/document.xml", bad_margin))
    (tmp_path / "B.docx").write_bytes(build("21 00 16", "KILO"))
    pb = profile_batch([tmp_path / "A.docx", tmp_path / "B.docx"])
    [pat] = pb["patterns"]
    assert pat["files"] == ["B.docx", "A.docx"] and "error" not in pat["profile"]
    assert pat["unprofiled"] and pat["unprofiled"][0].startswith("A.docx")
    assert _Batch([tmp_path / "A.docx", tmp_path / "B.docx"], None).representatives() == [tmp_path / "B.docx"]
    only = profile_batch([tmp_path / "A.docx"])                                # no file profiles: said, not raised
    assert only["patterns"][0]["profile"]["error"]
