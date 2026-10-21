"""remove_empty_paragraphs (docs/BUILTIN_OPS_BATCH2_DESIGN.md §8): collapse and all, tracked and untracked (the whole
suite must pass), what is not blank, the container rules (a cell, the document and a header keep a last paragraph; two
tables are never joined; the last section keeps one), nothing-to-do files, and planted defects the verifier must
catch."""
from __future__ import annotations

from typing import Any, Dict, List

import pytest
from docx import Document

from fx_removal import W, build_blank, build_blank_last_section, save
from spec_surgeon.editspec import apply_spec, plan_spec
from spec_surgeon.ops import JobException, get_op, merge_fragments
from spec_surgeon.ops import empty_paragraphs as empty_module
from spec_surgeon.package import DocxPackage
from spec_surgeon.verify import run_suite
from test_text_edits import accepted_text, failed, mutate, v10_errors

AUTHOR = "Abraham Borg"
DOC = "word/document.xml"


def run_op(src: bytes, params: Dict[str, Any], select: Dict[str, Any] = None):
    pkg = DocxPackage(src)
    frag = get_op("remove_empty_paragraphs").plan(pkg, select or {}, {"author": AUTHOR, **params})
    plan = merge_fragments([frag])
    work = pkg.copy()
    get_op("remove_empty_paragraphs").apply(work, frag)
    out = work.output_bytes()
    return plan, out, run_suite(src, out, plan.to_dict(), "src.docx", "out.docx")


def outline(data: bytes, part: str = DOC) -> List[str]:
    """The container's blocks as a reader sees them with every change accepted: the paragraph's text, "·" for a
    paragraph without text (blank or not: a number, a border, a bookmark or a break shows no text), "TBL" for a table."""
    root = DocxPackage(data).xml(part)
    container = root.find(W + "body") if root.find(W + "body") is not None else root
    out = []
    for c in container:
        if c.tag == W + "tbl":
            out.append("TBL")
        elif c.tag == W + "p":
            if c.find(f"{W}pPr/{W}rPr/{W}del") is not None:
                continue                                              # deleted (tracked): gone once accepted
            text = accepted_text(c).strip("  ")
            out.append(text if text else "·")
    return out


@pytest.fixture(scope="module")
def blank() -> bytes:
    return build_blank()


@pytest.fixture(scope="module")
def last_section() -> bytes:
    return build_blank_last_section()


LAST_SECTION_FLAG = "the last section is now empty and shows a blank page"


SOURCE = ["Para A", "·", "·", "·", "Para B", "·", "·", "Para C", "·", "·", "·", "·", "Para D", "\t", "·", "Para E",
          "TBL", "·", "TBL", "Para F", "·", "Para G", "·", "Para H", "TBL", "·", "·"]


def test_the_fixture_reads_as_the_tests_expect(blank):
    assert outline(blank) == SOURCE


@pytest.mark.parametrize("tracked", [True, False])
def test_collapse_keeps_the_first_of_each_run(blank, tracked):
    plan, out, rec = run_op(blank, {"tracked": tracked})
    assert rec.passed, rec.to_markdown()
    after = outline(out)
    assert after[:8] == ["Para A", "·", "Para B", "·", "Para C", "·", "·", "·"]  # three blanks → one; spaces + NBSP → one
    if tracked:
        assert after[-2:] == ["·", "·"]                                      # the document's last paragraph stays
        assert plan.rev_delta["del_mark"] == 3
    else:
        assert after[-2:] == ["TBL", "·"]
        assert len(after) == len(SOURCE) - 4


@pytest.mark.parametrize("tracked", [True, False])
def test_all_removes_every_blank_paragraph_but_what_must_stay(blank, tracked):
    plan, out, rec = run_op(blank, {"mode": "all", "tracked": tracked})
    assert rec.passed, rec.to_markdown()
    after = outline(out)
    # not blank: a numbered paragraph (direct and by style), a bookmark, a tab, a page break, a section break, a border
    assert after == ["Para A", "Para B", "Para C", "·", "·", "·", "Para D", "\t", "·", "Para E", "TBL", "·", "TBL",
                     "Para F", "·", "Para G", "·", "Para H", "TBL", "·"]
    #                                                    between the tables ↑               the document's end ↑
    assert plan.rev_delta.get("del_mark", 0) == (7 if tracked else 0)
    tcs = list(DocxPackage(out).xml(DOC).iter(W + "tc"))
    assert len(tcs) == 4
    # the cell "X" + blank loses the blank untracked (tracked: a cell's last paragraph stays); the cell holding only a
    # blank keeps it either way
    assert [len(tc.findall(W + "p")) for tc in tcs][-2:] == ([2, 1] if tracked else [1, 1])


def test_the_header_only_when_asked(blank):
    plan, out, rec = run_op(blank, {"tracked": False}, {"scope": ["document", "headers"]})
    assert rec.passed, rec.to_markdown()
    header = next(p for p in DocxPackage(out).text_parts() if "header" in p)
    assert outline(out, header) == ["Header text", "·"]
    plan, out, rec = run_op(blank, {"tracked": False})
    assert outline(out, header) == ["Header text", "·", "·"]


def test_through_a_spec(blank):
    spec = {"request": "collapse the blank lines", "summary": "s", "mode": "untracked",
            "ops": [{"op": "remove_empty_paragraphs", "params": {"tracked": False, "author": AUTHOR}}]}
    pkg = DocxPackage(blank)
    plan, frags = plan_spec(pkg, spec)
    out = apply_spec(pkg, frags).output_bytes()
    assert run_suite(blank, out, plan.to_dict(), "s.docx", "o.docx").passed


def test_nothing_to_do_is_a_no_op_by_default_and_not_found_when_asked():
    doc = Document()
    doc.add_paragraph("One")
    doc.add_paragraph("Two")
    src = save(doc)
    plan, out, rec = run_op(src, {})
    assert rec.passed and out == src and any("nothing to do" in n for n in plan.notes)
    with pytest.raises(JobException) as e:
        run_op(src, {"allow_missing": False})
    assert e.value.code == "not-found"


@pytest.mark.parametrize("params", [{"mode": "some"}, {"allow_missing": "yes"}, {"tracked": "no"}])
def test_params(blank, params):
    with pytest.raises(JobException) as e:
        run_op(blank, params)
    assert e.value.code == "bad-params"


def test_a_lone_blank_between_two_tables_stays_and_a_cell_keeps_its_paragraph():
    doc = Document()
    t1 = doc.add_table(rows=1, cols=1)
    t1.cell(0, 0).paragraphs[0].text = ""
    doc.add_paragraph("")
    doc.add_paragraph("")
    doc.add_table(rows=1, cols=1)
    doc.add_paragraph("")
    src = save(doc)
    plan, out, rec = run_op(src, {"mode": "all", "tracked": False})
    assert rec.passed, rec.to_markdown()
    assert outline(out) == ["TBL", "·", "TBL", "·"]


@pytest.mark.parametrize("tracked", [True, False])
def test_the_last_section_keeps_its_last_blank_and_its_blank_page_is_flagged(last_section, blank, tracked):
    """Every block of the last section (after the body's last section break) a blank paragraph: its last one stays —
    not the break's paragraph alone at the end of the body, which Word shows as a blank page with nothing on it to
    remove (docs/BUILTIN_OPS_BATCH6_DESIGN.md, F2 review fixes: keep rule 4)."""
    assert outline(last_section) == ["Cover", "·", "Body text", "·", "·", "·", "·", "·"]
    plan, out, rec = run_op(last_section, {"mode": "all", "tracked": tracked})
    assert rec.passed, rec.to_markdown()
    assert plan.property_assertions[0]["expected"][DOC] == [3, 5, 6]     # the body's blank, two of the last section's
    assert outline(out) == ["Cover", "·", "Body text", "·", "·"]           # the second break · the last section's blank
    assert any(f.startswith(LAST_SECTION_FLAG) and f.endswith("remove the section break before it to drop the page")
               for f in plan.flags)
    body = DocxPackage(out).xml(DOC).find(W + "body")
    if not tracked:
        assert [c.tag for c in body][-3:] == [W + "p", W + "p", W + "sectPr"]
        assert body[-3].find(f"{W}pPr/{W}sectPr") is not None and body[-2].find(f"{W}pPr/{W}sectPr") is None
    # collapse: the last section's first blank stays anyway (tracked: the document's last paragraph too)
    plan, out, rec = run_op(last_section, {"tracked": tracked})
    assert rec.passed, rec.to_markdown()
    assert outline(out) == ["Cover", "·", "Body text", "·", "·", "·"] + (["·"] if tracked else [])
    assert not any(LAST_SECTION_FLAG in f for f in plan.flags)
    plan, _, _ = run_op(blank, {"mode": "all", "tracked": tracked})           # the last section keeps text: no flag
    assert not any(LAST_SECTION_FLAG in f for f in plan.flags)


# ---- V2's pairing with predicted deletions (verify/textcheck.diff_blocks) ----------------------------------------

def test_v2_a_predicted_deletion_that_did_not_happen_still_fails(blank):
    plan, out, rec = run_op(blank, {"tracked": False})
    rec = run_suite(blank, blank, plan.to_dict(), "s", "o")             # nothing deleted at all
    assert "V2" in failed(rec)
    unpredicted = rec.get("V2").data["unpredicted"]
    assert unpredicted and all(u[1] is None for u in unpredicted)       # the blocks still there: reported as new


def test_v2_a_deleted_blank_never_pairs_with_another_blank_far_away():
    """A deleted blank paragraph beside edited ones, with a blank paragraph much further on: setting the predicted
    deletion aside before aligning keeps every block paired with its own (an alignment on text alone paired the two
    blanks and reported every block between them as deleted and inserted)."""
    from spec_surgeon.verify.textcheck import Block, diff_blocks
    # block 1 (blank) is deleted; block 4 ("x") is emptied by the same op — the output's last block is blank too
    before = [Block(0, "A one"), Block(1, ""), Block(2, "B two"), Block(3, "C three"), Block(4, "x")]
    after = [Block(0, "A"), Block(1, "B"), Block(2, "C three"), Block(3, "")]
    changes = diff_blocks(before, after, {1}, set())
    got = sorted(((c.before, c.after) for c in changes), key=repr)
    assert got == sorted([("", None), ("A one", "A"), ("B two", "B"), ("x", "")], key=repr)


# ---- planted defects ---------------------------------------------------------------------------------------------

def test_defect_the_first_blank_of_a_run_removed_too(blank):
    plan, out, rec = run_op(blank, {"tracked": False})

    def drop_first_blank(root):
        body = root.find(W + "body")
        ps = [c for c in body if c.tag == W + "p"]
        assert accepted_text(ps[0]) == "Para A" and accepted_text(ps[1]) == ""
        body.remove(ps[1])
    rec = run_suite(blank, mutate(out, DOC, drop_first_blank), plan.to_dict(), "s", "o")
    assert "V2" in failed(rec) and "V10" in failed(rec)


def test_shared_defect_a_bookmark_paragraph_taken_as_blank(blank, monkeypatch):
    real = empty_module.is_empty

    def loose(p, sheet):
        if p.find(W + "bookmarkStart") is not None and not accepted_text(p):
            return True
        return real(p, sheet)
    monkeypatch.setattr(empty_module, "is_empty", loose)
    plan, out, rec = run_op(blank, {"mode": "all", "tracked": False})
    assert DocxPackage(out).xml(DOC).find(f".//{W}bookmarkStart") is None   # deleted by plan and apply alike …
    assert failed(rec) == ["V10"], rec.to_markdown()                       # … not blank to the verifier
    assert any("removed_empty_paragraphs" in e and "the plan predicted" in e for e in v10_errors(rec))


def test_shared_defect_two_tables_joined(blank, monkeypatch):
    monkeypatch.setattr(empty_module, "keep_rules", lambda blocks, chosen, tracked: [])
    plan, out, rec = run_op(blank, {"mode": "all", "tracked": False})
    assert "V10" in failed(rec)
    assert any("removed_empty_paragraphs" in e for e in v10_errors(rec))


def test_shared_defect_collapse_keeps_the_last_of_a_run(blank, monkeypatch):
    real = empty_module.blank_paragraphs

    def last_kept(root, sheet, mode, tracked):
        got = real(root, sheet, "all", tracked)                          # every blank, minus …
        keep = []
        for c in {id(p.getparent()): p.getparent() for p in got["delete"]}.values():
            blocks = empty_module.blocks_of(c)
            for i, b in enumerate(blocks):                               # … the LAST of each run instead of the first
                if b in got["delete"] and (i + 1 == len(blocks) or blocks[i + 1] not in got["delete"]):
                    keep.append(b)
        return {**got, "delete": [p for p in got["delete"] if p not in keep]}
    monkeypatch.setattr(empty_module, "blank_paragraphs", last_kept)
    plan, out, rec = run_op(blank, {"tracked": True})
    assert failed(rec) == ["V10"], rec.to_markdown()


def test_defect_the_last_sections_blank_removed_too(last_section):
    plan, out, rec = run_op(last_section, {"mode": "all", "tracked": False})

    def empty_the_last_section(root):
        body = root.find(W + "body")
        assert body[-3].find(f"{W}pPr/{W}sectPr") is not None and not accepted_text(body[-2])
        body.remove(body[-2])                                            # the break's paragraph now ends the body
    rec = run_suite(last_section, mutate(out, DOC, empty_the_last_section), plan.to_dict(), "s", "o")
    assert "V10" in failed(rec)
    assert any("removed_empty_paragraphs" in e for e in v10_errors(rec))


def test_shared_defect_the_last_section_emptied(last_section, monkeypatch):
    monkeypatch.setattr(empty_module, "last_section_keep", lambda *a: [])
    plan, out, rec = run_op(last_section, {"mode": "all", "tracked": False})
    body = DocxPackage(out).xml(DOC).find(W + "body")
    assert body[-2].find(f"{W}pPr/{W}sectPr") is not None                # emptied by plan and apply alike …
    assert not any(LAST_SECTION_FLAG in f for f in plan.flags)
    assert failed(rec) == ["V10"], rec.to_markdown()                       # … the verifier keeps the last blank
    assert any("removed_empty_paragraphs" in e and "the plan predicted" in e for e in v10_errors(rec))
