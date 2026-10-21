"""Fixtures for the table grid and cell ops of built-in batch 6 (F6 — delete_table_columns, set_table_cells,
merge_table_cells): tests/test_table_grid.py.

``build_grid()`` — a specification whose tables cover what the ops decide on. Every table is found by its ordinal
(``root.iter(w:tbl)`` of its part) and its first row's text:

* body: "SECTION 23 21 13 - HYDRONIC PIPING"; "1.1 PUMP SCHEDULE";
  table 0 — "Table Grid", a FIXED width (``w:tblW`` dxa 8640) over a grid of 1440 / 2880 / 2160 / 2160; 5 rows
  "Tag | Description | Size | Notes" (a repeating header row: own ``w:tblHeader``, bold, cells shaded D9D9D9) /
  "P-1 | Base-mounted pump | 2 in | Duty" / "P-2 | Base-mounted pump | 2 in | Standby" (its Notes cell merges down:
  ``w:vMerge restart``) / "P-3 | Inline pump | 1 in | " (continue) / "Totals | 5 in | " (its first cell spans the first
  two columns, ``w:tcW`` 4320); every cell's ``w:tcW`` dxa is its columns' widths;
  "1.2 ZONE SCHEDULE"; table 1 — "Table Grid", a width in PERCENT (``w:tblW`` pct 5000), a grid of 3120 × 3; 4 rows
  "Zone | Area | Floor" / row 1 starts one column in (``w:gridBefore`` 1, ``w:wBefore`` 3120): "North | 2" / row 2 ends
  one column early (``w:gridAfter`` 1, ``w:wAfter`` 3120): "South | Annex" / "East | Ground | 1" whose cells' widths are
  in percent (``w:tcW`` pct 1667);
  "1.3 VALVE SCHEDULE"; table 2 — no style, automatic width, a grid of 2400 × 2; 3 rows "Valve | Detail" / "Gate | …"
  whose second cell holds table 3 — nested, "Table Grid", 2 rows "Size | Rating" / "4 in | 175 psi" — and a paragraph
  "See detail." after it / "Check | Swing";
  "1.4 EQUIPMENT"; table 4 — "Arch Shaded" (its style's ``w:tcPr`` shades every cell F2F2F2 and turns ``w:noWrap`` on;
  its ``w:trPr`` repeats every row as a header row), a grid of 2000 × 3, 3 rows "Item | Count | Remarks" /
  "Pumps | 3 | " / "Valves | 12 | ";
  "END OF SECTION".
* header: table 0 — no style, 1 row "SECTION | 23 21 13"; a paragraph after it.
"""
from __future__ import annotations

from typing import Any, List

from docx import Document
from docx.oxml.ns import qn
from lxml import etree

from make_fixtures import _add_style, _fill, _row_first, _tcpr_add, docx_bytes

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W = "{%s}" % W_NS
_DECL = f'xmlns:w="{W_NS}"'
TITLE = "SECTION 23 21 13 - HYDRONIC PIPING"


def _xml(text: str) -> List[Any]:
    return list(etree.fromstring(f"<w:x {_DECL}>{text}</w:x>"))


def _run_props(cell, xml: str) -> None:
    for p in cell.paragraphs:
        for r in p.runs:
            rpr = r._r.get_or_add_rPr()
            for el in _xml(xml):
                rpr.append(el)


def _grid(table, widths: List[int], table_width: str = "") -> None:
    """The grid's widths, every cell's ``w:tcW`` its columns' (dxa), and — when given — the table's ``w:tblW``
    (``'dxa:8640'``, ``'pct:5000'``)."""
    tbl = table._tbl
    for gc, w in zip(tbl.find(qn("w:tblGrid")).findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(w))
    for tr in tbl.findall(qn("w:tr")):
        for c, tc in enumerate(tr.findall(qn("w:tc"))):
            tcw = tc.get_or_add_tcPr().find(qn("w:tcW"))
            tcw.set(qn("w:w"), str(widths[c]))
            tcw.set(qn("w:type"), "dxa")
    if table_width:
        kind, w = table_width.split(":")
        tblw = tbl.tblPr.find(qn("w:tblW"))
        tblw.set(qn("w:type"), kind)
        tblw.set(qn("w:w"), w)


def _span(tr, first: int, span: int, width: str) -> None:
    """The row's cell ``first`` spans ``span`` columns (the next ones removed), its ``w:tcW`` set to ``width``."""
    tcs = tr.findall(qn("w:tc"))
    tcpr = tcs[first].get_or_add_tcPr()
    tcw = tcpr.find(qn("w:tcW"))
    tcw.set(qn("w:w"), width)
    tcw.addnext(_xml(f'<w:gridSpan w:val="{span}"/>')[0])
    for tc in tcs[first + 1:first + span]:
        tr.remove(tc)


def build_grid() -> bytes:
    doc = Document()
    _add_style(doc, '<w:style w:type="table" w:customStyle="1" w:styleId="ArchShaded"><w:name w:val="Arch Shaded"/>'
                    '<w:basedOn w:val="TableNormal"/><w:trPr><w:tblHeader/></w:trPr><w:tcPr><w:shd w:val="clear" '
                    'w:color="auto" w:fill="F2F2F2"/><w:noWrap/></w:tcPr></w:style>')
    sec = doc.sections[0]
    ht = sec.header.add_table(1, 2, sec.page_width - sec.left_margin - sec.right_margin)
    _fill(ht, [["SECTION", "23 21 13"]])
    sec.header.add_paragraph("")

    doc.add_paragraph(TITLE)
    doc.add_paragraph("1.1 PUMP SCHEDULE")
    t0 = doc.add_table(rows=5, cols=4)
    t0.style = doc.styles["Table Grid"]
    _fill(t0, [["Tag", "Description", "Size", "Notes"], ["P-1", "Base-mounted pump", "2 in", "Duty"],
               ["P-2", "Base-mounted pump", "2 in", "Standby"], ["P-3", "Inline pump", "1 in", ""],
               ["Totals", "", "5 in", ""]])
    _grid(t0, [1440, 2880, 2160, 2160], "dxa:8640")
    _row_first(t0.rows[0], "<w:trPr><w:tblHeader/></w:trPr>")
    for c in range(4):
        _tcpr_add(t0.cell(0, c), '<w:shd w:val="clear" w:color="auto" w:fill="D9D9D9"/>')
        _run_props(t0.cell(0, c), "<w:b/>")
    _tcpr_add(t0.cell(2, 3), '<w:vMerge w:val="restart"/>')
    _tcpr_add(t0.cell(3, 3), '<w:vMerge/>')
    _span(t0.rows[4]._tr, 0, 2, "4320")

    doc.add_paragraph("1.2 ZONE SCHEDULE")
    t1 = doc.add_table(rows=4, cols=3)
    t1.style = doc.styles["Table Grid"]
    _fill(t1, [["Zone", "Area", "Floor"], ["", "North", "2"], ["South", "Annex", ""], ["East", "Ground", "1"]])
    _grid(t1, [3120, 3120, 3120], "pct:5000")
    tr1 = t1.rows[1]._tr
    tr1.remove(tr1.findall(qn("w:tc"))[0])
    _row_first(t1.rows[1], '<w:trPr><w:gridBefore w:val="1"/><w:wBefore w:w="3120" w:type="dxa"/></w:trPr>')
    tr2 = t1.rows[2]._tr
    tr2.remove(tr2.findall(qn("w:tc"))[2])
    _row_first(t1.rows[2], '<w:trPr><w:gridAfter w:val="1"/><w:wAfter w:w="3120" w:type="dxa"/></w:trPr>')
    for tc in t1.rows[3]._tr.findall(qn("w:tc")):
        tcw = tc.find(qn("w:tcPr")).find(qn("w:tcW"))
        tcw.set(qn("w:type"), "pct")
        tcw.set(qn("w:w"), "1667")

    doc.add_paragraph("1.3 VALVE SCHEDULE")
    t2 = doc.add_table(rows=3, cols=2)
    _fill(t2, [["Valve", "Detail"], ["Gate", ""], ["Check", "Swing"]])
    _grid(t2, [2400, 2400])
    cell = t2.cell(1, 1)
    t3 = cell.add_table(2, 2)                          # python-docx: [the cell's empty p, the table, a new empty p]
    t3.style = doc.styles["Table Grid"]
    _fill(t3, [["Size", "Rating"], ["4 in", "175 psi"]])
    _grid(t3, [1200, 1200])
    cell.paragraphs[-1].text = "See detail."
    first = cell.paragraphs[0]._p
    first.getparent().remove(first)

    doc.add_paragraph("1.4 EQUIPMENT")
    t4 = doc.add_table(rows=3, cols=3)
    t4.style = doc.styles["Arch Shaded"]
    _fill(t4, [["Item", "Count", "Remarks"], ["Pumps", "3", ""], ["Valves", "12", ""]])
    _grid(t4, [2000, 2000, 2000])
    doc.add_paragraph("END OF SECTION")
    return docx_bytes(doc)
