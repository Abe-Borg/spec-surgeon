"""The dirty fixture family (``make_fixtures.build_dirty`` / ``DIRTY_KINDS``): one valid, recognizable specification
per kind, each carrying exactly one deliberate defect family — the shapes real project files bring (fragmented runs,
text-box headers, non-breaking spaces, compatibility mode, hundreds of tracked changes, protection, LibreOffice output,
an empty body) beside a well-styled one.

Three things are pinned. (1) Every kind builds, opens (``DocxPackage``, python-docx) and fingerprints as its defect
says — and carries none of the other kinds' defects. (2) A mechanical edit still plans, applies and verifies on every
kind it applies to: the shipped header/footer removal spec on every kind with a header, a whole-document font change on
every kind with body text. (3) One refusal the engine makes deterministically on a dirty shape stays a refusal: the
issue stamp inside a header text box (``set_issue_stamp`` → ``textbox``), at plan time and in the preview."""
from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from lxml import etree

import spec_surgeon
from make_fixtures import DIRTY_KINDS, DIRTY_STAMP, W, build_dirty
from spec_surgeon.editspec import load_spec, plan_spec
from spec_surgeon.fingerprint import fingerprint
from spec_surgeon.ops import JobException
from spec_surgeon.ops.selectors import DEFAULT_PRESERVE_FONTS
from spec_surgeon.package import DocxPackage
from spec_surgeon.preview import make_preview

SPECS = Path(spec_surgeon.__file__).parent / "examples" / "specs"

FONT_SPEC = {
    "request": "change all fonts to Arial",
    "summary": "Set the Latin and complex-script fonts to Arial on every run and paragraph mark of the body.",
    "mode": "formatting",
    "ops": [{"op": "set_font", "params": {"font": "Arial"}}],
}


@pytest.fixture(scope="module")
def dirty(tmp_path_factory) -> dict:
    """{kind: path} — every kind built once, written where the previews read them from."""
    folder = tmp_path_factory.mktemp("dirty")
    paths = {}
    for kind in DIRTY_KINDS:
        path = folder / f"dirty-{kind}.docx"
        path.write_bytes(build_dirty(kind))
        paths[kind] = path
    return paths


def _settings(pkg: DocxPackage):
    return pkg.xml("word/settings.xml")


def _compat_mode(pkg: DocxPackage):
    for cs in _settings(pkg).iter(W + "compatSetting"):
        if cs.get(W + "name") == "compatibilityMode":
            return cs.get(W + "val")
    return None


def _protection(pkg: DocxPackage):
    return _settings(pkg).find(W + "documentProtection")


def _revisions(fp: dict) -> int:
    return sum(d["ins"] + d["del"] for d in fp["parts"].values())


def _hf_textboxes(fp: dict) -> int:
    return sum(d["textboxes"] for part, d in fp["parts"].items() if "header" in part or "footer" in part)


def _rsid_sightings(pkg: DocxPackage) -> list:
    """Every rsid attribute or element, by member — what LibreOffice never writes."""
    out = []
    for name in pkg.member_names:
        if not name.endswith(".xml"):
            continue
        root = pkg.xml(name)
        for el in root.iter():
            if not isinstance(el.tag, str):
                continue
            if etree.QName(el).localname.startswith("rsid"):
                out.append((name, etree.QName(el).localname))
            out += [(name, "@" + etree.QName(k).localname) for k in el.attrib if etree.QName(k).localname.startswith("rsid")]
    return out


def _style(pkg: DocxPackage, sid: str):
    return next((s for s in pkg.xml("word/styles.xml").iter(W + "style") if s.get(W + "styleId") == sid), None)


def _body_text(path: Path) -> list:
    return [p.text for p in Document(str(path)).paragraphs]


def _has_headers(fp: dict) -> bool:
    return any(s["headers"] for s in fp["sections"])


# ---- (1) every kind builds, opens and fingerprints as its defect says ------------------------------------------------

@pytest.mark.parametrize("kind", DIRTY_KINDS)
def test_every_kind_builds_opens_and_fingerprints(dirty, kind):
    path = dirty[kind]
    pkg = DocxPackage.from_path(path)
    fp = fingerprint(pkg)
    assert fp["pattern_id"] and fp["main_part"] == "word/document.xml" and fp["sections"]
    doc = Document(str(path))                                          # python-docx opens it too
    assert doc.sections
    text = "\n".join(_body_text(path))
    if kind == "empty_body":
        assert not text.strip()
    else:
        assert "SECTION 21 13 13" in text and "WET-PIPE SPRINKLER SYSTEMS" in text and "PART 1 - GENERAL" in text


def test_the_kinds_are_the_published_tuple_and_an_unknown_kind_is_refused():
    assert DIRTY_KINDS == ("fragmented_runs", "textbox_headers", "nbsp_spaces", "compat_mode", "many_revisions",
                           "protected", "libreoffice_like", "empty_body", "styled")
    with pytest.raises(ValueError, match="unknown dirty kind"):
        build_dirty("tidy")


def test_fragmented_runs_fragment_the_runs_not_the_text(dirty):
    fp = fingerprint(DocxPackage.from_path(dirty["fragmented_runs"]))
    assert fp["body_hygiene"]["run_fragmentation_ratio"] > 4
    assert _body_text(dirty["fragmented_runs"]) == _body_text(dirty["protected"])   # the visible text is the plain spec's
    root = DocxPackage.from_path(dirty["fragmented_runs"]).xml("word/document.xml")
    rsids = {r.get(W + "rsidR") for r in root.iter(W + "r")}
    assert len(rsids) > 100                                             # rsid-like noise differs run to run


def test_textbox_headers_put_the_header_and_footer_in_text_boxes(dirty):
    pkg = DocxPackage.from_path(dirty["textbox_headers"])
    fp = fingerprint(pkg)
    assert _hf_textboxes(fp) > 0 and fp["pattern"]["has_textboxes_in_hf"] is True
    hf = [p for p in fp["parts"] if "header" in p or "footer" in p]
    assert all(fp["parts"][p]["textboxes"] > 0 for p in hf) and len(hf) == 2
    header = next(p for p in hf if "header" in p)
    boxed = "".join(t.text or "" for t in pkg.xml(header).iter(W + "txbxContent") for t in t.iter(W + "t"))
    assert DIRTY_STAMP in boxed and "SECTION 21 13 13" in boxed
    footer = next(p for p in hf if "footer" in p)
    assert "PAGE" in fp["parts"][footer]["fields"]


def test_nbsp_spaces_set_both_whitespace_flags(dirty):
    fp = fingerprint(DocxPackage.from_path(dirty["nbsp_spaces"]))
    assert fp["body_hygiene"]["has_nbsp"] is True and fp["body_hygiene"]["has_multi_space"] is True
    text = "\n".join(_body_text(dirty["nbsp_spaces"]))
    assert "175 psi" in text and "4 in" in text and "NFPA 13" in text and "  " in text


def test_compat_mode_is_word_2010_mode_in_settings(dirty):
    pkg = DocxPackage.from_path(dirty["compat_mode"])
    assert _compat_mode(pkg) == "14"
    app = pkg.xml("docProps/app.xml")
    assert app.find("{%s}AppVersion" % etree.QName(app).namespace).text == "14.0000"


def test_many_revisions_has_hundreds_pending_with_tracking_on(dirty):
    pkg = DocxPackage.from_path(dirty["many_revisions"])
    fp = fingerprint(pkg)
    assert _revisions(fp) >= 300 and fp["trackRevisions"] is True
    assert fp["body_hygiene"]["run_fragmentation_ratio"] <= 4          # the revisions alone never read as fragmentation
    main = fp["parts"]["word/document.xml"]
    assert main["ins"] >= 150 and main["del"] >= 150
    root = pkg.xml("word/document.xml")
    assert {e.get(W + "author") for e in root.iter(W + "ins")} == {"Jeremy"}
    assert {e.get(W + "author") for e in root.iter(W + "del")} == {"Jack"}
    ids = [e.get(W + "id") for e in root.iter(W + "ins", W + "del")]
    assert len(ids) == len(set(ids))                                    # every revision id unique


def test_protected_enforces_read_only_protection(dirty):
    dp = _protection(DocxPackage.from_path(dirty["protected"]))
    assert dp is not None and dp.get(W + "edit") == "readOnly" and dp.get(W + "enforcement") == "1"
    settings = _settings(DocxPackage.from_path(dirty["protected"]))
    names = [etree.QName(c).localname for c in settings if isinstance(c.tag, str)]
    assert names.index("documentProtection") < names.index("defaultTabStop")   # in CT_Settings order


def test_libreoffice_like_has_no_rsids_and_libreoffice_styles(dirty):
    pkg = DocxPackage.from_path(dirty["libreoffice_like"])
    assert _rsid_sightings(pkg) == []
    body = _style(pkg, "TextBody")
    assert body is not None and body.find(W + "name").get(W + "val") == "Text Body"
    assert _style(pkg, "Heading1") is not None
    app = pkg.xml("docProps/app.xml")
    assert app.find("{%s}Application" % etree.QName(app).namespace).text.startswith("LibreOffice/7.6")
    settings = _settings(pkg)
    assert len([c for c in settings if isinstance(c.tag, str)]) <= 6
    assert not any(n in pkg.member_names for n in ("word/stylesWithEffects.xml", "word/webSettings.xml",
                                                   "docProps/thumbnail.jpeg", "customXml/item1.xml"))
    assert all(r.resolved is None or pkg.has(r.resolved) or r.target_mode == "External"
               for _, r in pkg.relationships(source=True))            # no relationship dangles
    fp = fingerprint(pkg)
    assert fp["body_hygiene"]["dominant_style"] == "TextBody"


def test_empty_body_has_a_sectpr_and_one_empty_paragraph_only(dirty):
    pkg = DocxPackage.from_path(dirty["empty_body"])
    fp = fingerprint(pkg)
    body = pkg.xml("word/document.xml").find(W + "body")
    assert [etree.QName(c).localname for c in body] == ["p", "sectPr"]
    assert fp["parts"]["word/document.xml"]["paragraphs"] == 1
    h = fp["body_hygiene"]
    assert h["section_anchor_shape"] == "missing" and h["numbering_regime"] == "none" and h["run_fragmentation_ratio"] == 0
    assert _has_headers(fp)                                             # the header still names the section


def test_styled_is_the_styled_regime(dirty):
    fp = fingerprint(DocxPackage.from_path(dirty["styled"]))
    assert fp["body_hygiene"]["style_regime"] == "styled" and fp["body_hygiene"]["custom_styles"]


@pytest.mark.parametrize("kind", DIRTY_KINDS)
def test_each_kind_carries_only_its_own_defect(dirty, kind):
    pkg = DocxPackage.from_path(dirty[kind])
    fp = fingerprint(pkg)
    h = fp["body_hygiene"]
    assert (h["run_fragmentation_ratio"] > 4) == (kind == "fragmented_runs")
    assert (_hf_textboxes(fp) > 0) == (kind == "textbox_headers")
    assert (h["has_nbsp"] or h["has_multi_space"]) == (kind == "nbsp_spaces")
    assert (_compat_mode(pkg) == "14") == (kind == "compat_mode")
    assert (_revisions(fp) > 0) == (kind == "many_revisions")
    assert fp["trackRevisions"] == (kind == "many_revisions")
    assert (_protection(pkg) is not None) == (kind == "protected")
    assert (_rsid_sightings(pkg) == []) == (kind == "libreoffice_like")
    assert (fp["parts"]["word/document.xml"]["paragraphs"] == 1) == (kind == "empty_body")
    # a LibreOffice file is styled by nature (Text Body / Heading 1); the plain kinds are Normal-only typed specs
    assert (h["style_regime"] == "styled") == (kind in ("styled", "libreoffice_like"))
    if kind not in ("styled", "libreoffice_like", "empty_body"):
        assert (h["numbering_regime"], h["section_anchor_shape"]) == ("manual_text", "standalone")
    assert fp["junk_members"] == []


# ---- (2) a mechanical edit still verifies on every kind ------------------------------------------------------------

def _preview(spec, path: Path, out: Path):
    pv = make_preview(spec, [str(path)], out, renderer="none", do_render=False)
    assert not pv.batch_plan.exceptions, [f.exception for f in pv.batch_plan.files if f.exception]
    assert pv.samples and all(s.exception is None for s in pv.samples), [s.exception for s in pv.samples]
    assert all(s.passed for s in pv.samples), [s.checks for s in pv.samples]
    assert pv.token
    return pv


@pytest.mark.parametrize("kind", DIRTY_KINDS)
def test_removing_headers_and_footers_verifies_on_every_kind_with_a_header(dirty, kind, tmp_path):
    fp = fingerprint(DocxPackage.from_path(dirty[kind]))
    if not _has_headers(fp):
        pytest.skip(f"{kind} has no header")
    spec = load_spec(SPECS / "remove-all-headers-footers.json")
    pv = _preview(spec, dirty[kind], tmp_path / "out")
    out = DocxPackage(pv.samples[0].output.read_bytes())
    assert not any(s.header_refs or s.footer_refs for s in out.sections())
    assert not any("header" in n or "footer" in n for n in out.member_names)
    assert _body_text(pv.samples[0].output) == _body_text(dirty[kind])   # the body is untouched


@pytest.mark.parametrize("kind", DIRTY_KINDS)
def test_a_whole_document_font_change_verifies_on_every_kind_with_body_text(dirty, kind, tmp_path):
    if not any(t.strip() for t in _body_text(dirty[kind])):
        pytest.skip(f"{kind} has no body text")
    pv = _preview(FONT_SPEC, dirty[kind], tmp_path / "out")
    assert _body_text(pv.samples[0].output) == _body_text(dirty[kind])   # formatting adds no characters
    out = DocxPackage(pv.samples[0].output.read_bytes())
    fonts = {rf.get(W + "ascii") for rf in out.xml("word/document.xml").iter(W + "rFonts")}
    assert fonts <= {"Arial", None} | set(DEFAULT_PRESERVE_FONTS)     # symbol fonts stay: changing them changes glyphs


# ---- (3) one refusal, pinned ---------------------------------------------------------------------------------------

def test_the_issue_stamp_inside_a_header_text_box_is_refused(dirty, tmp_path):
    """The shipped update-issue-stamp spec finds the stamp in the header's text box (Word keeps two copies of it) and
    refuses with code ``textbox`` — at plan time, and the preview lists the file as that exception."""
    spec = load_spec(SPECS / "update-issue-stamp.json")
    pkg = DocxPackage.from_path(dirty["textbox_headers"])
    with pytest.raises(JobException) as e:
        plan_spec(pkg, spec)
    assert e.value.code == "textbox" and "set_issue_stamp" in e.value.reason
    pv = make_preview(spec, [str(dirty["textbox_headers"])], tmp_path / "out", renderer="none", do_render=False)
    refused = [f for f in pv.batch_plan.files if f.exception]
    assert len(refused) == 1 and refused[0].exception["code"] == "textbox"
    assert not pv.samples or all(s.exception is None for s in pv.samples)
    # the same spec plans cleanly where the stamp is ordinary header text
    plan, _ = plan_spec(DocxPackage.from_path(dirty["protected"]), spec)
    assert plan.intended_parts
