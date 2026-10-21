"""Approved comment additions: independent graph proof, retained revisions and authority."""
from __future__ import annotations

import copy
import hashlib
import io
import zipfile
from pathlib import Path

import pytest
from docx import Document
from lxml import etree

from fx_comments import COMMENTS, DOC, build_commented, build_plain
from make_fixtures import docx_bytes, patch_docx
from spec_surgeon.editspec import apply_spec, plan_spec, validate_spec
from spec_surgeon.ops import JobException
from spec_surgeon.ops import add_comments
from spec_surgeon.package import DocxPackage, W
from spec_surgeon.verify.reader import VPackage
from spec_surgeon.verify.suite import run_suite
from spec_surgeon.verify.textcheck import accept_view, reject_view, view_texts, blocks
from spec_surgeon.verify.effects import verify_effects
from spec_surgeon.verify.goals import _hash_json

AUTHOR = "Approved Reviewer"
REASON = "Retain the approved project requirement."


def spec(pattern="Nothing", **params):
    return {"request": "Include these reviewed reasons", "summary": "Add approved review comments", "mode": "untracked",
            "ops": [{"op": "add_review_comments", "select": {"paragraphs": {"text_regex": pattern}},
                     "params": {"text": REASON, "author": AUTHOR, **params}}]}


def run(data, recipe=None):
    recipe = recipe or spec()
    pkg = DocxPackage(data)
    validate_spec(recipe)
    plan, fragments = plan_spec(pkg, recipe)
    out = apply_spec(pkg, fragments).output_bytes()
    rec = run_suite(data, out, plan.to_dict(), "source.docx", "result.docx")
    return plan, out, rec


def failures(record):
    return {c.id: c.summary for c in record.checks if c.status == "FAIL"}


def tree(data, part):
    return VPackage(data).xml(part)


def comment_part(data):
    return VPackage(data).related(add_comments.COMMENTS)


@pytest.mark.parametrize("data,pattern", [
    (build_plain(), "Nothing"),
    (build_commented(side=False), "^G\\."),
    (build_commented(), "^G\\."),
], ids=["plain", "legacy-comments", "modern-comments"])
def test_complete_verification_with_new_legacy_and_modern_comments(data, pattern):
    plan, output, record = run(data, spec(pattern, source="Approved coordination memo, item 3.", initials="AR"))
    assert record.passed, failures(record)
    main = VPackage(data).main_part()
    assert [b.text for b in blocks(tree(data, main))] == [b.text for b in blocks(tree(output, main))]
    old = tree(data, comment_part(data)) if comment_part(data) else None
    new = tree(output, comment_part(output))
    old_count = len(old) if old is not None else 0
    assert len(new) == old_count + 1
    if old is not None:
        assert [etree.tostring(e) for e in old] == [etree.tostring(e) for e in list(new)[:old_count]]
    added = new[-1]
    assert added.get(W + "author") == AUTHOR
    assert [p.find(W + "r/" + W + "t").text for p in added] == [REASON, "Approved coordination memo, item 3."]
    assert plan.rev_delta == {"ins": 0, "del": 0}
    assert plan.untracked_diffs_expected


def test_comment_text_is_verbatim_including_spaces_and_unicode():
    literal = "  Reviewer’s exact requirement — unchanged.  "
    _, output, record = run(build_plain(), spec(text=literal))
    assert record.passed, failures(record)
    t = next(tree(output, comment_part(output)).iter(W + "t"))
    assert t.text == literal
    assert t.get("{http://www.w3.org/XML/1998/namespace}space") == "preserve"


def test_multiple_selected_paragraphs_have_independently_checked_counts():
    d = Document()
    d.add_paragraph("Match first")
    d.add_paragraph("Leave this alone")
    d.add_paragraph("Match second")
    data = docx_bytes(d)
    plan, output, record = run(data, spec("^Match", expected_count=2))
    assert record.passed, failures(record)
    assert [c.get(W + "id") for c in tree(output, comment_part(output))] == ["0", "1"]
    assert plan.property_assertions[0]["expected_count"] == 2
    with pytest.raises(JobException, match="expected 1"):
        run(data, spec("^Match", expected_count=1))


def test_table_cell_comments_and_revised_row_refusal():
    document = Document()
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Selected schedule cell"
    data = docx_bytes(document)
    _, _, result = run(data, spec("Selected"))
    assert result.passed, failures(result)
    def revise(root):
        row = next(root.iter(W + "tr"))
        properties = etree.Element(W + "trPr")
        row.insert(0, properties)
        etree.SubElement(properties, W + "ins", {W + "id": "15", W + "author": AUTHOR, W + "date": "2026-10-01T12:00:00Z"})
    with pytest.raises(JobException, match="revised table"):
        run(patch_docx(data, {DOC: revise}), spec("Selected"))


def test_creation_without_existing_main_relationships_or_rels_default():
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("_rels/.rels", '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="main" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="document.xml"/></Relationships>')
        archive.writestr("[Content_Types].xml", '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="xml" ContentType="application/xml"/>'
            '<Override PartName="/_rels/.rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Override PartName="/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>')
        archive.writestr("document.xml", '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            '<w:body><w:p><w:r><w:t>Nothing to remove here.</w:t></w:r></w:p></w:body></w:document>')
    plan, output, result = run(stream.getvalue())
    assert result.passed, failures(result)
    assert plan.members_added == ["review-comments.xml", "_rels/document.xml.rels"]
    assert comment_part(output) == "review-comments.xml"


def test_explicit_missing_selection_is_a_verified_byte_identical_noop():
    data = build_plain()
    plan, output, record = run(data, spec("MISSING", allow_missing=True, expected_count=0))
    assert output == data
    assert record.passed, failures(record)
    assert not plan.diffs and not plan.members_added and not plan.intended_parts
    with pytest.raises(JobException, match="no review-comment"):
        run(data, spec("MISSING"))


def revised_source():
    data = build_plain("Original ")
    def change(root):
        p = next(root.iter(W + "p"))
        for tag, text in (("del", "old"), ("ins", "new")):
            rev = etree.SubElement(p, W + tag, {W + "id": "9" if tag == "del" else "10", W + "author": "Earlier reviewer", W + "date": "2026-10-01T12:00:00Z"})
            run = etree.SubElement(rev, W + "r")
            etree.SubElement(run, W + ("delText" if tag == "del" else "t")).text = text
    return patch_docx(data, {DOC: change})


def test_anchors_outside_tracked_wrappers_preserve_accept_and_reject_views():
    data = revised_source()
    _, output, record = run(data, spec("^Original"))
    assert record.passed, failures(record)
    before, after = tree(data, DOC), tree(output, DOC)
    for simulator, view in ((accept_view, "accept"), (reject_view, "reject")):
        assert view_texts(simulator(before), view) == view_texts(simulator(after), view)
        simulated = simulator(after)
        for tag in ("commentRangeStart", "commentRangeEnd", "commentReference"):
            assert len(list(simulated.iter(W + tag))) == 1
    for anchor in after.iter(W + "commentRangeStart", W + "commentRangeEnd", W + "commentReference"):
        assert not any(a.tag in {W + "ins", W + "del"} for a in anchor.iterancestors())


@pytest.mark.parametrize("pattern", ["^A\\.", "^D\\.", "^F\\."])
def test_existing_comment_ranges_are_not_overlapped(pattern):
    with pytest.raises(JobException, match="overlaps"):
        run(build_commented(), spec(pattern))


@pytest.mark.parametrize("scenario", ["control", "textbox", "mark_revision", "parent_revision", "empty"])
def test_unsupported_anchor_contexts_fail_before_mutation(scenario):
    data = build_plain()
    def alter(root):
        p = next(root.iter(W + "p"))
        if scenario == "control":
            etree.SubElement(p, W + "sdt")
        elif scenario == "textbox":
            etree.SubElement(etree.SubElement(p, W + "r"), W + "txbxContent")
        elif scenario == "mark_revision":
            ppr = etree.Element(W + "pPr")
            p.insert(0, ppr)
            etree.SubElement(etree.SubElement(ppr, W + "rPr"), W + "ins", {W + "id": "1"})
        elif scenario == "parent_revision":
            parent = p.getparent()
            parent.remove(p)
            etree.SubElement(parent, W + "ins").append(p)
        else:
            for t in p.iter(W + "t"):
                t.text = ""
    data = patch_docx(data, {DOC: alter})
    recipe = spec(".*")
    with pytest.raises(JobException):
        run(data, recipe)


@pytest.mark.parametrize("change", ["missing_start", "moved_end", "wrong_id", "body_word", "body_property", "payload", "author", "extra_comment", "relationship"])
def test_independent_verifier_rejects_damaged_output(change):
    data = revised_source()
    plan, output, record = run(data, spec("^Original"))
    assert record.passed, failures(record)
    part = comment_part(output)
    def damage_doc(root):
        p = next(root.iter(W + "p"))
        if change == "missing_start":
            p.remove(p.find(W + "commentRangeStart"))
        elif change == "moved_end":
            end = p.find(W + "commentRangeEnd")
            p.remove(end)
            p.find(W + "ins").append(end)
        elif change == "wrong_id":
            p.find(W + "r/" + W + "commentReference").set(W + "id", "99")
        elif change == "body_word":
            next(p.iter(W + "t")).text = "Changed without approval"
        else:
            p.set(W + "rsidR", "BADC0DE1")
    def damage_comment(root):
        if change == "payload":
            next(root.iter(W + "t")).text = "Invented reason"
        elif change == "author":
            root[0].set(W + "author", "Different author")
        else:
            root.append(copy.deepcopy(root[0]))
    def damage_rels(root):
        next(e for e in root if e.get("Type") == add_comments.COMMENTS).set("Target", "wrong.xml")
    if change == "relationship":
        damaged = patch_docx(output, {"word/_rels/document.xml.rels": damage_rels})
    elif change in {"payload", "author", "extra_comment"}:
        damaged = patch_docx(output, {part: damage_comment})
    else:
        damaged = patch_docx(output, {DOC: damage_doc})
    result = run_suite(data, damaged, plan.to_dict(), "before", "after")
    assert result.get("V10").status == "FAIL", failures(result)


@pytest.mark.parametrize("member", [COMMENTS, "word/commentsExtended.xml", "word/commentsIds.xml", "word/commentsExtensible.xml"])
def test_existing_comments_and_side_metadata_cannot_be_changed(member):
    data = build_commented()
    plan, output, record = run(data, spec("^G\\."))
    assert record.passed, failures(record)
    damaged = patch_docx(output, {member: lambda root: root.remove(root[0])})
    result = run_suite(data, damaged, plan.to_dict(), "before", "after")
    assert result.get("V10").status == "FAIL"


def test_shared_plan_and_transform_wrong_target_is_not_accepted(monkeypatch):
    d = Document()
    d.add_paragraph("Target paragraph")
    d.add_paragraph("Wrong paragraph")
    data = docx_bytes(d)
    original = add_comments.derive
    def wrong(*args, **kwargs):
        result = original(*args, **kwargs)
        result["entries"][0]["paragraph"] = 1
        return result
    monkeypatch.setattr(add_comments, "derive", wrong)
    _, _, record = run(data, spec("^Target"))
    assert record.get("V10").status == "FAIL", failures(record)


def relocated(data, non_xml=False):
    """Fixture-only relocation retains all relationships while changing their names."""
    suffix = "data" if non_xml else "xml"
    names = {DOC: f"parts/main.{suffix}", "word/_rels/document.xml.rels": f"parts/_rels/main.{suffix}.rels"}
    old = VPackage(data)
    if COMMENTS in old.names:
        names[COMMENTS] = f"notes/reviewer.{suffix}"
        if non_xml:
            for name in ("commentsExtended", "commentsIds", "commentsExtensible"):
                if f"word/{name}.xml" in old.names:
                    names[f"word/{name}.xml"] = f"notes/{name}.data"
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as archive:
        for info in old.zip.infolist():
            content = old.read(info.filename)
            if info.filename in {"_rels/.rels", "word/_rels/document.xml.rels", "[Content_Types].xml"}:
                root = etree.fromstring(content)
                for entry in root:
                    if "PartName" in entry.attrib:
                        p = entry.get("PartName").lstrip("/")
                        if p in names:
                            entry.set("PartName", "/" + names[p])
                    elif info.filename == "_rels/.rels" and entry.get("Type", "").endswith("/officeDocument"):
                        entry.set("Target", names[DOC])
                    elif info.filename == "word/_rels/document.xml.rels" and entry.get("TargetMode", "Internal") == "Internal":
                        target = old.rels_targets(DOC)[entry.get("Id")]["target"]
                        entry.set("Target", "/" + names.get(target, target))
                content = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            changed = copy.copy(info)
            changed.filename = names.get(info.filename, info.filename)
            archive.writestr(changed, content)
    return output.getvalue()


@pytest.mark.parametrize("data,pattern", [(build_plain(), "Nothing"), (build_commented(), "^G\\.")],
                         ids=["plain", "modern-comments"])
@pytest.mark.parametrize("non_xml", [False, True])
def test_relationships_resolve_nonstandard_main_and_comment_part_paths(data, pattern, non_xml):
    data = relocated(data, non_xml=non_xml)
    plan, output, record = run(data, spec(pattern))
    assert record.passed, failures(record)
    suffix = "data" if non_xml else "xml"
    assert VPackage(output).main_part() == f"parts/main.{suffix}"
    assert comment_part(output) in {"parts/review-comments.xml", f"notes/reviewer.{suffix}"}
    ids = [c.get(W + "id") for c in tree(output, comment_part(output)).findall(W + "comment")]
    assert len(ids) == len(set(ids))


def contract_for(data, recipe=None):
    recipe = recipe or spec()
    op = recipe["ops"][0]
    goal = {"kind": "document_effects", "effects": [{"id": "comments", "type": "review_comments",
            "select": op["select"], "params": op["params"], "expected_count": 1}]}
    binding = {"source_sha256": hashlib.sha256(data).hexdigest(), "contract_sha256": _hash_json(goal),
               "effects": {"comments": {"targets": [], "choice": "unique", "source_count": 1}}}
    return goal, binding


@pytest.mark.parametrize("data,pattern", [(build_plain(), "Nothing"), (build_commented(), "^G\\.")],
                         ids=["plain", "modern-comments"])
def test_original_to_final_outcome_proof_uses_no_execution_plan(data, pattern):
    recipe = spec(pattern, source="Approved source reference")
    _, output, record = run(data, recipe)
    assert record.passed, failures(record)
    goal, binding = contract_for(data, recipe)
    result = verify_effects(data, output, goal, binding)
    assert result["passed"], result
    altered = patch_docx(output, {comment_part(output): lambda root: setattr(list(root.iter(W + "t"))[-1], "text", "Invented")})
    assert not verify_effects(data, altered, goal, binding)["passed"]
    altered_style = patch_docx(output, {"word/styles.xml": lambda root: root.set("unapproved", "change")})
    assert not verify_effects(data, altered_style, goal, binding)["passed"]
    assert not verify_effects(data, data, goal, binding)["passed"]
    binding["effects"]["comments"]["source_count"] = 2
    assert not verify_effects(data, output, goal, binding)["passed"]


def test_compilation_carries_only_approved_literal_comment_wording(tmp_path):
    from spec_surgeon.goals import approved_goal_words, compile_goal_candidate, resolve_goal
    data = build_plain()
    path = tmp_path / "source.docx"
    path.write_bytes(data)
    goal, _ = contract_for(data, spec(source="Approved source"))
    assert approved_goal_words(goal) == [REASON, "Approved source"]
    resolution = resolve_goal(path, goal)
    assert resolution["status"] == "resolved"
    candidate = compile_goal_candidate(path, goal, resolution["binding"])
    assert candidate["spec"]["mode"] == "untracked"
    assert candidate["spec"]["ops"][0]["params"]["text"] == REASON


def test_comment_effect_requires_its_own_approved_stage():
    from spec_surgeon.goals import GoalError, validate_contract
    goal, binding = contract_for(build_plain())
    goal["effects"].append({"id": "pages", "type": "page_setup", "margins": {"left": 1440}})
    with pytest.raises(GoalError, match="own approved contract"):
        validate_contract(goal)
    binding["contract_sha256"] = _hash_json(goal)
    assert not verify_effects(build_plain(), build_plain(), goal, binding)["passed"]


def test_whole_part_proof_refuses_same_stage_body_edits():
    recipe = spec()
    recipe["ops"].append({"op": "set_paragraph_props", "select": {"paragraphs": {"text_regex": "Nothing"}},
                          "params": {"jc": "center", "tracked": False}})
    with pytest.raises(JobException) as error:
        plan_spec(DocxPackage(build_plain()), recipe)
    assert error.value.code == "conflict"


def test_shipped_example_and_compiler_require_traceable_literal_wording():
    from spec_surgeon.catalog import REQUEST_ONLY_OPS, op_catalog
    from spec_surgeon.compiler import placed_words, untraceable
    from spec_surgeon.editspec import load_spec
    recipe = load_spec(Path(__file__).resolve().parents[1] / "examples/specs/add-review-comments.json")
    assert "add_review_comments" in REQUEST_ONLY_OPS
    assert op_catalog(examples=False)["ops"]["add_review_comments"]["mode"] == "untracked"
    assert placed_words(recipe) == [REASON, "Source: approved coordination memo, item 3.", "Abraham Borg", "AB"]
    assert untraceable(recipe, [recipe["request"]]) == []
    assert untraceable(recipe, ["Add some useful comments"]) == placed_words(recipe)
    # the reviewer Word shows is the user's too: a name the user never gave is refused like invented wording
    reviewer = copy.deepcopy(recipe)
    reviewer["ops"][0]["params"].update(author="Spec Reviewer", initials="SR")
    assert untraceable(reviewer, [recipe["request"]]) == ["Spec Reviewer", "SR"]
    _, _, result = run(build_plain("A. Provide listed sprinklers."), recipe)
    assert result.passed, failures(result)
