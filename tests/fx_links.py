"""Fixtures for remove_hyperlinks and unlink_fields (built-in batch 2, docs/BUILTIN_OPS_BATCH2_DESIGN.md §9, §10) —
every case is found by its text in the tests, never by index.

``build_links()``: link elements external (a web address, one shared by two links, one also used by a picture's click
link, one with direct blue/underline instead of the style, one in the FollowedHyperlink style, one in a character style
that is not a link style), internal (a bookmark), HYPERLINK fields external and internal (``\\l``), a link in the header
and in a footnote, and fields that are not links (REF, DOCPROPERTY, PAGE in the footer).

``build_fields()``: DOCPROPERTY (body and header), DATE (a simple field), REF, a HYPERLINK field, a DOCPROPERTY nested in
a HYPERLINK field's result, an XE field that shows nothing, a DOCPROPERTY with an empty result, two fields in one
paragraph, PAGE in the footer.

``awkward(kind)``: one refusal each (a text box, a field over two paragraphs, a field in a pending insertion, a field in
another field's code, a form field).
"""
from __future__ import annotations

import io
import zipfile
from typing import List, Optional

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from lxml import etree

from make_fixtures import _page_field, _ref_field, docx_bytes, png

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
W = "{%s}" % W_NS


def styles(doc) -> None:
    h = doc.styles.add_style("Hyperlink", WD_STYLE_TYPE.CHARACTER)
    h.font.color.rgb = RGBColor(0x05, 0x63, 0xC1)
    h.font.underline = True
    f = doc.styles.add_style("FollowedHyperlink", WD_STYLE_TYPE.CHARACTER)
    f.font.color.rgb = RGBColor(0x95, 0x4F, 0x72)
    f.font.underline = True
    e = doc.styles.add_style("Emphasis Two", WD_STYLE_TYPE.CHARACTER)        # not a link style
    e.font.italic = True


def run(text: str = "", style: Optional[str] = None, color: Optional[str] = None, underline: bool = False):
    r = OxmlElement("w:r")
    if style or color or underline:
        rpr = OxmlElement("w:rPr")
        if style:
            rs = OxmlElement("w:rStyle"); rs.set(qn("w:val"), style); rpr.append(rs)
        if color:
            c = OxmlElement("w:color"); c.set(qn("w:val"), color); rpr.append(c)
        if underline:
            u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rpr.append(u)
        r.append(rpr)
    if text:
        t = OxmlElement("w:t"); t.text = text; t.set(qn("xml:space"), "preserve"); r.append(t)
    return r


def link(paragraph, text: str, url: Optional[str] = None, anchor: Optional[str] = None, style: Optional[str] = "Hyperlink",
         **look):
    h = OxmlElement("w:hyperlink")
    if url:
        h.set(qn("r:id"), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
    if anchor:
        h.set(qn("w:anchor"), anchor)
    h.append(run(text, style, **look))
    paragraph._p.append(h)
    return h


def _marker(kind: str, ff: bool = False):
    r = OxmlElement("w:r")
    fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), kind)
    if ff:
        fc.append(OxmlElement("w:ffData"))
    r.append(fc)
    return r


def _instr(text: str):
    r = OxmlElement("w:r")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = text; r.append(it)
    return r


def field(paragraph, instruction: str, result: List, separate: bool = True, ff: bool = False) -> None:
    """A complex field appended to the paragraph: begin, the instruction, separate, the result runs, end."""
    p = paragraph._p if hasattr(paragraph, "_p") else paragraph
    p.append(_marker("begin", ff))
    p.append(_instr(instruction))
    if separate:
        p.append(_marker("separate"))
        for r in result:
            p.append(r)
    p.append(_marker("end"))


def simple(paragraph, instruction: str, result: List) -> None:
    fs = OxmlElement("w:fldSimple")
    fs.set(qn("w:instr"), instruction)
    for r in result:
        fs.append(r)
    paragraph._p.append(fs)


def _picture_with_link(paragraph, url: str) -> None:
    """A picture whose click link (a:hlinkClick) uses the same relationship as a text link to ``url``."""
    rid = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    paragraph.add_run().add_picture(io.BytesIO(png((20, 120, 40), size=6)), width=Pt(18))
    doc_pr = paragraph._p.findall(".//{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}docPr")[-1]
    click = etree.SubElement(doc_pr, "{%s}hlinkClick" % A_NS, nsmap={"a": A_NS})
    click.set("{%s}id" % R_NS, rid)


FOOTNOTES_CT = "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"
FOOTNOTES_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes"


def _with_footnote(data: bytes, url: str) -> bytes:
    """Add a footnotes part whose footnote 1 holds a link to ``url`` (its own rels), and a reference to it."""
    zin = zipfile.ZipFile(io.BytesIO(data))
    members = {i.filename: zin.read(i.filename) for i in zin.infolist()}
    order = [i.filename for i in zin.infolist()]
    members["word/footnotes.xml"] = (
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<w:footnotes xmlns:w="{W_NS}" xmlns:r="{R_NS}">'
        '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
        '<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:r><w:continuationSeparator/></w:r></w:p>'
        '</w:footnote><w:footnote w:id="1"><w:p><w:r><w:footnoteRef/></w:r><w:r><w:t xml:space="preserve"> See </w:t>'
        '</w:r><w:hyperlink r:id="rId1"><w:r><w:rPr><w:rStyle w:val="Hyperlink"/></w:rPr><w:t>nfpa.org</w:t></w:r>'
        '</w:hyperlink><w:r><w:t>.</w:t></w:r></w:p></w:footnote></w:footnotes>').encode()
    members["word/_rels/footnotes.xml.rels"] = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<Relationships '
        'xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" '
        f'Type="{RT.HYPERLINK}" Target="{url}" TargetMode="External"/></Relationships>').encode()
    order += ["word/footnotes.xml", "word/_rels/footnotes.xml.rels"]
    rels = etree.fromstring(members["word/_rels/document.xml.rels"])
    el = etree.SubElement(rels, "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship")
    el.set("Id", "rId900"); el.set("Type", FOOTNOTES_REL); el.set("Target", "footnotes.xml")
    members["word/_rels/document.xml.rels"] = etree.tostring(rels, xml_declaration=True, encoding="UTF-8",
                                                             standalone=True)
    ct = etree.fromstring(members["[Content_Types].xml"])
    ov = etree.SubElement(ct, "{http://schemas.openxmlformats.org/package/2006/content-types}Override")
    ov.set("PartName", "/word/footnotes.xml"); ov.set("ContentType", FOOTNOTES_CT)
    members["[Content_Types].xml"] = etree.tostring(ct, xml_declaration=True, encoding="UTF-8", standalone=True)
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as z:
        for name in order:
            zi = zipfile.ZipInfo(name, date_time=(2026, 1, 1, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            z.writestr(zi, members[name])
    return out.getvalue()


def build_links(footnote: bool = True) -> bytes:
    doc = Document()
    styles(doc)
    sec = doc.sections[0]
    hp = sec.header.paragraphs[0]
    hp.add_run("ACME DATA CENTERS — ")
    link(hp, "acme.example", url="https://acme.example/")
    fp = sec.footer.paragraphs[0]
    fp.add_run("Page ")
    _page_field(fp)
    doc.add_paragraph("SECTION 21 13 13 - WET-PIPE SPRINKLER SYSTEMS")
    p = doc.add_paragraph("A.\tSee the ")
    link(p, "Owner portal", url="https://example.com/owner")
    p.add_run(" for drawings.")
    p = doc.add_paragraph("B.\tRefer to ")
    link(p, "Article 1.02", anchor="_Ref1")
    p.add_run(".")
    p = doc.add_paragraph("C.\tThe spec is at ")
    field(p, ' HYPERLINK "https://example.com/spec" \\o "the spec" ', [run("example.com/spec", "Hyperlink")])
    p.add_run(".")
    p = doc.add_paragraph("D.\tContents: ")
    field(p, ' HYPERLINK \\l "_Toc1" ', [run("Part 1 General", "Hyperlink")])
    p = doc.add_paragraph("E.\tOne address twice: ")
    link(p, "first copy", url="https://example.com/shared")
    p.add_run(" and ")
    link(p, "second copy", url="https://example.com/shared")
    p = doc.add_paragraph("F.\tA hand-made link: ")
    link(p, "plain blue link", url="https://example.com/plain", style=None, color="0000FF", underline=True)
    p = doc.add_paragraph("G.\tA picture shares an address: ")
    _picture_with_link(p, "https://example.com/pic")
    link(p, "picture site", url="https://example.com/pic")
    p = doc.add_paragraph("H.\tVisited: ")
    link(p, "followed link", url="https://example.com/followed", style="FollowedHyperlink")
    p = doc.add_paragraph("I.\tOwn style: ")
    link(p, "emphasised link", url="https://example.com/em", style="EmphasisTwo")
    p = doc.add_paragraph("J.\tSee ")
    _ref_field(p)
    p = doc.add_paragraph("K.\tTitle: ")
    field(p, " DOCPROPERTY Title ", [run("Spec Title")])
    p = doc.add_paragraph("L.\tA quoted switch is an address: ")
    field(p, ' HYPERLINK "https://example.com/a\\l" ', [run("example.com/a", "Hyperlink")])
    p = doc.add_paragraph("N.\tPending: ")
    ins = OxmlElement("w:ins"); ins.set(qn("w:id"), "41"); ins.set(qn("w:author"), "Jeremy")
    ins.set(qn("w:date"), "2026-09-22T10:15:00Z")
    p._p.append(ins)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), p.part.relate_to("https://example.com/ins", RT.HYPERLINK, is_external=True))
    h.append(run("inserted link", "Hyperlink"))
    ins.append(h)                                                              # a link inside a pending insertion
    p = doc.add_paragraph("O.\tSimple: ")
    simple(p, ' HYPERLINK "https://example.com/simple" ', [run("simple link", "Hyperlink")])
    if footnote:
        p = doc.add_paragraph("M.\tFootnoted")
        fr = OxmlElement("w:r"); ref = OxmlElement("w:footnoteReference"); ref.set(qn("w:id"), "1"); fr.append(ref)
        p._p.append(fr)
    data = docx_bytes(doc)
    return _with_footnote(data, "https://nfpa.org/") if footnote else data


def build_fields() -> bytes:
    doc = Document()
    styles(doc)
    sec = doc.sections[0]
    hp = sec.header.paragraphs[0]
    hp.add_run("Header title: ")
    field(hp, " DOCPROPERTY Title ", [run("Spec Title")])
    fp = sec.footer.paragraphs[0]
    fp.add_run("Page ")
    _page_field(fp)
    doc.add_paragraph("SECTION 21 13 13 - WET-PIPE SPRINKLER SYSTEMS")
    p = doc.add_paragraph("A.\tTitle: ")
    field(p, " DOCPROPERTY Title \\* MERGEFORMAT ", [run("Spec Title")])
    p.add_run(" (from the properties).")
    p = doc.add_paragraph("B.\tYear: ")
    simple(p, ' DATE \\@ "yyyy" ', [run("2026")])
    p = doc.add_paragraph("C.\tSee ")
    _ref_field(p)
    p = doc.add_paragraph("D.\tLink: ")
    field(p, ' HYPERLINK "https://example.com/" ', [run("example.com", "Hyperlink")])
    p = doc.add_paragraph("E.\tCompany: ")
    p._p.append(_marker("begin"))
    p._p.append(_instr(' HYPERLINK "https://contoso.example/" '))
    p._p.append(_marker("separate"))
    field(p, " DOCPROPERTY Company ", [run("Contoso", "Hyperlink")])          # nested in the link's result
    p._p.append(_marker("end"))
    p = doc.add_paragraph("F.\tIndex entry: ")
    field(p, ' XE "sprinklers" ', [], separate=False)                          # shows nothing
    p.add_run("here.")
    p = doc.add_paragraph("G.\tManager: ")
    field(p, " DOCPROPERTY Manager ", [])                                       # an empty result
    p = doc.add_paragraph("H.\tTwo: ")
    field(p, " DOCPROPERTY Title ", [run("Spec Title")])
    p.add_run(" / ")
    field(p, " DOCPROPERTY Subject ", [run("Sprinklers")])
    doc.add_paragraph("I.\tNo field here.")
    return docx_bytes(doc)


def _textbox_with(paragraph, inner_xml: str) -> None:
    xml = (f'<w:r xmlns:w="{W_NS}" xmlns:r="{R_NS}" xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
           'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
           f'xmlns:a="{A_NS}" xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
           '<mc:AlternateContent><mc:Choice Requires="wps"><w:drawing><wp:inline><wp:extent cx="914400" cy="457200"/>'
           '<wp:docPr id="901" name="Text Box 1"/><a:graphic><a:graphicData '
           'uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"><wps:wsp><wps:txbx><w:txbxContent>'
           f'<w:p>{inner_xml}</w:p></w:txbxContent></wps:txbx><wps:bodyPr/></wps:wsp></a:graphicData></a:graphic>'
           '</wp:inline></w:drawing></mc:Choice><mc:Fallback/></mc:AlternateContent></w:r>')
    paragraph._p.append(etree.fromstring(xml))


def awkward(kind: str) -> bytes:
    """One case a selected link or field must be refused for."""
    doc = Document()
    styles(doc)
    doc.add_paragraph("Intro.")
    if kind == "textbox-link":
        p = doc.add_paragraph("A box: ")
        rid = p.part.relate_to("https://example.com/box", RT.HYPERLINK, is_external=True)
        _textbox_with(p, f'<w:hyperlink r:id="{rid}"><w:r><w:t>boxed link</w:t></w:r></w:hyperlink>')
    elif kind == "textbox-field":
        p = doc.add_paragraph("A box: ")
        _textbox_with(p, '<w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText> DOCPROPERTY Title '
                         '</w:instrText></w:r><w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:t>T</w:t></w:r>'
                         '<w:r><w:fldChar w:fldCharType="end"/></w:r>')
    elif kind == "span":
        p = doc.add_paragraph()
        p._p.append(_marker("begin"))
        p._p.append(_instr(' HYPERLINK "https://example.com/toc" '))
        p._p.append(_marker("separate"))
        p.add_run("first line")
        p2 = doc.add_paragraph("second line")
        p2._p.append(_marker("end"))
    elif kind == "revision":
        p = doc.add_paragraph("Inserted: ")
        ins = OxmlElement("w:ins"); ins.set(qn("w:id"), "5"); ins.set(qn("w:author"), "Jeremy")
        ins.set(qn("w:date"), "2026-09-22T10:15:00Z")
        p._p.append(ins)
        field(ins, ' HYPERLINK "https://example.com/ins" ', [run("ins link", "Hyperlink")])
        field(ins, " DOCPROPERTY Title ", [run("Spec Title")])
    elif kind == "nested":
        p = doc.add_paragraph("Nested: ")
        p._p.append(_marker("begin"))
        p._p.append(_instr(" IF "))
        field(p, " MERGEFIELD Owner ", [run("Contoso")])
        p._p.append(_instr(' = "Contoso" '))
        field(p, ' HYPERLINK "https://example.com/if" ', [run("yes", "Hyperlink")])
        p._p.append(_marker("separate"))
        p._p.append(run("yes"))
        p._p.append(_marker("end"))
    elif kind == "toc":                                                         # entries rebuilt by the TOC
        p = doc.add_paragraph()
        p._p.append(_marker("begin"))
        p._p.append(_instr(' TOC \\o "1-3" \\h '))
        p._p.append(_marker("separate"))
        link(p, "PART 1 - GENERAL", anchor="_Toc1")
        p2 = doc.add_paragraph()
        link(p2, "PART 2 - PRODUCTS", anchor="_Toc2")
        p2._p.append(_marker("end"))
        doc.add_paragraph("After the contents.")
    elif kind == "form":
        p = doc.add_paragraph("Fill in: ")
        field(p, " FORMTEXT ", [run("Name")], ff=True)
    else:
        raise ValueError(kind)
    return docx_bytes(doc)
