"""Fixtures for built-in batch 6, F3 — text and character formatting (docs/BUILTIN_OPS_BATCH6_DESIGN.md §3, §4): every
case is found by its text in the tests, never by index.

``build_text()``: a fire-protection section whose runs already carry what the new run properties set, so each one's
"already as asked", nearer-style and toggle cases are there —

* styles: ``SpecPart`` (the PART headings, plain), ``SpecHidden`` (a paragraph style hiding its text), ``SpecHiddenChar``
  (a character style hiding its text: inside a SpecHidden paragraph, w:vanish at two style levels), ``SpecDouble`` (a
  character style with double strikethrough), ``SpecSpaced`` (a paragraph style expanding its text by 1 pt),
  ``SpecSuper`` (a character style in superscript), ``SpecPartLow`` (based on SpecPart, its own baseline and light grey
  highlight: a nearer style's own values);
* body: "A." with a direct single and a direct double strikethrough, an "m2" whose 2 is already superscript and an
  "m3" whose 3 is superscript by SpecSuper;
  "B." with yellow and green highlighting; "C." with a hidden word; "D." with 2 pt of character spacing and with
  spacing written as a measure ("2pt"); "E." in SpecHidden with a SpecHiddenChar run; "F." with a SpecDouble run; "G."
  in SpecSpaced; "H." with a Wingdings bullet and a w:sym run; "I." with a run formatted earlier under tracked changes;
  "J." with a complex-script run; "K." an anchor with direct paragraph and run formatting (an indent, keep-with-next;
  Courier New, bold, 9 pt, a red single underline); "L." a w:t that itself holds a literal tab character beside the
  words "REPLACE ME"; "M." whose "hot-dip galvanized" sits in Jeremy's pending insertion; "N." a hyperlink; the last
  paragraph of the first section (it carries the section break); PART 2 in SpecPart, PART 3 in SpecPartLow; a table;
  "END OF SECTION";
* the header holds "PROJECT NAME"; trackRevisions is on.
"""
from __future__ import annotations

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from lxml import etree

from make_fixtures import _W_DECL, _hyperlink, _ppr_add, _rev, _rfonts, _rpr_add, docx_bytes

LITERAL_TAB_TEXT = "Item\tliteral, then REPLACE ME here"


def _style_rpr(style, xml: str) -> None:
    """Add rPr children written as XML (``<w:vanish/>``, the w: prefix undeclared) to a python-docx style."""
    rpr = style.element.get_or_add_rPr()
    for el in etree.fromstring(f"<w:x {_W_DECL}>{xml}</w:x>"):
        rpr.append(el)


def build_text(section_number: str = "21 13 13", title: str = "WET-PIPE SPRINKLER SYSTEMS") -> bytes:
    doc = Document()
    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)
    part = styles.add_style("SpecPart", WD_STYLE_TYPE.PARAGRAPH)
    part.base_style = styles["Normal"]
    hidden = styles.add_style("SpecHidden", WD_STYLE_TYPE.PARAGRAPH)
    hidden.base_style = styles["Normal"]
    _style_rpr(hidden, "<w:vanish/>")
    hidden_char = styles.add_style("SpecHiddenChar", WD_STYLE_TYPE.CHARACTER)
    _style_rpr(hidden_char, "<w:vanish/>")
    double = styles.add_style("SpecDouble", WD_STYLE_TYPE.CHARACTER)
    _style_rpr(double, "<w:dstrike/>")
    spaced = styles.add_style("SpecSpaced", WD_STYLE_TYPE.PARAGRAPH)
    spaced.base_style = styles["Normal"]
    _style_rpr(spaced, '<w:spacing w:val="20"/>')
    sup = styles.add_style("SpecSuper", WD_STYLE_TYPE.CHARACTER)
    _style_rpr(sup, '<w:vertAlign w:val="superscript"/>')
    low = styles.add_style("SpecPartLow", WD_STYLE_TYPE.PARAGRAPH)
    low.base_style = part
    _style_rpr(low, '<w:vertAlign w:val="baseline"/><w:highlight w:val="lightGray"/>')

    hp = doc.sections[0].header.paragraphs[0]
    hp.add_run(f"SECTION {section_number} - PROJECT NAME")
    doc.add_paragraph(f"SECTION {section_number}")
    doc.add_paragraph(title)
    doc.add_paragraph("PART 1 - GENERAL", style="SpecPart")
    doc.add_paragraph("1.01\tSUMMARY")
    p = doc.add_paragraph()
    p.add_run("A.\tStrike ")
    _rpr_add(p.add_run("struck"), "<w:strike/>")
    p.add_run(" and ")
    _rpr_add(p.add_run("double struck"), "<w:dstrike/>")
    p.add_run(" in an area of 10 m")
    _rpr_add(p.add_run("2"), '<w:vertAlign w:val="superscript"/>')
    p.add_run(" or a volume of 10 m")
    p.add_run("3").style = styles["SpecSuper"]                                      # superscript by its style
    p = doc.add_paragraph()
    p.add_run("B.\tHighlighted ")
    _rpr_add(p.add_run("yellow note"), '<w:highlight w:val="yellow"/>')
    p.add_run(" and ")
    _rpr_add(p.add_run("green note"), '<w:highlight w:val="green"/>')
    p = doc.add_paragraph()
    p.add_run("C.\tHidden ")
    _rpr_add(p.add_run("secret"), "<w:vanish/>")
    p.add_run(" visible.")
    p = doc.add_paragraph()
    p.add_run("D.\tSpaced ")
    _rpr_add(p.add_run("wide"), '<w:spacing w:val="40"/>')
    p.add_run(" and ")
    _rpr_add(p.add_run("measured"), '<w:spacing w:val="2pt"/>')
    p = doc.add_paragraph("E.\tHidden by its style ", style="SpecHidden")
    p.add_run("twice hidden").style = styles["SpecHiddenChar"]
    p = doc.add_paragraph("F.\tDouble by style ")
    p.add_run("double style").style = styles["SpecDouble"]
    doc.add_paragraph("G.\tSpaced by its style", style="SpecSpaced")
    p = doc.add_paragraph()
    p.add_run("H.\tSymbols ")
    _rfonts(p.add_run("l"), ascii="Wingdings", hAnsi="Wingdings")
    sym_run = p.add_run()
    sym = OxmlElement("w:sym"); sym.set(qn("w:font"), "Symbol"); sym.set(qn("w:char"), "F0B7"); sym_run._r.append(sym)
    p.add_run(" end")
    p = doc.add_paragraph()
    fmt = p.add_run("I.\tFormatted earlier by Jeremy.")
    fmt.bold = True
    chg = OxmlElement("w:rPrChange")
    chg.set(qn("w:id"), "303"); chg.set(qn("w:author"), "Jeremy"); chg.set(qn("w:date"), "2026-09-22T11:00:00Z")
    chg.append(OxmlElement("w:rPr"))
    fmt._r.rPr.append(chg)
    p = doc.add_paragraph()
    p.add_run("J.\t")
    p.add_run("Complex script").font.complex_script = True
    p = doc.add_paragraph()
    _ppr_add(p, '<w:keepNext/><w:ind w:left="720" w:hanging="720"/>')
    for text in ("K.", "\tAnchor with direct formatting"):
        r = p.add_run(text)
        _rpr_add(r, '<w:rFonts w:ascii="Courier New" w:hAnsi="Courier New"/><w:b/><w:sz w:val="18"/>'
                    '<w:u w:val="single" w:color="FF0000"/>')
    p = doc.add_paragraph()
    r = p.add_run()
    t = OxmlElement("w:t"); t.set(qn("xml:space"), "preserve"); t.text = "L.\t" + LITERAL_TAB_TEXT
    # the w:t holds a literal tab character itself (a typed tab Word keeps as text): replacing beside it never turns
    # it into a w:tab
    r._r.append(t)
    p = doc.add_paragraph()
    p.add_run("M.\tInserted by Jeremy: ")
    _rev("ins", p.add_run("hot-dip galvanized"), 501, "Jeremy")
    p.add_run(" steel.")
    p = doc.add_paragraph("N.\tSee the ")
    _hyperlink(p, "Owner standards")
    p.add_run(" online.")
    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.paragraphs[-1].add_run("Last paragraph of the first section")              # it carries the section break
    doc.add_paragraph("PART 2 - PRODUCTS", style="SpecPart")
    doc.add_paragraph("PART 3 - EXECUTION", style="SpecPartLow")
    t = doc.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "Tag"; t.cell(0, 1).text = "Value"
    doc.add_paragraph("END OF SECTION")
    doc.settings.element.insert(0, OxmlElement("w:trackRevisions"))
    return docx_bytes(doc)
