"""A failed recipe may change on retry; the user's reviewed outcome may not disappear."""
import copy
import json

import pytest
from docx import Document

from fake_claude import FakeClient, resp, tool
from spec_surgeon import recovery
from spec_surgeon.batch import run_batch
from spec_surgeon.chain import chain_from_specs, run_chain
from spec_surgeon.goals import resolve_goal
from spec_surgeon.recovery import RecoveryOptions


def margin(value):
    return {"request": "Set left margin to two inches", "summary": "Left margin", "mode": "formatting",
            "ops": [{"op": "set_page_setup", "params": {"margins": {"left": value}}}]}


def failed_run(tmp_path, *, chain=False, in_place=False):
    source = tmp_path / "source.docx"
    doc = Document()
    doc.add_paragraph("Preserve every word.")
    doc.save(source)
    contract = {"kind": "document_effects", "effects": [
        {"id": "margin", "type": "page_setup", "margins": {"left": 2880}}]}
    bindings = {str(source): resolve_goal(source, contract)["binding"]}
    out = tmp_path / "run"
    kwargs = dict(workers=1, do_render=False, in_place=in_place, goal_contract=contract, goal_bindings=bindings)
    if chain:
        result = run_chain([str(source)], chain_from_specs([("wrong margin", margin(1440))]), out, **kwargs)
    else:
        result = run_batch([str(source)], "edit-spec", {}, out, spec=margin(1440), **kwargs)
    assert result.counts()["FAIL"] == 1
    return source, result, contract, bindings


@pytest.mark.parametrize("chain,in_place", [(False, False), (True, False), (True, True)])
def test_retry_retains_outcome_and_delivers_only_a_corrected_candidate(tmp_path, chain, in_place):
    source, result, contract, bindings = failed_run(tmp_path, chain=chain, in_place=in_place)
    original = source.read_bytes()
    for filename in ("set_aside.json", "chain.json" if chain else "batch.json"):
        saved = json.loads((result.out_dir / filename).read_text(encoding="utf-8"))
        assert saved["goal_contract"] == contract
        assert saved["goal_bindings"] == bindings
    fake = FakeClient(resp(tool("submit_repair", {"candidate": {"kind": "spec", "spec": margin(2880)}})))
    recovered = recovery.recover_from_run(result.out_dir, tmp_path / "retry", client=fake,
        options=RecoveryOptions(render=False, use_memory=False, critic_review=False, max_attempts=1))
    assert recovered.key["adaptive_contract"] == contract
    assert recovered.key["adaptive_bindings"] == bindings
    assert recovered.ready, recovered.to_text()
    destination = recovery.deliver_recovery(recovered)
    assert Document(destination / source.name).sections[0].left_margin.twips == 2880
    assert source.read_bytes() == original
    if chain:
        saved = json.loads((result.out_dir / "chain.json").read_text(encoding="utf-8"))
        assert saved["atomic_final_delivery"] is in_place
        if in_place:
            assert "only after final outcome verification" in (result.out_dir / "CHAIN_REPORT.md").read_text(encoding="utf-8")


@pytest.mark.parametrize("damage", ["missing_map", "missing_file", "missing_effects", "missing_authority", "stale_source", "wrong_digest"])
def test_retry_cannot_fall_back_to_recipe_authority_when_binding_is_invalid(tmp_path, damage):
    source, result, _, _ = failed_run(tmp_path)
    path = result.out_dir / "set_aside.json"
    value = json.loads(path.read_text(encoding="utf-8"))
    if damage == "missing_map":
        del value["goal_bindings"]
    elif damage == "missing_file":
        value["goal_bindings"] = {}
    elif damage == "wrong_digest":
        value["goal_bindings"][str(source)]["source_sha256"] = "0" * 64
    elif damage == "missing_effects":
        value["goal_bindings"][str(source)]["effects"] = {}
    elif damage == "missing_authority":
        del value["goal_contract"]
        del value["goal_bindings"]
    else:
        doc = Document(source)
        doc.add_paragraph("Later user save")
        doc.save(source)
    path.write_text(json.dumps(value), encoding="utf-8")
    if damage == "missing_effects":
        # An exported manifest without its companion record must still refuse
        # validator defaults that would reconstruct the lost approval.
        (result.out_dir / "batch.json").unlink()
    with pytest.raises(ValueError):
        recovery.recover_from_run(result.out_dir, tmp_path / "retry", client=FakeClient())
    assert not (tmp_path / "retry").exists()


def test_retry_preserves_the_specific_choice_among_duplicate_titles(tmp_path, monkeypatch):
    from spec_surgeon import ordinary_outcomes
    source = tmp_path / "source.docx"
    doc = Document()
    doc.add_paragraph("TITLE")
    doc.add_paragraph("TITLE")
    doc.save(source)
    contract = {"kind": "document_effects", "effects": [{"id": "title", "type": "format",
        "target": {"locators": [{"exact_text": "TITLE"}], "cardinality": 1}, "run": {"b": True}}]}
    review = ordinary_outcomes.build_review("Bold the second TITLE", [source], contract)
    candidates = review["documents"][0]["resolution"]["effects"]["title"]["candidates"]
    review = ordinary_outcomes.choose_targets(review, source, [candidates[1]["candidate_id"]], effect_id="title")
    _, bindings = ordinary_outcomes.validate(review, [source])
    run = tmp_path / "run"
    run.mkdir()
    (run / "set_aside.json").write_text(json.dumps({"request": review["request"], "spec": margin(1440),
        "goal_contract": contract, "goal_bindings": bindings, "files": [{"source": str(source)}]}))
    captured = {}

    def record(*args, **kwargs):
        captured.update(copy.deepcopy(kwargs))
        return captured

    monkeypatch.setattr(recovery, "recover", record)
    recovery.recover_from_run(run, tmp_path / "retry")
    assert captured["adaptive_binding"] == bindings
    assert captured["adaptive_contract"] == contract


def test_old_partial_in_place_chain_cannot_claim_safe_retry_without_retained_outcome(tmp_path):
    source, result, _, _ = failed_run(tmp_path, chain=True)
    metadata = result.out_dir / "chain.json"
    data = json.loads(metadata.read_text(encoding="utf-8"))
    data.update(in_place=True, atomic_final_delivery=False)
    metadata.write_text(json.dumps(data))
    with pytest.raises(ValueError, match="partial stages"):
        recovery.recover_from_run(result.out_dir, tmp_path / "retry", client=FakeClient())


def test_recovery_export_rechecks_originals_after_final_outcome_proof(tmp_path, monkeypatch):
    source, result, contract, bindings = failed_run(tmp_path)
    recovered = recovery.prepare_recovery([source], margin(2880)["request"], tmp_path / "prepared",
        spec=margin(2880), adaptive_contract=contract, adaptive_binding=bindings,
        options=RecoveryOptions(render=False, use_memory=False, critic_review=False))
    assert recovered.ready
    original = recovery._assert_adaptive

    def proof_then_user_save(*args, **kwargs):
        original(*args, **kwargs)
        doc = Document(source)
        doc.add_paragraph("A new user save during verification.")
        doc.save(source)

    monkeypatch.setattr(recovery, "_assert_adaptive", proof_then_user_save)
    destination = tmp_path / "delivery"
    with pytest.raises(ValueError, match="changed during final verification"):
        recovery.deliver_recovery(recovered, destination)
    assert not destination.exists()
    assert Document(source).paragraphs[-1].text == "A new user save during verification."
