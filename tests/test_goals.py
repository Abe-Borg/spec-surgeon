"""Explicit adaptive title goals: heterogeneous structure, ambiguity and source pins."""
from __future__ import annotations

import copy

import pytest
from docx import Document
from lxml import etree

from spec_surgeon.editspec import apply_spec, plan_spec
from spec_surgeon.goals import (GoalError, NeedsInput, choose_goal, compile_goal, resolve_goal,
                                validate_contract, verify_goal)
from spec_surgeon.package import DocxPackage, W
from spec_surgeon.verify import run_suite


TITLE = "FIRE SUPPRESSION"


def goal(**changes):
    return {"kind": "format_section_title", "target": {"title": TITLE}, "run": {"b": True}, **changes}


def write_document(path, layout="manual", duplicate=False):
    doc = Document()
    if layout == "table":
        cells = doc.add_table(rows=1, cols=2).rows[0].cells
        cells[0].text = "SECTION 21 13 13"
        paragraph = cells[1].paragraphs[0]
    else:
        doc.add_paragraph("SECTION 21 13 13")
        paragraph = doc.add_paragraph(style="Heading 1" if layout == "styled" else "Normal")
    # Mixed run formatting is preserved; the title is evidence across both runs.
    paragraph.add_run("FIRE ").bold = False
    paragraph.add_run("SUPPRESSION").italic = True
    doc.add_paragraph("Use the specified assemblies. Keep every word.")
    if duplicate:
        doc.add_paragraph("SECTION 22 11 16")
        doc.add_paragraph(TITLE)
    doc.save(path)
    return path


def execute(source, contract, binding=None):
    if binding is None:
        binding = resolve_goal(source, contract)["binding"]
    spec = compile_goal(source, contract, binding)
    pkg = DocxPackage(source.read_bytes())
    plan, fragments = plan_spec(pkg, spec)
    output = apply_spec(pkg, fragments).output_bytes()
    record = run_suite(source.read_bytes(), output, plan.to_dict(), str(source), "candidate.docx")
    return output, record, verify_goal(source, output, contract, binding)


@pytest.mark.parametrize("layout", ["styled", "manual", "table"])
@pytest.mark.parametrize("tracked", [False, True])
def test_same_goal_adapts_to_each_representation(tmp_path, layout, tracked):
    source = write_document(tmp_path / "spec.docx", layout)
    contract = goal(target={"title": TITLE, "section_number": "21 13 13"},
                    paragraph={"keepNext": True}, tracked=tracked)
    resolved = resolve_goal(source, contract)
    assert resolved["status"] == "resolved"
    assert resolved["binding"]["target"]["in_table"] == (layout == "table")
    assert resolved["binding"]["target"]["section_anchor_text"] == "SECTION 21 13 13"
    output, record, final = execute(source, contract)
    assert record.passed, record.to_markdown()
    assert final["passed"], final
    assert output != source.read_bytes()


def test_duplicate_title_requires_deliberate_source_bound_choice(tmp_path):
    source = write_document(tmp_path / "spec.docx", "styled", duplicate=True)
    contract = goal()
    resolved = resolve_goal(source, contract)
    assert resolved["status"] == "needs_input" and resolved["binding"] is None
    assert len(resolved["candidates"]) == 2  # a heading style never silently beats an exact manual title
    with pytest.raises(NeedsInput) as error:
        compile_goal(source, contract)
    assert error.value.alternatives == resolved["candidates"]
    selected = resolved["candidates"][1]
    bound = choose_goal(source, contract, selected["candidate_id"])
    assert bound["choice"] == "user" and bound["target"] == selected
    _, record, final = execute(source, contract, bound)
    assert record.passed, record.to_markdown()
    assert final["passed"], final


def test_explicit_style_restriction_is_never_reinterpreted_as_role(tmp_path):
    source = write_document(tmp_path / "spec.docx", "manual")
    contract = goal(target={"title": TITLE, "style_id": "Heading1"})
    resolved = resolve_goal(source, contract)
    assert resolved["status"] == "needs_input" and resolved["candidates"] == []
    with pytest.raises(NeedsInput):
        compile_goal(source, contract)
    styled = write_document(tmp_path / "styled.docx", "styled", duplicate=True)
    assert len(resolve_goal(styled, contract)["candidates"]) == 1


def test_section_anchor_disambiguates_without_normalizing_number(tmp_path):
    source = write_document(tmp_path / "spec.docx", duplicate=True)
    restricted = goal(target={"title": TITLE, "section_number": "22 11 16"})
    bound = resolve_goal(source, restricted)["binding"]
    assert bound["target"]["paragraph_index"] == 4
    compact_number = goal(target={"title": TITLE, "section_number": "221116"})
    assert resolve_goal(source, compact_number)["candidates"] == []


@pytest.mark.parametrize("text", ["Fire Suppression", " FIRE SUPPRESSION", "FIRE  SUPPRESSION", "FIRE\u00a0SUPPRESSION"])
def test_exact_approved_title_does_not_gain_implicit_equivalences(tmp_path, text):
    source = write_document(tmp_path / "spec.docx")
    assert resolve_goal(source, goal(target={"title": text}))["status"] == "needs_input"


def test_changed_document_and_changed_goal_invalidate_binding(tmp_path):
    source = write_document(tmp_path / "spec.docx")
    contract = goal()
    binding = resolve_goal(source, contract)["binding"]
    changed_contract = goal(run={"b": False})
    with pytest.raises(GoalError, match="source or approved goal changed"):
        compile_goal(source, changed_contract, binding)
    doc = Document(source)
    doc.add_paragraph("A later user edit.")
    doc.save(source)
    with pytest.raises(GoalError, match="source or approved goal changed"):
        compile_goal(source, contract, binding)
    with pytest.raises(GoalError, match="not in the current source"):
        choose_goal(source, contract, binding["candidate_id"])


def test_binding_cannot_retarget_a_paragraph_or_claim_false_uniqueness(tmp_path):
    source = write_document(tmp_path / "spec.docx", duplicate=True)
    contract = goal()
    candidates = resolve_goal(source, contract)["candidates"]
    bound = choose_goal(source, contract, candidates[0]["candidate_id"])
    tampered = copy.deepcopy(bound)
    tampered["target"]["paragraph_index"] = 2
    with pytest.raises(GoalError, match="evidence"):
        compile_goal(source, contract, tampered)
    tampered = {**bound, "choice": "unique"}
    with pytest.raises(GoalError, match="ambiguous"):
        compile_goal(source, contract, tampered)
    tampered = copy.deepcopy(bound)
    tampered["target"]["in_table"] = 0  # False == 0 in Python; binding evidence still has an exact JSON type.
    with pytest.raises(GoalError, match="evidence"):
        compile_goal(source, contract, tampered)


def test_deleted_title_is_excluded_instead_of_selected(tmp_path):
    source = write_document(tmp_path / "spec.docx")
    pkg = DocxPackage(source.read_bytes())
    root = pkg.xml(pkg.main_part())
    paragraph = list(root.iter(W + "p"))[1]
    deleted = etree.SubElement(paragraph, W + "del")
    deleted.set(W + "id", "3")
    deleted.set(W + "author", "Original author")
    run = etree.SubElement(deleted, W + "r")
    etree.SubElement(run, W + "delText").text = " OLD"
    pkg.write_xml(pkg.main_part(), root)
    source.write_bytes(pkg.output_bytes())
    resolved = resolve_goal(source, goal())
    assert resolved["candidates"] == []
    assert "pending deleted content" in resolved["excluded"][0]["reason"]


def test_no_change_already_satisfies_goal_but_still_independently_verified(tmp_path):
    source = write_document(tmp_path / "spec.docx")
    contract = goal()
    output, record, result = execute(source, contract)
    assert record.passed, record.to_markdown()
    assert result["passed"], result
    source.write_bytes(output)
    second, record, result = execute(source, contract)
    assert record.passed, record.to_markdown()
    assert result["passed"], result
    assert second == output


def test_empty_paragraph_properties_do_not_hide_unapproved_changes(tmp_path):
    source = write_document(tmp_path / "spec.docx")
    contract = goal()
    spec = compile_goal(source, contract)
    pkg = DocxPackage(source.read_bytes())
    plan, fragments = plan_spec(pkg, spec)
    output = apply_spec(pkg, fragments)
    root = output.xml(output.main_part())
    paragraph = list(root.iter(W + "p"))[1]
    spacing = etree.SubElement(paragraph.find(W + "pPr"), W + "spacing")
    spacing.set(W + "before", "480")
    output.write_xml(output.main_part(), root)
    record = run_suite(source.read_bytes(), output.output_bytes(), plan.to_dict(), str(source), "bad.docx")
    assert record.get("V10").status == "FAIL"
    assert "unpredicted change" in record.get("V10").summary


@pytest.mark.parametrize("change", [
    {"kind": "guess_title"}, {"target": {"role": "title"}}, {"target": {"title": ""}},
    {"run": {"b": 1}}, {"run": {}}, {"run": {"b": True}, "force": True},
    {"run": {"caps": True, "smallCaps": True}}, {"tracked": "yes"},
])
def test_goal_contract_is_explicit_and_closed(change):
    with pytest.raises(GoalError):
        validate_contract(goal(**change))
