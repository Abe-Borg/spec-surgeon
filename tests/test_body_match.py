"""Safe mode's "match body formatting to template" (``template_roles``, ``body_match``): the
template's roles read without a model and checked; each document's own plan (its outline converted to the
template's styles and numbering with every number keeping its values, the section line, title and END OF SECTION
restyled, the page and font matched, paragraph spacing, indents and alignment cleared — bold words and tables left);
documents it cannot match left out with the reason; the whole run with no AI call, every copy through the independent
checks; and the refusals that keep a look from changing unasked."""
from __future__ import annotations

import hashlib
from pathlib import Path

import pytest
from lxml import etree

from make_fixtures import build_body_template, build_numbering, build_styled_outline, build_typed, patch_docx
from spec_surgeon import body_match as bm
from spec_surgeon import llm, preview, recipes, template_roles as tr, vision
from spec_surgeon.chain import ChainFile
from spec_surgeon.document_task import DocumentTask
from spec_surgeon.gui import Session
from spec_surgeon.ooxml import paragraphs
from spec_surgeon.ops import lists as tl
from spec_surgeon.ops.effective import Resolver, theme_fonts
from spec_surgeon.ops.selectors import StyleSheet
from spec_surgeon.package import DocxPackage, W

W_DECL = f'xmlns:w="{W[1:-1]}"'
MAIN = "word/document.xml"


@pytest.fixture()
def folder(tmp_path) -> Path:
    (tmp_path / "tpl").mkdir()
    (tmp_path / "tpl" / "OFFICE MASTER.docx").write_bytes(build_body_template())
    (tmp_path / "21 13 13.docx").write_bytes(build_styled_outline(table=True))
    (tmp_path / "21 13 16.docx").write_bytes(build_styled_outline("21 13 16", "DRY-PIPE SPRINKLER SYSTEMS", table=True))
    (tmp_path / "typed.docx").write_bytes(build_typed(masterspec=True))
    return tmp_path


def template(folder: Path) -> Path:
    return folder / "tpl" / "OFFICE MASTER.docx"


def roles_of(folder: Path, **chosen) -> tr.TemplateRoles:
    return tr.resolve(template(folder), {**tr.detect(template(folder))["roles"], **chosen})


def no_ai(monkeypatch):
    def never(*a, **k):
        raise AssertionError("a Safe-mode button makes no AI call")
    monkeypatch.setattr(vision, "review_samples", never)
    monkeypatch.setattr(llm, "client", never)
    monkeypatch.setattr(preview, "render_samples", lambda *a, **k: None)


def outline(data: bytes):
    """paragraph text → (style, number) of the main part, by the transform's own list model."""
    pkg = DocxPackage(data)
    root = pkg.xml(MAIN)
    pars = paragraphs(root)
    nums = {x.index: x.text for x in tl.count(root, StyleSheet(pkg), tl.ListDefs.load(pkg))}
    out = {}
    for i, p in enumerate(pars):
        ps = p.find(f"{W}pPr/{W}pStyle")
        text = "".join(t.text or "" for t in p.iter(W + "t"))
        out.setdefault(text, (ps.get(W + "val") if ps is not None else None, nums.get(i), p))
    return out


def own_ppr(p) -> list:
    ppr = p.find(W + "pPr")
    return [etree.QName(c).localname for c in ppr] if ppr is not None else []


# ---- the template's roles --------------------------------------------------------------------------------------------

def test_the_templates_roles_page_and_font_are_read_without_a_model(folder):
    found = tr.detect(template(folder))
    assert found["problems"] == []
    assert found["roles"] == {"section_id": "SCT", "section_title": "TTL", "part": "PRT", "article": "ART",
                              "paragraph": "PR1", "subparagraph": "PR2", "subparagraph_2": "PR3",
                              "end_of_section": "EOS"}
    assert found["page"] == "8.5 × 11 in portrait, margins 1 in" and found["font"] == "Arial 10 pt"
    assert {"id": "PRT", "name": "PRT"} in found["styles"] and found["order"][0] == "section_id"
    roles = roles_of(folder)
    assert roles.levels == ["PRT", "ART", "PR1", "PR2", "PR3"] and roles.level_of("paragraph") == 2
    assert roles.page["margins"]["header"] == 720 and roles.page["size"] == {"w": 12240, "h": 15840}


@pytest.mark.parametrize("chosen, words", [
    ({"part": ""}, "no style stands for PART"),
    ({"subparagraph": ""}, "skips Subparagraph (1.)"),
    ({"article": "PR1"}, "have the same style"),
    ({"section_title": "Normal"}, "default paragraph style"),
    ({"paragraph": "NoSuchStyle"}, "has no paragraph style “NoSuchStyle”"),
    ({"article": "TTL", "section_title": ""}, "do not form one numbering"),
    ({"nonsense": "PRT"}, "There is no role"),
])
def test_a_choice_that_cannot_be_used_is_refused_in_one_sentence(folder, chosen, words):
    with pytest.raises(ValueError, match=words.replace("(", r"\(").replace(")", r"\)")):
        roles_of(folder, **chosen)


def test_a_template_without_an_outline_says_why(tmp_path):
    from docx import Document
    doc = Document()
    doc.add_paragraph("Just some text, no outline.")
    path = tmp_path / "plain.docx"
    doc.save(path)
    found = tr.detect(path)
    assert found["problems"] and "no style stands for PART" in found["problems"][0]
    assert tr.detect(None)["problems"] == [] and tr.detect(None)["roles"] == {}


# ---- each document's own plan ----------------------------------------------------------------------------------------

def test_a_styled_outline_plans_page_styles_numbering_and_clearing(folder):
    plan = bm.plan_file(folder / "21 13 13.docx", roles_of(folder))
    assert plan.kind == "automatic" and plan.body_size == 11.0 and plan.notes == []
    ops = [o["op"] for o in plan.ops]
    assert ops == ["set_page_setup", "apply_style", "apply_style", "apply_style", "convert_numbering",
                   "clear_direct_formatting"]
    styles = [o["params"]["style"] for o in plan.ops if o["op"] == "apply_style"]
    assert styles == ["SCT", "TTL", "EOS"]
    assert all(o["params"]["existing"] == "replace" and o["params"]["tracked"] is False
               for o in plan.ops if o["op"] == "apply_style")
    conv = plan.ops[4]
    assert conv["select"] == {"include_tables": False}
    assert conv["params"]["counters"] == "keep" and "from" not in conv["params"]       # the main list
    assert conv["params"]["levels"] == ["PRT", "ART", "PR1", "PR2", "PR3"] and "level_map" not in conv["params"]
    clear = plan.ops[5]
    assert clear["params"]["run"] is False and clear["select"] == {"paragraphs": {"in_table": False}}
    assert not set(bm.CLEARED) & set(clear["params"]["keep"]) and "tabs" in clear["params"]["keep"]
    assert plan.ops[0]["params"]["margins"]["left"] == 1440


def test_a_typed_outline_plans_typed_numbers_by_value(folder):
    plan = bm.plan_file(folder / "typed.docx", roles_of(folder))
    conv = next(o for o in plan.ops if o["op"] == "convert_typed_numbers")
    assert plan.kind == "typed" and conv["params"]["match"] == "counters" and conv["params"]["tracked"] is False


def test_the_form_boxes_leave_out_the_page_and_the_clearing(folder):
    plan = bm.plan_file(folder / "21 13 13.docx", roles_of(folder), page=False, clear=False)
    assert [o["op"] for o in plan.ops] == ["apply_style"] * 3 + ["convert_numbering"]


def test_a_document_with_a_landscape_page_keeps_its_pages(folder):
    def landscape(root):
        sz = root.find(f"{W}body/{W}sectPr/{W}pgSz")
        w, h = sz.get(W + "w"), sz.get(W + "h")
        sz.set(W + "w", h)
        sz.set(W + "h", w)
    path = folder / "wide.docx"
    path.write_bytes(patch_docx(build_styled_outline(), {MAIN: landscape}))
    plan = bm.plan_file(path, roles_of(folder))
    assert "set_page_setup" not in [o["op"] for o in plan.ops] and "another orientation" in plan.notes[0]


def test_an_article_level_the_template_lacks_refuses(folder):
    """With no style chosen for the template's two deepest levels, the outline's subparagraph has nowhere to go."""
    roles = roles_of(folder, subparagraph="", subparagraph_2="")
    with pytest.raises(bm.Refused, match=r"goes down to Subparagraph \(1\.\)"):
        bm.plan_file(folder / "21 13 13.docx", roles)


def test_an_outline_mixing_typed_and_automatic_numbers_refuses(folder):
    def typed_tail(root):
        last = root.find(W + "body").findall(W + "p")[-1]
        for num, text in (("2.02", "FITTINGS"), ("A.", "Steel fittings.")):
            last.addprevious(etree.fromstring(
                f'<w:p {W_DECL}><w:r><w:t xml:space="preserve">{num}</w:t></w:r><w:r><w:tab/></w:r>'
                f'<w:r><w:t>{text}</w:t></w:r></w:p>'))
    path = folder / "mixed.docx"
    path.write_bytes(patch_docx(build_styled_outline(), {MAIN: typed_tail}))
    with pytest.raises(bm.Refused, match="mixes numbers Word writes with numbers typed"):
        bm.plan_file(path, roles_of(folder))


def test_a_style_the_template_would_replace_under_other_paragraphs_refuses(folder):
    """The document has a "PR1" of its own on a note the button does not restyle: the template's PR1 would change the
    note's look unasked — refused, naming it."""
    def own_pr1(root):
        root.append(etree.fromstring(f'<w:style {W_DECL} w:type="paragraph" w:customStyle="1" w:styleId="PR1">'
                                     f'<w:name w:val="PR1"/><w:basedOn w:val="Normal"/></w:style>'))

    def note(root):
        last = root.find(W + "body").findall(W + "p")[-1]
        last.addprevious(etree.fromstring(f'<w:p {W_DECL}><w:pPr><w:pStyle w:val="PR1"/></w:pPr>'
                                          f'<w:r><w:t>A note in PR1.</w:t></w:r></w:p>'))
    path = folder / "guard.docx"
    path.write_bytes(patch_docx(build_styled_outline(), {"word/styles.xml": own_pr1, MAIN: note}))
    with pytest.raises(bm.Refused, match="“PR1” would take the template's definition.*‘A note in PR1.’"):
        bm.plan_file(path, roles_of(folder))


def with_header(data: bytes, style: str = "") -> bytes:
    """The document with a header line ("ISSUED FOR PERMIT"), in ``style`` when given."""
    import io
    from docx import Document
    doc = Document(io.BytesIO(data))
    head = doc.sections[0].header
    head.is_linked_to_previous = False
    head.paragraphs[0].add_run("ISSUED FOR PERMIT")
    if style:
        head.paragraphs[0]._p.get_or_add_pPr().insert(0, etree.fromstring(f'<w:pStyle {W_DECL} w:val="{style}"/>'))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_a_style_the_template_would_replace_in_a_header_refuses_too(folder):
    """Codex review on Abe-Borg/spec-surgeon#100: the guard read the body only. A header line in the document's own
    "PR1" would take the template's PR1 unasked — refused, naming it."""
    def own_pr1(root):
        root.append(etree.fromstring(f'<w:style {W_DECL} w:type="paragraph" w:customStyle="1" w:styleId="PR1">'
                                     f'<w:name w:val="PR1"/><w:basedOn w:val="Normal"/></w:style>'))
    path = folder / "header.docx"
    path.write_bytes(with_header(patch_docx(build_styled_outline(), {"word/styles.xml": own_pr1}), "PR1"))
    with pytest.raises(bm.Refused, match="“PR1” would take the template's definition.*‘ISSUED FOR PERMIT’"):
        bm.plan_file(path, roles_of(folder))


def test_a_style_the_template_would_replace_in_a_table_cell_the_conversion_skips_refuses(folder):
    """The conversion leaves table cells out; a cell on the main list in the document's own "ART" (the template's
    article style id) would take the template's ART unasked — refused (it no longer counts as restyled)."""
    def own_art(root):
        root.append(etree.fromstring(f'<w:style {W_DECL} w:type="paragraph" w:customStyle="1" w:styleId="ART">'
                                     f'<w:name w:val="ART"/><w:basedOn w:val="SpArt"/></w:style>'))

    def cell_in_art(root):
        cell = root.find(f".//{W}tc/{W}p")
        cell.insert(0, etree.fromstring(f'<w:pPr {W_DECL}><w:pStyle w:val="ART"/></w:pPr>'))
        old = cell.find(W + "pPr")
        if old is not None and old is not cell[0]:
            cell.remove(old)
    path = folder / "cell.docx"
    path.write_bytes(patch_docx(build_styled_outline(table=True), {"word/styles.xml": own_art, MAIN: cell_in_art}))
    with pytest.raises(bm.Refused, match="“ART” would take the template's definition.*‘Product data’"):
        bm.plan_file(path, roles_of(folder))


def test_the_font_step_keeps_to_the_body_outside_tables(folder):
    """Codex review on Abe-Borg/spec-surgeon#100: the font step reached headers, footers, notes, every style and the
    document defaults, while the button promises that tables and headers and footers stay as they are. Now: the body's
    runs outside tables, the styles only they use, and the defaults only when nothing else could inherit them."""
    path = folder / "headed.docx"
    path.write_bytes(with_header(build_styled_outline(table=True)))
    plan = bm.plan_file(path, roles_of(folder), font=True)
    assert plan.font_styles == ["SpArt", "SpPara", "SpPart", "SpSub"]      # Normal: the table cell and header use it
    assert "Header" not in plan.font_styles and plan.defaults_free is False
    bare = bm.plan_file(folder / "21 13 13.docx", roles_of(folder), font=True)
    assert bare.defaults_free is False                                     # its table cell
    ops = bm.font_ops("Arial", 10, [11.0], plan.font_styles, plan.defaults_free)
    assert [o["op"] for o in ops] == ["set_font", "set_style_font"]
    assert ops[0]["select"]["scope"] == ["document"] and ops[0]["select"]["paragraphs"] == {"in_table": False}
    assert ops[1]["select"] == {"styles": {"ids": ["SpArt", "SpPara", "SpPart", "SpSub"]}}
    assert [o["op"] for o in bm.font_ops("Arial", 10, [11.0], [], True)] == ["set_font", "set_doc_defaults_font"]
    # the template's look, in a step of its own: each imported style pinned to its font and size in the template
    assert {o["params"]["font"] for o in plan.look} == {"Arial"} and all(o["params"]["only_where_defined"] is False
                                                                       for o in plan.look)
    pinned = sorted(i for o in plan.look for i in o["select"]["styles"]["ids"])
    assert pinned == sorted(["SCT", "TTL", "EOS", "PRT", "ART", "PR1", "PR2", "PR3"])
    assert bm.plan_file(path, roles_of(folder)).look == []                 # the font box off: none


def test_the_button_leaves_the_headers_and_tables_as_they_were(folder, monkeypatch):
    """End to end on a document with a header: the copy's header part is byte-identical, and the table cell keeps its
    font, while the outline takes the template's."""
    no_ai(monkeypatch)
    path = folder / "headed.docx"
    path.write_bytes(with_header(build_styled_outline(table=True)))
    task = DocumentTask(Session(out_root=folder / "out"), render=False, mode="safe")
    task.attach([path])
    task.attach([template(folder)], template=True)
    view = task.perform("recipe", recipe="body_formatting", inputs={})
    assert view["stage"] == "preview", [m["text"] for m in task.messages]
    view = task.perform("primary", token=view["token"])
    assert view["stage"] == "done" and all(f.status == "PASS" for f in task.session.last_batch.files)
    src, out = DocxPackage.from_path(path), DocxPackage(Path(task.results[str(path.resolve())]).read_bytes())
    heads = [part for part in src.text_parts() if "header" in part]
    assert heads and all(src.read(part) == out.read(part) for part in heads)
    assert src.read("word/styles.xml") != out.read("word/styles.xml")

    def header_look(pkg):
        p = paragraphs(pkg.xml(heads[0]))[0]
        r = p.find(f".//{W}r")
        got = Resolver(StyleSheet(pkg), theme_fonts(pkg)).run(r.find(W + "rPr"), p)
        return got["ascii"], got["sz"]
    assert header_look(out) == header_look(src) == ("Times New Roman", "22")


def test_a_document_without_an_outline_refuses(tmp_path, folder):
    from docx import Document
    doc = Document()
    doc.add_paragraph("A letter, not a specification.")
    path = tmp_path / "letter.docx"
    doc.save(path)
    with pytest.raises(bm.Refused, match="no outline"):
        bm.plan_file(path, roles_of(folder))


# ---- the button, end to end, with no AI call ---------------------------------------------------------------------------

@pytest.mark.render                 # the button that moves the most of a page: its copies pass V6/V7
def test_the_button_matches_a_mixed_batch_and_leaves_out_what_it_cannot_match(folder, monkeypatch):
    """Two styled outlines (one layout), a typed one (another) and the architect's messy file: three copies pass
    every independent check — the template's styles, "1.01" shown as "1.1", Letter with 1 in margins, Arial, the
    paragraphs' own spacing, indents and alignment gone, the bold word and the table cell as they were — and the
    messy file (a skipped list level) is left untouched with the reason, ready for Default mode. No AI call; the
    template is never written."""
    no_ai(monkeypatch)
    (folder / "messy.docx").write_bytes(build_numbering(shared_restart=False))
    tpl_md5 = hashlib.md5(template(folder).read_bytes()).hexdigest()
    task = DocumentTask(Session(out_root=folder / "out"), render=True, mode="safe")    # key "present" (conftest)
    docs = [folder / n for n in ("21 13 13.docx", "21 13 16.docx", "typed.docx", "messy.docx")]
    task.attach(docs)
    task.attach([template(folder)], template=True)
    assert task.template_info()["body"]["roles"]["part"] == "PRT"           # what the form starts from
    view = task.perform("recipe", recipe="body_formatting", inputs={})
    assert view["stage"] == "preview" and view["action"] == "Create edited copies", [m["text"] for m in task.messages]
    assert task.session.model_free and len(task.session.spec["layouts"]) == 2
    assert [(i["name"], i["reason"][:12]) for i in view["set_aside"]] == [("messy.docx", "Not matched:")]
    assert "‘Shop drawings.’" in view["set_aside"][0]["reason"] and "op #" not in view["set_aside"][0]["reason"]
    view = task.perform("primary", token=view["token"])
    assert view["stage"] == "done" and len(task.results) == 3
    assert [e["name"] for e in view["left_untouched"]] == ["messy.docx"]
    assert all(f.status == "PASS" for f in task.session.last_batch.files)
    assert hashlib.md5(template(folder).read_bytes()).hexdigest() == tpl_md5
    out = Path(task.results[str((folder / "21 13 13.docx").resolve())]).read_bytes()
    got = outline(out)
    assert got["SECTION 21 13 13"][0] == "SCT" and got["WET-PIPE SPRINKLER SYSTEMS"][0] == "TTL"
    assert got["SUMMARY"][:2] == ("ART", "1.1") and got["GENERAL"][:2] == ("PRT", "PART 1 -")
    assert got["END OF SECTION 21 13 13"][0] == "EOS"
    para = got[next(t for t in got if t.startswith("The installer"))][2]
    assert own_ppr(para) == ["pStyle"]                                  # its justification cleared
    assert any(r.find(f"{W}rPr/{W}b") is not None for r in para.iter(W + "r"))   # "shall" stays bold
    assert "ind" in own_ppr(got["Product data"][2])                    # the table cell as it was
    pkg = DocxPackage(out)
    mar = DocxPackage.sect_prs(pkg.xml(MAIN))[0].find(W + "pgMar")
    assert mar.get(W + "left") == "1440" and mar.get(W + "header") == "720"
    # the font: the outline takes the template's Arial 10 pt (its imported styles pinned to the template's look); the
    # table cell and the document defaults it inherits stay Times New Roman 11 pt — the button keeps out of tables
    look = Resolver(StyleSheet(pkg), theme_fonts(pkg))

    def seen(p):
        r = p.find(f".//{W}r")
        got_look = look.run(r.find(W + "rPr"), p)
        return got_look["ascii"], got_look["sz"]
    assert seen(got["SUMMARY"][2]) == ("Arial", "20") and seen(para) == ("Arial", "20")
    assert seen(got["SECTION 21 13 13"][2])[0] == "Arial"
    assert seen(got["Product data"][2]) == ("Times New Roman", "22")
    fonts = pkg.xml("word/styles.xml").find(f"{W}docDefaults/{W}rPrDefault/{W}rPr/{W}rFonts")
    assert fonts.get(W + "ascii") == "Times New Roman"
    typed = outline(Path(task.results[str((folder / "typed.docx").resolve())]).read_bytes())
    assert typed["SUMMARY"][:2] == ("ART", "1.1") and typed["GENERAL"][:2] == ("PRT", "PART 1 -")
    assert any("match the body formatting" in line.lower() or "style" in line.lower() for line in task.history)
    task.hand_off([str((folder / "messy.docx").resolve())])
    assert task.mode == "default" and "Shop drawings" in task.todo[str((folder / "messy.docx").resolve())]["reason"]


def test_the_layout_plans_each_pass_the_gate_as_ordered_steps(folder):
    ctx = recipes.Context([folder / "21 13 13.docx", folder / "typed.docx"], template(folder), [])
    plan, left, _ = bm.build_layouts(roles_of(folder), ctx, "Match", page=True, font=True, clear=True)
    assert left == {} and plan["kind"] == "layouts" and len(plan["layouts"]) == 2
    for layout in plan["layouts"]:
        stages = [[o["op"] for o in st["spec"]["ops"]] for st in layout["spec"]["stages"]]
        assert stages[0][:2] == ["set_page_setup", "set_font"]          # the page and the font come first
        assert set(stages[-1]) <= {"set_style_font", "set_font"} and "set_style_font" in stages[-1]  # the look, last
        for op in layout["spec"]["stages"][-1]["spec"]["ops"]:
            if op["op"] == "set_font":                                  # plain body text: by its style, no table
                assert op["select"]["paragraphs"]["in_table"] is False and op["select"]["paragraphs"]["style_in"]
        assert stages[-2][-1] == "clear_direct_formatting"              # the clearing before it
        first = layout["spec"]["stages"][0]["spec"]["ops"][1]
        assert first["select"] == {"scope": ["document"], "runs": "all", "paragraphs": {"in_table": False}}
        assert all(st["spec"]["request"] == "Match" for st in layout["spec"]["stages"])


def test_documents_of_one_kind_whose_plans_differ_are_left_out_but_the_majority(folder, monkeypatch):
    """Three files of one pattern: the largest group with one plan runs; the odd one out is left out with why."""
    for n in ("a", "b", "c"):
        (folder / f"{n}.docx").write_bytes(build_styled_outline())
    real = bm.plan_file

    def odd(path, roles, **k):
        plan = real(path, roles, **k)
        if Path(path).name == "c.docx":
            plan.ops = plan.ops[1:]
        return plan
    monkeypatch.setattr(bm, "plan_file", odd)
    ctx = recipes.Context([folder / "a.docx", folder / "b.docx", folder / "c.docx"], template(folder), [])
    plan, left, _ = bm.build_layouts(roles_of(folder), ctx, "Match", page=True, font=True, clear=True)
    assert list(left) == [str((folder / "c.docx").resolve())] and "built like a.docx" in next(iter(left.values()))
    assert len(plan["layouts"]) == 1


def test_a_batch_nothing_can_match_previews_nothing_and_says_why(tmp_path, folder, monkeypatch):
    no_ai(monkeypatch)
    from docx import Document
    doc = Document()
    doc.add_paragraph("A letter, not a specification.")
    doc.save(tmp_path / "letter.docx")
    task = DocumentTask(Session(out_root=tmp_path / "out", visual_check=False), render=False, mode="safe")
    task.attach([tmp_path / "letter.docx"])
    task.attach([template(folder)], template=True)
    view = task.perform("recipe", recipe="body_formatting", inputs={})
    assert view["stage"] == "empty" and task.session.spec is None
    assert [e["name"] for e in view["left_untouched"]] == ["letter.docx"]
    assert "no outline" in view["left_untouched"][0]["reason"] and "None of the documents" in task.messages[-1]["text"]


def test_a_wrong_role_choice_is_said_and_nothing_is_previewed(folder):
    task = DocumentTask(Session(out_root=folder / "out", visual_check=False), render=False, mode="safe")
    task.attach([folder / "21 13 13.docx"])
    task.attach([template(folder)], template=True)
    view = task.perform("recipe", recipe="body_formatting", inputs={"roles": {"article": "PR1"}})
    assert view["stage"] == "empty" and "have the same style" in task.messages[-1]["text"]
    assert "Default mode" in task.messages[-1]["text"]


def test_a_level_map_the_recipe_got_wrong_cannot_renumber_a_document(folder, monkeypatch):
    """A planted defect: every role one level deeper in the template than it is (articles as paragraphs). The numbers
    cannot keep their values (PART has no place left), and the conversion refuses: the document is left out, never
    delivered renumbered."""
    roles = roles_of(folder)
    real = roles.level_of
    monkeypatch.setattr(roles, "level_of", lambda role: None if real(role) is None else min(real(role) + 1, 4))
    ctx = recipes.Context([folder / "21 13 13.docx"], template(folder), [])
    plan, left, _ = bm.build_layouts(roles, ctx, "Match", page=False, font=False, clear=False)
    assert plan is None and next(iter(left.values())).startswith("Not matched: ‘GENERAL’")


# ---- a step's stop, read like a batch's ---------------------------------------------------------------------------------

def test_a_chain_files_stop_reads_like_a_batch_files():
    """A plan per layout mixes chain files and batch files in one result; its readers ask both for ``exception`` and
    ``record``."""
    cf = ChainFile(source=Path("a.docx"), out_name="a.docx")
    assert cf.exception is None and cf.record is None
    cf.status, cf.stopped_at, cf.code, cf.reason = "EXCEPTION", 2, "numbering-mismatch", "it would count differently"
    assert cf.exception == {"code": "numbering-mismatch", "kind": "refusal", "stage": "step 2",
                            "reason": "it would count differently", "details": {}}
    cf.status = "FAIL"
    assert cf.exception is None


def test_plain_body_text_on_a_shared_style_takes_the_body_font(folder, monkeypatch):
    """Codex review on Abe-Borg/spec-surgeon#103: a body paragraph the plan does not restyle, in Normal — which the table
    cell shares, so the font step leaves Normal alone — kept Times New Roman. Its runs now take the template's body
    font in the last step; the table cell keeps its font."""
    no_ai(monkeypatch)

    def plain_note(root):
        summary = next(p for p in root.iter(W + "p") if "".join(t.text or "" for t in p.iter(W + "t")) == "SUMMARY")
        summary.addnext(etree.fromstring(f'<w:p {W_DECL}><w:r><w:t>A plain note to the specifier.</w:t></w:r></w:p>'))
    path = folder / "plain.docx"
    path.write_bytes(patch_docx(build_styled_outline(table=True), {MAIN: plain_note}))
    plan = bm.plan_file(path, roles_of(folder), font=True)
    assert plan.plain_styles == ["Normal"] and "Normal" not in plan.font_styles
    task = DocumentTask(Session(out_root=folder / "out"), render=False, mode="safe")
    task.attach([path])
    task.attach([template(folder)], template=True)
    view = task.perform("recipe", recipe="body_formatting", inputs={})
    assert view["stage"] == "preview", [m["text"] for m in task.messages]
    view = task.perform("primary", token=view["token"])
    assert view["stage"] == "done" and all(f.status == "PASS" for f in task.session.last_batch.files)
    data = Path(task.results[str(path.resolve())]).read_bytes()
    out, got = DocxPackage(data), outline(data)
    look = Resolver(StyleSheet(out), theme_fonts(out))

    def seen(p):
        r = p.find(f".//{W}r")
        return look.run(r.find(W + "rPr"), p)["ascii"]
    assert seen(got["A plain note to the specifier."][2]) == "Arial"
    assert seen(got["Product data"][2]) == "Times New Roman"
    assert seen(got["SECTION 21 13 13"][2]) == "Arial" and got["SECTION 21 13 13"][0] == "SCT"
